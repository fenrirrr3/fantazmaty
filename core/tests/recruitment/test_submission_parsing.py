"""Rozbiór zgłoszeń tekstów z poczty i z wklejonych rekordów: format z etykietami (CF7), dawne CSV, zgody, granice wiadomości, załączniki, tożsamość wiadomości i paczki ZIP (przetwarzanie DOCX, potwierdzenie przebudowy, limity czasu)."""
from email.message import EmailMessage
from html import escape
from io import BytesIO
from itertools import chain, repeat
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.test import SimpleTestCase, TestCase, tag
from docx import Document

from core.models import MailboxConnection
from core.services.mailbox import MailboxError, polish_date, story_title
from core.services.document_converter import RebuildConfirmationRequired
from core.services.mailbox_import import _package_messages, package_messages, parse_message, prepare_forms, safe_name
from core.services.review_import_parser import encode_submission, parse_review_records
from core.tests.programs.support import sized_docx
from fantazmaty.testing import create_superuser
from texts.models import Anthology

from .support import (
    LEGACY_EXAMPLE, NAMED_BODY, NAMED_MESSAGE, SUBMISSION_FIELDS, csv_mail, identity_mail, named_mail,
)


class NamedSubmissionTests(SimpleTestCase):
    def test_multiple_pasted_named_submissions_preserve_messages(self):
        second = (NAMED_BODY.replace('Anna Maria', 'Jan').replace('Nowak-Kowalska', 'Autor')
                  .replace('ANNA@example.com', 'jan@example.com').replace('premierach, naborach', 'naborach'))
        records, errors = parse_review_records(NAMED_BODY + '\n\n' + second)
        self.assertEqual(errors, [])
        self.assertEqual(len(records), 2)
        self.assertEqual([row['author_first_name'] for row in records], ['Anna Maria', 'Jan'])
        self.assertEqual(records[0]['author_message'], NAMED_MESSAGE)
        self.assertEqual(records[1]['email'], 'jan@example.com')
        self.assertFalse(records[1]['newsletter_premieres'])
        self.assertTrue(records[1]['newsletter_recruitment'])

    def test_email_never_silently_imports_multiple_submissions(self):
        """Jedna wiadomość z kilkoma zgłoszeniami (z etykietami lub CSV) jest odrzucana, a nie importowana po cichu."""
        with self.assertRaises(MailboxError):
            parse_message(1, named_mail(NAMED_BODY + '\n\n' + NAMED_BODY))
        for line in ('Brak danych', 'A B;T;g;2;a@b.pl;;N\nC D;T;g;2;c@d.pl;;N'):
            with self.subTest(line=line), self.assertRaises(MailboxError):
                parse_message(1, csv_mail(line=line))

    def test_csv_error_identifies_person_instead_of_uid_and_row(self):
        body = 'Anna Nowak;"Niedomknięty tytuł;fantasy;1000;anna@example.com;;Test\n' + '\n'.join('Linia' for _ in range(86))
        with self.assertRaises(MailboxError) as error:
            parse_message(1402, named_mail(body, reply_to='Anna Nowak <anna@example.com>'))
        text = str(error.exception)
        self.assertTrue(text.startswith('Anna Nowak <anna@example.com>:'))
        self.assertIn('niepoprawny zapis cudzysłowów', text)
        self.assertNotIn('Wiadomość 1402', text)
        self.assertNotRegex(text, r'Wiersz\s+\d+')

    def test_attachment_errors_identify_person(self):
        message = EmailMessage()
        message['Reply-To'] = 'Anna Nowak <anna@example.com>'
        message.set_content(NAMED_BODY)
        message.add_attachment(b'', maintype='application', subtype='octet-stream', filename='pusty.docx')
        with self.assertRaises(MailboxError) as error:
            parse_message(1402, message.as_bytes())
        self.assertIn('Anna Nowak <anna@example.com>', str(error.exception))
        self.assertIn('pusty (0 bajtów)', str(error.exception))
        self.assertNotIn('1402', str(error.exception))

    def test_names_pseudonym_semicolons_and_multiline_message(self):
        parsed = parse_message(12, named_mail())
        rows, errors = parse_review_records(parsed['record'])
        self.assertEqual(errors, [])
        row = rows[0]
        self.assertEqual(row['author_first_name'], 'Anna Maria')
        self.assertEqual(row['author_last_name'], 'Nowak-Kowalska')
        self.assertEqual(row['author_pseudonym'], 'A. M. Nowak')
        self.assertEqual(row['title'], '[Opowieść; część druga]')
        self.assertEqual(row['content_warnings'], 'przemoc; śmierć')
        self.assertEqual(row['length'], 37930)
        self.assertEqual(row['email'], 'anna@example.com')
        self.assertTrue(row['newsletter_premieres'] and row['newsletter_recruitment'])
        self.assertEqual(row['author_message'], NAMED_MESSAGE)
        self.assertEqual(parsed['anthology'], 'Test')
        self.assertEqual(parsed['folder'], 'Opowieść')
        self.assertEqual(parsed['author_pseudonym'], 'A. M. Nowak')

    def test_manual_and_mail_share_named_parser(self):
        direct, errors = parse_review_records(NAMED_BODY)
        self.assertEqual(errors, [])
        mailed, errors = parse_review_records(parse_message(1, named_mail())['record'])
        self.assertEqual(errors, [])
        self.assertEqual(direct, mailed)

    def test_optional_fields_can_be_empty_or_omitted(self):
        body = '\n'.join(line for line in NAMED_BODY.splitlines()[:11]
                         if not line.startswith(('Pseudonim:', 'Ostrzeżenia o treści:', 'Numer telefonu:', 'Newsletter:')))
        row = parse_review_records(parse_message(1, named_mail(body))['record'])[0][0]
        for name in ('author_pseudonym', 'content_warnings', 'phone_number', 'author_message'):
            self.assertEqual(row[name], '')
        self.assertFalse(row['newsletter_premieres'] or row['newsletter_recruitment'])

    def test_case_whitespace_bom_and_crlf(self):
        body = '﻿\r\n' + NAMED_BODY.replace('Imię:', '  IMIĘ :').replace('\n', '\r\n')
        self.assertEqual(parse_review_records(body)[0][0]['author_first_name'], 'Anna Maria')

    def test_html_message(self):
        body = '<html><body>' + ''.join('<p>' + escape(line) + '</p>' for line in NAMED_BODY.splitlines()) + '</body></html>'
        row = parse_review_records(parse_message(1, named_mail(body, html=True))['record'])[0][0]
        self.assertEqual(row['author_first_name'], 'Anna Maria')
        self.assertEqual(row['title'], '[Opowieść; część druga]')
        self.assertIn('Dzień dobry; przesyłam tekst.', row['author_message'])

    def test_empty_html_is_readable_mailbox_error(self):
        message = EmailMessage()
        message['Subject'] = 'Test'
        message.set_content('')
        message.set_type('text/html')
        message.add_attachment(b'notempty', maintype='application', subtype='octet-stream', filename='tekst.docx')
        with self.assertRaisesMessage(MailboxError, 'treść HTML jest pusta'):
            parse_message(1, message.as_bytes())

    def test_missing_duplicate_and_unknown_metadata_fail(self):
        for body, expected in (
            (NAMED_BODY.replace('Nazwisko: Nowak-Kowalska', 'Nazwisko:'), 'nazwisko'),
            (NAMED_BODY.replace('Nazwisko: Nowak-Kowalska', ''), 'nazwisko'),
            (NAMED_BODY.replace('Nazwisko:', 'Imię: Jan\nNazwisko:'), 'więcej niż raz'),
            (NAMED_BODY.replace('Gatunek:', 'Gantunek:'), 'nieznana etykieta'),
            (NAMED_BODY.replace('ANNA@example.com', 'nie jest adresem'), 'adres e-mail'),
        ):
            with self.subTest(expected=expected), self.assertRaises(MailboxError) as error:
                parse_message(1, named_mail(body))
            self.assertIn(expected, str(error.exception).lower())

    def test_invalid_length_pseudonym_and_consent_are_rejected(self):
        for body in (
            NAMED_BODY.replace('37 930', '0'),
            NAMED_BODY.replace('37 930', 'tekst'),
            NAMED_BODY.replace('A. M. Nowak', 'x' * 101),
            NAMED_BODY.replace('premierach, naborach', 'wszystko'),
        ):
            with self.subTest(body=body):
                try:
                    parsed = parse_message(1, named_mail(body))
                except MailboxError:
                    continue
                rows, errors = parse_review_records(parsed['record'])
                self.assertTrue(errors)
                self.assertEqual(rows, [])

    def test_subject_supplies_anthology_and_is_required(self):
        with self.assertRaises(MailboxError) as error:
            parse_message(1, named_mail(subject='Autor – Opowieść'))
        self.assertIn('w temacie brakuje', str(error.exception))

    def test_optional_end_marker_excludes_form_footer(self):
        """Stopka formularza (np. IP) po znaczniku końca nie trafia do wiadomości autora – w obu formatach."""
        parsed = parse_message(1, named_mail(NAMED_BODY + '\n--- KONIEC WIADOMOŚCI AUTORA ---\nStopka techniczna'))
        self.assertNotIn('Stopka techniczna', parsed['author_message'])
        self.assertIn('"Dziękuję".', parsed['author_message'])

        def message(body):
            record = parse_message(1, named_mail(body, subject='Nabór: „Test”'))['record']
            return parse_review_records(record)[0][0]['author_message']

        line = 'Jan Autor;T;fantasy;;100;jan@example.com;;premierach;Dziękuję.'
        self.assertEqual(message(line + '\nDodatkowa treść'), 'Dziękuję.\nDodatkowa treść')
        with_marker = message(line + '\nDrugi wiersz.\n--- KONIEC WIADOMOŚCI AUTORA ---\nIP: 127.0.0.1')
        self.assertEqual(with_marker, 'Dziękuję.\nDrugi wiersz.')
        self.assertNotIn('127.0.0.1', with_marker)
        self.assertEqual(message(line), 'Dziękuję.')


class CsvSubmissionTests(SimpleTestCase):
    def test_old_mail_formats_still_parse(self):
        cases = (
            ('Jan Autor;Stary;fantasy;1000;jan@example.com;;Test', '', False),
            ('Jan Autor;Stary;fantasy;1000;jan@example.com;;Test;naborach', '', True),
            (encode_submission(['Jan Autor', 'Stary; drugi', 'fantasy', 'PRZEMOC', '1000', 'jan@example.com', '',
                                'naborach', 'Wiadomość; autora']), 'przemoc', True),
        )
        for body, warnings, recruitment in cases:
            with self.subTest(body=body):
                rows, errors = parse_review_records(parse_message(1, named_mail(body))['record'])
                self.assertEqual(errors, [])
                self.assertEqual(rows[0]['author_last_name'], 'Autor')
                self.assertEqual(rows[0]['content_warnings'], warnings)
                self.assertEqual(rows[0]['newsletter_recruitment'], recruitment)
                self.assertEqual(rows[0]['author_pseudonym'], '')

    def test_parser_new_and_legacy_and_unknown_consent(self):
        fields = list(SUBMISSION_FIELDS)
        rows, errors = parse_review_records(encode_submission(fields))
        self.assertEqual(errors, [])
        row = rows[0]
        self.assertEqual(row['length'], 1234)
        self.assertEqual(row['content_warnings'], 'przemoc')
        self.assertEqual(row['author_message'], fields[8])
        self.assertTrue(row['newsletter_premieres'])
        self.assertTrue(row['newsletter_recruitment'])
        legacy, errors = parse_review_records('Jan Autor;Stary;fantasy;1234;;jan@example.com;')
        self.assertFalse(legacy[0]['has_consent_fields'])
        self.assertFalse(errors)
        fields[7] = 'nieznane'
        self.assertTrue(parse_review_records(encode_submission(fields))[1])

    def test_invalid_legacy_consent_not_silently_dropped(self):
        with self.assertRaises(MailboxError):
            parse_message(1, named_mail('Jan Autor;T;fantasy;100;jan@example.com;;Test;premierach, literowka',
                                        subject='Nabór: „Test”'))

    def test_mail_new_fields_subject_and_end_marker(self):
        msg = EmailMessage()
        msg['Subject'] = 'Nabór: „Nabór testowy” – Opowieść'
        msg.set_content(';'.join(SUBMISSION_FIELDS) + '\n--- KONIEC WIADOMOŚCI AUTORA ---\nStopka techniczna')
        msg.add_attachment(b'not empty', maintype='application',
                           subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename='tekst.docx')
        parsed = parse_message(1, msg.as_bytes())
        self.assertEqual(parsed['anthology'], 'Nabór testowy')
        rows, errors = parse_review_records(parsed['record'])
        self.assertFalse(errors)
        self.assertEqual(rows[0]['author_message'], SUBMISSION_FIELDS[8])
        self.assertTrue(rows[0]['newsletter_recruitment'])

    def test_mail_csv_matches_manual_import_and_ignores_signature_images(self):
        fields = ['Jan Autor', 'Tytuł; część druga', 'fantasy', 'przemoc', '1000', 'jan@example.test', '123',
                  'premierach', 'Linia pierwsza\nLinia druga; też']
        mail = EmailMessage()
        mail['Subject'] = 'Nabór: „Test” – Tytuł'
        mail.set_content(encode_submission(fields))
        mail.add_attachment(b'docx', maintype='application',
                            subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename='tekst.docx')
        mail.add_attachment(b'logo', maintype='image', subtype='png', filename='logo.png', disposition='inline')
        result = parse_message(1, mail.as_bytes())
        self.assertEqual([name for name, data in result['files']], ['tekst.docx'])
        self.assertEqual(result['ignored_files'], ['logo.png'])
        direct = parse_review_records(encode_submission(fields))[0][0]
        imported = parse_review_records(result['record'])[0][0]
        self.assertEqual(imported['title'], direct['title'])
        self.assertEqual(imported['author_message'], direct['author_message'])

    def test_fields_ignore_trailing_and_body(self):
        parsed = parse_message(12, csv_mail())
        self.assertEqual(parsed['record'],
                         'Julia Moskalik;Potworna Przystan;dark fantasy;15204;;girl@example.com;885711216')
        self.assertEqual(parsed['folder'], 'Potworna Przystan')
        self.assertNotIn('IGNORUJ', str(parsed['record']))

    def test_exact_pasted_markdown_example(self):
        """Rekord wklejony z linkiem mailto i ucieczkami Markdown daje poprawne pola i zgody."""
        rows, errors = parse_review_records(LEGACY_EXAMPLE)
        self.assertFalse(errors)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['genre'], 'Grimdark fantasy')
        self.assertEqual(row['length'], 37930)
        self.assertEqual(row['email'], 'renata.osciak@gmail.com')
        self.assertTrue(row['newsletter_premieres'] and row['newsletter_recruitment'])
        self.assertEqual(row['phone_number'], '727908378')
        self.assertEqual(row['source_anthology'], 'Na pokład, psubraty')

    def test_legacy_mail_retains_consents(self):
        msg = EmailMessage()
        msg['Subject'] = 'Renata Ościak – Dwadzieścia rys na amulecie'
        msg.set_content(LEGACY_EXAMPLE)
        msg.add_attachment(b'data', maintype='application', subtype='octet-stream', filename='tekst.docx')
        parsed = parse_message(1, msg.as_bytes())
        rows, errors = parse_review_records(parsed['record'])
        self.assertFalse(errors)
        self.assertTrue(rows[0]['newsletter_premieres'])
        self.assertTrue(rows[0]['newsletter_recruitment'])
        self.assertEqual(parsed['anthology'], 'Na pokład, psubraty')

    def test_polish_date_and_title(self):
        self.assertEqual(polish_date('Tue, 12 May 2026 10:30:00 +0200'), '12 maja 2026, 10:30')
        self.assertEqual(polish_date('bad'), '–')
        self.assertEqual(story_title('Autor – Tytuł – część druga'), 'Tytuł – część druga')

    def test_submission_mail_also_keeps_inline_documents(self):
        msg = EmailMessage()
        msg.set_content('Treść')
        msg.add_attachment(b'doc', maintype='application', subtype='msword', filename='stary.doc', disposition='inline')
        self.assertEqual(parse_message(1, msg.as_bytes(), submission=False)['files'], [('stary.doc', b'doc')])


@tag('slow', 'docx')
class SubmissionPackagingTests(SimpleTestCase):
    def test_original_and_converted_zip(self):
        row = parse_message(12, csv_mail())
        row['folder'] = '../CON'
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            for clean, convert in ((False, False), (True, False), (False, True), (True, True)):
                with (
                    self.subTest(clean=clean, convert=convert),
                    _package_messages([row], clean, convert) as out,
                    ZipFile(out) as z,
                ):
                    self.assertEqual(len(z.namelist()), 3 if convert else 1)
                    self.assertTrue(all('..' not in p.split('/') for p in z.namelist()))
                    doc = Document(BytesIO(z.read(next(p for p in z.namelist() if p.endswith('.docx')))))
                    self.assertEqual(doc.paragraphs[0].text,
                                     'Koniec. Następne zdanie.' if clean else 'Koniec. następne zdanie.')

    def test_soft_line_breaks_become_paragraphs_except_lists_and_headings(self):
        """„Popraw pliki Odkurzaczem” (sam albo z przebudową i konwersją): miękki enter (Shift+Enter) w zwykłym
        akapicie i tabeli staje się nowym akapitem, w punkcie listy i nagłówku zostaje; bez Odkurzacza – bez zmian."""
        document = Document()
        document.add_heading('Rozdział', level=1).runs[0].add_break()
        document.paragraphs[-1].add_run('podtytuł')
        paragraph = document.add_paragraph('Pierwsza linia')
        paragraph.runs[0].add_break()
        paragraph.add_run('druga linia')
        item = document.add_paragraph('Punkt listy', style='List Bullet')
        item.runs[0].add_break()
        item.add_run('ciąg dalszy punktu')
        cell = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell.add_run('Komórka').add_break()
        cell.add_run('druga w komórce')
        stream = BytesIO()
        document.save(stream)
        row = {'uid': 12, 'folder': 'Opowieść', 'sender': 'Autor', 'files': [('tekst.docx', stream.getvalue())]}
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            for clean, convert, rebuild in ((True, False, False), (True, False, True), (True, True, False),
                                            (False, False, True)):
                with (
                    self.subTest(clean=clean, convert=convert, rebuild=rebuild),
                    # Przebudowa pomija tabele i punktory list (potwierdzone), tekst punktu zostaje.
                    _package_messages([row], clean, convert, rebuild, True) as out,
                    ZipFile(out) as archive,
                ):
                    result = Document(BytesIO(archive.read(next(p for p in archive.namelist() if p.endswith('.docx')))))
                    texts = [p.text for p in result.paragraphs if p.text.strip()]
                    cells = [p.text for p in result.tables[0].cell(0, 0).paragraphs] if not rebuild else None
                    if clean:
                        self.assertIn('Pierwsza linia', texts)
                        self.assertIn('druga linia', texts)
                        self.assertIn(cells, (None, ['Komórka', 'druga w komórce']))
                    else:
                        self.assertIn('Pierwsza linia\ndruga linia', texts)
                    self.assertIn('Rozdział\npodtytuł', texts)
                    self.assertIn('Punkt listy\nciąg dalszy punktu', texts)

    def test_partial_archive_keeps_equal_titles_in_separate_folders(self):
        messages = [parse_message(12, csv_mail()), parse_message(13, csv_mail())]
        accepted, errors = [], []
        with (
            package_messages(messages, False, False, False, skipped_errors=errors, accepted=accepted,
                             rebuild_warnings=[]) as output,
            ZipFile(output) as archive,
        ):
            self.assertEqual(len(archive.namelist()), 2)
            self.assertEqual(len({name.split('/')[0] for name in archive.namelist()}), 2)
        self.assertEqual([message['uid'] for message in accepted], [12, 13])
        self.assertFalse(errors)


class MessageIdentityTests(TestCase):
    def test_identity_survives_transport_headers_but_not_content_change(self):
        a = parse_message(1, identity_mail(), submission=False)
        b = parse_message(2, identity_mail(received='new server'), submission=False)
        self.assertEqual(a['fingerprint'], b['fingerprint'])
        self.assertNotEqual(a['fingerprint'],
                            parse_message(2, identity_mail(body='different'), submission=False)['fingerprint'])

    def test_identical_messages_with_two_uids_import_once(self):
        box = MailboxConnection.objects.create(name='teksty', host='imap.example.test', username='teksty@example.test')
        Anthology.objects.create(title='Nabór')
        raw = identity_mail(body='Jan Kowalski;Tytuł;fantasy;100;a@example.test;123456789;Nabór')
        rows = [parse_message(i, raw) for i in (1, 2)]
        forms = prepare_forms(rows, box, 1, create_superuser())
        self.assertEqual(len(forms), 1)
        self.assertTrue(rows[1]['downloaded'])
        self.assertEqual(forms[0].data['records'].count('\n'), 0)


class MessageStructureErrorTests(SimpleTestCase):
    """Wiadomości bez treści, bez plików, ze zbyt wieloma plikami albo bez naboru w temacie są odrzucane z opisem."""

    def test_attachment_count_and_missing_body(self):
        message = EmailMessage()
        message['From'] = 'Jan <jan@example.test>'
        message.set_content('Treść bez plików')
        with self.assertRaisesMessage(MailboxError, 'Jan <jan@example.test>: nie znaleziono załączników'):
            parse_message(1, message.as_bytes())
        for number in range(11):
            message.add_attachment(b'x', maintype='application', subtype='pdf', filename=f'{number}.pdf')
        with self.assertRaisesMessage(MailboxError, 'znaleziono 11 załączników; limit wynosi 10'):
            parse_message(1, message.as_bytes())
        only_file = EmailMessage()
        only_file['From'] = 'jan@example.test'
        only_file.add_attachment(b'x', maintype='application', subtype='pdf', filename='tekst.pdf')
        self.assertEqual(parse_message(1, only_file.as_bytes(), submission=False)['title'], '(bez tytułu)')
        with self.assertRaisesMessage(MailboxError, 'brak danych zgłoszenia w treści'):
            parse_message(1, only_file.as_bytes())

    def test_new_format_requires_anthology_in_subject(self):
        raw = named_mail(encode_submission(list(SUBMISSION_FIELDS)), subject='Jan Autor – Opowieść')
        with self.assertRaisesMessage(MailboxError, 'w temacie brakuje nazwy antologii'):
            parse_message(1, raw)

    def test_safe_names_avoid_reserved_and_empty_names(self):
        self.assertEqual(safe_name('CON'), 'opowiadanie_CON')
        self.assertEqual(safe_name('lpt1.docx'), 'opowiadanie_lpt1.docx')
        self.assertEqual(safe_name(' .. '), 'opowiadanie_')
        self.assertEqual(safe_name('a<b>:c?'), 'a_b__c_')


def converted(payload=b'docx', warnings=()):
    """Wynik convert_document: ZIP z DOCX, PDF i EPUB oraz opcjonalnymi uwagami konwersji."""
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        for ext in ('docx', 'pdf', 'epub'):
            archive.writestr('document.' + ext, payload + ext.encode())
    output.seek(0)
    output.conversion_warnings = list(warnings)
    return output, None, None


@tag('docx')
class PackagingErrorTests(SimpleTestCase):
    """Paczka ZIP: przetwarzanie tylko DOCX, potwierdzenie przebudowy, limity czasu i sprzątanie po błędzie."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.docx = parse_message(12, csv_mail())['files'][0][1]

    def message(self, uid=12, files=None, folder='Opowieść'):
        return {'uid': uid, 'folder': folder, 'sender': f'Autor {uid}', 'files': files or [('tekst.docx', self.docx)]}

    def names(self, output):
        with output, ZipFile(output) as archive:
            return sorted(archive.namelist())

    def test_processing_requires_valid_docx(self):
        with self.assertRaisesMessage(MailboxError, 'Odkurzacz i konwerter obsługują DOCX'):
            _package_messages([self.message(files=[('tekst.pdf', b'%PDF')])], True, False, False)
        with self.assertRaisesMessage(MailboxError, 'Nieprawidłowy DOCX: zly.docx. Wyłącz przetwarzanie'):
            _package_messages([self.message(files=[('zly.docx', b'not a zip')])], False, False, True)
        with self.assertRaisesMessage(MailboxError, 'Nieprawidłowy DOCX: zly.docx.'):
            _package_messages([self.message(files=[('zly.docx', b'not a zip')])], True, False, False)
        self.assertEqual(self.names(_package_messages([self.message(files=[('tekst.pdf', b'%PDF')])], False, False, False)),
                         ['Opowieść/01_tekst.pdf'])

    @patch('core.services.mailbox_import.convert_document', side_effect=lambda *args, **kwargs: (BytesIO(b'x'), None, None))
    def test_mail_documents_keep_ten_mb_limit(self, convert):
        """DOCX z poczty sprawdzamy jak pliki w programach, ale z limitem 10 MB zamiast 5 MB."""
        def package(size):
            return _package_messages([self.message(files=[('tekst.docx', sized_docx(size).read())])], True, False, False)

        self.assertEqual(self.names(package(5 * 1024 * 1024 + 1)), ['Opowieść/01_tekst.docx'])
        with self.assertRaisesMessage(MailboxError, 'Nieprawidłowy DOCX: tekst.docx. Plik jest za duży. Maksymalny rozmiar to 10 MB.'):
            package(10 * 1024 * 1024 + 1)

    @patch('core.services.mailbox_import.inspect_document', side_effect=RebuildConfirmationRequired('pominie tabelę'))
    def test_rebuild_omissions_need_confirmation(self, inspect):
        with self.assertRaisesMessage(RebuildConfirmationRequired, 'tekst.docx: pominie tabelę'):
            _package_messages([self.message()], False, False, True)
        warnings, accepted, errors = [], [], []
        with package_messages([self.message()], False, False, True, skipped_errors=errors, accepted=accepted,
                              rebuild_warnings=warnings) as output, ZipFile(output) as archive:
            self.assertEqual(archive.namelist(), [])
        self.assertEqual(warnings, ['Autor 12: tekst.docx: pominie tabelę'])
        self.assertEqual([message['uid'] for message in accepted], [12])
        self.assertEqual(errors, [])

    @patch('core.services.mailbox_import.convert_document')
    def test_processed_documents_lose_embedded_fonts(self, convert):
        """Każde przetworzenie DOCX ze skrzynki (Odkurzacz, konwersja, przebudowa) usuwa osadzone czcionki,
        a z Odkurzaczem także zamienia miękkie entery; bez przetwarzania zapisujemy oryginał bez zmian."""
        convert.side_effect = lambda *args, **kwargs: (BytesIO(b'przetworzony'), None, None)
        for options in ((True, False, False), (False, False, True)):
            with self.subTest(options=options):
                convert.reset_mock()
                _package_messages([self.message()], *options, True)
                self.assertIs(convert.call_args.kwargs['remove_fonts'], True)
                # Miękkie entery – tylko z Odkurzaczem (pierwsza opcja).
                self.assertIs(convert.call_args.kwargs['remove_soft_whitespace'], options[0])
        convert.side_effect = lambda *args, **kwargs: converted()
        _package_messages([self.message()], False, True, False)
        self.assertIs(convert.call_args.kwargs['remove_fonts'], True)
        convert.reset_mock()
        _package_messages([self.message()], False, False, False)
        convert.assert_not_called()

    @patch('core.services.mailbox_import.convert_document')
    def test_rebuild_only_and_conversion_with_notes(self, convert):
        convert.side_effect = lambda *args, **kwargs: (BytesIO(b'przebudowany'), None, None)
        self.assertEqual(self.names(_package_messages([self.message()], False, False, True, True)),
                         ['Opowieść/01_tekst.docx'])
        self.assertEqual(convert.call_args.args[1], [])
        self.assertTrue(convert.call_args.kwargs['allow_rebuild_omissions'])
        convert.side_effect = lambda *args, **kwargs: converted(warnings=['Brak przypisu 3'])
        self.assertEqual(self.names(_package_messages([self.message()], False, True, False)),
                         ['Opowieść/01_tekst.docx', 'Opowieść/01_tekst.epub', 'Opowieść/01_tekst.pdf',
                          'Opowieść/01_tekst_uwagi.txt'])
        convert.side_effect = RebuildConfirmationRequired('utrata obrazu')
        with self.assertRaisesMessage(RebuildConfirmationRequired, 'tekst.docx: utrata obrazu'):
            _package_messages([self.message()], False, True, True, True)

    def test_equal_folders_get_uid_suffix(self):
        messages = [self.message(12), self.message(13)]
        self.assertEqual(self.names(_package_messages(messages, False, False, False)),
                         ['Opowieść/01_tekst.docx', 'Opowieść_13/01_tekst.docx'])

    def test_time_limits(self):
        """Limit 90 sekund: w trakcie i po zakończeniu paczki jednej partii oraz całej paczki wiadomości."""
        def per_message(messages, *options):
            return package_messages(messages, *options, skipped_errors=[], accepted=[], rebuild_warnings=[])

        for package, ticks in (
            (_package_messages, [0, 100]),
            (_package_messages, [0, 0, 100]),
            (per_message, [0, 100]),
            (per_message, [0, 0, 0, 0, 0, 100]),
        ):
            with (
                self.subTest(package=package.__name__, ticks=ticks),
                patch('core.services.mailbox_import.time.monotonic', side_effect=chain(ticks, repeat(ticks[-1]))),
                self.assertRaisesMessage(MailboxError, '90 sekund'),
            ):
                package([self.message()], False, False, False)


LIMIT = 150 * 1024 * 1024
TOO_BIG = LIMIT + 1


def big_zip(name, size):
    """Archiwum z jednym plikiem zer o podanym rozmiarze po rozpakowaniu (mały po kompresji)."""
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
        archive.writestr(name, bytes(size))
    stream.seek(0)
    return stream


@tag('slow', 'docx')
class SubmissionPackageLimitTests(SimpleTestCase):
    """Limit 150 MB wyniku (przebudowa, konwersja, paczka wielu wiadomości); wolny: kompresja 150 MB zer."""

    def setUp(self):
        self.message = parse_message(12, csv_mail())

    def test_package_over_150_mb_is_refused_for_every_output_kind(self):
        """Przebudowany DOCX, wynik konwersji i paczka z wielu wiadomości mają ten sam limit 150 MB."""
        converted = big_zip('document.docx', 1)
        with ZipFile(converted, 'a', ZIP_DEFLATED) as archive:
            archive.writestr('document.pdf', bytes(TOO_BIG))
        converted.seek(0)
        cases = {
            'przebudowa bez konwersji': (dict(convert=False), (BytesIO(bytes(TOO_BIG)), 'docx', '')),
            'konwersja do PDF i EPUB': (dict(convert=True), (converted, 'zip', '')),
        }
        for name, (options, result) in cases.items():
            with self.subTest(name), patch('core.services.mailbox_import.convert_document', return_value=result):
                with self.assertRaisesMessage(MailboxError, 'Wynik przekracza 150 MB'):
                    _package_messages([self.message], clean=False, rebuild=True, allow_rebuild_omissions=True,
                                      **options)
        with self.subTest('paczka wielu wiadomości'), \
                patch('core.services.mailbox_import._package_messages',
                      return_value=big_zip('folder/01_tekst.docx', TOO_BIG)):
            skipped = []
            with self.assertRaisesMessage(MailboxError, 'Wynik przekracza 150 MB'):
                package_messages([self.message], skipped_errors=skipped, accepted=[], rebuild_warnings=[])
            self.assertEqual(skipped, [])
