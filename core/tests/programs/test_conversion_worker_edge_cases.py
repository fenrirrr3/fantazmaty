"""Proces roboczy konwertera: start (limit pamięci, brak pamięci, brak modułu resource), odnośniki między
rozdziałami EPUB, obrazy odrzucane przy konwersji, raport błędu, którego nie da się zapisać, oraz zapis
rzadszych edytorów: polskie nazwy stylów nagłówków i uwagi o elementach pominiętych w e-booku."""
import json
import os
import runpy
import struct
import sys
import zlib
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase, tag
from docx import Document
from PIL import Image

import core.services
from core.services import document_progress
from fpdf import FPDF

from core.services.document_conversion_worker import convert, cross_chapter_links, main, polish_warnings
from core.services.document_errors import DocumentInputError

from .editor_fixtures_exotic import localized_style_names
from .support import png

BOOTSTRAP = Path(core.services.__file__).with_name('document_worker_bootstrap.py')


class WorkerDirectoryMixin:
    def setUp(self):
        super().setUp()
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        # Proces roboczy kieruje raport postępu do katalogu zadania; po teście wyłączamy go.
        progress = patch.object(document_progress, '_target', None)
        progress.start()
        self.addCleanup(progress.stop)


class BootstrapTests(WorkerDirectoryMixin, SimpleTestCase):
    def run_bootstrap(self, **patches):
        """Uruchom bootstrap jako __main__ z podmienionym startem procesu roboczego; zwraca kod wyjścia lub None."""
        runner = runpy.run_path
        (self.directory / 'job.json').write_text(json.dumps({}), encoding='utf-8')
        with patch('sys.argv', ['bootstrap', str(self.directory)]), patch('runpy.run_path', **patches):
            try:
                runner(str(BOOTSTRAP), run_name='__main__')
            except SystemExit as stop:
                return stop.code
        return None

    def report(self):
        return json.loads((self.directory / 'error.json').read_text(encoding='utf-8'))

    def test_memory_error_is_reported_without_the_frames_of_later_errors(self):
        """Brak pamięci – także wtedy, gdy podczas jego obsługi wystąpił kolejny błąd – to raport MemoryError."""
        def nested(*args, **kwargs):
            try:
                raise MemoryError
            except MemoryError:
                raise ValueError('PRYWATNE dane z dokumentu') from None

        cases = {'bezpośrednio': MemoryError, 'w trakcie innego błędu': nested}
        for label, failure in cases.items():
            with self.subTest(label), patch('resource.setrlimit'):
                self.assertEqual(self.run_bootstrap(side_effect=failure), 3)
                report = self.report()
                self.assertEqual((report['stage'], report['error']), ('konwersja', 'MemoryError'))
                self.assertTrue(report['frames'])
                self.assertNotIn('PRYWATNE', json.dumps(report))

    def test_platform_without_resource_module_runs_without_a_limit(self):
        with patch.dict(sys.modules, {'resource': None}):
            self.assertIsNone(self.run_bootstrap())
        self.assertEqual(json.loads((self.directory / 'progress.json').read_text())['stage'], 'start')
        self.assertFalse((self.directory / 'error.json').exists())

    def test_finite_hard_limit_caps_the_memory_limit(self):
        import resource
        hard = 600 * 1024 * 1024
        with patch('resource.getrlimit', return_value=(hard, hard)), patch('resource.setrlimit') as setrlimit:
            self.assertIsNone(self.run_bootstrap())
        setrlimit.assert_called_once_with(resource.RLIMIT_AS, (hard, hard))


class CrossChapterLinkTests(SimpleTestCase):
    def test_only_links_to_other_chapters_get_a_file_name(self):
        chapters = [
            ('chapter-1.xhtml', '<p id="a">A <a href="#a">tutaj</a> <a href="#brak">nigdzie</a></p>'),
            ('chapter-2.xhtml', '<p>B <a href="#a">do pierwszego</a> <a href="#b">tutaj</a></p><p id="b">C</p>'),
        ]
        first, second = cross_chapter_links(chapters)
        self.assertEqual(first, chapters[0][1])  # rozdział bez odnośników do innych plików zostaje bez zmian
        self.assertIn('href="chapter-1.xhtml#a"', second)
        self.assertIn('href="#b"', second)
        self.assertEqual(cross_chapter_links(chapters[:1]), [chapters[0][1]])


def png_chunk(kind, data):
    return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xFFFFFFFF)


class RejectedImageTests(WorkerDirectoryMixin, SimpleTestCase):
    def source(self, *pictures):
        """DOCX z obrazami; ich bajty w paczce zastępujemy podanymi (kolejność plików word/media)."""
        document = Document()
        document.add_paragraph('Ilustracja')
        for _ in pictures:
            document.add_picture(png(4, 4))
        original = BytesIO()
        document.save(original)
        path = self.directory / 'source.docx'
        replacements = iter(pictures)
        with ZipFile(original) as source, ZipFile(path, 'w') as target:
            for item in source.infolist():
                data = source.read(item)
                if item.filename.startswith('word/media/'):
                    data = next(replacements)
                target.writestr(item, data)
        return path

    def test_damaged_or_oversized_images_get_public_codes(self):
        header = png(4, 4).getvalue()[:33]  # sygnatura i IHDR: Pillow otwiera plik, ale nie wczyta danych
        truncated = header + png_chunk(b'IDAT', zlib.compress(b'\x00' * 3)[:4])
        oversized = png(4, 4).getvalue() + b'\x00' * (20 * 1024 * 1024)
        for label, picture, code in (('ucięte dane obrazu', truncated, 'unsupported_image'),
                                     ('ponad 20 MB', oversized, 'image_too_large')):
            with self.subTest(label), self.assertRaises(DocumentInputError) as raised:
                convert(self.source(picture), self.directory, ['epub'], 'Obraz')
            self.assertEqual(raised.exception.public_code, code)

    @tag('slow')
    def test_images_over_the_total_limit_are_refused(self):
        """Każdy obraz mieści się w limicie, ale razem po rasteryzacji przekraczają 30 MB."""
        noise = BytesIO()
        Image.frombytes('RGB', (2150, 2150), os.urandom(2150 * 2150 * 3)).save(noise, 'PNG', compress_level=0)
        with self.assertRaises(DocumentInputError) as raised:
            convert(self.source(noise.getvalue(), noise.getvalue()), self.directory, ['epub'], 'Obraz')
        self.assertEqual(raised.exception.public_code, 'too_many_images')


class UnwritableReportTests(WorkerDirectoryMixin, SimpleTestCase):
    def test_failure_is_returned_even_when_the_report_cannot_be_saved(self):
        document = Document()
        document.add_paragraph('Tekst')
        document.save(self.directory / 'source.docx')
        (self.directory / 'job.json').write_text(json.dumps({'formats': ['epub'], 'title': 'T'}))
        (self.directory / 'error.json').mkdir()  # zapis raportu kończy się błędem systemu plików
        with patch('sys.argv', ['worker', str(self.directory)]), \
                patch('core.services.document_conversion_worker.convert', side_effect=RuntimeError):
            self.assertEqual(main(), 3)
        self.assertTrue((self.directory / 'error.json').is_dir())


class EditorMarkupConversionTests(WorkerDirectoryMixin, SimpleTestCase):
    def test_polish_heading_names_make_chapters_and_the_pdf_outline(self):
        """Regresja: styl „Nagłówek 1” (polska nazwa zamiast wbudowanej „heading 1”) był w e-booku zwykłym
        akapitem – bez rozdziałów, spisu treści i konspektu PDF."""
        source = self.directory / 'source.docx'
        source.write_bytes(localized_style_names().getvalue())
        sections, original = [], FPDF.start_section

        def section_spy(pdf, name, level=0, strict=True):
            sections.append((name, level))
            return original(pdf, name, level, strict)
        with patch.object(FPDF, 'start_section', section_spy):
            convert(source, self.directory, ['epub', 'pdf'], 'Książka')
        with ZipFile(self.directory / 'document.epub') as epub:
            chapters = sorted(name for name in epub.namelist() if name.startswith('EPUB/chapter-'))
            second = epub.read(chapters[1]).decode('utf-8')
        self.assertEqual(len(chapters), 3)  # strona tytułowa i dwa rozdziały
        self.assertIn('<h1 id="section-1" class="align-left">Rozdział pierwszy</h1>', second)
        self.assertIn('<h2 id="section-2" class="align-left">Podrozdział</h2>', second)
        self.assertEqual(sections, [('Rozdział pierwszy', 0), ('Podrozdział', 1), ('Rozdział drugi', 0)])

    def test_warnings_about_equations_attachments_shapes_and_linked_images(self):
        """Uwagi o równaniach, wstawkach altChunk i obrazach połączonych są po polsku; wygląd kształtów VML
        (wypełnienie, obrys) z zapasowej wersji pola tekstowego – np. z LibreOffice – nie jest zgłaszany."""
        unrecognised = 'An unrecognised element was ignored: '
        cases = {
            unrecognised + '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara':
                'Pominięto równanie (wzór matematyczny) – e-book i PDF go nie zawierają.',
            unrecognised + '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath':
                'Pominięto równanie (wzór matematyczny) – e-book i PDF go nie zawierają.',
            unrecognised + 'w:altChunk': 'Pominięto treść dołączoną z innego pliku (w:altChunk) – otwórz dokument '
                                          'w Wordzie i zapisz go ponownie, aby ją scalić.',
            "could not open external image 'file:///C:/Obrazy/a.jpg', external file access is disabled":
                'Pominięto obraz połączony z plikiem spoza dokumentu – wstaw obraz do DOCX.',
            unrecognised + 'v:fill': None,
            unrecognised + 'v:stroke': None,
            unrecognised + 'o:lock': None,
            unrecognised + 'w10:wrap': None,
            unrecognised + 'w:ruby': 'Pominięto nieobsługiwany element dokumentu (w:ruby).',
        }
        for message, expected in cases.items():
            with self.subTest(message):
                self.assertEqual(polish_warnings([message]), [expected] if expected else [])
