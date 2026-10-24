"""Przebudowa dokumentu do nowego DOCX: zachowana treść i formatowanie, usunięte metadane, ostrzeżenia o pominięciach;
zmiany śledzone i komentarze zostają (także po ujednoliceniu formatowania, zamianie miękkich enterów i Odkurzaczu)."""
import re
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile

from django.test import SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from core.program_test_support import run_program
from core.services import program_jobs as jobs
from core.services.document_converter import RebuildConfirmationRequired, convert_document
from core.services.document_formatting import _SEQUENCES
from core.services.document_preparation import DEFAULT_EDITORIAL_RULES, prepare_docx
from core.services.document_rebuild import (FORMAT_HISTORY_OMISSION, HYPERLINK_OMISSION, LOST_COMMENTS_OMISSION,
                                            PLACEHOLDER_OMISSION, UNKNOWN_FORMAT_OMISSION, RebuildUnsupported,
                                            inspect_docx, rebuild_docx)
from core.services.mailbox_import import package_messages
from fantazmaty.testing import create_superuser

from .editor_fixtures import _rels
from .editor_fixtures_exotic import (_comment_parts, google_docs_extras_docx, libreoffice_review_docx, note_reference,
                                     notes_part, p, r, review_docx, word_docx)
from .support import (
    NOTE_PARTS, IsolatedProgramDirsMixin, add_block, add_complex_field, add_text_box, content_control,
    downloaded_document, noted_document, noted_text, saved, text_run,
)

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{%s}' % NS['w']
W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'
W15 = '{http://schemas.microsoft.com/office/word/2012/wordml}'
REVISION_NAMES = ('ins', 'del', 'moveFrom', 'moveTo')
# Kolejność schematu: znak akapitu (w:pPr/w:rPr) zaczyna się od zmian, potem formatowanie jak w przebiegu.
SCHEMA_ORDER = {W + 'rPr': [W + name for name in REVISION_NAMES + tuple(_SEQUENCES['w:rPr'].split())],
                W + 'pPr': [W + name for name in _SEQUENCES['w:pPr'].split()]}


# --- Recenzja: zmiany śledzone i komentarze w pakiecie (pomocniki testów) ---------------------------------

def review_parts(stream):
    """Korzenie XML części pakietu: document, footnotes, comments, commentsExtended (brak części – None)."""
    with ZipFile(stream) as archive:
        names = set(archive.namelist())
        return {name: etree.fromstring(archive.read(f'word/{name}.xml')) if f'word/{name}.xml' in names else None
                for name in ('document', 'footnotes', 'comments', 'commentsExtended')}


def local(node):
    return etree.QName(node).localname


def text_of(node):
    """Tekst węzła: w:t i w:delText, tabulator i podział wiersza (bez tabulatorów z w:pPr/w:tabs)."""
    glyphs = {W + 'tab': '\t', W + 'br': '\n', W + 'cr': '\n'}
    return ''.join(glyphs.get(item.tag, item.text or '') for item in node.iter(W + 't', W + 'delText', *glyphs)
                   if item.getparent().tag == W + 'r')


def stories(parts):
    """Akapity kolejnych tekstów: treść główna (akapity wprost w w:body) i każdy zwykły przypis dolny."""
    result = [parts['document'].find(W + 'body').findall(W + 'p')]
    if parts['footnotes'] is not None:
        result += [note.findall(W + 'p') for note in parts['footnotes'] if note.get(W + 'type') is None]
    return result


def resolved(parts, accept):
    """Teksty akapitów (treści i każdego przypisu) po przyjęciu (accept=True) albo odrzuceniu wszystkich zmian,
    jak w Wordzie: przyjęcie usuwa w:del/w:moveFrom, odrzucenie – w:ins/w:moveTo; akapit, którego znak akapitu
    znika (usunięty przy przyjęciu, wstawiony przy odrzuceniu), łączy się z następnym."""
    gone = ('del', 'moveFrom') if accept else ('ins', 'moveTo')
    result = []
    for paragraphs in stories(parts):
        texts, carried = [], None
        for paragraph in paragraphs:
            paragraph = deepcopy(paragraph)
            for node in list(paragraph.iter(*(W + name for name in REVISION_NAMES))):
                parent = node.getparent()
                if parent is None or parent.tag == W + 'rPr':
                    continue
                if local(node) not in gone:
                    for child in list(node):
                        node.addprevious(child)
                parent.remove(node)
            text = (carried or '') + text_of(paragraph)
            mark = paragraph.find(f'{W}pPr/{W}rPr')
            if mark is not None and any(local(child) in gone for child in mark):
                carried = text
            else:
                texts.append(text)
                carried = None
        result.append(texts + ([carried] if carried is not None else []))
    return result


def revisions(parts):
    """Śledzone zmiany w kolejności tekstu: (rodzaj, autor, data, tekst; dla znaku akapitu None)."""
    return [(local(node), node.get(W + 'author'), node.get(W + 'date'),
             None if node.getparent().tag == W + 'rPr' else text_of(node))
            for paragraphs in stories(parts) for paragraph in paragraphs
            for node in paragraph.iter(*(W + name for name in REVISION_NAMES))]


def format_changes(parts):
    """Zmiany formatowania: (rodzaj, autor, data, tekst, obecne i poprzednie pogrubienie/kursywa/wyrównanie)."""
    def marks(properties, names=('b', 'i', 'jc')):
        if properties is None:
            return ()
        return tuple(f'{name}={node.get(W + "val", "1")}' for name in names for node in properties.findall(W + name))

    found = []
    for paragraphs in stories(parts):
        for paragraph in paragraphs:
            for change in paragraph.iter(W + 'rPrChange', W + 'pPrChange'):
                owner = change.getparent()
                text = text_of(owner.getparent() if local(change) == 'rPrChange' else paragraph)
                found.append((local(change), change.get(W + 'author'), change.get(W + 'date'), text, marks(owner),
                              marks(change.find(W + ('rPr' if local(change) == 'rPrChange' else 'pPr')))))
    return found


def comment_anchors(parts):
    """w:id komentarza → (tekst zakresu z ¶ między akapitami albo None, słowo tuż przed odwołaniem)."""
    anchors = {}
    for root in (parts['document'], parts['footnotes']):
        opened, before = {}, ''
        for node in root.iter() if root is not None else ():
            key = node.get(W + 'id')
            if node.tag == W + 'p':
                opened = {name: text + '¶' if text else text for name, text in opened.items()}
                before = ''
            elif node.tag == W + 'commentRangeStart':
                opened[key] = ''
            elif node.tag == W + 'commentRangeEnd':
                anchors.setdefault(key, [None, None])[0] = opened.pop(key, None)
            elif node.tag == W + 'commentReference':
                anchors.setdefault(key, [None, None])[1] = (before.split() or [''])[-1]
            elif node.tag in (W + 't', W + 'delText'):
                opened = {name: text + (node.text or '') for name, text in opened.items()}
                before += node.text or ''
    return {key: tuple(value) for key, value in anchors.items()}


def comments_of(parts):
    """Komentarze (posortowane): autor, inicjały, data, treść (akapity), pogrubione i pochylone fragmenty treści,
    tekst zakresu, słowo przed odwołaniem, treść komentarza nadrzędnego (odpowiedź) i stan „rozwiązany”."""
    if parts['comments'] is None:
        return []
    threads = {entry.get(W15 + 'paraId'): (entry.get(W15 + 'paraIdParent'), entry.get(W15 + 'done') in ('1', 'true'))
               for entry in (parts['commentsExtended'] if parts['commentsExtended'] is not None else ())}
    anchors, contents, found = comment_anchors(parts), {}, []
    for comment in parts['comments'].findall(W + 'comment'):
        paragraphs = comment.findall(W + 'p')
        contents[paragraphs[-1].get(W14 + 'paraId')] = '\n'.join(text_of(paragraph) for paragraph in paragraphs)
    for comment in parts['comments'].findall(W + 'comment'):
        paragraphs = comment.findall(W + 'p')
        parent, done = threads.get(paragraphs[-1].get(W14 + 'paraId'), (None, False))
        emphasis = tuple((text_of(run), local(flag)) for run in comment.iter(W + 'r')
                         for flag in run.findall(f'{W}rPr/*') if local(flag) in ('b', 'i', 'u'))
        found.append((comment.get(W + 'author'), comment.get(W + 'initials'), comment.get(W + 'date'),
                      '\n'.join(text_of(paragraph) for paragraph in paragraphs), emphasis,
                      *anchors.get(comment.get(W + 'id'), (None, None)), contents.get(parent) if parent else None, done))
    return sorted(found, key=repr)


def review_markup_errors(parts):
    """Błędy zapisu ważne dla Worda: powtórzone w:id (zmiany i komentarze razem), zakresy komentarzy i przeniesień
    bez pary, przeniesienie bez drugiej strony, odwołania do brakujących komentarzy i komentarze bez odwołania,
    kolejność elementów w:rPr i w:pPr (także w zmianach formatowania) niezgodna ze schematem."""
    roots = [root for name, root in parts.items() if name in ('document', 'footnotes') and root is not None]

    def found(tag, attribute=W + 'id'):
        return sorted(node.get(attribute) for root in roots for node in root.iter(W + tag))

    errors = []
    changed = [value for tag in (*REVISION_NAMES, 'rPrChange', 'pPrChange', 'moveFromRangeStart', 'moveToRangeStart')
               for value in found(tag)]
    comments = [node.get(W + 'id') for node in parts['comments']] if parts['comments'] is not None else []
    identifiers = changed + comments
    if len(identifiers) != len(set(identifiers)):
        errors.append(f'powtórzone w:id: {sorted(set(i for i in identifiers if identifiers.count(i) > 1))}')
    if found('commentReference') != sorted(comments):
        errors.append(f'odwołania {found("commentReference")} a komentarze {sorted(comments)}')
    for start, end in (('commentRangeStart', 'commentRangeEnd'), ('moveFromRangeStart', 'moveFromRangeEnd'),
                       ('moveToRangeStart', 'moveToRangeEnd')):
        if found(start) != found(end) or len(set(found(start))) != len(found(start)):
            errors.append(f'{start} {found(start)} a {end} {found(end)}')
    if found('moveFromRangeStart', W + 'name') != found('moveToRangeStart', W + 'name'):
        errors.append('przeniesienie bez drugiej strony')
    for root in roots:
        for node in root.iter(W + 'rPr', W + 'pPr'):
            order = SCHEMA_ORDER[node.tag]
            positions = [order.index(child.tag) if child.tag in order else -1 for child in node]
            if -1 in positions or positions != sorted(positions):
                errors.append(f'kolejność w {local(node)}: {[local(child) for child in node]}')
    return sorted(set(errors))



class RebuildTests(SimpleTestCase):
    def test_text_format_empty_paragraphs_and_metadata(self):
        doc = Document()
        doc.core_properties.author = 'Secret'
        doc.core_properties.title = 'Private'
        doc.sections[0].header.paragraphs[0].text = 'Private header'
        doc.styles['Normal'].font.italic = True
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.add_run('Tekst').bold = True
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph('Środek').alignment = WD_ALIGN_PARAGRAPH.CENTER
        with rebuild_docx(saved(doc), allow_omissions=True) as out:
            result = Document(out)
            self.assertEqual([p.text for p in result.paragraphs], ['Tekst', '', '', 'Środek'])
            self.assertTrue(result.paragraphs[0].runs[0].bold)
            self.assertTrue(result.paragraphs[0].runs[0].italic)
            self.assertEqual(result.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertEqual(result.paragraphs[-1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(result.core_properties.author, '')
            self.assertIsNone(result.core_properties.created)
            with ZipFile(out) as archive:
                self.assertNotIn('word/header1.xml', archive.namelist())
                self.assertNotIn(b'Secret', archive.read('docProps/core.xml'))
            self.assertAlmostEqual(result.sections[0].page_width.cm, 21, places=2)
            self.assertAlmostEqual(result.sections[0].left_margin.cm, 2.5, places=2)

    def test_doc_defaults_and_theme_fonts_are_kept(self):
        doc = Document()
        doc.add_paragraph('Test')
        defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
        defaults.find(qn('w:sz')).set(qn('w:val'), '34')
        font = defaults.find(qn('w:rFonts'))
        for key in list(font.attrib):
            del font.attrib[key]
        font.set(qn('w:ascii'), 'Courier New')
        font.set(qn('w:hAnsi'), 'Courier New')
        rebuilt = Document(rebuild_docx(saved(doc)))
        self.assertEqual(rebuilt.paragraphs[0].runs[0].font.size.pt, 17)
        self.assertEqual(rebuilt.paragraphs[0].runs[0].font.name, 'Courier New')

    def test_revisions_and_hidden_text(self):
        """Ukryty tekst znika; wstawka i usunięcie zostają śledzonymi zmianami (tekst po przyjęciu i odrzuceniu)."""
        doc = Document()
        p = doc.add_paragraph('A')
        p.add_run('hidden').font.hidden = True
        for tag_name, text in [('ins', 'B'), ('del', 'C')]:
            wrapper = OxmlElement('w:' + tag_name)
            wrapper.set(qn('w:author'), 'Redaktor')
            run = OxmlElement('w:r')
            t = OxmlElement('w:t' if tag_name == 'ins' else 'w:delText')
            t.text = text
            run.append(t)
            wrapper.append(run)
            p._p.append(wrapper)
        parts = review_parts(rebuild_docx(saved(doc)))
        self.assertEqual((resolved(parts, accept=True), resolved(parts, accept=False)), ([['AB']], [['AC']]))
        self.assertEqual(revisions(parts), [('ins', 'Redaktor', None, 'B'), ('del', 'Redaktor', None, 'C')])

    def test_tables_and_numbering_fail_explicitly(self):
        for kind in ('table', 'list'):
            doc = Document()
            if kind == 'table':
                doc.add_table(rows=1, cols=1)
            else:
                doc.add_paragraph('item', 'List Number')
            with self.assertRaises(RebuildUnsupported):
                rebuild_docx(saved(doc))

    def test_nonempty_header_requires_rebuild_acceptance(self):
        document = Document()
        document.add_paragraph('Treść')
        document.sections[0].header.paragraphs[0].text = 'Ważny nagłówek'
        with self.assertRaises(RebuildUnsupported) as error:
            inspect_docx(saved(document))
        self.assertIn('nagłówki', str(error.exception))
        with rebuild_docx(saved(document), allow_omissions=True) as result:
            self.assertEqual(Document(result).paragraphs[0].text, 'Treść')

    def test_loss_warning_lists_sections_and_links_and_keeps_headings(self):
        doc = Document()
        doc.add_heading('Rozdział', 1)
        doc.add_section()
        link = OxmlElement('w:hyperlink')
        run = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = 'Odnośnik'
        run.append(t)
        link.append(run)
        doc.add_paragraph()._p.append(link)
        with self.assertRaises(RebuildUnsupported) as error:
            rebuild_docx(saved(doc))
        self.assertIn('sekcje', str(error.exception))
        self.assertIn('hiperłączy', str(error.exception))
        self.assertIn('liczba: 1', str(error.exception))
        rebuilt = Document(rebuild_docx(saved(doc), allow_omissions=True))
        self.assertEqual(rebuilt.paragraphs[0].style.name, 'Heading 1')


class RebuildPageLayoutWarningTests(SimpleTestCase):
    """Przebudowa ustawia A4 – inny format książki, orientacja lub kolumny wymagają potwierdzenia."""

    def test_book_format_landscape_and_columns_are_reported(self):
        for name, change in (
            ('A5', lambda section: (setattr(section, 'page_width', Cm(14.8)),
                                    setattr(section, 'page_height', Cm(21)))),
            ('poziomo', lambda section: (setattr(section, 'page_width', Cm(29.7)),
                                         setattr(section, 'page_height', Cm(21)))),
            ('kolumny', lambda section: section._sectPr.find(qn('w:cols')).set(qn('w:num'), '2')),
        ):
            with self.subTest(name):
                document = Document()
                document.add_paragraph('Treść')
                change(document.sections[0])
                with self.assertRaises(RebuildUnsupported) as error:
                    inspect_docx(saved(document))
                self.assertIn('format strony, orientacja lub kolumny', str(error.exception))
                rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
                self.assertAlmostEqual(rebuilt.sections[0].page_width.cm, 21, places=1)

    def test_a4_and_letter_with_any_margins_need_no_confirmation(self):
        for width, height in ((21, 29.7), (21.59, 27.94)):
            with self.subTest(width=width):
                document = Document()
                document.add_paragraph('Treść')
                section = document.sections[0]
                section.page_width, section.page_height = Cm(width), Cm(height)
                section.left_margin = section.right_margin = Cm(2)
                self.assertIsNone(inspect_docx(saved(document)))  # bez RebuildUnsupported
                # Przebudowa bez zgody na pominięcia przechodzi i zachowuje treść; wynik ma format A4.
                rebuilt = Document(rebuild_docx(saved(document)))
                self.assertEqual([paragraph.text for paragraph in rebuilt.paragraphs], ['Treść'])
                self.assertAlmostEqual(rebuilt.sections[0].page_width.cm, 21, places=1)


class RebuildDetailTests(SimpleTestCase):
    """Elementy wewnątrz akapitów: zachowane (tabulatory, podziały, łączniki, kolory motywu) albo zgłoszone."""

    @staticmethod
    def element(tag, **attributes):
        node = OxmlElement(tag)
        for key, value in attributes.items():
            node.set(qn(key), value)
        return node

    def run_with(self, paragraph, text, *properties):
        run = paragraph.add_run(text)
        for node in properties:
            run._r.get_or_add_rPr().append(node)
        return run

    def detailed_document(self):
        element = self.element
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5))
        run = paragraph.add_run('A')
        run.add_tab()
        run.add_text('B')
        run.add_break(WD_BREAK.PAGE)
        run.add_break()
        run.add_text('C')
        for node in (element('w:noBreakHyphen'), element('w:softHyphen'), element('w:commentReference', **{'w:id': '1'}),
                     element('w:ptab')):
            run._r.append(node)
        self.run_with(paragraph, 'ciemny', element('w:color', **{'w:val': 'FF0000', 'w:themeShade': '80'}))
        self.run_with(paragraph, 'jasny', element('w:color', **{'w:val': '000000', 'w:themeTint': '80'}))
        self.run_with(paragraph, 'auto', element('w:color', **{'w:val': 'auto'}))
        self.run_with(paragraph, 'motyw', element('w:rFonts', **{'w:asciiTheme': 'majorHAnsi'}))
        self.run_with(paragraph, 'rozstrzelony', element('w:spacing', **{'w:val': '20'}))
        framed = document.add_paragraph('Ramka')
        properties = framed._p.get_or_add_pPr()
        properties.append(element('w:pBdr'))
        mark = element('w:rPr')
        mark.append(element('w:del', **{'w:id': '2', 'w:author': 'Ktoś'}))
        mark.append(element('w:ins', **{'w:id': '3', 'w:author': 'Ktoś'}))
        properties.append(mark)
        document.element.body.insert(0, element('w:bookmarkStart', **{'w:id': '9', 'w:name': 'zakładka'}))
        return document

    def test_inline_content_is_rebuilt_and_unsupported_parts_are_listed(self):
        document = self.detailed_document()
        with self.assertRaises(RebuildUnsupported) as caught:
            rebuild_docx(saved(document))
        # Usunięty i wstawiony znak akapitu (w:pPr/w:rPr) nie są już pomijane – zostają w nowym dokumencie.
        self.assertEqual(caught.exception.omissions, [
            'element ptab wewnątrz tekstu (liczba: 1)',
            'obramowania akapitów (liczba: 1)',
            'rozstrzelenie znaków (liczba: 1)',
        ])
        result = Document(rebuild_docx(saved(document), allow_omissions=True))
        self.assertEqual([etree.QName(node).localname for node in result.paragraphs[1]._p.xpath('w:pPr/w:rPr/*')][:2],
                         ['ins', 'del'])
        paragraph = result.paragraphs[0]
        self.assertEqual(paragraph.text, 'A\tB\nC\u2011\u00adciemnyjasnyautomotywrozstrzelony')
        self.assertEqual(paragraph._p.xpath('.//w:br/@w:type'), ['page'])
        self.assertAlmostEqual(paragraph.paragraph_format.tab_stops[0].position.cm, 5, places=2)
        runs = {run.text: run for run in paragraph.runs}
        self.assertEqual(runs['ciemny'].font.color.rgb, RGBColor(0x80, 0, 0))
        self.assertEqual(runs['jasny'].font.color.rgb, RGBColor(0x7F, 0x7F, 0x7F))
        self.assertIsNone(runs['auto'].font.color.rgb)
        self.assertNotEqual(runs['motyw'].font.name, runs['auto'].font.name)
        self.assertEqual(result.paragraphs[1].text, 'Ramka')

    def test_theme_fonts_without_a_theme_part_are_reported(self):
        document = Document()
        document.add_paragraph('Tekst')
        for key, relationship in list(document.part.rels.items()):
            if relationship.reltype.endswith('/theme'):
                document.part.drop_rel(key)
        with self.assertRaisesMessage(RebuildUnsupported, 'nierozpoznany krój czcionki z motywu'):
            inspect_docx(saved(document))

    def test_missing_theme_font_is_reported_once_per_document(self):
        """Regresja: brak czcionki motywu to jedna uwaga, nie jedna na każdy fragment tekstu."""
        document = Document()
        for _ in range(3):
            document.add_paragraph('Tekst').add_run(' dalej')
        for key, relationship in list(document.part.rels.items()):
            if relationship.reltype.endswith('/theme'):
                document.part.drop_rel(key)
        with self.assertRaises(RebuildUnsupported) as caught:
            inspect_docx(saved(document))
        self.assertEqual(caught.exception.omissions,
                         ['nierozpoznany krój czcionki z motywu (zastosujemy domyślny) (liczba: 1)'])

    def test_one_equation_is_counted_once(self):
        """Regresja: części równania (m:r, m:t) nie są liczone jako osobne równania."""
        document = Document()
        document.add_paragraph('Wzór')._p.append(
            parse_xml('<m:oMath %s><m:r><m:t>x</m:t></m:r></m:oMath>' % nsdecls('m')))
        with self.assertRaisesMessage(RebuildUnsupported, 'równania matematyczne (liczba: 1)'):
            inspect_docx(saved(document))


def omissions_of(document):
    """Lista ostrzeżeń przebudowy (pusta, gdy nic nie zostanie pominięte)."""
    try:
        inspect_docx(saved(document))
    except RebuildUnsupported as error:
        return error.omissions
    return []


def drop_doc_defaults(document):
    """Dokument bez ustawień domyślnych (jak z LibreOffice lub starszego Worda): Word stosuje własne wartości bazowe."""
    defaults = document.styles.element.find(qn('w:docDefaults'))
    for child in list(defaults):
        defaults.remove(child)


class RebuildFidelityTests(SimpleTestCase):
    """Nowy DOCX nie może wnosić formatowania z szablonu python-docx, którego nie było w źródle."""

    def test_template_defaults_do_not_add_spacing_size_or_font(self):
        document = Document()
        drop_doc_defaults(document)
        document.add_paragraph('Akapit bez własnych ustawień.')
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Akapit bez własnych ustawień.')
        styles = rebuilt.styles.element
        # Szablon python-docx ma 10 pt po akapicie, interlinię 1,15, 11 pt i czcionkę motywu.
        for path in ('w:docDefaults//w:spacing', 'w:docDefaults//w:sz', 'w:docDefaults//w:rFonts',
                     'w:style[@w:styleId="Normal"]/w:pPr', 'w:style[@w:styleId="Normal"]/w:rPr'):
            self.assertEqual(styles.xpath(path), [], path)
        # Brak ustawień domyślnych też jest odtworzony: LibreOffice przyjmuje wtedy wartości Worda 2007.
        self.assertEqual(len(styles.xpath('w:docDefaults')), 1)
        self.assertEqual(styles.xpath('w:docDefaults/*'), [])
        run = rebuilt.paragraphs[0].runs[0]
        self.assertIsNone(run.font.size)
        self.assertIsNone(rebuilt.paragraphs[0].paragraph_format.space_after)

    def test_heading_style_keeps_structure_without_template_look(self):
        document = Document()
        heading = document.styles['Heading 1']
        heading.font.color.rgb = None
        heading.paragraph_format.space_before = None
        document.add_heading('Rozdział pierwszy', 1)
        rebuilt = Document(rebuild_docx(saved(document)))
        paragraph = rebuilt.paragraphs[0]
        self.assertEqual(paragraph.style.name, 'Heading 1')
        self.assertIsNone(paragraph.runs[0].font.color.rgb)
        self.assertIsNone(paragraph.style.font.color.rgb)
        self.assertIsNone(paragraph.style.paragraph_format.space_before)
        self.assertIsNone(paragraph.paragraph_format.space_before)
        # Poziom konspektu zostaje: nawigacja i spis treści w Wordzie dalej widzą rozdział.
        self.assertEqual(paragraph.style.element.xpath('w:pPr/w:outlineLvl/@w:val'), ['0'])
        self.assertTrue(paragraph.runs[0].font.bold)
        self.assertTrue(paragraph.paragraph_format.keep_with_next)

    def test_document_and_run_language_are_kept(self):
        document = Document()
        language = document.styles.element.xpath('w:docDefaults//w:lang')[0]
        language.set(qn('w:val'), 'pl-PL')
        paragraph = document.add_paragraph('Zażółć gęślą jaźń. ')
        quote = paragraph.add_run('To be or not to be.')
        quote._r.get_or_add_rPr().append(RebuildDetailTests.element('w:lang', **{'w:val': 'en-GB'}))
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.styles.element.xpath('w:docDefaults//w:lang/@w:val'), ['pl-PL'])
        polish, english = rebuilt.paragraphs[0].runs
        self.assertEqual(polish._r.xpath('w:rPr/w:lang/@w:val'), [])
        self.assertEqual(english._r.xpath('w:rPr/w:lang/@w:val'), ['en-GB'])

    def test_empty_paragraph_keeps_its_line_height(self):
        """Puste akapity (odstępy między scenami) mają wysokość z formatowania znaku akapitu."""
        document = Document()
        document.styles['Normal'].font.size = Pt(14)
        document.styles['Normal'].font.name = 'Georgia'
        document.add_paragraph('Scena pierwsza.')
        document.add_paragraph()
        rebuilt = Document(rebuild_docx(saved(document)))
        mark = rebuilt.paragraphs[1]._p.xpath('w:pPr/w:rPr')
        self.assertEqual(len(mark), 1)
        self.assertEqual(mark[0].xpath('w:sz/@w:val'), ['28'])
        self.assertEqual(mark[0].xpath('w:rFonts/@w:ascii'), ['Georgia'])

    def test_default_tab_stop_is_kept(self):
        document = Document()
        document.settings.element.find(qn('w:defaultTabStop')).set(qn('w:val'), '709')
        document.add_paragraph('A\tB')
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.settings.element.xpath('w:defaultTabStop/@w:val'), ['709'])

    def test_invalid_theme_shade_does_not_stop_the_rebuild(self):
        document = Document()
        run = document.add_paragraph().add_run('Czerwony')
        run._r.get_or_add_rPr().append(RebuildDetailTests.element(
            'w:color', **{'w:val': 'FF0000', 'w:themeShade': 'zz', 'w:themeTint': ''}))
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.paragraphs[0].runs[0].font.color.rgb, RGBColor(0xFF, 0, 0))


class RebuildWarningAccuracyTests(SimpleTestCase):
    """Ostrzeżenia liczą to, co użytkownik widzi w dokumencie: bez podwójnego liczenia i bez fałszywych alarmów."""

    def test_word_text_box_is_one_item_and_surrounding_text_stays(self):
        document = Document()
        add_text_box(document.add_paragraph('Przed polem. '), 'W ramce')
        self.assertEqual(omissions_of(document), ['ilustracje i pola tekstowe wraz z ich treścią (liczba: 1)'])
        rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Przed polem. ')

    def test_alternate_content_without_fallback_keeps_plain_text(self):
        document = Document()
        paragraph = document.add_paragraph('Początek ')
        paragraph._p.append(parse_xml(
            '<mc:AlternateContent %s xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice Requires="w14"><w:r><w:t>środek</w:t></w:r></mc:Choice></mc:AlternateContent>'
            % nsdecls('w')))
        paragraph.add_run(' koniec')
        self.assertEqual(omissions_of(document), [])
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Początek środek koniec')

    def test_one_field_is_one_item_and_its_result_stays(self):
        document = Document()
        add_complex_field(document.add_paragraph('Strona '), 'PAGE', '7')
        self.assertEqual(omissions_of(document),
                         ['mechanizm pól automatycznych (zwykły tekst wyniku może pozostać) (liczba: 1)'])
        rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Strona 7')

    def test_nested_table_and_image_in_text_box_count_once(self):
        document = Document()
        document.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
        self.assertEqual(omissions_of(document), ['tabele wraz z ich treścią (liczba: 1)'])

    def test_unknown_formatting_values_do_not_stop_the_rebuild(self):
        """Wartości spoza słownika python-docx (np. w:jc „start”/„end” z OOXML Strict, LibreOffice, Google Docs)
        nie przerywają przebudowy: start → do lewej, end → do prawej bez ostrzeżenia; nieznane wyrównanie
        i nieznane podkreślenie – pominięte i policzone (także w stylu akapitu)."""
        def jc(value):
            return parse_xml(f'<w:jc {nsdecls("w")} w:val="{value}"/>')

        cases = (
            ('start w akapicie', 'start', None, WD_ALIGN_PARAGRAPH.LEFT, []),
            ('end w akapicie', 'end', None, WD_ALIGN_PARAGRAPH.RIGHT, []),
            ('nieznane w akapicie', 'zzz', None, None, [UNKNOWN_FORMAT_OMISSION + ' (liczba: 1)']),
            ('end w stylu', None, 'end', WD_ALIGN_PARAGRAPH.RIGHT, []),
            ('nieznane w stylu', None, 'zzz', None, [UNKNOWN_FORMAT_OMISSION + ' (liczba: 1)']),
        )
        for case, own, in_style, expected, omissions in cases:
            with self.subTest(case):
                document = Document()
                paragraph = document.add_paragraph('Tekst')
                if own:
                    paragraph._p.get_or_add_pPr().append(jc(own))
                if in_style:
                    document.styles['Normal'].element.get_or_add_pPr().append(jc(in_style))
                self.assertEqual(omissions_of(document), omissions)
                rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
                self.assertEqual(rebuilt.paragraphs[0].text, 'Tekst')
                self.assertEqual(rebuilt.paragraphs[0].paragraph_format.alignment, expected)
        with self.subTest('tabulator o nieznanym wyrównaniu'):
            document = Document()
            paragraph = document.add_paragraph('A\tB')
            paragraph._p.get_or_add_pPr().append(parse_xml(
                f'<w:tabs {nsdecls("w")}><w:tab w:val="zzz" w:pos="2835"/><w:tab w:val="right" w:pos="5670"/></w:tabs>'))
            self.assertEqual(omissions_of(document), [UNKNOWN_FORMAT_OMISSION + ' (liczba: 1)'])
            rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
            self.assertEqual([stop.position for stop in rebuilt.paragraphs[0].paragraph_format.tab_stops], [5670 * 635])
        with self.subTest('nieznane podkreślenie fragmentu'):
            document = Document()
            run = document.add_paragraph().add_run('Tekst')
            run._r.get_or_add_rPr().append(parse_xml(f'<w:u {nsdecls("w")} w:val="zzz"/>'))
            self.assertEqual(omissions_of(document), [UNKNOWN_FORMAT_OMISSION + ' (liczba: 1)'])
            rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
            self.assertEqual(rebuilt.paragraphs[0].runs[0].text, 'Tekst')

    def test_content_free_markers_between_paragraphs_are_not_reported(self):
        document = Document()
        document.add_paragraph('A')
        body = document.element.body
        for xml in ('<w:bookmarkStart %s w:id="0" w:name="_GoBack"/>', '<w:bookmarkEnd %s w:id="0"/>',
                    '<w:proofErr %s w:type="gramEnd"/>', '<w:commentRangeEnd %s w:id="4"/>',
                    '<w:permEnd %s w:id="1"/>'):
            body.insert(1, parse_xml(xml % nsdecls('w')))
        document.add_paragraph('B')
        self.assertEqual(omissions_of(document), [])
        self.assertEqual([p.text for p in Document(rebuild_docx(saved(document))).paragraphs], ['A', 'B'])

    def test_numbering_switched_off_is_not_reported(self):
        document = Document()
        paragraph = document.add_paragraph('Zwykły akapit', 'List Number')
        paragraph._p.get_or_add_pPr().append(parse_xml(
            '<w:numPr %s><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>' % nsdecls('w')))
        self.assertEqual(omissions_of(document), [])
        numbered = Document()
        numbered.add_paragraph('Pozycja', 'List Number')
        self.assertEqual(omissions_of(numbered),
                         ['automatyczne numerowanie i punktory list (tekst pozycji pozostanie) (liczba: 1)'])

    def test_formatting_history_is_kept_when_the_previous_formatting_can_be_rebuilt(self):
        """Zmiana formatowania zostaje (poprzednie formatowanie odtworzone jak obecne: kursywa, wyrównanie, styl
        nagłówka), a gdy poprzednie formatowanie ma to, czego przebudowa nie przenosi (rozstrzelenie, cieniowanie,
        obramowanie, numeracja), zmiana jest przyjmowana z jedną uwagą na zmianę."""
        history = FORMAT_HISTORY_OMISSION + ' (liczba: {})'
        cases = (
            # nazwa, poprzednie w:rPr, poprzednie w:pPr, pominięcia, XPath każdej zachowanej zmiany
            ('kursywa i wyrównanie', '<w:i/>', '<w:jc w:val="right"/>', [],
             ['w:r/w:rPr/w:rPrChange/w:rPr/w:i', 'w:pPr/w:pPrChange/w:pPr/w:jc[@w:val="right"]']),
            ('styl nagłówka', '', '<w:pStyle w:val="Heading1"/>', [],
             ['w:r/w:rPr/w:rPrChange', 'w:pPr/w:pPrChange/w:pPr/w:pStyle[@w:val="Heading1"]']),
            ('rozstrzelenie i obramowanie', '<w:spacing w:val="20"/><w:shd w:fill="FF0000"/>', '<w:pBdr/>',
             [history.format(2)], []),
            ('numeracja', '', '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>', [history.format(1)],
             ['w:r/w:rPr/w:rPrChange']),
            ('numeracja ze stylu', '', '<w:pStyle w:val="ListNumber"/>', [history.format(1)], ['w:r/w:rPr/w:rPrChange']),
        )
        for label, run_before, paragraph_before, omissions, kept in cases:
            with self.subTest(label):
                document = Document()
                paragraph = document.add_paragraph()
                run = paragraph.add_run('Tekst')
                run.bold = True
                run._r.get_or_add_rPr().append(parse_xml(
                    f'<w:rPrChange {nsdecls("w")} w:id="5" w:author="Redaktor"><w:rPr>{run_before}</w:rPr></w:rPrChange>'))
                paragraph._p.get_or_add_pPr().append(parse_xml(
                    f'<w:pPrChange {nsdecls("w")} w:id="6" w:author="Redaktor" w:date="2026-01-01T00:00:00Z">'
                    f'<w:pPr>{paragraph_before}</w:pPr></w:pPrChange>'))
                self.assertEqual(omissions_of(document), omissions)
                rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
                self.assertTrue(rebuilt.paragraphs[0].runs[0].bold)
                changes = rebuilt.paragraphs[0]._p.xpath('.//w:rPrChange | .//w:pPrChange')
                self.assertEqual(len(changes), len(kept))
                for path in kept:
                    self.assertTrue(rebuilt.paragraphs[0]._p.xpath(path), path)
                self.assertEqual(review_markup_errors(review_parts(rebuild_docx(saved(document),
                                                                                allow_omissions=True))), [])

    def test_link_inside_removed_table_is_not_promised_to_stay(self):
        """„Widoczny tekst pozostanie” nie może dotyczyć odnośnika, który znika razem z tabelą."""
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0]._p.append(parse_xml(
            '<w:hyperlink %s w:anchor="x"><w:r><w:t>link</w:t></w:r></w:hyperlink>' % nsdecls('w')))
        self.assertEqual(omissions_of(document), ['tabele wraz z ich treścią (liczba: 1)'])

    def test_header_with_only_a_vml_watermark_is_reported(self):
        document = Document()
        document.add_paragraph('Treść')
        header = document.sections[0].header.paragraphs[0]
        header._p.append(parse_xml(
            '<w:r %s xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape>'
            '<v:textpath string="POUFNE"/></v:shape></w:pict></w:r>' % nsdecls('w')))
        self.assertEqual(omissions_of(document), ['niepuste nagłówki lub stopki (liczba: 1)'])


def note_parts(stream, kind):
    """(tekst treści z odsyłaczami, teksty kolejnych przypisów, korzeń części przypisów albo None, pliki pakietu)."""
    with ZipFile(stream) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    notes = files.get(f'word/{kind}s.xml')
    root = etree.fromstring(notes) if notes is not None else None
    texts = [noted_text(note) for note in root.findall(qn('w:' + kind)) if note.get(qn('w:type')) is None] if notes else []
    return noted_text(etree.fromstring(files['word/document.xml'])), texts, root, files


class RebuildNotesTests(SimpleTestCase):
    """Przypisy dolne i końcowe: odsyłacz na swoim miejscu, treść w nowej części, sanityzowana jak treść główna."""

    def test_notes_are_recreated_in_a_new_part(self):
        """Nowa część z separatorami (w:id -1 i 0), relacją, typem zawartości, ustawieniami i stylami;
        odsyłacz we własnym przebiegu ze stylem odsyłacza, znak przypisu na początku treści, bez XML źródła."""
        for kind, (name, reltype) in NOTE_PARTS.items():
            title = kind.capitalize()
            with self.subTest(kind):
                source = noted_document(['Zdanie', 1, ' dalej.'], kind, note_runs=(
                    '<w:r w:rsidR="00ABCDEF"><w:t xml:space="preserve"> Treść przypisu.</w:t></w:r>'))
                inspect_docx(source)
                result = rebuild_docx(source)
                body, texts, root, files = note_parts(result, kind)
                self.assertEqual(body, 'Zdanie[1] dalej.')
                self.assertEqual(texts, ['[ref] Treść przypisu.'])
                self.assertNotIn(b'00ABCDEF', files[f'word/{name}.xml'])
                self.assertEqual([(note.get(qn('w:type')), note.get(qn('w:id'))) for note in root],
                                 [('separator', '-1'), ('continuationSeparator', '0'), (None, '1')])
                document = Document(result)
                reference = next(document.element.body.iter(qn(f'w:{kind}Reference')))
                self.assertEqual([child.tag for child in reference.getparent()], [qn('w:rPr'), reference.tag])
                self.assertEqual(reference.getparent().style, f'{title}Reference')
                self.assertEqual(root.xpath('./w:*[@w:id="1"]/w:p/w:pPr/w:pStyle/@w:val', namespaces=NS),
                                 [f'{title}Text'])
                styles = document.styles.element
                self.assertEqual(styles.get_by_id(f'{title}Text').type, WD_STYLE_TYPE.PARAGRAPH)
                self.assertEqual(styles.get_by_id(f'{title}Reference').xpath('./w:rPr/w:vertAlign/@w:val'),
                                 ['superscript'])
                self.assertEqual(document.settings.element.xpath(f'./w:{kind}Pr/w:{kind}/@w:id'), ['-1', '0'])
                self.assertIn(f'PartName="/word/{name}.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                              f'wordprocessingml.{name}+xml"'.encode(), files['[Content_Types].xml'])
                self.assertIn(f'Type="{reltype}"'.encode(), files['word/_rels/document.xml.rels'])
                other = 'endnote' if kind == 'footnote' else 'footnote'
                self.assertNotIn(f'word/{other}s.xml', files)

    def test_note_content_follows_the_rules_of_the_main_text(self):
        """Formatowanie, pominięcia (liczone razem z treścią), kolejność i numeracja, brakujący przypis,
        przypis bez odsyłacza i własny znak odsyłacza."""
        drawing = '<w:r><w:t>Tekst</w:t></w:r><w:r><w:drawing/></w:r>'
        table = ('<w:r><w:t>Tekst</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>komórka</w:t></w:r></w:p>'
                 '</w:tc></w:tr></w:tbl><w:p>')
        # nazwa: (przebiegi, opcje noted_document, pominięcia, treść, przypisy, XPath w części przypisów)
        cases = {
            'kursywa': (['A', 1], {'note_runs': '<w:r><w:rPr><w:i/></w:rPr><w:t>kursywa</w:t></w:r>'}, [],
                        'A[1]', ['[ref]kursywa'], './/w:r[w:t="kursywa"]/w:rPr/w:i'),
            'ilustracje w dwóch przypisach': (
                ['A', 1, 'B', 2], {'note_runs': drawing},
                ['ilustracje i pola tekstowe wraz z ich treścią (liczba: 2)'], 'A[1]B[2]', ['[ref]Tekst'] * 2, None),
            'tabela w przypisie': (['A', 1], {'note_runs': table}, ['tabele wraz z ich treścią (liczba: 1)'],
                                   'A[1]', ['[ref]Tekst'], None),
            'kolejność i numeracja': (['A', 3, 'B', 1, 'C', 2], {'note_text': 'Przypis {id}'}, [], 'A[1]B[2]C[3]',
                                      ['[ref]Przypis 3', '[ref]Przypis 1', '[ref]Przypis 2'], None),
            'odsyłacz do brakującego przypisu': (
                ['A', 1, 2], {'notes': [1], 'note_text': 'Przypis {id}'},
                ['odsyłacze do nieistniejących przypisów (liczba: 1)'], 'A[1]', ['[ref]Przypis 1'], None),
            'przypis bez odsyłacza': (['A', 2], {'notes': [1, 2, 3], 'note_text': 'Przypis {id}'}, [], 'A[1]',
                                      ['[ref]Przypis 2'], None),
            'własny znak bez znaku przypisu w treści': (
                ['A', 1, '*'], {'reference': 'w:customMarkFollows="1"', 'note_text': '* Gwiazdka', 'mark': False},
                [], 'A[1]*', ['* Gwiazdka'], './/w:*[@w:id="1"]'),
            'hiperłącze w przypisie': (
                ['A', 1], {'note_runs': '<w:hyperlink><w:r><w:t>link</w:t></w:r></w:hyperlink>'},
                ['odnośniki hiperłączy (widoczny tekst pozostanie) (liczba: 1)'], 'A[1]', ['[ref]link'], None),
            'dwa odsyłacze do tego samego przypisu': (['A', 1, 'B', 1], {'note_text': 'Jeden'}, [], 'A[1]B[2]',
                                                      ['[ref]Jeden', '[ref]Jeden'], None),
            'znak przypisu dopiero po tekście': (
                ['A', 1], {'mark': False, 'note_runs': '<w:r><w:t>x</w:t><w:footnoteRef/></w:r>'},
                [], 'A[1]', ['[ref]x'], None),
        }
        for label, (parts, options, omissions, body, notes, xpath) in cases.items():
            with self.subTest(label):
                source = noted_document(parts, **options)
                try:
                    inspect_docx(source)
                    found = []
                except RebuildUnsupported as error:
                    found = error.omissions
                self.assertEqual(found, omissions)
                result = rebuild_docx(source, allow_omissions=True)
                text, texts, root, files = note_parts(result, 'footnote')
                self.assertEqual((text, texts), (body, notes))
                if xpath:
                    self.assertTrue(root.xpath(xpath, namespaces=NS))
                custom = Document(result).element.body.xpath('.//w:footnoteReference/@w:customMarkFollows')
                self.assertEqual(custom, ['1'] if 'reference' in options else [])

    def test_text_in_the_reference_run_and_rebuilt_separators(self):
        """Tekst w tym samym przebiegu przed odsyłaczem i za nim trafia do osobnych przebiegów; ponowna przebudowa
        wyniku (część z separatorami w:type) odtwarza tylko zwykłe przypisy."""
        source = noted_document(['A', 1], note_text='Przypis')
        output = BytesIO()
        with ZipFile(source) as archive, ZipFile(output, 'w') as target:
            for item in archive.infolist():
                data = archive.read(item.filename)
                if item.filename == 'word/document.xml':
                    root = etree.fromstring(data)
                    reference = root.xpath('.//w:footnoteReference', namespaces=NS)[0]
                    reference.addprevious(parse_xml(f'<w:t {nsdecls("w")}>przed</w:t>'))
                    reference.getparent().append(parse_xml(f'<w:t {nsdecls("w")}>po</w:t>'))
                    data = etree.tostring(root)
                target.writestr(item, data)
        output.seek(0)
        first = rebuild_docx(output)
        self.assertEqual(note_parts(first, 'footnote')[:2], ('Aprzed[1]po', ['[ref]Przypis']))
        runs = Document(first).paragraphs[0].runs
        self.assertEqual([run.text for run in runs], ['A', 'przed', '', 'po'])
        second = rebuild_docx(BytesIO(first.getvalue()))
        _, texts, root, _ = note_parts(second, 'footnote')
        self.assertEqual(texts, ['[ref]Przypis'])
        self.assertEqual(root.xpath('./w:footnote/@w:type', namespaces=NS), ['separator', 'continuationSeparator'])

    def test_notes_survive_the_whole_preparation_with_the_cleaner(self):
        """Przebudowa → ujednolicenie → miękkie entery → Odkurzacz: treść przypisu poprawiona,
        odsyłacz przesunięty przed kropkę (note_references), bez pominięć."""
        for kind in NOTE_PARTS:
            with self.subTest(kind):
                result = prepare_docx(noted_document(['Tekst.', 1], kind), rebuild=True, normalize_formatting=True,
                                      remove_soft_whitespace=True, cleaner_rules=list(DEFAULT_EDITORIAL_RULES))
                body, texts, _, _ = note_parts(result, kind)
                self.assertEqual((body, texts), ('Tekst[1].', ['[ref] Treść przypisu…']))


REVIEW_FIXTURES = {
    'Word 365': review_docx,
    'LibreOffice 24.2': libreoffice_review_docx,
    'Google Docs': google_docs_extras_docx,
}
# Kroki przygotowania bez Odkurzacza: opcje prepare_docx i czy zaczynają od przebudowy (nowe, poprawne w:id).
REVIEW_STEPS = {
    'przebudowa': ({'rebuild': True}, True),
    'przebudowa i ujednolicenie': ({'rebuild': True, 'normalize_formatting': True}, True),
    'przebudowa, ujednolicenie i miękkie entery': (
        {'rebuild': True, 'normalize_formatting': True, 'remove_soft_whitespace': True}, True),
    'ujednolicenie': ({'normalize_formatting': True}, False),
    'miękkie entery': ({'remove_soft_whitespace': True}, False),
}


def review_docx_with(body, comments='', **options):
    """DOCX „jak z Worda 365” z podaną treścią i komentarzami (w:comment) – bez części Worda 365 poza comments.xml."""
    parts, relationships, overrides = _comment_parts(comments, word365=False)
    return word_docx(body, parts={**parts, **options.pop('parts', {})},
                     relationships=[*relationships, *options.pop('relationships', [])],
                     overrides=[*overrides, *options.pop('overrides', [])], **options)


def comment_xml(number, text, author='Recenzentka'):
    return (f'<w:comment w:id="{number}" w:author="{author}" w:date="2026-02-01T10:00:00Z" w:initials="R">'
            f'<w:p>{text}</w:p></w:comment>')


def comment_reference(number):
    return f'<w:r><w:commentReference w:id="{number}"/></w:r>'


def table(*paragraphs):
    return ('<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid><w:tr><w:tc>' + ''.join(paragraphs)
            + '</w:tc></w:tr></w:tbl>')


class RebuildReviewTests(SimpleTestCase):
    """Zmiany śledzone (wstawki, usunięcia, przeniesienia, znaki akapitu, zmiany formatowania) i komentarze (zakresy,
    odpowiedzi, rozwiązanie) zostają po przebudowie i kolejnych krokach przygotowania bez Odkurzacza."""

    def test_review_survives_every_preparation_step(self):
        """Pliki z Worda 365, LibreOffice i Google Docs: te same zmiany (rodzaj, autor, data, tekst), ten sam tekst
        po przyjęciu i po odrzuceniu wszystkich zmian (treść i przypisy), te same zmiany formatowania i komentarze
        (autor, inicjały, data, treść z pogrubieniem i kursywą, zakres na tych samych słowach, komentarz punktowy,
        odpowiedź, rozwiązanie, komentarz w przypisie). Wynik otwiera python-docx (z komentarzami), a zapis jest
        poprawny: po przebudowie bez błędów, bez przebudowy – bez nowych błędów (identyfikatory źródła zostają)."""
        for editor, build in REVIEW_FIXTURES.items():
            source = review_parts(build())
            for step, (options, rebuilt) in REVIEW_STEPS.items():
                with self.subTest(editor, step=step):
                    output = prepare_docx(build(), allow_omissions=True, **options)
                    parts = review_parts(output)
                    for accept in (True, False):
                        self.assertEqual(resolved(parts, accept), resolved(source, accept))
                    self.assertEqual(revisions(parts), revisions(source))
                    self.assertEqual(format_changes(parts), format_changes(source))
                    self.assertEqual(comments_of(parts), comments_of(source))
                    self.assertEqual(len(Document(output).comments), len(comments_of(source)))
                    if rebuilt:
                        self.assertEqual(review_markup_errors(parts), [])
                    else:
                        self.assertLessEqual(set(review_markup_errors(parts)), set(review_markup_errors(source)))

    def test_word_review_fixture_has_the_expected_texts(self):
        """Pomocniki przyjęcia i odrzucenia zmian na pliku z Worda: oczekiwane teksty, zmiany i wątki komentarzy."""
        parts = review_parts(rebuild_docx(review_docx()))
        self.assertEqual(resolved(parts, accept=True), [[
            'Początek komentarza, wstawiony i  oraz  tekst.', 'Dalszy ciąg pogrubiony i koniec komentarza.',
            'Punkt bez zakresu i przypis.', 'Akapit z nowym znakiem końca', 'i dalszy ciąg, połączony po przyjęciu.',
            'Przeniesiony akapit.', 'Rozwiązany komentarz.', 'Koniec.'], [' Przypis z uwagą i wstawką.']])
        self.assertEqual(resolved(parts, accept=False), [[
            'Początek komentarza,  i usunięty oraz  tekst.', 'Dalszy ciąg pogrubiony i koniec komentarza.',
            'Punkt bez zakresu i przypis.', 'Przeniesiony akapit.', 'Akapit z nowym znakiem końcai dalszy ciąg, ',
            'połączony po przyjęciu.', 'Rozwiązany komentarz.', 'Cały usunięty akapit.', 'Koniec.'],
            [' Przypis z uwagą.']])
        self.assertEqual([change[:2] for change in revisions(parts)[:4]],
                         [('ins', 'Recenzentka'), ('del', 'Recenzentka'), ('ins', 'Autorka'), ('del', 'Recenzentka')])
        threads = {comment[3]: comment[-2:] for comment in comments_of(parts)}
        self.assertEqual(threads['Odpowiedź.'], ('Pytanie?', False))
        self.assertEqual(threads['Rozwiązane.'], (None, True))
        self.assertEqual(comment_anchors(parts)['1'], (None, 'Punkt'))

    def test_review_markup_is_not_reported_as_lost(self):
        """Przegląd nie zgłasza już zmian i komentarzy (zostają); Word 365 i LibreOffice – bez żadnych uwag."""
        for editor, build in REVIEW_FIXTURES.items():
            with self.subTest(editor):
                try:
                    inspect_docx(build())
                    omissions = []
                except RebuildUnsupported as error:
                    omissions = error.omissions
                if editor != 'Google Docs':
                    self.assertEqual(omissions, [])
                self.assertFalse([item for item in omissions if 'komentarz' in item or 'zmian' in item], omissions)

    def test_word_365_extension_parts_are_not_carried(self):
        """commentsIds, commentsExtensible i people (identyfikatory kont autorów) nie są potrzebne do odpowiedzi ani
        rozwiązania – nowy dokument ich nie ma, a nowe części komentarzy mają w14:paraId z mc:Ignorable."""
        with ZipFile(rebuild_docx(review_docx())) as archive:
            names = archive.namelist()
            payload = b''.join(archive.read(name) for name in names)
            comments = archive.read('word/comments.xml')
        self.assertIn('word/commentsExtended.xml', names)
        for name in ('commentsIds', 'commentsExtensible', 'people'):
            self.assertNotIn(f'word/{name}.xml', names)
        self.assertNotIn(b'wydawnictwo.test', payload)
        self.assertNotIn(b'dateUtc', payload)
        self.assertIn(b'mc:Ignorable="w14"', comments)

    def test_comments_and_changes_next_to_omitted_content(self):
        """Komentarz w pomijanej tabeli – uwaga; odwołanie do brakującego komentarza i znacznik bez pary znikają bez
        uwagi; koniec zakresu w tabeli – koniec przed odwołaniem; początek w tabeli – komentarz punktowy; odnośnik
        w komentarzu – tekst zostaje z uwagą; zmiana śledzona w tabeli znika z tabelą (bez osobnej uwagi); przypis
        z komentarzem wskazany dwoma odsyłaczami (treść odtwarzana dwa razy) – znaczniki komentarza tylko raz."""
        tables = 'tabele wraz z ich treścią (liczba: 1)'
        lost = LOST_COMMENTS_OMISSION + ' (liczba: 1)'
        link = ('<w:hyperlink r:id="rIdLink">' + r('strona') + '</w:hyperlink>')
        cases = (
            # nazwa, treść, komentarze, opcje, pominięcia, oczekiwane komentarze (treść, zakres, słowo przed odwołaniem)
            ('komentarz w tabeli', p(r('Tekst.')) + table(p('<w:commentRangeStart w:id="0"/>', r('Komórka'),
                                                            '<w:commentRangeEnd w:id="0"/>', comment_reference(0))),
             comment_xml(0, r('W tabeli.')), {}, [lost, tables], []),
            ('brakujący komentarz', p('<w:commentRangeStart w:id="7"/>', r('Tekst'), '<w:commentRangeEnd w:id="7"/>',
                                      comment_reference(7)), comment_xml(0, r('Inny.')), {}, [], []),
            ('koniec zakresu w tabeli', p('<w:commentRangeStart w:id="0"/>', r('Początek'))
             + table(p(r('Komórka'), '<w:commentRangeEnd w:id="0"/>')) + p(r('Dalej'), comment_reference(0)),
             comment_xml(0, r('Uwaga.')), {}, [tables], [('Uwaga.', 'Początek¶Dalej', 'Dalej')]),
            ('początek zakresu w tabeli', table(p('<w:commentRangeStart w:id="0"/>', r('Komórka')))
             + p(r('Dalej'), '<w:commentRangeEnd w:id="0"/>', comment_reference(0)),
             comment_xml(0, r('Uwaga.')), {}, [tables], [('Uwaga.', None, 'Dalej')]),
            ('odnośnik w komentarzu', p('<w:commentRangeStart w:id="0"/>', r('Tekst'), '<w:commentRangeEnd w:id="0"/>',
                                        comment_reference(0)), comment_xml(0, r('Zobacz ') + link),
             {'parts': {'word/_rels/comments.xml.rels': _rels_with_link()}}, [HYPERLINK_OMISSION + ' (liczba: 1)'],
             [('Zobacz strona', 'Tekst', 'Tekst')]),
            ('przeniesienie bez pary', '<w:moveFromRangeStart w:id="1" w:author="A" w:name="m"/>'
             + p('<w:moveFrom w:id="2" w:author="A">' + r('Stąd.') + '</w:moveFrom>')
             + '<w:moveFromRangeEnd w:id="1"/>' + table(p('<w:moveToRangeStart w:id="3" w:author="A" w:name="m"/>',
                                                          '<w:moveTo w:id="4" w:author="A">' + r('Stąd.') + '</w:moveTo>',
                                                          '<w:moveToRangeEnd w:id="3"/>')), '', {}, [tables], []),
            ('usunięcie w tabeli', p(r('Tekst.')) + table(p('<w:del w:id="1" w:author="A"><w:r><w:delText>usunięte'
                                                            '</w:delText></w:r></w:del>')), '', {}, [tables], []),
            ('przypis z komentarzem i dwoma odsyłaczami',
             p(r('A'), note_reference('footnote', 1), r(' B'), note_reference('footnote', 1)), comment_xml(0, r('K.')),
             {'parts': {'word/footnotes.xml': notes_part('footnote', {1: '<w:commentRangeStart w:id="0"/>' + r('uwaga')
                                                                     + '<w:commentRangeEnd w:id="0"/>'
                                                                     + comment_reference(0)})},
              'relationships': [('rIdFoot', NOTE_PARTS['footnote'][1], 'footnotes.xml')],
              'overrides': [('/word/footnotes.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.'
                                                    'footnotes+xml')]},
             [], [('K.', 'uwaga', 'uwaga')]),
        )
        for label, body, comments, options, omissions, expected in cases:
            with self.subTest(label):
                try:
                    inspect_docx(review_docx_with(body, comments, **dict(options)))
                    found = []
                except RebuildUnsupported as error:
                    found = error.omissions
                self.assertEqual(found, omissions)
                parts = review_parts(rebuild_docx(review_docx_with(body, comments, **dict(options)),
                                                  allow_omissions=True))
                self.assertEqual([comment[3:4] + comment[5:7] for comment in comments_of(parts)], expected)
                self.assertEqual(review_markup_errors(parts), [])
                self.assertEqual(parts['comments'] is None, not expected)


def words_of(texts):
    """Słowa (małymi literami) kolejnych akapitów kolejnych tekstów – bez odstępów i interpunkcji."""
    return [[re.findall(r'\w+', text.casefold()) for text in story] for story in texts]


def cleaned_review(build, track_changes):
    """Pełne przygotowanie: przebudowa, ujednolicenie, miękkie entery i Odkurzacz (reguły domyślne)."""
    return prepare_docx(build(), rebuild=True, normalize_formatting=True, remove_soft_whitespace=True,
                        use_cleaner=True, cleaner_rules=list(DEFAULT_EDITORIAL_RULES), track_changes=track_changes,
                        allow_omissions=True)


class ReviewCleanerChainTests(SimpleTestCase):
    """Pełny łańcuch prepare_docx z Odkurzaczem (ze śledzeniem zmian i bez) na plikach z recenzją."""

    def test_full_preparation_keeps_the_review(self):
        """Zmiany i komentarze innych osób zostają (rodzaj, autor, data, tekst; treść, autor, wątek, rozwiązanie),
        słowa tekstu po przyjęciu zmian są te same co w źródle (Odkurzacz zmienia odstępy i interpunkcję), a tekst
        po przyjęciu jest taki sam ze śledzeniem i bez; ze śledzeniem odrzucenie wszystkich zmian (także Odkurzacza)
        daje dokładnie tekst źródła. Zapis poprawny, python-docx otwiera wynik z komentarzami."""
        for editor, build in REVIEW_FIXTURES.items():
            source = review_parts(build())
            accepted = {}
            for track_changes in (False, True):
                with self.subTest(editor, track_changes=track_changes):
                    output = cleaned_review(build, track_changes)
                    parts = review_parts(output)
                    accepted[track_changes] = resolved(parts, accept=True)
                    self.assertEqual(words_of(accepted[track_changes]), words_of(resolved(source, accept=True)))
                    if track_changes:
                        self.assertEqual(resolved(parts, accept=False), resolved(source, accept=False))
                    self.assertEqual([change for change in revisions(parts) if change[1] != 'Odkurzacz'],
                                     revisions(source))
                    self.assertEqual([comment[:5] + comment[7:] for comment in comments_of(parts)],
                                     [comment[:5] + comment[7:] for comment in comments_of(source)])
                    self.assertEqual(review_markup_errors(parts), [])
                    self.assertEqual(len(Document(output).comments), len(comments_of(source)))
            with self.subTest(editor, compare='ze śledzeniem i bez'):
                self.assertEqual(accepted[True], accepted[False])

    def test_rejecting_changes_after_cleaning_without_tracking_keeps_the_words(self):
        """Bez śledzenia zmian Odkurzacz nie zbija spacji wokół cudzego usunięcia („ i ” + usunięty + „ oraz ”),
        więc po odrzuceniu zmian słowa się nie sklejają („usunięty oraz”, nie „usuniętyoraz”)."""
        for editor, build in REVIEW_FIXTURES.items():
            with self.subTest(editor):
                parts = review_parts(cleaned_review(build, track_changes=False))
                self.assertEqual(words_of(resolved(parts, accept=False)),
                                 words_of(resolved(review_parts(build()), accept=False)))


def _rels_with_link():
    """Relacje części komentarzy z odnośnikiem zewnętrznym rIdLink."""
    return _rels([('rIdLink', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                   'https://example.test/', 'External')])


class RebuildConversionTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_real_worker_docx_and_mail_default(self):
        doc = Document()
        doc.add_paragraph('Tekst')
        doc.core_properties.author = 'Secret'
        payload = saved(doc).getvalue()
        out, extension, mime = convert_document(
            SimpleUploadedFile('a.docx', payload), [], include_docx=True, rebuild=True, normalize=False)
        with out:
            self.assertEqual(extension, 'docx')
            self.assertIn('wordprocessingml', mime)
            self.assertEqual(Document(out).core_properties.author, '')
        row = {'uid': 1, 'folder': 'Story', 'files': [('story.docx', payload)]}
        for rebuild in (True, False):
            with package_messages([row], clean=False, convert=False, rebuild=rebuild) as archive, \
                    ZipFile(archive) as z:
                content = z.read(z.namelist()[0])
                if rebuild:
                    self.assertEqual(Document(BytesIO(content)).core_properties.author, '')
                else:
                    self.assertEqual(content, payload)

    def test_original_invalid_docx_is_downloadable_from_mail(self):
        data = b'not a valid docx'
        rows = [{'uid': 1, 'folder': 'Text', 'files': [('text.docx', data)]}]
        with package_messages(rows, clean=False, convert=False, rebuild=False) as out, ZipFile(out) as archive:
            self.assertEqual(archive.read(archive.namelist()[0]), data)

    def test_confirmation_still_builds_new_docx(self):
        doc = Document()
        doc.core_properties.author = 'Private'
        doc.add_paragraph('Zachowaj').runs[0].italic = True
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Pomiń tabelę'
        doc.add_paragraph('Lista', 'List Number')
        payload = saved(doc).getvalue()
        with self.assertRaises(RebuildConfirmationRequired) as caught:
            convert_document(BytesIO(payload), [], include_docx=True, rebuild=True)
        self.assertIn('tabele wraz z ich treścią', str(caught.exception))
        self.assertIn('tekst pozycji pozostanie', str(caught.exception))
        out, ext, _ = convert_document(
            BytesIO(payload), [], include_docx=True, rebuild=True, allow_rebuild_omissions=True)
        with out:
            result = Document(out)
            self.assertEqual(ext, 'docx')
            self.assertEqual(result.core_properties.author, '')
            self.assertEqual(len(result.tables), 0)
            self.assertEqual([p.text for p in result.paragraphs], ['Zachowaj', 'Lista'])
            self.assertTrue(result.paragraphs[0].runs[0].italic)

    def test_mail_batch_collects_warnings_for_all_files(self):
        rows = []
        for uid, kind in enumerate(('table', 'list'), 1):
            doc = Document()
            doc.add_paragraph('Tekst')
            if kind == 'table':
                doc.add_table(rows=1, cols=1)
            else:
                doc.add_paragraph('Pozycja', 'List Number')
            rows.append({'uid': uid, 'folder': kind, 'files': [(kind + '.docx', saved(doc).getvalue())]})
        with self.assertRaises(RebuildConfirmationRequired) as caught:
            package_messages(rows, clean=False, convert=False)
        self.assertIn('table.docx', str(caught.exception))
        self.assertIn('list.docx', str(caught.exception))
        with package_messages(rows, clean=False, convert=False, allow_rebuild_omissions=True) as out, \
                ZipFile(out) as archive:
            self.assertEqual(len(archive.namelist()), 2)


class RebuildPageTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = create_superuser('rebuild')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)

    def test_rebuild_checkbox_is_off_by_default_and_combines_with_rules(self):
        page = self.client.get(reverse('core:programs'))
        self.assertFalse(page.context['form'].fields['rebuild'].initial)
        self.assertNotContains(page, 'value="rebuild"')
        for enabled in (False, True):
            doc = Document()
            doc.add_paragraph('Zdanie. następne zdanie...')
            doc.core_properties.author = 'Private'
            data = {
                'program_action': 'clean',
                'rules': ['sentence_case'],
                'document': SimpleUploadedFile('a.docx', saved(doc).getvalue()),
            }
            if enabled:
                data['rebuild'] = 'on'
            response = run_program(self.client, data)
            self.assertEqual(response.status_code, 200)
            result = downloaded_document(response)
            self.assertEqual(result.paragraphs[0].text, 'Zdanie. Następne zdanie...')
            self.assertEqual(result.core_properties.author, '' if enabled else 'Private')

    @tag('slow')
    def test_button_download_and_unsupported_error(self):
        for table in (False, True):
            doc = Document()
            doc.add_paragraph('Tekst')
            doc.core_properties.author = 'Private'
            if table:
                doc.add_table(rows=1, cols=1)
            response = run_program(self.client, {
                'program_action': 'clean',
                'rebuild': 'on',
                'document': SimpleUploadedFile('story.docx', saved(doc).getvalue()),
            })
            self.assertEqual(response.status_code, 200)
            if not table:
                self.assertIn('_nowy.docx', response['Content-Disposition'])
                result = downloaded_document(response)
                self.assertEqual(result.core_properties.author, '')
                self.assertEqual(result.paragraphs[0].text, 'Tekst')
                continue
            # Zadanie czeka na potwierdzenie; oryginał zostaje na serwerze tylko do tej chwili.
            self.assertContains(response, 'W nowym DOCX zostaną pominięte')
            self.assertContains(response, 'value="confirm"')
            token = response.context['token']
            with patch.object(jobs, 'launch', side_effect=jobs.execute):
                accepted = self.client.post(reverse('core:program_job', args=[token]),
                                            {'action': 'confirm', 'view': 'page'})
            self.assertEqual(accepted.status_code, 302)
            page = self.client.get(accepted['Location'])
            self.assertEqual(page.context['state']['state'], 'done')
            rebuilt = downloaded_document(self.client.get(page.context['download_url']))
            self.assertEqual(len(rebuilt.tables), 0)
            self.assertEqual(rebuilt.core_properties.author, '')
            self.assertFalse(list(Path(self.work_dir).glob('convert/program_jobs/*/source.docx')))


FIELD_WARNING = 'mechanizm pól automatycznych (zwykły tekst wyniku może pozostać) (liczba: 1)'
TOC_WARNING = 'spis treści (wstaw go ponownie w Wordzie: Odwołania → Spis treści) (liczba: 1)'


def placeholder_run(text, style='PlaceholderText'):
    """Tekst zastępczy kontrolki tak, jak zapisuje go Word (styl znakowy „Tekst zastępczy”)."""
    return f'<w:r><w:rPr><w:rStyle w:val="{style}"/></w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def texts(document, **options):
    return [p.text for p in Document(rebuild_docx(saved(document), **options)).paragraphs]


class RebuildContentControlTests(SimpleTestCase):
    """Kontrolki zawartości znikają, ale ich tekst i formatowanie zostają; tekst zastępczy i spis treści nie."""

    def test_inline_control_keeps_text_and_formatting_without_warning(self):
        document = Document()
        paragraph = document.add_paragraph('Imię: ')
        paragraph._p.append(content_control('<w:r><w:rPr><w:b/><w:i/></w:rPr><w:t>Anna</w:t></w:r>'))
        paragraph.add_run(' i dalej.')
        self.assertEqual(omissions_of(document), [])
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Imię: Anna i dalej.')
        name = rebuilt.paragraphs[0].runs[1]
        self.assertEqual(name.text, 'Anna')
        self.assertTrue(name.bold)
        self.assertTrue(name.italic)
        self.assertEqual(rebuilt.element.body.xpath('.//w:sdt'), [])

    def test_block_control_keeps_paragraphs_in_place(self):
        document = Document()
        document.add_paragraph('Przed')
        add_block(document, content_control(
            '<w:p><w:pPr><w:jc w:val="center"/></w:pPr>' + text_run('Pierwszy') + '</w:p>'
            '<w:p>' + text_run('Drugi') + '</w:p>'))
        document.add_paragraph('Po')
        self.assertEqual(omissions_of(document), [])
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual([p.text for p in rebuilt.paragraphs], ['Przed', 'Pierwszy', 'Drugi', 'Po'])
        self.assertEqual(rebuilt.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_nested_controls_are_all_unwrapped(self):
        document = Document()
        outer = add_block(document, content_control('<w:p>' + text_run('Zewnątrz, ') + '</w:p>'))
        outer.find(qn('w:sdtContent')).find(qn('w:p')).append(content_control(text_run('środek')))
        self.assertEqual(omissions_of(document), [])
        self.assertEqual(texts(document), ['Zewnątrz, środek'])

    def test_placeholder_text_is_not_carried(self):
        document = Document()
        paragraph = document.add_paragraph('Data: ')
        paragraph._p.append(content_control(placeholder_run('Kliknij tutaj, aby wprowadzić datę.'), placeholder=True))
        add_block(document, content_control(
            '<w:p>' + placeholder_run('Kliknij lub naciśnij tutaj, aby wprowadzić tekst.') + '</w:p>',
            placeholder=True))
        document.add_paragraph('Koniec')
        self.assertEqual(omissions_of(document), [])
        self.assertEqual(texts(document), ['Data: ', 'Koniec'])

    def test_polish_placeholder_style_is_recognized_by_its_name(self):
        document = Document()
        style = document.styles.add_style('Tekstzastpczy', WD_STYLE_TYPE.CHARACTER)
        style.element.find(qn('w:name')).set(qn('w:val'), 'Placeholder Text')
        style_id = style.element.get(qn('w:styleId'))
        paragraph = document.add_paragraph('Data: ')
        paragraph._p.append(content_control(placeholder_run('Kliknij tutaj.', style_id), placeholder=True))
        self.assertEqual(omissions_of(document), [])
        self.assertEqual(texts(document), ['Data: '])

    def test_placeholder_flag_with_real_content_keeps_the_text_with_a_warning(self):
        """Generator dokumentów wpisał treść, a oznaczenie w:showingPlcHdr zostało – treść nie może zniknąć."""
        document = Document()
        paragraph = document.add_paragraph('Imię: ')
        paragraph._p.append(content_control(text_run('Jan'), placeholder=True))
        add_block(document, content_control('<w:p>' + text_run('TEKST WYPEŁNIONY') + '</w:p>', placeholder=True))
        self.assertEqual(omissions_of(document), [PLACEHOLDER_OMISSION + ' (liczba: 2)'])
        self.assertEqual(texts(document, allow_omissions=True), ['Imię: Jan', 'TEKST WYPEŁNIONY'])

    def test_placeholder_group_with_filled_fields_is_unwrapped(self):
        document = Document()
        from lxml import etree
        inner = etree.tostring(content_control(text_run('KOWALSKI')), encoding=str)
        add_block(document, content_control('<w:p>' + text_run('Nazwisko: ') + inner + '</w:p>', placeholder=True))
        self.assertEqual(omissions_of(document), [PLACEHOLDER_OMISSION + ' (liczba: 1)'])
        self.assertEqual(texts(document, allow_omissions=True), ['Nazwisko: KOWALSKI'])

    def test_table_of_contents_is_removed_with_one_warning(self):
        document = Document()
        entry = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                 '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" </w:instrText></w:r>'
                 '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                 '<w:hyperlink w:anchor="_Toc1"><w:r><w:t>Rozdział 1</w:t></w:r></w:hyperlink>'
                 '<w:fldSimple w:instr="PAGEREF _Toc1"><w:r><w:t>5</w:t></w:r></w:fldSimple>'
                 '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
        add_block(document, content_control(
            '<w:p>' + text_run('Spis treści') + '</w:p>' + entry, gallery='Table of Contents'))
        document.add_heading('Rozdział 1', 1)
        # Pola i odnośniki spisu nie są liczone osobno: znikają razem z nim.
        self.assertEqual(omissions_of(document), [TOC_WARNING])
        self.assertEqual(texts(document, allow_omissions=True), ['Rozdział 1'])

    def test_other_gallery_blocks_are_unwrapped(self):
        document = Document()
        add_block(document, content_control('<w:p>' + text_run('Tytuł książki') + '</w:p>', gallery='Cover Pages'))
        self.assertEqual(omissions_of(document), [])
        self.assertEqual(texts(document), ['Tytuł książki'])

    def test_controls_in_table_rows_and_cells_leave_only_the_table_warning(self):
        document = Document()
        document.add_paragraph('Tekst')
        table = document.add_table(rows=1, cols=1)
        row = table.rows[0]._tr
        cell = row.tc_lst[0]
        cell.append(content_control('<w:p>' + text_run('w komórce') + '</w:p>'))
        table._tbl.append(content_control(
            '<w:tr><w:tc><w:p>' + text_run('wiersz') + '</w:p></w:tc></w:tr>'))
        row.append(content_control('<w:tc><w:p>' + text_run('komórka') + '</w:p></w:tc>'))
        self.assertEqual(omissions_of(document), ['tabele wraz z ich treścią (liczba: 1)'])
        self.assertEqual(texts(document, allow_omissions=True), ['Tekst'])

    def test_section_break_inside_block_control_is_reported(self):
        document = Document()
        add_block(document, content_control(
            '<w:p><w:pPr><w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:pPr>'
            + text_run('Część pierwsza') + '</w:p>'))
        document.add_paragraph('Część druga')
        self.assertEqual(omissions_of(document),
                         ['podziały na sekcje i ustawienia ich układu (pozostanie jedna sekcja A4) (liczba: 1)'])


class RebuildSimpleFieldTests(SimpleTestCase):
    """Pole proste (w:fldSimple) zachowuje się jak złożone: wynik zostaje zwykłym tekstem."""

    def test_simple_field_result_stays_as_text(self):
        document = Document()
        paragraph = document.add_paragraph('Autor: ')
        paragraph._p.append(parse_xml(
            '<w:fldSimple %s w:instr=" AUTHOR "><w:r><w:rPr><w:i/></w:rPr><w:t>Jan</w:t></w:r>'
            '<w:fldData>AAAA</w:fldData></w:fldSimple>' % nsdecls('w')))
        self.assertEqual(omissions_of(document), [FIELD_WARNING])
        rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Autor: Jan')
        self.assertTrue(rebuilt.paragraphs[0].runs[1].italic)
        self.assertEqual(rebuilt.element.body.xpath('.//w:fldSimple | .//w:fldData'), [])

    def test_nested_simple_fields_count_once_and_fields_in_tables_are_not_counted(self):
        document = Document()
        document.add_paragraph()._p.append(parse_xml(
            '<w:fldSimple %s w:instr="IF"><w:fldSimple w:instr="PAGE"><w:r><w:t>3</w:t></w:r></w:fldSimple>'
            '</w:fldSimple>' % nsdecls('w')))
        self.assertEqual(omissions_of(document), [FIELD_WARNING])
        self.assertEqual(texts(document, allow_omissions=True), ['3'])
        tabled = Document()
        tabled.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]._p.append(parse_xml(
            '<w:fldSimple %s w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>' % nsdecls('w')))
        self.assertEqual(omissions_of(tabled), ['tabele wraz z ich treścią (liczba: 1)'])


def book_document():
    """Książka A5 z lustrzanymi marginesami, odległością nagłówka i dwiema kolumnami o różnych szerokościach."""
    document = Document()
    document.add_paragraph('Treść')
    section = document.sections[0]
    section.page_width, section.page_height = Cm(14.8), Cm(21)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2.2)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.5)
    section.header_distance, section.footer_distance = Cm(0.9), Cm(1.1)
    section.gutter = Cm(0.5)
    columns = section._sectPr.find(qn('w:cols'))
    columns.set(qn('w:num'), '2')
    columns.set(qn('w:equalWidth'), '0')
    columns.set(qn('w:sep'), '1')
    columns.append(parse_xml('<w:col %s w:w="3000" w:space="400"/>' % nsdecls('w')))
    columns.append(parse_xml('<w:col %s w:w="2500"/>' % nsdecls('w')))
    document.settings.element.insert(1, parse_xml('<w:mirrorMargins %s/>' % nsdecls('w')))
    return document


class RebuildKeepPageFormatTests(SimpleTestCase):
    """keep_page_format=True: wynik ma format strony, marginesy i kolumny pierwszej sekcji źródła."""

    def test_book_layout_is_kept_without_warning(self):
        document = book_document()
        inspect_docx(saved(document), keep_page_format=True)
        self.assertEqual(omissions_of(document),
                         ['format strony, orientacja lub kolumny (wynik: A4 pionowo, jedna kolumna) (liczba: 1)'])
        rebuilt = Document(rebuild_docx(saved(document), keep_page_format=True))
        section = rebuilt.sections[0]
        for name, expected in (('page_width', 14.8), ('page_height', 21), ('top_margin', 2),
                               ('bottom_margin', 2.2), ('left_margin', 1.8), ('right_margin', 1.5),
                               ('header_distance', 0.9), ('footer_distance', 1.1), ('gutter', 0.5)):
            self.assertAlmostEqual(getattr(section, name).cm, expected, places=2, msg=name)
        columns = section._sectPr.find(qn('w:cols'))
        self.assertEqual((columns.get(qn('w:num')), columns.get(qn('w:equalWidth')), columns.get(qn('w:sep'))),
                         ('2', '0', '1'))
        self.assertEqual([dict(c.attrib) for c in columns.findall(qn('w:col'))],
                         [{qn('w:w'): '3000', qn('w:space'): '400'}, {qn('w:w'): '2500'}])
        self.assertEqual(len(rebuilt.settings.element.xpath('w:mirrorMargins')), 1)
        self.assertEqual(rebuilt.paragraphs[0].text, 'Treść')

    def test_default_rebuild_still_uses_a4(self):
        rebuilt = Document(rebuild_docx(saved(book_document()), allow_omissions=True))
        section = rebuilt.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 21, places=1)
        self.assertAlmostEqual(section.left_margin.cm, 2.5, places=2)
        self.assertIsNone(section._sectPr.find(qn('w:cols')).get(qn('w:num')))
        self.assertEqual(rebuilt.settings.element.xpath('w:mirrorMargins'), [])

    def test_landscape_orientation_is_kept(self):
        document = Document()
        document.add_paragraph('Treść')
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        inspect_docx(saved(document), keep_page_format=True)
        rebuilt = Document(rebuild_docx(saved(document), keep_page_format=True)).sections[0]
        self.assertEqual(rebuilt.orientation, WD_ORIENT.LANDSCAPE)
        self.assertAlmostEqual(rebuilt.page_width.cm, 29.7, places=1)

    def test_missing_source_values_fall_back_to_a4(self):
        document = Document()
        document.add_paragraph('Treść')
        sect_pr = document.sections[0]._sectPr
        for tag_name in ('w:pgSz', 'w:pgMar', 'w:cols'):
            sect_pr.remove(sect_pr.find(qn(tag_name)))
        rebuilt = Document(rebuild_docx(saved(document), keep_page_format=True)).sections[0]
        self.assertAlmostEqual(rebuilt.page_width.cm, 21, places=1)
        self.assertAlmostEqual(rebuilt.top_margin.cm, 2.5, places=2)

    def test_mirror_margins_switched_off_are_not_copied(self):
        document = Document()
        document.add_paragraph('Treść')
        document.settings.element.insert(1, parse_xml('<w:mirrorMargins %s w:val="0"/>' % nsdecls('w')))
        rebuilt = Document(rebuild_docx(saved(document), keep_page_format=True))
        self.assertEqual(rebuilt.settings.element.xpath('w:mirrorMargins'), [])

    def test_several_sections_keep_the_first_layout(self):
        document = book_document()
        document.add_section()
        document.sections[1].page_width = Cm(21)
        with self.assertRaises(RebuildUnsupported) as caught:
            inspect_docx(saved(document), keep_page_format=True)
        self.assertEqual(caught.exception.omissions, [
            'podziały na sekcje i ustawienia ich układu (pozostanie jedna sekcja o układzie pierwszej sekcji) '
            '(liczba: 1)'])
        rebuilt = Document(rebuild_docx(saved(document), allow_omissions=True, keep_page_format=True))
        self.assertEqual(len(rebuilt.sections), 1)
        self.assertAlmostEqual(rebuilt.sections[0].page_width.cm, 14.8, places=1)
