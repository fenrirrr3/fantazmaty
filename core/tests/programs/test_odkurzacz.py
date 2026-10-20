"""Odkurzacz: czyszczenie DOCX we wszystkich częściach dokumentu oraz strona programu z pobraniem wyniku."""
from io import BytesIO
from zipfile import ZipFile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from lxml import etree, html

from core.program_test_support import run_program
from core.services.document_preparation import prepare_docx
from core.services.odkurzacz import (
    ALL_EDITORIAL_RULES, DEFAULT_EDITORIAL_RULES, EDITORIAL_RULES, FORM_DEFAULT_EDITORIAL_RULES, clean_docx,
    correct_editorial_text,
)
from fantazmaty.testing import create_member

from .support import (
    NOTE_PARTS, IsolatedProgramDirsMixin, add_internal_link, add_page_break, add_text_box, docx_bytes, downloaded,
    downloaded_document, footnote_text, formatted_docx, noted_document, noted_text, saved, set_default_language,
    set_language, tracked_changes_document, upload, whole_document, with_footnote,
)


def soft_whitespace_document():
    """Twarde spacje, miękki enter i podział strony w akapicie, tabeli i nagłówku."""
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run('Ala\u00a0ma')
    run.bold = True
    run.add_break()
    run.add_text('kota\u202f!')
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Drugi akapit')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela\u00a0A\nB'
    doc.sections[0].header.paragraphs[0].text = 'Nagłówek\u00a0A\nB'
    return saved(doc).getvalue()


class CleanDocxTests(SimpleTestCase):
    def test_correction_across_runs_also_covers_tables_and_headers(self):
        document = Document()
        para = document.add_paragraph()
        para.add_run('Ala ').bold = True
        para.add_run(' ma kota...').underline = True
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Bez  zmian...'
        document.sections[0].header.paragraphs[0].text = 'Nagłówek  bez zmian...'
        result = Document(clean_docx(saved(document), ['spaces', 'ellipsis']))
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota…')
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].underline)
        self.assertEqual(result.tables[0].cell(0, 0).text, 'Bez zmian…')
        self.assertEqual(result.sections[0].header.paragraphs[0].text, 'Nagłówek bez zmian…')
        self.assertFalse(result.paragraphs[0]._p.xpath('.//w:color | .//w:highlight'))

    def test_default_rules_cover_tables_headers_and_footnotes(self):
        result = clean_docx(whole_document(), sorted(DEFAULT_EDITORIAL_RULES))
        document = Document(BytesIO(result.getvalue()))
        self.assertEqual(document.paragraphs[0].text, 'Akapit główny.')
        self.assertEqual(document.tables[0].cell(0, 0).text, 'Lampa świeciła, lampa gasła.')
        self.assertEqual(document.sections[0].header.paragraphs[0].text, 'Nagłówek strony')
        self.assertIn('Przypis z lampą i lampą.', footnote_text(BytesIO(result.getvalue())))

    def test_all_rules_run(self):
        result = Document(clean_docx(BytesIO(docx_bytes('  Ala  ma kota...  ')), [key for key, _ in EDITORIAL_RULES]))
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota…')

    def test_oversize_text_is_rejected(self):
        with self.assertRaises(ValueError):
            clean_docx(BytesIO(docx_bytes('a' * 20001)), ['spaces'])

    def test_tabs_are_replaced_with_single_spaces(self):
        text = '\tAla\tma \t kota\t\t'
        result = Document(clean_docx(BytesIO(docx_bytes(text)), ['tabs']))
        self.assertEqual(result.paragraphs[0].text, ' Ala ma kota ')
        self.assertFalse(result.paragraphs[0]._p.xpath('.//w:tab'))
        unchanged = Document(clean_docx(BytesIO(docx_bytes(text)), []))
        self.assertEqual(unchanged.paragraphs[0].text, text)

    def test_dialogue_spaces_survive_internal_docx_markers(self):
        """Regresja: znaczniki proofErr/bookmark między przebiegami gubiły spację przed półpauzą."""
        cases = (
            ('viraptorka', ' – pochwalił.'),
            ('rozumiemy.', ' – W jej głos'),
            ('pszczół.', ' – Roześmiała się'),
            ('viraptorka – ', 'pochwalił.'),
        )
        rules = ['hyphen_dash', 'dash_spaces', 'trim', 'before_punct', 'after_punct']
        for marker in ('proofErr', 'bookmarkStart', 'bookmarkEnd'):
            for left, right in cases:
                with self.subTest(marker=marker, left=left):
                    document = Document()
                    paragraph = document.add_paragraph()
                    paragraph.add_run(left).bold = True
                    paragraph._p.append(OxmlElement('w:' + marker))
                    paragraph.add_run(right).italic = True
                    result = Document(clean_docx(saved(document), rules))
                    self.assertEqual(result.paragraphs[0].text, left + right)
                    self.assertTrue(result.paragraphs[0].runs[0].bold)
                    self.assertTrue(result.paragraphs[0].runs[-1].italic)

    def test_soft_whitespace_keeps_paragraphs_page_breaks_formatting_and_cleans_tables(self):
        with prepare_docx(BytesIO(soft_whitespace_document()), use_cleaner=True, cleaner_rules=['spaces'],
                          remove_soft_whitespace=True) as stream:
            doc = Document(stream)
        self.assertEqual([p.text for p in doc.paragraphs[:2]], ['Ala ma', 'kota !'])
        self.assertEqual(len(doc.paragraphs), 3)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertEqual(len(doc.element.xpath('//w:br[@w:type="page"]')), 1)
        self.assertEqual(len(doc.element.xpath('//w:br[not(@w:type) or @w:type="textWrapping"]')), 0)
        self.assertEqual([p.text for p in doc.tables[0].cell(0, 0).paragraphs], ['Tabela A', 'B'])
        self.assertEqual([p.text for p in doc.sections[0].header.paragraphs], ['Nagłówek A', 'B'])


class CleanDocxMarkerAndOrderTests(SimpleTestCase):
    """Znaczniki Worda bez treści nie dzielą zdania, a Odkurzacz działa po usunięciu miękkich spacji."""

    def test_zero_width_word_markers_do_not_split_the_sentence(self):
        """Regresja: proofErr/zakładka/lastRenderedPageBreak przerywały korektę („domu  .Potem”)."""
        for marker, inside_run in (('proofErr', False), ('bookmarkStart', False), ('commentRangeStart', False),
                                   ('lastRenderedPageBreak', True)):
            with self.subTest(marker=marker):
                document = Document()
                paragraph = document.add_paragraph()
                paragraph.add_run('Wrócił do domu ').bold = True
                element = OxmlElement('w:' + marker)
                if inside_run:
                    run = paragraph.add_run()
                    run._r.append(element)
                    run.add_text(' .potem')
                else:
                    paragraph._p.append(element)
                    paragraph.add_run(' .potem')
                result = Document(clean_docx(saved(document), FORM_DEFAULT_EDITORIAL_RULES))
                self.assertEqual(result.paragraphs[0].text, 'Wrócił do domu. Potem')
                self.assertTrue(result.paragraphs[0].runs[0].bold)
                self.assertEqual(len(result.paragraphs[0]._p.xpath('.//w:' + marker)), 1)

    def test_cleaner_after_soft_whitespace_removal_uses_only_plain_spaces(self):
        """Decyzja właściciela: także po zamianie miękkich enterów nigdzie nie ma spacji nierozdzielających."""
        document = Document()
        run = document.add_paragraph().add_run('Był w\u00a0domu i ogrodzie,')
        run.add_break()
        run.add_text('miał\u00a012345 zł.')
        with prepare_docx(saved(document), use_cleaner=True, cleaner_rules=list(FORM_DEFAULT_EDITORIAL_RULES),
                          remove_soft_whitespace=True) as stream:
            result = Document(stream)
        self.assertEqual([p.text for p in result.paragraphs],
                         ['Był w domu i ogrodzie,', 'miał 12 345 zł.'])


class CleanDocxStructureTests(SimpleTestCase):
    """Odkurzacz nie zmienia pól Worda ani pustych akapitów, które niosą układ dokumentu."""

    def test_field_code_and_result_are_left_alone(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Przed  polem ')
        for kind, text in (('begin', None), (None, ' HYPERLINK  "x" '), ('separate', None), (None, 'Wynik  pola'),
                           ('end', None)):
            run = paragraph.add_run()
            if kind:
                marker = OxmlElement('w:fldChar')
                marker.set(qn('w:fldCharType'), kind)
                run._r.append(marker)
            else:
                code = OxmlElement('w:instrText' if 'HYPERLINK' in text else 'w:t')
                code.text = text
                code.set(qn('xml:space'), 'preserve')
                run._r.append(code)
        paragraph.add_run(' i  po nim')
        result = Document(clean_docx(saved(document), ['spaces']))
        self.assertEqual(result.paragraphs[0].text, 'Przed polem Wynik  pola i po nim')
        self.assertEqual(result.paragraphs[0]._p.xpath('.//w:instrText/text()'), [' HYPERLINK  "x" '])

    def test_empty_paragraphs_with_layout_are_kept(self):
        document = Document()
        document.add_paragraph('A')
        document.add_paragraph(' \u00a0').add_run().add_tab()
        document.add_paragraph()                                            # usuwany: drugi pusty z rzędu
        document.add_paragraph().paragraph_format.page_break_before = True  # podział strony
        document.add_paragraph('', style='List Number')                      # numeracja
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)          # treść inna niż tekst
        bookmarked = document.add_paragraph()
        bookmarked._p.append(OxmlElement('w:bookmarkStart'))
        document.add_paragraph()
        document.add_table(rows=1, cols=1)
        document.add_paragraph()
        document.add_paragraph('B')
        result = Document(clean_docx(saved(document), ['empty_paragraphs']))
        self.assertEqual(len(result.paragraphs), len(document.paragraphs) - 1)
        self.assertEqual(result.paragraphs[2].paragraph_format.page_break_before, True)

    def test_unknown_rule_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'Nieznana reguła'):
            clean_docx(BytesIO(docx_bytes()), ['nieznana'])


W_DEL, W_INS = qn('w:del'), qn('w:ins')


def _resolve_revisions(root, accept):
    """Programowe „Zaakceptuj wszystkie” / „Odrzuć wszystkie” zmiany w drzewie XML jednej części."""
    remove, unwrap = (W_DEL, W_INS) if accept else (W_INS, W_DEL)
    for paragraph in list(root.iter(qn('w:p'))):
        mark = paragraph.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{W_DEL}")
        if mark is None:
            continue
        mark.getparent().remove(mark)
        following = paragraph.getnext()
        if accept and following is not None and following.tag == qn('w:p'):
            # Usunięty znak akapitu: treść łączy się z następnym akapitem.
            position = 1 if following.find(qn('w:pPr')) is not None else 0
            for child in [child for child in paragraph if child.tag != qn('w:pPr')]:
                following.insert(position, child)
                position += 1
            paragraph.getparent().remove(paragraph)
    for element in list(root.iter(remove)):
        element.getparent().remove(element)
    for element in list(root.iter(unwrap)):
        for child in list(element):
            element.addprevious(child)
        element.getparent().remove(element)
    for node in root.iter(qn('w:delText')):
        node.tag = qn('w:t')


def _story_texts(stream, resolve=None):
    """Teksty akapitów wszystkich części (treść, tabele, pola tekstowe, nagłówki, stopki, przypisy)."""
    from core.services.document_stories import document_stories
    from docx.text.paragraph import Paragraph
    stream.seek(0)
    stories, _ = document_stories(Document(stream))
    texts = []
    for story in stories:
        if resolve is not None:
            _resolve_revisions(story.root, accept=resolve == 'accept')
        texts.append([Paragraph(element, None).text for element in story.root.iter(qn('w:p'))])
    return texts


def _revisions(stream):
    stream.seek(0)
    with ZipFile(stream) as archive:
        parts = {name: archive.read(name) for name in archive.namelist() if name.endswith('.xml')}
    from lxml import etree
    found = []
    for name, data in parts.items():
        root = etree.fromstring(data)
        found.extend((name, element) for element in root.iter(W_DEL, W_INS))
    return found


class TrackChangesTests(SimpleTestCase):
    """track_changes=True: poprawki jako śledzone zmiany Worda we wszystkich częściach dokumentu."""

    rules = sorted(ALL_EDITORIAL_RULES)

    def test_accept_all_gives_the_untracked_result_and_reject_all_the_original(self):
        source = tracked_changes_document().getvalue()
        untracked = clean_docx(BytesIO(source), self.rules)
        tracked = clean_docx(BytesIO(source), self.rules, track_changes=True)
        Document(BytesIO(tracked.getvalue()))  # wynik otwiera się ponownie
        expected = _story_texts(untracked, 'accept')
        self.assertEqual(_story_texts(BytesIO(tracked.getvalue()), 'accept'), expected)
        self.assertEqual(_story_texts(BytesIO(tracked.getvalue()), 'reject'),
                         _story_texts(BytesIO(source), 'reject'))
        # Zmieniły się wszystkie części: treść, tabela, pole tekstowe, nagłówek, stopka, przypis.
        changed_parts = {name for name, element in _revisions(BytesIO(tracked.getvalue()))
                         if element.get(qn('w:author')) == 'Odkurzacz'}
        self.assertTrue({'word/document.xml', 'word/footnotes.xml'} <= changed_parts)
        self.assertTrue(any('header' in name for name in changed_parts))
        self.assertTrue(any('footer' in name for name in changed_parts))
        self.assertIn('Tabela z kropkami…', expected[0])
        self.assertIn('Pole tekstowe…', expected[0])

    def test_revision_ids_are_unique_and_attributes_complete(self):
        tracked = clean_docx(tracked_changes_document(), self.rules, track_changes=True)
        revisions = _revisions(tracked)
        ours = [element for _, element in revisions if element.get(qn('w:author')) == 'Odkurzacz']
        self.assertGreater(len(ours), 10)
        ids = [element.get(qn('w:id')) for _, element in revisions]
        self.assertEqual(len(ids), len(set(ids)))
        # Także względem innych identyfikatorów (zakładki, przypisy, istniejąca wstawka w:id=9).
        self.assertTrue(all(int(element.get(qn('w:id'))) > 9 for element in ours))
        for element in ours:
            self.assertRegex(element.get(qn('w:date')), r'^\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ$')
        self.assertTrue(any(element.get(qn('w:author')) == 'Redakcja' for _, element in revisions))
        with ZipFile(tracked) as archive:
            self.assertNotIn('word/comments.xml', archive.namelist())

    def test_changes_keep_run_formatting_and_are_word_granular(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Ala  ma ').bold = True
        paragraph.add_run('kota... Potem pickupami.').italic = True
        result = Document(clean_docx(saved(document), ['spaces', 'ellipsis', 'user_word_corrections'],
                                     track_changes=True))
        p = result.paragraphs[0]._p
        deleted = [(node.text, node.getparent().find(qn('w:rPr'))) for node in p.iter(qn('w:delText'))]
        inserted = [(node.text, node.getparent().find(qn('w:rPr'))) for node in p.xpath('./w:ins//w:t')]
        self.assertEqual([text for text, _ in deleted], [' ', '...', 'pickupami'])
        self.assertEqual([text for text, _ in inserted], ['…', 'pikapami'])
        self.assertIsNotNone(deleted[0][1].find(qn('w:b')))
        self.assertTrue(all(props.find(qn('w:i')) is not None for _, props in deleted[1:] + inserted))
        # Niezmieniony tekst zostaje w zwykłych przebiegach z formatowaniem źródła.
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota Potem .')
        self.assertTrue(result.paragraphs[0].runs[0].bold)

    def test_single_character_change_marks_only_that_word(self):
        result = Document(clean_docx(BytesIO(docx_bytes('Długi akapit, w którym Ona coś robi.')), ['pronouns_lower'],
                                     track_changes=True))
        p = result.paragraphs[0]._p
        self.assertEqual(p.xpath('.//w:delText/text()'), ['Ona'])
        self.assertEqual(p.xpath('./w:ins//w:t/text()'), ['ona'])

    def test_tabs_and_text_in_one_run_are_tracked(self):
        result = clean_docx(BytesIO(docx_bytes('\tAla\t \tma')), ['tabs', 'trim'], track_changes=True)
        self.assertEqual(_story_texts(result, 'accept')[0], ['Ala ma'])
        self.assertEqual(_story_texts(result, 'reject')[0], ['\tAla\t \tma'])
        self.assertTrue(Document(BytesIO(result.getvalue())).paragraphs[0]._p.xpath('./w:del/w:r/w:tab'))

    def test_empty_paragraph_merge_is_a_tracked_paragraph_deletion(self):
        document = Document()
        for text in ('A', '', '  ', '', 'B', '', 'C'):
            document.add_paragraph(text)
        source = saved(document).getvalue()
        rules = ['empty_paragraphs', 'trim']
        untracked = clean_docx(BytesIO(source), rules)
        tracked = clean_docx(BytesIO(source), rules, track_changes=True)
        self.assertEqual(_story_texts(tracked, 'accept'), _story_texts(untracked))
        self.assertEqual(_story_texts(tracked, 'accept')[0], ['A', '', 'B', '', 'C'])
        self.assertEqual(_story_texts(tracked, 'reject')[0], ['A', '', '  ', '', 'B', '', 'C'])
        paragraphs = Document(BytesIO(tracked.getvalue())).paragraphs
        self.assertEqual(len(paragraphs), 7)
        self.assertEqual([bool(p._p.xpath('./w:pPr/w:rPr/w:del')) for p in paragraphs],
                         [False, True, True, False, False, False, False])

    def test_without_changes_the_document_has_no_revisions(self):
        tracked = clean_docx(BytesIO(docx_bytes('Tekst bez zmian.')), self.rules, track_changes=True)
        self.assertEqual(_revisions(tracked), [])


class CleanerSummaryTests(SimpleTestCase):
    def test_summary_counts_changes_per_rule_without_changing_the_result(self):
        source = tracked_changes_document().getvalue()
        rules = sorted(ALL_EDITORIAL_RULES)
        summary = {'stare': 1}
        counted = clean_docx(BytesIO(source), rules, summary=summary)
        plain = clean_docx(BytesIO(source), rules)
        self.assertEqual(_story_texts(counted), _story_texts(plain))
        # Słownik jest uzupełniany (inne kroki mogą dopisać swoje klucze).
        self.assertEqual(summary.pop('stare'), 1)
        self.assertEqual(summary['empty_paragraphs'], 3)
        self.assertEqual(summary['user_word_corrections'], 1)
        self.assertEqual(summary['dates_times'], 1)
        self.assertEqual(summary['ellipsis'], 5)  # akapit, tabela, pole tekstowe (kopia zapasowa liczona raz), stopka, przypis
        self.assertEqual(summary.pop('replacements'), [{'source': 'pickupami', 'target': 'pikapami', 'count': 1}])
        self.assertTrue(all(value > 0 for value in summary.values()))
        self.assertLessEqual(set(summary), set(rules))

    def test_counts_for_simple_text(self):
        summary = {}
        clean_docx(BytesIO(docx_bytes('Ala  ma  kota... i psa.')), ['spaces', 'ellipsis', 'quotes'],
                   summary=summary)
        self.assertEqual(summary, {'spaces': 2, 'ellipsis': 1})
        summary = {}
        clean_docx(BytesIO(docx_bytes('Bez zmian.')), ['spaces'], summary=summary)
        self.assertEqual(summary, {})

    def test_summary_also_counts_tracked_changes(self):
        summary = {}
        clean_docx(BytesIO(docx_bytes('Dnia 1.2.2025 o 9:5 wrócił.')), ['dates_times'], track_changes=True,
                   summary=summary)
        self.assertEqual(summary, {'dates_times': 1})


class LanguageTests(SimpleTestCase):
    """Decyzja właściciela: cały dokument jest polski – oznaczenia innego języka niczego nie wyłączają."""

    rules = ['spaces', 'ellipsis', 'user_word_corrections', 'pronouns_lower']

    def test_text_marked_with_another_language_is_cleaned_too(self):
        document = Document()
        set_default_language(document, 'pl-PL')
        set_language(document.styles['Quote'].element, 'en-GB')
        paragraph = document.add_paragraph()
        paragraph.add_run('Mówił  wtedy... ')
        set_language(paragraph.add_run('Pickup  truck...'), 'en-US')
        paragraph.add_run(' i  pickupem...')
        document.add_paragraph('To  be...', style='Quote')
        header = document.sections[0].header.paragraphs[0]
        set_language(header.add_run('Running  head...'), 'en-US')
        for track_changes in (False, True):
            with self.subTest(track_changes=track_changes):
                stream = clean_docx(saved(document), self.rules, track_changes=track_changes)
                texts = _story_texts(stream, 'accept')
                self.assertEqual(texts[0], ['Mówił wtedy… Pikap truck… i pikapem…', 'To be…'])
                self.assertIn('Running head…', [text for story in texts[1:] for text in story])


class CleanerR7Tests(SimpleTestCase):
    """Poprawki po przeglądzie r6: spacje nierozdzielające, podsumowanie zamian, śledzone zmiany, podział strony."""

    @staticmethod
    def nbsp_document():
        document = Document()
        paragraph = document.add_paragraph('Był\u00a0w\u202fdomu\u2007i ')
        add_internal_link(paragraph, 'x', 'w\u00a0łączu')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela\u00a0z'
        add_text_box(document.add_paragraph('Przed.'), 'Pole\u00a0z')
        document.sections[0].header.paragraphs[0].text = 'Nagłówek\u00a0a'
        document.sections[0].footer.paragraphs[0].text = 'Stopka\u202fo'
        return with_footnote(document, 'Przypis\u00a0w')

    def test_existing_nonbreaking_spaces_become_plain_spaces_in_every_part(self):
        source = self.nbsp_document().getvalue()
        for rules, track_changes in ((['spaces'], False), (['spaces'], True), ([], False),
                                     (sorted(ALL_EDITORIAL_RULES), True)):
            with self.subTest(rules=rules[:1], track_changes=track_changes):
                summary = {}
                result = clean_docx(BytesIO(source), rules, track_changes=track_changes, summary=summary)
                with ZipFile(result) as archive:
                    xml = ''.join(archive.read(name).decode() for name in archive.namelist() if name.endswith('.xml'))
                for character in '\u00a0\u202f\u2007':
                    self.assertNotIn(character, xml)
                texts = [text for story in _story_texts(result, 'accept') for text in story]
                for expected in ('Był w domu i w łączu', 'Tabela z', 'Pole z', 'Nagłówek a', 'Stopka o', 'Przypis w'):
                    self.assertIn(expected, texts)
                # Zamiana jest cicha (bez tysięcy niewidocznych zmian) i liczona raz, bez kopii zapasowej pola.
                self.assertFalse(_revisions(result))
                self.assertEqual(summary.get('nonbreaking_spaces'), 9)

    def test_special_spaces_dashes_and_hyphens_in_docx(self):
        """Zadanie 15: w wyniku nie zostaje żaden znak specjalny; spacje zamieniane cicho i liczone
        w nonbreaking_spaces (także wąska, półfiretowa itd.), kreski jako zmiany dash_style."""
        spaces = '\u00a0\u202f\u2009\u2002\u2003\u2008\u200a\u205f'
        document = Document()
        document.add_paragraph('\u2014 Mróz\u2009\u22125\u202f°C \u2015 rzekł, lata 1990\u20121995.')
        document.add_paragraph('Biało\u2010czerwony i\u2002PRL\u2011u, 5\u00a0kg,\u2003J.\u2008Kowalski\u200a'
                               'i\u205f2025 r., odku\u00adrza\u200bcz.')
        source = saved(document).getvalue()
        expected = ['– Mróz –5 °C – rzekł, lata 1990–1995.',
                    'Biało-czerwony i PRL-u, 5 kg, J. Kowalski i 2025 r., odkurzacz.']
        for track_changes in (False, True):
            with self.subTest(track_changes=track_changes):
                summary = {}
                result = clean_docx(BytesIO(source), FORM_DEFAULT_EDITORIAL_RULES, track_changes=track_changes,
                                    summary=summary)
                self.assertEqual(_story_texts(result, 'accept')[0], expected)
                self.assertEqual(summary['nonbreaking_spaces'], 8)
                self.assertEqual(summary['dash_style'], 6)
                self.assertEqual(summary['invisible_chars'], 2)
                revised = ''.join(node.text or '' for _, element in _revisions(result)
                                  for node in element.iter(qn('w:delText'), qn('w:t')))
                # Śledzone zmiany obejmują kreski i niewidoczne znaki, ale nie spacje (te są zamieniane cicho).
                self.assertEqual(bool(revised), track_changes)
                self.assertFalse(set(revised) & set(spaces))
        # Bez reguł spacje i tak stają się zwykłe (decyzja właściciela); kreski zostają.
        summary = {}
        texts = _story_texts(clean_docx(BytesIO(source), [], summary=summary))[0]
        self.assertEqual(texts[0], '\u2014 Mróz \u22125 °C \u2015 rzekł, lata 1990\u20121995.')
        self.assertEqual(summary, {'nonbreaking_spaces': 8})

    def test_summary_lists_dictionary_replacements(self):
        document = Document()
        document.add_paragraph('Pickupem jechał, Pickupem wrócił, pickupem też. Niechcąco.')
        add_text_box(document.add_paragraph('Pole'), 'Pickupem w polu.')
        summary = {}
        clean_docx(saved(document), ['user_word_corrections'], summary=summary)
        self.assertEqual(summary['replacements'], [
            {'source': 'Pickupem', 'target': 'Pikapem', 'count': 3},
            {'source': 'Niechcąco', 'target': 'Niechcący', 'count': 1},
            {'source': 'pickupem', 'target': 'pikapem', 'count': 1},
        ])
        self.assertEqual(summary['user_word_corrections'], 5)
        summary = {}
        clean_docx(saved(document), ['spaces'], summary=summary)
        self.assertNotIn('replacements', summary)

    def test_paragraph_mark_deleted_earlier_is_not_deleted_twice(self):
        document = Document()
        for text in ('A', '', '', 'B'):
            document.add_paragraph(text)
        mark = document.paragraphs[1]._p.get_or_add_pPr()
        properties = OxmlElement('w:rPr')
        properties.append(OxmlElement('w:del', {qn('w:id'): '3', qn('w:author'): 'Redakcja',
                                                qn('w:date'): '2026-01-01T00:00:00Z'}))
        mark.append(properties)
        result = clean_docx(saved(document), ['empty_paragraphs'], track_changes=True)
        paragraphs = Document(result).paragraphs
        self.assertEqual([len(p._p.xpath('./w:pPr/w:rPr/w:del')) for p in paragraphs], [0, 1, 0, 0])

    def test_split_runs_get_unique_formatting_revision_ids(self):
        document = Document()
        run = document.add_paragraph().add_run('Ala  ma  kota...')
        change = OxmlElement('w:rPrChange', {qn('w:id'): '5', qn('w:author'): 'Redakcja',
                                              qn('w:date'): '2026-01-01T00:00:00Z'})
        change.append(OxmlElement('w:rPr'))
        run._r.get_or_add_rPr().append(change)
        result = clean_docx(saved(document), ['spaces', 'ellipsis'], track_changes=True)
        with ZipFile(result) as archive:
            root = etree.fromstring(archive.read('word/document.xml'))
        ids = [node.get(qn('w:id')) for node in root.iter() if node.get(qn('w:id')) is not None]
        self.assertGreater(len(root.findall('.//' + qn('w:rPrChange'))), 1)
        self.assertEqual(len(ids), len(set(ids)))

    def test_text_after_a_page_break_at_paragraph_start_is_trimmed_like_a_paragraph_start(self):
        for track_changes in (False, True):
            with self.subTest(track_changes=track_changes):
                document = Document()
                paragraph = document.add_paragraph()
                paragraph.add_run().add_break(WD_BREAK.PAGE)
                paragraph.add_run('  - Wiem  tak. ')
                add_page_break(document.add_paragraph('Koniec. '))
                result = clean_docx(saved(document), FORM_DEFAULT_EDITORIAL_RULES, track_changes=track_changes)
                texts = _story_texts(result, 'accept')[0]
                self.assertEqual([text.strip('\n') for text in texts], ['– Wiem tak.', 'Koniec.'])
                self.assertEqual(len(Document(result).element.xpath('//w:br[@w:type="page"]')), 2)


class OdkurzaczPageTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('member', 'Redaktor')
        cls.outsider = get_user_model().objects.create_user('outsider', 'outsider@example.com')

    def setUp(self):
        super().setUp()
        self.url = reverse('core:programs')
        self.client.force_login(self.member)

    def test_page_has_all_rules_and_submit_above_them(self):
        response = self.client.get(self.url)
        self.assertContains(response, 'Odkurzacz')
        self.assertContains(response, 'name="rules"', count=len(EDITORIAL_RULES))
        checked = sum(' checked' in str(checkbox.tag()) for checkbox in response.context['form']['rules'])
        self.assertEqual(checked, len(FORM_DEFAULT_EDITORIAL_RULES))
        self.assertNotContains(response, 'gifrific')
        page = response.content.decode()
        self.assertLess(page.index('Odkurz i pobierz DOCX'), page.index('<legend>Opcje korekty'))
        self.assertIn('odkurzacz-upload', page)

    def test_access_and_csrf(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.assertEqual(self.client.post(self.url, {'document': upload()}).status_code, 302)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.post(self.url, {'document': upload()}).status_code, 403)
        secure_client = Client(enforce_csrf_checks=True)
        secure_client.force_login(self.member)
        self.assertEqual(secure_client.post(self.url, {'document': upload()}).status_code, 403)

    def test_invalid_uploads_and_rules(self):
        for data in ({}, {'document': upload(b'bad')}, {'document': upload(name='tekst.txt')},
                     {'document': upload(), 'rules': ['unknown']}):
            with self.subTest(keys=list(data)):
                response = self.client.post(self.url, data)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['form'].errors)
                self.assertFalse(response.streaming)

    def test_download_uses_selected_rules(self):
        response = run_program(self.client, {'document': upload(), 'rules': ['spaces']})
        self.assertEqual(response.status_code, 200)
        self.assertIn('tekst_odkurzony.docx', response['Content-Disposition'])
        self.assertIn('private', response['Cache-Control'])
        self.assertIn('no-store', response['Cache-Control'])
        self.assertEqual(downloaded_document(response).paragraphs[0].text, 'Ala ma kota...')

    def test_no_rules_preserves_text_and_formatting(self):
        document = Document()
        run = document.add_paragraph().add_run('Ala  ma kota...')
        run.bold = True
        run.underline = True
        run.font.color.rgb = RGBColor(12, 34, 56)
        response = run_program(self.client, {'document': upload(saved(document).getvalue())})
        run = downloaded_document(response).paragraphs[0].runs[0]
        self.assertEqual(run.text, 'Ala  ma kota...')
        self.assertTrue(run.bold)
        self.assertTrue(run.underline)
        self.assertEqual(run.font.color.rgb, RGBColor(12, 34, 56))

    def test_without_formatting_option_font_and_alignment_stay(self):
        response = run_program(self.client, {'document': upload(formatted_docx()), 'rules': ['spaces']})
        doc = downloaded_document(response)
        self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 17)
        self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_formatting_option_applies_mailbox_typography_with_and_without_rebuild(self):
        for rebuild in (False, True):
            with self.subTest(rebuild=rebuild):
                data = {'document': upload(formatted_docx()), 'normalize_formatting': 'on', 'rules': ['spaces']}
                if rebuild:
                    data['rebuild'] = 'on'
                response = run_program(self.client, data)
                self.assertEqual(response.status_code, 200)
                first, center = downloaded_document(response).paragraphs
                self.assertEqual(first.text, 'Ala ma kota.')
                self.assertEqual(first.runs[0].font.name, 'Times New Roman')
                self.assertEqual(first.runs[0].font.size.pt, 12)
                self.assertTrue(first.runs[0].italic)
                self.assertEqual(first.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
                self.assertEqual(first.paragraph_format.line_spacing, 1.5)
                self.assertEqual(center.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                self.assertEqual(center.paragraph_format.first_line_indent, 0)

    def test_failure_is_reported_on_the_job_page(self):
        with patch('core.services.program_jobs.convert_document', side_effect=ValueError('bad package')):
            response = run_program(self.client, {'document': upload(), 'rules': ['spaces']})
        self.assertContains(response, 'Nie udało się przetworzyć dokumentu')
        self.assertContains(response, 'Wróć do programów')
        self.assertNotIn('bad package', response.content.decode())

    @tag('slow')
    def test_soft_whitespace_checkbox_reaches_cleaner_and_conversion_worker(self):
        response = run_program(self.client, {
            'document': SimpleUploadedFile('tekst.docx', soft_whitespace_document()),
            'remove_soft_whitespace': 'on',
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('\u00a0', downloaded_document(response).paragraphs[0].text)
        response = run_program(self.client, {
            'program_action': 'convert',
            'convert-document': SimpleUploadedFile('tekst.docx', soft_whitespace_document()),
            'convert-formats': ['epub'],
            'convert-use_cleaner': 'on',
        })
        self.assertEqual(response.status_code, 200)
        with ZipFile(BytesIO(downloaded(response))) as archive:
            content = archive.read('EPUB/content.xhtml')
        self.assertEqual(
            [p.text_content().strip() for p in html.fromstring(content).xpath('//p')][:2],
            ['Ala ma', 'kota!'],  # Odkurzacz usuwa też spację przed wykrzyknikiem.
        )


class CleanerR9Tests(SimpleTestCase):
    """Runda 9: korekta widzi tekst w cudzych wstawkach, hiperłączach i kontrolkach; usunięty tekst
    nie przerywa zdania; miękki enter to brzeg linii."""

    W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

    @staticmethod
    def text_run(text):
        return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'

    def document(self, inner):
        from docx.oxml import parse_xml
        document = Document()
        body = document.element.body
        body.insert(len(body) - 1, parse_xml(f'<w:p {self.W}>{inner}</w:p>'))
        return document

    @staticmethod
    def visible_text(document):
        parts = []
        for node in document.element.body.iter():
            if node.tag == qn('w:t'):
                parts.append(node.text or '')
            elif node.tag in (qn('w:br'), qn('w:cr')):
                parts.append('\n')
        return ''.join(parts)

    def test_insertions_links_controls_and_deleted_text_are_part_of_the_sentence(self):
        revision = 'w:id="9" w:author="X" w:date="2020-01-01T00:00:00Z"'
        cases = {
            'insertion': (self.text_run('Powiedział ,,tak') + f'<w:ins {revision}>{self.text_run(" i nie")}</w:ins>'
                          + self.text_run('" .potem wyszedł'), 'Powiedział „tak i nie”. Potem wyszedł'),
            'whole paragraph inserted': (f'<w:ins {revision}>{self.text_run("Cały akapit ,tak .")}</w:ins>',
                                         'Cały akapit, tak.'),
            'deleted text': (self.text_run('Powiedział "tak') + f'<w:del {revision}><w:r><w:delText> stare</w:delText>'
                             '</w:r></w:del>' + self.text_run('" i poszedł .'), 'Powiedział „tak” i poszedł.'),
            'hyperlink': (self.text_run('Zobacz "') + f'<w:hyperlink r:id="rId99">{self.text_run("stronę")}</w:hyperlink>'
                          + self.text_run('" ,tak.'), 'Zobacz „stronę”, tak.'),
            'content control': (self.text_run('Przed ') + f'<w:sdt><w:sdtPr><w:id w:val="5"/></w:sdtPr><w:sdtContent>'
                                f'{self.text_run("w kontrolce ,tak")}</w:sdtContent></w:sdt>' + self.text_run(' po .koniec'),
                                'Przed w kontrolce, tak po. Koniec'),
            'soft break': (self.text_run('Pierwsza linia ') + '<w:r><w:br/></w:r>' + self.text_run('  - Druga linia  '),
                           'Pierwsza linia\n– Druga linia'),
        }
        for name, (inner, expected) in cases.items():
            with self.subTest(name):
                document = self.document(inner)
                result = Document(clean_docx(saved(document), FORM_DEFAULT_EDITORIAL_RULES))
                self.assertEqual(self.visible_text(result), expected)
                # Cudze zmiany i hiperłącza zostają w dokumencie.
                for tag in ('w:ins', 'w:del', 'w:hyperlink', 'w:sdt'):
                    self.assertEqual(len(list(result.element.body.iter(qn(tag)))),
                                     len(list(document.element.body.iter(qn(tag)))))

    def test_tracked_mode_inside_a_foreign_insertion(self):
        revision = 'w:id="9" w:author="X" w:date="2020-01-01T00:00:00Z"'
        document = self.document(self.text_run('Tekst ') + f'<w:ins {revision}>{self.text_run("z ,,cytatem"  "\"")}</w:ins>')
        result = Document(clean_docx(saved(document), FORM_DEFAULT_EDITORIAL_RULES, track_changes=True))
        insertion = next(result.element.body.iter(qn('w:ins')))
        self.assertTrue(any(node.tag == qn('w:del') for node in insertion.iter()))
        self.assertEqual(self.visible_text(result).replace('cytatem"', ''), 'Tekst z „cytatem”')


def _noted_texts(stream, part='word/document.xml'):
    """Tekst części po przyjęciu zmian; odsyłacz jako [w:id], odsyłacz wewnątrz przypisu jako [ref]."""
    with ZipFile(BytesIO(stream.getvalue())) as archive:
        root = etree.fromstring(archive.read(part))
    _resolve_revisions(root, accept=True)
    return noted_text(root)


class NoteReferenceTests(SimpleTestCase):
    """Odsyłacze przypisów: reguły widzą zdanie ciągłe, a reguła note_references stawia odsyłacz wg norm."""

    on = FORM_DEFAULT_EDITORIAL_RULES
    off = FORM_DEFAULT_EDITORIAL_RULES - {'note_references'}

    # nazwa: (przebiegi, wynik z regułą, wynik bez reguły, rodzaj przypisu)
    CASES = {
        'cudzysłów przez odsyłacz': (['Powiedział "cytat', 1, '" i wyszedł'],
                                     'Powiedział „cytat”[1] i wyszedł', 'Powiedział „cytat[1]” i wyszedł'),
        'wielka litera po odsyłaczu': (['To zdanie.', 1, ' następne zdanie.'],
                                       'To zdanie[1]. Następne zdanie.', 'To zdanie.[1] Następne zdanie.'),
        'spacja przed interpunkcją za odsyłaczem': (['Słowo', 1, ' , dalej.'], 'Słowo[1], dalej.', 'Słowo[1], dalej.'),
        'spacja przed odsyłaczem': (['Słowo ', 1, ' dalej.'], 'Słowo[1] dalej.', 'Słowo[1] dalej.'),
        'kropka': (['Tekst.', 1], 'Tekst[1].', 'Tekst.[1]'),
        'przecinek': (['Tekst,', 1, ' dalej.'], 'Tekst[1], dalej.', 'Tekst,[1] dalej.'),
        'kropka skrótu': (['I tak dalej itd.', 1], 'I tak dalej itd.[1]', 'I tak dalej itd.[1]'),
        'cudzysłów zamykający': (['„Cytat', 1, '”'], '„Cytat”[1]', '„Cytat[1]”'),
        'kropka w cudzysłowie': (['„Cytat.', 1, '”'], '„Cytat”[1].', '„Cytat”.[1]'),
        'nawias zamykający': (['(dopisek', 1, ')'], '(dopisek)[1]', '(dopisek[1])'),
        'kropka w nawiasie': (['(Całe zdanie.', 1, ')'], '(Całe zdanie.)[1]', '(Całe zdanie.[1])'),
        'za pytajnikiem': (['Naprawdę?', 1], 'Naprawdę?[1]', 'Naprawdę?[1]'),
        'przed pytajnikiem': (['Naprawdę', 1, '?'], 'Naprawdę[1]?', 'Naprawdę[1]?'),
        'dwa odsyłacze obok siebie': (['Tekst.', 1, 2], 'Tekst[1][2].', 'Tekst.[1][2]'),
        'początek akapitu': ([1, ' Tekst na początku.'], '[1] Tekst na początku.', '[1] Tekst na początku.'),
        'przypis końcowy': (['Koniec.', 1], 'Koniec[1].', 'Koniec.[1]', 'endnote'),
    }

    @tag('slow')
    def test_note_references_keep_continuity_and_follow_polish_rules(self):
        """Każdy przypadek bez śledzenia i ze śledzeniem (po przyjęciu zmian), z regułą i bez niej."""
        for name, (parts, with_rule, without_rule, *kind) in self.CASES.items():
            kind = kind[0] if kind else 'footnote'
            source = noted_document(parts, kind).getvalue()
            ids = [str(part) for part in parts if isinstance(part, int)]
            for rules, expected in ((self.on, with_rule), (self.off, without_rule)):
                for track_changes in (False, True):
                    with self.subTest(name, rule='note_references' in rules, track_changes=track_changes):
                        summary = {}
                        result = clean_docx(BytesIO(source), rules, track_changes=track_changes, summary=summary)
                        self.assertEqual(_noted_texts(result), expected)
                        moved = 'note_references' in rules and with_rule != without_rule
                        self.assertEqual(summary.get('note_references'), 1 if moved else None)
                        # Odsyłacz zostaje tym samym przebiegiem: ten sam w:id i styl, poza śledzonymi zmianami.
                        body = Document(BytesIO(result.getvalue())).element.body
                        references = list(body.iter(qn(f'w:{kind}Reference')))
                        self.assertEqual([node.get(qn('w:id')) for node in references], ids)
                        for node in references:
                            run = node.getparent()
                            self.assertEqual([child.tag for child in run], [qn('w:rPr'), node.tag])
                            self.assertEqual(run.getparent().tag, qn('w:p'))

    def test_custom_mark_reference_stays_with_its_mark(self):
        """Odsyłacz z własnym znakiem (customMarkFollows, znak „*” w osobnym przebiegu) jest barierą: reguła
        note_references nie oddziela go od znaku."""
        source = noted_document(['Tekst.', 1, '* dalej.'], reference='w:customMarkFollows="1"')
        self.assertEqual(_noted_texts(clean_docx(source, self.on)), 'Tekst.[1]* dalej.')

    def test_rule_alone_and_plain_text(self):
        """Odstęp przed odsyłaczem znika zawsze (bez reguły liczy go „spaces”); bez quote_punct odsyłacz idzie
        za cudzysłów zamykający razem z kropką w środku; w tekście bez odsyłaczy reguła nic nie robi."""
        for name, rules, parts, expected, counted in (
            ('sama reguła', ['note_references'], ['Słowo ', 1, ' dalej.'], 'Słowo[1] dalej.', {'note_references': 1}),
            ('bez reguły', ['ellipsis'], ['Słowo ', 1], 'Słowo[1]', {'spaces': 1}),
            ('bez quote_punct', self.on - {'quote_punct'}, ['„Cytat.', 1, '”'], '„Cytat.”[1]', {'note_references': 1}),
        ):
            with self.subTest(name):
                summary = {}
                result = clean_docx(noted_document(parts), rules, summary=summary)
                self.assertEqual(_noted_texts(result), expected)
                # Treść przypisu daje własne liczniki – odejmujemy je (ten sam dokument bez akapitu treści).
                baseline = {}
                clean_docx(noted_document([1]), rules, summary=baseline)
                self.assertEqual({key: summary.get(key, 0) - baseline.get(key, 0)
                                  for key in ('note_references', 'spaces')
                                  if summary.get(key, 0) != baseline.get(key, 0)}, counted)
        self.assertEqual(correct_editorial_text('Tekst. „Cytat.” (dopisek)', {'note_references'}),
                         'Tekst. „Cytat.” (dopisek)')
        self.assertIn('note_references', DEFAULT_EDITORIAL_RULES)

    def test_note_text_is_corrected_and_its_mark_is_untouched(self):
        """Treść przypisu nadal jest poprawiana, a w:footnoteRef / w:endnoteRef zostaje na początku."""
        for kind in NOTE_PARTS:
            for track_changes in (False, True):
                with self.subTest(kind, track_changes=track_changes):
                    result = clean_docx(noted_document(['Tekst.', 1], kind), self.on, track_changes=track_changes)
                    self.assertEqual(_noted_texts(result, f'word/{kind}s.xml'), '[ref] Treść przypisu…')
