"""Analiza powtórzeń: kolorowanie odmian słów, dodatkowe kontrole, palety kolorów, walidacja parametrów, model
językowy i pamięć lematów, nietypowe przebiegi tekstu, formularz i strona programu."""
import unicodedata
from collections import OrderedDict
from contextlib import contextmanager
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

from core.odkurzacz_forms import RepetitionsForm
from core.services import document_repetitions as repetitions
from core.services.document_converter import convert_document
from core.services.document_errors import DocumentInputError
from core.services.document_repetitions import LEXICAL_RE, _get_nlp, _unused_pipes, color_document
from fantazmaty.testing import create_member

from .support import (
    REPETITION_FIELDS, REPETITION_PARAMETERS, IsolatedProgramDirsMixin, add_complex_field, add_internal_link,
    add_tracked_insertion, docx_with_paragraph_xml, downloaded_document, footnote_text, run_program, saved,
    text_run, whole_document,
)

ONLY_REPETITIONS = {'duplicates': False, 'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}


def payload(document):
    return saved(document).getvalue()


def formatted_repetition():
    """„Koty koty.” – pogrubione pierwsze słowo, kursywa w drugim przebiegu."""
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Koty').bold = True
    paragraph.add_run(' koty.').italic = True
    return payload(document)


def font_colors(paragraph):
    return [run.font.color.rgb for run in paragraph.runs if run.font.color.rgb]


@tag('docx')
class RepeatAnalysisTests(SimpleTestCase):
    def test_real_polish_model_flexion_and_run_preservation(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run('Kot').bold = True
        paragraph.add_run(' spaceruje obok ')
        paragraph.add_run('kota').italic = True
        paragraph.add_run('.  Bez zmian!')
        document.add_paragraph('')
        marked = Document(color_document(BytesIO(payload(document)), min_word_length=3, analysis_options={
            'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}))
        self.assertEqual([p.text for p in marked.paragraphs], [p.text for p in document.paragraphs])
        self.assertEqual(marked.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        runs = marked.paragraphs[0].runs
        self.assertEqual(runs[0].font.color.rgb, runs[2].font.color.rgb)
        self.assertIsNotNone(runs[0].font.color.rgb)
        self.assertTrue(runs[0].bold)
        self.assertTrue(runs[2].italic)

    def test_cross_run_word_and_ignored_inflection(self):
        document = Document()
        p = document.add_paragraph()
        p.add_run('Ko').bold = True
        p.add_run('ty').italic = True
        p.add_run(' kot kota')
        data = payload(document)
        options = {'long_sentences': False, 'long_paragraphs': False}
        marked = Document(color_document(BytesIO(data), min_word_length=3, ignored_words='kota',
                                         analysis_options=options))
        self.assertEqual(marked.paragraphs[0].text, p.text)
        self.assertTrue(all(r.font.color.rgb is None and r.font.highlight_color is None
                            for r in marked.paragraphs[0].runs))
        marked = Document(color_document(BytesIO(data), min_word_length=3, analysis_options=options))
        self.assertEqual(marked.paragraphs[0].runs[0].font.color.rgb, marked.paragraphs[0].runs[1].font.color.rgb)
        self.assertIsNotNone(marked.paragraphs[0].runs[0].font.color.rgb)

    def test_priority_tracked_words_and_existing_underline(self):
        document = Document()
        document.add_paragraph().add_run('Kot kota').underline = True
        marked = Document(color_document(BytesIO(payload(document)), min_word_length=3, tracked_words='kot',
                                         analysis_options={'long_sentences': False, 'paragraph_limit': 1}))
        runs = marked.paragraphs[0].runs
        self.assertEqual(runs[0].font.highlight_color, WD_COLOR_INDEX.BRIGHT_GREEN)
        self.assertTrue(all(r.underline for r in runs))

    def test_extra_checks_real_sentence_parser(self):
        document = Document()
        document.add_paragraph('Kot spokojnie chodzi po domu. []')
        marked = Document(color_document(BytesIO(payload(document)),
                                         analysis_options={'sentence_limit': 2, 'paragraph_limit': 2}))
        highlights = {r.font.highlight_color for r in marked.paragraphs[0].runs}
        self.assertIn(WD_COLOR_INDEX.TURQUOISE, highlights)
        self.assertIn(WD_COLOR_INDEX.PINK, highlights)

    def test_repetitions_are_marked_in_tables_and_footnotes(self):
        result = color_document(whole_document())
        document = Document(BytesIO(result.getvalue()))
        cell = document.tables[0].cell(0, 0).paragraphs[0]
        self.assertTrue(cell._p.xpath('.//w:color'))
        self.assertEqual(cell.text, 'Lampa świeciła,  lampa gasła.')
        self.assertIn('w:color', footnote_text(BytesIO(result.getvalue())))


class ModelPipelineTests(SimpleTestCase):
    """Pomijane składniki modelu (dla szybkości) nie zmieniają lematów ani granic zdań."""

    TEXT = ('Kot spał na oknie, a koty sąsiadów biegały po ogrodzie. – Czy widziałaś Annę? – zapytał Piotr. '
            'Ludzie mówili, że lata mijały szybko; potem przyszły gorsze dni i jeszcze gorsze noce.')

    def test_reduced_pipeline_matches_full_pipeline(self):
        nlp = _get_nlp()
        words = sorted({word.lower() for word in LEXICAL_RE.findall(self.TEXT)})
        full = [doc[0].lemma_ for doc in nlp.pipe(words, disable=['parser'])]
        reduced = [doc[0].lemma_ for doc in nlp.pipe(words, disable=_unused_pipes(nlp, 'lemmas'))]
        self.assertEqual(reduced, full)
        # Słowa podane bez tokenizera (gotowe dokumenty jednowyrazowe) dają te same lematy.
        from spacy.tokens import Doc
        direct = [doc[0].lemma_ for doc in nlp.pipe((Doc(nlp.vocab, words=[word]) for word in words),
                                                    disable=_unused_pipes(nlp, 'lemmas'))]
        self.assertEqual(direct, full)
        # Jeden przebieg „context” daje te same granice zdań co przebieg tylko dla zdań.
        self.assertEqual([s.text for s in nlp(self.TEXT, disable=_unused_pipes(nlp, 'context')).sents],
                         [s.text for s in nlp(self.TEXT, disable=['lemmatizer']).sents])
        self.assertEqual([s.text for s in nlp(self.TEXT, disable=_unused_pipes(nlp, 'sentences')).sents],
                         [s.text for s in nlp(self.TEXT, disable=['lemmatizer']).sents])


@tag('docx')
class WordMarkupTests(SimpleTestCase):
    """Niewidoczne znaczniki Worda (sprawdzanie pisowni, zakładki, łamanie stron) nie dzielą tekstu."""

    def marked_paragraph(self, *fragments, **options):
        source = docx_with_paragraph_xml(*fragments)
        original = Document(source).paragraphs[0].text
        source.seek(0)
        marked = Document(color_document(source, min_word_length=3, analysis_options={
            'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}, **options)).paragraphs[0]
        self.assertEqual(marked.text, original)
        return marked

    def test_markers_between_words_do_not_hide_neighbouring_duplicates(self):
        """Word otacza powtórzone słowo znacznikami proofErr – właśnie to powtórzenie trzeba wykryć."""
        cases = {
            'proofErr': (text_run('Ala ma '), '<w:proofErr w:type="spellStart"/>', text_run('kota'),
                         '<w:proofErr w:type="spellEnd"/>', text_run(' '), '<w:proofErr w:type="spellStart"/>',
                         text_run('kota'), '<w:proofErr w:type="spellEnd"/>', text_run('.')),
            'lastRenderedPageBreak': (text_run('Ala ma kota '),
                                      '<w:r><w:lastRenderedPageBreak/><w:t>kota.</w:t></w:r>'),
            'komentarz': (text_run('Ala ma kota'), '<w:commentRangeStart w:id="1"/>',
                          '<w:r><w:commentReference w:id="1"/></w:r>', text_run(' kota.')),
            'miękki enter': ('<w:r><w:t>Ala ma kota</w:t><w:br/><w:t>kota.</w:t></w:r>',),
        }
        for name, fragments in cases.items():
            with self.subTest(name):
                paragraph = self.marked_paragraph(*fragments)
                highlights = [run.font.highlight_color for run in paragraph.runs if 'kota' in run.text]
                self.assertEqual(highlights, [WD_COLOR_INDEX.BRIGHT_GREEN] * 2)

    def test_marker_inside_a_word_keeps_it_one_word(self):
        """Zakładka _GoBack lub łącznik miękki w środku słowa nie może go rozdzielić."""
        for name, marker in (('zakładka', '<w:bookmarkStart w:id="0" w:name="_GoBack"/>'),
                             ('łącznik miękki', '<w:r><w:softHyphen/></w:r>')):
            with self.subTest(name):
                paragraph = self.marked_paragraph(text_run('Ko'), marker, text_run('ty i koty'))
                colors = {run.font.color.rgb for run in paragraph.runs if run.text in ('Ko', 'ty', 'koty')}
                self.assertEqual(len(colors), 1)
                self.assertIsNotNone(colors.pop())

    def test_decomposed_polish_letters_stay_inside_words(self):
        """Tekst w postaci NFD (np. z Pages lub PDF): „ę” to „e” + ogonek; słowo nie może się rozpaść."""
        decomposed = unicodedata.normalize('NFD', 'Pięść i pięść, a potem Pięść.')
        paragraph = self.marked_paragraph(text_run(decomposed), tracked_words='pięść')
        colored = [run.text for run in paragraph.runs if run.font.color.rgb]
        self.assertEqual(colored, [unicodedata.normalize('NFD', word) for word in ('Pięść', 'pięść', 'Pięść')])
        self.assertEqual(len({run.font.color.rgb for run in paragraph.runs if run.font.color.rgb}), 1)
        self.assertTrue(all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW
                            for run in paragraph.runs if run.font.color.rgb))

    def test_text_box_in_a_run_with_marked_text_keeps_its_own_marks(self):
        """Regresja: podział runu z polem tekstowym kopiował pole przed oznaczeniem jego treści."""
        shape = ('<w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
                 '<wp:extent cx="900000" cy="400000"/><wp:docPr id="9" name="Pole"/>'
                 '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                 '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
                 '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:spPr/>'
                 '<wps:txbx><w:txbxContent><w:p><w:r><w:t>Psa szczeka i psa</w:t></w:r></w:p></w:txbxContent>'
                 '</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>')
        source = docx_with_paragraph_xml(
            f'<w:r><w:t xml:space="preserve">Kota widać </w:t>{shape}<w:t xml:space="preserve"> i kota.</w:t></w:r>')
        result = color_document(source, min_word_length=3, analysis_options=ONLY_REPETITIONS)
        body = Document(result).element.body
        self.assertEqual(len(body.xpath('.//w:txbxContent')), 1)
        self.assertEqual(''.join(body.xpath('.//w:txbxContent//w:t/text()')), 'Psa szczeka i psa')
        self.assertEqual(len(body.xpath('.//w:txbxContent//w:color')), 2)
        self.assertEqual(len(body.xpath('./w:p/w:r/w:rPr/w:color')), 2)


@tag('docx')
class ColorPaletteTests(SimpleTestCase):
    def test_both_palettes_keep_text_and_formatting_with_correct_rgb_range(self):
        for palette, low, high in [('dark', 0, 150), ('light', 160, 220)]:
            with self.subTest(palette=palette), patch(
                    'core.services.document_repetitions._lemma_map',
                    side_effect=lambda words: {word.lower(): word.lower() for word in words}):
                document = Document(color_document(BytesIO(formatted_repetition()), color_palette=palette,
                                                   analysis_options=ONLY_REPETITIONS))
                self.assertEqual(document.paragraphs[0].text, 'Koty koty.')
                colors = font_colors(document.paragraphs[0])
                self.assertTrue(colors)
                self.assertEqual(len(set(colors)), 1)
                self.assertTrue(all(low <= channel <= high for color in colors for channel in color))
                self.assertTrue(document.paragraphs[0].runs[0].bold)
                self.assertTrue(any(run.italic for run in document.paragraphs[0].runs))


@tag('docx')
class RepetitionsWorkerTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_worker_roundtrip_keeps_text_and_applies_light_palette(self):
        """Proces roboczy koloruje powtórzenia jasną paletą i nie wprowadza poprawek Odkurzacza."""
        document = Document()
        document.add_paragraph('Koty  koty...')
        output, extension, _ = convert_document(
            SimpleUploadedFile('test.docx', payload(document)), [], include_docx=True, normalize=False,
            repetitions={'color_palette': 'light', 'analysis_options': {'long_sentences': False}})
        with output:
            marked = Document(output)
        self.assertEqual(extension, 'docx')
        self.assertEqual(marked.paragraphs[0].text, 'Koty  koty...')
        self.assertIsNotNone(marked.paragraphs[0].runs[0].font.color.rgb)
        colors = font_colors(marked.paragraphs[0])
        self.assertTrue(colors)
        self.assertTrue(all(160 <= channel <= 220 for color in colors for channel in color))


class RepetitionsFormTests(SimpleTestCase):
    def upload(self):
        return {'document': SimpleUploadedFile('test.docx', formatted_repetition())}

    def test_form_rejects_bad_parameters_and_defaults_match_desktop(self):
        form = RepetitionsForm({**REPETITION_PARAMETERS, 'window_size': '0'}, self.upload())
        self.assertFalse(form.is_valid())
        self.assertIn('window_size', form.errors)
        self.assertTrue(RepetitionsForm().fields['duplicates'].initial)

    def test_form_accepts_light_and_defaults_missing_palette_to_dark(self):
        for palette in ('light', ''):
            form = RepetitionsForm({**REPETITION_PARAMETERS, 'color_palette': palette}, self.upload())
            self.assertTrue(form.is_valid(), form.errors)
            self.assertEqual(form.analysis_config()['color_palette'], palette or 'dark')
        form = RepetitionsForm({**REPETITION_PARAMETERS, 'color_palette': 'invalid'}, self.upload())
        self.assertFalse(form.is_valid())
        self.assertIn('color_palette', form.errors)


class RepetitionsPageTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('reader', 'Redaktor')

    def test_member_download_and_separate_form(self):
        self.client.force_login(self.member)
        response = self.client.get(reverse('core:programs'))
        self.assertContains(response, 'Kolorowanie powtórzeń')
        self.assertContains(response, 'name="repetitions-window_size"')
        self.assertContains(response, 'name="repetitions-color_palette"')
        document = Document()
        document.add_paragraph('Koty koty')
        with patch('core.services.program_jobs.convert_document',
                   return_value=(BytesIO(payload(document)), 'docx', 'application/octet-stream')) as convert:
            response = run_program(self.client, {
                'program_action': 'repetitions',
                'repetitions-document': SimpleUploadedFile('test.docx', payload(document)),
                'repetitions-color_palette': 'light',
                **REPETITION_FIELDS,
            })
        self.assertEqual(response.status_code, 200)
        self.assertIn('_powtorzenia.docx', response['Content-Disposition'])
        self.assertFalse(convert.call_args.kwargs['normalize'])
        self.assertEqual(convert.call_args.kwargs['repetitions']['color_palette'], 'light')
        self.assertNotIn('use_cleaner', convert.call_args.kwargs)
        self.assertEqual(downloaded_document(response).paragraphs[0].text, 'Koty koty')


@tag('docx')
class ContextLemmaTests(SimpleTestCase):
    """Lemat z kontekstu zdania to drugi klucz grupowania; granice zdań pochodzą z tego samego przebiegu."""

    def test_forms_joined_only_by_the_context_lemma_share_a_colour(self):
        document = Document()
        document.add_paragraph('Pręty leżały obok, a jeden pręt był zgięty.')
        text = document.paragraphs[0].text
        starts = {word: text.index(word) for word in ('Pręty', 'pręt', 'leżały', 'zgięty')}
        identity = lambda words: {word.lower(): word.lower() for word in words}  # noqa: E731
        context = [({starts['Pręty']: 'pręt', starts['pręt']: 'pręt', starts['leżały']: 'leżeć'},
                    [(0, len(text))])]
        with patch('core.services.document_repetitions._lemma_map', side_effect=identity), \
                patch('core.services.document_repetitions._context_analysis', return_value=context) as analysis:
            marked = Document(color_document(BytesIO(payload(document)), analysis_options={
                **ONLY_REPETITIONS, 'long_sentences': True, 'sentence_limit': 3}))
        analysis.assert_called_once()
        self.assertTrue(analysis.call_args.kwargs['sentences'])
        runs = marked.paragraphs[0].runs
        coloured = [run.text for run in runs if run.font.color.rgb]
        self.assertEqual(coloured, ['Pręty', 'pręt'])
        self.assertTrue(all(run.font.highlight_color == WD_COLOR_INDEX.TURQUOISE for run in runs))

    def test_sentence_pass_is_skipped_when_long_sentences_are_off(self):
        document = Document()
        document.add_paragraph('Kot i kot.')
        with patch('core.services.document_repetitions._context_analysis', return_value=[({}, [])]) as analysis:
            color_document(BytesIO(payload(document)), analysis_options=ONLY_REPETITIONS)
        self.assertFalse(analysis.call_args.kwargs['sentences'])


@contextmanager
def simple_language():
    """Lemat = słowo małymi literami, bez lematów z kontekstu i z jednym zdaniem na akapit (bez modelu)."""
    def context(texts, *, sentences):
        return [({}, [(0, len(text))]) for text in texts]

    with patch.object(repetitions, '_lemma_map', side_effect=lambda words: {w.lower(): w.lower() for w in words}), \
            patch.object(repetitions, '_context_analysis', side_effect=context):
        yield


def marked(source, **options):
    with simple_language():
        return Document(color_document(source, **options))


def coloured(paragraph):
    return [run.text for run in paragraph.runs if run.font.color.rgb]


class ParameterTests(SimpleTestCase):
    def test_parameters_out_of_range_are_refused(self):
        source = saved(Document())
        cases = {
            'nieznana paleta': ({'color_palette': 'invalid'}, ValueError),
            'zakres jako wartość logiczna': ({'window_size': True}, ValueError),
            'zero słów': ({'min_word_length': '0'}, ValueError),
            'zakres ponad 500': ({'window_size': 501}, DocumentInputError),
            'nieznana kontrola': ({'analysis_options': {'nieznana': True}}, DocumentInputError),
            'próg zdania ponad 10 000': ({'analysis_options': {'sentence_limit': 10_001}}, DocumentInputError),
        }
        for label, (options, error) in cases.items():
            with self.subTest(label), self.assertRaises(error):
                color_document(source, **options)

    def test_document_with_tracked_changes_is_refused(self):
        document = Document()
        add_tracked_insertion(document.add_paragraph('Tekst '), 'dopisany')
        with self.assertRaises(DocumentInputError) as raised:
            color_document(saved(document))
        self.assertEqual(raised.exception.public_code, 'tracked_changes')


class LanguageModelTests(SimpleTestCase):
    def test_missing_polish_model_is_an_import_error(self):
        with patch('spacy.load', side_effect=OSError('brak pakietu')), \
                self.assertRaisesMessage(ImportError, 'Brakuje polskiego modelu analizy.'):
            repetitions._get_nlp.__wrapped__()  # bez pamięci podręcznej prawdziwego modelu

    def test_rule_based_lemmatizer_keeps_the_tagger(self):
        """Lematyzator regułowy potrzebuje części mowy: pomijamy tylko parser, senter i NER."""
        nlp = SimpleNamespace(pipe_names=['tok2vec', 'tagger', 'parser', 'lemmatizer', 'ner'],
                              get_pipe=lambda name: object())
        self.assertEqual(repetitions._unused_pipes(nlp, 'lemmas'), ['parser', 'ner'])
        self.assertEqual(repetitions._lemma_from_doc([], 'słowo'), 'słowo')

    def test_word_lemmas_with_a_bounded_cache(self):
        with patch.object(repetitions, '_LEMMA_CACHE', OrderedDict()) as cache, \
                patch.object(repetitions, '_LEMMA_CACHE_LIMIT', 2):
            self.assertEqual(repetitions._lemma_map(['KOTAMI']), {'kotami': 'kot'})
            self.assertEqual(repetitions._lemma_map(['psami', 'domami']), {'psami': 'pies', 'domami': 'dom'})
            self.assertEqual(list(cache), [(repetitions.MODEL_NAME, 'psami'), (repetitions.MODEL_NAME, 'domami')])


@tag('docx')
class GroupingTests(SimpleTestCase):
    def test_similar_beginnings_and_repetitions_outside_the_window(self):
        document = Document()
        document.add_paragraph('Domownik domowy.')
        document.add_paragraph('Lasy szumią wokół, potem lasy.')
        cases = {'bez podobnego początku, okno 1': ({'window_size': 1}, [], []),
                 'podobny początek, okno 35': ({'include_prefix_matches': True}, ['Domownik', 'domowy'],
                                               ['Lasy', 'lasy'])}
        for label, (options, first, second) in cases.items():
            with self.subTest(label):
                result = marked(saved(document), analysis_options=ONLY_REPETITIONS, **options)
                self.assertEqual([coloured(p) for p in result.paragraphs], [first, second])


    def test_neighbouring_duplicate_across_a_paragraph_boundary(self):
        document = Document()
        document.add_paragraph('Widzę kota')
        document.add_paragraph('kota widzę.')
        result = marked(saved(document), analysis_options={**ONLY_REPETITIONS, 'duplicates': True})
        self.assertEqual([[run.text for run in p.runs if run.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN]
                          for p in result.paragraphs], [['kota'], ['kota']])


@tag('docx')
class UnusualRunTests(SimpleTestCase):
    def test_empty_text_nodes_keep_their_place(self):
        source = docx_with_paragraph_xml(
            '<w:r><w:t/><w:t xml:space="preserve">Koty i koty, </w:t></w:r>',
            '<w:r><w:t xml:space="preserve">lisy </w:t><w:t/><w:t xml:space="preserve">i lisy.</w:t></w:r>')
        paragraph = marked(source, analysis_options=ONLY_REPETITIONS).paragraphs[0]
        self.assertEqual(paragraph.text, 'Koty i koty, lisy i lisy.')
        self.assertEqual(coloured(paragraph), ['Koty', 'koty', 'lisy', 'lisy'])
        self.assertEqual(len(paragraph._p.xpath('.//w:t[not(text())]')), 2)

    def test_links_and_field_results_are_not_words_of_the_text(self):
        document = Document()
        paragraph = document.add_paragraph('Lampa ')
        add_internal_link(paragraph, '_Toc1', 'lampa')
        add_complex_field(paragraph, 'AUTHOR', 'lampa')
        paragraph.add_run(' i lampa.')
        result = marked(saved(document), analysis_options=ONLY_REPETITIONS).paragraphs[0]
        # Kolor dostają tylko słowa tekstu: ani odnośnik, ani wynik pola (między separate i end).
        self.assertEqual(result._p.xpath('.//w:r[w:rPr/w:color]/w:t/text()'), ['Lampa', 'lampa'])
        self.assertEqual(result.runs[-2].text, 'lampa')
        self.assertIsNotNone(result.runs[-2].font.color.rgb)

    def test_runs_with_backgrounds_only(self):
        """Puste nawiasy w osobnym przebiegu dostają samo tło; sąsiednie pary – jeden przebieg."""
        source = docx_with_paragraph_xml(text_run('Tekst '), text_run('()'), text_run(' i '), text_run('[]{}'))
        runs = marked(source, analysis_options={**ONLY_REPETITIONS, 'empty_pairs': True}).paragraphs[0].runs
        self.assertEqual([(run.text, run.font.highlight_color, run.font.color.rgb) for run in runs], [
            ('Tekst ', None, None), ('()', WD_COLOR_INDEX.PINK, None), (' i ', None, None),
            ('[]{}', WD_COLOR_INDEX.PINK, None)])
