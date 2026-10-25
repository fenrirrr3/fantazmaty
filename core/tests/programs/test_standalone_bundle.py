"""Program lokalny (jeden plik): budowa z modułów core/services i działanie wszystkich czterech programów
bez Django."""
import ast
import importlib.util
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from unittest.mock import patch

from django.test import SimpleTestCase, tag
from docx import Document

from core.services.document_errors import MESSAGES
from core.services.document_images import linked_images_note

from .editor_fixtures_exotic import foreign_file
from .support import docx_with_linked_images, external_images, saved

BUILD = Path(__file__).resolve().parents[3] / 'tools' / 'redaktor_docx' / 'build.py'


def load_bundle(path):
    spec = importlib.util.spec_from_file_location('redaktor_docx_bundle', path)
    module = importlib.util.module_from_spec(spec)
    sys.modules['redaktor_docx_bundle'] = module
    spec.loader.exec_module(module)
    return module


def load_build():
    spec = importlib.util.spec_from_file_location('redaktor_docx_build', BUILD)
    build = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(build)
    return build


class BundleNamesTests(SimpleTestCase):
    def test_bundled_modules_define_each_global_name_once(self):
        """Scalone moduły nie definiują tej samej nazwy globalnej (funkcji, klasy, stałej) dwa razy: w jednym
        pliku późniejsza definicja nadpisałaby wcześniejszą – tak przebudowa przypisów i komentarzy w programie
        lokalnym wołała kiedyś funkcję _story_blocks z modułu PDF."""
        build = load_build()
        owners = {}
        for module in build.MODULES:
            tree = ast.parse((build.SERVICES / f'{module}.py').read_text(encoding='utf-8'))
            names = {node.name for node in tree.body if isinstance(node, (ast.FunctionDef, ast.ClassDef))}
            names |= {target.id for node in tree.body if isinstance(node, (ast.Assign, ast.AnnAssign))
                      for target in getattr(node, 'targets', [getattr(node, 'target', None)])
                      if isinstance(target, ast.Name)}
            renames = build.RENAMES.get(module, {})
            for name in names:
                owners.setdefault(renames.get(name, name), []).append(module)
        self.assertEqual({name: modules for name, modules in owners.items() if len(modules) > 1}, {})


@tag('slow', 'docx')
class StandaloneBundleTests(SimpleTestCase):
    def test_bundle_builds_and_runs_every_program_without_django(self):
        build = load_build()
        with TemporaryDirectory() as folder:
            target = Path(folder) / 'redaktor_docx.py'
            build.build(target)
            source = target.read_text(encoding='utf-8')
            self.assertNotIn('from .document_', source)
            self.assertNotIn("Path(__file__).with_name('fonts')", source)
            bundle = load_bundle(target)
            import os
            os.environ['XDG_CONFIG_HOME'] = str(Path(folder) / 'config')
            document = Document()
            document.add_paragraph('Koty i koty "siedzą"  obok ,tak. Karnister stał - powiedział.')
            document.add_paragraph('Drugi akapit z kotem.')
            document.add_paragraph('Czyby to był nie najlepszy niby-artysta, pół żartem superPolak.')
            docx = Path(folder) / 'tekst.docx'
            docx.write_bytes(saved(document).getvalue())
            settings, warning = bundle.load_settings()
            self.assertEqual(warning, '')
            progress = []
            results = {}
            for program in bundle.PROGRAMS:
                result, notes = bundle.run_program(program, docx, settings, progress.append)
                results[program] = (result, notes)
            cleaned = Document(results['Odkurzacz'][0]).paragraphs[0].text
            self.assertEqual(cleaned, 'Koty i koty „siedzą” obok, tak. Kanister stał – powiedział.')
            self.assertEqual(Document(results['Odkurzacz'][0]).paragraphs[2].text,
                             'Czy by to był nienajlepszy nibyartysta, półżartem super-Polak.')
            self.assertIn('Podsumowanie zmian', results['Odkurzacz'][1][0])
            pdf, epub = results['Konwersja'][0]
            self.assertTrue(pdf.stat().st_size > 1000 and epub.stat().st_size > 1000)
            coloured = Document(results['Powtórzenia'][0])
            self.assertTrue(coloured.element.body.xpath('.//w:color'))
            self.assertEqual(Document(results['Przebudowa'][0]).paragraphs[1].text, 'Drugi akapit z kotem.')
            self.assertTrue(progress)
            # Ustawienia i czcionki trafiają do katalogu konfiguracji; słownik zamian ma format tekstowy.
            bundle.save_settings(settings)
            stored = json.loads((Path(folder) / 'config' / 'RedaktorDOCX' / 'ustawienia.json').read_text(encoding='utf-8'))
            self.assertEqual(stored['rules'], settings['rules'])
            # Ustawienia zapisane przed regułami „Ortografia 2026”: nowe reguły włączone, odznaczona zostaje.
            old_rules = {key: key != 'spaces' for key in stored['rules'] if key not in bundle.ORTHOGRAPHY_2026_KEYS}
            bundle.save_settings({**stored, 'rules': old_rules})
            rules = bundle.load_settings()[0]['rules']
            self.assertTrue(all(rules[key] for key in bundle.ORTHOGRAPHY_2026_KEYS))
            self.assertFalse(rules['spaces'])
            self.assertEqual(bundle.parse_replacements('kot = pies\nre:\\bpsa\\b = kota\n# uwaga')[1],
                             {'kind': 'pattern', 'source': r'\bpsa\b', 'target': 'kota'})
            self.assertTrue((Path(folder) / 'config' / 'RedaktorDOCX' / 'fonts' / 'NimbusRoman-Regular.ttf').exists())

    def test_bundle_names_files_of_other_formats(self):
        """Plik .doc albo ODT z rozszerzeniem .docx: program lokalny pokazuje ten sam komunikat co CMS."""
        build = load_build()
        with TemporaryDirectory() as folder, patch.dict('os.environ', {'XDG_CONFIG_HOME': str(Path(folder) / 'config')}):
            target = Path(folder) / 'redaktor_docx.py'
            build.build(target)
            bundle = load_bundle(target)
            settings, _ = bundle.load_settings()
            for kind, code in (('doc', 'format_doc'), ('odt', 'format_odt')):
                with self.subTest(kind):
                    source = Path(folder) / 'tekst.docx'
                    source.write_bytes(foreign_file(kind))
                    with self.assertRaisesMessage(ValueError, MESSAGES[code]):
                        bundle.run_program('Konwersja', source, settings, lambda text: None)

    def test_bundle_unlinks_linked_images_with_and_without_image_removal(self):
        """Program lokalny jak CMS: link obrazu do pliku zewnętrznego znika zawsze (z uwagą), także gdy obrazy
        nie są usuwane – wtedy obraz osadzony i połączony zostaje jako osadzony."""
        from zipfile import ZipFile
        build = load_build()
        with TemporaryDirectory() as folder, patch.dict('os.environ', {'XDG_CONFIG_HOME': str(Path(folder) / 'config')}):
            target = Path(folder) / 'redaktor_docx.py'
            build.build(target)
            bundle = load_bundle(target)
            settings, _ = bundle.load_settings()
            source = Path(folder) / 'obrazy.docx'
            source.write_bytes(docx_with_linked_images().getvalue())
            for strip in (True, False):
                with self.subTest(strip_images=strip):
                    settings['strip_images'] = strip
                    result, notes = bundle.run_program('Odkurzacz', source, settings, lambda text: None)
                    self.assertIn(linked_images_note(7), notes)
                    self.assertEqual('Usunięto obrazy: 1.' in notes, strip)
                    with result.open('rb') as stream:
                        self.assertEqual(external_images(stream), [])
                    with ZipFile(result) as archive:
                        self.assertEqual(any(name.startswith('word/media/') for name in archive.namelist()), not strip)
