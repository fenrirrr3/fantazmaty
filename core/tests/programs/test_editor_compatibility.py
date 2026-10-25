"""Zgodność z edytorami: DOCX „jak z LibreOffice” (Transitional i Strict) i „jak z Google Docs” przechodzą
cały łańcuch programów (przegląd, przebudowa, przygotowanie z Odkurzaczem, miękkie entery, powtórzenia,
konwersja do EPUB/PDF, podział i scalanie) bez wyjątku, a generator wiernie odtwarza cechy próbek.

Scenariusze skrajne (editor_fixtures_exotic: Word 365/2007, WPS, OnlyOffice, Pages, Google Docs, minimalne
eksportery i elementy problematyczne): przegląd i przebudowa bez wyjątku i bez cichej utraty tekstu,
konwersja bez utraty widocznego tekstu, a pliki innych formatów i niebezpieczne pakiety – z czytelnym
komunikatem przed konwersją."""
import json
import re
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django import forms
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

from core.odkurzacz_forms import validate_docx_upload
from core.services import document_progress
from core.services.document_conversion_worker import convert, main
from core.services.document_converter import DAMAGED_MESSAGE, ConversionError, convert_document, validate_resources
from core.services.document_errors import MESSAGES
from core.services.document_images import force_polish, strip_images
from core.services.document_merge import merge_documents
from core.services.document_preparation import prepare_docx
from core.services.document_rebuild import (
    HYPERLINK_OMISSION, UNKNOWN_FORMAT_OMISSION, RebuildUnsupported, inspect_docx, rebuild_docx,
)
from core.services.document_repetitions import color_document
from core.services.document_split import split_document
from core.services.document_whitespace import normalize_spacing_docx

from .editor_fixtures import (
    ALIGNED, HEADER_TEXT, LAST, LINK_TEXT, NOTE_TEXT, TEXT_BOX, expected_texts, google_docs_docx, libreoffice_docx,
)
from .editor_fixtures_exotic import EDITOR_SCENARIOS, FEATURE_SCENARIOS, REFUSED_SCENARIOS, foreign_file
from .support import IsolatedProgramDirsMixin
from .test_pdf_rendering import pdf_html_spy

LISTS_OMISSION = 'automatyczne numerowanie i punktory list (tekst pozycji pozostanie)'
TEXT_BOX_OMISSION = 'ilustracje i pola tekstowe wraz z ich treścią'
MARGIN_OMISSION = 'niepuste nagłówki lub stopki'
CLEANER_RULES = ['quotes', 'spaces']
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
ALIGNMENTS = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER,
              WD_ALIGN_PARAGRAPH.JUSTIFY]

# Edytor → (generator, pominięcia zgłaszane przez przegląd dla pliku z listami).
EDITORS = {
    'libreoffice_transitional': (lambda **options: libreoffice_docx('transitional', **options),
                                 [f'{LISTS_OMISSION} (liczba: 2)']),
    'libreoffice_strict': (lambda **options: libreoffice_docx('strict', **options),
                           [f'{LISTS_OMISSION} (liczba: 2)']),
    'google_docs': (google_docs_docx, [f'{LISTS_OMISSION} (liczba: 2)', f'{HYPERLINK_OMISSION} (liczba: 1)']),
}


def variants(footnotes=(False, True)):
    """(nazwa przypadku, generator z ustawionymi przypisami, oczekiwane pominięcia) dla każdego edytora."""
    for editor, (build, omissions) in EDITORS.items():
        for noted in footnotes:
            yield f'{editor}, przypisy: {noted}', (lambda build=build, noted=noted: build(footnotes=noted)), omissions


def texts_of(editor, **options):
    """Oczekiwane teksty akapitów; w pliku Google Docs ostatni akapit ma jeszcze tekst hiperłącza."""
    texts = expected_texts(**options)
    if editor.startswith('google'):
        texts[-1] = f'{LAST} {LINK_TEXT}'
    return texts


def tracked_text(paragraph):
    """Tekst akapitu z przyjętymi śledzonymi zmianami (wstawki w w:ins liczą się, usunięcia nie)."""
    glyphs = {qn('w:tab'): '\t', qn('w:br'): '\n'}
    return ''.join(
        glyphs.get(node.tag, node.text or '')
        for node in paragraph._p.iter(qn('w:t'), *glyphs)
        if node.getparent().tag == qn('w:r')  # nie tabulatory z w:pPr/w:tabs
    )


def omissions_of(stream):
    try:
        inspect_docx(stream)
    except RebuildUnsupported as error:
        return error.omissions
    return []


def part(stream, name):
    with ZipFile(stream) as archive:
        return archive.read(name).decode('utf-8')


class EditorFixtureFidelityTests(SimpleTestCase):
    """Generator odtwarza cechy próbek LibreOffice 24.2 (i znane cechy Google Docs) – lista w docstringu
    editor_fixtures; sprawdzamy te, od których zależy tolerancja programów."""

    def test_libreoffice_variants_match_the_samples(self):
        cases = {
            'transitional': {'jc_left': 'left', 'jc_right': 'right', 'ind': 'left', 'compat': '12',
                             'charset': '<w:charset w:val="00"/>', 'core_rel': 'officeDocument/2006'},
            'strict': {'jc_left': 'start', 'jc_right': 'end', 'ind': 'start', 'compat': '15',
                       'charset': '<w:charset w:val="00" w:characterSet="windows-1252"/>',
                       'core_rel': 'officedocument/2006'},
        }
        for variant, expected in cases.items():
            with self.subTest(variant):
                stream = libreoffice_docx(variant, footnotes=True)
                document, styles = part(stream, 'word/document.xml'), part(stream, 'word/styles.xml')
                numbering, settings = part(stream, 'word/numbering.xml'), part(stream, 'word/settings.xml')
                other = {'left': 'start', 'start': 'left'}[expected['jc_left']]
                self.assertIn(f'<w:jc w:val="{expected["jc_left"]}"/>', document)
                self.assertIn(f'<w:jc w:val="{expected["jc_right"]}"/>', document)
                self.assertNotIn(f'<w:jc w:val="{other}"/>', document + styles)
                self.assertIn('<w:jc w:val="center"/>', document)
                self.assertIn('<w:jc w:val="both"/>', document)
                self.assertIn(f'<w:ind w:{expected["ind"]}="720"/>', styles)
                self.assertIn(f'<w:lvlJc w:val="{expected["jc_left"]}"/>', numbering)
                self.assertIn(f'w:val="{expected["compat"]}"/></w:compat>', settings)
                self.assertIn(expected['charset'], part(stream, 'word/fontTable.xml'))
                self.assertIn(f'http://schemas.openxmlformats.org/{expected["core_rel"]}/relationships/metadata/'
                              'core-properties', part(stream, '_rels/.rels'))
                # Cechy wspólne obu odmian.
                self.assertIn('<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Cambria"', styles)
                self.assertIn('<w:style w:type="paragraph" w:styleId="Normal"><w:name w:val="Normal"/>', styles)
                self.assertNotIn('w:default="1"', styles)
                self.assertIn('<w:pStyle w:val="Normal"/>', document)
                self.assertIn('<w:keepNext w:val="true"/><w:keepLines/><w:numPr><w:ilvl w:val="0"/>'
                              '<w:numId w:val="0"/></w:numPr>', document)
                self.assertIn('<w:tab w:val="clear" w:pos="720"/><w:tab w:val="right" w:pos="2835" w:leader="dot"/>',
                              document)
                self.assertIn('<w:widowControl/><w:bidi w:val="0"/><w:spacing w:lineRule="auto" w:line="276"',
                              document)
                self.assertIn('<w:numFmt w:val="none"/><w:suff w:val="nothing"/>', numbering)
                self.assertIn('<w:numFmt w:val="bullet"/><w:lvlText w:val=""/>', numbering)
                self.assertIn('<w:pStyle w:val="ListBullet"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/>',
                              document)
                self.assertIn('<w:type w:val="nextPage"/><w:pgSz w:w="12240" w:h="15840"/>', document)
                self.assertIn('<w:docGrid w:type="default" w:linePitch="360" w:charSpace="4096"/>', document)
                self.assertNotIn('w:rsid', document + settings)
                self.assertNotIn('w:proofErr', document)
                self.assertNotIn('_GoBack', document)
                self.assertNotIn('w:proofState', settings)
                self.assertIn('<Application>LibreOffice/24.2', part(stream, 'docProps/app.xml'))
                footnotes = part(stream, 'word/footnotes.xml')
                self.assertIn('<w:footnote w:id="0" w:type="separator">', footnotes)
                self.assertIn('<w:footnote w:id="1" w:type="continuationSeparator">', footnotes)
                self.assertIn('<w:footnote w:id="2"><w:p><w:pPr><w:pStyle w:val="Footnote"/>', footnotes)
                self.assertIn('<w:footnoteReference w:id="2"/>', document)
                self.assertIn('<w:sectPr><w:footnotePr><w:numFmt w:val="decimal"/></w:footnotePr>', document)
                self.assertIn('<w:footnotePr><w:footnote w:id="0"/><w:footnote w:id="1"/></w:footnotePr>', settings)

    def test_libreoffice_options_add_margin_content_and_text_box(self):
        stream = libreoffice_docx(header_footer=True, text_box=True, lists=False)
        document = part(stream, 'word/document.xml')
        self.assertIn('<w:headerReference w:type="default" r:id="rId7"/>', document)
        self.assertIn('<w:footerReference w:type="default" r:id="rId8"/>', document)
        self.assertIn('w:header="720"', document)
        self.assertIn('<mc:AlternateContent><mc:Choice Requires="wps">', document)
        self.assertIn('<mc:Fallback><w:pict><v:rect', document)
        self.assertNotIn('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/>', document)
        self.assertIn('<w:hdr ', part(stream, 'word/header1.xml'))
        self.assertIn('<w:pStyle w:val="Footer"/>', part(stream, 'word/footer1.xml'))

    def test_google_docs_fixture_has_the_known_export_traits(self):
        stream = google_docs_docx(footnotes=True)
        document, styles = part(stream, 'word/document.xml'), part(stream, 'word/styles.xml')
        with ZipFile(stream) as archive:
            self.assertNotIn('docProps/app.xml', archive.namelist())
        self.assertIn('<w:style w:type="paragraph" w:styleId="Normal" w:default="1"><w:name w:val="normal"/>', styles)
        self.assertIn('w:styleId="Heading1"><w:name w:val="heading 1"/>', styles)
        self.assertIn('<w:name w:val="Title"/>', styles)
        self.assertIn('<w:name w:val="Subtitle"/>', styles)
        self.assertIn('<w:rFonts w:ascii="Arial" w:cs="Arial" w:eastAsia="Arial" w:hAnsi="Arial"/><w:sz w:val="22"/>',
                      styles)
        self.assertIn('<w:spacing w:after="0" w:before="0" w:line="276" w:lineRule="auto"/>', styles)
        self.assertIn('<w:pStyle w:val="Heading1"/><w:keepNext w:val="1"/><w:keepLines w:val="1"/>', document)
        self.assertIn('<w:bookmarkStart w:colFirst="0" w:colLast="0" w:name="_heading=h.', document)
        self.assertIn('<w:b w:val="1"/>', document)
        self.assertIn('<w:u w:val="single"/>', document)
        self.assertIn('<w:color w:val="000000"/><w:rtl w:val="0"/>', document)
        self.assertIn('<w:hyperlink r:id="rId5">', document)
        self.assertIn('TargetMode="External"', part(stream, 'word/_rels/document.xml.rels'))
        self.assertIn('<w:pgSz w:w="11909" w:h="16834"/>', document)
        self.assertIn('<w:lvlText w:val="●"/>', part(stream, 'word/numbering.xml'))
        self.assertNotIn('<w:pStyle w:val="ListBullet"/>', document)
        self.assertNotIn('w:pBdr', document)
        self.assertIn('<w:footnote w:id="1"><w:p>', part(stream, 'word/footnotes.xml'))

    def test_python_docx_reads_every_fixture(self):
        """Każdy plik otwiera się w python-docx; nazwę „heading 1” Google Docs python-docx tłumaczy na
        „Heading 1” (więc nagłówek jest rozpoznawany), ale „normal” już zostaje małymi literami."""
        for case, build, _ in variants():
            with self.subTest(case):
                document = Document(build())
                self.assertEqual([p.text for p in document.paragraphs], texts_of(case))
                self.assertEqual(document.paragraphs[0].style.name, 'Heading 1')
                self.assertEqual(document.paragraphs[1].style.name.casefold(), 'normal')
                self.assertTrue(document.paragraphs[5].runs[0].bold)
                self.assertTrue(document.paragraphs[5].runs[1].italic)


@tag('slow', 'docx')
class EditorChainTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Każdy wariant przechodzi łańcuch programów działających w tym procesie."""

    def test_inspection_reports_only_the_expected_omissions(self):
        """Przegląd zgłasza tylko listy (i hiperłącze Google Docs) – bez „nierozpoznanej wartości”
        dla w:jc start/end i bez wyjątku."""
        for case, build, omissions in variants():
            with self.subTest(case):
                self.assertEqual(omissions_of(build()), omissions)
        for editor, (build, _) in EDITORS.items():
            with self.subTest(editor + ', bez list'):
                omissions = ' '.join(omissions_of(build(footnotes=True, lists=False)))
                self.assertNotIn(LISTS_OMISSION, omissions)
                self.assertNotIn(UNKNOWN_FORMAT_OMISSION, omissions)

    def test_rebuild_keeps_text_alignment_and_formatting(self):
        """start → do lewej, end → do prawej, center i both zachowane; nagłówek, pogrubienie i przypis też."""
        for case, build, _ in variants():
            with self.subTest(case):
                rebuilt = Document(rebuild_docx(build(), allow_omissions=True))
                self.assertEqual([p.text for p in rebuilt.paragraphs], texts_of(case))
                self.assertEqual([p.alignment for p in rebuilt.paragraphs[1:5]], ALIGNMENTS)
                self.assertEqual(rebuilt.paragraphs[0].style.name, 'Heading 1')
                runs = rebuilt.paragraphs[5].runs
                self.assertEqual((runs[0].bold, runs[1].italic, runs[2].underline), (True, True, True))
                noted = case.endswith('True')
                self.assertEqual(bool(rebuilt.element.body.xpath('.//w:footnoteReference')), noted)

    def test_preparation_with_rebuild_cleans_the_text(self):
        """Przebudowa + ujednolicanie + miękkie entery + Odkurzacz (ze śledzeniem zmian i bez):
        polskie cudzysłowy, jedna spacja, miękki enter jako osobny akapit, wyrównania zachowane."""
        for case, build, _ in variants():
            for track_changes in (False, True):
                with self.subTest(case, track_changes=track_changes):
                    result = Document(prepare_docx(
                        build(), rebuild=True, normalize_formatting=True, remove_soft_whitespace=True,
                        use_cleaner=True, cleaner_rules=CLEANER_RULES, allow_omissions=True,
                        track_changes=track_changes))
                    self.assertEqual([tracked_text(p) for p in result.paragraphs],
                                     texts_of(case, soft_break_split=True, cleaned=True))
                    self.assertEqual(result.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
                    self.assertEqual(result.paragraphs[3].alignment, WD_ALIGN_PARAGRAPH.CENTER)
                    self.assertEqual(bool(result.element.body.xpath('.//w:ins')), track_changes)

    def test_cleaner_alone_works_on_the_original_file(self):
        """Odkurzacz bez przebudowy na oryginalnym zapisie edytora (także w:jc start/end) – ze śledzeniem
        zmian i bez; przypis też jest odkurzany."""
        for case, build, _ in variants():
            for track_changes in (False, True):
                with self.subTest(case, track_changes=track_changes):
                    output = prepare_docx(build(), use_cleaner=True, cleaner_rules=CLEANER_RULES,
                                          track_changes=track_changes)
                    result = Document(output)
                    self.assertEqual([tracked_text(p) for p in result.paragraphs], texts_of(case, cleaned=True))
                    if case.endswith('True'):
                        notes = etree.fromstring(part(output, 'word/footnotes.xml').encode('utf-8'))
                        self.assertIn('Treść „przypisu”.', ''.join(notes.xpath('.//w:t/text()', namespaces=NS)))
                    if case.startswith('libreoffice_strict'):
                        self.assertIn('<w:jc w:val="start"/>', part(output, 'word/document.xml'))

    def test_soft_breaks_become_paragraphs(self):
        for case, build, _ in variants(footnotes=(False,)):
            with self.subTest(case):
                result = Document(normalize_spacing_docx(build()))
                self.assertEqual([p.text for p in result.paragraphs], texts_of(case, soft_break_split=True))

    def test_split_and_merge_keep_the_footnotes(self):
        """Podział pliku z przypisem daje opowiadanie z tym samym przypisem (otwiera się w python-docx);
        przebudowane części scalają się, a wynik ma oba przypisy."""
        for case, build, _ in variants(footnotes=(True,)):
            with self.subTest(case):
                archive = BytesIO()
                report = split_document(build(), archive)
                self.assertEqual(report['titles'], ['Nagłówek 1'])
                with ZipFile(archive) as parts:
                    story = parts.read(report['files'][0])
                self.assertEqual([p.text for p in Document(BytesIO(story)).paragraphs], texts_of(case))
                self.assertEqual(part(BytesIO(story), 'word/footnotes.xml').count(NOTE_TEXT), 1)
                rebuilt = rebuild_docx(BytesIO(story), allow_omissions=True).getvalue()
                merged = merge_documents([rebuilt, rebuilt])
                result = Document(BytesIO(merged))
                self.assertEqual([p.text for p in result.paragraphs], texts_of(case) * 2)
                self.assertEqual(part(BytesIO(merged), 'word/footnotes.xml').count(NOTE_TEXT), 2)
                self.assertEqual(len(result.element.body.xpath('.//w:footnoteReference')), 2)

    def test_margin_content_and_text_boxes_are_reported_and_dropped(self):
        """Nagłówek, stopka i pole tekstowe LibreOffice: przegląd je zgłasza, przebudowa zostawia sam tekst."""
        for variant in ('transitional', 'strict'):
            with self.subTest(variant):
                stream = libreoffice_docx(variant, header_footer=True, text_box=True, lists=False, footnotes=True)
                self.assertEqual(omissions_of(stream), [f'{TEXT_BOX_OMISSION} (liczba: 1)',
                                                        f'{MARGIN_OMISSION} (liczba: 2)'])
                stream.seek(0)
                output = rebuild_docx(stream, allow_omissions=True)
                self.assertEqual([p.text for p in Document(output).paragraphs], expected_texts(lists=False))
                self.assertNotIn(TEXT_BOX, part(output, 'word/document.xml'))
                self.assertNotIn(HEADER_TEXT, ''.join(part(output, name) for name in ZipFile(output).namelist()
                                                      if name.endswith('.xml')))


@tag('slow', 'docx')
class EditorConversionTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Konwersja do EPUB/PDF (w tym procesie i przez proces roboczy) oraz analiza powtórzeń z modelem."""

    in_process_worker = False  # test_full_conversion_through_the_worker_process: prawdziwy podproces

    def test_worker_converts_every_variant_to_epub_and_pdf(self):
        """Regresja: PDF/EPUB bez przebudowy padały na w:jc „start” (ValueError python-docx) – teraz
        wyrównania trafiają do klas HTML, a PDF powstaje."""
        for case, build, _ in variants(footnotes=(True,)):
            with self.subTest(case), TemporaryDirectory() as folder:
                directory = Path(folder)
                (directory / 'source.docx').write_bytes(build().getvalue())
                convert(directory / 'source.docx', directory, ['epub', 'pdf'], 'Test')
                with ZipFile(directory / 'document.epub') as epub:
                    content = epub.read('EPUB/content.xhtml').decode('utf-8')
                for marker in ('<h1', 'class="align-right"', 'class="align-center"', 'class="align-justify"',
                               ALIGNED['right'], 'Treść'):
                    self.assertIn(marker, content)
                self.assertGreater((directory / 'document.pdf').stat().st_size, 1000)

    def test_full_conversion_through_the_worker_process(self):
        """convert_document (osobny proces): przebudowa, Odkurzacz, miękkie entery, EPUB, PDF i DOCX."""
        for case, build, _ in variants(footnotes=(True,)):
            with self.subTest(case):
                upload = SimpleUploadedFile('tekst.docx', build().getvalue())
                output, extension, _ = convert_document(
                    upload, ['epub', 'pdf'], include_docx=True, rebuild=True, allow_rebuild_omissions=True,
                    use_cleaner=True, cleaner_rules=CLEANER_RULES, remove_soft_whitespace=True)
                with output, ZipFile(output) as archive:
                    self.assertEqual(extension, 'zip')
                    names = sorted(Path(name).suffix for name in archive.namelist())
                    self.assertEqual(names, ['.docx', '.epub', '.pdf'])
                    docx = next(name for name in archive.namelist() if name.endswith('.docx'))
                    result = Document(BytesIO(archive.read(docx)))
                self.assertEqual([p.text for p in result.paragraphs],
                                 texts_of(case, soft_break_split=True, cleaned=True))

    def test_repetitions_colour_every_variant(self):
        """Analiza powtórzeń (prawdziwy model spaCy) na oryginalnym zapisie edytora nie zmienia tekstu."""
        for case, build, _ in variants(footnotes=(True,)):
            with self.subTest(case):
                result = Document(color_document(build(), min_word_length=3))
                self.assertEqual([p.text for p in result.paragraphs], texts_of(case))
                self.assertTrue(result.element.body.xpath('.//w:highlight | .//w:color[@w:val!="000000"]'))


# --- Scenariusze skrajne (editor_fixtures_exotic) ----------------------------------------------------------

W = '{%s}' % NS['w']
SKIPPED = {W + 'del', W + 'moveFrom', '{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice'}
NOT_NOTES = ('separator', 'continuationSeparator', 'continuationNotice')
EXOTIC = {**EDITOR_SCENARIOS, **FEATURE_SCENARIOS}
# Pliki odrzucane przed pracą programów: (komunikat formularza albo None, komunikat konwertera).
FORMAT_ERRORS = {
    'strict': 'format_strict', 'docm': 'format_template', 'dotx': 'format_template', 'dotm': 'format_template',
    'foreign_doc': 'format_doc', 'foreign_encrypted': 'format_encrypted', 'foreign_rtf': 'format_rtf',
    'foreign_odt': 'format_odt', 'foreign_word2003': 'format_word_xml', 'foreign_html': 'format_html',
}
REFUSED = {
    **{name: (MESSAGES[code], MESSAGES[code]) for name, code in FORMAT_ERRORS.items()},
    **{f'doctype_{kind}': (None, 'Nieobsługiwany dokument DOCX.')
       for kind in ('internal', 'external', 'expansion', 'styles')},
    'package_backslash_target': (None, 'Dokument zawiera nieprawidłowe odwołanie do pliku.'),
    'package_backslash_entry': (None, 'Nieprawidłowa struktura dokumentu DOCX.'),
    'package_duplicate': (None, 'Dokument zawiera powielone elementy. Zapisz go ponownie w Wordzie.'),
    'package_case_mismatch': (None, DAMAGED_MESSAGE),
    'package_missing_part': (None, DAMAGED_MESSAGE),
}
# Treść, której e-book i PDF nie mają (równania, wykresy, wstawki altChunk, odsyłacz donikąd) – z uwagą.
CONVERSION_NOTES = {
    'math': 'Pominięto równanie (wzór matematyczny)',
    'chart_and_smart_art': 'Pominięto wykres lub diagram SmartArt',
    'alt_chunk': 'Pominięto treść dołączoną z innego pliku',
    'unusual_notes': 'Pominięto odsyłacz do przypisu, którego nie ma w dokumencie.',
}


# Koniec bloku i podział wiersza w HTML rozdzielają słowa (tekst sąsiednich akapitów nie może się skleić).
BLOCK_END = re.compile(r'<(?:br|hr|/(?:p|h[1-6]|li|td|th))\b[^>]*>')


def words(text):
    """Słowa tekstu (litery i cyfry); miękki dywiz nie dzieli słowa."""
    return set(re.findall(r'\w+', text.replace('\u00ad', '')))


def markup_words(markup):
    """Słowa tekstu z HTML/XHTML (e-book albo HTML przekazany do składu PDF)."""
    tree = etree.HTML('<div>' + BLOCK_END.sub(lambda match: ' ' + match[0], markup) + '</div>')
    return words(etree.tostring(tree, method='text', encoding='unicode'))


def hidden(node):
    vanish = node.getparent().find(f'{W}rPr/{W}vanish')
    return vanish is not None and vanish.get(W + 'val', 'true') not in ('0', 'false')


def visible_words(stream):
    """Słowa widocznego tekstu treści i przypisów: bez usuniętych poprawek, ukrytych przebiegów, kopii mc:Choice
    i separatorów przypisów; inicjał (w:dropCap) jest początkiem następnego akapitu."""
    found, carried = set(), ''
    with ZipFile(stream) as archive:
        for name in ('word/document.xml', 'word/footnotes.xml', 'word/endnotes.xml'):
            if name not in archive.namelist():
                continue
            for paragraph in etree.fromstring(archive.read(name)).iter(W + 'p'):
                if any(node.tag in SKIPPED or node.get(W + 'type') in NOT_NOTES for node in paragraph.iterancestors()):
                    continue
                text, carried = carried, ''
                for node in paragraph.iter(W + 't', W + 'tab', W + 'ptab', W + 'br', W + 'cr', W + 'noBreakHyphen'):
                    if (next(node.iterancestors(W + 'p')) is not paragraph or hidden(node)
                            or any(ancestor.tag in SKIPPED for ancestor in node.iterancestors())):
                        continue
                    text += (node.text or '') if node.tag == W + 't' else '-' if 'Hyphen' in node.tag else ' '
                if paragraph.xpath('./w:pPr/w:framePr[@w:dropCap="drop"]', namespaces=NS):
                    carried = text
                    continue
                found |= words(text)
    return found


def prepared(build):
    """Plik po pierwszym kroku procesu roboczego: bez obrazów, w języku polskim."""
    stripped, _ = strip_images(build())
    return force_polish(stripped)


class ExoticFixtureFidelityTests(SimpleTestCase):
    """Generatory scenariuszy skrajnych odtwarzają cechy opisane w docstringach editor_fixtures_exotic."""

    def test_editors_and_features_carry_the_described_markup(self):
        traits = {
            'word365': ('word/document.xml', '<w16se:symEx w16se:font="Segoe UI Emoji"'),
            'word2007': ('word/document.xml', '<w:smartTag w:uri='),
            'wps': ('word/_rels/document.xml.rels', 'Target="../customXml/item1.xml"'),
            'onlyoffice': ('word/footer1.xml', '<w:fldSimple w:instr="PAGE \\* MERGEFORMAT">'),
            'pages': ('word/styles.xml', 'w:styleId="Body"><w:name w:val="Body"/>'),
            'google_docs_extras': ('word/document.xml', 'w:name="id.x1y2z3"'),
            'minimal': ('[Content_Types].xml', '/word/document.xml'),
            'unusual_notes': ('word/footnotes.xml', '<w:footnote w:type="separator" w:id="0">'),
            'frames_and_drop_cap': ('word/document.xml', 'w:dropCap="drop"'),
            'image_svg': ('word/document.xml', '<asvg:svgBlip r:embed="rIdSvg"/>'),
        }
        for name, (part_name, marker) in traits.items():
            with self.subTest(name):
                self.assertIn(marker, part(EXOTIC[name](), part_name))
        with ZipFile(EXOTIC['minimal']()) as archive:
            self.assertEqual(sorted(archive.namelist()), ['[Content_Types].xml', '_rels/.rels',
                                                          'word/_rels/document.xml.rels', 'word/document.xml'])
        self.assertIn('xmlns:w="http://purl.oclc.org/ooxml/wordprocessingml/main"',
                      part(REFUSED_SCENARIOS['strict'](), 'word/document.xml'))

    def test_foreign_files_have_the_signatures_of_the_real_samples(self):
        """Sygnatury jak w plikach z LibreOffice 24.2: OLE (z EncryptedPackage dla hasła), RTF, ODT, XML 2003."""
        self.assertTrue(foreign_file('doc').startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
        self.assertIn('WordDocument'.encode('utf-16-le'), foreign_file('doc'))
        self.assertIn('EncryptedPackage'.encode('utf-16-le'), foreign_file('encrypted'))
        self.assertTrue(foreign_file('rtf').startswith(b'{\\rtf1'))
        with ZipFile(BytesIO(foreign_file('odt'))) as archive:
            self.assertEqual(archive.namelist()[0], 'mimetype')
        self.assertIn(b'<?mso-application progid="Word.Document"?><w:wordDocument', foreign_file('word2003'))


@tag('slow', 'docx')
class ExoticChainTests(SimpleTestCase):
    """Przegląd, przebudowa, Odkurzacz i miękkie entery na każdym scenariuszu – w tym procesie."""

    def test_rebuild_never_fails_and_never_drops_text_silently(self):
        """Regresja: dokument bez w:sectPr (minimalny eksport) przerywał przegląd i przebudowę (IndexError).
        Utracony tekst zawsze ma pozycję w liście pominięć."""
        for name, build in EXOTIC.items():
            for keep_page_format in (False, True):
                with self.subTest(name, keep_page_format=keep_page_format):
                    source = prepared(build)
                    try:
                        inspect_docx(source, keep_page_format=keep_page_format)
                        omissions = []
                    except RebuildUnsupported as error:
                        omissions = error.omissions
                    rebuilt = rebuild_docx(source, allow_omissions=True, keep_page_format=keep_page_format)
                    Document(rebuilt)
                    lost = visible_words(source) - visible_words(rebuilt)
                    if lost:
                        self.assertTrue(omissions, f'utracono bez uwagi: {sorted(lost)}')

    def test_cleaner_and_soft_breaks_keep_the_text_of_every_scenario(self):
        """Odkurzacz (cudzysłowy, spacje) i zamiana miękkich enterów na oryginalnym zapisie każdego edytora."""
        for name, build in EXOTIC.items():
            with self.subTest(name):
                source = prepared(build)
                output = prepare_docx(source, use_cleaner=True, cleaner_rules=CLEANER_RULES,
                                      remove_soft_whitespace=True)
                Document(output)
                self.assertEqual(visible_words(output), visible_words(source))


class ExoticRefusedFileTests(SimpleTestCase):
    """Pliki innych formatów, nieobsługiwane odmiany DOCX i niebezpieczne pakiety: czytelny komunikat
    formularza (gdy rozpoznaje go już formularz) i konwertera, zanim plik trafi do programów."""

    def form_error(self, data):
        try:
            validate_docx_upload(SimpleUploadedFile('tekst.docx', data))
        except forms.ValidationError as error:
            return ' '.join(error.messages)
        return None

    def test_refused_files_get_a_clear_message(self):
        """Regresja: .doc, ODT, RTF, DOCX z hasłem, Strict, szablon i dokument z makrami kończyły się ogólnym
        „Nie można odczytać pliku” albo błędem procesu (ValueError, KeyError); DOCTYPE poza .rels i brak części
        pakietu – błędem procesu."""
        self.assertEqual(set(REFUSED), set(REFUSED_SCENARIOS))
        for name, (form_message, message) in REFUSED.items():
            with self.subTest(name):
                data = REFUSED_SCENARIOS[name]().getvalue()
                self.assertEqual(self.form_error(data), form_message)
                with self.assertRaises(ConversionError) as raised:
                    validate_resources(BytesIO(data))
                self.assertEqual(str(raised.exception), message)

    def test_other_scenarios_are_accepted(self):
        """Pozostałe scenariusze (także brak pliku obrazu, który i tak zostanie usunięty) przechodzą sprawdzenie."""
        for name, build in EXOTIC.items():
            with self.subTest(name):
                data = build().getvalue()
                validate_resources(BytesIO(data))
                # Część główna o innej nazwie (word/document2.xml) – formularz wymaga word/document.xml.
                if name != 'package_renamed_main':
                    self.assertIsNone(self.form_error(data))


@tag('slow', 'docx')
class ExoticConversionTests(SimpleTestCase):
    """Proces roboczy (main w tym procesie, jak przy konwersji z formularza): EPUB i PDF z każdego scenariusza."""

    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        progress = patch.object(document_progress, '_target', None)
        progress.start()
        self.addCleanup(progress.stop)

    def convert(self, source, **config):
        """(słowa EPUB, słowa HTML przekazanego do składu PDF, uwagi) albo błąd testu; config – dodatkowe opcje
        zadania (np. przebudowa)."""
        for leftover in self.directory.iterdir():
            leftover.unlink()
        (self.directory / 'source.docx').write_bytes(source.getvalue())
        (self.directory / 'job.json').write_text(json.dumps(
            {'formats': ['epub', 'pdf'], 'prepare': True, 'normalize': True, 'title': 'Test', **config}),
            encoding='utf-8')
        captured = []
        with patch('sys.argv', ['worker', str(self.directory)]), pdf_html_spy(captured):
            code = main()
        error = self.directory / 'error.json'
        self.assertEqual(code, 0, error.read_text(encoding='utf-8') if error.exists() else '')
        with ZipFile(self.directory / 'document.epub') as epub:
            chapters = ' '.join(epub.read(name).decode('utf-8') for name in epub.namelist()
                                if name.endswith('.xhtml') and 'nav' not in name)
        notes = self.directory / 'warnings.json'
        return (markup_words(chapters), markup_words(' '.join(content for content, _ in captured)),
                json.loads(notes.read_text(encoding='utf-8')) if notes.exists() else [])

    def test_every_scenario_keeps_its_visible_text(self):
        """Regresja: e-book i PDF gubiły wynik w:fldSimple (numer podpisu, pole korespondencji) i tekst w w:dir/
        w:bdo, sklejały słowa wokół w:cr i w:ptab, a odsyłacz do nieistniejącego przypisu przerywał konwersję."""
        for name, build in EXOTIC.items():
            with self.subTest(name):
                expected = visible_words(prepared(build))
                epub, pdf, notes = self.convert(build())
                self.assertEqual(expected - epub, set(), 'EPUB')
                self.assertEqual(expected - pdf, set(), 'PDF')
                if name in CONVERSION_NOTES:
                    self.assertTrue(any(note.startswith(CONVERSION_NOTES[name]) for note in notes), notes)

    def test_review_markup_stays_out_of_ebook_and_pdf(self):
        """Plik ze zmianami i komentarzami (Word 365, LibreOffice), bez przebudowy i z nią: konwersja się udaje,
        e-book i PDF mają tekst po przyjęciu zmian – bez usuniętego i przeniesionego skąd tekstu oraz bez treści
        komentarzy, ze wstawkami."""
        deleted = {'usunięty', 'cofnięty', 'Cały'}
        comments = {'Uwaga', 'Pytanie', 'Odpowiedź', 'Rozwiązane', 'punktowy'}
        for name in ('review', 'libreoffice_review'):
            for rebuild in (False, True):
                with self.subTest(name, rebuild=rebuild):
                    epub, pdf, _ = self.convert(EXOTIC[name](), rebuild=rebuild, allow_rebuild_omissions=rebuild)
                    for words in (epub, pdf):
                        self.assertEqual(words & (deleted | comments), set())
                        self.assertLessEqual({'wstawiony', 'Przeniesiony', 'Koniec'}, words)
