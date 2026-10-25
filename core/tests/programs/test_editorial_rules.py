"""Reguły korekty Odkurzacza na poziomie tekstu oraz pola, limity i odrzucane pakiety formularzy programów."""
import json
import sys
import time
from difflib import SequenceMatcher
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from django.test import SimpleTestCase
from docx import Document

from core.odkurzacz_forms import (
    REPLACEMENT_WORD_MAX, WORD_SOURCE_RE, DocumentConversionForm, OdkurzaczForm, RepetitionsForm,
)
from core.services.odkurzacz import (
    ABBREVIATION_GROUPS, ALL_EDITORIAL_RULES, CONTEXT, DEFAULT_EDITORIAL_RULES, DEFAULT_REPLACEMENTS,
    EDITORIAL_RULES, FORM_DEFAULT_EDITORIAL_RULES, MAX_PATTERN_LENGTH, ORTHOGRAPHY_2026_KEYS,
    ORTHOGRAPHY_2026_RULES, SENTENCE_ABBREVIATIONS, USER_WORD_EXACT, USER_WORD_PATTERNS, _change_regions,
    _compile_pattern, _edit_opcodes, clean_docx, compile_replacements, correct_editorial_text,
)

from .support import formatted_docx, saved, upload


def package(entries, encrypted=False):
    """Archiwum ZIP z wpisami (nazwa, bajty); encrypted – bit szyfrowania w katalogu centralnym (jak ZipCrypto)."""
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        for name, data in entries:
            archive.writestr(ZipInfo(name), data)
    data = bytearray(stream.getvalue())
    if encrypted:
        central = data.rindex(b'PK\x01\x02')  # ostatni wpis katalogu centralnego
        data[central + 8] |= 1
    return bytes(data)


class EditorialTextRuleTests(SimpleTestCase):
    def fix(self, text):
        return correct_editorial_text(text, DEFAULT_EDITORIAL_RULES)

    def test_protected_abbreviations_keep_internal_spacing(self):
        """Regresja: skróty typu „o.o.” były rozbijane albo dostawały odstęp jak po zdaniu."""
        for abbreviation in ('ś.p.', 'c.d.n.', 'o.o.', 'j.w.'):
            with self.subTest(abbreviation=abbreviation):
                self.assertIn(abbreviation, self.fix(f'To {abbreviation} przykład.'))
        self.assertEqual(correct_editorial_text('To o.o. przykład.', ['after_punct']), 'To o.o. przykład.')

    def test_period_and_comma_always_follow_the_closing_quote(self):
        """Polska interpunkcja: kropka za cudzysłowem; gdy zdanie trwa dalej, kropka znika."""
        cases = {
            'Napisał „tak.” i wyszedł.': 'Napisał „tak” i wyszedł.',
            'Powiedział „tak.”': 'Powiedział „tak”.',
            'Powiedział „tak.”.': 'Powiedział „tak”.',
            'Napisał „tak.” Potem wyszedł.': 'Napisał „tak”. Potem wyszedł.',
            '„Tak.” – powiedział.': '„Tak” – powiedział.',
            'Powiedział: „Tak. Nie. Może.” I wyszedł.': 'Powiedział: „Tak. Nie. Może”. I wyszedł.',
            'Mówił: „Idź!”. Potem „co?” i „hm…”.': 'Mówił: „Idź!”. Potem „co?” i „hm…”.',
            'Słowo „tak,” i „nie;” oraz „może:” tak.': 'Słowo „tak”, i „nie”; oraz „może”: tak.',
            '„Cytat «w środku.» koniec.”': '„Cytat «w środku» koniec”.',
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.fix(text), expected)

    def test_letter_marks_after_numbers_are_not_units(self):
        self.assertEqual(self.fix('Klasy 2B i 3A wyszły.'), 'Klasy 2B i 3A wyszły.')
        self.assertEqual(self.fix('Telewizor 4K.'), 'Telewizor 4K.')
        self.assertEqual(self.fix('Kabel 2m, 5kg i 20 K.'), 'Kabel 2 m, 5 kg i 20 K.')

    def test_gender_forms_in_brackets_stay_joined(self):
        self.assertEqual(self.fix('Drogi(a) Kliencie, zrobił(a) to.'), 'Drogi(a) Kliencie, zrobił(a) to.')
        self.assertEqual(self.fix('Pies(kundel) szczekał.'), 'Pies (kundel) szczekał.')

    def test_score_keeps_spaces_around_colon(self):
        self.assertEqual(self.fix('Wynik 3 : 1 , a potem :'), 'Wynik 3 : 1, a potem:')

    def test_partial_fragment_is_not_a_new_dialogue(self):
        rules = ['hyphen_dash', 'dash_spaces', 'trim']
        self.assertEqual(correct_editorial_text(' – narracja', rules, trim_start=False), ' – narracja')
        self.assertEqual(correct_editorial_text(' - narracja', rules, trim_start=False), ' – narracja')
        self.assertEqual(correct_editorial_text('tekst – ', rules, trim_end=False), 'tekst – ')
        self.assertEqual(correct_editorial_text('  –Hę?', rules), '– Hę?')

    def test_dash_context_looks_past_a_long_run_of_spaces(self):
        """Sąsiedztwo myślnika: gdy w oknie kontekstu są same spacje, decyduje najbliższy znak poza oknem,
        więc fragment bez przycinania nie jest brany za początek dialogu."""
        gap = ' ' * (CONTEXT + 5)
        for text, expected in (
            (f'a{gap}–b', 'a – b'),
            (f'a–{gap}b', 'a – b'),
            (f'{gap}–b', f'{gap}–b'),
        ):
            with self.subTest(text=text.strip()):
                self.assertEqual(correct_editorial_text(text, ['dash_spaces'], trim_start=False, trim_end=False), expected)

    def test_typographic_rules_are_available_and_can_be_disabled(self):
        added = {'unit_nbsp', 'reference_nbsp', 'dates_times', 'decimal', 'digit_groups'}
        self.assertTrue(added <= {key for key, _ in EDITORIAL_RULES})
        source = '2025-1-2 2/1/2025 o 9.30 9:30 12.50 kg 123456 kg A.Kowalski 20°C'
        self.assertEqual(correct_editorial_text(source, []), source)
        self.assertEqual(correct_editorial_text(source),
                         '02.01.2025 02.01.2025 o 09:30 09:30 12,50 kg 123 456 kg A. Kowalski 20 °C')

    def test_defaults_and_explicit_empty_paragraph_rule(self):
        self.assertEqual(correct_editorial_text('kot\tkota', {'tabs'}), 'kot kota')
        self.assertEqual(correct_editorial_text('1.02.2025', {'dates_times'}), '01.02.2025')
        self.assertNotIn('empty_paragraphs', DEFAULT_EDITORIAL_RULES)
        self.assertNotIn('user_word_corrections', DEFAULT_EDITORIAL_RULES)
        self.assertEqual(
            correct_editorial_text('np. kot. pies! – powiedział.', {'sentence_case'}),
            'np. kot. Pies! – powiedział.',
        )
        document = Document()
        document.add_paragraph('A')
        document.add_paragraph()
        document.add_paragraph()
        document.add_paragraph('B')
        payload = saved(document).getvalue()
        marked = Document(clean_docx(BytesIO(payload), ['empty_paragraphs']))
        self.assertEqual([p.text for p in marked.paragraphs], ['A', '', 'B'])
        marked = Document(clean_docx(BytesIO(payload), DEFAULT_EDITORIAL_RULES))
        self.assertEqual(len(marked.paragraphs), 4)

    def test_edit_opcodes_handle_the_longest_paragraph_quickly(self):
        source = ('Ala  ma kota a kot ma Alę... ' * 700)[:20000]
        corrected = self.fix(source)
        started = time.monotonic()
        opcodes = _edit_opcodes(source, corrected)
        self.assertLess(time.monotonic() - started, 2)
        rebuilt = ''.join(source[a:b] if tag == 'equal' else corrected[c:d]
                          for tag, a, b, c, d in opcodes if tag != 'delete')
        self.assertEqual(rebuilt, corrected)


class OptionalAndEdgeRuleTests(SimpleTestCase):
    """Reguły spoza domyślnego zestawu oraz przypadki, które reguły celowo zostawiają bez zmian."""

    def fix(self, text):
        return correct_editorial_text(text, DEFAULT_EDITORIAL_RULES)

    def test_personal_pronouns_are_lowered_only_inside_sentences(self):
        self.assertEqual(
            correct_editorial_text('Spotkałem Go wczoraj. On przyszedł. Dałem Mu to: Ona wie. '
                                   'Mówił np. Ona przyjdzie! Ja też. Widziałem JEGO.', {'pronouns_lower'}),
            'Spotkałem go wczoraj. On przyszedł. Dałem mu to: Ona wie. '
            'Mówił np. ona przyjdzie! Ja też. Widziałem JEGO.')
        self.assertEqual(correct_editorial_text('Pisał – Ona wie, https://example.com/On', {'pronouns_lower'}),
                         'Pisał – ona wie, https://example.com/On')

    def test_user_word_list_keeps_the_original_letter_case(self):
        self.assertEqual(
            correct_editorial_text('Mentlikiem jechał niechcąco, MENTLIKAMI; europejczykiem i azjatami; '
                                   'komedia dellarte; Przekonywujący oddziaływującego', {'user_word_corrections'}),
            'Mętlikiem jechał niechcący, MĘTLIKAMI; Europejczykiem i Azjatami; '
            'komedia dell’arte; Przekonujący oddziałującego')

    def test_impossible_dates_and_times_stay_unchanged(self):
        self.assertEqual(correct_editorial_text('31.02.2025, 2025-13-01, godz. 25:00, 24:61 i o 9.30',
                                                {'dates_times'}),
                         '31.02.2025, 2025-13-01, godz. 25:00, 24:61 i o 09:30')

    def test_sentence_case_skips_ellipsis_ordinals_initials_and_file_names(self):
        self.assertEqual(correct_editorial_text('Koniec... dalej. 3. punkt. A. kowalski wrócił. Patrz np. tutaj. '
                                                'dalej', {'sentence_case'}),
                         'Koniec... dalej. 3. punkt. A. kowalski wrócił. Patrz np. tutaj. Dalej')

    def test_sentence_case_does_not_change_addresses_starting_a_sentence(self):
        """Regresja: adres na początku zdania nie dostaje wielkiej litery („Https://…”, „Jan@…”)."""
        for text in ('Zobacz to. https://example.com/x', 'Koniec! jan@example.com napisał'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, {'sentence_case'}), text)

    def test_dashes_between_numbers_and_minus_signs_keep_their_spacing(self):
        self.assertEqual(correct_editorial_text('10–20 i (–5) oraz =–3, a dalej–tak', {'dash_spaces'}),
                         '10–20 i (–5) oraz =–3, a dalej – tak')

    def test_identifiers_are_not_grouped_like_quantities(self):
        self.assertEqual(
            correct_editorial_text('012345 osób, nr 123456, 123456789 i 123456789 osób, 12345 kg',
                                   {'digit_groups'}),
            '012345 osób, nr 123456, 123456789 i 123 456 789 osób, 12 345 kg')

    def test_technical_text_is_protected_from_default_rules(self):
        source = 'Adres: www.example.pl/strona?a=1. Napisz: jan@example.com, wersja v1.2.3, IP 192.168.0.1'
        self.assertEqual(self.fix(source), source)
        # Znak z obszaru prywatnego Unicode w tekście nie myli masek chroniących adresy.
        marked = '\U000f0000 www.example.pl'
        self.assertEqual(self.fix(marked), marked)

    def test_unknown_rule_is_refused(self):
        with self.assertRaisesMessage(ValueError, 'Nieznane reguły korekty'):
            correct_editorial_text('x', {'nieznana'})

    def test_edit_opcodes_fall_back_to_difflib(self):
        with patch.dict(sys.modules, {'rapidfuzz.distance': None}):
            opcodes = _edit_opcodes('Ala  ma', 'Ala ma')
        self.assertEqual(opcodes, SequenceMatcher(None, 'Ala  ma', 'Ala ma', autojunk=False).get_opcodes())


class ReplacementListTests(SimpleTestCase):
    """Słownik zamian jako dane: DEFAULT_REPLACEMENTS i własna lista `replacements`."""

    def fix(self, text, replacements=None):
        return correct_editorial_text(text, {'user_word_corrections'}, replacements=replacements)

    def test_default_list_is_data_equivalent_to_the_builtin_list(self):
        self.assertTrue(all(set(entry) == {'kind', 'source', 'target'} and entry['kind'] in ('word', 'pattern')
                            for entry in DEFAULT_REPLACEMENTS))
        self.assertEqual(len(DEFAULT_REPLACEMENTS), len(USER_WORD_EXACT) + len(USER_WORD_PATTERNS) + 1)
        json.dumps(DEFAULT_REPLACEMENTS)  # nadaje się do zapisania w CMS / konfiguracji zadania
        samples = ('Mentlikiem jechał niechcąco, MENTLIKAMI; europejczykiem, Europejczykiem i EUROPEJCZYKIEM; '
                   'azjatami; komedia dellarte, KOMEDIA DELLARTE, KoMeDiA dEllArte, komediami Dellarte; '
                   'ZACZYM, Zaczym; MASSMEDIAMI; miałczeniem; nieprzekonywującego; wymyśleć; kanister')
        expected = ('Mętlikiem jechał niechcący, MĘTLIKAMI; Europejczykiem, Europejczykiem i EUROPEJCZYKIEM; '
                    'Azjatami; komedia dell’arte, KOMEDIA DELL’ARTE, KoMeDiA dell’arte, komediami Dell’arte; '
                    'ZA CZYM, Za czym; MASS MEDIAMI; miauczeniem; nieprzekonującego; wymyślić; kanister')
        self.assertEqual(self.fix(samples), expected)
        self.assertEqual(self.fix(samples, DEFAULT_REPLACEMENTS), expected)

    def test_custom_list_replaces_the_builtin_one(self):
        replacements = [{'kind': 'word', 'source': 'kot', 'target': 'pies'},
                        {'kind': 'word', 'source': 'nie wiem', 'target': 'nie mam pojęcia'}]
        self.assertEqual(self.fix('Kot, KOT i kot; kotek; mentlik. Nie wiem.', replacements),
                         'Pies, PIES i pies; kotek; mentlik. Nie mam pojęcia.')
        self.assertEqual(self.fix('mentlik', []), 'mentlik')

    def test_patterns_use_named_groups_and_capital_targets(self):
        replacements = [
            {'kind': 'pattern', 'source': r'\bwarszawiak(?P<ending>iem|owi|a|u)?\b', 'target': r'Warszawiak\g<ending>'},
            {'kind': 'pattern', 'source': r'\bwziąść\b', 'target': 'wziąć'},
            {'kind': 'word', 'source': 'polak', 'target': 'Polak'},
        ]
        self.assertEqual(self.fix('Z warszawiakiem, WARSZAWIAKOWI; Wziąść i WZIĄŚĆ; polak', replacements),
                         'Z Warszawiakiem, WARSZAWIAKOWI; Wziąć i WZIĄĆ; Polak')

    def test_invalid_and_costly_entries_are_skipped_without_stopping_the_cleaner(self):
        good = {'kind': 'word', 'source': 'kot', 'target': 'pies'}
        broken = (
            {'kind': 'pattern', 'source': '(kot', 'target': 'x'},                  # błąd kompilacji
            {'kind': 'pattern', 'source': r'\bkot\b', 'target': r'\g<brak>'},      # nieznana grupa
            {'kind': 'pattern', 'source': r'\bkot\b', 'target': '\\q'},            # zły znak ucieczki
            {'kind': 'pattern', 'source': r'(a+)+$', 'target': 'x'},               # zagnieżdżone powtórzenie
            {'kind': 'pattern', 'source': r'(?:ab|a)*c', 'target': 'x'},           # alternatywa w powtórzeniu
            {'kind': 'pattern', 'source': r'x*', 'target': 'y'},                   # pasuje do pustego napisu
            {'kind': 'pattern', 'source': 'a' * 501, 'target': 'x'},               # za długi wzorzec
            {'kind': 'word', 'source': 'ala', 'target': 'x' * 201},                # za długi target
            {'kind': 'regex', 'source': 'ala', 'target': 'x'},                     # nieznany rodzaj
            {'kind': 'word', 'source': '', 'target': 'x'},
            {'kind': 'word', 'source': 'ala', 'target': None},
            'ala=x',
        )
        text = 'Ala ma kota, kot ' + 'a' * 30 + '!'
        start = time.perf_counter()
        for entry in broken:
            with self.subTest(entry=str(entry)[:40]):
                self.assertEqual(self.fix(text, [entry, good]), text.replace('kot ', 'pies '))
        self.assertLess(time.perf_counter() - start, 2)
        # Bezpieczne powtórzenia i opcjonalne alternatywy są dozwolone.
        self.assertEqual(self.fix('kotów kocur', [
            {'kind': 'pattern', 'source': r'\bkot(?P<e>ów|a)?\b', 'target': r'pies\g<e>'},
            {'kind': 'pattern', 'source': r'\bkoc\w+', 'target': 'pies'}]), 'piesów pies')

    def test_replacements_reach_the_docx_cleaner(self):
        document = Document()
        document.add_paragraph('Mentlik i kot.')
        result = Document(clean_docx(saved(document), ['user_word_corrections'],
                                     replacements=[{'kind': 'word', 'source': 'kot', 'target': 'pies'}]))
        self.assertEqual(result.paragraphs[0].text, 'Mentlik i pies.')
        result = Document(clean_docx(saved(document), ['user_word_corrections']))
        self.assertEqual(result.paragraphs[0].text, 'Mętlik i kot.')


class ChangeRegionTests(SimpleTestCase):
    """Granularność zmian (śledzenie zmian i podsumowanie): cały wyraz albo sama interpunkcja/odstęp."""

    def regions(self, original, corrected):
        return [(original[a:b], corrected[c:d]) for a, b, c, d in _change_regions(original, corrected)]

    def test_word_changes_cover_the_whole_word_and_spacing_changes_stay_minimal(self):
        self.assertEqual(self.regions('Był w mentliku.', 'Był w mętliku.'), [('mentliku', 'mętliku')])
        self.assertEqual(self.regions('Ala  ma kota...', 'Ala ma kota…'), [(' ', ''), ('...', '…')])
        self.assertEqual(self.regions('w 2025r. i a b', 'w 2025 r. i a b'), [('', ' '), (' ', ' ')])
        self.assertEqual(self.regions('1.2.2025', '01.02.2025'), [('1', '01'), ('2', '02')])
        self.assertEqual(self.regions('bez zmian', 'bez zmian'), [])

    def test_summary_counting_does_not_change_the_text(self):
        for text in ('Ala  ma kota... „tak” – rzekł. On  przyszedł,a potem 1.2.2025 r. 5-7 kg.', '   ', 'a'):
            counts = {}
            self.assertEqual(correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES, counts=counts),
                             correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES))


class EditorialRuleRegressionTests(SimpleTestCase):
    """Regresje znalezione w przeglądzie Odkurzacza (październik 2026)."""

    def fix(self, text, rules=FORM_DEFAULT_EDITORIAL_RULES):
        return correct_editorial_text(text, rules)

    def test_pronouns_in_capitals_stay_in_titles_shouting_and_emphasis(self):
        """Wersaliki to napis, krzyk albo wyróżnienie (decyzja właściciela) – zostają."""
        rules = {'pronouns_lower'}
        for text in ('ROZDZIAŁ PIERWSZY. ON I ONA', 'Krzyknął: NIE CHCĘ GO WIDZIEĆ!', 'Powiedziałem NIE.',
                     'Widziałem JEGO, nie ją.', 'Daj TOBIE to.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, rules), text)

    def test_straight_quote_inside_polish_quote_becomes_second_level(self):
        self.assertEqual(self.fix('„Powiedział: "nie"” – zaśmiał się.', {'quotes'}),
                         '„Powiedział: «nie»” – zaśmiał się.')
        self.assertEqual(self.fix('"Tak" i "nie".', {'quotes'}), '„Tak” i „nie”.')

    def test_postal_code_is_not_a_number_range(self):
        self.assertEqual(self.fix('Ul. Polna 3, 00-950 Warszawa.', {'range_dash', 'range_spaces'}),
                         'Ul. Polna 3, 00-950 Warszawa.')
        self.assertEqual(self.fix('Od 10-200 osób.', {'range_dash'}), 'Od 10–200 osób.')

    def test_missing_space_before_capitalized_word_is_not_a_domain(self):
        """Regresja: „domu.Co”, „kinie.De” wyglądały jak adresy .co/.de i nie dostawały spacji."""
        for text, expected in (('Wrócił do domu.Co dalej?', 'Wrócił do domu. Co dalej?'),
                               ('Był w kinie.De facto nic.', 'Był w kinie. De facto nic.')):
            with self.subTest(text=text):
                self.assertEqual(self.fix(text, {'after_punct'}), expected)
        self.assertEqual(self.fix('Strona fantazmaty.pl i onet.com.pl.', {'after_punct', 'sentence_case'}),
                         'Strona fantazmaty.pl i onet.com.pl.')

    def test_dotted_technical_names_are_protected(self):
        text = 'Zobacz core.views.tags. dalej'
        self.assertEqual(self.fix(text, {'after_punct'}), text)
        self.assertEqual(self.fix(text, {'sentence_case'}), 'Zobacz core.views.tags. Dalej')

    def test_period_after_time_ends_the_sentence(self):
        self.assertEqual(self.fix('Wyszedł o 9:30. potem wrócił.', {'sentence_case'}),
                         'Wyszedł o 9:30. Potem wrócił.')
        self.assertEqual(self.fix('Rozdział 3. punkt 2.', {'sentence_case'}), 'Rozdział 3. punkt 2.')

    def test_ellipsis_between_words_gets_a_space(self):
        self.assertEqual(self.fix('Pół godziny…a potem cisza.', {'after_punct'}), 'Pół godziny… a potem cisza.')
        self.assertEqual(self.fix('…i wtedy, co?…', {'after_punct'}), '…i wtedy, co?…')

    def test_typewriter_double_hyphen_becomes_a_dash(self):
        rules = {'hyphen_dash', 'dash_spaces', 'range_dash'}
        for text, expected in (('Tak--powiedział.', 'Tak – powiedział.'), ('-- Wiem -- rzekł.', '– Wiem – rzekł.'),
                               ('Strony 10--20.', 'Strony 10–20.')):
            with self.subTest(text=text):
                self.assertEqual(self.fix(text, rules), expected)

    def test_typewriter_comma_quotes_become_polish_quotes(self):
        """Regresja: „,,tak''” było psute na „, tak''” przez regułę powtórzonych przecinków."""
        rules = {'quotes', 'duplicate_punct', 'after_punct'}
        self.assertEqual(self.fix("Mówił ,,tak'' i ,,nie\".", rules), 'Mówił „tak” i „nie”.')
        self.assertEqual(self.fix('Liczby a,, b.', rules), 'Liczby a, b.')

    def test_long_paragraph_with_many_numbers_pronouns_and_dashes_is_fast(self):
        """Reguły patrzą tylko na krótki kontekst; wcześniej koszt rósł z kwadratem długości akapitu."""
        source = ('– Nie wiem – powiedział Jan. On go widział nr 12345 i 67890 kg, o 9.30, ' * 300)[:20000]
        started = time.monotonic()
        result = self.fix(source)
        self.assertLess(time.monotonic() - started, 1.5)
        self.assertIn('nr 12345', result)
        self.assertIn('67 890 kg', result)
        self.assertIn('. On go', result)
        self.assertIn('Jan. On go widział', result)


class NonBreakingSpaceDecisionTests(SimpleTestCase):
    """Decyzja właściciela: Odkurzacz nie wstawia spacji nierozdzielających, a istniejące zamienia na zwykłe."""

    def test_no_rule_inserts_nonbreaking_spaces(self):
        source = ('W domu i z nim: 5 kg, 12.50 zł, 2025 r., s. 15, nr 3, 1234567 osób, A. B. Kowalski, '
                  'A.Kowalski, 20 °C i 300 K, o 9.30 u nas.')
        for rules in (ALL_EDITORIAL_RULES, FORM_DEFAULT_EDITORIAL_RULES, DEFAULT_EDITORIAL_RULES,
                      *({key} for key in ALL_EDITORIAL_RULES)):
            with self.subTest(rules=sorted(rules)[:3]):
                result = correct_editorial_text(source, rules)
                self.assertFalse(set(result) & {'\u00a0', '\u202f', '\u2007'}, result)
        self.assertEqual(correct_editorial_text(source, ALL_EDITORIAL_RULES),
                         'W domu i z nim: 5 kg, 12,50 zł, 2025 r., s. 15, nr 3, 1 234 567 osób, A. B. Kowalski, '
                         'A. Kowalski, 20 °C i 300 K, o 09:30 u nas.')

    def test_existing_nonbreaking_spaces_become_plain_spaces(self):
        source = 'Był\u00a0w\u202fdomu\u2007o 5\u00a0kg.'
        for rules in ({'spaces'}, {'quotes'}, ALL_EDITORIAL_RULES):
            with self.subTest(rules=sorted(rules)[:2]):
                self.assertEqual(correct_editorial_text(source, rules), 'Był w domu o 5 kg.')
        self.assertEqual(correct_editorial_text(source, []), source)

    def test_single_letter_rule_is_gone_and_labels_do_not_promise_nonbreaking_spaces(self):
        self.assertNotIn('single_nbsp', ALL_EDITORIAL_RULES)
        for key, label in EDITORIAL_RULES:
            with self.subTest(rule=key):
                self.assertNotIn('nierozdzielaj', label.lower())
        with self.assertRaisesMessage(ValueError, 'Nieznane reguły korekty'):
            correct_editorial_text('a b', {'single_nbsp'})


SPECIAL_SPACES = '\u00a0\u202f\u2009\u2002\u2003\u2007\u2008\u200a\u205f'
SPECIAL_DASHES = '\u2012\u2014\u2015\u2212'
UNICODE_HYPHENS = '\u2010\u2011'
SPECIAL_CHARACTERS = frozenset(SPECIAL_SPACES + SPECIAL_DASHES + UNICODE_HYPHENS + '\u00ad\u200b')


class SpecialCharacterTests(SimpleTestCase):
    """Zadanie 15: spacje specjalne → zwykła spacja, kreski → półpauza, łączniki Unicode → dywiz,
    miękki dywiz i znak o zerowej szerokości → usunięte."""

    def fix(self, text, rules=FORM_DEFAULT_EDITORIAL_RULES):
        return correct_editorial_text(text, rules)

    def test_special_spaces_become_plain_spaces_in_typical_contexts(self):
        """Każda spacja specjalna w środku zdania, przy jednostce, roku, skrócie, inicjale i myślniku."""
        for space in SPECIAL_SPACES:
            source = (f'Ala{space}ma kota, 5{space}kg i 2025{space}r., s.{space}15, J.{space}Kowalski, '
                      f'tak{space}–{space}nie.')
            with self.subTest(space=f'U+{ord(space):04X}'):
                result = self.fix(source)
                self.assertEqual(result, 'Ala ma kota, 5 kg i 2025 r., s. 15, J. Kowalski, tak – nie.')
                self.assertFalse(set(result) & SPECIAL_CHARACTERS)

    def test_special_spaces_feed_the_space_rules(self):
        """Po zamianie działają dalsze reguły: podwójna spacja, spacja przed przecinkiem i na brzegu."""
        for space in SPECIAL_SPACES:
            with self.subTest(space=f'U+{ord(space):04X}'):
                self.assertEqual(self.fix(f'{space}Ala {space}ma{space},kota{space}'), 'Ala ma, kota')

    def test_special_dashes_become_en_dash_in_every_context(self):
        """Pauza cyfrowa, pauza, kreska pozioma i minus: myślnik, dialog, zakres, liczba ujemna, wzór."""
        for dash in SPECIAL_DASHES:
            cases = {
                f'Był tam {dash} i wyszedł.': 'Był tam – i wyszedł.',
                f'Był tam{dash}i wyszedł.': 'Był tam – i wyszedł.',
                f'{dash} Idziemy {dash} powiedział.': '– Idziemy – powiedział.',
                f'Lata 1990{dash}1995 i s. 5 {dash} 7.': 'Lata 1990–1995 i s. 5–7.',
                f'Wynik 5 {dash} 3 = 2.': 'Wynik 5 – 3 = 2.',
                f'Słowo {dash}{dash} słowo.': 'Słowo – słowo.',
            }
            for source, expected in cases.items():
                with self.subTest(dash=f'U+{ord(dash):04X}', source=source):
                    result = self.fix(source)
                    self.assertEqual(result, expected)
                    self.assertFalse(set(result) & SPECIAL_CHARACTERS)

    def test_minus_and_figure_dash_before_number_stay_attached(self):
        """Minus i pauza cyfrowa przed liczbą → półpauza bez odstępu (znak liczby); pauza i kreska
        pozioma to myślnik, więc dostają odstępy jak dotąd półpauza w tym miejscu."""
        for dash in SPECIAL_DASHES:
            minus = dash in '\u2012\u2212'
            cases = {
                f'Mróz {dash}5 °C, ({dash}3) i x = {dash}2.':
                    'Mróz –5 °C, (–3) i x = –2.' if minus else 'Mróz – 5 °C, (–3) i x = –2.',
                f'{dash}12345 kg i spadek o {dash}5%.':
                    '–12 345 kg i spadek o –5%.' if minus else '–12 345 kg i spadek o – 5%.',
                # Ścisły zapis z minusem lub pauzą cyfrową (numer) zostaje bez odstępów.
                f'Tel. 555{dash}1234.': 'Tel. 555–1234.' if minus else 'Tel. 555 – 1234.',
            }
            for source, expected in cases.items():
                with self.subTest(dash=f'U+{ord(dash):04X}', source=source):
                    self.assertEqual(self.fix(source), expected)

    def test_unicode_hyphens_become_plain_hyphen(self):
        """Łącznik Unicode i nierozdzielający → dywiz, także przy liczbach i skrótach."""
        for hyphen in UNICODE_HYPHENS:
            with self.subTest(hyphen=f'U+{ord(hyphen):04X}'):
                self.assertEqual(self.fix(f'Biało{hyphen}czerwony, PRL{hyphen}u i 18{hyphen}latek.'),
                                 'Biało-czerwony, PRL-u i 18-latek.')

    def test_soft_hyphen_and_zero_width_space_are_removed(self):
        """Miękki dywiz i znak o zerowej szerokości znikają (reguła invisible_chars)."""
        for character in '\u00ad\u200b':
            with self.subTest(character=f'U+{ord(character):04X}'):
                self.assertEqual(self.fix(f'Odku{character}rzacz i{character} tak.'), 'Odkurzacz i tak.')

    def test_counts_per_rule(self):
        """Podsumowanie: każda zamieniona kreska liczy się w dash_style, miękki dywiz w invisible_chars."""
        counts = {}
        correct_editorial_text('A \u2014 b, c \u2015 d, \u22125 °C, 1990\u20121995, e\u2010f, g\u00adh.',
                               FORM_DEFAULT_EDITORIAL_RULES, counts=counts)
        self.assertEqual(counts.get('dash_style'), 5)
        self.assertEqual(counts.get('invisible_chars'), 1)

    def test_disabled_dash_rule_keeps_dashes_outside_ranges(self):
        """Bez dash_style kreski zostają; zakres liczb nadal poprawia range_dash (jak dotąd)."""
        rules = FORM_DEFAULT_EDITORIAL_RULES - {'dash_style'}
        self.assertEqual(self.fix('Tak \u2014 nie, \u22125 °C, a\u2010b i 1990\u22121995.', rules),
                         'Tak \u2014 nie, \u22125 °C, a\u2010b i 1990–1995.')
        self.assertEqual(self.fix('Ala\u2009ma\u00adkota.', set()), 'Ala\u2009ma\u00adkota.')


class EditorialRuleR7Tests(SimpleTestCase):
    """Poprawki po przeglądzie r6: zaimki, słownik zamian, skróty, nazwy techniczne, zakresy godzin."""

    def fix(self, text, rules=FORM_DEFAULT_EDITORIAL_RULES, replacements=None):
        return correct_editorial_text(text, rules, replacements=replacements)

    def test_pronoun_right_after_an_opening_quote_keeps_its_capital(self):
        for text in ('Czytał „Ja, robot” Asimova.', 'Film "On i ona" wszedł.', 'Tytuł «Ona» znam.',
                     'Napisał (cytuję) „Ona wie”.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, {'pronouns_lower'}), text)
        self.assertEqual(correct_editorial_text('Spotkałem Go i dałem Mu to.', {'pronouns_lower'}),
                         'Spotkałem go i dałem mu to.')

    def test_honorific_second_person_forms_keep_their_capital(self):
        text = 'Piszę do Ciebie, bo Tobie ufam, z Tobą chcę być, kocham Cię i dam Ci to, a Ty wiesz. Was też, Wam i Wami.'
        self.assertEqual(correct_editorial_text(text, {'pronouns_lower'}), text)

    def test_replacement_target_spelling_is_kept_unless_the_source_is_all_capitals(self):
        entries = [{'kind': 'word', 'source': 'nato', 'target': 'NATO'},
                   {'kind': 'word', 'source': 'iphone', 'target': 'iPhone'},
                   {'kind': 'word', 'source': 'mcdonalds', 'target': 'McDonald’s'},
                   {'kind': 'word', 'source': 'pdf', 'target': 'PDF'},
                   {'kind': 'pattern', 'source': r'\bmac(?P<e>a|iem)?\b', 'target': r'MacBook\g<e>'}]
        self.assertEqual(
            self.fix('W nato i Nato; mam iphone, Iphone i IPHONE; w mcdonalds; plik pdf; z maciem.',
                     {'user_word_corrections'}, entries),
            'W NATO i NATO; mam iPhone, iPhone i IPHONE; w McDonald’s; plik PDF; z MacBookiem.')

    def test_replacement_words_may_start_or_end_with_punctuation_and_contain_spaces(self):
        entries = [{'kind': 'word', 'source': 'wg.', 'target': 'według'},
                   {'kind': 'word', 'source': 'itd.', 'target': 'itp.'},
                   {'kind': 'word', 'source': "'ten", 'target': 'ten'},
                   {'kind': 'word', 'source': 'e-book', 'target': 'e-książka'},
                   {'kind': 'word', 'source': 'na prawdę', 'target': 'naprawdę'}]
        self.assertEqual(
            self.fix("Zrobione wg. planu, koty itd. A 'ten' kot; e-book i e-booki. Na  prawdę, na\u00a0prawdę.",
                     {'user_word_corrections'}, entries),
            'Zrobione według planu, koty itp. A ten\' kot; e-książka i e-booki. Naprawdę, naprawdę.')
        # Granice wyrazu nadal obowiązują: „wg.” nie pasuje wewnątrz „swg.”.
        self.assertEqual(self.fix('swg. i wgx.', {'user_word_corrections'}, entries), 'swg. i wgx.')

    def test_slow_patterns_are_rejected_quickly(self):
        started = time.perf_counter()
        for source in (r'\w*\w*\w*\w*!', r'(a?){22}a{22}', r'\s*\s*\s*\s*x', r'(\w+\s?)*$',
                       r'(.*a){12}', r'(x+x+)+y', r'\d*\d*\d*\d*\d*z'):
            with self.subTest(source=source):
                self.assertIsNone(_compile_pattern(source))
                self.assertIsNone(_compile_pattern(source, 'Z'))
        self.assertLess(time.perf_counter() - started, 3)
        # Wzorce domyślnej listy i zwykłe wzorce przechodzą (także z jednym argumentem).
        for entry in DEFAULT_REPLACEMENTS:
            if entry['kind'] == 'pattern':
                with self.subTest(source=entry['source'][:30]):
                    self.assertIsNotNone(_compile_pattern(entry['source'], entry['target']))
        self.assertIsNotNone(_compile_pattern(r'\bkot(?P<e>ów|a)?\b'))
        self.assertIsNotNone(_compile_pattern(r'\b(?:bardzo )+dobry'))
        self.assertIsNone(_compile_pattern('(kot'))

    def test_more_abbreviations_do_not_start_a_sentence(self):
        for text in ('Jan, ur. w Krakowie.', 'Kierownik ds. kadr.', 'Pismo dot. sprawy.', 'Dnia dn. piątego.',
                     'We wt. rano.', 'Podaj temp. wody.', 'Mój tel. komórkowy.', 'Zrobione wg. planu.',
                     'Rozmowa nt. prozy.', 'W pt. wieczorem.', 'W śr. i czw. rano.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, {'sentence_case', 'after_punct'}), text)
        self.assertEqual(correct_editorial_text('Zapłacił 5 zł. potem wyszedł.', {'sentence_case'}),
                         'Zapłacił 5 zł. Potem wyszedł.')

    def test_file_and_technology_names_are_protected_by_every_rule(self):
        text = 'Plik raport.docx, biblioteka Node.js i Vue.js, platforma ASP.NET.'
        for rules in ({'after_punct'}, FORM_DEFAULT_EDITORIAL_RULES):
            with self.subTest(rules=len(rules)):
                self.assertEqual(self.fix(text, rules), text)
        self.assertEqual(self.fix('Wrócił do domu.Co dalej?', {'after_punct'}), 'Wrócił do domu. Co dalej?')

    def test_hour_ranges_are_normalized_at_both_ends(self):
        for text in ('Czynne 8.00-16.00.', 'Czynne 8:00-16:00.', 'Czynne 8:00 - 16:00.', 'Czynne 08:00–16:00.'):
            with self.subTest(text=text):
                self.assertEqual(self.fix(text), 'Czynne 08:00–16:00.')
        self.assertEqual(self.fix('Czynne 8:00-16:00.', {'range_dash'}), 'Czynne 8:00–16:00.')
        self.assertEqual(self.fix('Czynne 8:00 - 16:00.', {'range_spaces'}), 'Czynne 8:00-16:00.')
        # Kwoty i inne liczby z kropką nie są godzinami.
        for text in ('Cena 1.50-2.50 zł.', 'Wzrost 1.20-1.80 punktu.', 'Wersje 2.10-3.10 zostają.'):
            with self.subTest(text=text):
                self.assertNotIn(':', self.fix(text))

    def test_inch_marks_are_not_quotes(self):
        self.assertEqual(self.fix('Zdjęcie 5" x 7" leży.', {'inside_quotes'}), 'Zdjęcie 5" x 7" leży.')
        self.assertEqual(self.fix('Mówił " tak " i wyszedł.', {'inside_quotes'}), 'Mówił "tak" i wyszedł.')


class ProgramFormTests(SimpleTestCase):
    def test_cleaner_rule_defaults_and_order(self):
        self.assertEqual(EDITORIAL_RULES[-1][0], 'empty_paragraphs')
        form = OdkurzaczForm()
        initial = set(form.fields['rules'].initial)
        self.assertNotIn('empty_paragraphs', initial)
        self.assertTrue({key for key, *_ in EDITORIAL_RULES} - {'empty_paragraphs'} <= initial)
        self.assertTrue(form.fields['remove_soft_whitespace'].initial)
        # Automatyczne czyszczenie poczty zachowuje dotychczasowe reguły.
        self.assertTrue(DEFAULT_EDITORIAL_RULES)

    def test_only_the_cleaner_offers_formatting_and_soft_whitespace_options(self):
        self.assertNotIn('remove_soft_whitespace', DocumentConversionForm(prefix='convert').fields)
        self.assertNotIn('normalize_formatting', DocumentConversionForm().fields)
        self.assertNotIn('normalize_formatting', RepetitionsForm().fields)

    def test_damaged_packages_are_refused(self):
        """Pakiet rozpakowujący się do ponad 50 MB, z zaszyfrowanym wpisem albo bez word/document.xml."""
        with ZipFile(BytesIO(formatted_docx())) as archive:
            parts = [(item.filename, archive.read(item)) for item in archive.infolist()]
        expanding = BytesIO()
        with ZipFile(expanding, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', b'a' * (51 * 1024 * 1024))
            archive.writestr('[Content_Types].xml', b'')
        cases = (
            ('ponad 50 MB po rozpakowaniu', expanding.getvalue(), '50 MB'),
            ('zaszyfrowany wpis', package(parts, encrypted=True), 'Dokument nie może być zaszyfrowany.'),
            ('bez word/document.xml', package([part for part in parts if part[0] != 'word/document.xml']),
             'Ten plik nie jest prawidłowym dokumentem DOCX.'),
        )
        for label, data, message in cases:
            with self.subTest(label):
                form = OdkurzaczForm({}, {'document': upload(data)})
                self.assertFalse(form.is_valid())
                self.assertIn(message, str(form.errors['document']))

    def test_stored_text_settings_of_another_type_are_skipped(self):
        stored = {'ignored_words': 5, 'tracked_words': 'kot pies', 'color_palette': ['dark']}
        self.assertEqual(RepetitionsForm.initial_from_settings(stored), {'tracked_words': 'kot pies'})


class EditorialRuleR9Tests(SimpleTestCase):
    """Runda 9: przypadki z przeglądu (cudzysłowy, kreski, maski, skróty, listy, niewidoczne znaki)."""

    def fix(self, text, rules=FORM_DEFAULT_EDITORIAL_RULES):
        return correct_editorial_text(text, rules)

    def test_review_cases_are_corrected(self):
        cases = {
            # cudzysłowy
            'Koniec cytatu."': 'Koniec cytatu."',
            '”tak” i ”nie”.': '„tak” i „nie”.',
            '„tak" i "nie„.': '„tak” i „nie”.',
            '"Cytat z "zagnieżdżonym" cudzysłowem".': '„Cytat z «zagnieżdżonym» cudzysłowem”.',
            '„Cytat „zagnieżdżony” polski”.': '„Cytat «zagnieżdżony» polski”.',
            '„Cytat z "zagnieżdżonym" cudzysłowem”.': '„Cytat z «zagnieżdżonym» cudzysłowem”.',
            "Powiedział ``tak'' i ''nie''.": 'Powiedział „tak” i „nie”.',
            '«Cytat» i »odwrotny«.': '„Cytat” i „odwrotny”.',
            'To było "5" x 7" zdjęcie.': 'To było "5" x 7" zdjęcie.',
            "Harry'ego i Jacques'a, rok '98.": 'Harry’ego i Jacques’a, rok \'98.',
            "It's a 'test' i ‘angielski’ cytat.": 'It’s a „test” i „angielski” cytat.',
            # kreski
            'Słowo – – słowo, słowo –– słowo, słowo -- -- słowo.': 'Słowo – słowo, słowo – słowo, słowo – słowo.',
            '„– Tak – powiedział.”': '„– Tak – powiedział”.',
            'Zdanie.- Nie. Tak?- Nie!- Może.': 'Zdanie. – Nie. Tak? – Nie! – Może.',
            '- 5 osób przyszło.': '– 5 osób przyszło.',
            '-5 stopni.': '-5 stopni.',
            'Lata 1990‒1995 i 1990―1995 i 1990−1995.': 'Lata 1990–1995 i 1990–1995 i 1990–1995.',
            'Czarno‐biały i czarno‑biały.': 'Czarno-biały i czarno-biały.',
            '― Idziemy ― powiedział.': '– Idziemy – powiedział.',
            '2 - 1 = 1, a 3 + 4 - 2.': '2 - 1 = 1, a 3 + 4 - 2.',
            # interpunkcja
            'Kwota 1.234.567,89 zł i wersja 1.2.3:tak.': 'Kwota 1.234.567,89 zł i wersja 1.2.3: tak.',
            'Nie, , nie, nie,, nie.': 'Nie, nie, nie, nie.',
            'Jak to?.. Tak!.. Nie.': 'Jak to?… Tak!… Nie.',
            'Tak … – powiedział. Co …tak.': 'Tak… – powiedział. Co… tak.',
            '– …i tak dalej.': '– …i tak dalej.',
            # daty, skróty, inicjały
            'Dnia 5.X.2024 r. tak.': 'Dnia 5.X.2024 r. tak.',
            'S.A. i P.S. i U.S.A. i N.N. i J.R.R. Tolkien i A.B. Nowak.':
                'S.A. i P.S. i U.S.A. i N.N. i J. R. R. Tolkien i A. B. Nowak.',
            'Lista: a) jeden. b) dwa. c. trzy.': 'Lista: a) jeden. b) dwa. c. trzy.',
            'Zrobił to. a potem wyszedł.': 'Zrobił to. A potem wyszedł.',
            'W 5tys. zł i 5mln i 2mld i 5tej.': 'W 5 tys. zł i 5 mln i 2 mld i 5tej.',
            'Zobacz nr.5, nr. 5, t.3, cz.2, art.5, ok.5 i s.5.': 'Zobacz nr 5, nr 5, t. 3, cz. 2, art. 5, ok. 5 i s. 5.',
            # jednostki
            'Wzrost o 5 % i 5 ‰.': 'Wzrost o 5% i 5‰.',
            'Mieszkanie 50m2, 50 m2 i 5km2 i 2m3.': 'Mieszkanie 50 m², 50 m² i 5 km² i 2 m³.',
            # niewidoczne znaki i nietypowe spacje
            'Ze​ro­wid﻿th⁠ i‎ tak‏.': 'Zerowidth i tak.',
            'Spacja wąska, em,　ideograficzna, 5 kg.': 'Spacja wąska, em, ideograficzna, 5 kg.',
            # separatory
            '* * *': '* * *',
            '- - -': '- - -',
            '-----': '-----',
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.fix(text), expected)

    def test_new_rules_can_be_disabled_and_are_counted(self):
        self.assertEqual(self.fix('Ze​ro i Harry\'ego i 5 %', set()), 'Ze​ro i Harry\'ego i 5 %')
        self.assertEqual(self.fix("Harry'ego", {'invisible_chars'}), "Harry'ego")
        counts = {}
        correct_editorial_text('Ze​ro­ "tak" i Harry\'ego 5 %', FORM_DEFAULT_EDITORIAL_RULES, counts=counts)
        self.assertEqual({key: counts[key] for key in ('invisible_chars', 'quotes', 'apostrophe', 'unit_nbsp')},
                         {'invisible_chars': 2, 'quotes': 2, 'apostrophe': 1, 'unit_nbsp': 1})
        self.assertIn('invisible_chars', DEFAULT_EDITORIAL_RULES)

    def test_checked_dictionary_entries_skip_the_timing_probe(self):
        from core.services.odkurzacz import CompiledReplacements
        with patch('core.services.odkurzacz._too_slow', side_effect=AssertionError('probe')):
            compiled = CompiledReplacements([{'kind': 'pattern', 'source': r'\bkot(?P<e>a|y)?\b',
                                              'target': r'pies\g<e>', 'checked': True}])
        self.assertEqual(compiled.apply('kot i koty'), 'pies i piesy')
        with patch('core.services.odkurzacz._too_slow', return_value=False) as probe:
            CompiledReplacements([{'kind': 'pattern', 'source': r'\bkot\b', 'target': 'pies'}])
        probe.assert_called_once()


class AbbreviationListTests(SimpleTestCase):
    """Lista skrótów Odkurzacza: kropki wewnątrz skrótu bez spacji, mała litera po skrócie,
    zachowany zapis ze spacją, scalanie jednoznacznych skrótów i zwykłe wyrazy na końcu zdania."""

    # grupa listy: zdania, w których skrót stoi w środku zdania (żadna reguła domyślna ich nie zmienia)
    SENTENCES = {
        'ogólne': ('Ta w.w. osoba wymieniła Piotra i in. świętych.', 'Patrz s. 5 i nast. strony, zał. nr 2.',
                   'Dom z c.o. i wew. instalacją, odp. kwota.'),
        # p.p. (per procura) nie ma na liście: ciąg liter z kropkami nie rozpada się na „p. p.”.
        'tytuły, stopnie i funkcje': ('Pismo podpisał p.o. dyrektora, w.z. kierownika, p.p. Jan Nowak.',
                                      'Był prof. nadzw. uczelni.',
                                      'Podpisał z up. kierownika Sz.P. Jan, r.pr. Nowak.'),
        'kościelne': ('Kościół p.w. Najświętszej Marii Panny i pw. świętych Piotra.',
                      'Parafia rz.-kat. w mieście.', 'Modlili się za śś. męczenników.'),
        'wojskowe i służbowe': ('Awansował na st. szer. rezerwy.', 'Był rtm. kawalerii i adm. floty.'),
        'adresowe': ('Mieszka w m.st. Warszawie, w bl. czwartym, pok. siódmy.', 'Wieś leży na płn. od miasta.'),
        'firmy i instytucje': ('Firma s.c. i spółka sp.j. oraz sp. z o.o.', 'Spółka w likw. ogłosiła upadłość.'),
        'prawo i urzędy': ('Zgodnie z art. 5 k.p.c., k.r.o. i k.s.h. oraz rozp. ministra, z późn. zm.',
                           'Ogłoszono w Dz.U. i M.P. z 2020 r., l.dz. 5, sygn. akt II K 1/20.',
                           'Sąd wydał wyr. zaoczny.'),
        'bibliografia i wydawnictwa': ('Wyd. 2, popr. i uzup., b.m.w. i b.r.w.', 'Tom 2, tabl. i wykr. w aneksie.'),
        'czas i miary': ('W maju ub.r. i b.r. było ciepło, a ub.m. też.', 'Szczyt ma 300 m n.p.m., ok. 2 km wys. i jest stromy.',
                         'Trwało 5 sek. dłużej.'),
        'językowe': ('Słowo przestarz. i rzad. używane, gw. regionalne.', 'Forma zdr. od imienia, jęz. polski.'),
    }

    def test_list_has_unique_entries_in_named_groups(self):
        """Każdy wpis listy jest jeden raz, małymi literami i kończy się kropką; grupy pokrywają testy."""
        entries = [entry for _, group in ABBREVIATION_GROUPS for entry in group.split()]
        duplicates = sorted({entry for entry in entries if entries.count(entry) > 1})
        self.assertEqual(duplicates, [])
        self.assertEqual(len(SENTENCE_ABBREVIATIONS), len(entries))
        for entry in entries:
            with self.subTest(entry=entry):
                self.assertTrue(entry.endswith('.') and entry == entry.lower())
        self.assertEqual({name for name, _ in ABBREVIATION_GROUPS}, set(self.SENTENCES))

    def test_abbreviations_keep_internal_dots_and_lowercase_after_them(self):
        """after_punct nie rozdziela kropek wewnątrz skrótu, a sentence_case nie zaczyna po nim zdania."""
        for group, sentences in self.SENTENCES.items():
            for text in sentences:
                for rules in ({'after_punct'}, {'sentence_case'}, FORM_DEFAULT_EDITORIAL_RULES):
                    with self.subTest(group=group, text=text, rules=len(rules)):
                        self.assertEqual(correct_editorial_text(text, rules), text)

    def test_missing_space_after_an_abbreviation_is_added_once(self):
        """Brakująca spacja po całym skrócie (także przed cudzysłowem); kropki w środku skrótu zostają."""
        cases = {
            'Był p.o.Dyrektora i np.w domu.': 'Był p.o. Dyrektora i np. w domu.',
            'Kościół p.w.Świętej Anny, firma s.c.Kowalski.': 'Kościół p.w. Świętej Anny, firma s.c. Kowalski.',
            'Książka pt.„Lalka” i Dz.U.Nr 5.': 'Książka pt. „Lalka” i Dz.U. Nr 5.',
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES), expected)

    def test_existing_spaced_forms_stay_unchanged(self):
        """Zapis ze spacją wewnątrz skrótu (Dz. U., p. o.) i bez niej zostaje taki, jaki był."""
        for text in ('Był p. o. dyrektora.', 'Ogłoszono w Dz. U. z 2020 r. i Dz. Urz. woj.',
                     'Mieszka w m. st. Warszawie.', 'Spotkanie ub. r. i b. r. minęło.', 'Art. 5 k. p. c. stanowi.',
                     'Ogłoszono w Dz.U. z 2020 r. i Dz.Urz. woj.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES), text)

    def test_only_unambiguous_spaced_abbreviations_are_merged(self):
        """Reguła abbreviations scala n. p. m., p. p. m. i p. n. e.; „t.j.” przed Dz.U. to tekst jednolity."""
        cases = {
            'Rok 50 p. n. e., 300 m n. p. m. i 2 m p.p. m.': 'Rok 50 p.n.e., 300 m n.p.m. i 2 m p.p.m.',
            '(t.j. Dz.U. z 2020 r.) i (t. j. Dz. U.)': '(t.j. Dz.U. z 2020 r.) i (t. j. Dz. U.)',
            'Był t. j. kierownikiem, m. in. tam.': 'Był tj. kierownikiem, m.in. tam.',
            'Był p. o. dyrektora, b. r. i N. P. M. Kowalski.': 'Był p. o. dyrektora, b. r. i N. P. M. Kowalski.',
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, {'abbreviations'}), expected)

    def test_words_equal_to_skipped_abbreviations_still_end_a_sentence(self):
        """Wyrazy równe pominiętym skrótom (kat., im., rzecz., sport. …) kończą zdanie: dalej wielka litera."""
        for word in ('kat', 'im', 'rzecz', 'sport', 'lis', 'nauk', 'czas', 'post', 'sąd', 'mat', 'ser', 'pułk',
                     'marsz', 'miejsc', 'dom', 'las', 'pan', 'bal'):
            with self.subTest(word=word):
                self.assertNotIn(f'{word}.', SENTENCE_ABBREVIATIONS)
                self.assertEqual(correct_editorial_text(f'Był tam {word}. potem wyszedł.', {'sentence_case'}),
                                 f'Był tam {word}. Potem wyszedł.')

    def test_abbreviation_at_the_end_of_a_sentence_keeps_the_next_capital(self):
        """Po skrócie kończącym zdanie wielka litera i zaimek zostają (itd. Potem, i in. On)."""
        for text in ('Mówił o kotach, psach itd. Potem wyszedł.', 'Wymienił Jana i in. On przyszedł.',
                     'Tekst ustawy z późn. zm. Ona obowiązuje.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES), text)

    def test_numbers_after_reference_abbreviations_get_a_space(self):
        """Brakująca spacja między skrótem a liczbą także po rys., tab., zał., sygn., pkt."""
        self.assertEqual(
            correct_editorial_text('Patrz rys.3, ryc.4, tab.5, il.2, fot.1, zał.3, sygn.12, pkt3, zd.2, wyd.2 i nr3.',
                                   {'reference_space'}),
            'Patrz rys. 3, ryc. 4, tab. 5, il. 2, fot. 1, zał. 3, sygn. 12, pkt 3, zd. 2, wyd. 2 i nr 3.')
        self.assertEqual(correct_editorial_text('Opkt3 i tabela5.', {'reference_space'}), 'Opkt3 i tabela5.')

    def test_long_paragraph_full_of_abbreviations_is_fast(self):
        """Kilkaset skrótów w jednym wzorcu: najdłuższy akapit nadal szybko (wzorzec grupowany po literze)."""
        source = ('Był p.o. dyrektora, m.in. prof. Kowalski z ul. Polnej 5, pt.„Lalka”, wyd. 2, popr. i uzup., '
                  's. 15 i nast. Potem ks. proboszcz parafii pw. św. Anny (Dz.U. z 2020 r. poz. 5) i in. ' * 120)[:20000]
        started = time.monotonic()
        result = correct_editorial_text(source, FORM_DEFAULT_EDITORIAL_RULES)
        self.assertLess(time.monotonic() - started, 1.5)
        self.assertIn('Był p.o. dyrektora, m.in. prof. Kowalski z ul. Polnej 5, pt. „Lalka”', result)
        self.assertIn('poz. 5) i in. Był', result)


class OwnerReplacementListTests(SimpleTestCase):
    """Lista automatycznych zamian właściciela (październik 2026): pełna fleksja, wielkość liter jak w tekście
    (nazwy mieszkańców i ras – wielką literą), poprawne wyrazy-sąsiedzi bez zmian; pickup → pikap usunięte."""

    # nazwa: ({błędna forma: poprawna}, [formy, które muszą zostać])
    CASES = {
        'absorpcja': ({'absorbcja': 'absorpcja', 'absorbcji': 'absorpcji', 'absorbcją': 'absorpcją',
                       'absorbcjami': 'absorpcjami', 'absorbcyjny': 'absorpcyjny', 'absorbcyjnych': 'absorpcyjnych',
                       'absorbcyjność': 'absorpcyjność'}, ['absorpcja', 'absorbować', 'absorbuje', 'absorbent']),
        'Azjata, Azjatka': ({'azjata': 'Azjata', 'azjacie': 'Azjacie', 'azjatów': 'Azjatów', 'azjatka': 'Azjatka',
                             'azjatce': 'Azjatce', 'azjatek': 'Azjatek', 'azjatkami': 'Azjatkami'},
                            ['azjatycki', 'azjatyckiej']),
        'cętkowany, cętka': ({'centkowana': 'cętkowana', 'centkowanego': 'cętkowanego', 'centkowani': 'cętkowani',
                              'centka': 'cętka', 'centki': 'cętki', 'centce': 'cętce', 'centek': 'cętek',
                              'centkami': 'cętkami'},
                             ['cent', 'centa', 'centów', 'centem', 'cencie', 'Centkiewicz', 'Centkiewiczowie',
                              'procent', 'centrum']),
        'Europejczyk, Europejka': ({'europejczykiem': 'Europejczykiem', 'europejczycy': 'Europejczycy',
                                    'europejka': 'Europejka', 'europejce': 'Europejce', 'europejek': 'Europejek'},
                                   ['europejski', 'europejska', 'europejskie']),
        'greckokatolicki': ({'grekokatolicki': 'greckokatolicki', 'grekokatolickiej': 'greckokatolickiej',
                             'grekokatoliccy': 'greckokatoliccy', 'grekokatolickimi': 'greckokatolickimi'},
                            ['grekokatolik', 'grekokatolika', 'grekokatolicy', 'grekokatolików', 'greckokatolicki']),
        'indywidualny': ({'indywidulany': 'indywidualny', 'indywidulanej': 'indywidualnej',
                          'indywidulani': 'indywidualni', 'indywidulanie': 'indywidualnie'}, ['indywidualista']),
        'dell’arte': ({'dellarte': 'dell’arte', 'komedia dellarte': 'komedia dell’arte',
                       'komedii dellarte': 'komedii dell’arte'}, ['dell’arte']),
        'Kreol, Kreolka': ({'kreol': 'Kreol', 'kreola': 'Kreola', 'kreolowie': 'Kreolowie', 'kreolami': 'Kreolami',
                            'kreolka': 'Kreolka', 'kreolce': 'Kreolce'},
                           ['kreolski', 'kreolska', 'kreolskich', 'kreolszczyzna', 'kreolszczyzny']),
        'księżmi': ({'księżami': 'księżmi'}, ['księżmi', 'księży', 'księża', 'księżom']),
        'kurewstwo': ({'kurestwo': 'kurewstwo', 'kurestwa': 'kurewstwa', 'kurestwem': 'kurewstwem',
                       'kurestwie': 'kurewstwie', 'kurestw': 'kurewstw'}, ['kurewstwo']),
        'mityng, mityngowy': ({'miting': 'mityng', 'mitingu': 'mityngu', 'mitingiem': 'mityngiem',
                               'mitingów': 'mityngów', 'mitingowy': 'mityngowy', 'mitingowej': 'mityngowej'},
                              ['mityng', 'mitygować']),
        'Mulat, Mulatka': ({'mulat': 'Mulat', 'mulata': 'Mulata', 'mulacie': 'Mulacie', 'mulaci': 'Mulaci',
                            'mulatów': 'Mulatów', 'mulatka': 'Mulatka', 'mulatce': 'Mulatce', 'mulatek': 'Mulatek'},
                           ['mulacki', 'mulacka', 'mulackie', 'muła', 'mułach']),
        'nabazgrzę': ({'nabazgram': 'nabazgrzę', 'nabazgrasz': 'nabazgrzesz', 'nabazgra': 'nabazgrze',
                       'nabazgramy': 'nabazgrzemy', 'nabazgracie': 'nabazgrzecie', 'nabazgrają': 'nabazgrzą'},
                      ['nabazgrał', 'nabazgrała', 'nabazgrali', 'nabazgrać', 'nabazgrzę', 'nabazgrany']),
        'nie dowierzać': ({'niedowierza': 'nie dowierza', 'niedowierzam': 'nie dowierzam',
                           'niedowierzają': 'nie dowierzają', 'niedowierzać': 'nie dowierzać',
                           'niedowierzał': 'nie dowierzał', 'niedowierzałaby': 'nie dowierzałaby',
                           'niedowierzaliśmy': 'nie dowierzaliśmy', 'niedowierzając': 'nie dowierzając'},
                          ['niedowierzanie', 'niedowierzaniem', 'niedowierzający', 'niedowierzająca',
                           'niedowierzającym', 'niedowiarek', 'nie dowierza']),
        # Reforma 2026 (decyzja właściciela): „nie” ze stopniem wyższym i najwyższym łącznie – słownik nie rozdziela.
        'nie ze stopniem wyższym i najwyższym': ({}, ['niemniejszy', 'niemniejsza', 'niemniejsi', 'nienajgorszy',
                                                       'nienajgorzej', 'nienajlepsza', 'nienajlepsi', 'nienajlepiej',
                                                       'niemniej', 'nie mniej niż', 'nienajedzony']),
        'nie wiadomo': ({'niewiadomo': 'nie wiadomo'}, ['niewiadoma', 'niewiadomą', 'niewiadome', 'niewiadomych']),
        'oddziałujący': ({'oddziaływujący': 'oddziałujący', 'nieoddziaływujących': 'nieoddziałujących',
                          'oddziaływuje': 'oddziałuje', 'oddziaływują': 'oddziałują', 'oddziaływując': 'oddziałując'},
                         ['oddziaływać', 'oddziaływał', 'oddziaływanie', 'oddziałujący', 'oddziaływający']),
        'oldbojów': ({'oldboi': 'oldbojów'}, ['oldbojów', 'oldboje', 'oldboj']),
        'przeciwwskazany': ({'przeciwskazana': 'przeciwwskazana', 'przeciwskazanych': 'przeciwwskazanych',
                             'przeciwskazanie': 'przeciwwskazanie', 'przeciwskazań': 'przeciwwskazań'},
                            ['przeciwwskazana', 'wskazana']),
        'przekonujący': ({'przekonywujący': 'przekonujący', 'nieprzekonywującego': 'nieprzekonującego',
                          'przekonywująco': 'przekonująco', 'przekonywuje': 'przekonuje',
                          'przekonywujesz': 'przekonujesz'},
                         ['przekonywać', 'przekonywał', 'przekonywający', 'przekonujący', 'przekonywanie']),
        'przysiągł': ({'przysięgł': 'przysiągł', 'przysięgłem': 'przysiągłem', 'przysięgłeś': 'przysiągłeś',
                       'przysięgłby': 'przysiągłby', 'zaprzysięgł': 'zaprzysiągł'},
                      ['przysięgła', 'przysięgły', 'przysięgłe', 'przysięgłego', 'przysięgłych', 'przysięgli',
                       'przysięgłam', 'zaprzysięgły', 'zaprzysięgła']),
        'kategoryzować (nie ma „skategoryzować”)': (
            {'skategoryzować': 'kategoryzować', 'skategoryzował': 'kategoryzował', 'skategoryzuję': 'kategoryzuję',
             'skategoryzowany': 'kategoryzowany', 'skategoryzowanie': 'kategoryzowanie',
             'nieskategoryzowane': 'niekategoryzowane'}, ['kategoryzacja', 'kategoria']),
        'średnio zaawansowany': ({'średniozaawansowany': 'średnio zaawansowany',
                                  'średniozaawansowanych': 'średnio zaawansowanych',
                                  'średniozaawansowani': 'średnio zaawansowani'}, ['zaawansowany']),
        'tonik': ({'tonic': 'tonik', 'tonicu': 'toniku', 'tonicem': 'tonikiem', 'tonicy': 'toniki'}, ['toniku']),
        'wykrztusić': ({'wyksztusić': 'wykrztusić', 'wyksztusił': 'wykrztusił', 'wyksztuszę': 'wykrztuszę',
                        'wyksztuś': 'wykrztuś', 'wyksztuszony': 'wykrztuszony', 'wyksztusiła': 'wykrztusiła'},
                       ['wykrztusić', 'wykształcić', 'wykształcony']),
        'zadowalający': ({'zadawalający': 'zadowalający', 'niezadawalających': 'niezadowalających',
                          'zadawalająco': 'zadowalająco', 'zadawala': 'zadowala', 'zadawalał': 'zadowalał'},
                         ['zadawał', 'zadawała', 'zadawali', 'zadowalający', 'zadawać']),
        'zapewne': ({'zapewnie': 'zapewne'}, ['zapewnię', 'zapewni', 'zapewne']),
        'zarządzanie': ({'zarzadzaniu': 'zarządzaniu', 'zarzadzanie': 'zarządzanie', 'zarzadza': 'zarządza',
                         'zarzadzający': 'zarządzający', 'zarzadzenie': 'zarządzenie'}, ['zarządzaniu', 'zarząd']),
        'scentralizowany': ({'zcentralizowana': 'scentralizowana', 'zcentralizować': 'scentralizować',
                             'niezcentralizowanych': 'niescentralizowanych'},
                            ['scentralizowana', 'zdecentralizowana', 'centralizować']),
        'miauczeć': ({'miałczeć': 'miauczeć', 'miałczał': 'miauczał', 'zamiałczała': 'zamiauczała',
                      'miałczenie': 'miauczenie', 'miałcząc': 'miaucząc', 'pomiałczmy': 'pomiauczmy'}, ['miauczeć']),
        'pozostałe pozycje': ({'niechcąco': 'niechcący', 'zaczym': 'za czym', 'władnie': 'władny',
                               'terefere': 'tere-fere', 'inąd': 'skądinąd', 'wszechczasów': 'wszech czasów',
                               'karnister': 'kanister', 'karnistrem': 'kanistrem', 'massmediach': 'mass mediach',
                               'mentlikiem': 'mętlikiem', 'wymyśleć': 'wymyślić', 'domyśleć': 'domyślić',
                               'dopatrzeć': 'dopatrzyć', 'dopatrzał': 'dopatrzył'},
                              ['skądinąd', 'wymyślał', 'wymyślę', 'domyślał', 'dopatrzył']),
        'pickup zostaje (decyzja właściciela)': ({}, ['pickup', 'pickupem', 'Pickupami', 'PICKUP']),
    }

    @staticmethod
    def fix(text):
        return correct_editorial_text(text, {'user_word_corrections'})

    @staticmethod
    def capital(text):
        return text[:1].upper() + text[1:]

    def test_every_form_is_changed_with_its_letter_case_and_neighbours_stay(self):
        """Każda błędna forma małymi literami, od wielkiej litery i WERSALIKAMI (wynik w tym samym zapisie;
        nazwy mieszkańców zawsze wielką literą), także w zdaniu; poprawne wyrazy podobne zostają."""
        for name, (changed, kept) in self.CASES.items():
            with self.subTest(name):
                for source, target in changed.items():
                    self.assertEqual(self.fix(source), target)
                    self.assertEqual(self.fix(self.capital(source)), self.capital(target))
                    self.assertEqual(self.fix(source.upper()), target.upper())
                    self.assertEqual(self.fix(f'Było to, {source}, a potem.'), f'Było to, {target}, a potem.')
                for word in kept:
                    self.assertEqual(self.fix(word), word)
                    self.assertEqual(self.fix(word.upper()), word.upper())

    def test_default_entries_pass_the_dictionary_checks_and_stay_fast(self):
        """Wpisy domyślne przechodzą tę samą kontrolę co formularz słownika (składnia, zagnieżdżone powtórzenia,
        próba czasu, limity długości), a cała lista na długim akapicie działa szybko."""
        for entry in DEFAULT_REPLACEMENTS:
            with self.subTest(entry['source'][:40]):
                self.assertNotIn('pickup', entry['source'])
                if entry['kind'] == 'word':
                    self.assertTrue(WORD_SOURCE_RE.fullmatch(entry['source']))
                    self.assertLessEqual(len(entry['source']), REPLACEMENT_WORD_MAX)
                else:
                    self.assertLessEqual(len(entry['source']), MAX_PATTERN_LENGTH)
                    self.assertIsNotNone(_compile_pattern(entry['source'], entry['target']))
        words = [source for changed, kept in self.CASES.values() for source in [*changed, *kept]]
        paragraph = ' '.join(words * (20_000 // len(' '.join(words)) + 1))[:20_000]
        compiled = compile_replacements()
        start = time.perf_counter()
        compiled.apply(paragraph)
        self.assertLess(time.perf_counter() - start, 1)


class Orthography2026RuleTests(SimpleTestCase):
    """Grupa „Ortografia 2026” (komunikat RJP z 10 maja 2024 r., część A): -by po spójnikach i partykułach,
    niby/quasi/pseudo, łącznik po przedrostku przed wielką literą, pół oraz „nie” ze stopniem wyższym
    i najwyższym. Każda reguła działa sama, w każdej wielkości liter i w zdaniu; wyjątki zostają."""

    @staticmethod
    def fix(text, rules):
        return correct_editorial_text(text, rules)

    def assertChanges(self, rule, changed, kept):
        """Zmieniane formy (tak jak w tabeli, od wielkiej litery, w zdaniu) i formy, które zostają."""
        for source, target in changed.items():
            with self.subTest(rule=rule, source=source):
                self.assertEqual(self.fix(source, {rule}), target)
                self.assertEqual(self.fix(f'Był to, {source}, a potem.', {rule}), f'Był to, {target}, a potem.')
                if source.islower():
                    capital = source[:1].upper() + source[1:]
                    self.assertEqual(self.fix(capital, {rule}), target[:1].upper() + target[1:])
        for text in kept:
            with self.subTest(rule=rule, kept=text):
                self.assertEqual(self.fix(text, {rule}), text)
                self.assertEqual(self.fix(text, FORM_DEFAULT_EDITORIAL_RULES - {'sentence_case', 'pronouns_lower'}),
                                 self.fix(text, FORM_DEFAULT_EDITORIAL_RULES - ORTHOGRAPHY_2026_KEYS
                                          - {'sentence_case', 'pronouns_lower'}))

    def test_rules_form_one_group_switched_on_by_default(self):
        """Pięć osobnych reguł jedna po drugiej, przed wielką literą po kropce; etykiety z nazwą grupy
        i przykładem; domyślnie włączone w formularzu i w automatycznym czyszczeniu zgłoszeń z poczty."""
        keys = [key for key, _ in EDITORIAL_RULES]
        group = [key for key, _ in ORTHOGRAPHY_2026_RULES]
        self.assertEqual(group, ['by_separate', 'niby_quasi', 'prefix_hyphen', 'pol_compounds', 'nie_comparative'])
        start = keys.index(group[0])
        self.assertEqual(keys[start:start + len(group)], group)
        self.assertLess(start, keys.index('sentence_case'))
        for key, label in ORTHOGRAPHY_2026_RULES:
            with self.subTest(key=key):
                self.assertTrue(label.startswith('Ortografia 2026: ') and '→' in label)
        self.assertTrue(ORTHOGRAPHY_2026_KEYS <= DEFAULT_EDITORIAL_RULES)
        self.assertTrue(ORTHOGRAPHY_2026_KEYS <= FORM_DEFAULT_EDITORIAL_RULES)
        self.assertTrue(ORTHOGRAPHY_2026_KEYS <= set(OdkurzaczForm().fields['rules'].initial))

    def test_by_is_written_separately_after_conjunctions_and_particles(self):
        """Zmiana 3: czyby → czy by we wszystkich końcówkach; zrosty (aby, żeby, gdyby, jakby, czyżby, niechby…),
        niejednoznaczne toby, jeśliby, niżeliby i rzeczowniki boby, nimby zostają."""
        changed = {
            'czyby': 'czy by', 'czybym': 'czy bym', 'czybyś': 'czy byś', 'czybyśmy': 'czy byśmy',
            'czybyście': 'czy byście', 'alebym': 'ale bym', 'ależby': 'ależ by', 'alboby': 'albo by',
            'albowiemby': 'albowiem by', 'aniby': 'ani by', 'ażby': 'aż by', 'bowiemby': 'bowiem by',
            'chybaby': 'chyba by', 'czyliby': 'czyli by', 'dopókibyś': 'dopóki byś', 'jednakby': 'jednak by',
            'jednakżeby': 'jednakże by', 'jeszczebym': 'jeszcze bym', 'jużbyśmy': 'już byśmy', 'leczby': 'lecz by',
            'lubby': 'lub by', 'możeby': 'może by', 'natomiastby': 'natomiast by', 'nawetby': 'nawet by',
            'nużby': 'nuż by', 'orazby': 'oraz by', 'ponieważby': 'ponieważ by', 'pókiby': 'póki by',
            'przecieżbyś': 'przecież byś', 'przetoby': 'przeto by', 'skoroby': 'skoro by', 'tedyby': 'tedy by',
            'teżbym': 'też bym', 'toteżby': 'toteż by', 'tylkoby': 'tylko by', 'więcby': 'więc by',
            'wszakby': 'wszak by', 'wszakżeby': 'wszakże by', 'zanimby': 'zanim by', 'zatemby': 'zatem by',
            'zaśbyście': 'zaś byście', 'bobym': 'bo bym', 'bobyście': 'bo byście', 'nimbyś': 'nim byś',
            'CZYBY': 'CZY BY', 'Zastanawiam się, czyby nie pojechać w góry.': 'Zastanawiam się, czy by nie pojechać w góry.',
        }
        kept = ['aby', 'abym', 'ażeby', 'żeby', 'żebyś', 'iżby', 'gdyby', 'gdybyśmy', 'oby', 'niby', 'jakby', 'jakbym',
                'jakoby', 'coby', 'byleby', 'choćby', 'chociażby', 'czyżby', 'czyliżby', 'gdzieżby', 'jakżeby', 'niżby',
                'bodajby', 'niechby', 'niechajby', 'niechżeby', 'toby', 'jeśliby', 'jeżeliby', 'aniżeliby', 'niżeliby',
                'niźliby', 'boby', 'nimby', 'zęby', 'osoby', 'mógłby', 'chciałabym', 'czy by', 'e-czyby', 'czyby-',
                'www.czyby.pl', 'plik czyby.docx', 'jan.czyby@example.com']
        self.assertChanges('by_separate', changed, kept)

    def test_niby_quasi_pseudo_are_joined_before_lower_case_and_hyphenated_before_capitals(self):
        """Zmiana 10: niby-artysta → nibyartysta (także łącznik Unicode i wersaliki), nibyPolak → niby-Polak;
        osobne „niby”, niby-Polak, niby-e-book, nazwy od wielkiej litery i człon o końcówce formy czasownika
        (niby-żyję, quasi-myślę, niby-spał – a także quasi-naukę, niby-artystą) zostają."""
        changed = {
            'niby-artysta': 'nibyartysta', 'quasi-naukowy': 'quasinaukowy', 'pseudo-naukowy': 'pseudonaukowy',
            'niby‑romantycznie': 'nibyromantycznie', 'quasi‐opiekun': 'quasiopiekun',
            'NIBY-ARTYSTA': 'NIBYARTYSTA', 'nibyPolak': 'niby-Polak', 'quasiAnglia': 'quasi-Anglia',
            'pseudoKopernik': 'pseudo-Kopernik', 'quasiNATO': 'quasi-NATO', 'quasiGMO': 'quasi-GMO',
            'nibyPRL-u': 'niby-PRL-u',
        }
        kept = ['niby-żyję', 'quasi-myślę', 'niby-spał', 'niby-spała', 'quasi-istnieć', 'niby-robisz', 'quasi-naukę',
                'niby-artystą', 'quasi-pracujemy', 'niby-mówicie', 'quasi-wierzą',
                'nibyartysta', 'quasinauka', 'niby-Polak', 'quasi-Anglia', 'Niby-Polak', 'niby tak', 'Niby nic',
                'niby e-book', 'niby-e-book', 'pseudonim', 'Nibylandia', 'NibyPolak', 'niby-', 'quasi -naukowy',
                'pseudo-Kowalski', 'quasiNarodowy']
        self.assertChanges('niby_quasi', changed, kept)

    def test_prefix_before_a_capital_letter_gets_a_hyphen(self):
        """Zmiana 9a: superPolak → super-Polak, także przed skrótowcem i z końcówką (antyPiS, anty-NATO, post-PRL-u).
        Bez zmian: utrwalone terminy i nazwy z jawnej listy (mikroRNA, miniUSB, proALPHA, także z końcówką),
        przymiotnik na -ski/-cki/-dzki (proRosyjski, antyPolski – albo nazwisko; decyzja koordynatora) i na -owy
        (postSolidarnościowy), nazwy od wielkiej litery (MiniMax) oraz „pan”/„para” + nazwisko (brak spacji)."""
        changed = {'superPolak': 'super-Polak', 'antyPiS': 'anty-PiS', 'proPiS-u': 'pro-PiS-u', 'antyNATO': 'anty-NATO',
                   'neoPO': 'neo-PO', 'postPRL-u': 'post-PRL-u', 'eksMiss': 'eks-Miss',
                   'proEuropejczyk': 'pro-Europejczyk', 'antyPolacy': 'anty-Polacy', 'wiceMarszałek': 'wice-Marszałek',
                   'megaGwiazdor': 'mega-Gwiazdor', 'postSolidarność': 'post-Solidarność', 'antyGMO': 'anty-GMO'}
        kept = ['superX', 'miniB-2', 'mikroRNA', 'miniDV', 'proALPHA', 'miniUSB', 'miniUSB-a', 'mikroUSB', 'miniSD', 'mikroSDHC', 'miniHDMI',
                'microSD', 'antyPolski', 'antyPolskiego', 'proRosyjski', 'proRosyjskich', 'antyNiemiecki',
                'antyNiemieccy', 'proAmerykańską', 'postRadzieccy', 'arcyKowalski', 'MiniMax', 'SuperNova', 'EkoPark',
                'panKowalski', 'paraPolicja', 'postSolidarnościowy', 'antyhitlerowski', 'anty-Polak', 'super Polak',
                'iPhone', 'eBay', 'nanoFantazje', 'super-Polak', 'www.superPolak.pl', 'plik superPolak.docx']
        self.assertChanges('prefix_hyphen', changed, kept)

    def test_pol_is_joined_in_fixed_expressions_and_hyphenated_before_people(self):
        """Zmiana 6: pół żartem, pół serio → półżartem, półserio; pół Polak, pół Szkot → pół-Polak, pół-Szkot
        (mianownik i narzędnik, także pojedynczo); „pół” ilościowe zostaje, a zapis równy dopełniaczowi nazwy
        krainy (pół Niemiec, pół Włoch) zmienia się tylko w jednym zdaniu z pewną nazwą osoby."""
        changed = {
            'pół żartem': 'półżartem', 'pół serio': 'półserio', 'pół żartem, pół serio': 'półżartem, półserio',
            'PÓŁ ŻARTEM': 'PÓŁŻARTEM', 'pół Polak, pół Szkot': 'pół-Polak, pół-Szkot',
            'pół Polka, pół Francuzka': 'pół-Polka, pół-Francuzka', 'pół Polakiem, pół Niemcem': 'pół-Polakiem, pół-Niemcem',
            'pół Niemiec, pół Polak': 'pół-Niemiec, pół-Polak', 'pół Polką i pół Włoszką': 'pół-Polką i pół-Włoszką',
            'pół Azjatą': 'pół-Azjatą', 'półPolak': 'pół-Polak', 'pół Ukraińcem': 'pół-Ukraińcem',
            'pół Turkiem': 'pół-Turkiem', 'pół Węgrem': 'pół-Węgrem',
            'Zwiedził pół Niemiec. Pół Polak, pół Szkot': 'Zwiedził pół Niemiec. Pół-Polak, pół-Szkot',
        }
        kept = ['pół Polak, pół byk', 'pół byk, pół Polak', 'pół godziny', 'pół roku', 'pół na pół', 'pół litra',
                'Pół Polski, pół Europy', 'pół Warszawy',
                'pół Niemiec', 'pół Niemiec i pół Włoch', 'pół Kaszub', 'pół-Polak', 'półżartem', 'pół Polaka',
                'pół Żartem', 'pół Elon Musk', 'pół Polakowi', 'półtora']
        self.assertChanges('pol_compounds', changed, kept)

    def test_nie_is_joined_with_comparatives_and_superlatives(self):
        """Zmiana 11, tylko zapisy jednoznaczne: „nie” łącznie z przymiotnikiem w stopniu wyższym lub najwyższym
        (-ejszy, tematy na -szy, naj-) i z przysłówkiem lepiej, gorzej, prędzej, na -niej, -ściej, -ciej (lista)."""
        changed = {
            'nie najlepszy': 'nienajlepszy', 'nie najlepiej': 'nienajlepiej', 'NIE NAJLEPSZY': 'NIENAJLEPSZY',
            'nie lepiej niż ty': 'nielepiej niż ty', 'nie gorzej': 'niegorzej', 'nie prędzej': 'nieprędzej',
            'nie najprędzej': 'nienajprędzej', 'nie ładniejszy': 'nieładniejszy', 'nie ładniejsi': 'nieładniejsi',
            'nie mniejszy': 'niemniejszy', 'nie mniejszą': 'niemniejszą', 'nie większe': 'niewiększe',
            'nie lepsi': 'nielepsi', 'nie gorsza': 'niegorsza', 'nie najgorszych': 'nienajgorszych',
            'nie tańszych': 'nietańszych', 'nie droższym': 'niedroższym', 'nie milszy': 'niemilszy',
            'nie najmilszy': 'nienajmilszy', 'nie najstaranniej': 'nienajstaranniej', 'nie ładniej': 'nieładniej',
            'nie szybciej': 'nieszybciej', 'nie częściej': 'nieczęściej', 'nie najczęściej': 'nienajczęściej',
            'nie wyraźniej': 'niewyraźniej', 'nie taniej': 'nietaniej', 'nie cierpliwsza': 'niecierpliwsza',
            'nie najciekawszy': 'nienajciekawszy', 'nie nowszy': 'nienowszy', 'nie bystrzejszy': 'niebystrzejszy',
            'nie najpiękniejsza': 'nienajpiękniejsza', 'nie najszczęśliwsi': 'nienajszczęśliwsi',
            'nie obfitszy': 'nieobfitszy', 'nie\tnajlepszy': 'nienajlepszy',
            'To był nie najlepszy dzień.': 'To był nienajlepszy dzień.',
            'Nie najlepiej się czuł.': 'Nienajlepiej się czuł.',
            'Nie najlepszy, a nawet najgorszy.': 'Nienajlepszy, a nawet najgorszy.',
            'Głos sprawiał nie mniejszy ból niż światło.': 'Głos sprawiał niemniejszy ból niż światło.',
        }
        self.assertChanges('nie_comparative', changed, [])

    def test_nie_stays_separate_in_exceptions(self):
        """Decyzja właściciela – każdy zapis sporny zostaje: miara, czas, miejsce i znaczenia niestopniowe (dalej,
        wcześniej, później, dłuższy, dalszy, późniejszy, starszy…), tryb rozkazujący (siej, wyśmiej, miej), liczebniki
        i przysłówki nieodprzymiotnikowe, wyrazy o podobnej budowie, zaimek „nie” po przyimku, całe zdanie pytające,
        „czy nie”, warunek (jeśli, o ile, gdyby), przeciwstawienie, stan pośredni (ale i nie, choć nie), „nie”
        po „wcale”, „bynajmniej”, „ani” oraz zapis mieszany wielkimi i małymi literami."""
        kept = [
            # miara, czas, miejsce, znaczenia niestopniowe
            'Nie dalej jak dwa lata temu.', 'nie dalej niż dwa metry', 'nie najdalej', 'nie bliżej', 'nie najbliżej',
            'nie wcześniej niż w dziesiątym i nie później niż w dwunastym dniu', 'nie najwcześniej', 'nie najpóźniej',
            'nie dawniej niż', 'nie dłużej niż', 'nie krócej', 'nie wyżej', 'nie niżej', 'nie najwyżej',
            'nie dłuższym niż pięć dni', 'nie krótszy', 'nie późniejszą niż dzisiejsza', 'nie wcześniejszą',
            'nie dalszy ciąg', 'nie bliższe informacje', 'nie wyższa uczelnia', 'nie niższy', 'nie najwyżsi',
            'nie starszy pan', 'nie młodszy brat', 'nie dawniejszy',
            # tryb rozkazujący i inne przysłówki spoza listy
            'nie siej zamętu', 'nie wyśmiej go', 'nie miej pretensji', 'nie śmiej się', 'nie zasiej', 'nie rozwiej',
            'nie zechciej', 'nie chciej', 'nie czerwieniej', 'nie łatwiej', 'nie ciszej', 'nie zdrowiej', 'nie milej',
            # liczebniki, przysłówki nieodprzymiotnikowe, partykuły, wyrazy o podobnej budowie
            'nie mniej niż pięć', 'niemniej ciekawy', 'nie więcej', 'nie najwięcej', 'nie bardziej', 'nie najbardziej',
            'nie najmniej', 'nie pierwszy', 'nie najpierw', 'nie najwyraźniej', 'nie najprawdopodobniej',
            'nie dzisiejszy', 'nie tutejsza', 'nie niniejszy', 'nie ostatniej', 'nie średniej', 'nie kilkuletniej',
            'nie zachodniej', 'nie trzeciej', 'nie złodziej', 'nie turniej', 'nie najadłszy się', 'nie określiwszy',
            'nie zmniejsza', 'to nie pomniejszy zasług', 'nie najem',
            # zapis
            'nie Najlepszy', 'nie NAJLEPSZY', 'Nie Lepiej', 'nie-najlepszy', 'nienajlepszy', 'www.nie-najlepszy.pl',
            # kontekst
            'Patrzył na nie łaskawiej.', 'Liczył za nie ładniej.', 'Czy to nie najlepszy pomysł?',
            'A nie lepiej odejść?', 'Nie lepiej odpocząć?', '– Nie lepiej odstrzelić go teraz? – mruknął.',
            'Czy nie lepiej…?', 'Czuł się nie najlepiej?', 'Sam nie wiem, czy nie przyjemniejsza niż sen.',
            'Skończyłby z siniakiem, jeśli nie gorzej.', 'jeżeli nie lepiej', 'o ile nie gorzej', 'gdyby nie lepiej',
            'Nie najmądrzejszy, ale mądry.', 'Nie wyższy, tylko najwyższy.', 'nie lepszy, lecz gorszy',
            'nie lepiej, a gorzej', 'Mądry, a nie najmądrzejszy.', 'Nie najlepiej, ale i nie najgorzej.',
            'nie najlepszy, ale nie najgorszy', 'nie lepszy, choć nie gorszy', 'nie ładniejsza, ale też nie brzydsza',
            'Wcale nie lepsza ode mnie.', 'wcale a wcale nie lepszy', 'Bynajmniej nie najmądrzejszy.',
            'ani nie lepszy, ani nie gorszy',
        ]
        self.assertChanges('nie_comparative', {}, kept)

    def test_interaction_with_other_rules(self):
        """Zaimki małą literą działają przed „nie” (Nie w środku zdania), wielka litera po kropce – po nim;
        skróty, słownik zamian i maski adresów współpracują z nowymi regułami (wszystkie reguły naraz)."""
        rules = FORM_DEFAULT_EDITORIAL_RULES
        cases = {
            'Zobaczyłem Nie najlepszy film.': 'Zobaczyłem nienajlepszy film.',
            'Tak. nie najlepiej, czyby nie.': 'Tak. Nienajlepiej, czy by nie.',
            'Był m. in. nie najlepszy i pół żartem karnistrem.': 'Był m.in. nienajlepszy i półżartem kanistrem.',
            'Adres www.czyby.pl i superPolak pisał:  nie lepiej.': 'Adres www.czyby.pl i super-Polak pisał: nielepiej.',
            '„Niby-artysta” – rzekł pół Polak, pół Szkot...': '„Nibyartysta” – rzekł pół-Polak, pół-Szkot…',
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, rules), expected)
        self.assertEqual(correct_editorial_text('Zobaczyłem Nie najlepszy film.', {'nie_comparative'}),
                         'Zobaczyłem Nienajlepszy film.')
        # Lista słownikowa sprzed cofnięcia (nienajlepszy → nie najlepszy) nie wygrywa z regułą.
        old_entry = [{'kind': 'pattern', 'source': r'\bnienaj(?P<e>lepsz\w*)\b', 'target': r'nie naj\g<e>'}]
        self.assertEqual(correct_editorial_text('Nienajlepszy dzień.', {'user_word_corrections', 'nie_comparative'},
                                                replacements=old_entry), 'Nienajlepszy dzień.')

    def test_counts_and_disabled_rules(self):
        """Podsumowanie liczy zmiany każdej reguły osobno; wyłączona reguła niczego nie zmienia."""
        text = 'Czyby to nie najlepszy niby-artysta, superPolak i pół żartem, pół serio pół Polak.'
        counts = {}
        result = correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES, counts=counts)
        self.assertEqual(result, 'Czy by to nienajlepszy nibyartysta, super-Polak i półżartem, półserio pół-Polak.')
        self.assertEqual({key: counts.get(key) for key in ORTHOGRAPHY_2026_KEYS},
                         {'by_separate': 1, 'niby_quasi': 1, 'prefix_hyphen': 1, 'pol_compounds': 3,
                          'nie_comparative': 1})
        self.assertEqual(correct_editorial_text(text, FORM_DEFAULT_EDITORIAL_RULES - ORTHOGRAPHY_2026_KEYS), text)

    def test_long_paragraph_with_many_candidates_is_fast(self):
        """Wzorce bez kosztownych konstrukcji: najdłuższy akapit pełen kandydatów i wyjątków działa szybko."""
        source = ('Czyby to nie najlepszy niby-artysta. Nie lepiej, a gorzej. Pół żartem pół Polakiem, '
                  'pół Niemiec, superPolak, na nie łaskawiej, nie dzisiejszy złodziej, aby gdyby, ' * 200)[:20000]
        started = time.monotonic()
        result = correct_editorial_text(source, FORM_DEFAULT_EDITORIAL_RULES)
        self.assertLess(time.monotonic() - started, 1.5)
        self.assertIn('Czy by to nienajlepszy nibyartysta.', result)
