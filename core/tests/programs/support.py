"""Pomocniki testów programów: pliki DOCX w pamięci, izolowane katalogi robocze i zadania bez podprocesu."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image

from core.services import program_jobs as jobs

# Poprawne parametry formularza powtórzeń (prefiks dodaje formularz strony programów).
REPETITION_PARAMETERS = {
    'window_size': 35,
    'min_word_length': 4,
    'sentence_limit': 35,
    'paragraph_limit': 150,
}
REPETITION_FIELDS = {f'repetitions-{key}': value for key, value in REPETITION_PARAMETERS.items()}


def run_program(client, data, **extra):
    """Wysyła formularz programu i zwraca odpowiedź z pobraniem wyniku albo stronę z powodem, dla którego go nie ma.

    Widok zawsze tworzy zadanie w tle; tu zadanie biegnie w procesie testu (z prawdziwym konwerterem),
    a test przechodzi przez stronę stanu do pobrania – jak przeglądarka bez JavaScriptu.
    """
    with patch.object(jobs, 'launch', side_effect=jobs.execute):
        response = client.post(reverse('core:programs'), data, **extra)
    if response.status_code != 302:
        return response
    page = client.get(response['Location'])
    if page.status_code == 200 and page.context['state'].get('state') == 'done':
        return client.get(page.context['download_url'])
    return page


def saved(document):
    """Dokument python-docx zapisany do strumienia ustawionego na początek."""
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def docx_bytes(text='Ala  ma kota...'):
    document = Document()
    document.add_paragraph(text)
    return saved(document).getvalue()


def formatted_docx():
    """DOCX z kursywą 17 pt wyrównaną do lewej i wyśrodkowanym akapitem."""
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Ala  ma kota.').italic = True
    paragraph.runs[0].font.size = Pt(17)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph('Środek').alignment = WD_ALIGN_PARAGRAPH.CENTER
    return saved(document).getvalue()


def sized_docx(size):
    """Poprawny DOCX o dokładnie zadanym rozmiarze w bajtach."""
    stream = BytesIO(formatted_docx())
    with ZipFile(stream, 'a', compression=ZIP_STORED) as archive:
        overhead = 30 + 46 + 2 * len('padding.bin')
        archive.writestr('padding.bin', b'0' * (size - len(stream.getvalue()) - overhead))
    assert len(stream.getvalue()) == size
    return SimpleUploadedFile('tekst.docx', stream.getvalue())


def png(width, height):
    """Jednobarwny obraz PNG o zadanych wymiarach (strumień ustawiony na początek)."""
    stream = BytesIO()
    Image.new('RGB', (width, height), 'red').save(stream, 'PNG')
    stream.seek(0)
    return stream


def upload(data=None, name='tekst.docx'):
    return SimpleUploadedFile(name, docx_bytes() if data is None else data)


def downloaded(response):
    """Treść odpowiedzi strumieniowej (klient testowy sam ją zamyka)."""
    return b''.join(response.streaming_content)


def downloaded_document(response):
    return Document(BytesIO(downloaded(response)))


def typography_sample():
    """Akapit z kursywą, pogrubieniem i własnymi odstępami, tytuł wyśrodkowany stylem, akapit do prawej i tabela."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('Kursywa ').italic = True
    p.add_run('pogrubienie').bold = True
    p.paragraph_format.space_after = Pt(20)
    p.paragraph_format.line_spacing = 2
    doc.add_paragraph('Środek', style='Title')
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Prawo').alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela'
    return saved(doc)


NUMBERING_OFF = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'
# Reguła nagłówków i list (document_styles.ParagraphOutline): przypadek (tekst akapitu) → (styl akapitu, własne
# właściwości w:pPr, oczekiwany poziom nagłówka, czy punkt listy). Style przypadków tworzy outline_document().
OUTLINE_CASES = {
    'heading 1': ('Heading 1', '', 1, False),
    'heading 2': ('Heading 2', '', 2, False),
    'Nagłówek 1 bez poziomu konspektu': ('Nagłówek 1', '', 1, False),
    'Naglowek 1': ('Naglowek 1', '', 1, False),
    'Heading bez numeru': ('Heading', '', None, False),
    'Heading Custom': ('Heading Custom', '', None, False),
    'Nagłówek tabeli': ('Nagłówek tabeli', '', None, False),
    'styl oparty na heading 1': ('Tytuł opowiadania', '', 1, False),
    'styl Rozdział z poziomem konspektu 0': ('Rozdział', '', 1, False),
    'id Heading1, inna nazwa, bez poziomu': ('Tytuł rozdziału', '', 1, False),
    'zwykły akapit z poziomem konspektu 0': (None, '<w:outlineLvl w:val="0"/>', 1, False),
    'heading 1 z numId 0': ('Heading 1', NUMBERING_OFF, 1, False),
    'heading 1 z poziomem 9': ('Heading 1', '<w:outlineLvl w:val="9"/>', None, False),
    'numId 5': (None, '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="5"/></w:numPr>', None, True),
    'numPr bez numId': (None, '<w:numPr><w:ilvl w:val="0"/></w:numPr>', None, False),
    'styl listy z numId': ('Lista punktowana', '', None, True),
    'styl z numId 0 oparty na stylu listy': ('Lista bez numeracji', '', None, False),
    'zwykły akapit': (None, '', None, False),
    'pętla basedOn': ('Pętla A', '', None, False),
}


def outline_document(*, soft_break=False):
    """Dokument z akapitem na każdy przypadek OUTLINE_CASES (tekst akapitu = nazwa przypadku, z miękkim enterem
    i „drugi wiersz” przy soft_break). Wbudowany „heading 1” ma id Heading1orig – id Heading1 ma styl
    „Tytuł rozdziału” bez poziomu konspektu."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def properties(xml):
        return parse_xml(f'<w:pPr {nsdecls("w")}>{xml}</w:pPr>')

    document = Document()
    document.styles['Heading 1'].style_id = 'Heading1orig'
    styles = {}
    for name, style_id, based, own in (
            ('Nagłówek 1', 'Nagwek1', None, ''), ('Naglowek 1', 'Naglowek1', None, ''),
            ('Heading', 'Heading', None, ''), ('Heading Custom', 'HeadingCustom', None, ''),
            ('Nagłówek tabeli', 'Nagwektabeli', None, ''), ('Tytuł opowiadania', 'Tytulopowiadania', 'Heading 1', ''),
            ('Rozdział', 'Rozdzial', None, '<w:outlineLvl w:val="0"/>'), ('Tytuł rozdziału', 'Heading1', None, ''),
            ('Lista punktowana', 'Listapunktowana', None, '<w:numPr><w:numId w:val="7"/></w:numPr>'),
            ('Lista bez numeracji', 'Listabeznumeracji', 'Lista punktowana', '<w:numPr><w:numId w:val="0"/></w:numPr>'),
            ('Pętla A', 'PetlaA', None, ''), ('Pętla B', 'PetlaB', 'Pętla A', '')):
        style = styles[name] = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.style_id = style_id
        if based:
            style.base_style = document.styles[based]
        if own:
            style.element.append(properties(own))
    styles['Pętla A'].base_style = styles['Pętla B']
    for text, (style, own, _, _) in OUTLINE_CASES.items():
        paragraph = document.add_paragraph(text, style=style)
        for child in list(properties(own)):
            paragraph._p.get_or_add_pPr().append(child)
        if soft_break:
            paragraph.runs[0].add_break()
            paragraph.add_run('drugi wiersz')
    return document


FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">{}</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
)


def with_footnote(document, text):
    """DOCX z przypisem dolnym (python-docx nie umie ich tworzyć)."""
    raw = saved(document)
    output = BytesIO()
    with ZipFile(raw) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdFoot" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            target.writestr(item, data)
        target.writestr('word/footnotes.xml', FOOTNOTES.format(text))
    output.seek(0)
    return output


NOTE_PARTS = {
    'footnote': ('footnotes', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'),
    'endnote': ('endnotes', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes'),
}


def noted_document(parts, kind='footnote', note_text=' Treść  przypisu...', *, notes=None, note_runs=None,
                   reference='', mark=True):
    """DOCX z jednym akapitem: napisy to przebiegi tekstu, liczby – odsyłacze przypisu (w:id) z rStyle.

    notes – identyfikatory przypisów w części (domyślnie te, do których są odsyłacze); note_runs – surowy
    OOXML treści akapitu przypisu za znakiem przypisu (domyślnie przebieg z note_text); „{id}” w treści
    zastępuje numer przypisu; reference – dodatkowe atrybuty odsyłacza (np. w:customMarkFollows);
    mark=False – przypis bez znaku przypisu (w:footnoteRef) na początku treści.
    """
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    name, reltype = NOTE_PARTS[kind]
    style = 'FootnoteReference' if kind == 'footnote' else 'EndnoteReference'
    document = Document()
    paragraph = document.add_paragraph()
    for part in parts:
        if isinstance(part, int):
            paragraph._p.append(parse_xml(
                f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="{style}"/></w:rPr>'
                f'<w:{kind}Reference {reference} w:id="{part}"/></w:r>'))
        else:
            paragraph.add_run(part)
    if notes is None:
        notes = sorted({part for part in parts if isinstance(part, int)})
    content = note_runs if note_runs is not None else f'<w:r><w:t xml:space="preserve">{note_text}</w:t></w:r>'
    if mark:
        content = f'<w:r><w:rPr><w:rStyle w:val="{style}"/></w:rPr><w:{kind}Ref/></w:r>' + content
    notes_xml = ''.join(
        f'<w:{kind} w:id="{number}"><w:p>{content.replace("{id}", str(number))}</w:p></w:{kind}>'
        for number in notes)
    raw, output = saved(document), BytesIO()
    with ZipFile(raw) as source, ZipFile(output, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', (
                    f'<Override PartName="/word/{name}.xml" ContentType="application/vnd.openxmlformats-'
                    f'officedocument.wordprocessingml.{name}+xml"/></Types>').encode())
            if item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', (
                    f'<Relationship Id="rIdNotes" Type="{reltype}" Target="{name}.xml"/></Relationships>').encode())
            target.writestr(item, data)
        target.writestr(f'word/{name}.xml', (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:{name} {nsdecls("w")}>{notes_xml}</w:{name}>'))
    output.seek(0)
    return output


def noted_text(root):
    """Tekst drzewa XML części; odsyłacz jako [w:id], znak przypisu w jego treści jako [ref]."""
    from docx.oxml.ns import qn
    pieces = []
    for node in root.iter():
        if node.tag == qn('w:t'):
            pieces.append(node.text or '')
        elif node.tag in (qn('w:footnoteReference'), qn('w:endnoteReference')):
            pieces.append(f'[{node.get(qn("w:id"))}]')
        elif node.tag in (qn('w:footnoteRef'), qn('w:endnoteRef')):
            pieces.append('[ref]')
    return ''.join(pieces)


def footnote_text(stream):
    with ZipFile(stream) as archive:
        return archive.read('word/footnotes.xml').decode()


def whole_document():
    """Akapit, tabela, nagłówek i przypis – z podwójnymi spacjami i powtórzeniem „lampa”."""
    document = Document()
    document.add_paragraph('Akapit  główny.')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Lampa świeciła,  lampa gasła.'
    document.sections[0].header.paragraphs[0].text = 'Nagłówek  strony'
    return with_footnote(document, 'Przypis  z lampą i lampą.')


class IsolatedProgramDirsMixin:
    """Zadania programów i dziennik aktywności zapisują pliki w prywatnym katalogu testu."""

    def setUp(self):
        super().setUp()
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.work_dir = Path(temporary.name)
        override = self.settings(
            DOCUMENT_CONVERSION_DIR=self.work_dir / 'convert',
            ACTIVITY_SPOOL_DIR=self.work_dir / 'activity',
        )
        override.enable()
        self.addCleanup(override.disable)


def docx_with_paragraph_xml(*fragments):
    """DOCX z jednym akapitem złożonym z surowych fragmentów OOXML (runy, znaczniki proofErr, zakładki…).

    Fragmenty piszemy bez deklaracji przestrzeni nazw, np. ``'<w:r><w:t>kot</w:t></w:r>'``.
    """
    import re

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph()
    for fragment in fragments:
        paragraph._p.append(parse_xml(re.sub(r'^<(w:\w+)', r'<\1 ' + nsdecls('w'), fragment, count=1)))
    return saved(document)


def text_run(text):
    """Fragment OOXML runu z tekstem (spacje zachowane)."""
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def add_text_box(paragraph, text, *, fallback=True):
    """Pole tekstowe tak, jak zapisuje je Word: mc:AlternateContent z wersją DrawingML i (opcjonalnie) zapasową VML."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    box = f'<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>'
    choice = (
        '<mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic><a:graphicData>'
        '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:txbx>{box}</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
    )
    spare = f'<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>'
    paragraph._p.append(parse_xml(
        '<w:r %s xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent>%s%s</mc:AlternateContent></w:r>'
        % (nsdecls('w', 'wp', 'a'), choice, spare if fallback else '')))


def add_complex_field(paragraph, instruction, result):
    """Pole Worda w postaci złożonej: begin, instrukcja, separate, wynik, end."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    for body in ('<w:fldChar w:fldCharType="begin"/>',
                 f'<w:instrText xml:space="preserve"> {instruction} </w:instrText>',
                 '<w:fldChar w:fldCharType="separate"/>',
                 f'<w:t xml:space="preserve">{result}</w:t>',
                 '<w:fldChar w:fldCharType="end"/>'):
        paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}>{body}</w:r>'))


def add_footnote_reference(paragraph, note_id=1):
    """Odnośnik do przypisu dolnego (treść przypisu dodaje with_footnote)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="{note_id}"/></w:r>'))


def add_bookmark(paragraph, name, text):
    """Tekst objęty zakładką Worda (np. _Toc… albo _Hlk… tworzone przez Worda)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    for body in (f'<w:bookmarkStart {nsdecls("w")} w:id="0" w:name="{name}"/>',
                 f'<w:r {nsdecls("w")}><w:t xml:space="preserve">{text}</w:t></w:r>',
                 f'<w:bookmarkEnd {nsdecls("w")} w:id="0"/>'):
        paragraph._p.append(parse_xml(body))


def add_internal_link(paragraph, anchor, text):
    """Hiperłącze do zakładki w tym samym dokumencie."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}"><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:hyperlink>'))


def add_tracked_insertion(paragraph, text):
    """Tekst wstawiony przy włączonym śledzeniu zmian (w:ins)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    paragraph._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="9" w:author="Redakcja" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:ins>'))


def content_control(content, *, placeholder=False, gallery=None):
    """Kontrolka zawartości Worda (w:sdt) jako element OOXML do wstawienia w treść, akapit lub tabelę.

    ``content`` to surowe fragmenty bez deklaracji przestrzeni nazw (akapity, runy, wiersze albo komórki);
    ``placeholder`` – kontrolka pokazuje tekst zastępczy; ``gallery`` – blok z galerii Worda (np. spis treści).
    """
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    properties = '<w:alias w:val="Pole"/><w:id w:val="101"/>'
    if placeholder:
        properties += '<w:showingPlcHdr/>'
    if gallery:
        properties += (f'<w:docPartObj><w:docPartGallery w:val="{gallery}"/>'
                       '<w:docPartUnique/></w:docPartObj>')
    return parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr>{properties}</w:sdtPr>'
                     f'<w:sdtEndPr/><w:sdtContent>{content}</w:sdtContent></w:sdt>')


def add_block(document, element):
    """Element blokowy (np. kontrolka zawartości) na końcu treści dokumentu, przed ustawieniami sekcji."""
    document.element.body.insert_element_before(element, 'w:sectPr')
    return element


def add_simple_field(paragraph, instruction, result):
    """Pole Worda w postaci prostej (w:fldSimple) z zapisanym wynikiem."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    paragraph._p.append(parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr=" {instruction} "><w:r><w:t>{result}</w:t></w:r></w:fldSimple>'))


def add_gallery_page_number(story):
    """Numer strony tak, jak wstawia go galeria Worda: akapit z polem PAGE w kontrolce zawartości (w:sdt)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    story._element.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Page Numbers (Bottom of Page)"/>'
        '</w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:pPr><w:jc w:val="right"/></w:pPr>'
        '<w:fldSimple w:instr="PAGE   \\* MERGEFORMAT"><w:r><w:t>7</w:t></w:r></w:fldSimple>'
        '</w:p></w:sdtContent></w:sdt>'))


def add_page_break(paragraph, after=''):
    """Podział strony (Ctrl+Enter, w:br w:type="page") na końcu akapitu, opcjonalnie z dalszym tekstem."""
    from docx.enum.text import WD_BREAK
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    if after:
        paragraph.add_run(after)
    return paragraph


def set_language(properties_owner, value):
    """Oznacz język (w:lang w:val) przebiegu, stylu albo rPr – jak Word „Ustaw język sprawdzania”."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    element = getattr(properties_owner, '_r', None)
    if element is None:
        element = getattr(properties_owner, 'element', properties_owner)
    properties = element.get_or_add_rPr()
    lang = properties.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        properties.append(lang)
    lang.set(qn('w:val'), value)


def set_default_language(document, value):
    """Język w docDefaults (python-docx ma tam domyślnie en-US)."""
    from docx.oxml.ns import qn
    for lang in document.styles.element.iter(qn('w:lang')):
        if lang.getparent().getparent().tag == qn('w:rPrDefault'):
            lang.set(qn('w:val'), value)


def tracked_changes_document():
    """Wszystkie części dokumentu z poprawkami do śledzenia.

    Akapit z kilkoma przebiegami (pogrubienie, kursywa, tabulator), zastępowane słowo,
    istniejąca śledzona wstawka (w:id 9) i zakładka, seria pustych akapitów ze spacjami,
    tabela, pole tekstowe, nagłówek, stopka i przypis.
    """
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Ala  ma ').bold = True
    paragraph.add_run('kota...\tPotem karnistrami jechał.').italic = True
    add_bookmark(paragraph, '_Hlk1', ' Ona  wie.')
    add_tracked_insertion(paragraph, 'Dawna  wstawka')
    for text in ('', '  ', '\t', ''):
        document.add_paragraph(text)
    document.add_paragraph('Dnia 1.2.2025 r. wrócił.')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela  z kropkami...'
    add_text_box(document.add_paragraph('Przed polem.'), 'Pole  tekstowe...')
    document.sections[0].header.paragraphs[0].text = 'Nagłówek  strony'
    document.sections[0].footer.paragraphs[0].text = 'Stopka...'
    return with_footnote(document, 'Przypis  z tekstem...')


PICTURE_DRAWING = (
    '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="190500" cy="95250"/>'
    '<wp:docPr id="{id}" name="Obraz {id}"/><a:graphic><a:graphicData '
    'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr>'
    '<pic:cNvPr id="{id}" name="obraz.png"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
    '<a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr>'
    '<a:xfrm><a:off x="0" y="0"/><a:ext cx="190500" cy="95250"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
    '</a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing>'
)
IMAGE_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
)
IMAGE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'


def docx_with_images():
    """DOCX z obrazami w treści, nagłówku, przypisie i polu tekstowym oraz obrazem VML i obiektem OLE.

    Liczba obrazów do usunięcia: 6 (treść, pole tekstowe, VML, OLE, nagłówek, przypis). Pole tekstowe ma
    wersję zapasową VML z tym samym obrazem – liczy się raz. Tekst „Pole z obrazem” w polu tekstowym zostaje.
    """
    from docx.oxml import parse_xml
    document = Document()
    document.add_paragraph('Przed obrazem.')
    document.add_picture(png(20, 10))
    body_rid = next(rid for rid, rel in document.part.rels.items() if rel.reltype == IMAGE_REL)
    ole_rid = document.part.relate_to('embeddings/oleObject1.bin',
                                      'http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject',
                                      is_external=True)
    paragraph = document.add_paragraph('Pole: ')
    picture = PICTURE_DRAWING.format(id=7, rid=body_rid)
    paragraph._p.append(parse_xml(
        f'<w:r {IMAGE_NAMESPACES}><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline>'
        '<wp:docPr id="8" name="Pole"/><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent>'
        f'<w:p><w:r><w:t>Pole z obrazem</w:t></w:r><w:r>{picture}</w:r></w:p></w:txbxContent></wps:txbx>'
        '</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback><w:pict>'
        '<v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Pole z obrazem</w:t></w:r><w:r><w:pict><v:shape>'
        f'<v:imagedata r:id="{body_rid}"/></v:shape></w:pict></w:r></w:p></w:txbxContent></v:textbox></v:shape>'
        '</w:pict></mc:Fallback></mc:AlternateContent></w:r>'))
    paragraph._p.append(parse_xml(
        f'<w:r {IMAGE_NAMESPACES}><w:pict><v:shape style="width:10pt;height:10pt">'
        f'<v:imagedata r:id="{body_rid}" o:title=""/></v:shape></w:pict></w:r>'))
    paragraph._p.append(parse_xml(
        f'<w:r {IMAGE_NAMESPACES}><w:object><v:shape style="width:10pt;height:10pt"><v:imagedata r:id="{body_rid}"/>'
        f'</v:shape><o:OLEObject Type="Embed" ProgID="Package" ShapeID="_x0000_i1025" r:id="{ole_rid}"/></w:object></w:r>'))
    document.add_paragraph('Po obrazie.')
    header = document.sections[0].header
    header.paragraphs[0].text = 'Nagłówek '
    header.paragraphs[0].add_run().add_picture(png(20, 10))
    raw = with_footnote(document, 'Przypis z obrazem')
    output = BytesIO()
    with ZipFile(raw) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'word/document.xml.rels' or item.filename == 'word/_rels/document.xml.rels':
                # Obiekt OLE jest osadzony w pakiecie (relacja wewnętrzna).
                data = data.replace(b'TargetMode="External"', b'')
            if item.filename == 'word/footnotes.xml':
                data = data.replace(b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"',
                                    IMAGE_NAMESPACES.encode())
                data = data.replace(b'Przypis z obrazem</w:t></w:r>', b'Przypis z obrazem</w:t></w:r><w:r>'
                                    + PICTURE_DRAWING.format(id=9, rid='rIdFootImage').encode() + b'</w:r>')
            target.writestr(item, data)
        target.writestr('word/_rels/footnotes.xml.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdFootImage" '
            f'Type="{IMAGE_REL}" Target="media/footnote.png"/></Relationships>'))
        target.writestr('word/media/footnote.png', png(5, 5).getvalue())
        target.writestr('word/embeddings/oleObject1.bin', b'OLE')
    output.seek(0)
    return output


def add_leading_break(paragraph, kind='page'):
    """Podział strony albo kolumny (w:br w:type) na początku akapitu – tak Word zapisuje rozdział od nowej strony."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0,
                        parse_xml(f'<w:r {nsdecls("w")}><w:br w:type="{kind}"/></w:r>'))
    return paragraph


LINKED_IMAGE = 'file:///C:/Obrazy/zdjecie.jpg'
HYPERLINK_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'


def docx_with_linked_images(heading=False):
    """DOCX z obrazami połączonymi z plikami zewnętrznymi (relacje obrazów z TargetMode="External").

    Obrazy tylko połączone (7): rysunek wp:inline i wp:anchor z a:blip r:link, kształt VML (v:imagedata r:id
    z o:href), pole tekstowe z obrazem w mc:Choice i zapasowym VML (liczone raz), obraz w nagłówku, w stopce
    (VML) i w przypisie. Obraz osadzony i połączony w jednym a:blip (r:embed + r:link) – po usunięciu linku
    zostaje osadzony. Hiperłącze zewnętrzne „strona” obrazem nie jest. heading – pierwszy akapit to
    Nagłówek 1 (plik do podziału).
    """
    from docx.oxml import parse_xml
    document = Document()
    if heading:
        document.add_heading('Opowiadanie', level=1)
    document.add_paragraph('Przed obrazami.')
    document.add_picture(png(20, 10))
    linked = document.part.relate_to(LINKED_IMAGE, IMAGE_REL, is_external=True)
    document.paragraphs[-1]._p.xpath('.//a:blip')[0].set(
        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link',
        document.part.relate_to('file:///C:/Obrazy/oryginal.png', IMAGE_REL, is_external=True))
    inline = PICTURE_DRAWING.format(id=7, rid=linked).replace('r:embed=', 'r:link=')
    anchor = inline.replace('<wp:inline distT="0" distB="0" distL="0" distR="0">', (
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" '
        'locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom='
        '"column"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph">'
        '<wp:posOffset>0</wp:posOffset></wp:positionV>')).replace('</wp:inline>', '</wp:anchor>').replace(
        '<wp:docPr id="7" name="Obraz 7"/>', '<wp:wrapNone/><wp:docPr id="8" name="Obraz 8"/>')
    vml = (f'<w:pict><v:shape style="width:10pt;height:10pt"><v:imagedata r:id="{linked}" o:href="{LINKED_IMAGE}" '
           'o:title=""/></v:shape></w:pict>')
    paragraph = document.add_paragraph('Obrazy: ')
    for content in (inline, anchor, vml):
        paragraph._p.append(parse_xml(f'<w:r {IMAGE_NAMESPACES}>{content}</w:r>'))
    link = document.part.relate_to('https://example.test/', HYPERLINK_REL, is_external=True)
    paragraph._p.append(parse_xml(f'<w:hyperlink {IMAGE_NAMESPACES} r:id="{link}"><w:r><w:t>strona</w:t></w:r>'
                                  '</w:hyperlink>'))
    box = (f'<w:p><w:r><w:t>Pole z obrazem</w:t></w:r><w:r>{PICTURE_DRAWING.format(id=9, rid=linked)}</w:r></w:p>'
           ).replace('r:embed=', 'r:link=')
    spare = f'<w:p><w:r><w:t>Pole z obrazem</w:t></w:r><w:r>{vml}</w:r></w:p>'
    document.add_paragraph()._p.append(parse_xml(
        f'<w:r {IMAGE_NAMESPACES}><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline>'
        '<wp:docPr id="10" name="Pole"/><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/'
        f'wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent>{box}</w:txbxContent></wps:txbx></wps:wsp>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox>'
        f'<w:txbxContent>{spare}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>'
        '</w:r>'))
    section = document.sections[0]
    for story, text, content in ((section.header, 'Nagłówek ', PICTURE_DRAWING.format(id=11, rid='{rid}').replace(
            'r:embed=', 'r:link=')), (section.footer, 'Stopka ', vml.replace(linked, '{rid}'))):
        story.paragraphs[0].text = text
        rid = story.part.relate_to(LINKED_IMAGE, IMAGE_REL, is_external=True)
        story.paragraphs[0]._p.append(parse_xml(f'<w:r {IMAGE_NAMESPACES}>{content.replace("{rid}", rid)}</w:r>'))
    raw, output = with_footnote(document, 'Przypis z obrazem'), BytesIO()
    with ZipFile(raw) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'word/footnotes.xml':
                data = data.replace(b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"',
                                    IMAGE_NAMESPACES.encode())
                data = data.replace(b'Przypis z obrazem</w:t></w:r>', b'Przypis z obrazem</w:t></w:r><w:r>' + (
                    PICTURE_DRAWING.format(id=12, rid='rIdNoteImage').replace('r:embed=', 'r:link=').encode())
                    + b'</w:r>')
            target.writestr(item, data)
        target.writestr('word/_rels/footnotes.xml.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.'
            f'openxmlformats.org/package/2006/relationships"><Relationship Id="rIdNoteImage" Type="{IMAGE_REL}" '
            f'Target="{LINKED_IMAGE}" TargetMode="External"/></Relationships>'))
    output.seek(0)
    return output


def external_images(stream):
    """Relacje obrazów z TargetMode="External" we wszystkich częściach pakietu: [(plik .rels, adres)]."""
    from lxml import etree
    found = []
    with ZipFile(stream) as archive:
        for name in archive.namelist():
            if name.endswith('.rels'):
                for relationship in etree.fromstring(archive.read(name)):
                    if relationship.get('TargetMode') == 'External' and relationship.get('Type').endswith('/image'):
                        found.append((name, relationship.get('Target')))
    stream.seek(0)
    return found
