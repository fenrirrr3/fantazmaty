"""Przygotowanie dokumentów: dziedziczenie stylów, reguła nagłówków i list, normalizacja typografii i strony,
wyjustowanie, kolejność kroków przygotowania oraz zamiana miękkich enterów i twardych spacji (z kontrolkami,
zmianami śledzonymi i komentarzami)."""
import json
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase, tag
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

from core.services.document_conversion_worker import main as conversion_worker_main
from core.services.document_formatting import normalize_docx
from core.services.document_preparation import prepare_docx
from core.services.document_styles import ParagraphOutline, document_outline, paragraph_properties
from core.services.document_whitespace import normalize_spacing_docx
from core.services.mailbox_import import _package_messages
from core.services.odkurzacz import clean_docx

from .editor_fixtures_exotic import p, r
from .support import (
    OUTLINE_CASES, IsolatedProgramDirsMixin, content_control, outline_document, saved, typography_sample,
)
from .test_document_rebuild import (
    NS, comment_anchors, comment_reference, comment_xml, resolved, review_docx_with, review_markup_errors,
    review_parts,
)


PARA_ID = '{http://schemas.microsoft.com/office/word/2010/wordml}paraId'
W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'


def justification_sample():
    document = Document()
    paragraph = document.add_paragraph('Zwykły tekst ')
    paragraph.add_run('kursywa').italic = True
    paragraph.add_run(' gruby').bold = True
    document.add_paragraph('Do prawej').alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph('Pośrodku').alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('Środek ze stylu', style='Title')
    document.add_paragraph('')
    return saved(document)


def indent(paragraph):
    node = paragraph._p.find(f'{qn("w:pPr")}/{qn("w:ind")}')
    return None if node is None else node.get(qn('w:firstLine'))


def alignment(paragraph):
    node = paragraph._p.find(f'{qn("w:pPr")}/{qn("w:jc")}')
    return None if node is None else node.get(qn('w:val'))


def xml_paragraphs(stream):
    with ZipFile(stream) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    return root.find(qn('w:body')).findall(qn('w:p'))


def text(paragraph):
    return ''.join(paragraph.itertext())


class ParagraphStyleTests(SimpleTestCase):
    def test_direct_zero_and_false_override_inherited_values(self):
        document = Document()
        base = document.styles.add_style('Parent', WD_STYLE_TYPE.PARAGRAPH)
        base.paragraph_format.space_after = Pt(18)
        base.paragraph_format.page_break_before = True
        base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        child = document.styles.add_style('Child', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        child.paragraph_format.line_spacing = 2
        paragraph = document.add_paragraph('Test', style=child)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.page_break_before = False
        expected = {'space_after': Pt(0), 'page_break_before': False,
                    'alignment': WD_ALIGN_PARAGRAPH.RIGHT, 'line_spacing': 2}
        self.assertEqual(paragraph_properties(paragraph, expected), expected)

    def test_alignment_values_unknown_to_python_docx_are_resolved(self):
        """Regresja: w:jc „start”/„end” (OOXML Strict z LibreOffice, Google Docs) w akapicie albo stylu
        przerywało konwersję do EPUB/PDF wyjątkiem python-docx; nieznana wartość liczy się jako brak."""
        cases = (
            ('start w akapicie', 'start', None, WD_ALIGN_PARAGRAPH.LEFT),
            ('end w akapicie', 'end', None, WD_ALIGN_PARAGRAPH.RIGHT),
            ('end w stylu', None, 'end', WD_ALIGN_PARAGRAPH.RIGHT),
            ('nieznane w akapicie, styl do środka', 'zzz', 'center', WD_ALIGN_PARAGRAPH.CENTER),
            ('nieznane w stylu', None, 'zzz', None),
        )
        for case, own, in_style, expected in cases:
            with self.subTest(case):
                document = Document()
                style = document.styles.add_style('Własny', WD_STYLE_TYPE.PARAGRAPH)
                if in_style:
                    style.element.get_or_add_pPr().append(OxmlElement('w:jc', {qn('w:val'): in_style}))
                paragraph = document.add_paragraph('Tekst', style=style)
                if own:
                    paragraph._p.get_or_add_pPr().append(OxmlElement('w:jc', {qn('w:val'): own}))
                self.assertEqual(paragraph_properties(paragraph, ('alignment', 'page_break_before')),
                                 {'alignment': expected, 'page_break_before': None})

    def test_heading_level_and_list_item_follow_one_rule(self):
        """Reguła nagłówków i list (ParagraphOutline) dla przypadków OUTLINE_CASES: poziom konspektu akapitu, potem
        styl z poziomem konspektu albo nazwą lub id nagłówka z numerem (z dziedziczeniem, bez pętli); poziom 9 znosi
        nagłówek, numeracja na niego nie wpływa; lista – pierwszy numId akapitu albo stylu, numId 0 jej nie daje.
        Bez części stylów nagłówkiem jest tylko styl, którego id pasuje do wzorca."""
        document = outline_document()
        outline = document_outline(document)
        for paragraph in document.paragraphs:
            _, _, level, listed = OUTLINE_CASES[paragraph.text]
            with self.subTest(paragraph.text):
                self.assertEqual((outline.heading_level(paragraph._p), outline.is_list_item(paragraph._p)),
                                 (level, listed))
        for style_id, level in (('Heading1', 1), ('Nagwek2', 2), ('NAGLOWEK3', 3), ('Tytul', None), ('', None)):
            with self.subTest(bez_stylów=style_id):
                paragraph = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="{style_id}"/></w:pPr></w:p>')
                self.assertEqual(ParagraphOutline().heading_level(paragraph), level)

    def test_missing_properties_in_cyclic_styles_remain_unset(self):
        document = Document()
        one = document.styles.add_style('CycleOne', WD_STYLE_TYPE.PARAGRAPH)
        two = document.styles.add_style('CycleTwo', WD_STYLE_TYPE.PARAGRAPH)
        one.base_style = two
        two.base_style = one
        two.paragraph_format.line_spacing = 2
        paragraph = document.add_paragraph('Test', style=one)
        self.assertEqual(paragraph_properties(paragraph, ('line_spacing', 'page_break_before')),
                         {'line_spacing': 2, 'page_break_before': None})

    def test_cyclic_style_chain_is_bounded_during_normalization(self):
        """Regresja: cykl stylów bazowych zapętlał normalizację dokumentu."""
        document = Document()
        a = document.styles.add_style('CycleA', 1)
        b = document.styles.add_style('CycleB', 1)
        a.base_style = b
        b.base_style = a
        document.add_paragraph('Tekst', style=a)
        with normalize_docx(saved(document)) as result:
            self.assertEqual(Document(result).paragraphs[0].text, 'Tekst')


@tag('docx')
class NormalizationTests(SimpleTestCase):
    def test_typography_preserves_runs_and_inherited_alignment(self):
        with normalize_docx(typography_sample()) as out:
            doc = Document(out)
        first, center, right = doc.paragraphs
        self.assertTrue(first.runs[0].italic)
        self.assertTrue(first.runs[1].bold)
        self.assertEqual(first.runs[0].font.name, 'Times New Roman')
        self.assertEqual(first.runs[0].font.size.pt, 12)
        self.assertEqual(first.paragraph_format.line_spacing, 1.5)
        self.assertEqual(first.paragraph_format.space_after.pt, 0)
        self.assertEqual(first.paragraph_format.space_before.pt, 0)
        self.assertAlmostEqual(first.paragraph_format.first_line_indent.cm, 1.25, places=2)
        self.assertEqual(center.paragraph_format.first_line_indent, 0)
        self.assertEqual(center.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(right.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(doc.tables[0].cell(0, 0).paragraphs[0].runs[0].font.size.pt, 12)

    def test_added_properties_follow_the_word_schema_order(self):
        """Regresja: odstępy, wcięcie i rozmiar czcionki trafiały za w:rPr/w:lang, a Word odrzuca taki DOCX."""
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run('Tekst')
        run.font.underline = True
        run._r.get_or_add_rPr().append(OxmlElement('w:lang'))
        paragraph._p.get_or_add_pPr().append(OxmlElement('w:rPr'))
        with normalize_docx(saved(doc), justify=True) as output:
            result = Document(output).paragraphs[0]._p
        names = [child.tag.rsplit('}', 1)[1] for child in result.pPr]
        self.assertEqual(names, ['spacing', 'ind', 'contextualSpacing', 'jc', 'rPr'])
        names = [child.tag.rsplit('}', 1)[1] for child in result.r_lst[0].rPr]
        self.assertEqual(names, ['rFonts', 'sz', 'szCs', 'u', 'lang'])

    def test_page_normalization_all_sections(self):
        doc = Document()
        doc.add_paragraph('A')
        doc.add_section()
        for section in doc.sections:
            section.page_width = Cm(40)
            section.page_height = Cm(20)
            section.top_margin = section.bottom_margin = Cm(1)
        result = Document(normalize_docx(saved(doc)))
        for section in result.sections:
            self.assertAlmostEqual(section.page_width.cm, 21, places=2)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=2)
            for name in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
                self.assertAlmostEqual(getattr(section, name).cm, 2.5, places=2)

    def test_formatting_happens_before_cleaning(self):
        observed = []

        def cleaner(source, rules, **options):
            observed.append(Document(source).paragraphs[0].runs[0].font.size.pt)
            source.seek(0)
            return clean_docx(source, rules, **options)

        with (
            TemporaryDirectory() as tmp,
            patch('core.services.document_preparation.clean_docx', side_effect=cleaner),
        ):
            directory = Path(tmp)
            (directory / 'source.docx').write_bytes(typography_sample().getvalue())
            (directory / 'job.json').write_text(
                json.dumps({'formats': ['epub'], 'prepare': True, 'clean': True, 'title': 'Test'}))
            with patch('sys.argv', ['worker', tmp]):
                self.assertEqual(conversion_worker_main(), 0)
        self.assertEqual(observed, [12])

    def test_style_of_unknown_type_is_not_a_heading(self):
        """Styl z błędnym w:type (uszkodzony DOCX): zwykły akapit – z własnym wyrównaniem albo bez niego
        (wcześniej akapit bez własnego wyrównania przerywał ujednolicanie błędem)."""
        cases = {'own alignment': (WD_ALIGN_PARAGRAPH.RIGHT, 'right'), 'no own alignment': (None, 'both')}
        for case, (own, expected) in cases.items():
            document = Document()
            style = document.styles.add_style('Uszkodzony', WD_STYLE_TYPE.PARAGRAPH)
            style.element.set(qn('w:type'), 'nieznany')
            paragraph = document.add_paragraph('Tekst')
            paragraph.alignment = own
            paragraph._p.get_or_add_pPr().insert(0, OxmlElement('w:pStyle', {qn('w:val'): 'Uszkodzony'}))
            with self.subTest(case):
                result = Document(normalize_docx(saved(document), justify=True)).paragraphs[0]
                self.assertEqual((alignment(result), indent(result)), (expected, '709'))

    def test_theme_fonts_are_replaced_by_times_new_roman(self):
        document = Document()
        run = document.add_paragraph().add_run('Motyw')
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:asciiTheme'), 'minorHAnsi')
        fonts = Document(normalize_docx(saved(document))).paragraphs[0].runs[0]._r.rPr.rFonts
        self.assertIsNone(fonts.get(qn('w:asciiTheme')))
        self.assertEqual(fonts.get(qn('w:ascii')), 'Times New Roman')

    def test_previous_formatting_of_tracked_changes_is_normalized_too(self):
        """Poprzednie formatowanie zmiany dostaje to samo ujednolicenie (czcionka, odstępy, wcięcie), więc zmiana
        pokazuje tylko pracę redaktora. Regresja: w:rFonts dopisywany za w:rPrChange (python-docx nie zna jej
        miejsca w schemacie) – Word uznaje taki plik za uszkodzony; pusty w:pPrChange z LibreOffice dostaje w:pPr,
        a jego miejsce przed w:rPr w w:pPr (zapis LibreOffice) jest poprawiane."""
        change = f'w:id="1" w:author="Redaktor" {nsdecls("w")}'
        cases = (
            # nazwa, w:pPr akapitu, w:rPr przebiegu, oczekiwane dzieci w:pPr, w:rPr i poprzednich właściwości
            ('sama zmiana formatowania znaków', '', f'<w:rPrChange {change}><w:rPr><w:i/></w:rPr></w:rPrChange>',
             None, ['rFonts', 'sz', 'szCs', 'rPrChange'], ['rFonts', 'i', 'sz', 'szCs']),
            ('zmiana akapitu z LibreOffice', f'<w:jc w:val="center"/><w:pPrChange {change}/><w:rPr/>', '',
             ['spacing', 'ind', 'contextualSpacing', 'jc', 'rPr', 'pPrChange'], None,
             ['spacing', 'ind', 'contextualSpacing']),
        )
        for label, paragraph_xml, run_xml, paragraph_children, run_children, previous_children in cases:
            with self.subTest(label):
                document = Document()
                paragraph = document.add_paragraph()
                run = paragraph.add_run('Tekst')
                if paragraph_xml:
                    properties = parse_xml(f'<w:pPr {nsdecls("w")}>{paragraph_xml}</w:pPr>')
                    paragraph._p.insert(0, properties)
                if run_xml:
                    run._r.insert(0, parse_xml(f'<w:rPr {nsdecls("w")}>{run_xml}</w:rPr>'))
                result = Document(normalize_docx(saved(document))).paragraphs[0]

                def names(node):
                    return [child.tag.split('}')[1] for child in node]

                if paragraph_children is not None:
                    self.assertEqual(names(result._p.pPr), paragraph_children)
                    previous = result._p.pPr.find(f'{qn("w:pPrChange")}/{qn("w:pPr")}')
                    self.assertEqual(previous.find(qn('w:ind')).get(qn('w:firstLine')), '709')
                else:
                    self.assertEqual(names(result.runs[0]._r.rPr), run_children)
                    previous = result.runs[0]._r.rPr.find(f'{qn("w:rPrChange")}/{qn("w:rPr")}')
                    self.assertEqual(previous.find(qn('w:rFonts')).get(qn('w:ascii')), 'Times New Roman')
                self.assertEqual(names(previous), previous_children)


@tag('docx')
class NormalizedAlignmentTests(SimpleTestCase):
    """Wyjustowanie tylko tekstu wyrównanego do lewej, bez nagłówków i stopek; pełne marginesy strony."""

    def test_only_left_or_unaligned_paragraphs_are_justified_and_headers_are_skipped(self):
        document = Document()
        document.add_paragraph('Bez wyrównania')
        document.add_paragraph('Do lewej').alignment = WD_ALIGN_PARAGRAPH.LEFT
        document.add_paragraph('Do prawej').alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.add_paragraph('Wyjustowany').alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        document.styles['Quote'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.add_paragraph('Prawo ze stylu', style='Quote')
        document.sections[0].header.paragraphs[0].text = 'Nagłówek'
        document.sections[0].footer.paragraphs[0].text = 'Stopka'
        document.sections[0].footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        result = Document(normalize_docx(saved(document), justify=True))
        self.assertEqual([paragraph.alignment for paragraph in result.paragraphs],
                         [WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT,
                          WD_ALIGN_PARAGRAPH.JUSTIFY, None])
        self.assertIsNone(result.sections[0].header.paragraphs[0].alignment)
        self.assertEqual(result.sections[0].footer.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_alignment_values_unknown_to_python_docx_are_understood(self):
        """w:jc „start”/„end” (nowsze edytory) w akapicie albo w stylu nie przerywa ujednolicania: „end” to
        wyrównanie do prawej (zostaje), „start” – do lewej (justowane); wcięcie pierwszego wiersza w obu."""
        document = Document()
        for value in ('start', 'end'):
            document.add_paragraph(f'akapit {value}')._p.get_or_add_pPr().append(
                OxmlElement('w:jc', {qn('w:val'): value}))
            style = document.styles.add_style(f'Styl {value}', WD_STYLE_TYPE.PARAGRAPH)
            style.element.get_or_add_pPr().append(OxmlElement('w:jc', {qn('w:val'): value}))
            document.add_paragraph(f'styl {value}', style=style)
        result = Document(normalize_docx(saved(document), justify=True))
        self.assertEqual({paragraph.text: (alignment(paragraph), indent(paragraph)) for paragraph in result.paragraphs},
                         {'akapit start': ('both', '709'), 'akapit end': ('end', '709'),
                          'styl start': ('both', '709'), 'styl end': (None, '709')})

    def test_created_page_margins_have_header_and_footer_distance(self):
        document = Document()
        document.add_paragraph('Tekst')
        section = document.sections[0]._sectPr
        section.remove(section.find(qn('w:pgMar')))
        result = Document(normalize_docx(saved(document)))
        margins = result.sections[0]._sectPr.find(qn('w:pgMar'))
        self.assertEqual((margins.get(qn('w:header')), margins.get(qn('w:footer'))), ('708', '708'))
        self.assertEqual(margins.get(qn('w:top')), '1417')
        # Istniejące odległości nagłówka i stopki zostają.
        document = Document()
        document.sections[0].header_distance = Cm(2)
        margins = Document(normalize_docx(saved(document))).sections[0]._sectPr.find(qn('w:pgMar'))
        self.assertEqual(margins.get(qn('w:header')), str(Cm(2).twips))


@tag('docx')
class JustificationTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def check_document(self, stream):
        doc = Document(stream)
        self.assertEqual(len(doc.paragraphs), 5)
        self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        # Wyrównanie do prawej jest zamierzone (podpis, data) – zostaje.
        self.assertEqual(doc.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        for paragraph in doc.paragraphs[2:4]:
            self.assertEqual(paragraph.paragraph_format.first_line_indent, 0)
            self.assertTrue(
                paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
                or paragraph.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
            )
        self.assertTrue(doc.paragraphs[0].runs[1].italic)
        self.assertTrue(doc.paragraphs[0].runs[2].bold)
        self.assertEqual(doc.paragraphs[4].text, '')

    def test_shared_preparation_only_justifies_when_requested(self):
        with prepare_docx(justification_sample(), normalize_formatting=True, justify=True) as output:
            self.check_document(output)
        with prepare_docx(justification_sample(), normalize_formatting=True) as output:
            self.assertEqual(Document(output).paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_mail_conversion_justifies_docx_pdf_and_epub(self):
        row = {'uid': 1, 'folder': 'Tekst', 'files': [('tekst.docx', justification_sample().getvalue())]}
        with _package_messages([row], clean=True, convert=True) as result, ZipFile(result) as archive:
            self.check_document(BytesIO(archive.read('Tekst/01_tekst.docx')))
            self.assertTrue(archive.read('Tekst/01_tekst.pdf').startswith(b'%PDF'))
            with ZipFile(BytesIO(archive.read('Tekst/01_tekst.epub'))) as epub:
                content = epub.read('EPUB/content.xhtml').decode()
                self.assertIn('align-justify', content)
                self.assertIn('align-center', content)


@tag('docx')
class PreparationTests(SimpleTestCase):
    def test_unknown_cleaner_rules_are_refused(self):
        for rules in (['nieznana'], 'spaces'):
            with self.subTest(rules=rules), self.assertRaisesMessage(ValueError, 'Invalid cleaner rules'):
                prepare_docx(typography_sample(), cleaner_rules=rules)

    def test_rebuild_runs_first_and_drops_private_metadata(self):
        document = Document()
        document.core_properties.author = 'Prywatny'
        document.add_paragraph('Ala  ma kota')
        with prepare_docx(saved(document), rebuild=True, cleaner_rules=['spaces']) as output:
            result = Document(output)
        self.assertEqual(result.core_properties.author, '')
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota')


@tag('docx')
class SoftWhitespaceTests(SimpleTestCase):
    """Miękki enter (także U+2028) tworzy nowy akapit; właściwości akapitu trafiają tylko tam, gdzie mają sens."""

    def normalized(self, document):
        result = Document(normalize_spacing_docx(saved(document)))
        return [Paragraph(element, result) for element in result.element.body.iter(qn('w:p'))]

    def test_line_separator_splits_paragraph_and_keeps_section_and_page_break_once(self):
        document = Document()
        paragraph = document.add_paragraph('Pierwszy\u2028Drugi\u00a0wiersz')
        paragraph.paragraph_format.page_break_before = True
        paragraph._p.set(PARA_ID, '1234ABCD')
        paragraph._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        first, second = self.normalized(document)[:2]
        self.assertEqual((first.text, second.text), ('Pierwszy', 'Drugi wiersz'))
        self.assertTrue(first.paragraph_format.page_break_before)
        self.assertIsNone(second.paragraph_format.page_break_before)
        self.assertIsNone(first._p.pPr.find(qn('w:sectPr')))
        self.assertIsNotNone(second._p.pPr.find(qn('w:sectPr')))
        # Identyfikator akapitu Worda zostaje tylko przy pierwszej części.
        self.assertEqual((first._p.get(PARA_ID), second._p.get(PARA_ID)), ('1234ABCD', None))

    def test_text_box_paragraph_inside_a_split_paragraph_stays_whole(self):
        document = Document()
        outer = document.add_paragraph('Przed')
        outer.add_run().add_break()
        box = OxmlElement('w:txbxContent')
        inner = OxmlElement('w:p')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Pole tekstowe'
        run.append(text)
        inner.append(run)
        box.append(inner)
        outer.add_run('Po')._r.append(box)
        before, after, boxed = self.normalized(document)
        self.assertEqual([before.text, after.text, boxed.text], ['Przed', 'Po', 'Pole tekstowe'])
        self.assertEqual(len(after._p.xpath('.//w:txbxContent/w:p')), 1)

    def test_soft_break_in_deleted_text_does_not_split_the_paragraph(self):
        """Regresja: miękki enter w usuniętym (śledzonym) tekście tworzył nowy akapit z reszty zdania."""
        document = Document()
        paragraph = document.add_paragraph('Początek ')
        deleted = OxmlElement('w:del', {qn('w:id'): '1', qn('w:author'): 'Redakcja'})
        run = OxmlElement('w:r')
        for name, text in (('w:delText', 'usunięte'), ('w:br', None), ('w:delText', 'zdanie')):
            child = OxmlElement(name)
            child.text = text
            run.append(child)
        deleted.append(run)
        paragraph._p.append(deleted)
        paragraph.add_run('koniec')
        (result,) = self.normalized(document)
        self.assertEqual(result.text, 'Początek koniec')
        self.assertEqual(len(result._p.xpath('.//w:del//w:br')), 1)

    def test_soft_break_stays_in_headings_and_list_items(self):
        """Miękki enter zostaje w nagłówku i punkcie listy według reguły (OUTLINE_CASES) – także w nagłówku
        z wyłączoną numeracją (numId 0, zapis LibreOffice), stylu opartym na nagłówku i stylu z poziomem konspektu;
        w nagłówku z poziomem 9, stylu „Heading” bez numeru i akapicie z numPr bez numId dzieli akapit."""
        document = outline_document(soft_break=True)
        plain = document.add_paragraph('Zwykły')
        plain.runs[0].add_break()
        plain.add_run('akapit\u2007tekstu')
        expected = []
        for text, (_, _, level, listed) in OUTLINE_CASES.items():
            expected += [f'{text}\ndrugi wiersz'] if level or listed else [text, 'drugi wiersz']
        self.assertEqual([paragraph.text for paragraph in self.normalized(document)],
                         expected + ['Zwykły', 'akapit tekstu'])

    def test_split_tracked_insertion_and_paragraph_mark_get_unique_ids(self):
        document = Document()
        paragraph = document.add_paragraph('Początek ')
        mark = OxmlElement('w:rPr')
        mark.append(OxmlElement('w:ins', {qn('w:id'): '4', qn('w:author'): 'Redakcja',
                                          qn('w:date'): '2026-01-01T00:00:00Z'}))
        paragraph._p.get_or_add_pPr().append(mark)
        inserted = OxmlElement('w:ins', {qn('w:id'): '7', qn('w:author'): 'Redakcja',
                                         qn('w:date'): '2026-01-01T00:00:00Z'})
        run = OxmlElement('w:r')
        for name, text in (('w:t', 'wstawka'), ('w:br', None), ('w:t', 'dalej'), ('w:br', None), ('w:t', 'koniec')):
            child = OxmlElement(name)
            child.text = text
            run.append(child)
        inserted.append(run)
        paragraph._p.append(inserted)
        result = self.normalized(document)
        self.assertEqual([''.join(p._p.xpath('.//w:t/text()')) for p in result], ['Początek wstawka', 'dalej', 'koniec'])
        ids = [node.get(qn('w:id')) for p in result for node in p._p.iter() if node.get(qn('w:id')) is not None]
        self.assertEqual(len(ids), 6)
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue({'4', '7'} <= set(ids))

    def test_oversized_package_is_refused(self):
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            for index in range(2001):
                archive.writestr(f'part{index}.xml', b'')
        with self.assertRaisesMessage(ValueError, 'limit rozpakowanej zawartości'):
            normalize_spacing_docx(stream)

    def test_content_control_split_by_a_soft_break_gets_a_new_id(self):
        """Kopia kontrolki w nowym akapicie dostaje identyfikator większy od wszystkich w dokumencie – także od
        identyfikatorów innych kontrolek (<w:id w:val>), a nie tylko od atrybutów w:id."""
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(content_control('<w:r><w:t>Przed</w:t><w:br/><w:t>po</w:t></w:r>'))
        other = content_control('<w:r><w:t>Inna</w:t></w:r>')
        other.find(f'{qn("w:sdtPr")}/{qn("w:id")}').set(qn('w:val'), '1')
        document.add_paragraph()._p.append(other)
        first, second, third = xml_paragraphs(normalize_spacing_docx(saved(document)))
        self.assertEqual((text(first), text(second), text(third)), ('Przed', 'po', 'Inna'))
        ids = [p.find(f'{qn("w:sdt")}/{qn("w:sdtPr")}/{qn("w:id")}').get(qn('w:val')) for p in (first, second, third)]
        self.assertEqual((ids[0], ids[2]), ('101', '1'))
        self.assertEqual(ids[1], '102')

    def test_package_without_styles_and_paragraph_identity(self):
        """Bez word/styles.xml żaden styl nie jest nagłówkiem; druga część traci tylko w14:paraId."""
        document = Document()
        paragraph = document.add_paragraph('Pierwsza')
        paragraph.add_run().add_break()
        paragraph.add_run('druga')
        paragraph._p.set(f'{{{W14}}}paraId', '1A2B3C4D')
        paragraph._p.set(qn('w:rsidR'), '00AB12CD')
        source, stripped = saved(document), BytesIO()
        with ZipFile(source) as archive, ZipFile(stripped, 'w') as target:
            for item in archive.infolist():
                if item.filename != 'word/styles.xml':
                    target.writestr(item, archive.read(item))
        first, second = xml_paragraphs(normalize_spacing_docx(stripped))
        self.assertEqual((text(first), text(second)), ('Pierwsza', 'druga'))
        self.assertEqual(first.get(f'{{{W14}}}paraId'), '1A2B3C4D')
        self.assertIsNone(second.get(f'{{{W14}}}paraId'))
        self.assertEqual(second.get(qn('w:rsidR')), '00AB12CD')

    def test_tracked_changes_and_comments_survive_the_split(self):
        """Regresja: miękki enter we wstawce dawał zwykły znak akapitu (odrzucenie wstawki zostawiało pusty akapit) –
        teraz znak akapitu jest wstawiony przez tę samą zmianę; usunięty znak akapitu źródła był kopiowany do każdej
        części (przyjęcie zmian sklejało podzielone akapity) – teraz zostaje przy ostatniej; spacja nierozdzielająca
        w tekście przeniesionym skąd (w:moveFrom ma w:t) była zamieniana – teraz zostaje, a we wstawce jest zamieniana.
        Miękki enter w usuniętym tekście nadal nie dzieli akapitu."""
        meta = 'w:author="Autorka" w:date="2026-01-01T00:00:00Z"'
        cases = (
            # nazwa, treść, (teksty po przyjęciu, po odrzuceniu), XPath w wyniku
            ('we wstawce', p(r('A'), f'<w:ins w:id="1" {meta}>' + r('B') + '<w:r><w:br/></w:r>' + r('C') + '</w:ins>'),
             ([['AB', 'C']], [['A']]), './w:body/w:p[1]/w:pPr/w:rPr/w:ins[@w:author="Autorka"]'),
            ('na początku wstawki', p(r('A'), f'<w:ins w:id="1" {meta}><w:r><w:br/></w:r>' + r('B') + '</w:ins>'),
             ([['A', 'B']], [['A']]), 'not(./w:body/w:p[1]/w:ins)'),  # bez pustej wstawki w pierwszym akapicie
            ('usunięty znak akapitu', p(r('A'), '<w:r><w:br/></w:r>', r('B'), properties=f'<w:rPr><w:del w:id="1" '
                                        f'{meta}/></w:rPr>') + p(r('C')),
             ([['A', 'BC']], [['A', 'B', 'C']]), './w:body/w:p[2]/w:pPr/w:rPr/w:del[@w:id="1"]'),
            ('usunięty tekst', p(r('A'), f'<w:del w:id="1" {meta}><w:r><w:delText>B</w:delText><w:br/>'
                                 '<w:delText>C</w:delText></w:r></w:del>'),
             ([['A']], [['AB\nC']]), './w:body/w:p[1]/w:del/w:r/w:br'),
            ('spacje nierozdzielające', p(r('A\u00a0B'), f'<w:moveFrom w:id="1" {meta}>' + r('E\u00a0F') + '</w:moveFrom>',
                                          f'<w:ins w:id="2" {meta}>' + r('G\u00a0H') + '</w:ins>'),
             ([['A BG H']], [['A BE\u00a0F']]), './w:body/w:p[1]/w:moveFrom/w:r/w:t[.="E\u00a0F"]'),
        )
        for label, body, (accepted, rejected), xpath in cases:
            with self.subTest(label):
                parts = review_parts(normalize_spacing_docx(review_docx_with(body)))
                self.assertEqual((resolved(parts, accept=True), resolved(parts, accept=False)), (accepted, rejected))
                self.assertTrue(parts['document'].xpath(xpath, namespaces=NS), xpath)
                self.assertEqual(review_markup_errors(parts), [])

    def test_comment_range_split_by_a_soft_break_keeps_its_words(self):
        """Zakres komentarza z miękkim enterem obejmuje po podziale oba akapity: te same słowa i to samo odwołanie."""
        body = p('<w:commentRangeStart w:id="0"/>', r('Pierwsza linia'), '<w:r><w:br/></w:r>', r('druga linia'),
                 '<w:commentRangeEnd w:id="0"/>', comment_reference(0))
        parts = review_parts(normalize_spacing_docx(review_docx_with(body, comment_xml(0, r('Uwaga.')))))
        self.assertEqual(len(parts['document'].findall(f'{qn("w:body")}/{qn("w:p")}')), 2)
        self.assertEqual(comment_anchors(parts), {'0': ('Pierwsza linia¶druga linia', 'linia')})
        self.assertEqual(review_markup_errors(parts), [])


@tag('docx')
class NormalizedLayoutTests(SimpleTestCase):
    """Listy, nagłówki i komórki tabel zachowują układ; fonty symboliczne zostają."""

    def test_headings_and_list_items_keep_their_indent(self):
        """Bez wcięcia pierwszego wiersza zostają nagłówki i punkty listy według reguły (OUTLINE_CASES), także
        w poprzednim formatowaniu zmiany śledzonej; nagłówek z numId 0 jest nagłówkiem, a z poziomem 9, styl
        „Heading Custom” czy „Nagłówek tabeli” – zwykłym akapitem z wcięciem."""
        document = outline_document()
        for paragraph in document.paragraphs:
            paragraph._p.get_or_add_pPr().append(parse_xml(
                f'<w:pPrChange {nsdecls("w")} w:id="1" w:author="Redaktor">{paragraph._p.pPr.xml}</w:pPrChange>'))
        for paragraph in Document(normalize_docx(saved(document))).paragraphs:
            _, _, level, listed = OUTLINE_CASES[paragraph.text]
            expected = None if level or listed else '709'
            with self.subTest(paragraph.text):
                for properties in (paragraph._p.pPr, paragraph._p.pPr.find(f'{qn("w:pPrChange")}/{qn("w:pPr")}')):
                    indent = properties.find(qn('w:ind'))
                    self.assertEqual(None if indent is None else indent.get(qn('w:firstLine')), expected)

    def test_lists_headings_tables_and_symbol_fonts_keep_their_layout(self):
        document = Document()
        document.add_paragraph('Zwykły akapit')
        document.add_paragraph('Punkt listy', style='List Number')
        document.add_paragraph('Nagłówek', style='Heading 1')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'W tabeli'
        symbol = document.add_paragraph().add_run(' strzałka')
        symbol.font.name = 'Wingdings'
        result = Document(normalize_docx(saved(document), justify=True))
        plain, listed, heading, symbols = result.paragraphs[0], result.paragraphs[1], result.paragraphs[2], result.paragraphs[3]
        self.assertEqual(plain.paragraph_format.first_line_indent.twips, 709)
        self.assertEqual(plain.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        for paragraph in (listed, heading):
            with self.subTest(paragraph.text):
                self.assertIsNone(paragraph._p.find(f'{qn("w:pPr")}/{qn("w:ind")}'))
                self.assertEqual(paragraph.paragraph_format.line_spacing, 1.5)
        cell = result.tables[0].cell(0, 0).paragraphs[0]
        self.assertIsNone(cell._p.find(f'{qn("w:pPr")}/{qn("w:ind")}'))
        self.assertIsNone(cell._p.find(f'{qn("w:pPr")}/{qn("w:spacing")}'))
        self.assertIsNone(cell.alignment)
        self.assertEqual(symbols.runs[0].font.name, 'Wingdings')
        self.assertEqual(result.paragraphs[0].runs[0].font.name, 'Times New Roman')
