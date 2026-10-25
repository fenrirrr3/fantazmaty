"""Konwerter DOCX → PDF/EPUB po stronie serwera: walidacja pakietu DOCX, wybór interpretera, uruchomienie procesu
roboczego z limitem czasu i wolny konwerter (blokada), tłumaczenie jego błędów, wynik (plik lub ZIP, pliki
pomocnicze summary.json i images.json, obrazy połączone) i strona programu."""
import json
import os
import subprocess
import sys
import warnings
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

import django
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from lxml import etree
from lxml import html as html_parser

from core.services.document_conversion_worker import convert
from core.services.document_converter import (
    ConversionError, RebuildConfirmationRequired, conversion_filename, conversion_slot,
    convert_document, converter_python, run_converter, validate_resources,
)
from core.services.document_errors import MESSAGES
from core.services.document_images import linked_images_note
from fantazmaty.testing import create_member

from .support import (
    IsolatedProgramDirsMixin, docx_with_linked_images, downloaded, formatted_docx, run_program, saved,
    typography_sample,
)


RELATIONSHIPS = 'word/_rels/document.xml.rels'


def repacked(replace=None, extra=()):
    """Poprawny DOCX z podmienionymi częściami {nazwa: funkcja(bajty)} i dopisanymi wpisami (nazwa, bajty)."""
    output = BytesIO()
    with ZipFile(BytesIO(formatted_docx())) as source, ZipFile(output, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            target.writestr(item, (replace or {}).get(item.filename, lambda value: value)(data))
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)  # celowo powielona nazwa wpisu
            for name, data in extra:
                target.writestr(name, data)
    output.seek(0)
    return output


def with_relationship(target, mode='', kind='image'):
    attributes = f' TargetMode="{mode}"' if mode else ''
    relationship = (f'<Relationship Id="rIdTest" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    f'relationships/{kind}" Target="{target}"{attributes}/></Relationships>').encode()
    return repacked({RELATIONSHIPS: lambda data: data.replace(b'</Relationships>', relationship)})


class Process:
    """Proces roboczy: kod wyjścia albo (code=None) proces, który wciąż działa."""

    def __init__(self, code=None):
        self.code, self.pid, self.waited = code, 4242, 0

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def wait(self, timeout=None):
        self.waited += 1
        if self.code is None and timeout is not None:
            raise subprocess.TimeoutExpired('worker', timeout)
        return self.code

    def poll(self):
        return self.code


class Stop(Exception):
    """Przerwanie zadania zgłoszone przez funkcję kontrolną (np. użytkownik anulował pracę)."""


class FakeMsvcrt:
    """Blokada plikowa Windows: pierwszy LK_NBLCK zajmuje bajt, kolejny zgłasza OSError do LK_UNLCK."""
    LK_NBLCK, LK_UNLCK = 'nblck', 'unlck'

    def __init__(self):
        self.locked, self.calls = False, []

    def locking(self, fd, mode, length):
        self.calls.append((mode, length))
        if mode == self.LK_NBLCK:
            if self.locked:
                raise OSError('zajęte')
            self.locked = True
        else:
            self.locked = False


@tag('docx')
class RealWorkerProcessTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Prawdziwy podproces konwertera (bootstrap, środowisko bez danych projektu) – pozostałe testy uruchamiają
    proces roboczy w procesie testu (support.InProcessWorker)."""

    in_process_worker = False

    def test_worker_exposes_only_known_reason_not_document_content(self):
        """Prawdziwy proces roboczy: użytkownik widzi stały komunikat, a nie treść dokumentu ani wyjątku."""
        doc = Document()
        doc.add_paragraph('PRYWATNY_FRAGMENT')
        doc.element.body.append(OxmlElement('w:ins'))
        directory = self.work_dir
        doc.save(directory / 'source.docx')
        (directory / 'job.json').write_text(
            json.dumps({'formats': [], 'include_docx': True, 'repetitions': {}}), encoding='utf-8')
        with self.assertLogs('core.services.document_converter', 'ERROR') as logs, \
                self.assertRaisesMessage(ConversionError, MESSAGES['tracked_changes']):
            run_converter(directory, 30)
        self.assertNotIn('PRYWATNY_FRAGMENT', '\n'.join(logs.output))
        report = (directory / 'error.json').read_text(encoding='utf-8')
        self.assertEqual(json.loads(report)['public_code'], 'tracked_changes')
        self.assertNotIn('PRYWATNY_FRAGMENT', report)


@tag('docx')
class ConverterServiceTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_warnings_are_returned_with_a_single_file_not_an_archive(self):
        """Regresja: uwagi konwertera zamieniały pojedynczy plik w archiwum ZIP."""
        cases = (
            ('epub', b'EPUB test output', lambda: BytesIO(formatted_docx()), {}),
            ('pdf', b'%PDF-test', lambda: SimpleUploadedFile('Zażółć.docx', formatted_docx()),
             {'preserve_filename': True}),
        )
        for extension, content, source, options in cases:
            def worker(directory, timeout, **_):
                (directory / f'document.{extension}').write_bytes(content)
                (directory / 'warnings.json').write_text(json.dumps(['Pominięto nieobsługiwany element.']))

            with self.subTest(extension=extension), \
                    patch('core.services.document_converter.run_converter', side_effect=worker):
                result, result_extension, _ = convert_document(source(), [extension], **options)
                with result:
                    self.assertEqual(result_extension, extension)
                    self.assertEqual(result.read(), content)
                self.assertEqual(result.conversion_warnings, ['Pominięto nieobsługiwany element.'])

    def test_names_cannot_escape_zip_or_inject_response_headers(self):
        self.assertEqual(conversion_filename(r'C:\folder\Tytuł.v1.docx', 'pdf'), 'Tytuł.v1.pdf')
        self.assertEqual(conversion_filename('../../tekst\r\n.docx', 'epub'), 'tekst.epub')
        self.assertEqual(conversion_filename('..docx', 'pdf'), 'Dokument.pdf')

    def test_conversion_epub_and_pdf_preserve_formatting(self):
        result, _, _ = convert_document(typography_sample(), ['pdf', 'epub'])
        with result, ZipFile(result) as archive:
            self.assertTrue(archive.read('document.pdf').startswith(b'%PDF'))
            with ZipFile(BytesIO(archive.read('document.epub'))) as epub:
                content = epub.read('EPUB/content.xhtml').decode()
                css = epub.read('EPUB/style.css').decode()
        self.assertIn('align-center', content)
        self.assertIn('<strong>', content)
        self.assertIn('<em>', content)
        self.assertIn('1.25cm', css)
        self.assertIn('Times New Roman', css)

    def test_epub_preserves_empty_paragraphs_and_removes_work_folder(self):
        doc = Document()
        doc.add_paragraph('A')
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph('B')
        out, _, _ = convert_document(saved(doc), ['epub'])
        with out, ZipFile(out) as archive:
            root = etree.fromstring(archive.read('EPUB/content.xhtml'))
        paragraphs = root.findall('.//{http://www.w3.org/1999/xhtml}p')
        self.assertEqual([''.join(p.itertext()).strip() for p in paragraphs], ['A', '', '', 'B'])
        self.assertFalse(list((self.work_dir / 'convert').glob('document-*')))

    def test_empty_paragraphs_survive_cleaning_and_html_conversion(self):
        """Regresja: Odkurzacz i konwersja do HTML zjadały puste akapity przed PDF."""
        doc = Document()
        doc.add_paragraph('Pierwszy.')
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph('Drugi.').alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        def inspect_worker(directory, timeout, **options):
            self.assertEqual([p.text for p in Document(directory / 'source.docx').paragraphs],
                             ['Pierwszy.', '', '', 'Drugi.'])
            return run_converter(directory, timeout, **options)

        with patch('core.services.document_converter.run_converter', side_effect=inspect_worker):
            output, _, _ = convert_document(saved(doc), ['pdf'], use_cleaner=True)
        with output:
            self.assertTrue(output.read().startswith(b'%PDF'))
        source = self.work_dir / 'input.docx'
        doc.save(source)
        with patch('core.services.document_conversion_worker.render_pdf') as pdf:
            convert(source, self.work_dir, ['pdf'], 'Test')
        html = pdf.call_args.args[2]
        paragraphs = html_parser.fragment_fromstring(html, create_parent='div').xpath('.//p')
        self.assertEqual(sum(not p.text_content().strip() for p in paragraphs), 2)
        self.assertIn('align-justify', html)


@tag('docx')
class ConversionPageTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('programy', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def test_conversion_preserves_name_for_single_and_multiple_formats(self):
        name = 'Żółw i smok.v2.docx'
        for formats, ext in ((['pdf'], 'pdf'), (['pdf', 'epub'], 'zip')):
            with self.subTest(formats=formats):
                response = run_program(self.client, {
                    'program_action': 'convert',
                    'convert-formats': formats,
                    'convert-document': SimpleUploadedFile(name, formatted_docx()),
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn('Żółw i smok.v2.' + ext, unquote(response['Content-Disposition']))
                payload = downloaded(response)
                if ext == 'zip':
                    with ZipFile(BytesIO(payload)) as archive:
                        self.assertEqual(set(archive.namelist()), {'Żółw i smok.v2.pdf', 'Żółw i smok.v2.epub'})
                else:
                    self.assertTrue(payload.startswith(b'%PDF'))


class ResourceValidationTests(SimpleTestCase):
    """Konwerter nie pobiera plików spoza DOCX i nie przyjmuje podejrzanie zbudowanych pakietów."""

    def test_safe_document_external_hyperlink_and_linked_images_are_accepted(self):
        """Regresja: plik z obrazem połączonym z plikiem zewnętrznym był odrzucany („Osadź go w DOCX”) – proces
        roboczy usuwa link (decyzja właściciela); inne podłączone pliki odrzuca test_dangerous_packages_are_refused."""
        validate_resources(repacked())
        for source in (with_relationship('https://example.test', 'External', 'hyperlink'), docx_with_linked_images()):
            validate_resources(source)
            self.assertEqual(source.tell(), 0)

    def test_attached_template_and_absolute_internal_targets_are_accepted(self):
        """Regresja: dokument z szablonu Worda (attachedTemplate) i ścieżki „/word/…” były odrzucane."""
        validate_resources(with_relationship('file:///C:/Szablony/Normal.dotm', 'External', 'attachedTemplate'))
        absolute = repacked({RELATIONSHIPS: lambda data: data.replace(b'Target="styles.xml"',
                                                                     b'Target="/word/styles.xml"')})
        validate_resources(absolute)
        with ZipFile(absolute) as archive:
            self.assertIn(b'Target="/word/styles.xml"', archive.read(RELATIONSHIPS))
        with self.assertRaisesMessage(ConversionError, 'nieprawidłowe odwołanie'):
            validate_resources(with_relationship('/../../etc/passwd'))
        with self.assertRaisesMessage(ConversionError, 'nieprawidłowe odwołanie'):
            validate_resources(with_relationship('//serwer/udział/obraz.png'))

    def test_dangerous_packages_are_refused(self):
        doctype = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "x">]><Relationships/>'
        cases = (
            (repacked(extra=[('../poza.xml', b'x')]), 'Nieprawidłowa struktura'),
            (repacked(extra=[('word\\plik.xml', b'x')]), 'Nieprawidłowa struktura'),
            (repacked(extra=[('[Content_Types].xml', b'x')]), 'powielone elementy'),
            (repacked(extra=[('word/_rels/extra.xml.rels', doctype)]), 'Nieobsługiwany dokument'),
            (with_relationship('../../../etc/passwd'), 'nieprawidłowe odwołanie'),
            (with_relationship('C:/obraz.png'), 'nieprawidłowe odwołanie'),
            (with_relationship('file:///C:/Arkusz.xlsx', 'External', 'oleObject'), 'podłączony plik zewnętrzny'),
        )
        for source, message in cases:
            with self.subTest(message=message), self.assertRaisesMessage(ConversionError, message):
                validate_resources(source)

    def test_unreadable_packages_get_a_conversion_message(self):
        """Regresja: plik spoza ZIP lub uszkodzone .rels kończyły się ogólnym błędem zadania zamiast komunikatu."""
        truncated = BytesIO(formatted_docx()[:-200])
        cases = (
            (BytesIO(b'To nie jest DOCX'), 'Nie można odczytać pliku'),
            (truncated, 'Nie można odczytać pliku'),
            (repacked({RELATIONSHIPS: lambda data: data[:-30]}), 'Nieprawidłowa struktura'),
        )
        for source, message in cases:
            with self.subTest(message=message), self.assertRaisesMessage(ConversionError, message):
                validate_resources(source)


class ConverterPythonTests(SimpleTestCase):
    def test_configured_interpreter_must_be_an_existing_absolute_path(self):
        with self.settings(DOCUMENT_CONVERTER_PYTHON=sys.executable):
            self.assertEqual(converter_python(), sys.executable)
        for configured in ('python3', '/nie/ma/pythona'):
            with self.subTest(configured=configured), self.settings(DOCUMENT_CONVERTER_PYTHON=configured), \
                    self.assertRaisesMessage(ConversionError, 'DOCUMENT_CONVERTER_PYTHON'):
                converter_python()

    def test_without_virtualenv_the_running_python_is_used_but_not_a_wsgi_server(self):
        python = Path(sys.executable).resolve()
        with self.settings(DOCUMENT_CONVERTER_PYTHON=''), \
                patch.object(django, '__file__', '/nie/ma/django/__init__.py'), \
                patch('core.services.document_converter.sys') as fake_sys:
            fake_sys.prefix = '/nie/ma/prefiksu'
            fake_sys.executable = str(python)
            self.assertEqual(converter_python(), str(python))
            fake_sys.executable = str(python.with_name('uwsgi'))
            with self.assertRaisesMessage(ConversionError, 'Nie znaleziono Pythona'):
                converter_python()


class RunConverterTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def setUp(self):
        super().setUp()
        self.directory = self.work_dir

    def finished(self, code, report=None):
        if report is not None:
            (self.directory / 'error.json').write_text(json.dumps(report), encoding='utf-8')
        with patch('core.services.document_converter.subprocess.Popen', return_value=Process(code)):
            return run_converter(self.directory, 5)

    def test_exhausted_time_or_unstartable_process(self):
        with patch('core.services.document_converter.subprocess.Popen') as popen, \
                self.assertRaisesMessage(ConversionError, 'limit 90 sekund'):
            run_converter(self.directory, 0)
        popen.assert_not_called()
        with patch('core.services.document_converter.subprocess.Popen', side_effect=OSError), \
                self.assertRaisesMessage(ConversionError, 'Nie można uruchomić konwertera'):
            run_converter(self.directory, 5)

    def test_worker_failures_become_public_messages(self):
        self.assertIsNone(self.finished(0))
        cases = (
            (-9, None, 'Nie udało się przygotować plików (start: ProcessTerminated)'),
            (2, {'stage': 'PDF', 'error': 'ImportError'}, 'Brakuje bibliotek konwertera. Zainstaluj requirements.txt. PDF: ImportError'),
            (3, {'stage': '<b>', 'error': 'Zły <tag>'}, '(konwersja: błąd procesu)'),
            (3, {'error': 'DocumentInputError', 'public_code': 'nieznany'}, 'DocumentInputError'),
            (3, {'stage': 'DOCX', 'error': 'MemoryError'}, 'wymaga zbyt dużo pamięci'),
        )
        for code, report, message in cases:
            with self.subTest(message=message), self.assertLogs('core.services.document_converter', 'ERROR'), \
                    self.assertRaisesMessage(ConversionError, message):
                self.finished(code, report)
            (self.directory / 'error.json').unlink(missing_ok=True)

    def test_rebuild_omissions_ask_for_confirmation_without_logging_an_error(self):
        with self.assertNoLogs('core.services.document_converter', 'ERROR'), \
                self.assertRaises(RebuildConfirmationRequired) as caught:
            self.finished(3, {'error': 'RebuildUnsupported', 'omissions': ['tabele (liczba: 1)']})
        self.assertIn('W nowym DOCX zostaną pominięte: tabele (liczba: 1).', str(caught.exception))
        # Scalanie bierze same pominięcia z wyjątku, bez rozbierania komunikatu.
        self.assertEqual(caught.exception.omissions, ['tabele (liczba: 1)'])

    def test_process_over_time_limit_is_killed(self):
        sleeper = [sys.executable, '-c', 'import time; time.sleep(30)']
        actual = subprocess.Popen
        for control in (None, lambda directory: None):
            processes = []

            def spawn(command, **kwargs):
                processes.append(actual(sleeper, **kwargs))
                return processes[-1]

            with self.subTest(control=control), \
                    patch('core.services.document_converter.subprocess.Popen', side_effect=spawn), \
                    self.assertRaisesMessage(ConversionError, 'limit 90 sekund'):
                run_converter(self.directory, 0.3, control=control)
            self.assertIsNotNone(processes[0].poll())

    def start(self, process, control):
        with patch('core.services.document_converter.subprocess.Popen', return_value=process), \
                patch('core.services.document_converter.os.killpg', side_effect=ProcessLookupError) as kill:
            try:
                return run_converter(self.directory, 5, control=control)
            finally:
                self.kill = kill

    def test_time_used_by_the_control_function_counts_against_the_limit(self):
        """Gdy kontrola postępu zużyje cały czas, proces jest zatrzymywany (nawet jeśli już zniknął)."""
        process = Process()
        clock = iter([0, 100])
        with patch('core.services.document_converter.time.monotonic', side_effect=lambda: next(clock)), \
                self.assertRaisesMessage(ConversionError, 'limit 90 sekund'):
            self.start(process, lambda directory: None)
        self.kill.assert_called_once()
        self.assertEqual(process.waited, 1)

    def test_interrupted_job_kills_only_a_running_process(self):
        def stop(directory):
            raise Stop

        for code, killed in ((None, True), (0, False)):
            process = Process(code)
            with self.subTest(running=code is None), self.assertRaises(Stop):
                self.start(process, stop)
            self.assertEqual(self.kill.called, killed)
            self.assertEqual(process.waited, 1)

    def test_oversized_error_report_is_not_read(self):
        (self.directory / 'error.json').write_text(
            json.dumps({'error': 'DocumentInputError', 'public_code': 'tracked_changes', 'x': 'y' * 40000}))
        with patch('core.services.document_converter.subprocess.Popen', return_value=Process(3)), \
                self.assertLogs('core.services.document_converter', 'ERROR'), \
                self.assertRaisesMessage(ConversionError, '(start: ProcessTerminated)'):
            run_converter(self.directory, 5)


class ConvertDocumentResultTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_unknown_format_is_refused_before_conversion(self):
        for formats, options in ((['mobi'], {}), ([], {})):
            with self.subTest(formats=formats), \
                    patch('core.services.document_converter.run_converter') as worker, \
                    self.assertRaisesMessage(ConversionError, 'Wybierz co najmniej jeden'):
                convert_document(BytesIO(formatted_docx()), formats, **options)
            worker.assert_not_called()

    def test_missing_output_is_an_error_and_several_outputs_are_zipped(self):
        with patch('core.services.document_converter.run_converter'), \
                self.assertRaisesMessage(ConversionError, 'nie utworzyła poprawnego pliku'):
            convert_document(BytesIO(formatted_docx()), ['pdf'])

        def worker(directory, timeout, **_):
            (directory / 'document.pdf').write_bytes(b'%PDF')

        source = formatted_docx()  # jeden egzemplarz – DOCX zbudowany sekundę później ma inne znaczniki czasu w ZIP-ie
        with patch('core.services.document_converter.run_converter', side_effect=worker):
            result, extension, mime = convert_document(BytesIO(source), ['pdf'], include_docx=True)
        with result, ZipFile(result) as archive:
            self.assertEqual((extension, mime), ('zip', 'application/zip'))
            self.assertEqual(sorted(archive.namelist()), ['document.docx', 'document.pdf'])
            self.assertEqual(archive.read('document.docx'), source)

    def test_failed_pdf_does_not_take_away_the_ready_epub(self):
        """Raport błędu PDF trafia do dziennika; gdy nie da się go odczytać, dziennik jest pomijany – w obu
        przypadkach użytkownik dostaje gotowy EPUB."""
        def worker(directory, timeout, **_):
            (directory / 'document.epub').write_bytes(b'PK epub')
            (directory / 'pdf_error.json').write_text('{"error": "ValueError"}', encoding='utf-8')

        original = Path.read_text

        def unreadable(path, *args, **kwargs):
            if path.name == 'pdf_error.json':
                raise OSError('brak dostępu')
            return original(path, *args, **kwargs)

        cases = (('raport czytelny', self.assertLogs('core.services.document_converter', 'ERROR')),
                 ('raport nieczytelny', patch.object(Path, 'read_text', unreadable)))
        for label, context in cases:
            with self.subTest(label), patch('core.services.document_converter.run_converter', side_effect=worker), \
                    context:
                result, extension, mime = convert_document(BytesIO(formatted_docx()), ['pdf', 'epub'])
            with result:
                self.assertEqual((extension, mime), ('epub', 'application/epub+zip'))
                self.assertEqual(result.read(), b'PK epub')

    def test_result_finished_after_the_deadline_is_discarded(self):
        def worker(directory, timeout, **_):
            (directory / 'document.pdf').write_bytes(b'%PDF')

        clock = iter([0, 0, 0, 10_000])
        with patch('core.services.document_converter.run_converter', side_effect=worker), \
                patch('core.services.document_converter.time.monotonic', side_effect=lambda: next(clock)), \
                self.assertRaisesMessage(ConversionError, 'przekroczyło limit czasu'):
            convert_document(BytesIO(formatted_docx()), ['pdf'])

    def test_unexpected_summary_and_image_counts_are_ignored(self):
        """Granica zaufania: konwerter przepuszcza z summary.json i images.json tylko dodatnie liczby całkowite
        i poprawne zamiany słownikowe – program_jobs nie sprawdza ich ponownie."""
        cases = {
            'lista zamiast słownika': ('[1, 2]', '[3]'),
            'uszkodzony JSON': ('{"spaces": 2}', '{uszkodzony'),
            'liczba jako tekst': ('{"spaces": 2, "replacements": "x"}', '{"removed_images": "3"}'),
            'błędne zamiany': ('{"spaces": 2, "zero": 0, "flaga": true, "replacements": ['
                               '{"source": "a", "target": "b", "count": 0}, {"source": "a", "target": "b", "count": true},'
                               '{"source": "a", "target": "b", "count": 2}]}', '{}'),
        }
        expected = {'lista zamiast słownika': None, 'uszkodzony JSON': {'spaces': 2},
                    'liczba jako tekst': {'spaces': 2},
                    'błędne zamiany': {'spaces': 2, 'replacements': [{'source': 'a', 'target': 'b', 'count': 2}]}}
        for label, (summary, images) in cases.items():
            def worker(directory, timeout, summary=summary, images=images, **_):
                (directory / 'document.epub').write_bytes(b'PK epub')
                (directory / 'summary.json').write_text(summary, encoding='utf-8')
                (directory / 'images.json').write_text(images, encoding='utf-8')

            with self.subTest(label), patch('core.services.document_converter.run_converter', side_effect=worker):
                result, extension, _ = convert_document(BytesIO(formatted_docx()), ['epub'], cleaner_summary=True)
                result.close()
            self.assertEqual(extension, 'epub')
            self.assertEqual(result.cleaner_summary, expected[label])
            self.assertEqual(result.removed_images, 0)

    def test_linked_images_get_a_note_before_other_warnings(self):
        """Liczba obrazów tylko połączonych (images.json → linked_images) daje osobną uwagę na początku listy;
        liczba usuniętych obrazów osadzonych zostaje w removed_images."""
        for label, images, expected in (
                ('obrazy połączone', {'removed_images': 1, 'linked_images': 2}, [linked_images_note(2), 'Uwaga']),
                ('bez obrazów połączonych', {'removed_images': 1}, ['Uwaga'])):
            def worker(directory, timeout, images=images, **_):
                (directory / 'document.epub').write_bytes(b'PK epub')
                (directory / 'warnings.json').write_text(json.dumps(['Uwaga']), encoding='utf-8')
                (directory / 'images.json').write_text(json.dumps(images), encoding='utf-8')

            with self.subTest(label), patch('core.services.document_converter.run_converter', side_effect=worker):
                result, _, _ = convert_document(BytesIO(formatted_docx()), ['epub'])
                result.close()
                self.assertEqual(result.conversion_warnings, expected)
                self.assertEqual(result.removed_images, 1)


class ConversionSlotTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_windows_lock_is_taken_once_and_released_after_the_slot(self):
        """Na Windows (msvcrt) zajęty bajt blokady oznacza zajęty konwerter; wyjście ze slotu zwalnia blokadę."""
        msvcrt = FakeMsvcrt()
        windows = type('WindowsOs', (), {'name': 'nt', '__getattr__': lambda self, name: getattr(os, name)})()
        with patch('core.services.document_converter.os', windows), patch.dict(sys.modules, {'msvcrt': msvcrt}):
            with conversion_slot():
                with self.assertRaisesMessage(ConversionError, 'przetwarza teraz inny dokument'):
                    with conversion_slot():
                        pass
            self.assertEqual(msvcrt.calls, [('nblck', 1), ('nblck', 1), ('unlck', 1)])
            self.assertFalse(msvcrt.locked)
            with conversion_slot():
                pass
        self.assertEqual(msvcrt.calls[-2:], [('nblck', 1), ('unlck', 1)])

    def test_slot_is_retried_every_second_without_a_wait_callback(self):
        with patch('core.services.document_converter._try_lock', side_effect=[False, False, True]) as lock, \
                patch('core.services.document_converter.time.sleep') as sleep:
            with conversion_slot(wait=60) as directory:
                self.assertTrue(directory.is_dir())
        self.assertEqual(lock.call_count, 3)
        self.assertEqual([call.args for call in sleep.call_args_list], [(1,), (1,)])
