"""Przebudowa DOCX: dokument bez ustawień domyślnych, kolory motywu spoza palety i błędne kolory, puste
kontrolki tekstu zastępczego, nieznane bloki treści, pusty nagłówek, znak akapitu i błędna liczba kolumn."""
from django.test import SimpleTestCase
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

from core.services.document_rebuild import RebuildUnsupported, inspect_docx, rebuild_docx

from .support import add_block, saved


def omissions_of(document):
    try:
        inspect_docx(saved(document))
    except RebuildUnsupported as error:
        return error.omissions
    return []


def color(**attributes):
    return OxmlElement('w:color', {qn('w:' + key): value for key, value in attributes.items()})


class RebuildGapTests(SimpleTestCase):
    def test_document_without_defaults_and_default_tab_stop(self):
        document = Document()
        document.add_paragraph('Tekst')
        styles = document.styles.element
        styles.remove(styles.find(qn('w:docDefaults')))
        settings = document.settings.element
        settings.remove(settings.find(qn('w:defaultTabStop')))
        rebuilt = Document(rebuild_docx(saved(document)))
        self.assertEqual(rebuilt.paragraphs[0].text, 'Tekst')
        # Źródło nie miało ustawień domyślnych – wynik też ich nie ma; odstęp tabulatora zostaje z szablonu.
        self.assertIsNone(rebuilt.styles.element.find(qn('w:docDefaults')))
        self.assertEqual(rebuilt.settings.element.xpath('w:defaultTabStop/@w:val'), ['720'])

    def test_colors_outside_the_theme_palette_and_invalid_values(self):
        """Kolor motywu spoza palety: wartość zapisana w w:val; błędny kolor przebiegu: kolor ze stylu."""
        document = Document()
        style = document.styles['Normal']
        style.element.get_or_add_rPr().append(color(val='112233'))
        document.add_paragraph().add_run('Motyw')._r.get_or_add_rPr().append(color(val='0563C1', themeColor='zzz'))
        document.add_paragraph().add_run('Błędny')._r.get_or_add_rPr().append(color(val='niebieski'))
        first, second = (p.runs[0] for p in Document(rebuild_docx(saved(document))).paragraphs)
        self.assertEqual(first.font.color.rgb, RGBColor(0x05, 0x63, 0xC1))
        self.assertEqual(second.font.color.rgb, RGBColor(0x11, 0x22, 0x33))

    def test_empty_placeholder_controls_disappear_without_a_warning(self):
        document = Document()
        document.add_paragraph('Początek')
        # Kontrolka bez treści i kontrolka z odstępem obok tekstu zastępczego.
        add_block(document, parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:showingPlcHdr/></w:sdtPr></w:sdt>'))
        add_block(document, parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:showingPlcHdr/></w:sdtPr><w:sdtContent><w:p>'
            '<w:r><w:t xml:space="preserve"> </w:t></w:r><w:r><w:rPr><w:rStyle w:val="PlaceholderText"/></w:rPr>'
            '<w:t>Kliknij tutaj.</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        self.assertEqual(omissions_of(document), [])
        self.assertEqual([p.text for p in Document(rebuild_docx(saved(document))).paragraphs], ['Początek'])

    def test_unknown_blocks_empty_header_and_damaged_columns(self):
        document = Document()
        document.add_paragraph('Tekst')
        # Pusty nagłówek (bez tekstu i grafiki) i liczba kolumn, której nie da się odczytać: bez uwag.
        document.sections[0].header.is_linked_to_previous = False
        document.sections[0]._sectPr.find(qn('w:cols')).set(qn('w:num'), 'dwie')
        self.assertEqual(omissions_of(document), [])
        add_block(document, parse_xml(f'<w:customXml {nsdecls("w")} w:element="dane"><w:p><w:r><w:t>Ukryte</w:t>'
                                      '</w:r></w:p></w:customXml>'))
        self.assertEqual(omissions_of(document), ['blok customXml wraz z treścią (liczba: 1)'])
        self.assertEqual([p.text for p in Document(rebuild_docx(saved(document), allow_omissions=True)).paragraphs],
                         ['Tekst'])

    def test_paragraph_mark_with_only_a_size_or_only_a_font(self):
        """Znak akapitu przenosi tylko to, co określono: sam rozmiar albo sam krój (gdy domyślnych brak)."""
        document = Document()
        styles = document.styles.element
        styles.remove(styles.find(qn('w:docDefaults')))
        sized = document.add_paragraph()._p.get_or_add_pPr()
        sized.append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="32"/></w:rPr>'))
        named = document.add_paragraph()._p.get_or_add_pPr()
        named.append(parse_xml(f'<w:rPr {nsdecls("w")}><w:rFonts w:ascii="Georgia" w:hAnsi="Georgia"/></w:rPr>'))
        marks = [p._p.pPr.find(qn('w:rPr')) for p in Document(rebuild_docx(saved(document))).paragraphs]
        self.assertEqual([child.tag for child in marks[0]], [qn('w:sz')])
        self.assertEqual(marks[0].find(qn('w:sz')).get(qn('w:val')), str(int(Pt(16).pt * 2)))
        self.assertEqual([child.tag for child in marks[1]], [qn('w:rFonts')])
        self.assertEqual(marks[1].find(qn('w:rFonts')).get(qn('w:ascii')), 'Georgia')
