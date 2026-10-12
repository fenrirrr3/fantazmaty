"""Proces roboczy konwertera uruchomiony tak jak na serwerze: jako skrypt (bez pakietu core.services).

Bootstrap i proces roboczy importują wtedy moduły po samej nazwie pliku (gałąź „else” przy
``if __package__``). Testy uruchamiają je w tym procesie przez runpy, a po teście usuwają z
sys.modules kopie modułów załadowane spoza pakietu.
"""
import json
import runpy
import sys
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase, tag
from docx import Document

import core.services

SERVICES = Path(core.services.__file__).parent
BOOTSTRAP = SERVICES / 'document_worker_bootstrap.py'
SCRIPT_MODULES = {path.stem for path in SERVICES.glob('*.py')}


@contextmanager
def script_mode(directory):
    """Katalog core/services na sys.path (jak dla skryptu) i argv procesu; po wyjściu sprzątamy kopie modułów."""
    before = set(sys.modules)
    with patch.object(sys, 'path', [str(SERVICES), *sys.path]), \
            patch.object(sys, 'argv', ['document_worker_bootstrap.py', str(directory)]):
        try:
            yield
        finally:
            for name in set(sys.modules) - before:
                if name in SCRIPT_MODULES:
                    del sys.modules[name]


@tag('slow')
class ScriptModeWorkerTests(SimpleTestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)

    def run_job(self, document, **config):
        """Uruchom bootstrap jako __main__; zwraca kod wyjścia procesu roboczego."""
        document.save(self.directory / 'source.docx')
        # memory_limit_mb = 0: bez limitu RLIMIT_AS, który dotyczyłby całego procesu testów.
        (self.directory / 'job.json').write_text(json.dumps({'title': 'Skrypt', 'memory_limit_mb': 0, **config}),
                                                 encoding='utf-8')
        with script_mode(self.directory), self.assertRaises(SystemExit) as caught:
            runpy.run_path(str(BOOTSTRAP), run_name='__main__')
        return caught.exception.code

    def test_full_job_runs_every_step_without_the_package(self):
        """Odkurzacz, miękkie entery, ujednolicenie, powtórzenia, EPUB i PDF – wszystko z importów skryptu."""
        document = Document()
        paragraph = document.add_paragraph('Kot  i kot "siedzą" obok')
        paragraph.add_run().add_break()
        paragraph.add_run('drugiej linii.')
        code = self.run_job(document, formats=['epub', 'pdf'], prepare=True, clean=True, cleaner_rules=None,
                            remove_soft_whitespace=True, include_docx=True, cleaner_summary=True,
                            repetitions={'min_word_length': 3, 'analysis_options': {
                                'long_sentences': False, 'long_paragraphs': False}})
        self.assertEqual(code, 0)
        self.assertFalse((self.directory / 'error.json').exists())
        marked = Document(self.directory / 'source.docx')
        # Miękki enter dzieli akapit; reguły domyślne poprawiają spacje i cudzysłowy.
        self.assertEqual([p.text for p in marked.paragraphs], ['Kot i kot „siedzą” obok', 'drugiej linii.'])
        self.assertTrue(any(run.font.color.rgb for run in marked.paragraphs[0].runs))
        self.assertGreater(json.loads((self.directory / 'summary.json').read_text(encoding='utf-8'))
                           ['spaces'], 0)
        with ZipFile(self.directory / 'document.epub') as epub:
            self.assertIn('drugiej linii.', epub.read('EPUB/content.xhtml').decode())
        self.assertTrue((self.directory / 'document.pdf').read_bytes().startswith(b'%PDF'))
        self.assertEqual(json.loads((self.directory / 'images.json').read_text()), {'removed_images': 0})

    def test_inspection_and_public_errors_without_the_package(self):
        """Inspekcja przebudowy (z formatem strony) kończy się sukcesem, a błąd dokumentu ma publiczny kod."""
        self.assertEqual(self.run_job(Document(), formats=[], inspect=True, keep_page_format=True), 0)
        document = Document()
        document.add_paragraph('a' * 20_001)
        self.assertEqual(self.run_job(document, formats=['epub'], prepare=True, clean=True, cleaner_rules=[]), 3)
        report = json.loads((self.directory / 'error.json').read_text(encoding='utf-8'))
        self.assertEqual((report['error'], report['public_code']), ('DocumentInputError', 'cleaner_paragraph_limit'))
