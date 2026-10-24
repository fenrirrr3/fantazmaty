"""Zamiana miękkich enterów na akapity: kontrolki zawartości dostają nowe identyfikatory, poziom konspektu
akapitu decyduje o nagłówku, a dokument bez części stylów też jest dzielony; zmiany śledzone i komentarze
zostają (tekst po przyjęciu i odrzuceniu zmian taki, jak gdyby miękki enter był podziałem akapitu od początku)."""
from io import BytesIO
from zipfile import ZipFile

from django.test import SimpleTestCase
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from core.services.document_whitespace import normalize_spacing_docx

from .editor_fixtures_exotic import p, r
from .support import content_control, saved
from .test_document_rebuild import (NS, comment_anchors, comment_reference, comment_xml, resolved,
                                    review_docx_with, review_markup_errors, review_parts)

W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'


def body_paragraphs(stream):
    with ZipFile(stream) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    return root.find(qn('w:body')).findall(qn('w:p'))


def text(paragraph):
    return ''.join(paragraph.itertext())


class SoftBreakTests(SimpleTestCase):
    def test_content_control_split_by_a_soft_break_gets_a_new_id(self):
        """Kopia kontrolki w nowym akapicie dostaje identyfikator większy od wszystkich w dokumencie – także od
        identyfikatorów innych kontrolek (<w:id w:val>), a nie tylko od atrybutów w:id."""
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(content_control('<w:r><w:t>Przed</w:t><w:br/><w:t>po</w:t></w:r>'))
        other = content_control('<w:r><w:t>Inna</w:t></w:r>')
        other.find(f'{qn("w:sdtPr")}/{qn("w:id")}').set(qn('w:val'), '1')
        document.add_paragraph()._p.append(other)
        first, second, third = body_paragraphs(normalize_spacing_docx(saved(document)))
        self.assertEqual((text(first), text(second), text(third)), ('Przed', 'po', 'Inna'))
        ids = [p.find(f'{qn("w:sdt")}/{qn("w:sdtPr")}/{qn("w:id")}').get(qn('w:val')) for p in (first, second, third)]
        self.assertEqual((ids[0], ids[2]), ('101', '1'))
        self.assertEqual(ids[1], '102')

    def test_direct_outline_level_decides_whether_the_line_break_stays(self):
        """Poziom konspektu 1–9 to nagłówek (miękki enter zostaje); w:val="9" to zwykły tekst."""
        for level, expected in (('0', ['Tytuł rozdziału']), ('9', ['Tytuł', ' rozdziału'])):
            document = Document()
            paragraph = document.add_paragraph('Tytuł')
            paragraph.add_run().add_break()
            paragraph.add_run(' rozdziału')
            paragraph._p.get_or_add_pPr().append(OxmlElement('w:outlineLvl', {qn('w:val'): level}))
            with self.subTest(level=level):
                self.assertEqual([text(p) for p in body_paragraphs(normalize_spacing_docx(saved(document)))],
                                 expected)

    def test_package_without_styles_and_paragraph_identity(self):
        """Bez word/styles.xml żaden styl nie jest nagłówkiem; druga część traci tylko w14:paraId."""
        document = Document()
        paragraph = document.add_paragraph('Pierwsza')
        paragraph.add_run().add_break()
        paragraph.add_run('druga')
        paragraph._p.set(f'{{{W14}}}paraId', '1A2B3C4D')
        paragraph._p.set(qn('w:rsidR'), '00AB12CD')
        source, stripped = saved(document), BytesIO()
        with ZipFile(source) as archive, ZipFile(stripped, 'w') as target:
            for item in archive.infolist():
                if item.filename != 'word/styles.xml':
                    target.writestr(item, archive.read(item))
        first, second = body_paragraphs(normalize_spacing_docx(stripped))
        self.assertEqual((text(first), text(second)), ('Pierwsza', 'druga'))
        self.assertEqual(first.get(f'{{{W14}}}paraId'), '1A2B3C4D')
        self.assertIsNone(second.get(f'{{{W14}}}paraId'))
        self.assertEqual(second.get(qn('w:rsidR')), '00AB12CD')

    def test_tracked_changes_and_comments_survive_the_split(self):
        """Regresja: miękki enter we wstawce dawał zwykły znak akapitu (odrzucenie wstawki zostawiało pusty akapit) –
        teraz znak akapitu jest wstawiony przez tę samą zmianę; usunięty znak akapitu źródła był kopiowany do każdej
        części (przyjęcie zmian sklejało podzielone akapity) – teraz zostaje przy ostatniej; spacja nierozdzielająca
        w tekście przeniesionym skąd (w:moveFrom ma w:t) była zamieniana – teraz zostaje, a we wstawce jest zamieniana.
        Miękki enter w usuniętym tekście nadal nie dzieli akapitu."""
        meta = 'w:author="Autorka" w:date="2026-01-01T00:00:00Z"'
        cases = (
            # nazwa, treść, (teksty po przyjęciu, po odrzuceniu), XPath w wyniku
            ('we wstawce', p(r('A'), f'<w:ins w:id="1" {meta}>' + r('B') + '<w:r><w:br/></w:r>' + r('C') + '</w:ins>'),
             ([['AB', 'C']], [['A']]), './w:body/w:p[1]/w:pPr/w:rPr/w:ins[@w:author="Autorka"]'),
            ('na początku wstawki', p(r('A'), f'<w:ins w:id="1" {meta}><w:r><w:br/></w:r>' + r('B') + '</w:ins>'),
             ([['A', 'B']], [['A']]), 'not(./w:body/w:p[1]/w:ins)'),  # bez pustej wstawki w pierwszym akapicie
            ('usunięty znak akapitu', p(r('A'), '<w:r><w:br/></w:r>', r('B'), properties=f'<w:rPr><w:del w:id="1" '
                                        f'{meta}/></w:rPr>') + p(r('C')),
             ([['A', 'BC']], [['A', 'B', 'C']]), './w:body/w:p[2]/w:pPr/w:rPr/w:del[@w:id="1"]'),
            ('usunięty tekst', p(r('A'), f'<w:del w:id="1" {meta}><w:r><w:delText>B</w:delText><w:br/>'
                                 '<w:delText>C</w:delText></w:r></w:del>'),
             ([['A']], [['AB\nC']]), './w:body/w:p[1]/w:del/w:r/w:br'),
            ('spacje nierozdzielające', p(r('A\u00a0B'), f'<w:moveFrom w:id="1" {meta}>' + r('E\u00a0F') + '</w:moveFrom>',
                                          f'<w:ins w:id="2" {meta}>' + r('G\u00a0H') + '</w:ins>'),
             ([['A BG H']], [['A BE\u00a0F']]), './w:body/w:p[1]/w:moveFrom/w:r/w:t[.="E\u00a0F"]'),
        )
        for label, body, (accepted, rejected), xpath in cases:
            with self.subTest(label):
                parts = review_parts(normalize_spacing_docx(review_docx_with(body)))
                self.assertEqual((resolved(parts, accept=True), resolved(parts, accept=False)), (accepted, rejected))
                self.assertTrue(parts['document'].xpath(xpath, namespaces=NS), xpath)
                self.assertEqual(review_markup_errors(parts), [])

    def test_comment_range_split_by_a_soft_break_keeps_its_words(self):
        """Zakres komentarza z miękkim enterem obejmuje po podziale oba akapity: te same słowa i to samo odwołanie."""
        body = p('<w:commentRangeStart w:id="0"/>', r('Pierwsza linia'), '<w:r><w:br/></w:r>', r('druga linia'),
                 '<w:commentRangeEnd w:id="0"/>', comment_reference(0))
        parts = review_parts(normalize_spacing_docx(review_docx_with(body, comment_xml(0, r('Uwaga.')))))
        self.assertEqual(len(parts['document'].findall(f'{qn("w:body")}/{qn("w:p")}')), 2)
        self.assertEqual(comment_anchors(parts), {'0': ('Pierwsza linia¶druga linia', 'linia')})
        self.assertEqual(review_markup_errors(parts), [])
