"""Scalanie DOCX (tylko superuser): kolejność plików, podział strony, opcje Odkurzacza, limity, pominięcia
i style szablonu Fantazmatów."""
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from unittest import skipUnless
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from lxml import etree

from core.odkurzacz_forms import MERGE_MAX_FILES, MergeForm, plural
from core.program_test_support import run_program
from core.services import program_jobs as jobs
from core.services.document_converter import ConversionError, RebuildConfirmationRequired
from core.services.document_merge import MergeError, merge_documents
from core.services.document_rebuild import rebuild_docx
from core.services.document_split import split_document
from core.services.document_template import TEMPLATE_PATH, apply_template_styles
from fantazmaty.testing import create_superuser

from .editor_fixtures_exotic import libreoffice_review_docx, review_docx, review_stories_docx
from .support import IsolatedProgramDirsMixin, docx_bytes, downloaded, noted_document, noted_text, saved
from .test_document_rebuild import comments_of, resolved, review_markup_errors, review_parts, revisions

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def story(title, *paragraphs, style=None):
    """DOCX z tytułem w stylu Nagłówek 1 i akapitami (opcjonalnie w podanym stylu, np. lista numerowana)."""
    document = Document()
    document.add_heading(title, level=1)
    for text in paragraphs:
        document.add_paragraph(text, style=style)
    return saved(document).getvalue()


def paragraph_docx(fragment):
    """DOCX z jednym akapitem z surowego OOXML (przestrzenie w: i r: zadeklarowane na akapicie)."""
    document = Document()
    document.element.body.insert(0, etree.fromstring(
        f'<w:p xmlns:w="{W_NS}" xmlns:r="{R_NS}">{fragment}</w:p>'))
    return saved(document).getvalue()


def replace_document(payload, change):
    """Kopia DOCX z word/document.xml zmienionym funkcją change(korzeń)."""
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'word/document.xml':
                root = etree.fromstring(data)
                change(root)
                data = etree.tostring(root)
            target.writestr(item, data)
    return output.getvalue()


def replace_part(payload, part, change, drop=()):
    """Kopia DOCX z częścią part zmienioną funkcją change(korzeń) i bez części z drop."""
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, 'w') as target:
        for item in source.infolist():
            if item.filename in drop:
                continue
            data = source.read(item.filename)
            if item.filename == part:
                root = etree.fromstring(data)
                change(root)
                data = etree.tostring(root)
            target.writestr(item, data)
    return output.getvalue()


def uploads(*parts):
    return [SimpleUploadedFile(name, data) for name, data in parts]


def body_paragraphs(payload):
    document = Document(BytesIO(payload))
    return [(paragraph.style.name, paragraph.text,
             paragraph._p.find(qn('w:pPr') + '/' + qn('w:pageBreakBefore')) is not None)
            for paragraph in document.paragraphs]


class MergeFormTests(SimpleTestCase):
    def form(self, files):
        return MergeForm(data={}, files={'merge-documents': files}, prefix='merge')

    def test_file_count_size_and_package_are_checked_per_file(self):
        """Brak plików, za mało i za dużo plików, zły plik (z nazwą w komunikacie) – formularz odrzuca przed zadaniem."""
        cases = (
            ('brak plików', [], 'To pole jest wymagane.'),
            ('jeden plik', uploads(('a.docx', docx_bytes())), 'co najmniej dwa pliki'),
            ('za dużo', uploads(*[(f'{n}.docx', docx_bytes()) for n in range(MERGE_MAX_FILES + 1)]),
             f'najwyżej {MERGE_MAX_FILES} plików'),
            ('nie DOCX', uploads(('a.docx', docx_bytes()), ('zly.txt', b'tekst')),
             '„zly.txt”: Wybierz plik z rozszerzeniem .docx.'),
            ('uszkodzony', uploads(('a.docx', docx_bytes()), ('zly.docx', b'to nie zip')),
             '„zly.docx”: Nie można odczytać pliku.'),
        )
        for label, files, message in cases:
            with self.subTest(label):
                form = self.form(files)
                self.assertFalse(form.is_valid())
                self.assertIn(message, ' '.join(form.errors['documents']))

    def test_polish_plural_of_file_count(self):
        """Odmiana liczby plików w komunikatach: 1 plik, 2–4 i 22–24 pliki, 5–21 i 25 plików."""
        for count, expected in ((1, 'plik'), (2, 'pliki'), (4, 'pliki'), (5, 'plików'), (12, 'plików'),
                                (21, 'plików'), (22, 'pliki'), (25, 'plików'), (112, 'plików')):
            with self.subTest(count=count):
                self.assertEqual(plural(count, 'plik', 'pliki', 'plików'), expected)

    def test_total_size_limit(self):
        """Łączny rozmiar plików ma własny limit z komunikatem po polsku."""
        with patch('core.odkurzacz_forms.MERGE_MAX_TOTAL_BYTES', 10):
            form = self.form(uploads(('a.docx', docx_bytes()), ('b.docx', docx_bytes())))
            self.assertFalse(form.is_valid())
        self.assertIn('Wybrane pliki są razem za duże', ' '.join(form.errors['documents']))

    def test_files_keep_the_selection_order_and_settings_skip_files(self):
        """Lista plików zachowuje kolejność wyboru; zapamiętane ustawienia nie zawierają plików."""
        form = self.form(uploads(('b.docx', docx_bytes()), ('a.docx', docx_bytes())))
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual([item.name for item in form.cleaned_data['documents']], ['b.docx', 'a.docx'])
        self.assertNotIn('documents', form.settings_values())
        self.assertNotIn('rebuild', form.fields)


class MergeDocumentsTests(SimpleTestCase):
    def test_content_the_rebuild_never_creates_stops_the_merge(self):
        """Odsyłacze do nieistniejących przypisów i komentarzy, numeracja list, obrazy i hiperłącza w którymkolwiek
        pliku przerywają scalanie; numeracja wyłączona (numId 0) i komentarz XML w treści nie przeszkadzają
        (komentarze z treścią są scalane – ReviewMergeTests)."""
        notes = 'bez numeracji list'
        missing = 'odsyłacz przypisu bez treści przypisu'
        relations = 'bez obrazów i hiperłączy'
        cases = (
            ('przypis', '<w:r><w:footnoteReference w:id="1"/></w:r>', missing),
            ('przypis końcowy', '<w:r><w:endnoteReference w:id="1"/></w:r>', missing),
            ('komentarz bez treści', '<w:r><w:commentReference w:id="0"/></w:r>',
             'odwołanie do komentarza bez treści komentarza'),
            ('lista', '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr></w:pPr>', notes),
            ('hiperłącze', '<w:hyperlink r:id="rId99"><w:r><w:t>link</w:t></w:r></w:hyperlink>', relations),
            ('obraz', '<w:r><w:pict><v:imagedata xmlns:v="urn:schemas-microsoft-com:vml" r:id="rId98"/></w:pict></w:r>',
             relations),
            ('wyłączona numeracja', '<w:pPr><w:numPr><w:numId w:val="0"/></w:numPr></w:pPr>', None),
            ('komentarz XML', '<!-- uwaga --><w:r><w:t>tekst</w:t></w:r>', None),
        )
        for label, fragment, message in cases:
            payload = paragraph_docx(fragment)
            for order in ((payload, docx_bytes()), (docx_bytes(), payload)):
                with self.subTest(label, first=order[0] is payload):
                    if message:
                        with self.assertRaisesMessage(MergeError, message):
                            merge_documents(list(order))
                    else:
                        self.assertEqual(len(Document(BytesIO(merge_documents(list(order)))).paragraphs), 2)

    def test_notes_of_all_files_get_new_numbers(self):
        """Przypisy dolne i końcowe kolejnych plików: numery od 1 w kolejności plików, część przypisów (z relacją,
        typem zawartości, ustawieniami i stylami) także wtedy, gdy pierwszy plik jej nie ma; w:id zmian
        w treści przypisów pozostają unikalne w całym dokumencie."""
        footnotes = rebuild_docx(noted_document(['A', 1, 'B', 2], note_text='Przypis {id}')).getvalue()
        endnotes = rebuild_docx(noted_document(['C', 1], 'endnote', note_text='Końcowy')).getvalue()
        tracked = noted_document(['D', 1], note_runs='<w:ins w:id="1" w:author="A"><w:r><w:t>wstawka</w:t></w:r></w:ins>'
                                 '<w:ins w:id="2" w:author="A"><w:r><w:t>!</w:t></w:r></w:ins>').getvalue()
        merged = merge_documents([docx_bytes('Bez przypisów'), footnotes, endnotes, footnotes, tracked, tracked])
        with ZipFile(BytesIO(merged)) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        body = etree.fromstring(files['word/document.xml'])
        self.assertEqual(noted_text(body), 'Bez przypisówA[1]B[2]C[1]A[3]B[4]D[5]D[6]')
        for kind, expected in (('footnote', ['[ref]Przypis 1', '[ref]Przypis 2', '[ref]Przypis 1', '[ref]Przypis 2',
                                             '[ref]wstawka!', '[ref]wstawka!']),
                               ('endnote', ['[ref]Końcowy'])):
            with self.subTest(kind):
                root = etree.fromstring(files[f'word/{kind}s.xml'])
                self.assertEqual([noted_text(note) for note in root.findall(qn('w:' + kind))
                                  if note.get(qn('w:type')) is None], expected)
                self.assertIn(f'/word/{kind}s.xml'.encode(), files['[Content_Types].xml'])
                self.assertIn(f'relationships/{kind}s"'.encode(), files['word/_rels/document.xml.rels'])
                self.assertIn(f'<w:{kind}Pr>'.encode(), files['word/settings.xml'])
                for style in ('Text', 'Reference'):
                    self.assertIn(f'w:styleId="{kind.capitalize()}{style}"'.encode(), files['word/styles.xml'])
        changes = [node.get(qn('w:id')) for name in ('word/document.xml', 'word/footnotes.xml')
                   for node in etree.fromstring(files[name]).iter(qn('w:ins'))]
        self.assertEqual(len(changes), len(set(changes)))
        self.assertEqual(len(changes), 4)
        Document(BytesIO(merged))

    def test_note_part_is_taken_from_the_donor_only_when_the_base_lacks_it(self):
        """Pakiet bez relacji dokumentu albo z relacją do nieistniejącej części przypisów nie ma przypisów
        (odsyłacz to błąd); pierwszy plik z własną częścią przypisów zachowuje ją; przy dodawaniu części z dawcy
        zajęty rId jest omijany, a ustawienia przypisów trafiają na koniec, gdy nie ma elementu kotwiczącego,
        i nie są dublowane, gdy już są; bez word/settings.xml albo bez stylów dawcy scalanie też działa."""
        footnotes = rebuild_docx(noted_document(['A', 1], note_text='Przypis')).getvalue()
        rels = 'word/_rels/document.xml.rels'
        missing = 'odsyłacz przypisu bez treści przypisu'
        for label, broken in (('bez relacji dokumentu', replace_part(footnotes, rels, lambda root: None, drop=(rels,))),
                              ('relacja do brakującej części', replace_part(footnotes, rels, lambda root: None,
                                                                            drop=('word/footnotes.xml',)))):
            with self.subTest(label), self.assertRaisesMessage(MergeError, missing):
                merge_documents([docx_bytes(), broken])

        def take_next_rid(root):
            root.append(etree.Element(f'{{{root.nsmap[None]}}}Relationship', Id=f'rId{len(root) + 2}',
                                      Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings',
                                      Target='settings.xml'))

        def keep_only_zoom(root):
            for node in list(root):
                if etree.QName(node).localname != 'zoom':
                    root.remove(node)

        def add_footnote_properties(root):
            root.insert(0, etree.Element(qn('w:footnotePr')))

        plain = docx_bytes('Bez przypisów')
        cases = (
            ('przypisy w pierwszym pliku', [footnotes, plain, footnotes], 'A[1]Bez przypisówA[2]', 1),
            ('zajęty rId', [replace_part(plain, rels, take_next_rid), footnotes], 'Bez przypisówA[1]', 1),
            ('ustawienia bez kotwicy', [replace_part(plain, 'word/settings.xml', keep_only_zoom), footnotes],
             'Bez przypisówA[1]', 1),
            ('ustawienia już z przypisami', [replace_part(plain, 'word/settings.xml', add_footnote_properties), footnotes],
             'Bez przypisówA[1]', 1),
            ('bez ustawień', [replace_part(plain, 'word/settings.xml', lambda root: None, drop=('word/settings.xml',)),
                              footnotes], 'Bez przypisówA[1]', None),
            ('dawca bez stylów', [plain, replace_part(footnotes, 'word/styles.xml', lambda root: None,
                                                      drop=('word/styles.xml',))], 'Bez przypisówA[1]', 1),
        )
        for label, parts, text, properties in cases:
            with self.subTest(label):
                merged = merge_documents(parts)
                with ZipFile(BytesIO(merged)) as archive:
                    files = {name: archive.read(name) for name in archive.namelist()}
                self.assertEqual(noted_text(etree.fromstring(files['word/document.xml'])), text)
                relationships = etree.fromstring(files[rels])
                ids = [node.get('Id') for node in relationships]
                self.assertEqual(len(ids), len(set(ids)))
                self.assertEqual(sum(node.get('Type', '').endswith('/footnotes') for node in relationships), 1)
                if properties is None:
                    self.assertNotIn('word/settings.xml', files)
                else:
                    self.assertEqual(files['word/settings.xml'].count(b'<w:footnotePr'), properties)
                    Document(BytesIO(merged))

    def test_missing_input_or_body_is_an_error(self):
        """Pusta lista plików albo dokument bez w:body – MergeError zamiast uszkodzonego wyniku."""
        no_body = replace_document(docx_bytes(), lambda root: root.remove(root.find(qn('w:body'))))
        for label, payloads, message in (('brak plików', [], 'Brak dokumentów'),
                                         ('bez treści', [docx_bytes(), no_body], 'nie ma treści')):
            with self.subTest(label), self.assertRaisesMessage(MergeError, message):
                merge_documents(payloads)

    def test_track_change_and_bookmark_ids_are_unique_after_merge(self):
        """Śledzone zmiany i zakładki kolejnych plików dostają nowe, unikalne w:id; para start–koniec zakładki – wspólne."""
        fragment = ('<w:bookmarkStart w:id="0" w:name="_Toc{n}"/><w:r><w:t>tekst</w:t></w:r><w:bookmarkEnd w:id="0"/>'
                    '<w:ins w:id="1" w:author="A"><w:r><w:t>wstawka</w:t></w:r></w:ins>')
        merged = Document(BytesIO(merge_documents([paragraph_docx(fragment.format(n=n)) for n in (1, 2, 3)])))
        body = merged.element.body
        starts = {node.get(qn('w:name')): node.get(qn('w:id')) for node in body.iter(qn('w:bookmarkStart'))}
        ends = [node.get(qn('w:id')) for node in body.iter(qn('w:bookmarkEnd'))]
        changes = [node.get(qn('w:id')) for node in body.iter(qn('w:ins'))]
        self.assertEqual(list(starts.values()), ends)
        self.assertEqual(len(set(ends + changes)), 6)
        self.assertEqual(starts['_Toc1'], '0')

    def test_page_format_comes_from_the_first_file(self):
        """Ustawienia strony (w:sectPr) bierzemy z pierwszego pliku; pierwszy bez w:sectPr – wynik też go nie ma."""
        from docx.enum.section import WD_ORIENT
        landscape = Document()
        landscape.sections[0].orientation = WD_ORIENT.LANDSCAPE
        landscape.sections[0].page_width, landscape.sections[0].page_height = (
            landscape.sections[0].page_height, landscape.sections[0].page_width)
        landscape.add_paragraph('Poziomo')
        landscape = saved(landscape).getvalue()
        no_section = replace_document(docx_bytes('Bez sekcji'), lambda root: root.find(qn('w:body')).remove(
            root.find(qn('w:body')).find(qn('w:sectPr'))))
        portrait = Document(BytesIO(merge_documents([docx_bytes('Pionowo'), landscape])))
        self.assertEqual(len(portrait.sections), 1)
        self.assertLess(portrait.sections[0].page_width, portrait.sections[0].page_height)
        self.assertEqual(portrait.element.body[-1].tag, qn('w:sectPr'))
        merged = Document(BytesIO(merge_documents([no_section, landscape])))
        self.assertEqual(merged.element.body.findall(qn('w:sectPr')), [])
        self.assertEqual([paragraph.text for paragraph in merged.paragraphs], ['Bez sekcji', 'Poziomo'])

    def test_each_next_file_starts_on_a_new_page_without_duplicate_breaks(self):
        """Podział strony: przed tabelą i w pustym pliku – osobny akapit; istniejący pageBreakBefore nie jest podwajany,
        a nowy stoi w w:pPr po pStyle/keepNext/keepLines (kolejność schematu)."""
        table = Document()
        table.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela'
        empty = replace_document(docx_bytes(), lambda root: root.find(qn('w:body')).remove(
            root.find(qn('w:body')).find(qn('w:p'))))
        styled = paragraph_docx('<w:pPr><w:pStyle w:val="Heading1"/><w:keepNext/><w:pageBreakBefore/>'
                                '<w:jc w:val="center"/></w:pPr><w:r><w:t>Tytuł</w:t></w:r>')
        merged = Document(BytesIO(merge_documents([docx_bytes('Start'), saved(table).getvalue(), empty, styled])))
        body = [child for child in merged.element.body if child.tag != qn('w:sectPr')]
        self.assertEqual([etree.QName(child).localname for child in body], ['p', 'p', 'tbl', 'p', 'p'])
        for index in (1, 3):
            self.assertIsNotNone(body[index].find(qn('w:pPr') + '/' + qn('w:pageBreakBefore')))
        properties = [etree.QName(child).localname for child in body[4].find(qn('w:pPr'))]
        self.assertEqual(properties, ['pStyle', 'keepNext', 'pageBreakBefore', 'jc'])


def merged_review(payloads):
    """(części recenzji scalonego pliku, scalone bajty) dla listy plików DOCX."""
    merged = merge_documents(payloads)
    return review_parts(BytesIO(merged)), merged


class ReviewMergeTests(SimpleTestCase):
    """Komentarze (z odpowiedziami i rozwiązaniem) i zmiany śledzone wszystkich scalanych plików zostają."""

    def test_comments_and_tracked_changes_of_all_files_are_merged(self):
        """Pliki z Worda 365 i LibreOffice po przebudowie, także z plikiem bez komentarzy na początku lub na końcu:
        tekst po przyjęciu i po odrzuceniu zmian (treść i przypisy), te same zmiany i komentarze (treść, zakres,
        odpowiedź, rozwiązanie); w:id komentarzy od 0, zmian – unikalne za nimi, pary zakresów wspólne, nazwy
        przeniesień unikalne; część komentarzy (z relacją, typem i stylami) także wtedy, gdy pierwszy plik jej nie ma."""
        word = rebuild_docx(review_docx()).getvalue()
        libreoffice = rebuild_docx(libreoffice_review_docx()).getvalue()
        plain = docx_bytes('Bez uwag.')
        cases = (
            ('Word i LibreOffice', [word, libreoffice], ['move1', 'move1_2']),
            ('najpierw plik bez komentarzy', [plain, libreoffice], ['move1']),
            ('na końcu plik bez komentarzy', [libreoffice, plain], ['move1']),
        )
        for label, payloads, names in cases:
            with self.subTest(label):
                sources = [review_parts(BytesIO(payload)) for payload in payloads]
                parts, merged = merged_review(payloads)
                for accept in (True, False):
                    bodies = [resolved(source, accept) for source in sources]
                    self.assertEqual(resolved(parts, accept), [[text for body in bodies for text in body[0]],
                                                               *[note for body in bodies for note in body[1:]]])
                self.assertEqual(sorted(revisions(parts), key=repr),
                                 sorted([change for source in sources for change in revisions(source)], key=repr))
                self.assertEqual(comments_of(parts), sorted([comment for source in sources
                                                             for comment in comments_of(source)], key=repr))
                self.assertEqual(review_markup_errors(parts), [])
                self.assertEqual(sorted(parts['document'].xpath('//w:moveFromRangeStart/@w:name', namespaces={
                    'w': W_NS})), names)
                document = Document(BytesIO(merged))
                self.assertEqual(len(document.comments), len(comments_of(parts)))
                self.assertIsNotNone(document.styles.element.get_by_id('CommentReference'))

    def test_rebuild_split_and_merge_keep_the_review(self):
        """Łańcuch przebudowa → podział → scalanie: teksty po przyjęciu i odrzuceniu zmian i zmiany jak w źródle,
        wszystkie komentarze z odpowiedzią i rozwiązaniem; komentarz przecięty granicą opowiadań obejmuje drugie
        opowiadanie od początku, a przeniesienie między opowiadaniami nie ma już zakresów."""
        source = review_parts(review_stories_docx())
        archive = BytesIO()
        report = split_document(rebuild_docx(review_stories_docx()), archive)
        with ZipFile(archive) as stories:
            payloads = [stories.read(name) for name in report['files']]
        parts, merged = merged_review(payloads)
        for accept in (True, False):
            self.assertEqual(resolved(parts, accept), resolved(source, accept))
        self.assertEqual(sorted(revisions(parts), key=repr), sorted(revisions(source), key=repr))
        self.assertEqual({comment[3]: comment[5:] for comment in comments_of(parts)}, {
            'Przez granicę.': ('Drugie¶Koniec komentarza', 'komentarza', None, False),
            'Punktowy.': (None, 'punkt', None, False),
            'Pytanie?': ('Pytanie', 'Pytanie', None, False),
            'Odpowiedź.': ('Pytanie', 'Pytanie', 'Pytanie?', False),
            'W przypisie.': ('uwagą', 'uwagą', None, True),
        })
        self.assertEqual(review_markup_errors(parts), [])
        self.assertEqual(parts['document'].xpath('//w:moveFromRangeStart | //w:moveToRangeStart',
                                                 namespaces={'w': W_NS}), [])
        self.assertEqual(len(Document(BytesIO(merged)).comments), 5)


class MergeReportTests(SimpleTestCase):
    def test_summaries_of_all_files_are_added_up(self):
        """Liczby reguł i zamiany słownikowe z kilku plików sumują się (wartości nieliczbowe pomijamy), zamiany od najczęstszej."""
        def row(source, target, count):
            return {'source': source, 'target': target, 'count': count}
        total = jobs._sum_summaries([
            {'spaces': 2, 'dashes': True, 'replacements': [row('kot', 'pies', 1)]},
            None,
            {'spaces': 1, 'quotes': 4, 'replacements': [row('kot', 'pies', 2), row('ala', 'ola', 5)]},
        ])
        self.assertEqual(total, {'spaces': 3, 'quotes': 4,
                                 'replacements': [row('ala', 'ola', 5), row('kot', 'pies', 3)]})
        self.assertEqual(jobs._sum_summaries([{'spaces': 1}, {}]), {'spaces': 1})

    def test_rebuild_omissions_are_cut_out_of_the_converter_message(self):
        """Z komunikatu o pominięciach zostaje sama lista elementów – z wstępem i zakończeniem albo bez nich."""
        from core.services.document_converter import REBUILD_WARNING
        for message in (REBUILD_WARNING + 'tabele (liczba: 1)' + jobs.REBUILD_SUFFIX, 'tabele (liczba: 1)'):
            with self.subTest(message=message):
                self.assertEqual(jobs._rebuild_omissions(RebuildConfirmationRequired(message)), 'tabele (liczba: 1)')


@tag('slow')
class MergeProgramTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = create_superuser('scalanie-admin')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def post(self, files, **extra):
        return run_program(self.client, {'program_action': 'merge', 'merge-documents': files,
                                         'merge-rules': ['spaces'], **extra})

    def test_files_are_merged_in_order_on_new_pages_with_cleaner_options(self):
        """Trzy pliki: kolejność wyboru, każdy kolejny od nowej strony, nagłówki zachowane, podwójne spacje usunięte."""
        response = self.post(uploads(('2.docx', story('Drugie', 'Kot  pije mleko.')),
                                     ('1.docx', story('Pierwsze', 'Ala  ma kota.')),
                                     ('3.docx', story('Trzecie', 'Pies  szczeka.'))))
        self.assertEqual(response.status_code, 200, getattr(response, 'context', None) and response.context.get('state'))
        self.assertIn('Scalone – 3 pliki.docx', unquote(response['Content-Disposition']))
        self.assertEqual(body_paragraphs(downloaded(response)), [
            ('Heading 1', 'Drugie', False), ('Normal', 'Kot pije mleko.', False),
            ('Heading 1', 'Pierwsze', True), ('Normal', 'Ala ma kota.', False),
            ('Heading 1', 'Trzecie', True), ('Normal', 'Pies szczeka.', False),
        ])

    def test_summary_lists_the_order_and_sums_the_changes(self):
        """Podsumowanie zadania podaje kolejność scalania i zsumowane liczby zmian z obu plików."""
        with patch.object(jobs, 'launch', side_effect=jobs.execute):
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'merge', 'merge-rules': ['spaces'], 'merge-cleaner_summary': 'on',
                'merge-documents': uploads(('b.docx', story('B', 'Raz  dwa.')), ('a.docx', story('A', 'Trzy  cztery.')))},
                HTTP_X_PROGRAM_JOB='1')
        state = self.client.get(response.json()['url']).json()
        self.assertEqual(state['state'], 'done', state)
        self.assertEqual(state['warnings'][0], 'Kolejność scalania: 1. „b.docx”, 2. „a.docx”.')
        self.assertEqual([(row['rule'], row['count']) for row in state['summary']], [('spaces', 2)])

    def test_numbered_lists_need_confirmation_naming_the_files(self):
        """Listy numerowane z obu plików: przebudowa pyta o pominięcie numeracji (z nazwami plików), potem scala tekst pozycji."""
        files = uploads(('lista1.docx', story('Jeden', 'pierwszy punkt', 'drugi punkt', style='List Number')),
                        ('lista2.docx', story('Dwa', 'trzeci punkt', style='List Number')))
        with patch.object(jobs, 'launch'):
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'merge', 'merge-documents': files}, HTTP_X_PROGRAM_JOB='1')
        url, folder = response.json()['url'], next(jobs.root().iterdir())
        jobs.execute(folder)
        state = self.client.get(url).json()
        self.assertEqual(state['state'], 'confirmation')
        self.assertIn('„lista1.docx” – automatyczne numerowanie', state['message'])
        self.assertIn('„lista2.docx” – automatyczne numerowanie', state['message'])
        with patch.object(jobs, 'launch'):
            url = self.client.post(url, {'action': 'confirm'}).json()['url']
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'done')
        payload = downloaded(self.client.get(url + '?download=1'))
        self.assertEqual([text for _, text, _ in body_paragraphs(payload)],
                         ['Jeden', 'pierwszy punkt', 'drugi punkt', 'Dwa', 'trzeci punkt'])
        with ZipFile(BytesIO(payload)) as archive:
            self.assertNotIn(b'w:numPr', archive.read('word/document.xml'))

    def start(self, files, **extra):
        """Zadanie scalania utworzone bez uruchamiania; zwraca (adres stanu, folder zadania)."""
        with patch.object(jobs, 'launch'):
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'merge', 'merge-documents': files, **extra}, HTTP_X_PROGRAM_JOB='1')
        return response.json()['url'], next(jobs.root().iterdir())

    def test_after_confirmation_only_files_that_needed_it_are_processed_again(self):
        """Plik bez pominięć jest przebudowany raz; po akceptacji konwerter dostaje tylko plik z listą numerowaną."""
        url, folder = self.start(uploads(('zwykly.docx', story('Zwykłe', 'Tekst.')),
                                         ('lista.docx', story('Lista', 'punkt', style='List Number'))))
        with patch.object(jobs, 'convert_document', wraps=jobs.convert_document) as convert:
            jobs.execute(folder)
            self.assertEqual([call.args[0].name for call in convert.call_args_list], ['zwykly.docx', 'lista.docx'])
            state = self.client.get(url).json()
            self.assertEqual(state['state'], 'confirmation')
            self.assertNotIn('zwykly.docx', state['message'])
            with patch.object(jobs, 'launch'):
                url = self.client.post(url, {'action': 'confirm'}).json()['url']
            convert.reset_mock()
            jobs.execute(folder)
            self.assertEqual([call.args[0].name for call in convert.call_args_list], ['lista.docx'])
        self.assertEqual(self.client.get(url).json()['state'], 'done')
        self.assertEqual([text for _, text, _ in body_paragraphs(downloaded(self.client.get(url + '?download=1')))],
                         ['Zwykłe', 'Tekst.', 'Lista', 'punkt'])
        self.assertFalse((folder / 'parts').exists())
        self.assertFalse((folder / 'sources').exists())

    def test_errors_name_the_file_or_the_size_limit(self):
        """Błąd konwertera wskazuje plik; za duży wynik scalania – osobny komunikat; części zadania są usuwane."""
        files = (('a.docx', story('A', 'Raz.')), ('b.docx', story('B', 'Dwa.')))
        cases = (
            ('błąd pliku', patch.object(jobs, 'convert_document', side_effect=ConversionError('Plik uszkodzony.')),
             '„a.docx”: Plik uszkodzony.'),
            ('za duży wynik', patch.object(jobs, 'MAX_OUTPUT', 10),
             'Scalony dokument przekroczył 50 MB. Scal mniej plików naraz.'),
        )
        for label, problem, message in cases:
            with self.subTest(label):
                url, folder = self.start(uploads(*files))
                with problem, self.assertLogs('core.services.program_jobs', 'ERROR'):
                    jobs.execute(folder)
                state = self.client.get(url).json()
                self.assertEqual((state['state'], state['message']), ('error', message))
                self.assertEqual(sorted(path.name for path in folder.iterdir() if path.is_dir()), [])
                shutil.rmtree(folder)

    def merged(self, files, **options):
        """Wynik scalania z podanymi opcjami formularza (bez prefiksu merge-)."""
        response = run_program(self.client, {'program_action': 'merge', 'merge-documents': files,
                                             'merge-rules': ['spaces'],
                                             **{f'merge-{key}': value for key, value in options.items()}})
        if not response.streaming:
            self.fail(response.content.decode()[-1500:])
        return downloaded(response)

    def test_template_styles_with_other_merge_options(self):
        """Z flagą i bez, z różnymi opcjami scalania: nagłówki wskazują style szablonu tylko z flagą, treść
        i formatowanie bezpośrednie są takie same jak bez flagi, a śledzone zmiany i format strony zostają."""
        from docx.enum.section import WD_ORIENT

        def files(landscape):
            first = Document()
            if landscape:  # bez zachowania formatu strony orientacja wymagałaby akceptacji pominięcia
                section = first.sections[0]
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width, section.page_height = section.page_height, section.page_width
            first.add_heading('Pierwsze', level=1)
            first.add_paragraph('Ala  ma kota.').runs[0].italic = True
            return uploads(('1.docx', saved(first).getvalue()), ('2.docx', story('Drugie', 'Kot  pije.')))

        cases = (
            ('domyślne', {}),
            ('śledzenie zmian', {'track_changes': 'on'}),
            ('format strony', {'keep_page_format': 'on'}),
            ('ujednolicenie', {'normalize_formatting': 'on', 'remove_soft_whitespace': 'on'}),
        )
        for label, options in cases:
            with self.subTest(label):
                landscape = 'keep_page_format' in options
                plain = self.merged(files(landscape), **options)
                styled = self.merged(files(landscape), use_template='on', **options)
                self.assertNotIn('Gwiazdki', styles_of(plain))
                self.assertIn('Gwiazdki', styles_of(styled))
                self.assertEqual([paragraph.style.style_id for paragraph in Document(BytesIO(styled)).paragraphs
                                  if paragraph.text in ('Pierwsze', 'Drugie')], ['Nagwek1', 'Nagwek1'])
                self.assertEqual(body_by_style_names(styled), body_by_style_names(plain))
                document = etree.fromstring(package(styled)['word/document.xml'])
                self.assertEqual(bool(list(document.iter(qn('w:ins'), qn('w:del')))), 'track_changes' in options)
                section = Document(BytesIO(styled)).sections[0]
                self.assertEqual(section.page_width > section.page_height, landscape)

    @skipUnless(shutil.which('soffice'), 'Brak LibreOffice')
    def test_merged_document_with_template_styles_opens_in_libreoffice(self):
        """Scalony plik ze stylami szablonu otwiera python-docx i LibreOffice zamienia go na PDF bez błędów."""
        payload = self.merged(uploads(('a.docx', story('A', 'Raz.')), ('b.docx', story('B', 'Dwa.'))), use_template='on')
        self.assertEqual(len(Document(BytesIO(payload)).paragraphs), 4)
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'scalone.docx'
            source.write_bytes(payload)
            subprocess.run(['soffice', f'-env:UserInstallation=file://{directory}/profil', '--headless',
                            '--convert-to', 'pdf', '--outdir', directory, str(source)],
                           check=True, capture_output=True, timeout=120)
            self.assertTrue((Path(directory) / 'scalone.pdf').read_bytes().startswith(b'%PDF'))


def package(payload):
    """Części pakietu DOCX/DOTX: nazwa → bajty."""
    with ZipFile(BytesIO(payload) if isinstance(payload, bytes) else payload) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def styles_of(payload):
    """Style z word/styles.xml: id → (typ, nazwa, element)."""
    root = etree.fromstring(package(payload)['word/styles.xml'])
    return {style.get(qn('w:styleId')): (style.get(qn('w:type')), style.find(qn('w:name')).get(qn('w:val'))
                                         if style.find(qn('w:name')) is not None else None, style)
            for style in root.iter(qn('w:style'))}


def template_parts():
    return package(TEMPLATE_PATH)


def body_by_style_names(payload):
    """Treść dokumentu z odwołaniami do stylów zastąpionymi nazwą stylu (małe litery) – do porównania wersji."""
    names = {style_id: (name or '').lower() for style_id, (_, name, _) in styles_of(payload).items()}
    body = etree.fromstring(package(payload)['word/document.xml']).find(qn('w:body'))
    for node in body.iter(qn('w:pStyle'), qn('w:rStyle'), qn('w:tblStyle')):
        node.set(qn('w:val'), names.get(node.get(qn('w:val')), '?'))
    for node in body.iter():
        node.attrib.pop(qn('w:date'), None)  # czas śledzonej zmiany zależy od chwili uruchomienia
    return etree.tostring(body, method='c14n')


def template_sample():
    """Dwa pliki: nagłówki 1–4, tytuł (brak w szablonie), kursywa i odstępy bezpośrednie, style znakowe."""
    first = Document()
    for level in (1, 2, 3, 4):
        first.add_heading(f'Nagłówek {level}', level=level)
    paragraph = first.add_paragraph()
    paragraph.add_run('Kursywa').italic = True
    paragraph.add_run(' i wyróżnienie', style='Emphasis')
    # Scalanie bierze style z pierwszego pliku, więc własny styl znakowy jest zdefiniowany tutaj.
    first.add_paragraph('Odnośnik').runs[0].style = first.styles.add_style('KURSYWA I BOLD', WD_STYLE_TYPE.CHARACTER)
    paragraph.paragraph_format.space_after = Pt(20)
    second = Document()
    second.add_paragraph('Tytuł zbioru', style='Title')
    return merge_documents([saved(first).getvalue(), saved(second).getvalue()])


class TemplatePackageTests(SimpleTestCase):
    def test_template_in_repository_is_a_valid_template_package(self):
        """Szablon w repo to poprawny DOTX: główna część typu template.main, style z docDefaults i stylami wydawnictwa."""
        parts = template_parts()
        self.assertIn('wordprocessingml.template.main+xml', parts['[Content_Types].xml'].decode())
        styles = {name for _, name, _ in styles_of(TEMPLATE_PATH.read_bytes()).values()}
        self.assertTrue({'Normal', 'heading 1', 'heading 2', 'heading 3', 'Podstawowy tekst', 'Gwiazdki',
                         'Cytat (motto)', 'Autor cytat', 'Kursywa', 'Rozstrzelenie'} <= styles)
        self.assertIsNotNone(etree.fromstring(parts['word/styles.xml']).find(qn('w:docDefaults')))

    def test_template_styles_need_neither_numbering_nor_theme(self):
        """Decyzja: style szablonu nie używają numeracji (w:numPr) ani motywu (czcionki, kolory, cieniowanie motywu),
        więc numbering.xml i theme1.xml szablonu nie są przenoszone. Nowy szablon z takimi odwołaniami wymaga zmian."""
        root = etree.fromstring(template_parts()['word/styles.xml'])
        self.assertEqual(list(root.iter(qn('w:numPr'))), [])
        theme = [name for node in root.iter() if isinstance(node.tag, str)
                 for name in node.attrib if 'theme' in etree.QName(name).localname.lower()]
        self.assertEqual(theme, [])


class TemplateStylesTests(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.plain = template_sample()
        cls.styled = apply_template_styles(cls.plain)

    def test_all_template_styles_defaults_and_latent_styles_are_present(self):
        """Wynik ma wszystkie style szablonu (te same id, nazwy i definicje), jego docDefaults i latentStyles."""
        template = etree.fromstring(template_parts()['word/styles.xml'])
        result = etree.fromstring(package(self.styled)['word/styles.xml'])
        styles = styles_of(self.styled)
        for style in template.iter(qn('w:style')):
            with self.subTest(style.get(qn('w:styleId'))):
                self.assertEqual(etree.tostring(styles[style.get(qn('w:styleId'))][2], method='c14n'),
                                 etree.tostring(style, method='c14n'))
        for name in ('docDefaults', 'latentStyles'):
            with self.subTest(name):
                self.assertEqual(etree.tostring(result.find(qn(f'w:{name}')), method='c14n'),
                                 etree.tostring(template.find(qn(f'w:{name}')), method='c14n'))
        self.assertNotIn('Gwiazdki', styles_of(self.plain))

    def test_headings_point_to_template_styles(self):
        """Nagłówki 1–3 wskazują na Nagwek1–Nagwek3 szablonu; Nagłówek 4 (brak w szablonie) zostaje z własną definicją."""
        document = Document(BytesIO(self.styled))
        styles = styles_of(self.styled)
        found = {paragraph.text: paragraph.style.style_id for paragraph in document.paragraphs}
        for level, expected in ((1, 'Nagwek1'), (2, 'Nagwek2'), (3, 'Nagwek3'), (4, 'Heading4')):
            with self.subTest(level=level):
                self.assertEqual(found[f'Nagłówek {level}'], expected)
                self.assertEqual(styles[expected][1], f'heading {level}')
        heading = styles['Heading4'][2]
        self.assertEqual(heading.find(qn('w:basedOn')).get(qn('w:val')), 'Normalny')
        self.assertEqual(heading.find(qn('w:next')).get(qn('w:val')), 'Normalny')
        self.assertIsNone(heading.find(qn('w:link')))  # styl znakowy „Heading 4 Char” nie jest używany

    def test_styles_missing_from_template_are_kept(self):
        """Tytuł i wyróżnienie (brak w szablonie) zostają z łańcuchem w:basedOn; „KURSYWA I BOLD” trafia do stylu
        szablonu „Kursywa i bold” (nazwy bez wielkości liter); jedynym domyślnym stylem akapitu jest Normalny szablonu."""
        styles = styles_of(self.styled)
        runs = {run.text: run.style.style_id for paragraph in Document(BytesIO(self.styled)).paragraphs
                for run in paragraph.runs}
        self.assertEqual((runs[' i wyróżnienie'], runs['Odnośnik']), ('Emphasis', 'Kursywaibold'))
        self.assertEqual(styles['Emphasis'][2].find(qn('w:basedOn')).get(qn('w:val')), 'Domylnaczcionkaakapitu')
        self.assertEqual(styles['Title'][1], 'Title')
        self.assertEqual(styles['Title'][2].find(qn('w:basedOn')).get(qn('w:val')), 'Normalny')
        self.assertEqual([style_id for style_id, (_, name, _) in styles.items() if name == 'Title'], ['Title'])
        defaults = [style_id for style_id, (kind, _, element) in styles.items()
                    if kind == 'paragraph' and element.get(qn('w:default')) == '1']
        self.assertEqual(defaults, ['Normalny'])

    def test_content_and_direct_formatting_do_not_change(self):
        """Teksty, w:pPr i w:rPr akapitów oraz ustawienia sekcji są takie same jak bez szablonu (poza id stylów)."""
        self.assertEqual(body_by_style_names(self.styled), body_by_style_names(self.plain))

    def test_nothing_but_styles_and_font_metadata_comes_from_the_template(self):
        """Wynik to zwykły .docx: bez attachedTemplate, typu template.main, docProps, ustawień, motywu i numeracji
        szablonu; fontTable dostaje tylko czcionki stylów szablonu, a stylesWithEffects (stara kopia stylów) znika."""
        before, after, template = package(self.plain), package(self.styled), template_parts()
        self.assertEqual(set(after), set(before) - {'word/stylesWithEffects.xml'})
        for name in set(after) - {'word/styles.xml', 'word/fontTable.xml', 'word/document.xml', 'word/numbering.xml',
                                  '[Content_Types].xml', 'word/_rels/document.xml.rels'}:
            with self.subTest(name):
                self.assertEqual(after[name], before[name])
        types = after['[Content_Types].xml'].decode()
        self.assertIn('/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                      'wordprocessingml.document.main+xml"', types)
        self.assertNotIn('template.main', types)
        self.assertNotIn('stylesWithEffects', types + after['word/_rels/document.xml.rels'].decode())
        self.assertNotIn(b'attachedTemplate', b''.join(after.values()))
        self.assertNotIn(b'Dawid Wiktorski', after['docProps/core.xml'])
        self.assertNotEqual(after['word/theme/theme1.xml'], template['word/theme/theme1.xml'])
        fonts = [font.get(qn('w:name')) for font in etree.fromstring(after['word/fontTable.xml'])]
        self.assertTrue({'Calibri Light', 'Segoe UI', 'Times New Roman', 'Calibri'} <= set(fonts))
        self.assertNotIn('DokChampa', fonts)  # w szablonie, ale żaden styl jej nie używa
        self.assertEqual(len(fonts), len(set(fonts)))
        body = etree.fromstring(after['word/document.xml']).find(qn('w:body'))
        self.assertEqual(etree.tostring(body.find(qn('w:sectPr')), method='c14n'),
                         etree.tostring(etree.fromstring(before['word/document.xml']).find(qn('w:body'))
                                        .find(qn('w:sectPr')), method='c14n'))

    def test_numbering_links_to_styles_stay_valid(self):
        """Poziomy list w numbering.xml wskazują style, które istnieją w wyniku (skopiowane z dotychczasowej definicji)."""
        styles = styles_of(self.styled)
        linked = [node.get(qn('w:val')) for node in
                  etree.fromstring(package(self.styled)['word/numbering.xml']).iter(qn('w:pStyle'))]
        self.assertTrue(linked)
        self.assertTrue(set(linked) <= set(styles), set(linked) - set(styles))

    def test_unusual_style_definitions(self):
        """Kolizja id ze stylem szablonu, styl bez nazwy, w:basedOn do nieistniejącego stylu, w:next/w:link do stylu
        szablonu i do brakującego, styl bez powiązań, odwołanie bez definicji (znane szablonowi albo nie) – wynik ma
        spójne odwołania."""
        def styles_change(root):
            def style(style_id, name=None, **links):
                node = etree.SubElement(root, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): style_id})
                if name:
                    etree.SubElement(node, qn('w:name'), {qn('w:val'): name})
                for key, value in links.items():
                    etree.SubElement(node, qn(f'w:{key}'), {qn('w:val'): value})
            style('Gwiazdki', 'Moje gwiazdki', basedOn='Brak', next='Heading1', link='Brak')
            style('BezNazwy', basedOn='Gwiazdki', next='BezNazwy')
            style('Samodzielny', 'Samodzielny')

        def document_change(root):
            body = root.find(qn('w:body'))
            for index, style_id in enumerate(('Gwiazdki', 'BezNazwy', 'Podstawowytekst', 'Nieznany', 'Samodzielny')):
                paragraph = etree.Element(qn('w:p'))
                etree.SubElement(etree.SubElement(paragraph, qn('w:pPr')), qn('w:pStyle'), {qn('w:val'): style_id})
                etree.SubElement(etree.SubElement(paragraph, qn('w:r')), qn('w:t')).text = style_id
                body.insert(index, paragraph)

        payload = replace_part(replace_part(docx_bytes(), 'word/styles.xml', styles_change),
                               'word/document.xml', document_change)
        result = apply_template_styles(payload)
        styles = styles_of(result)
        paragraphs = {paragraph.text: paragraph for paragraph in Document(BytesIO(result)).paragraphs}
        self.assertEqual(paragraphs['Gwiazdki'].style.style_id, 'Gwiazdki2')
        self.assertEqual(styles['Gwiazdki2'][1], 'Moje gwiazdki')
        self.assertEqual(styles['Gwiazdki'][1], 'Gwiazdki')  # styl szablonu bez zmian
        mine = styles['Gwiazdki2'][2]
        self.assertIsNone(mine.find(qn('w:basedOn')))
        self.assertIsNone(mine.find(qn('w:link')))
        self.assertEqual(mine.find(qn('w:next')).get(qn('w:val')), 'Nagwek1')
        nameless = styles['BezNazwy'][2]
        self.assertEqual((nameless.find(qn('w:basedOn')).get(qn('w:val')), nameless.find(qn('w:next')).get(qn('w:val'))),
                         ('Gwiazdki2', 'BezNazwy'))
        self.assertEqual(paragraphs['Podstawowytekst'].style.style_id, 'Podstawowytekst')
        self.assertEqual((paragraphs['Samodzielny'].style.style_id, len(styles['Samodzielny'][2])), ('Samodzielny', 1))
        self.assertIsNone(paragraphs['Nieznany']._p.find(qn('w:pPr') + '/' + qn('w:pStyle')))

    def test_package_without_font_table_or_styles_with_effects(self):
        """Bez fontTable.xml i stylesWithEffects.xml wynik dostaje style szablonu i nie dostaje nowych części."""
        payload = replace_part(docx_bytes(), 'word/styles.xml', lambda root: None,
                               drop=('word/fontTable.xml', 'word/stylesWithEffects.xml'))
        result = package(apply_template_styles(payload))
        self.assertEqual(set(result), set(package(payload)))
        self.assertIn('Gwiazdki', styles_of(apply_template_styles(payload)))
