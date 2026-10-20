"""Podział DOCX według Nagłówków 1 (tylko superuser): pliki opowiadań w ZIP, tylko własne przypisy, obrazy i komentarze;
wysyłka plików autorom (dopasowanie do tekstów antologii, lista kontrolna, SMTP, ochrona przed podwójną wysyłką)."""
import json
import smtplib
import time
from contextlib import nullcontext
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZIP_DEFLATED, ZipFile

from cryptography.fernet import InvalidToken
from django.core import mail, signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.mail import get_connection
from django.test import Client, SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree, html

from core.models import MailboxConnection
from core.program_test_support import run_program
from core.services import program_jobs as jobs
from core.services import split_mailing as mailing
from core.services.document_split import NoHeadingError, file_name, heading_style_ids, split_document
from fantazmaty.testing import create_member, create_superuser

from .support import IsolatedProgramDirsMixin, add_footnote_reference, docx_bytes, downloaded, png, saved

STORIES = (('Pierwsze opowiadanie', (20, 10)), ('Drugie: „cytat”?', (30, 15)), ('Trzecie', None))
EXPECTED = ['01 Pierwsze opowiadanie.docx', '02 Drugie „cytat”.docx', '03 Trzecie.docx']


def with_footnotes(payload, texts):
    """Dodaje do DOCX część przypisów dolnych o identyfikatorach 1, 2, … (python-docx ich nie tworzy)."""
    notes = ''.join(f'<w:footnote w:id="{number}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>'
                    for number, text in enumerate(texts, 1))
    footnotes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes '
                 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 f'{notes}</w:footnotes>')
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdFoot" Type="http://schemas.openxml'
                                    b'formats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
                                    b'</Relationships>')
            target.writestr(item, data)
        target.writestr('word/footnotes.xml', footnotes)
    return output.getvalue()


def anthology(intro=True):
    """Strona tytułowa i trzy opowiadania; dwa pierwsze mają obraz, każde przypis i komentarz."""
    document = Document()
    if intro:
        document.add_paragraph('Antologia „Zbiór” – strona tytułowa')
    for number, (title, image) in enumerate(STORIES, 1):
        document.add_heading(title, level=1)
        paragraph = document.add_paragraph(f'Treść opowiadania {number}.')
        add_footnote_reference(paragraph, number)
        if image:
            paragraph.add_run().add_picture(png(*image))
        document.add_comment(paragraph.runs[0], text=f'Komentarz opowiadania {number}', author='Redakcja')
    return with_footnotes(saved(document).getvalue(), [f'Przypis opowiadania {n}' for n in range(1, 4)])


def parts(payload):
    with ZipFile(BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def rewrite(payload, replace=None, add=None, drop=()):
    """Kopia DOCX ze zmienionymi (funkcja bajty → bajty), dodanymi i usuniętymi częściami pakietu."""
    replace, add = replace or {}, add or {}
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename not in drop:
                data = source.read(item.filename)
                target.writestr(item, replace[item.filename](data) if item.filename in replace else data)
        for name, data in add.items():
            target.writestr(name, data)
    return output.getvalue()


def add_relationships(*relationships):
    """Funkcja dopisująca relacje (Id, Type, Target[, TargetMode]) do pliku .rels."""
    def change(data):
        extra = ''.join(f'<Relationship Id="{rid}" Type="{kind}" Target="{target}"'
                        + (f' TargetMode="{mode[0]}"' if mode else '') + '/>'
                        for rid, kind, target, *mode in relationships)
        return data.replace(b'</Relationships>', extra.encode() + b'</Relationships>')
    return change


def split(payload):
    """split_document wywołany w tym procesie (coverage widzi kod, który zwykle działa w podprocesie)."""
    output = BytesIO()
    report = split_document(BytesIO(payload), output)
    return report, parts(output.getvalue())


def stories(*titles, intro=None):
    """DOCX z akapitem wstępu (opcjonalnie) i opowiadaniami: Nagłówek 1 z tytułem, akapit treści."""
    document = Document()
    if intro is not None:
        document.add_paragraph(intro)
    for number, title in enumerate(titles, 1):
        document.add_heading(title, level=1)
        document.add_paragraph(f'Treść {number}.')
    return document


def assert_story_files(test, files):
    """Pliki z anthology(): każdy ma tylko własną treść, przypis, obraz i komentarz (i żadnej miniatury)."""
    test.assertEqual(list(files), EXPECTED)
    for number, name in enumerate(EXPECTED, 1):
        with test.subTest(name=name):
            package = parts(files[name])
            document = Document(BytesIO(files[name]))
            test.assertEqual([paragraph.text for paragraph in document.paragraphs],
                             [STORIES[number - 1][0], f'Treść opowiadania {number}.'])
            everything = b''.join(package.values()).decode('utf-8', 'replace')
            test.assertNotIn('strona tytułowa', everything)
            for other in {1, 2, 3} - {number}:
                test.assertNotIn(f'Przypis opowiadania {other}', package['word/footnotes.xml'].decode())
                test.assertNotIn(f'opowiadania {other}', everything)
            test.assertIn(f'Przypis opowiadania {number}', package['word/footnotes.xml'].decode())
            test.assertIn(f'Komentarz opowiadania {number}', package['word/comments.xml'].decode())
            media = [part for part in package if part.startswith('word/media/')]
            test.assertEqual(len(media), 1 if STORIES[number - 1][1] else 0)
            if media:
                from PIL import Image
                test.assertEqual(Image.open(BytesIO(package[media[0]])).size, STORIES[number - 1][1])
            test.assertEqual(package['word/_rels/document.xml.rels'].count(b'media/'), len(media))
            # Miniatura pierwszej strony pokazuje początek całej antologii.
            test.assertNotIn('docProps/thumbnail.jpeg', list(package))
            test.assertNotIn(b'thumbnail', package['_rels/.rels'])


class SplitHelpersTests(SimpleTestCase):
    def test_heading_one_is_recognised_by_style_name_or_id(self):
        """Nagłówek 1 rozpoznajemy po nazwie stylu (bez względu na wielkość liter) albo po id Heading1/Nagwek1."""
        def styles(style_id, name):
            return (f'<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    f'<w:style w:type="paragraph" w:styleId="{style_id}"><w:name w:val="{name}"/></w:style>'
                    f'</w:styles>').encode()
        for style_id, name, expected in (('Tytuopowiadania', 'heading 1', True),
                                         ('Moj', 'Nagłówek 1', True), ('NAGWEK1', 'cokolwiek', True),
                                         ('Heading2', 'heading 2', False)):
            with self.subTest(style_id=style_id, name=name):
                self.assertEqual(style_id in heading_style_ids(styles(style_id, name)), expected)

    def test_file_names_are_numbered_cleaned_and_shortened(self):
        """Nazwa pliku: numer z zerem wiodącym, tytuł bez znaków niedozwolonych, skrócony; pusty tytuł zastąpiony."""
        for number, width, title, expected in ((1, 2, 'A/B: C?', '01 A B C.docx'), (7, 3, '  ...  ', '007 Opowiadanie.docx'),
                                               (2, 2, 'x' * 200, '02 ' + 'x' * 80 + '.docx')):
            with self.subTest(title=title):
                self.assertEqual(file_name(number, width, title), expected)


W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MS_REL = 'http://schemas.microsoft.com/office'


class SplitDocumentTests(SimpleTestCase):
    """split_document() w tym procesie: rozpoznanie nagłówków, nazwy plików, części pakietu i przypadki brzegowe."""

    def test_anthology_is_split_into_story_files(self):
        """Antologia z przypisami, obrazami i komentarzami: pliki opowiadań i informacja o pominiętym wstępie."""
        for intro in (True, False):
            with self.subTest(intro=intro):
                report, files = split(anthology(intro))
                self.assertEqual(report, {'files': EXPECTED, 'titles': [title for title, _ in STORIES],
                                          'skipped_intro': intro})
                assert_story_files(self, files)

    def test_heading_one_is_found_by_style_name_id_or_default_id(self):
        """Nagłówek 1: własny styl o nazwie „Nagłówek 1”, polskie id Nagwek1 albo Heading1 przy braku części stylów."""
        def renamed(style_id, name):
            return {'word/styles.xml': lambda data: data.replace(b'w:styleId="Heading1"', f'w:styleId="{style_id}"'.encode())
                    .replace(b'<w:name w:val="heading 1"/>', f'<w:name w:val="{name}"/>'.encode()),
                    'word/document.xml': lambda data: data.replace(b'w:val="Heading1"', f'w:val="{style_id}"'.encode())}
        payload = saved(stories('Jeden', 'Dwa')).getvalue()
        cases = (
            ('własny styl', rewrite(payload, renamed('TytulOpowiadania', 'Nagłówek 1'))),
            ('polski Word', rewrite(payload, renamed('Nagwek1', 'Nagłówek 1'))),
            ('bez stylów', rewrite(payload, {'word/_rels/document.xml.rels': lambda data: data.replace(
                b'relationships/styles"', b'relationships/x-styles"')})),
        )
        for label, variant in cases:
            with self.subTest(label):
                report, files = split(variant)
                self.assertEqual(report['files'], ['01 Jeden.docx', '02 Dwa.docx'])
                self.assertEqual(Document(BytesIO(files['02 Dwa.docx'])).paragraphs[1].text, 'Treść 2.')

    def test_document_without_heading_one_or_main_part_is_rejected(self):
        """Bez Nagłówka 1 (także bez treści) – NoHeadingError; pakiet bez głównej części – błąd zamiast pustego ZIP."""
        heading_two = Document()
        heading_two.add_heading('Rozdział', level=2)
        heading_two.add_paragraph('Tekst.')
        cases = (
            ('same akapity', docx_bytes('Sam tekst.'), NoHeadingError),
            ('Nagłówek 2', saved(heading_two).getvalue(), NoHeadingError),
            ('bez w:body', rewrite(docx_bytes(), {'word/document.xml': lambda data: data.replace(
                b'<w:body>', b'<w:x>').replace(b'</w:body>', b'</w:x>')}), NoHeadingError),
            ('bez relacji pakietu', rewrite(docx_bytes(), drop={'_rels/.rels'}), ValueError),
        )
        for label, payload, error in cases:
            with self.subTest(label), self.assertRaises(error):
                split(payload)

    def test_file_names_for_empty_duplicate_long_and_tabbed_titles(self):
        """Pusty nagłówek, powtórzony tytuł, bardzo długi tytuł i tabulator – unikalne, poprawne nazwy plików."""
        document = stories('', 'Ten sam', 'Ten sam', 'Ż' * 120)
        document.add_heading('Z\ttabem', level=1)
        report, files = split(saved(document).getvalue())
        self.assertEqual(report['files'], ['01 Opowiadanie.docx', '02 Ten sam.docx', '03 Ten sam.docx',
                                           '04 ' + 'Ż' * 80 + '.docx', '05 Z tabem.docx'])
        self.assertEqual(report['titles'], ['', 'Ten sam', 'Ten sam', 'Ż' * 120, 'Z tabem'])
        self.assertEqual(list(files), report['files'])
        self.assertEqual(Document(BytesIO(files['03 Ten sam.docx'])).paragraphs[1].text, 'Treść 3.')

    def test_intro_counts_as_skipped_only_when_visible(self):
        """Wstęp z samym obrazem jest pomijaną treścią; puste akapity przed pierwszym nagłówkiem – nie."""
        empty = stories('Jedno', intro='')
        empty.paragraphs[0].add_run('\t')
        picture = stories('Jedno', intro='')
        picture.paragraphs[0].add_run().add_picture(png(4, 4))
        for label, document, skipped in (('puste', empty, False), ('obraz', picture, True)):
            with self.subTest(label):
                report, files = split(saved(document).getvalue())
                self.assertEqual(report['skipped_intro'], skipped)
                self.assertFalse([name for name in parts(files['01 Jedno.docx']) if name.startswith('word/media/')])

    def test_story_keeps_its_table_link_and_page_header_only(self):
        """Tabela, hiperłącze i komentarz XML zostają w swoim opowiadaniu; nagłówek strony (z ustawień sekcji) – w każdym pliku."""
        document = stories('Z tabelą', 'Bez tabeli')
        document.sections[0].header.paragraphs[0].text = 'Żywa pagina'
        heading = document.paragraphs[2]
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Komórka'
        heading._p.addprevious(table._tbl)
        link = document.part.relate_to('https://example.com/opowiadanie', f'{REL}/hyperlink', is_external=True)
        document.paragraphs[1]._p.append(parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{link}"><w:r><w:t>link</w:t></w:r></w:hyperlink>'))
        document.paragraphs[1]._p.append(etree.Comment(' komentarz XML '))
        _, files = split(saved(document).getvalue())
        first, second = (parts(files[name]) for name in ('01 Z tabelą.docx', '02 Bez tabeli.docx'))
        self.assertIn(b'Kom\xc3\xb3rka', first['word/document.xml'])
        self.assertNotIn(b'<w:tbl>', second['word/document.xml'])
        self.assertIn(b'example.com/opowiadanie', first['word/_rels/document.xml.rels'])
        self.assertNotIn(b'example.com', second['word/_rels/document.xml.rels'])
        self.assertIn(b'<!-- komentarz XML -->', first['word/document.xml'])
        for package in (first, second):
            header = next(name for name in package if name.startswith('word/header'))
            self.assertIn('Żywa pagina', package[header].decode())

    def test_absolute_targets_and_document_without_section_properties(self):
        """Relacje z bezwzględną ścieżką (/word/…) i dokument bez końcowego w:sectPr dzielą się poprawnie."""
        payload = rewrite(saved(stories('Jedno', 'Dwa')).getvalue(), {
            '_rels/.rels': lambda data: data.replace(b'Target="word/document.xml"', b'Target="/word/document.xml"'),
            'word/_rels/document.xml.rels': lambda data: data.replace(b'Target="styles.xml"',
                                                                      b'Target="/word/styles.xml"'),
            'word/document.xml': lambda data: etree.tostring(_without_section(data)),
        })
        report, files = split(payload)
        self.assertEqual(report['files'], ['01 Jedno.docx', '02 Dwa.docx'])
        package = parts(files['02 Dwa.docx'])
        self.assertIn('word/styles.xml', package)
        self.assertNotIn(b'sectPr', package['word/document.xml'])
        self.assertEqual([p.style.name for p in Document(BytesIO(files['02 Dwa.docx'])).paragraphs],
                         ['Heading 1', 'Normal'])

    def test_endnotes_and_modern_comment_parts_keep_only_own_entries(self):
        """Przypisy końcowe i części komentarzy Worda 2013+ (commentsExtended/Ids/Extensible) – tylko wpisy opowiadania;
        relacja do brakującej części nie przerywa podziału."""
        document = stories('Jedno', 'Dwa')
        for number, paragraph in ((1, document.paragraphs[1]), (2, document.paragraphs[3])):
            paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:endnoteReference w:id="{number}"/></w:r>'))
            document.add_comment(paragraph.runs[0], text=f'Uwaga {number}', author='Redakcja')
        endnotes = ''.join(f'<w:endnote w:id="{n}"><w:p><w:r><w:t>Końcowy {n}</w:t></w:r></w:p></w:endnote>'
                           for n in (1, 2))

        def para_ids(data):
            root = etree.fromstring(data)
            for comment in root.iter(f'{{{W_NS}}}comment'):
                comment.find(f'{{{W_NS}}}p').set('{http://schemas.microsoft.com/office/word/2010/wordml}paraId',
                                                 f'0000000{comment.get(f"{{{W_NS}}}id")}')
            return etree.tostring(root)
        w15, cid, cex = (f'{MS_REL}/word/2012/wordml', f'{MS_REL}/word/2016/wordml/cid',
                         f'{MS_REL}/word/2018/wordml/cex')
        payload = rewrite(saved(document).getvalue(), {
            'word/comments.xml': para_ids,
            'word/_rels/document.xml.rels': add_relationships(
                ('rIdEnd', f'{REL}/endnotes', 'endnotes.xml'),
                ('rIdEx', f'{MS_REL}/2011/relationships/commentsExtended', 'commentsExtended.xml'),
                ('rIdIds', f'{MS_REL}/2016/09/relationships/commentsIds', 'commentsIds.xml'),
                ('rIdCex', f'{MS_REL}/2018/08/relationships/commentsExtensible', 'commentsExtensible.xml'),
                ('rIdMissing', f'{MS_REL}/2016/09/relationships/x-commentsIds', 'brak.xml'),
                ('rIdGlossary', f'{REL}/glossaryDocument', 'glossary/brak.xml')),
        }, add={
            'word/endnotes.xml': f'<w:endnotes xmlns:w="{W_NS}"><w:endnote w:type="separator" w:id="-1"><w:p/>'
                                 f'</w:endnote>{endnotes}</w:endnotes>',
            'word/commentsExtended.xml': f'<w15:commentsEx xmlns:w15="{w15}"><w15:commentEx w15:paraId="00000000"/>'
                                         '<w15:commentEx w15:paraId="00000001"/></w15:commentsEx>',
            'word/commentsIds.xml': f'<w16cid:commentsIds xmlns:w16cid="{cid}"><!-- Word -->'
                                    '<w16cid:commentId w16cid:paraId="00000000" w16cid:durableId="0AAA0000"/>'
                                    '<w16cid:commentId w16cid:paraId="00000001" w16cid:durableId="0BBB0000"/>'
                                    '</w16cid:commentsIds>',
            'word/commentsExtensible.xml': f'<w16cex:commentsExtensible xmlns:w16cex="{cex}"><!-- Word -->'
                                           '<w16cex:commentExtensible w16cex:durableId="0AAA0000"/>'
                                           '<w16cex:commentExtensible w16cex:durableId="0BBB0000"/>'
                                           '</w16cex:commentsExtensible>',
        })
        _, files = split(payload)
        for name, own, other in (('01 Jedno.docx', ('00000000', '0AAA0000', 'Uwaga 1', 'Końcowy 1'),
                                  ('00000001', '0BBB0000', 'Uwaga 2', 'Końcowy 2')),
                                 ('02 Dwa.docx', ('00000001', '0BBB0000', 'Uwaga 2', 'Końcowy 2'),
                                  ('00000000', '0AAA0000', 'Uwaga 1', 'Końcowy 1'))):
            with self.subTest(name):
                package = parts(files[name])
                text = ''.join(package[part].decode() for part in (
                    'word/comments.xml', 'word/commentsExtended.xml', 'word/commentsIds.xml',
                    'word/commentsExtensible.xml', 'word/endnotes.xml'))
                for value in own:
                    self.assertIn(value, text)
                for value in other:
                    self.assertNotIn(value, text)
                self.assertIn('w:type="separator"', package['word/endnotes.xml'].decode())


def _without_section(data):
    root = etree.fromstring(data)
    body = root.find(f'{{{W_NS}}}body')
    body.remove(body.find(f'{{{W_NS}}}sectPr'))
    return root


@tag('slow')
class SplitProgramTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = create_superuser('podzial-admin')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def run_split(self, payload):
        return run_program(self.client, {'program_action': 'split',
                                         'split-document': SimpleUploadedFile('Antologia.docx', payload)})

    def test_each_story_gets_only_its_own_content_notes_images_and_comments(self):
        """3 opowiadania + strona tytułowa → 3 pliki w ZIP; w każdym tylko własna treść, przypis, obraz i komentarz."""
        response = self.run_split(anthology())
        self.assertIn('Antologia – podział.zip', unquote(response['Content-Disposition']))
        assert_story_files(self, parts(downloaded(response)))

    def test_skipped_title_page_and_file_list_are_reported(self):
        """Podsumowanie wymienia utworzone pliki i informuje o pominiętej treści przed pierwszym Nagłówkiem 1."""
        for intro in (True, False):
            with self.subTest(intro=intro), patch.object(jobs, 'launch', side_effect=jobs.execute):
                response = self.client.post(reverse('core:programs'), {
                    'program_action': 'split', 'split-document': SimpleUploadedFile('A.docx', anthology(intro))},
                    HTTP_X_PROGRAM_JOB='1')
                state = self.client.get(response.json()['url']).json()
                self.assertEqual(state['state'], 'done', state)
                self.assertEqual(state['warnings'][0], 'Utworzono 3 pliki: ' + ', '.join(EXPECTED) + '.')
                self.assertEqual(any('Pominięto treść przed pierwszym Nagłówkiem 1' in text
                                     for text in state['warnings']), intro)
                downloaded(self.client.get(response.json()['url'] + '?download=1'))

    def test_converter_result_is_checked_before_download(self):
        """Wynik podprocesu: brak lub uszkodzony split.json, pusta lista albo brak archiwum – błąd; jeden plik – liczba pojedyncza."""
        def converter(report, archive=b'PK'):
            def run(directory, limit, control):
                if report is not None:
                    (directory / 'split.json').write_text(report, encoding='utf-8')
                if archive:
                    (directory / 'document.zip').write_bytes(archive)
            return run
        broken = 'Podział nie utworzył poprawnego archiwum lub wynik przekroczył 50 MB.'
        one = '{"files": ["01 A.docx"], "skipped_intro": false}'
        cases = (
            ('bez raportu', converter(None), 'error', broken),
            ('uszkodzony raport', converter('{'), 'error', broken),
            ('pusta lista', converter('{"files": []}'), 'error', broken),
            ('bez archiwum', converter(one, archive=None), 'error', broken),
            ('jeden plik', converter(one), 'done', 'Utworzono 1 plik: 01 A.docx.'),
        )
        for label, run, expected, message in cases:
            logs = self.assertLogs('core.services.program_jobs', 'ERROR') if expected == 'error' else nullcontext()
            with self.subTest(label), logs, patch.object(jobs, 'run_converter', side_effect=run), \
                    patch.object(jobs, 'launch', side_effect=jobs.execute):
                response = self.client.post(reverse('core:programs'), {
                    'program_action': 'split', 'split-document': SimpleUploadedFile('A.docx', anthology())},
                    HTTP_X_PROGRAM_JOB='1')
                state = self.client.get(response.json()['url']).json()
                self.assertEqual(state['state'], expected, state)
                self.assertEqual(state['message'] if expected == 'error' else state['warnings'][0], message)
                if expected == 'done':
                    self.assertEqual(downloaded(self.client.get(response.json()['url'] + '?download=1')), b'PK')

    def test_document_without_heading_one_is_an_error(self):
        """Bez akapitu w stylu Nagłówek 1 nie powstaje żaden plik – komunikat po polsku."""
        response = self.run_split(docx_bytes('Sam tekst bez nagłówków.'))
        self.assertContains(response, 'nie zawiera akapitów ze stylem „Nagłówek 1”', status_code=200)


# ===== Wysyłka plików autorom po podziale =====

PARTS = (('01 Pierwsze opowiadanie.docx', 'Pierwsze opowiadanie'), ('02 Drugie „cytat”.docx', 'Drugie: „cytat”?'),
         ('03 Trzecie.docx', 'Trzecie'), ('04 Wstęp redakcji.docx', 'Wstęp redakcji'), ('05 Czwarte.docx', 'Czwarte'))


def locmem_connection(mailbox):
    """Zamiast prawdziwego SMTP – skrzynka testowa Django (mail.outbox)."""
    return get_connection('django.core.mail.backends.locmem.EmailBackend')


class SplitMailingData(TestCase):
    """Antologia z autorami (z adresem, dwoje autorów, bez adresu), tekst wycofany i skrzynka wysyłkowa."""

    @classmethod
    def setUpTestData(cls):
        from authors.models import Author
        from texts.models import Anthology, Text, Translator
        from workflow.models import WorkflowStage

        cls.admin = create_superuser('wysylka-admin')
        cls.mailbox = MailboxConnection(name='Wysyłka', username='wysylka@example.test',
                                        purpose=MailboxConnection.Purpose.SPLIT_MAILING,
                                        smtp_host='smtp.example.test', smtp_port=587,
                                        smtp_security=MailboxConnection.Security.STARTTLS)
        cls.mailbox.set_password('tajne')
        cls.mailbox.save()
        cls.book = Anthology.objects.create(title='Zbiór')
        jan = Author.objects.create(first_name='Jan', last_name='Kowalski', email='jan@example.test')
        anna = Author.objects.create(first_name='Anna', last_name='Nowak', pseudonym='A. Pióro', email='anna@example.test')
        cls.silent = Author.objects.create(first_name='Ewa', last_name='Cicha', email=None)
        cls.texts = {}
        for title, writers in (('Pierwsze opowiadanie', [jan]), ('DRUGIE – "Cytat"', [jan, anna]),
                               ('Trzecie', [cls.silent]), ('Czwarte', [jan])):
            cls.texts[title] = Text.objects.create(title=title, anthology=cls.book, length=100)
            cls.texts[title].authors.add(*writers)
        WorkflowStage.objects.create(text=cls.texts['Czwarte'], stage_type='withdrawn', is_released=True,
                                     is_completed=False)
        cls.translated = Anthology.objects.create(title='Przekłady', is_translated=True)
        cls.translator = Translator.objects.create(first_name='Olga', last_name='Tłumacz', email='olga@example.test')
        foreign = Text.objects.create(title='Pierwsze opowiadanie', anthology=cls.translated, length=100)
        foreign.authors.add(jan)
        foreign.translation.translators.add(cls.translator)
        cls.untranslated = Text.objects.create(title='Trzecie', anthology=cls.translated, length=100)
        cls.extracts = cls.extract_volume()

    @classmethod
    def extract_volume(cls):
        """Antologia Ekstraktów i wpisy rejestru naboru (zmyślone): różne źródła klucza autora i adresu,
        dwa wpisy o tym samym nazwisku, wpis bez przyjętych tytułów, czarna lista i wpis innego naboru."""
        from authors.models import Author, BlacklistEntry
        from texts.models import Anthology, Extract

        book = Anthology.objects.create(title='Ekstrakty Próbne', is_extracts=True)
        entries = (
            # imię, nazwisko, pseudonim, e-mail autora, full_name wpisu, e-mail wpisu, nadesłane, przyjęte, nabór
            ('Marta', 'Lis', '', 'marta@example.test', 'Marta Lis', 'marta.wpis@example.test', 'Kot; Pies; Mysz',
             'Kot; Pies', 'ekstrakty  próbne'),
            ('Piotr', 'Zając', 'P. Zet', 'zet@example.test', 'Piotr Zając', '', 'Zając', 'Zając', 'Ekstrakty Próbne'),
            ('Karol', 'Wiśniewski-Nowy', '', 'karol@example.test', 'Karol Wiśniewski', 'karol@example.test',
             'Wiśnia; Śliwa; Grusza', 'Wiśnia; Śliwa; Grusza', 'Ekstrakty Próbne'),
            ('Ola', 'Bez', '', None, 'Ola Bez', '', 'Cisza', 'Cisza', 'Ekstrakty Próbne'),
            ('Adam', 'Mały', '', 'adam1@example.test', 'Adam Mały', 'adam1@example.test', 'A1', 'A1', 'Ekstrakty Próbne'),
            ('Adam', 'Mały', '', 'adam2@example.test', 'Adam Mały', 'adam2@example.test', 'A2', 'A2', 'Ekstrakty Próbne'),
            ('Zofia', 'Rzadka', '', 'zofia@example.test', 'Zofia Rzadka', 'zofia@example.test', 'Odrzut', '',
             'Ekstrakty Próbne'),
            ('Tomasz', 'Obcy', '', 'tomasz@example.test', 'Tomasz Obcy', 'tomasz@example.test', 'Obce', 'Obce',
             'Ekstrakty Inne'),
            ('Czarny', 'Kruk', '', 'kruk@example.test', 'Czarny Kruk', 'kruk@example.test', 'Kra', 'Kra',
             'Ekstrakty Próbne'),
            ('Biały', 'Gołąb', '', 'golab@example.test', 'Biały Gołąb', 'golab.wpis@example.test', 'Gru', 'Gru',
             'Ekstrakty Próbne'),
        )
        for first, last, pseudonym, author_email, full_name, email, title, accepted, recruitment in entries:
            author = Author.objects.create(first_name=first, last_name=last, pseudonym=pseudonym, email=author_email,
                                           is_blacklisted=last == 'Kruk')
            Extract.objects.create(author=author, full_name=full_name, email=email, title=title,
                                   accepted_titles=accepted, rejected_titles='Odrzut' if not accepted else '',
                                   recruitment=recruitment)
        BlacklistEntry.objects.create(email='golab.wpis@example.test')
        return book


class SplitMailingMatchTests(SplitMailingData):
    """Dopasowanie nagłówków do tekstów antologii i adresy odbiorców."""

    def test_title_normalization_ignores_case_spaces_quotes_and_punctuation_only(self):
        """Porównanie bez wielkości liter, białych znaków, cudzysłowów i interpunkcji – ale bez zgadywania."""
        cases = (
            ('Drugie: „cytat”?', 'DRUGIE – "Cytat"', True),
            ('Dom  nad\tjeziorem', 'Dom nad jeziorem', True),
            ('Rozdział 1.', 'rozdział 1', True),
            ('Żółw', 'Zolw', False),
            ('Ala', 'Ala 2', False),
            ('Kot', 'Koty', False),
        )
        for heading, title, same in cases:
            with self.subTest(heading=heading, title=title):
                self.assertEqual(mailing.normalized_title(heading) == mailing.normalized_title(title), same)

    def test_rows_report_match_address_and_reason(self):
        """Lista kontrolna: gotowe, brak adresu, brak dopasowania (także tekst wycofany, powtórzony nagłówek i kilka tekstów o tym tytule)."""
        from texts.models import Text

        parts = [{'file': name, 'title': title} for name, title in PARTS]
        cases = (
            ('zwykła lista', parts, [
                ('ready', 'Pierwsze opowiadanie', ['jan@example.test'], ''),
                ('ready', 'DRUGIE – "Cytat"', ['jan@example.test', 'anna@example.test'], ''),
                ('no_address', 'Trzecie', [], ''),
                ('no_match', '', [], ''),
                ('no_match', '', [], '')]),
            ('powtórzony nagłówek', parts[:1] * 2, [('no_match', '', [], 'Ten sam nagłówek ma kilka części dokumentu.')] * 2),
            ('stary wynik bez tytułów', [{'file': '01 Pierwsze opowiadanie.docx'}],
             [('ready', 'Pierwsze opowiadanie', ['jan@example.test'], '')]),
        )
        for label, rows, expected in cases:
            with self.subTest(label):
                result = mailing.build_rows(rows, self.book)
                self.assertEqual([(row['status'], row['title'], row['emails'], row['note']) for row in result], expected)
        Text.objects.create(title='Trzecie!', anthology=self.book, length=100).authors.add(self.silent)
        with self.subTest('kilka tekstów o tym tytule'):
            row = mailing.build_rows(parts[2:3], self.book)[0]
            self.assertEqual((row['status'], row['note']), ('no_match', 'Kilka tekstów antologii ma ten tytuł.'))

    def test_recipients_are_authors_or_translators(self):
        """Autor, kilku autorów (jedna wiadomość do wszystkich), tłumacz w antologii tłumaczonej, brak adresu
        i tekst bez rekordu tłumaczenia (brak odbiorców i podpisu)."""
        from authors.models import Author

        partial = self.texts['Trzecie']
        cases = (
            ('jeden autor', self.texts['Pierwsze opowiadanie'], False, (['Jan Kowalski'], ['jan@example.test'], [])),
            ('dwoje autorów', self.texts['DRUGIE – "Cytat"'], False,
             (['Jan Kowalski', 'A. Pióro'], ['jan@example.test', 'anna@example.test'], [])),
            ('tłumacz', self.translator.translations.get().text, True, (['Olga Tłumacz'], ['olga@example.test'], [])),
            ('tłumaczenie bez tłumacza', self.untranslated, True, ([], [], [])),
            ('brak adresu', partial, False, (['Ewa Cicha'], [], ['Ewa Cicha'])),
            ('antologia tłumaczona, tekst bez rekordu tłumaczenia', partial, True, ([], [], [])),
        )
        for label, text, translated, expected in cases:
            with self.subTest(label):
                self.assertEqual(mailing.recipients(text, translated), expected)
        with self.subTest('podpis tekstu bez rekordu tłumaczenia'):
            self.assertEqual(mailing.signature(partial, True), [])
        partial.authors.add(Author.objects.create(first_name='Piotr', last_name='Głośny', email='piotr@example.test'))
        with self.subTest('jeden z autorów bez adresu – wysyłka do pozostałych'):
            row = mailing.build_rows([{'file': '03 Trzecie.docx', 'title': 'Trzecie'}], self.book)[0]
            self.assertEqual((row['status'], row['emails'], row['missing']), ('ready', ['piotr@example.test'], ['Ewa Cicha']))

    def test_smtp_connection_uses_mailbox_settings(self):
        """Połączenie SMTP z danych skrzynki: serwer, port, login i hasło IMAP, SSL albo STARTTLS, limit czasu."""
        for security, ssl in ((MailboxConnection.Security.SSL, True), (MailboxConnection.Security.STARTTLS, False)):
            with self.subTest(security=security):
                self.mailbox.smtp_security = security
                backend = mailing.smtp_connection(self.mailbox)
                self.assertEqual((backend.host, backend.port, backend.username, backend.password, backend.use_ssl,
                                  backend.use_tls, backend.timeout, backend.fail_silently),
                                 ('smtp.example.test', 587, 'wysylka@example.test', 'tajne', ssl, not ssl,
                                  mailing.SMTP_TIMEOUT, False))

    def test_translated_anthology_rows_go_to_translators(self):
        """W antologii tłumaczonej plik idzie do tłumacza, [autor] to autor oryginału; tekst bez tłumacza – brak adresu."""
        from texts.models import ForeignAuthor

        self.translator.translations.get().foreign_authors.add(
            ForeignAuthor.objects.create(first_name='Hans', last_name='Christian', pseudonym='H. C. Andersen'))
        rows = mailing.build_rows([{'file': name, 'title': title} for name, title in (PARTS[0], PARTS[2])], self.translated)
        self.assertEqual([(row['status'], row['emails'], row['note'], row['authors']) for row in rows],
                         [('ready', ['olga@example.test'], '', ['H. C. Andersen']),
                          ('no_address', [], 'Tekst nie ma przypisanych tłumaczy.', [])])


    def test_extract_rows_match_heading_to_recruitment_authors(self):
        """Ekstrakty: nagłówek ↔ autor naboru (full_name, imię i nazwisko autora, pseudonim; bez wielkości liter
        i interpunkcji), adres z wpisu albo autora, przyjęte tytuły; niejednoznaczne, bez przyjętych, inny nabór,
        czarna lista i powtórzony nagłówek – bez wysyłki."""
        cases = (
            ('full_name, adres z wpisu', 'marta LIS.', ('ready', 'Kot i Pies', ['marta.wpis@example.test'], ['Marta Lis'], '')),
            ('pseudonim, adres autora', '„p zet”', ('ready', 'Zając', ['zet@example.test'], ['P. Zet'], '')),
            ('imię i nazwisko autora', 'Piotr Zając', ('ready', 'Zając', ['zet@example.test'], ['P. Zet'], '')),
            ('full_name innego niż autor', 'Karol Wiśniewski',
             ('ready', 'Wiśnia, Śliwa i Grusza', ['karol@example.test'], ['Karol Wiśniewski-Nowy'], '')),
            ('nazwisko autora', 'KAROL WIŚNIEWSKI–NOWY',
             ('ready', 'Wiśnia, Śliwa i Grusza', ['karol@example.test'], ['Karol Wiśniewski-Nowy'], '')),
            ('ogonki zostają', 'Karol Wisniewski', ('no_match', '', [], [], '')),
            ('brak adresu', 'Ola Bez', ('no_address', 'Cisza', [], ['Ola Bez'], '')),
            ('dwa wpisy', 'Adam Mały', ('no_match', '', [], [], 'Nagłówek pasuje do kilku autorów naboru.')),
            ('bez przyjętych', 'Zofia Rzadka',
             ('no_match', '', [], [], 'Autor nie ma przyjętych tekstów w tym naborze.')),
            ('inny nabór', 'Tomasz Obcy', ('no_match', '', [], [], '')),
            ('autor na czarnej liście', 'Czarny Kruk', ('blacklisted', 'Kra', ['kruk@example.test'], ['Czarny Kruk'],
                                                         'Autor jest na czarnej liście – pliku nie wysyłamy.')),
            ('adres na niezależnej czarnej liście', 'Biały Gołąb',
             ('blacklisted', 'Gru', ['golab.wpis@example.test'], ['Biały Gołąb'],
              'Autor jest na czarnej liście – pliku nie wysyłamy.')),
            ('tytuł zamiast autora', 'Kot', ('no_match', '', [], [], '')),
        )
        for label, heading, expected in cases:
            with self.subTest(label):
                row = mailing.build_rows([{'file': '01 x.docx', 'title': heading}], self.extracts)[0]
                self.assertEqual((row['status'], row['title'], row['emails'], row['people'], row['note']), expected)
        with self.subTest('powtórzony nagłówek'):
            rows = mailing.build_rows([{'file': '01 a.docx', 'title': 'Marta Lis'},
                                       {'file': '02 b.docx', 'title': 'marta lis'}], self.extracts)
            self.assertEqual({(row['status'], row['note']) for row in rows},
                             {('no_match', 'Ten sam nagłówek ma kilka części dokumentu.')})

    def test_extract_tags_use_author_signature_and_accepted_titles(self):
        """Ekstrakty: [autor] – podpis autora (pseudonim), [tytuł-tekstu] – przyjęte tytuły wpisu, także w grupie."""
        rows = mailing.build_rows([{'file': '01 a.docx', 'title': 'P. Zet'}, {'file': '02 b.docx', 'title': 'Marta Lis'}],
                                  self.extracts)
        batch = {'anthology_title': 'Ekstrakty Próbne', 'subject': '[autor]: [tytuł-tekstu] ([antologia])', 'message': ''}
        cases = (('pseudonim', rows[:1], 'P. Zet: Zając (Ekstrakty Próbne)'),
                 ('dwa przyjęte', rows[1:], 'Marta Lis: Kot i Pies (Ekstrakty Próbne)'),
                 ('kilka wierszy w mailu', rows, 'P. Zet i Marta Lis: Zając, Kot i Pies (Ekstrakty Próbne)'))
        for label, group, expected in cases:
            with self.subTest(label):
                self.assertEqual(mailing.mail_content(batch, group)[0], expected)


class SplitMailingGroupTests(SimpleTestCase):
    """Jeden mail na zbiór odbiorców: załączniki tekstów grupy, limit rozmiaru, raport i błąd całej grupy."""

    MAILBOX = MailboxConnection(username='wysylka@example.test')

    def send(self, rows, selected=None, sizes=None, connection=None):
        """Wysyła partię z wierszy (plik, adresy, stan) przez locmem; zwraca (maile, raport)."""
        folder = Path(self.enterContext(TemporaryDirectory()))
        rows = [{'index': index, 'file': f'{index + 1:02} T{index}.docx', 'title': f'T{index}', 'emails': emails,
                 'status': status} for index, (emails, status) in enumerate(rows)]
        with ZipFile(folder / 'result', 'w') as archive:
            for row in rows:
                archive.writestr(row['file'], b'x' * (sizes or {}).get(row['index'], 1))
        batch = {'subject': 'Temat', 'message': 'Treść', 'anthology_title': 'Zbiór', 'rows': rows}
        mail.outbox = []
        selected = {str(row['index']) for row in rows} if selected is None else selected
        with patch.object(mailing, 'smtp_connection', side_effect=lambda box: connection or locmem_connection(box)):
            report = mailing.send(folder, batch, selected, self.MAILBOX)
        return [(message.to, [name for name, _, _ in message.attachments]) for message in mail.outbox], report

    def test_rows_with_the_same_recipients_share_one_mail(self):
        """Jeden autor i dwa teksty – jeden mail; A+B i A – dwa maile; wielkość liter i kolejność adresów bez znaczenia."""
        a, b = 'a@example.test', 'b@example.test'
        cases = (
            ('autor z dwoma tekstami', [[a], [a]], None, [([a], ['T0.docx', 'T1.docx'])], [1, 1]),
            ('A+B i sam A', [[a, b], [a]], None, [([a, b], ['T0.docx']), ([a], ['T1.docx'])], [1, 2]),
            ('wielkość liter i kolejność', [[a, b], ['B@Example.test ', 'A@example.TEST'], [b]], None,
             [([a, b], ['T0.docx', 'T1.docx']), ([b], ['T2.docx'])], [1, 1, 2]),
            ('odznaczony wiersz', [[a], [b], [a]], {'0', '1'}, [([a], ['T0.docx']), ([b], ['T1.docx'])], [1, 2]),
        )
        for label, emails, selected, expected, numbers in cases:
            with self.subTest(label):
                mails, report = self.send([(row, mailing.READY) for row in emails], selected)
                self.assertEqual(mails, expected)
                self.assertEqual([(row['sent'], row['mail']) for row in report['rows']],
                                 [(True, number) for number in numbers])
                self.assertEqual((report['mails'], report['sent'], report['failed']), (len(expected), len(numbers), 0))

    def test_group_over_size_limit_is_split_into_several_mails(self):
        """Suma załączników ponad limit – kolejne maile, każdy w limicie; większy plik idzie sam; limit to 14 MB."""
        self.assertEqual(mailing.MAX_MAIL_BYTES, 14 * 1024 * 1024)
        with patch.object(mailing, 'MAX_MAIL_BYTES', 10):
            mails, report = self.send([(['a@example.test'], mailing.READY)] * 5, sizes={0: 6, 1: 4, 2: 3, 3: 11, 4: 1})
        self.assertEqual([names for _, names in mails],
                         [['T0.docx', 'T1.docx'], ['T2.docx'], ['T3.docx'], ['T4.docx']])
        self.assertEqual([row['mail'] for row in report['rows']], [1, 1, 2, 3, 4])

    def test_failed_group_mail_marks_all_its_rows(self):
        """Odrzucony mail grupy – błąd we wszystkich jej wierszach; inne grupy wysłane, raport liczy maile i pliki."""
        connection = locmem_connection(None)
        original = connection.send_messages

        def refuse_b(messages):
            if 'b@example.test' in messages[0].to:
                raise smtplib.SMTPRecipientsRefused({'b@example.test': (550, b'no')})
            return original(messages)
        connection.send_messages = refuse_b
        rows = [(['a@example.test'], mailing.READY), (['b@example.test'], mailing.READY),
                (['b@example.test'], mailing.READY), (['a@example.test'], mailing.NO_ADDRESS)]
        mails, report = self.send(rows, connection=connection)
        self.assertEqual(mails, [(['a@example.test'], ['T0.docx'])])
        self.assertEqual([(row['index'], row['sent'], row['mail'], row['error']) for row in report['rows']],
                         [(0, True, 1, ''), (1, False, 2, 'Serwer SMTP odrzucił adres odbiorcy.'),
                          (2, False, 2, 'Serwer SMTP odrzucił adres odbiorcy.')])
        self.assertEqual((report['mails'], report['sent'], report['failed']), (1, 1, 2))

    def test_invalid_recipient_address_fails_only_its_mail(self):
        """Adres, którego backend SMTP nie zakoduje (stare dane, np. „jan@”) – błąd tylko tego maila; pozostałe
        wychodzą, a partia kończy się raportem (wcześniej wyjątek przerywał wysyłkę i zostawiał znacznik SENT bez raportu)."""
        from django.core.mail.backends.smtp import EmailBackend

        connection = EmailBackend(host='smtp.example.test', port=25, use_tls=False, use_ssl=False)
        with patch('django.core.mail.backends.smtp.smtplib.SMTP') as server:
            _, report = self.send([(['jan@'], mailing.READY), (['a@example.test'], mailing.READY)], connection=connection)
        self.assertEqual([call.args[1] for call in server.return_value.sendmail.call_args_list], [['a@example.test']])
        self.assertEqual([(row['sent'], row['error']) for row in report['rows']],
                         [(False, mailing.INVALID_ADDRESS), (True, '')])
        self.assertEqual((report['mails'], report['sent'], report['failed']), (1, 1, 1))

    def test_checklist_numbers_planned_mails_for_ready_rows(self):
        """Lista kontrolna: maile gotowych wierszy (te same adresy – jeden mail) z tematem po uzupełnieniu tagów."""
        rows = [{'index': 0, 'title': 'A', 'emails': ['a@x.test'], 'status': mailing.READY, 'size': 1},
                {'index': 1, 'title': 'N', 'emails': [], 'status': mailing.NO_ADDRESS},
                {'index': 2, 'title': 'B', 'emails': ['b@x.test'], 'status': mailing.READY},
                {'index': 3, 'title': 'C', 'emails': ['A@x.test'], 'status': mailing.READY, 'size': 1}]
        batch = {'subject': '[tytuł-tekstu]', 'message': '', 'rows': rows}
        self.assertEqual([(number, [row['index'] for row in group], subject)
                          for number, group, subject in mailing.planned_mails(batch)],
                         [(1, [0, 3], 'A i C'), (2, [2], 'B')])

    def test_tags_are_filled_for_each_mail(self):
        """Tagi w temacie i treści: jeden, dwa i trzy teksty, wielkość liter tagu, brak tagów, tekst w nawiasach."""
        rows = [{'title': 'Dziewczynka z zapałkami', 'authors': ['Jan Kowalski']},
                {'title': 'Calineczka', 'authors': ['Jan Kowalski', 'A. Pióro']},
                {'title': 'Brzydkie kaczątko', 'authors': ['A. Pióro']}]
        cases = (
            ('jeden tekst', '[tytuł-tekstu] z antologii [antologia]', rows[:1],
             'Dziewczynka z zapałkami z antologii Bajki dla dzieci'),
            ('dwa teksty', '[tytuł-tekstu] – [autor]', rows[:2],
             'Dziewczynka z zapałkami i Calineczka – Jan Kowalski i A. Pióro'),
            ('trzy teksty', '[tytuł-tekstu]', rows, 'Dziewczynka z zapałkami, Calineczka i Brzydkie kaczątko'),
            ('wielkość liter', '[Antologia]: [TYTUŁ-TEKSTU]', rows[2:], 'Bajki dla dzieci: Brzydkie kaczątko'),
            ('bez tagów', 'Twój tekst [w załączniku]', rows[:1], 'Twój tekst [w załączniku]'),
        )
        for label, template, group, expected in cases:
            with self.subTest(label):
                batch = {'anthology_title': 'Bajki dla dzieci', 'subject': template, 'message': f'Treść:\n{template}'}
                self.assertEqual(mailing.mail_content(batch, group), (expected, f'Treść:\n{expected}'))

    def test_subject_is_one_line_within_limit(self):
        """Temat po uzupełnieniu tagów: bez nowych linii, najwyżej 255 znaków, ucięty na granicy słowa (także z „[TEST] ”)."""
        long_title = ' '.join(['słowo'] * 60)
        cases = (
            ('nowa linia w tytule', 'Tekst: [tytuł-tekstu]', 'A\r\nB', '', 'Tekst: A B'),
            ('długi tytuł', '[tytuł-tekstu]', long_title, '', ' '.join(['słowo'] * 42) + '…'),
            ('test z przedrostkiem', '[tytuł-tekstu]', long_title, '[TEST] ', '[TEST] ' + ' '.join(['słowo'] * 41) + '…'),
        )
        for label, template, title, prefix, expected in cases:
            with self.subTest(label):
                subject, _ = mailing.mail_content({'subject': template, 'message': ''}, [{'title': title}], prefix)
                self.assertEqual(subject, expected)
                self.assertLessEqual(len(subject), mailing.SUBJECT_LIMIT)


class SplitMailingServiceErrorTests(SimpleTestCase):
    """Serwis wysyłki przy uszkodzonym wyniku podziału, po wysłaniu partii i przy błędach SMTP."""

    MAILBOX = MailboxConnection(username='wysylka@example.test')
    ROW = {'index': 0, 'file': '01 A.docx', 'title': 'A', 'emails': ['a@example.test'], 'status': mailing.READY}

    def folder(self, archive=b'zip', parts=None, sent=False):
        """Katalog zadania: archiwum (b'zip' – poprawne z plikiem wiersza, inne bajty – uszkodzone, None – brak)."""
        folder = Path(self.enterContext(TemporaryDirectory()))
        if archive == b'zip':
            with ZipFile(folder / 'result', 'w') as result:
                result.writestr(self.ROW['file'], b'x' * 5)
        elif archive is not None:
            (folder / 'result').write_bytes(archive)
        if parts is not None:
            (folder / jobs.SPLIT_PARTS).write_text(parts, encoding='utf-8')
        if sent:
            (folder / mailing.SENT).touch()
        return folder

    def test_missing_or_broken_part_list_is_an_error(self):
        """Brak, uszkodzony lub niepełny split_parts.json – komunikat o ponownym podziale."""
        for label, parts in (('brak pliku', None), ('nie JSON', '{'), ('pusta lista', '[]'), ('nie lista', '{"file": "a"}'),
                             ('wiersz bez pliku', '[{"title": "A"}]'), ('wiersz nie słownik', '["01 A.docx"]')):
            with self.subTest(label):
                with self.assertRaisesMessage(mailing.MailingError, 'Uruchom podział ponownie.'):
                    mailing.read_parts(self.folder(parts=parts))

    def test_broken_archive_has_no_sizes(self):
        """Uszkodzone lub brakujące archiwum – rozmiary nieznane (0), lista kontrolna powstaje mimo to."""
        for archive in (b'nie zip', None):
            with self.subTest(archive=archive):
                self.assertEqual(mailing._sizes(self.folder(archive)), {})
        self.assertEqual(mailing._sizes(self.folder()), {'01 A.docx': 5})

    def test_sent_batch_cannot_be_prepared_or_tested_again(self):
        """Po wysłaniu (znacznik SENT) nie można przygotować nowej partii ani wysłać testu."""
        folder = self.folder(sent=True)
        for label, call, message in (
                ('nowa partia', lambda: mailing.prepare(folder, None, 'T', 'M'), 'zostały już wysłane.'),
                ('test', lambda: mailing.send_test(folder, {'rows': [self.ROW]}, 'ja@example.test', self.MAILBOX),
                 'test nie jest już możliwy')):
            with self.subTest(label):
                with self.assertRaisesMessage(mailing.MailingError, message):
                    call()

    def test_smtp_error_messages(self):
        """Błędy połączenia SMTP po polsku, bez haseł; nieznany błąd serwera – ucięty do 200 znaków."""
        cases = (
            (smtplib.SMTPAuthenticationError(535, b'no'), 'Serwer SMTP odrzucił login lub hasło skrzynki.'),
            (smtplib.SMTPRecipientsRefused({}), 'Serwer SMTP odrzucił adres odbiorcy.'),
            (InvalidToken(), 'Nie można odczytać hasła skrzynki. Zapisz je ponownie w panelu administratora.'),
            (TimeoutError(), 'Serwer SMTP nie odpowiada. Sprawdź serwer, port i szyfrowanie.'),
            (ConnectionRefusedError(), 'Serwer SMTP nie odpowiada. Sprawdź serwer, port i szyfrowanie.'),
            (smtplib.SMTPDataError(552, 'x' * 300), f'Błąd serwera SMTP: {str(smtplib.SMTPDataError(552, "x" * 300))[:200]}'),
            (OSError('gaierror'), 'Nie można połączyć się z serwerem SMTP. Sprawdź serwer, port i szyfrowanie.'),
        )
        for error, expected in cases:
            with self.subTest(error=type(error).__name__):
                self.assertEqual(mailing.smtp_error(error), expected)

    def test_send_with_missing_or_broken_files(self):
        """Plik z listy, którego nie ma w archiwum – błąd tego wiersza; uszkodzone archiwum – błąd wszystkich wierszy.
        W obu przypadkach partia nie jest zamknięta (bez znacznika SENT i SENDING), a archiwum zostaje do ponowienia."""
        missing = {**self.ROW, 'index': 1, 'file': '02 B.docx'}
        cases = (
            ('brak pliku w archiwum', b'zip', [(0, True, ''), (1, False, 'Brak pliku w archiwum podziału.')]),
            ('uszkodzone archiwum', b'nie zip', [(index, False, 'Nie można odczytać plików podziału. Uruchom podział ponownie.')
                                                 for index in (0, 1)]),
        )
        for label, archive, expected in cases:
            with self.subTest(label):
                folder = self.folder(archive)
                mail.outbox = []
                with patch.object(mailing, 'smtp_connection', side_effect=locmem_connection):
                    report = mailing.send(folder, {'subject': 'T', 'message': 'M', 'rows': [self.ROW, missing]},
                                          {'0', '1'}, self.MAILBOX)
                self.assertEqual([(row['index'], row['sent'], row['error']) for row in report['rows']], expected)
                self.assertEqual(len(mail.outbox), int(expected[0][1]))
                self.assertEqual([(folder / name).exists() for name in (mailing.SENT, mailing.SENDING, 'result')],
                                 [False, False, True])

    def test_test_mail_errors(self):
        """Test: uszkodzone archiwum i odrzucenie maila przez serwer – MailingError z komunikatem, bez maila."""
        refused = locmem_connection(None)
        refused.send_messages = lambda messages: (_ for _ in ()).throw(smtplib.SMTPRecipientsRefused({}))
        cases = (('uszkodzone archiwum', b'nie zip', locmem_connection(None), 'Uruchom podział ponownie.'),
                 ('serwer odrzucił adres', b'zip', refused, 'Serwer SMTP odrzucił adres odbiorcy.'))
        for label, archive, connection, message in cases:
            with self.subTest(label):
                mail.outbox = []
                with patch.object(mailing, 'smtp_connection', return_value=connection), \
                        self.assertRaisesMessage(mailing.MailingError, message):
                    mailing.send_test(self.folder(archive), {'subject': 'T', 'message': 'M', 'rows': [self.ROW]},
                                      'ja@example.test', self.MAILBOX)
                self.assertEqual(mail.outbox, [])


class SplitMailingRetryTests(SimpleTestCase):
    """Ponowienie częściowej wysyłki: tylko niewysłane wiersze, znacznik „trwa wysyłka”, archiwum do pełnej wysyłki."""

    MAILBOX = MailboxConnection(username='wysylka@example.test')

    def setUp(self):
        self.folder = Path(self.enterContext(TemporaryDirectory()))
        rows = [{'index': index, 'file': f'0{index + 1} T{index}.docx', 'title': f'T{index}', 'emails': [address],
                 'status': mailing.READY} for index, address in enumerate(('a@example.test', 'b@example.test'))]
        with ZipFile(self.folder / 'result', 'w') as archive:
            for row in rows:
                archive.writestr(row['file'], b'x')
        self.batch = {'batch': 'b1', 'subject': 'Temat', 'message': 'Treść', 'anthology_title': 'Zbiór', 'rows': rows}
        mail.outbox = []

    def send(self, selected=('0', '1'), refuse=()):
        """Wysyłka przez locmem; adresy z `refuse` serwer odrzuca. Zwraca adresatów tej próby i raport."""
        connection = locmem_connection(None)
        original = connection.send_messages

        def deliver(messages):
            if set(messages[0].to) & set(refuse):
                raise smtplib.SMTPRecipientsRefused({})
            return original(messages)
        connection.send_messages = deliver
        before = len(mail.outbox)
        with patch.object(mailing, 'smtp_connection', return_value=connection):
            report = mailing.send(self.folder, self.batch, set(selected), self.MAILBOX)
        return [message.to for message in mail.outbox[before:]], report

    def exists(self):
        return [(self.folder / name).exists() for name in ('result', mailing.SENT, mailing.SENDING)]

    def test_retry_sends_only_failed_rows_then_closes_batch(self):
        """Pierwsza próba: wysłany A, odrzucony B – archiwum zostaje. Ponowienie z zaznaczeniem obu wysyła tylko B,
        numer maila rośnie dalej, raport jest zbiorczy, a pełna wysyłka zamyka partię i usuwa archiwum."""
        sent, report = self.send(refuse=['b@example.test'])
        self.assertEqual((sent, report['sent'], report['failed'], self.exists()),
                         ([['a@example.test']], 1, 1, [True, False, False]))
        self.assertEqual(mailing.failed_indices(mailing.read_report(self.folder)), {1})
        with self.subTest('sam wysłany wiersz'):
            with self.assertRaisesMessage(mailing.MailingError, 'nie został jeszcze wysłany'):
                self.send(selected=['0'])
        sent, report = self.send()
        self.assertEqual(sent, [['b@example.test']])
        self.assertEqual([(row['index'], row['sent'], row['mail']) for row in report['rows']], [(0, True, 1), (1, True, 3)])
        self.assertEqual((report['sent'], report['failed'], report['mails'], report['last']),
                         (2, 0, 2, {'sent': 1, 'failed': 0, 'mails': 1}))
        self.assertEqual(self.exists(), [False, True, False])
        with self.assertRaisesMessage(mailing.MailingError, 'nie wysyłam jej drugi raz'):
            self.send()
        self.assertEqual(len(mail.outbox), 2)

    def test_send_in_progress_or_other_batch_sends_nothing(self):
        """Trwająca wysyłka (druga karta, podwójne kliknięcie) i raport innej partii – nic nie wychodzi; cudzy znacznik
        „trwa wysyłka” zostaje nietknięty."""
        (self.folder / mailing.SENDING).touch()
        with self.assertRaisesMessage(mailing.MailingError, 'Wysyłka tej partii już trwa'):
            self.send()
        self.assertTrue((self.folder / mailing.SENDING).exists())
        (self.folder / mailing.SENDING).unlink()
        jobs.write_json(self.folder / mailing.REPORT, {'batch': 'inna', 'rows': []})
        with self.assertRaisesMessage(mailing.MailingError, 'Lista kontrolna wygasła'):
            self.send()
        self.assertEqual((mail.outbox, self.exists()), ([], [True, False, False]))

    def test_unsaved_report_keeps_the_batch_locked(self):
        """Nie udało się zapisać raportu po wysłaniu – znacznik „trwa wysyłka” zostaje, więc nic nie wyjdzie drugi raz."""
        with patch.object(jobs, 'write_json', side_effect=OSError('dysk pełny')), \
                self.assertRaisesMessage(mailing.MailingError, 'Partia jest zablokowana'):
            self.send()
        self.assertEqual(self.exists(), [True, False, True])
        with self.assertRaisesMessage(mailing.MailingError, 'już trwa'):
            self.send()
        self.assertEqual(len(mail.outbox), 2)


class SplitMailingViewTests(IsolatedProgramDirsMixin, SplitMailingData):
    """Formularz, lista kontrolna, wysyłka (SMTP podmienione), raport, uprawnienia i ochrona przed podwójną wysyłką."""

    FORM = {'action': 'prepare', 'subject': 'Twój tekst [tytuł-tekstu] ([antologia])',
            'message': 'Dzień dobry, [autor],\nw załączniku plik.'}

    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)
        patcher = patch.object(mailing, 'smtp_connection', side_effect=locmem_connection)
        self.smtp = patcher.start()
        self.addCleanup(patcher.stop)

    def split_job(self, client=None, parts=PARTS, files=None, **state):
        """Zakończone zadanie podziału bez konwertera (ZIP części i lista nagłówków); zwraca adres wysyłki i folder."""
        with patch.object(jobs, 'launch'):
            response = (client or self.client).post(reverse('core:programs'), {
                'program_action': 'split', 'split-document': SimpleUploadedFile('A.docx', docx_bytes())},
                HTTP_X_PROGRAM_JOB='1')
        token = response.json()['url'].rstrip('/').rsplit('/', 1)[-1]
        folder = jobs.root() / signing.loads(token, salt=jobs.SALT)['key']
        output = BytesIO()
        with ZipFile(output, 'w') as archive:
            for name in (files if files is not None else [name for name, _ in parts]):
                archive.writestr(name, f'plik {name}'.encode())
        (folder / 'result').write_bytes(output.getvalue())
        jobs.write_json(folder / jobs.SPLIT_PARTS, [{'file': name, 'title': title} for name, title in parts])
        jobs.write_json(folder / 'state.json', {'state': 'done', 'split': True, 'mime': 'application/zip',
                                                'filename': 'A – podział.zip', 'warnings': [], **state})
        return reverse('core:split_mailing', args=[token]), folder

    def prepare(self, url, anthology=None):
        response = self.client.post(url, {**self.FORM, 'anthology': (anthology or self.book).pk})
        self.assertEqual(response.status_code, 302, response.content.decode()[:500])
        return response['Location']

    def checklist(self, url):
        page = self.client.get(url)
        return page, page.context['batch_token']

    def test_checklist_then_each_author_gets_only_own_file(self):
        """Lista kontrolna z tytułem, autorami, plikiem, adresami i stanem; każdy e-mail ma tylko plik swojego tekstu."""
        url, folder = self.split_job()
        self.assertContains(self.client.get(url), 'Przygotuj listę kontrolną')
        url = self.prepare(url)
        page, token = self.checklist(url)
        tree = html.fromstring(page.content.decode())
        cells = [[' '.join(cell.text_content().split()) for cell in row.xpath('./td')[1:]]
                 for row in tree.xpath('//table//tbody/tr')]
        self.assertEqual(cells[1], ['DRUGIE – "Cytat"', 'Jan Kowalski, A. Pióro', '02 Drugie „cytat”.docx',
                                    'jan@example.test, anna@example.test', 'Gotowe',
                                    'mail 2Temat: Twój tekst DRUGIE – "Cytat" (Zbiór)'])
        self.assertEqual([row[-2:] for row in cells], [
            ['Gotowe', 'mail 1Temat: Twój tekst Pierwsze opowiadanie (Zbiór)'], ['Gotowe', cells[1][-1]],
            ['Brak adresu', '–'], ['Brak dopasowania', '–'], ['Brak dopasowania', '–']])
        self.assertEqual(tree.xpath('//input[@name="rows"]/@value'), ['0', '1'])
        response = self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0', '1', '2', '3']}, follow=True)
        self.assertContains(response, 'Wysłane pliki: 2 z 2, wiadomości: 2')
        self.assertEqual([(message.to, message.from_email, message.subject, message.body,
                           [(name, content) for name, content, _ in message.attachments]) for message in mail.outbox], [
            (['jan@example.test'], 'wysylka@example.test', 'Twój tekst Pierwsze opowiadanie (Zbiór)',
             'Dzień dobry, Jan Kowalski,\nw załączniku plik.',
             [('Pierwsze opowiadanie.docx', b'plik 01 Pierwsze opowiadanie.docx')]),
            (['jan@example.test', 'anna@example.test'], 'wysylka@example.test', 'Twój tekst DRUGIE – "Cytat" (Zbiór)',
             'Dzień dobry, Jan Kowalski i A. Pióro,\nw załączniku plik.',
             [('Drugie „cytat”.docx', 'plik 02 Drugie „cytat”.docx'.encode())]),
        ])
        self.assertFalse((folder / 'result').exists())
        entries = [json.loads(path.read_text(encoding='utf-8')) for path in (self.work_dir / 'activity').glob('*.json')]
        self.assertIn(('Wysyłka plików autorom', 'Zbiór: wiadomości 2, pliki 2 z 2'),
                      [(entry['action'], entry['target']) for entry in entries])
        self.assertNotIn('Dzień dobry', json.dumps(entries, ensure_ascii=False))
        download = self.client.get(reverse('core:program_job', args=[url.rstrip('/').split('/')[-2]]) + '?download=1')
        self.assertContains(download, 'wysłano już autorom', status_code=404)

    def test_extracts_checklist_and_each_author_gets_only_own_part(self):
        """Ekstrakty: opis trybu nad listą, „Tytuł” = przyjęte teksty, podpis i adres autora; każdy autor dostaje
        tylko swoją część, a wiersze bez dopasowania i z czarnej listy nie wychodzą."""
        parts = (('01 Marta Lis.docx', 'Marta Lis'), ('02 P. Zet.docx', 'P. Zet'), ('03 Adam Mały.docx', 'Adam Mały'),
                 ('04 Czarny Kruk.docx', 'Czarny Kruk'), ('05 Wstęp.docx', 'Wstęp'))
        url, _ = self.split_job(parts=parts)
        url = self.prepare(url, self.extracts)
        page, token = self.checklist(url)
        self.assertContains(page, 'Tryb Ekstraktów: nagłówek części to autor.')
        tree = html.fromstring(page.content.decode())
        cells = [[' '.join(cell.text_content().split()) for cell in row.xpath('./td')[1:]]
                 for row in tree.xpath('//table//tbody/tr')]
        self.assertEqual([row[:2] + row[3:5] for row in cells], [
            ['Kot i Pies', 'Marta Lis', 'marta.wpis@example.test', 'Gotowe'],
            ['Zając', 'P. Zet', 'zet@example.test', 'Gotowe'],
            ['Nagłówek: Adam Mały', '–', '–', 'Brak dopasowaniaNagłówek pasuje do kilku autorów naboru.'],
            ['Kra', 'Czarny Kruk', 'kruk@example.test',
             'Czarna listaAutor jest na czarnej liście – pliku nie wysyłamy.'],
            ['Nagłówek: Wstęp', '–', '–', 'Brak dopasowania']])
        self.assertEqual(cells[0][-1], 'mail 1Temat: Twój tekst Kot i Pies (Ekstrakty Próbne)')
        self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0', '1', '2', '3', '4']})
        self.assertEqual([(message.to, message.body, [name for name, _, _ in message.attachments])
                          for message in mail.outbox], [
            (['marta.wpis@example.test'], 'Dzień dobry, Marta Lis,\nw załączniku plik.', ['Marta Lis.docx']),
            (['zet@example.test'], 'Dzień dobry, P. Zet,\nw załączniku plik.', ['P. Zet.docx'])])

    def test_only_checked_ready_rows_are_sent(self):
        """Odznaczony wiersz nie wychodzi; wiersze bez dopasowania lub adresu nie wychodzą nigdy."""
        cases = ((['1'], [['jan@example.test', 'anna@example.test']]), (['2', '3', '4'], []))
        for rows, expected in cases:
            with self.subTest(rows=rows):
                mail.outbox = []
                url = self.prepare(self.split_job()[0])
                _, token = self.checklist(url)
                response = self.client.post(url, {'action': 'send', 'batch': token, 'rows': rows})
                self.assertEqual([message.to for message in mail.outbox], expected)
                if not expected:
                    self.assertContains(response, 'Zaznacz co najmniej jeden wiersz', status_code=400)

    def test_second_send_of_the_same_batch_sends_nothing(self):
        """Ponowne zatwierdzenie (podwójne kliknięcie), wysyłka w toku w innej karcie i stara lista kontrolna – bez wysyłki."""
        url, folder = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        data = {'action': 'send', 'batch': token, 'rows': ['0']}
        self.client.post(url, data)
        self.assertContains(self.client.post(url, data), 'Wysłane pliki: 1 z 1')
        self.assertEqual(len(mail.outbox), 1)
        url, folder = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        (folder / mailing.SENT).touch()
        self.assertContains(self.client.post(url, {**data, 'batch': token}), 'nie wysyłam jej drugi raz', status_code=400)
        (folder / mailing.SENT).unlink()
        self.prepare(url)
        self.assertContains(self.client.post(url, {**data, 'batch': token}), 'Lista kontrolna wygasła', status_code=400)
        self.assertContains(self.client.post(url, {**data, 'batch': 'zmyślony'}), 'Lista kontrolna wygasła', status_code=400)
        self.assertEqual(len(mail.outbox), 1)

    def test_missing_mailbox_is_explained_and_nothing_is_sent(self):
        """Bez aktywnej skrzynki wysyłkowej: ostrzeżenie na formularzu i liście, wysyłka zablokowana."""
        MailboxConnection.objects.update(is_active=False)
        url, folder = self.split_job()
        self.assertContains(self.client.get(url), 'Brak skrzynki do wysyłki')
        url = self.prepare(url)
        page, token = self.checklist(url)
        self.assertContains(page, 'Brak skrzynki do wysyłki')
        response = self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0']})
        self.assertContains(response, 'Brak skrzynki do wysyłki', status_code=400)
        self.assertEqual((mail.outbox, (folder / mailing.SENT).exists()), ([], False))

    def test_smtp_failures_are_reported_and_batch_can_be_retried(self):
        """Błąd logowania SMTP – nic nie wyszło, partię można wysłać ponownie; odrzucony adres – błąd tylko w tym wierszu."""
        url, folder = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        data = {'action': 'send', 'batch': token, 'rows': ['0', '1']}
        broken = locmem_connection(None)
        broken.open = lambda: (_ for _ in ()).throw(smtplib.SMTPAuthenticationError(535, b'no'))
        self.smtp.side_effect = lambda mailbox: broken
        response = self.client.post(url, data)
        self.assertContains(response, 'Serwer SMTP odrzucił login lub hasło skrzynki.')
        self.assertEqual((mail.outbox, (folder / mailing.SENT).exists(), (folder / 'result').exists()), ([], False, True))
        partial = locmem_connection(None)
        original = partial.send_messages

        def refuse_anna(messages):
            if 'anna@example.test' in messages[0].to:
                raise smtplib.SMTPRecipientsRefused({'anna@example.test': (550, b'no')})
            return original(messages)
        partial.send_messages = refuse_anna
        self.smtp.side_effect = lambda mailbox: partial
        response = self.client.post(url, data, follow=True)
        self.assertContains(response, 'Wysłane pliki: 1 z 2, wiadomości: 1')
        self.assertContains(response, 'Serwer SMTP odrzucił adres odbiorcy.')
        self.assertEqual(len(mail.outbox), 1)

    def test_partial_send_can_be_retried_for_failed_rows_only(self):
        """Po częściowej wysyłce: raport z „Wyślij ponownie nieudane”, archiwum zostaje, test jest dozwolony, nowa lista
        kontrolna i ponowna wysyłka całej partii nic nie wysyłają; ponowienie wysyła tylko nieudany wiersz, ma osobny wpis
        w dzienniku, zamyka partię i usuwa archiwum – potem test i kolejne ponowienie są zablokowane."""
        url, folder = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        partial = locmem_connection(None)
        original = partial.send_messages

        def refuse_anna(messages):
            if 'anna@example.test' in messages[0].to:
                raise smtplib.SMTPRecipientsRefused({'anna@example.test': (550, b'no')})
            return original(messages)
        partial.send_messages = refuse_anna
        self.smtp.side_effect = lambda mailbox: partial
        page = self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0', '1']}, follow=True)
        self.assertContains(page, 'Wyślij ponownie nieudane (1)')
        self.assertEqual([(folder / name).exists() for name in ('result', mailing.SENT, mailing.SENDING)],
                         [True, False, False])
        retry_token = page.context['batch_token']
        with self.subTest('mail testowy po częściowej wysyłce'):
            self.smtp.side_effect = locmem_connection
            response = self.client.post(url, {'action': 'test', 'batch': retry_token, 'test_address': 'ja@example.test'},
                                        follow=True)
            self.assertContains(response, 'Wysłano test na ja@example.test')
        with self.subTest('bez nowej listy kontrolnej i bez ponownej wysyłki całej partii'):
            for data in ({**self.FORM, 'anthology': self.book.pk}, {'action': 'send', 'batch': token, 'rows': ['0', '1']}):
                self.assertContains(self.client.post(url, data), 'Wysłane pliki: 1 z 2')
        self.assertEqual(len(mail.outbox), 2)
        response = self.client.post(url, {'action': 'retry', 'batch': retry_token}, follow=True)
        self.assertContains(response, 'Ponowiono wysyłkę: pliki 1 z 1.')
        self.assertContains(response, 'Wysłane pliki: 2 z 2, wiadomości: 2')
        self.assertNotContains(response, 'Wyślij ponownie nieudane')
        self.assertEqual([message.to for message in mail.outbox[2:]], [['jan@example.test', 'anna@example.test']])
        self.assertEqual([(folder / name).exists() for name in ('result', mailing.SENT)], [False, True])
        entries = [json.loads(path.read_text(encoding='utf-8')) for path in (self.work_dir / 'activity').glob('*.json')]
        self.assertTrue({('Wysyłka plików autorom', 'Zbiór: wiadomości 1, pliki 1 z 2'),
                         ('Ponowienie wysyłki plików autorom', 'Zbiór: wiadomości 1, pliki 1 z 1')}
                        <= {(entry['action'], entry['target']) for entry in entries})
        for data, expected, status in (({'action': 'retry', 'batch': retry_token}, 'Wysłane pliki: 2 z 2', 200),
                                       ({'action': 'test', 'batch': retry_token, 'test_address': 'ja@example.test'},
                                        'test nie jest już możliwy', 400)):
            with self.subTest(action=data['action']):
                self.assertContains(self.client.post(url, data), expected, status_code=status)
        self.assertEqual(len(mail.outbox), 3)

    def test_test_mail_goes_to_owner_and_leaves_batch_ready(self):
        """„Wyślij test”: jeden mail „[TEST]” z plikiem losowego gotowego wiersza; adres w sesji, wpis w dzienniku,
        partia nie jest oznaczona jako wysłana, a właściwa wysyłka działa dalej."""
        url, folder = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        batch = mailing.read_batch(folder)
        with patch.object(mailing.random, 'choice', side_effect=lambda rows: rows[-1]) as choice:
            response = self.client.post(url, {'action': 'test', 'batch': token, 'test_address': 'ja@example.test'},
                                        follow=True)
        self.assertEqual([row['index'] for row in choice.call_args.args[0]], [0, 1])
        self.assertContains(response, 'Wysłano test na ja@example.test (plik: Drugie „cytat”.docx).')
        self.assertEqual([(message.to, message.subject, message.body, [name for name, _, _ in message.attachments])
                          for message in mail.outbox],
                         [(['ja@example.test'], '[TEST] Twój tekst DRUGIE – "Cytat" (Zbiór)',
                           'Dzień dobry, Jan Kowalski i A. Pióro,\nw załączniku plik.',
                           ['Drugie „cytat”.docx'])])
        self.assertEqual(response.context['test_form'].initial['test_address'], 'ja@example.test')
        self.assertEqual((mailing.read_batch(folder), (folder / mailing.SENT).exists()), (batch, False))
        entries = [json.loads(path.read_text(encoding='utf-8')) for path in (self.work_dir / 'activity').glob('*.json')]
        self.assertIn(('Test wysyłki plików autorom', 'Zbiór: test na ja@example.test'),
                      [(entry['action'], entry['target']) for entry in entries])
        self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0', '1']})
        self.assertEqual(len(mail.outbox), 3)
        response = self.client.post(url, {'action': 'test', 'batch': token, 'test_address': 'ja@example.test'})
        self.assertContains(response, 'test nie jest już możliwy', status_code=400)
        self.assertEqual(len(mail.outbox), 3)

    def test_test_mail_file_choice_and_errors(self):
        """Bez gotowych wierszy – losowanie z wierszy z plikiem; bez plików, zły adres, błąd SMTP – komunikat, bez maila."""
        def refuse(mailbox):
            broken = locmem_connection(None)
            broken.open = lambda: (_ for _ in ()).throw(smtplib.SMTPAuthenticationError(535, b'no'))
            return broken

        cases = (
            ('brak gotowych', PARTS[2:4], None, {}, 'ja@example.test', [0, 1], 'Wysłano test na ja@example.test'),
            ('brak plików', PARTS, [], {}, 'ja@example.test', None, 'Brak plików do wysłania w teście.'),
            ('zły adres', PARTS, None, {}, 'nie-adres', None, 'Wprowadź poprawny adres email.'),
            ('błąd SMTP', PARTS, None, {'side_effect': refuse}, 'ja@example.test', [0, 1],
             'Serwer SMTP odrzucił login lub hasło skrzynki.'),
        )
        for label, parts, files, smtp, address, candidates, expected in cases:
            with self.subTest(label):
                mail.outbox = []
                url, folder = self.split_job(parts=parts, files=files)
                url = self.prepare(url)
                _, token = self.checklist(url)
                with patch.object(mailing.random, 'choice', side_effect=lambda rows: rows[0]) as choice, \
                        patch.object(mailing, 'smtp_connection', **(smtp or {'side_effect': locmem_connection})):
                    response = self.client.post(url, {'action': 'test', 'batch': token, 'test_address': address},
                                                follow=True)
                self.assertContains(response, expected, status_code=200 if label == 'brak gotowych' else 400)
                self.assertEqual([row['index'] for row in choice.call_args.args[0]] if choice.called else None,
                                 candidates)
                self.assertEqual(len(mail.outbox), int(label == 'brak gotowych'))
                self.assertNotIn('tajne', response.content.decode())
                self.assertFalse((folder / mailing.SENT).exists())

    def test_test_mail_only_for_superuser_with_valid_checklist(self):
        """Test jak wysyłka: redaktor – 403, inny superuser – 404, nieaktualna lista kontrolna – komunikat."""
        url, _ = self.split_job()
        url = self.prepare(url)
        _, token = self.checklist(url)
        data = {'action': 'test', 'batch': token, 'test_address': 'ja@example.test'}
        for label, user, expected in (('redaktor', create_member('test-redaktor', 'Redaktor'), 403),
                                      ('inny superuser', create_superuser('test-inny'), 404)):
            with self.subTest(label):
                client = Client()
                client.force_login(user)
                self.assertEqual(client.post(url, data).status_code, expected)
        with self.subTest('stara lista kontrolna'):
            self.prepare(url)
            self.assertContains(self.client.post(url, data), 'Lista kontrolna wygasła', status_code=400)
        self.assertEqual(mail.outbox, [])

    def test_form_validation_and_changing_settings(self):
        """Temat w jednym wierszu, wymagana antologia, tylko znane tagi, treść bez znaków sterujących;
        „Zmień” wraca do formularza z poprzednimi wartościami."""
        url, _ = self.split_job()
        unknown = 'Nieznane tagi: [tytul]. Dostępne: [antologia], [tytuł-tekstu], [autor].'
        cases = (({'subject': 'Temat\nBcc: ktoś@example.test'}, 'subject', None), ({'anthology': ''}, 'anthology', None),
                 ({'message': ''}, 'message', None), ({'subject': '[tytul] z [ANTOLOGIA]'}, 'subject', unknown),
                 ({'message': 'Plik [tytul], [Autor]'}, 'message', unknown),
                 ({'message': 'Treść\x07z dzwonkiem'}, 'message', 'Treść: usuń znaki sterujące.'))
        for change, field, error in cases:
            with self.subTest(field=field, error=error):
                response = self.client.post(url, {**self.FORM, 'anthology': self.book.pk, **change})
                self.assertEqual(response.status_code, 400)
                self.assertIn(field, response.context['form'].errors)
                if error:
                    self.assertEqual(response.context['form'].errors[field], [error])
        self.assertContains(self.client.get(url), '[tytuł-tekstu] – tytuł tekstu')
        url = self.prepare(url)
        response = self.client.post(url, {'action': 'change'})
        self.assertEqual(response.context['form'].initial['subject'], self.FORM['subject'])
        self.assertIsNone(mailing.read_batch(next(jobs.root().iterdir())))

    def test_access_only_for_superuser_owner_of_finished_split(self):
        """Członek zespołu – 403; inny superuser (inna sesja) – 404; zadanie nie-podziału lub w toku – 404."""
        url, folder = self.split_job()
        member = create_member('wysylka-redaktor', 'Redaktor')
        other = create_superuser('wysylka-inny')
        for label, user, expected in (('redaktor', member, 403), ('inny superuser', other, 404)):
            with self.subTest(label):
                client = Client()
                client.force_login(user)
                self.assertEqual(client.post(url, {**self.FORM, 'anthology': self.book.pk}).status_code, expected)
        for label, state in (('w toku', {'state': 'running', 'stage': 'x', 'heartbeat': time.time()}),
                             ('nie podział', {'state': 'done', 'filename': 'a.docx', 'mime': 'x', 'warnings': []})):
            with self.subTest(label):
                jobs.write_json(folder / 'state.json', state)
                self.assertContains(self.client.get(url), 'po zakończonym podziale DOCX', status_code=404)
        self.assertEqual(mail.outbox, [])

    def test_errors_without_checklist_broken_job_or_unknown_action(self):
        """Wysyłka i test bez listy kontrolnej – formularz z komunikatem; nieczytelny stan zadania – 503;
        brak listy plików podziału przy przygotowaniu, test bez skrzynki i nieznana akcja – 400. Nic nie wychodzi."""
        def no_parts(folder):
            (folder / jobs.SPLIT_PARTS).unlink()

        def no_mailbox(folder):
            self.prepare(self.url)
            MailboxConnection.objects.update(is_active=False)

        cases = (
            ('wysyłka bez listy', None, lambda url: {'action': 'send', 'batch': 'x', 'rows': ['0']}, 400,
             'Lista kontrolna wygasła'),
            ('test bez listy', None, lambda url: {'action': 'test', 'batch': 'x', 'test_address': 'ja@example.test'},
             400, 'Lista kontrolna wygasła'),
            ('uszkodzony stan zadania', lambda folder: (folder / 'state.json').write_text('{', encoding='utf-8'),
             lambda url: {}, 503, 'Nie można odczytać zadania. Spróbuj ponownie.'),
            ('brak listy plików', no_parts, lambda url: {**self.FORM, 'anthology': self.book.pk}, 400,
             'Brak listy plików z tego podziału. Uruchom podział ponownie.'),
            ('test bez skrzynki', no_mailbox,
             lambda url: {'action': 'test', 'batch': self.checklist(url)[1], 'test_address': 'ja@example.test'}, 400,
             'Brak skrzynki do wysyłki'),
            ('nieznana akcja', None, lambda url: {'action': 'usuń'}, 400, 'Nieprawidłowa akcja.'),
        )
        for label, setup, data, status, message in cases:
            with self.subTest(label):
                MailboxConnection.objects.update(is_active=True)
                self.url, folder = self.split_job()
                if setup:
                    setup(folder)
                response = self.client.post(self.url, data(self.url)) if data(self.url) else self.client.get(self.url)
                self.assertContains(response, message, status_code=status)
                if label.endswith('bez listy'):
                    self.assertIn('form', response.context)
        self.assertEqual(mail.outbox, [])

    def test_activity_log_failure_does_not_break_sending(self):
        """Błąd zapisu dziennika aktywności (OSError) trafia do logu, a wysyłka kończy się raportem."""
        url = self.prepare(self.split_job()[0])
        _, token = self.checklist(url)
        with patch('core.activity_spool.enqueue_activity', side_effect=OSError('dysk pełny')), \
                self.assertLogs(level='ERROR') as logs:  # także middleware aktywności – ten sam zapis
            response = self.client.post(url, {'action': 'send', 'batch': token, 'rows': ['0']}, follow=True)
        self.assertContains(response, 'Wysłane pliki: 1 z 1')
        self.assertTrue(any(line.startswith('ERROR:core.views.split_mailing:') for line in logs.output), logs.output)
        self.assertEqual(len(mail.outbox), 1)

    def test_job_page_offers_mailing_and_split_result_survives_download(self):
        """Strona i JSON zadania podziału mają „Wyślij autorom”; po pobraniu ZIP archiwum zostaje dla wysyłki."""
        url, folder = self.split_job()
        job_url = reverse('core:program_job', args=[url.rstrip('/').split('/')[-2]])
        self.assertEqual(self.client.get(job_url).json()['mail_url'], url)
        self.assertContains(self.client.get(job_url + '?view=page'), f'href="{url}"')
        downloaded(self.client.get(job_url + '?download=1'))
        self.assertTrue((folder / 'result').exists())
        self.assertEqual(self.client.get(url).status_code, 200)


@tag('slow')
class SplitMailingEndToEndTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        from authors.models import Author
        from texts.models import Anthology, Text

        cls.admin = create_superuser('wysylka-e2e')
        box = MailboxConnection(name='Wysyłka', username='wysylka@example.test',
                                purpose=MailboxConnection.Purpose.SPLIT_MAILING, smtp_host='smtp.example.test')
        box.set_password('tajne')
        box.save()
        cls.book = Anthology.objects.create(title='Zbiór')
        for number, (title, _) in enumerate(STORIES, 1):
            author = Author.objects.create(first_name='Autor', last_name=str(number), email=f'a{number}@example.test')
            Text.objects.create(title=title.replace(':', ''), anthology=cls.book, length=100).authors.add(author)

    def test_real_split_files_reach_their_authors(self):
        """Prawdziwy podział antologii: nagłówki z dokumentu pasują do tytułów, każdy autor dostaje swój plik."""
        self.client.force_login(self.admin)
        with patch.object(jobs, 'launch', side_effect=jobs.execute):
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'split', 'split-document': SimpleUploadedFile('A.docx', anthology())},
                HTTP_X_PROGRAM_JOB='1')
        url = self.client.get(response.json()['url']).json()['mail_url']
        self.client.post(url, {'action': 'prepare', 'anthology': self.book.pk, 'subject': 'Plik', 'message': 'Treść'})
        page = self.client.get(url)
        with patch.object(mailing, 'smtp_connection', side_effect=locmem_connection):
            self.client.post(url, {'action': 'send', 'batch': page.context['batch_token'], 'rows': ['0', '1', '2']})
        self.assertEqual([(message.to, message.attachments[0][0]) for message in mail.outbox],
                         [([f'a{number}@example.test'], name.split(' ', 1)[1]) for number, name in enumerate(EXPECTED, 1)])
        for message, (title, _) in zip(mail.outbox, STORIES):
            self.assertEqual(Document(BytesIO(message.attachments[0][1])).paragraphs[0].text, title)
