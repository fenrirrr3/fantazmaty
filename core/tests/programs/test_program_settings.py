"""Nowe opcje programów: zapamiętane ustawienia, śledzenie zmian, podsumowanie, format strony, reguły w konwerterze i limity superusera."""
import json
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, tag
from django.urls import reverse
from docx import Document

from core.models import ProgramPreferences, ReplacementEntry
from core.odkurzacz_forms import (
    FORM_DEFAULT_RULE_LIST, DocumentConversionForm, OdkurzaczForm,
)
from core.services.odkurzacz import EDITORIAL_RULES
from core.services import document_conversion_worker as worker
from core.services import document_preparation as preparation
from core.services import document_repetitions as repetitions
from core.services import program_jobs as jobs
from core.services.document_converter import TIME_LIMIT, ConversionError, convert_document
from core.services.document_errors import DocumentInputError
from fantazmaty.testing import create_member, create_superuser

from .support import REPETITION_FIELDS, IsolatedProgramDirsMixin, formatted_docx, saved


def upload(name='tekst.docx'):
    doc = Document()
    doc.add_paragraph('Ala  ma kota.')
    return SimpleUploadedFile(name, saved(doc).getvalue())


class ProgramViewMixin(IsolatedProgramDirsMixin):
    url = reverse('core:programs')

    def post_job(self, data):
        """Wysyła formularz programu bez uruchamiania procesu; zwraca opcje zapisanego zadania."""
        with patch.object(jobs, 'launch'):
            response = self.client.post(self.url, data, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 202, response.content)
        folder = max(jobs.root().iterdir(), key=lambda path: path.stat().st_mtime_ns)
        request = jobs.read_json(folder / 'request.json')
        jobs.write_json(folder / 'state.json', {'state': 'cancelled'})
        return request['options']


class RememberedSettingsTests(ProgramViewMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('ustawienia', 'Redaktor')
        cls.other = create_member('ustawienia-inna', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def stored(self, user=None):
        return ProgramPreferences.objects.get(user=user or self.member).data

    def test_valid_submission_remembers_settings_without_the_file(self):
        self.post_job({'program_action': 'clean', 'document': upload(), 'rules': ['spaces', 'quotes'],
                       'track_changes': 'on', 'rebuild': 'on', 'keep_page_format': 'on'})
        data = self.stored()['clean']
        self.assertNotIn('document', data)
        self.assertEqual(data['rules'], ['spaces', 'quotes'])
        self.assertIs(data['track_changes'], True)
        self.assertIs(data['remove_soft_whitespace'], False)
        self.assertIs(data['cleaner_summary'], False)
        form = self.client.get(self.url).context['form']
        self.assertEqual(form.initial['rules'], ['spaces', 'quotes'])
        self.assertIs(form.initial['track_changes'], True)
        self.assertIs(form.initial['keep_page_format'], True)
        self.assertNotIn('document', form.initial)

    def test_each_program_is_remembered_separately_and_per_user(self):
        self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['pdf'],
                       'convert-use_cleaner': 'on', 'convert-rules_shown': 'True', 'convert-rules': ['spaces']})
        self.post_job({'program_action': 'repetitions', 'repetitions-document': upload(),
                       **REPETITION_FIELDS, 'repetitions-window_size': '50', 'repetitions-color_palette': 'light'})
        data = self.stored()
        self.assertEqual(set(data), {'convert', 'repetitions'})
        self.assertEqual(data['convert'], {'formats': ['pdf'], 'use_cleaner': True, 'rules': ['spaces'],
                                           'known_rules': [key for key, _ in EDITORIAL_RULES]})
        self.assertEqual(data['repetitions']['window_size'], 50)
        context = self.client.get(self.url).context
        self.assertEqual(context['conversion_form'].initial, {'formats': ['pdf'], 'use_cleaner': True, 'rules': ['spaces']})
        self.assertEqual(context['repetitions_form'].initial['color_palette'], 'light')
        self.assertEqual(context['form'].initial, {})  # Odkurzacz: nic nie zapisano – wartości domyślne
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url).context['conversion_form'].initial, {})

    def test_rendered_page_shows_the_remembered_values(self):
        self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['pdf'],
                       'convert-rules_shown': 'True', 'convert-rules': ['spaces']})
        html = self.client.get(self.url).content.decode()
        self.assertIn('name="convert-formats" value="pdf" id="id_convert-formats_0" checked', html)
        self.assertNotIn('name="convert-formats" value="epub" id="id_convert-formats_1" checked', html)
        self.assertIn('name="convert-rules" value="spaces" id="id_convert-rules_1" checked', html)
        self.assertNotIn('name="convert-rules" value="trim" id="id_convert-rules_2" checked', html)

    def test_rule_added_after_saving_follows_its_default(self):
        """Reguła spoza listy znanych w chwili zapisu jest zaznaczona tak jak domyślnie; znane – jak zapisano."""
        default_rule = FORM_DEFAULT_RULE_LIST[0]
        known = [key for key, _ in EDITORIAL_RULES if key != default_rule]
        ProgramPreferences.objects.create(user=self.member, data={
            'clean': {'rules': ['spaces'], 'known_rules': known},
            'convert': {'rules': ['spaces'], 'known_rules': [key for key, _ in EDITORIAL_RULES]},
        })
        context = self.client.get(self.url).context
        self.assertEqual(context['form'].initial['rules'], ['spaces', default_rule] if default_rule != 'spaces'
                         else ['spaces'])
        self.assertEqual(context['conversion_form'].initial['rules'], ['spaces'])
        # Zapis bez listy znanych reguł (sprzed tej zmiany) – jak dotąd.
        ProgramPreferences.objects.filter(user=self.member).update(data={'clean': {'rules': ['trim']}})
        self.assertEqual(self.client.get(self.url).context['form'].initial['rules'], ['trim'])

    def test_settings_are_remembered_only_when_the_job_is_accepted(self):
        with patch.object(jobs, 'create', side_effect=jobs.JobError('Kolejka programów jest pełna.')):
            response = self.client.post(self.url, {'program_action': 'clean', 'document': upload(),
                                                   'rules': ['spaces']}, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(ProgramPreferences.objects.filter(user=self.member).exists())

    def test_saving_one_program_locks_the_record_and_keeps_the_others(self):
        """Dwie karty: zapis jednego programu blokuje rekord (select_for_update) i nie gubi drugiego."""
        from django.db.models import QuerySet
        ProgramPreferences.objects.create(user=self.member, data={'convert': {'formats': ['pdf']}})
        original = QuerySet.select_for_update
        with patch.object(QuerySet, 'select_for_update', autospec=True, side_effect=original) as locked:
            self.post_job({'program_action': 'clean', 'document': upload(), 'rules': ['spaces']})
        self.assertTrue(any(call.args[0].model is ProgramPreferences for call in locked.call_args_list))
        self.assertEqual(set(self.stored()), {'clean', 'convert'})
        self.assertEqual(self.stored()['convert'], {'formats': ['pdf']})

    def test_invalid_submission_is_not_remembered(self):
        response = self.client.post(self.url, {'program_action': 'convert', 'convert-document': upload(),
                                               'convert-formats': []}, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(ProgramPreferences.objects.filter(user=self.member).exists())

    def test_stored_values_are_validated_when_read(self):
        """Lista reguł może się zmienić: nieznane reguły, klucze i złe typy są pomijane."""
        ProgramPreferences.objects.create(user=self.member, data={
            'clean': {'rules': ['spaces', 'usunieta_regula', 7], 'track_changes': 'tak', 'unknown': True,
                      'document': 'plik.docx', 'cleaner_summary': True},
            'convert': {'formats': [], 'use_cleaner': True, 'rules': 'spaces'},
            'repetitions': {'window_size': 100000, 'min_word_length': True, 'ignored_words': 'kot',
                            'color_palette': 'neon', 'sentence_limit': 40},
            'nieznany': {'x': 1},
        })
        context = self.client.get(self.url).context
        self.assertEqual(context['form'].initial, {'rules': ['spaces'], 'cleaner_summary': True})
        self.assertEqual(context['conversion_form'].initial, {'use_cleaner': True})
        self.assertEqual(context['repetitions_form'].initial, {'ignored_words': 'kot', 'sentence_limit': 40})

    def test_damaged_preferences_fall_back_to_defaults(self):
        for data in ([], {'clean': 'tekst'}, {'clean': None}):
            with self.subTest(data=data):
                ProgramPreferences.objects.update_or_create(user=self.member, defaults={'data': data})
                response = self.client.get(self.url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context['form'].initial, {})
        self.post_job({'program_action': 'clean', 'document': upload()})
        self.assertIn('clean', self.stored())

    def test_restore_defaults_forgets_only_that_program(self):
        ProgramPreferences.objects.create(user=self.member, data={
            'clean': {'rules': ['spaces']}, 'convert': {'formats': ['pdf']}})
        response = self.client.post(self.url, {'program_action': 'reset_settings', 'program': 'clean'}, follow=True)
        self.assertRedirects(response, self.url + '#program-clean', fetch_redirect_response=False)
        self.assertContains(response, 'Odkurzacz: przywrócono ustawienia domyślne.')
        self.assertEqual(self.stored(), {'convert': {'formats': ['pdf']}})
        self.assertEqual(response.context['form'].initial, {})
        self.assertEqual(response.context['form']['rules'].value(), FORM_DEFAULT_RULE_LIST)
        # Bez zapisanych ustawień przycisk niczego nie psuje; nieznany program jest odrzucany.
        self.assertEqual(self.client.post(self.url, {'program_action': 'reset_settings', 'program': 'clean'}).status_code, 302)
        self.assertEqual(self.client.post(self.url, {'program_action': 'reset_settings', 'program': 'inny'}).status_code, 400)
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(self.url, {'program_action': 'reset_settings', 'program': 'convert'}).status_code, 302)
        self.assertEqual(self.stored(), {'convert': {'formats': ['pdf']}})

    def test_every_program_section_offers_restoring_defaults(self):
        html = self.client.get(self.url).content.decode()
        for program in ('clean', 'convert', 'repetitions'):
            self.assertIn(f'name="program" value="{program}"', html)
        self.assertEqual(html.count('data-program-job'), 3)  # formularz przywracania nie jest zadaniem


class CleanerOptionsTests(ProgramViewMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('opcje', 'Redaktor')
        cls.admin = create_superuser('opcje-admin')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def test_new_cleaner_options_are_off_by_default(self):
        options = self.post_job({'program_action': 'clean', 'document': upload(), 'rules': ['spaces']})
        self.assertEqual((options['track_changes'], options['cleaner_summary'], options['keep_page_format']),
                         (False, False, False))
        self.assertNotIn('analysis_unlimited', options)
        self.assertNotIn('timeout', options)
        form = OdkurzaczForm()
        self.assertEqual([form[name].value() for name in ('track_changes', 'cleaner_summary', 'keep_page_format')],
                         [False, False, False])

    def test_cleaner_options_reach_the_job(self):
        options = self.post_job({'program_action': 'clean', 'document': upload(), 'track_changes': 'on',
                                 'cleaner_summary': 'on', 'rebuild': 'on', 'keep_page_format': 'on'})
        self.assertEqual((options['track_changes'], options['cleaner_summary'], options['keep_page_format']),
                         (True, True, True))

    def test_page_format_is_kept_only_together_with_rebuild(self):
        options = self.post_job({'program_action': 'clean', 'document': upload(), 'keep_page_format': 'on'})
        self.assertIs(options['keep_page_format'], False)

    def test_conversion_offers_cleaner_rules_selected_like_the_cleaner(self):
        form = DocumentConversionForm(prefix='convert')
        self.assertEqual(form['rules'].value(), FORM_DEFAULT_RULE_LIST)
        self.assertEqual([key for key, _ in form.fields['rules'].choices], [key for key, _ in OdkurzaczForm().fields['rules'].choices])
        for name in ('track_changes', 'cleaner_summary', 'keep_page_format', 'rebuild'):
            self.assertNotIn(name, form.fields)
        html = self.client.get(self.url).content.decode()
        self.assertNotIn('Ustawienia w osobnej sekcji Odkurzacza nie mają tu wpływu', html)
        self.assertIn('name="convert-rules_shown" value="True"', html)

    def test_conversion_uses_only_the_selected_rules(self):
        options = self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['epub'],
                                 'convert-use_cleaner': 'on', 'convert-rules_shown': 'True',
                                 'convert-rules': ['spaces', 'quotes']})
        self.assertEqual(options['cleaner_rules'], ['spaces', 'quotes'])
        self.assertTrue(options['use_cleaner'])
        # Wszystkie reguły odznaczone: Odkurzacz działa bez reguł korekty (tylko miękkie entery).
        options = self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['epub'],
                                 'convert-use_cleaner': 'on', 'convert-rules_shown': 'True'})
        self.assertEqual(options['cleaner_rules'], [])
        # Bez znacznika listy (np. starszy formularz) – reguły domyślne jak w Odkurzaczu.
        options = self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['epub'],
                                 'convert-use_cleaner': 'on'})
        self.assertEqual(options['cleaner_rules'], FORM_DEFAULT_RULE_LIST)

    def test_conversion_without_cleaner_sends_no_rules_or_replacements(self):
        options = self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['epub'],
                                 'convert-rules_shown': 'True', 'convert-rules': ['spaces']})
        self.assertFalse(options['use_cleaner'])
        self.assertNotIn('cleaner_rules', options)
        self.assertNotIn('replacements', options)

    def test_unknown_conversion_rule_is_rejected(self):
        response = self.client.post(self.url, {'program_action': 'convert', 'convert-document': upload(),
                                               'convert-formats': ['epub'], 'convert-rules': ['nie_ma']},
                                    HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 400)
        self.assertIn('rules', response.json()['errors'])

    def test_superuser_limits_are_set_by_the_server_only(self):
        forged = {'analysis_unlimited': 'on', 'repetitions-analysis_unlimited': 'on', 'timeout': '9999',
                  'repetitions-timeout': '9999'}
        options = self.post_job({'program_action': 'repetitions', 'repetitions-document': upload(),
                                 **REPETITION_FIELDS, **forged})
        self.assertNotIn('analysis_unlimited', options)
        self.assertNotIn('timeout', options)
        self.assertNotIn('analysis_unlimited', options['repetitions'])
        self.client.force_login(self.admin)
        for data in ({'program_action': 'repetitions', 'repetitions-document': upload(), **REPETITION_FIELDS},
                     {'program_action': 'clean', 'document': upload()}):
            with self.subTest(action=data['program_action']):
                options = self.post_job(data)
                self.assertIs(options['analysis_unlimited'], True)
                self.assertEqual(options['timeout'], jobs.SUPERUSER_TIME_LIMIT)
        self.assertContains(self.client.get(self.url), 'nie masz limitu wielkości analizy')

    def test_superuser_time_limit_setting_is_bounded(self):
        for value, expected in ((30, TIME_LIMIT), (450, 450), (10_000, 600)):
            with self.subTest(value=value), self.settings(PROGRAM_SUPERUSER_TIME_LIMIT=value):
                self.assertEqual(jobs.superuser_time_limit(), expected)


class ChangeSummaryTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('podsumowanie', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def run_job(self, data, summary):
        def converted(upload, formats, **options):
            result = BytesIO(formatted_docx())
            result.conversion_warnings = ['Uwaga konwertera.']
            result.cleaner_summary = summary
            result.removed_images = 3
            self.options = options
            return result, 'docx', 'application/octet-stream'

        with patch.object(jobs, 'launch', side_effect=jobs.execute), \
                patch.object(jobs, 'convert_document', side_effect=converted):
            response = self.client.post(reverse('core:programs'), {'program_action': 'clean', 'document': upload(), **data})
        return self.client.get(response['Location'])

    def test_summary_is_listed_in_rule_order_next_to_warnings(self):
        page = self.run_job({'cleaner_summary': 'on'}, {'quotes': 2, 'spaces': 5, 'user_word_corrections': 1})
        state = page.context['state']
        self.assertEqual([(row['rule'], row['count']) for row in state['summary']],
                         [('spaces', 5), ('quotes', 2), ('user_word_corrections', 1)])
        self.assertEqual(state['summary'][0]['label'], 'Podwójne i wielokrotne spacje')
        self.assertEqual(state['summary'][2]['label'], 'Własne zamiany słownikowe (słownik zamian)')
        self.assertContains(page, 'Podsumowanie zmian')
        self.assertContains(page, '<td>Podwójne i wielokrotne spacje</td><td>5</td>', html=False)
        self.assertContains(page, 'Uwaga konwertera.')
        self.assertTrue(self.options['cleaner_summary'])

    def test_replacements_and_removed_images_reach_the_job_state(self):
        replacements = [{'source': 'pickupem', 'target': 'pikapem', 'count': 2},
                        {'source': 'Zapewnie', 'target': 'Zapewne', 'count': 1}]
        page = self.run_job({'cleaner_summary': 'on'}, {'user_word_corrections': 3, 'replacements': replacements})
        state = page.context['state']
        self.assertEqual(state['replacements'], replacements)
        self.assertEqual(state['removed_images'], 3)
        self.assertEqual([row['rule'] for row in state['summary']], ['user_word_corrections'])
        url = page.request['PATH_INFO']
        data = self.client.get(url).json()
        self.assertEqual((data['removed_images'], data['replacements']), (3, replacements))
        # Bez podsumowania liczba usuniętych obrazów i tak trafia do stanu.
        state = self.run_job({}, {}).context['state']
        self.assertEqual(state['removed_images'], 3)
        self.assertNotIn('replacements', state)

    def test_damaged_replacements_and_image_count_in_state_are_ignored(self):
        folder = jobs.root() / ('b' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'state.json', {'state': 'done', 'replacements': 'x', 'removed_images': 'dużo'})
        state = jobs.status(folder)
        self.assertEqual((state['replacements'], state['removed_images']), ([], 0))
        self.assertEqual(jobs.replacement_rows({'replacements': [{'source': 'a', 'target': 'b', 'count': 0},
                                                                  {'source': 'a', 'target': 'b', 'count': True},
                                                                  {'source': 'a', 'target': 'b', 'count': 2}]}),
                         [{'source': 'a', 'target': 'b', 'count': 2}])

    def test_no_changes_and_summary_off(self):
        page = self.run_job({'cleaner_summary': 'on'}, {})
        self.assertEqual(page.context['state']['summary'], [])
        self.assertContains(page, 'Reguły korekty nie wprowadziły zmian.')
        page = self.run_job({}, {'spaces': 3})
        self.assertNotIn('summary', page.context['state'])
        self.assertNotContains(page, 'Podsumowanie zmian')

    def test_damaged_summary_in_state_is_ignored(self):
        folder = jobs.root() / ('a' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'state.json', {'state': 'done', 'summary': 'zepsute'})
        self.assertEqual(jobs.status(folder)['summary'], [])

    def test_summary_rows_skip_empty_counts_and_keep_unknown_rules_last(self):
        self.assertIsNone(jobs.summary_rows(None))
        self.assertEqual(jobs.summary_rows({'nowa': 1, 'spaces': 0, 'trim': 2}),
                         [{'rule': 'trim', 'label': 'Spacje na początku i końcu akapitu', 'count': 2},
                          {'rule': 'nowa', 'label': 'nowa', 'count': 1}])


class PreparationTests(SimpleTestCase):
    """prepare_docx przekazuje nowe opcje do Odkurzacza i przebudowy (interfejsy 1 i 2)."""

    def run_prepare(self, **options):
        calls = []

        def step(name):
            def operation(source, **kwargs):
                calls.append((name, kwargs))
                return BytesIO(source.read())
            operation.__name__ = name
            return operation

        with patch.object(preparation, 'clean_docx', step('clean_docx')), \
                patch.object(preparation, 'rebuild_docx', step('rebuild_docx')), \
                patch.object(preparation, 'normalize_docx', step('normalize_docx')):
            with preparation.prepare_docx(BytesIO(b'docx'), **options) as result:
                self.assertEqual(result.read(), b'docx')
        return dict(calls)

    def test_defaults_are_not_passed_on(self):
        calls = self.run_prepare(rebuild=True, use_cleaner=True, cleaner_rules=['spaces'])
        self.assertEqual(calls['rebuild_docx'], {'allow_omissions': False})
        self.assertEqual(calls['clean_docx'], {'rules': ['spaces']})

    def test_new_options_reach_the_cleaner_and_rebuild(self):
        summary = {}
        replacements = [{'kind': 'word', 'source': 'kot', 'target': 'pies'}]
        calls = self.run_prepare(rebuild=True, use_cleaner=True, cleaner_rules=['spaces'], track_changes=True,
                                 replacements=replacements, summary=summary, keep_page_format=True)
        self.assertEqual(calls['rebuild_docx'], {'allow_omissions': False, 'keep_page_format': True})
        self.assertEqual(calls['clean_docx'], {'rules': ['spaces'], 'track_changes': True,
                                               'replacements': replacements, 'summary': summary})
        self.assertIs(calls['clean_docx']['summary'], summary)
        # Pusta lista to „żadnych zamian”, a nie lista domyślna.
        self.assertEqual(self.run_prepare(use_cleaner=True, replacements=[])['clean_docx']['replacements'], [])

    def test_without_cleaner_the_cleaner_options_do_nothing(self):
        self.assertNotIn('clean_docx', self.run_prepare(track_changes=True, summary={}))

    def test_invalid_replacements_and_summary_are_rejected(self):
        for replacements in ('kot', [{'kind': 'regex', 'source': 'a', 'target': 'b'}],
                             [{'kind': 'word', 'source': '', 'target': 'b'}],
                             [{'kind': 'word', 'source': 'a', 'target': 1}],
                             [{'kind': 'word', 'source': 'a', 'target': 'b', 'extra': 1}], ['a']):
            with self.subTest(replacements=replacements), self.assertRaises(ValueError):
                self.run_prepare(use_cleaner=True, replacements=replacements)
        with self.assertRaises(ValueError):
            self.run_prepare(use_cleaner=True, summary=[])


class ConverterConfigurationTests(SimpleTestCase):
    def test_new_keys_reach_the_worker_and_summary_comes_back(self):
        seen = {}

        def worker_run(directory, timeout):
            seen.update(json.loads((directory / 'job.json').read_text(encoding='utf-8')))
            (directory / 'summary.json').write_text(json.dumps({
                'spaces': 3, 'trim': 0, 'x': 'tekst', 'y': True,
                'replacements': [{'source': 'pickupem', 'target': 'pikapem', 'count': 2},
                                 {'source': 'zły', 'target': 3, 'count': 1}, 'tekst']}))
            (directory / 'images.json').write_text(json.dumps({'removed_images': 4}))

        replacements = [{'kind': 'word', 'source': 'kot', 'target': 'pies'}]
        with TemporaryDirectory() as temporary, self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary)), \
                patch('core.services.document_converter.run_converter', side_effect=worker_run):
            result, extension, _ = convert_document(BytesIO(formatted_docx()), [], include_docx=True, use_cleaner=True,
                                                    track_changes=True, cleaner_summary=True, replacements=replacements,
                                                    keep_page_format=True, analysis_unlimited=True)
            result.close()
        self.assertEqual(extension, 'docx')
        self.assertEqual({key: seen[key] for key in ('track_changes', 'cleaner_summary', 'replacements',
                                                       'keep_page_format', 'analysis_unlimited')},
                         {'track_changes': True, 'cleaner_summary': True, 'replacements': replacements,
                          'keep_page_format': True, 'analysis_unlimited': True})
        self.assertEqual(result.cleaner_summary, {
            'spaces': 3, 'replacements': [{'source': 'pickupem', 'target': 'pikapem', 'count': 2}]})
        self.assertEqual(result.removed_images, 4)

    def test_too_large_document_xml_is_refused_before_the_worker(self):
        from core.services.document_converter import validate_resources
        from core.services.document_errors import MESSAGES
        with self.settings(DOCUMENT_XML_LIMIT=2000), patch('core.services.document_converter.run_converter') as run:
            with self.assertRaisesMessage(ConversionError, MESSAGES['document_too_large']):
                convert_document(BytesIO(formatted_docx()), [], include_docx=True)
            run.assert_not_called()
        validate_resources(BytesIO(formatted_docx()))  # domyślny limit 15 MB

    def test_defaults_keep_the_previous_configuration(self):
        seen = {}

        def worker_run(directory, timeout):
            seen.update(json.loads((directory / 'job.json').read_text(encoding='utf-8')))
            seen['timeout'] = timeout

        with TemporaryDirectory() as temporary, self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary)), \
                patch('core.services.document_converter.run_converter', side_effect=worker_run):
            result, _, _ = convert_document(BytesIO(formatted_docx()), [], include_docx=True)
            result.close()
        self.assertEqual((seen['track_changes'], seen['cleaner_summary'], seen['replacements'],
                          seen['keep_page_format'], seen['analysis_unlimited']), (False, False, None, False, False))
        self.assertIsNone(result.cleaner_summary)
        self.assertEqual(result.removed_images, 0)
        self.assertLessEqual(seen['timeout'], TIME_LIMIT)

    def test_superuser_time_limit_is_used_and_reported(self):
        seen = {}

        def too_slow(directory, timeout):
            seen['timeout'] = timeout
            raise ConversionError('Konwersja przekroczyła limit 90 sekund. Podziel dokument lub wybierz mniej formatów.')

        with TemporaryDirectory() as temporary, self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary)), \
                patch('core.services.document_converter.run_converter', side_effect=too_slow), \
                self.assertRaisesMessage(ConversionError, 'limit 300 sekund'):
            convert_document(BytesIO(formatted_docx()), [], include_docx=True, timeout=300)
        self.assertGreater(seen['timeout'], TIME_LIMIT)
        with TemporaryDirectory() as temporary, self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary)), \
                patch('core.services.document_converter.run_converter', side_effect=too_slow):
            for timeout in (TIME_LIMIT, 30):
                with self.subTest(timeout=timeout), self.assertRaisesMessage(ConversionError, 'limit 90 sekund'):
                    convert_document(BytesIO(formatted_docx()), [], include_docx=True, timeout=timeout)
            with self.assertRaisesMessage(ConversionError, 'limit 600 sekund'):  # górna granica
                convert_document(BytesIO(formatted_docx()), [], include_docx=True, timeout=10_000)
        self.assertLessEqual(seen['timeout'], 600)


class WorkerConfigurationTests(SimpleTestCase):
    """main() podprocesu: odczyt nowych kluczy i zapis summary.json."""

    def run_main(self, config):
        calls = {}

        def prepare(document, **options):
            calls['prepare'] = options
            if options.get('summary') is not None:
                options['summary']['spaces'] = 4
            return BytesIO(b'prepared')

        def color(document, **options):
            calls['color'] = options
            return BytesIO(b'colored')

        with TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / 'source.docx').write_bytes(formatted_docx())
            (directory / 'job.json').write_text(json.dumps({'formats': [], 'include_docx': True, 'title': 'T', **config}))
            with patch.object(sys, 'argv', ['worker', str(directory)]), \
                    patch('core.services.document_preparation.prepare_docx', side_effect=prepare), \
                    patch('core.services.document_repetitions.color_document', side_effect=color):
                code = worker.main()
            summary = directory / 'summary.json'
            calls['summary'] = json.loads(summary.read_text()) if summary.exists() else None
        self.assertEqual(code, 0)
        return calls

    def test_new_options_are_read_and_summary_is_written(self):
        replacements = [{'kind': 'pattern', 'source': 'a+', 'target': 'a'}]
        calls = self.run_main({'prepare': True, 'clean': True, 'cleaner_rules': ['spaces'], 'track_changes': True,
                               'cleaner_summary': True, 'replacements': replacements, 'rebuild': True,
                               'keep_page_format': True})
        options = calls['prepare']
        self.assertEqual((options['track_changes'], options['replacements'], options['keep_page_format']),
                         (True, replacements, True))
        self.assertEqual(calls['summary'], {'spaces': 4})

    def test_replacement_list_is_written_to_the_summary(self):
        def prepare(document, **options):
            options['summary'].update(spaces=1, replacements=[
                {'source': f's{index}', 'target': 't', 'count': 1} for index in range(300)])
            return BytesIO(formatted_docx())

        with TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / 'source.docx').write_bytes(formatted_docx())
            (directory / 'job.json').write_text(json.dumps({'formats': [], 'include_docx': True, 'title': 'T',
                                                            'prepare': True, 'clean': True, 'cleaner_summary': True,
                                                            'cleaner_rules': ['user_word_corrections']}))
            with patch.object(sys, 'argv', ['worker', str(directory)]), \
                    patch('core.services.document_preparation.prepare_docx', side_effect=prepare):
                self.assertEqual(worker.main(), 0)
            summary = json.loads((directory / 'summary.json').read_text())
        self.assertEqual(summary['spaces'], 1)
        self.assertEqual(len(summary['replacements']), 200)
        self.assertEqual(summary['replacements'][0], {'source': 's0', 'target': 't', 'count': 1})

    def test_old_configuration_keeps_defaults_and_writes_no_summary(self):
        calls = self.run_main({'prepare': True, 'clean': True, 'cleaner_rules': ['spaces']})
        options = calls['prepare']
        self.assertEqual((options['track_changes'], options['replacements'], options['summary'], options['keep_page_format']),
                         (False, None, None, False))
        self.assertIsNone(calls['summary'])

    def test_unlimited_analysis_is_passed_only_when_set(self):
        settings = {'window_size': 35}
        self.assertNotIn('analysis_unlimited', self.run_main({'repetitions': settings})['color'])
        self.assertIs(self.run_main({'repetitions': settings, 'analysis_unlimited': True})['color']['analysis_unlimited'], True)


class AnalysisLimitTests(SimpleTestCase):
    def long_paragraph_document(self):
        doc = Document()
        doc.add_paragraph('słowo ' * 4000)  # 24 000 znaków w jednym akapicie
        return saved(doc)

    def test_limit_applies_unless_unlimited(self):
        with self.assertRaises(DocumentInputError) as caught:
            repetitions.color_document(self.long_paragraph_document())
        self.assertEqual(caught.exception.public_code, 'analysis_limit')

        class Passed(Exception):
            pass

        # Bez limitu analiza przechodzi dalej (zatrzymujemy ją na etapie odmian, aby test był szybki).
        with patch.object(repetitions, '_lemma_map', side_effect=Passed), self.assertRaises(Passed):
            repetitions.color_document(self.long_paragraph_document(), analysis_unlimited=True)


class ReplacementsInJobsTests(ProgramViewMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('zamiany-zadania', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def test_active_entries_reach_cleaner_and_conversion_jobs(self):
        ReplacementEntry.objects.all().delete()
        ReplacementEntry.objects.create(kind='word', source='kot', target='pies', position=20)
        ReplacementEntry.objects.create(kind='word', source='mysz', target='szczur', position=10)
        ReplacementEntry.objects.create(kind='word', source='ptak', target='ryba', is_active=False)
        # Wpisy z CMS są już sprawdzone przy zapisie („checked”) – Odkurzacz nie powtarza próby czasu.
        expected = [{'kind': 'word', 'source': 'mysz', 'target': 'szczur', 'checked': True},
                    {'kind': 'word', 'source': 'kot', 'target': 'pies', 'checked': True}]
        self.assertEqual(self.post_job({'program_action': 'clean', 'document': upload()})['replacements'], expected)
        options = self.post_job({'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['pdf'],
                                 'convert-use_cleaner': 'on'})
        self.assertEqual(options['replacements'], expected)
        self.assertNotIn('replacements', self.post_job({'program_action': 'repetitions',
                                                       'repetitions-document': upload(), **REPETITION_FIELDS}))
        ReplacementEntry.objects.all().delete()
        self.assertEqual(self.post_job({'program_action': 'clean', 'document': upload()})['replacements'], [])



class RealProgramOptionsTests(IsolatedProgramDirsMixin, TestCase):
    """Całe zadanie z prawdziwym Odkurzaczem i przebudową (bez procesu w tle)."""

    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('opcje-prawdziwe', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def run_clean(self, document, state='done', **data):
        with patch.object(jobs, 'launch', side_effect=jobs.execute):
            response = self.client.post(self.url_programs, {
                'document': SimpleUploadedFile('tekst.docx', saved(document).getvalue()), **data})
        page = self.client.get(response['Location'])
        self.assertEqual(page.context['state']['state'], state, page.context['state'])
        return page

    url_programs = reverse('core:programs')

    @tag('slow')
    def test_tracked_changes_and_summary(self):
        doc = Document()
        doc.add_paragraph('Ala  ma "kota".')
        page = self.run_clean(doc, rules=['spaces', 'quotes'], track_changes='on', cleaner_summary='on')
        counts = {row['rule']: row['count'] for row in page.context['state']['summary']}
        self.assertEqual(set(counts), {'spaces', 'quotes'})
        self.assertTrue(all(count > 0 for count in counts.values()))
        result = self.client.get(page.context['download_url'])
        body = Document(BytesIO(b''.join(result.streaming_content))).element.body
        namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        self.assertTrue(body.findall(f'.//{namespace}ins'))
        self.assertTrue(body.findall(f'.//{namespace}del'))

    @tag('slow')
    def test_rebuild_keeps_the_page_format_on_request(self):
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm
        doc = Document()
        section = doc.sections[0]
        section.orientation, section.page_width, section.page_height = WD_ORIENT.LANDSCAPE, Cm(21), Cm(14.8)
        section.left_margin = Cm(1)
        doc.add_paragraph('Tekst.')
        page = self.run_clean(doc, rebuild='on', keep_page_format='on', rules=['spaces'])
        result = self.client.get(page.context['download_url'])
        rebuilt = Document(BytesIO(b''.join(result.streaming_content))).sections[0]
        source = Document(saved(doc)).sections[0]  # wartości po zapisie (zaokrąglenie do twipów)
        self.assertEqual((rebuilt.orientation, rebuilt.page_width, rebuilt.page_height, rebuilt.left_margin),
                         (WD_ORIENT.LANDSCAPE, source.page_width, source.page_height, source.left_margin))
        # Bez opcji przebudowa jak dotąd ostrzega o zmianie formatu strony i czeka na potwierdzenie.
        page = self.run_clean(doc, state='confirmation', rebuild='on', rules=['spaces'])
        self.assertIn('format strony', page.context['state']['message'])
