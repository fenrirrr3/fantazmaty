"""Kolejka zadań programów: tworzenie i walidacja, dostęp, anulowanie, potwierdzenie, pobranie i sprzątanie."""
import json
import os
import shutil
import subprocess
import threading
import time
from io import BytesIO
from itertools import count
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import Client, RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse
from docx import Document

from core.odkurzacz_forms import PROGRAM_MAX_UPLOAD_BYTES
from core.program_test_support import run_program
from core.services import program_jobs as jobs
from core.services.document_converter import ConversionError, conversion_slot, run_converter
from fantazmaty.testing import create_superuser

from .support import REPETITION_FIELDS, IsolatedProgramDirsMixin, docx_with_images, downloaded, saved, sized_docx


def document(header=False):
    doc = Document()
    doc.add_paragraph('Ala  ma kota. Kot odwiedził ogród.')
    if header:
        doc.sections[0].header.paragraphs[0].text = 'Nagłówek'
    return SimpleUploadedFile('Zażółć.docx', saved(doc).getvalue())


class ProgramJobTests(IsolatedProgramDirsMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = create_superuser('job-owner')
        cls.other = create_superuser('job-other')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.owner)
        self.url = reverse('core:programs')

    def create(self, **extra):
        """Zakłada zadanie Odkurzacza bez uruchamiania procesu; zwraca adres stanu i folder zadania."""
        with patch.object(jobs, 'launch'):
            response = self.client.post(self.url, {'program_action': 'clean', 'document': document(), **extra},
                                        HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 202, response.content)
        return response.json()['url'], next(jobs.root().iterdir())

    def test_opening_the_programs_page_removes_expired_results(self):
        """Regresja: niepobrane wyniki czekały na serwerze do następnego zgłoszenia dokumentu."""
        expired = jobs.root() / ('e' * 32)
        expired.mkdir(parents=True)
        jobs.write_json(expired / 'state.json', {'state': 'done', 'heartbeat': time.time()})
        old = time.time() - jobs.TTL - 60
        os.utime(expired, (old, old))
        self.assertEqual(self.client.get(self.url).status_code, 200)
        self.assertFalse(expired.exists())

    def test_all_three_forms_create_jobs_and_validate_before_launch(self):
        forms = [
            {'program_action': 'clean', 'document': document()},
            {'program_action': 'convert', 'convert-document': document(), 'convert-formats': ['pdf', 'epub']},
            {'program_action': 'repetitions', 'repetitions-document': document(), **REPETITION_FIELDS},
        ]
        for data in forms:
            with patch.object(jobs, 'launch') as launch:
                response = self.client.post(self.url, data, HTTP_X_PROGRAM_JOB='1')
                self.assertEqual(response.status_code, 202, response.content)
                launch.assert_called_once()
            folder = max(jobs.root().iterdir(), key=lambda p: p.stat().st_mtime_ns)
            jobs.write_json(folder / 'state.json', {'state': 'cancelled'})
        with patch.object(jobs, 'launch') as launch:
            bad = self.client.post(self.url, {'program_action': 'convert'}, HTTP_X_PROGRAM_JOB='1')
            self.assertEqual(bad.status_code, 400)
            launch.assert_not_called()

    def test_oversized_http_upload_never_starts_conversion(self):
        for action, field, extra in (
            ('clean', 'document', {}),
            ('convert', 'convert-document', {'convert-formats': ['pdf']}),
            ('repetitions', 'repetitions-document', REPETITION_FIELDS),
        ):
            with self.subTest(action=action), patch('core.services.program_jobs.convert_document') as converter:
                response = run_program(
                    self.client, {'program_action': action, field: sized_docx(PROGRAM_MAX_UPLOAD_BYTES + 1), **extra})
                self.assertContains(response, 'Maksymalny rozmiar to 5 MB')
                converter.assert_not_called()

    def test_five_megabyte_upload_is_accepted_by_every_program(self):
        """Limit pliku wynosi 5 MB we wszystkich programach (wcześniej 2 MB)."""
        self.assertEqual(PROGRAM_MAX_UPLOAD_BYTES, 5 * 1024 * 1024)
        self.assertContains(self.client.get(self.url), 'Maksymalnie 5 MB', count=3)
        for action, field, extra in (
            ('clean', 'document', {}),
            ('convert', 'convert-document', {'convert-formats': ['pdf']}),
            ('repetitions', 'repetitions-document', REPETITION_FIELDS),
        ):
            with self.subTest(action=action), patch.object(jobs, 'launch') as launch:
                response = self.client.post(self.url, {'program_action': action, field: sized_docx(PROGRAM_MAX_UPLOAD_BYTES), **extra},
                                            HTTP_X_PROGRAM_JOB='1')
                self.assertEqual(response.status_code, 202, response.content)
                launch.assert_called_once()
            for folder in jobs.root().iterdir():
                jobs.write_json(folder / 'state.json', {'state': 'cancelled'})

    def test_owner_session_expiration_and_csrf(self):
        url, folder = self.create()
        other = Client()
        other.force_login(self.other)
        same = Client()
        same.force_login(self.owner)
        for client in (other, same):
            self.assertEqual(client.get(url).status_code, 404)
            self.assertEqual(client.post(url, {'action': 'cancel'}).status_code, 404)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.owner)
        self.assertEqual(csrf_client.post(url, {'action': 'cancel'}).status_code, 403)
        with patch('django.core.signing.time.time', return_value=time.time() + jobs.TTL + 2):
            self.assertEqual(self.client.get(url).status_code, 404)
        self.assertFalse((folder / 'cancel').exists())

    def test_real_cleaner_result_download_and_duplicate_rejected(self):
        url, folder = self.create(rules=['spaces'])
        with patch.object(jobs, 'launch') as launch:
            duplicate = self.client.post(self.url, {'program_action': 'clean', 'document': document()},
                                         HTTP_X_PROGRAM_JOB='1')
            self.assertEqual(duplicate.status_code, 400)
            launch.assert_not_called()
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'done')
        result = downloaded(self.client.get(url + '?download=1'))
        self.assertEqual(Document(BytesIO(result)).paragraphs[0].text, 'Ala ma kota. Kot odwiedził ogród.')
        self.assertFalse((folder / 'source.docx').exists())

    def test_real_job_removes_images_and_reports_how_many(self):
        with patch.object(jobs, 'launch'):
            response = self.client.post(self.url, {'program_action': 'clean', 'rules': ['spaces'], 'document':
                                                   SimpleUploadedFile('obrazy.docx', docx_with_images().getvalue())},
                                        HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 202, response.content)
        url, folder = response.json()['url'], next(jobs.root().iterdir())
        jobs.execute(folder)
        state = self.client.get(url).json()
        self.assertEqual((state['state'], state['removed_images']), ('done', 6), state)
        with ZipFile(BytesIO(downloaded(self.client.get(url + '?download=1')))) as archive:
            self.assertFalse([name for name in archive.namelist() if 'media/' in name])

    def test_result_warnings_are_shown_and_job_is_removed_after_download(self):
        url, folder = self.create()
        (folder / 'result').write_bytes(b'wynik')
        jobs.write_json(folder / 'state.json', {
            'state': 'done', 'filename': 'Plik_odkurzony.docx', 'mime': 'application/octet-stream',
            'warnings': ['Uwaga konwertera: test'],
        })
        self.assertContains(self.client.get(url + '?view=page'), 'Uwaga konwertera: test')
        self.assertEqual(self.client.get(url).json()['warnings'], ['Uwaga konwertera: test'])
        # Klient testowy zamyka odpowiedź po odczycie; ręczne close() zamknęłoby połączenie MySQL.
        self.assertEqual(downloaded(self.client.get(url + '?download=1')), b'wynik')
        self.assertFalse(folder.exists())
        again = self.client.get(url + '?download=1')
        self.assertContains(again, 'Plik został już pobrany i usunięty z serwera', status_code=404)

    def test_streaming_download_closes_file_without_closing_test_connection(self):
        """Regresja: podwójne close() odpowiedzi psuło transakcję TestCase na MySQL."""
        response = run_program(self.client, {'program_action': 'clean', 'document': document()})
        self.assertTrue(response.streaming)
        with patch.object(connection, 'close_if_unusable_or_obsolete') as close_db:
            content = downloaded(response)
        close_db.assert_not_called()
        self.assertTrue(response.closed)
        # Bez zaznaczonych reguł tekst wraca bez zmian.
        self.assertEqual(Document(BytesIO(content)).paragraphs[0].text, 'Ala  ma kota. Kot odwiedził ogród.')

    def test_old_state_with_a_warning_count_still_renders(self):
        """Regresja: stan zapisany przez starszą wersję (liczba uwag zamiast listy) psuł stronę."""
        folder = jobs.root() / ('a' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'state.json', {'state': 'done', 'warnings': 2})
        self.assertEqual(jobs.status(folder)['warnings'], [])
        json.dumps(jobs.status(folder))

    def test_cancel_before_start_never_runs_converter(self):
        url, folder = self.create()
        self.assertEqual(self.client.post(url, {'action': 'cancel'}).json()['state'], 'cancelling')
        with patch.object(jobs, 'convert_document') as converter:
            jobs.execute(folder)
            converter.assert_not_called()
        self.assertEqual(self.client.get(url).json()['state'], 'cancelled')
        self.assertFalse((folder / 'source.docx').exists())
        self.assertEqual(self.client.get(url + '?download=1').status_code, 404)

    def test_real_rebuild_confirmation_reuses_original_without_upload(self):
        url, folder = self.create(document=document(header=True), rebuild='on')
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'confirmation')
        self.assertTrue((folder / 'source.docx').exists())
        with patch.object(jobs, 'launch') as launch:
            self.assertEqual(self.client.post(url, {'action': 'confirm'}).status_code, 202)
            self.assertEqual(self.client.post(url, {'action': 'confirm'}).status_code, 404)
            launch.assert_called_once()
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'done')

    def test_confirmation_renews_access_for_processing_and_download(self):
        url, folder = self.create()
        jobs.write_json(folder / 'state.json', {'state': 'confirmation', 'message': 'Potwierdź'})
        now = time.time()
        with (
            patch('django.core.signing.time.time', return_value=now + jobs.TTL - 60),
            patch.object(jobs, 'launch'),
        ):
            response = self.client.post(url, {'action': 'confirm'})
            self.assertEqual(response.status_code, 202)
            renewed = response.json()['url']
            self.assertNotEqual(url, renewed)
        with patch('django.core.signing.time.time', return_value=now + jobs.TTL + 10):
            self.assertEqual(self.client.get(url).status_code, 404)
            self.assertEqual(self.client.get(renewed).status_code, 200)

    def test_without_javascript_the_form_redirects_to_an_auto_refreshing_page(self):
        with patch.object(jobs, 'launch') as launch:
            response = self.client.post(self.url, {'program_action': 'clean', 'document': document()})
        launch.assert_called_once()
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response['Location'].endswith('?view=page'))
        page = self.client.get(response['Location'])
        self.assertContains(page, 'http-equiv="refresh"')
        self.assertContains(page, 'value="cancel"')
        cancelled = self.client.post(response['Location'].split('?')[0], {'action': 'cancel', 'view': 'page'})
        self.assertEqual(cancelled.status_code, 302)
        folder = next(jobs.root().iterdir())
        self.assertTrue((folder / 'cancel').exists())

    def test_full_queue_rejects_new_jobs_until_one_finishes(self):
        other = Client()
        other.force_login(self.other)

        def post_as_other():
            return other.post(self.url, {'program_action': 'clean', 'document': document()},
                              HTTP_X_PROGRAM_JOB='1')

        with self.settings(PROGRAM_MAX_QUEUED_JOBS=1):
            self.create()
            with patch.object(jobs, 'launch') as launch:
                response = post_as_other()
            self.assertEqual(response.status_code, 400)
            self.assertIn('Kolejka programów jest pełna', response.json()['message'])
            launch.assert_not_called()
            folder = next(jobs.root().iterdir())
            jobs.write_json(folder / 'state.json', {'state': 'done', 'filename': 'a.docx', 'mime': 'x'})
            with patch.object(jobs, 'launch'):
                self.assertEqual(post_as_other().status_code, 202)


    def test_unknown_program_and_server_problems_are_explained(self):
        response = self.client.post(self.url, {'program_action': 'inne'})
        self.assertContains(response, 'Wybierz narzędzie', status_code=400)
        response = self.client.post(self.url, {'program_action': 'inne'}, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual((response.status_code, response.json()), (400, {'message': 'Wybierz program.'}))
        for error, status, text in ((jobs.JobError('Kolejka programów jest pełna.'), 400, 'Kolejka programów'),
                                    (OSError('dysk'), 503, 'Nie można zapisać zadania')):
            with self.subTest(error=type(error).__name__), patch.object(jobs, 'create', side_effect=error):
                page = self.client.post(self.url, {'program_action': 'clean', 'document': document()})
                self.assertContains(page, text, status_code=status)
                data = self.client.post(self.url, {'program_action': 'clean', 'document': document()},
                                        HTTP_X_PROGRAM_JOB='1')
                self.assertEqual(data.status_code, status)
                self.assertIn(text, data.json()['message'])

    def test_job_page_rejects_unknown_action_and_unreadable_state(self):
        url, folder = self.create()
        response = self.client.post(url, {'action': 'inne'})
        self.assertEqual((response.status_code, response.json()), (404, {'message': 'Nieprawidłowa akcja.'}))
        (folder / 'state.json').write_text('{uszkodzony', encoding='utf-8')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()['message'], 'Nie można odczytać zadania. Spróbuj ponownie.')
        self.assertContains(self.client.get(url + '?view=page'), 'Nie można odczytać zadania', status_code=503)


class JobFolderTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def test_cleanup_removes_expired_jobs_and_converter_folders_but_keeps_live_ones(self):
        jobs.root().mkdir(parents=True)
        foreign = jobs.root() / 'not-a-job'
        expired = jobs.root() / ('b' * 32)
        running = jobs.root() / ('c' * 32)
        fresh = jobs.root() / ('d' * 32)
        finished = jobs.root() / ('a' * 32)
        stale_work = jobs.root().parent / 'document-stale'
        fresh_work = jobs.root().parent / 'document-fresh'
        for folder in (foreign, expired, running, fresh, finished, stale_work, fresh_work):
            folder.mkdir()
        # Zakończone dawno zadanie (stan bez znaku życia) jest usuwane.
        jobs.write_json(finished / 'state.json', {'state': 'done', 'heartbeat': time.time()})
        # Zadanie sprzed TTL, które wciąż daje znak życia, nie może zniknąć.
        jobs.write_json(running / 'state.json', {'state': 'running', 'heartbeat': time.time()})
        old = time.time() - jobs.TTL - 60
        for folder in (foreign, expired, running, finished, stale_work):
            os.utime(folder, (old, old))
        jobs.cleanup()
        # Folder o obcej nazwie nie jest zadaniem i nie wolno go usuwać.
        self.assertTrue(foreign.exists())
        self.assertFalse(expired.exists())
        self.assertFalse(finished.exists())
        self.assertFalse(stale_work.exists())
        self.assertTrue(running.exists())
        self.assertTrue(fresh.exists())
        self.assertTrue(fresh_work.exists())

    def test_active_jobs_ignore_stray_files(self):
        jobs.root().mkdir(parents=True)
        (jobs.root() / 'stray.tmp').write_text('x')
        job = jobs.root() / ('e' * 32)
        job.mkdir()
        jobs.write_json(job / 'request.json', {'user': 7})
        jobs.write_json(job / 'state.json', {'state': 'queued', 'heartbeat': time.time()})
        self.assertEqual([folder for _, folder in jobs.active_jobs()], [job])

    def test_worker_start_is_reported_as_starting_the_converter(self):
        """Etap „start” zapisany przez proces roboczy jest pokazywany jako uruchamianie konwertera."""
        folder = jobs.root() / ('f' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'request.json', {'name': 'a.docx', 'action': 'clean', 'options': {}})
        (folder / 'source.docx').write_bytes(b'docx')
        stages = []

        def converter(source, control, **options):
            directory = self.work_dir / 'worker'
            directory.mkdir()
            (directory / 'progress.json').write_text(json.dumps({'stage': 'start'}))
            control(directory)
            stages.append(jobs.read_json(folder / 'state.json')['stage'])
            return BytesIO(b'wynik'), 'docx', 'application/octet-stream'

        with patch.object(jobs, 'convert_document', side_effect=converter):
            jobs.execute(folder)
        self.assertEqual(stages, ['Uruchamianie konwertera'])
        self.assertEqual(jobs.status(folder)['state'], 'done')


class CancellationTests(SimpleTestCase):
    def test_cancel_kills_and_reaps_the_actual_converter_process(self):
        with TemporaryDirectory() as tmp:
            directory = Path(tmp)
            directory.joinpath('job.json').write_text('{"formats": ["pdf"]}')
            actual = subprocess.Popen
            children = []

            def spawn(*args, **kwargs):
                # Prawdziwy, długo działający proces potomny czyni anulowanie deterministycznym.
                process = actual([args[0][0], '-c', 'import time; time.sleep(30)'], **kwargs)
                children.append(process)
                return process

            calls = []

            def control(_):
                calls.append(1)
                if len(calls) >= 2:
                    raise jobs.JobCancelled()

            with patch('core.services.document_converter.subprocess.Popen', side_effect=spawn):
                started = time.monotonic()
                with self.assertRaises(jobs.JobCancelled):
                    run_converter(directory, 10, control=control)
            self.assertLess(time.monotonic() - started, 5)
            self.assertIsNotNone(children[0].poll())


class ConversionQueueTests(IsolatedProgramDirsMixin, SimpleTestCase):
    def hold_slot(self, release):
        def hold():
            with conversion_slot():
                release.wait(5)
        holder = threading.Thread(target=hold)
        holder.start()
        time.sleep(0.2)
        return holder

    def test_slot_waits_for_the_other_document(self):
        released = threading.Event()
        holder = self.hold_slot(released)
        with self.assertRaises(ConversionError):
            with conversion_slot():
                pass
        calls = []

        def on_wait():
            calls.append(1)
            released.set()

        with conversion_slot(wait=10, on_wait=on_wait):
            self.assertTrue(calls)
        holder.join()

    def test_waiting_can_be_cancelled(self):
        stop = threading.Event()
        holder = self.hold_slot(stop)

        def cancel():
            raise jobs.JobCancelled()

        with self.assertRaises(jobs.JobCancelled):
            with conversion_slot(wait=10, on_wait=cancel):
                pass
        stop.set()
        holder.join()

    def test_queue_wait_covers_the_time_limits_of_jobs_ahead(self):
        """Zadanie superusera (limit MAX_TIME_LIMIT) przed nami nie może wyczerpać naszego czasu oczekiwania."""
        from core.services.document_converter import MAX_TIME_LIMIT, TIME_LIMIT
        mine = self.job_folder('a' * 32, {})
        self.assertEqual(jobs.queue_wait(mine), jobs.QUEUE_WAIT)
        self.job_folder('b' * 32, {'timeout': MAX_TIME_LIMIT})
        self.job_folder('c' * 32, {})
        self.assertEqual(jobs.queue_wait(mine), MAX_TIME_LIMIT + TIME_LIMIT + 2 * jobs.QUEUE_MARGIN)
        self.assertGreater(jobs.queue_wait(mine), MAX_TIME_LIMIT)
        self.assertEqual(jobs.job_time_limit({'timeout': 10 ** 6}), MAX_TIME_LIMIT)
        self.assertEqual(jobs.job_time_limit({'timeout': 'x'}), TIME_LIMIT)

    def job_folder(self, key, options):
        folder = jobs.root() / key
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'request.json', {'user': 1, 'options': options})
        jobs.write_json(folder / 'state.json', {'state': 'queued', 'started': time.time()})
        return folder

    def test_queued_job_with_recent_heartbeat_is_alive(self):
        now = time.time()
        self.assertTrue(jobs.alive({'state': 'queued', 'started': now - 500, 'heartbeat': now}))
        self.assertFalse(jobs.alive({'state': 'running', 'started': now - 500}))


class JobStorageTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Pliki stanu: zapis atomowy z ponowieniem (Windows blokuje otwarte pliki) i odporność na śmieci w katalogu."""

    def setUp(self):
        super().setUp()
        jobs.root().mkdir(parents=True)
        sleep = patch('core.services.program_jobs.time.sleep')
        sleep.start()
        self.addCleanup(sleep.stop)

    def test_locked_state_file_is_retried_and_finally_reported(self):
        path = jobs.root() / 'state.json'
        replace, attempts = os.replace, []

        def locked_once(source, target):
            attempts.append(target)
            if len(attempts) == 1:
                raise PermissionError
            replace(source, target)

        with patch('core.services.program_jobs.os.replace', side_effect=locked_once):
            jobs.write_json(path, {'state': 'queued'})
        self.assertEqual(len(attempts), 2)
        with patch.object(Path, 'read_text', side_effect=[PermissionError, '{"state": "done"}']):
            self.assertEqual(jobs.read_json(path), {'state': 'done'})
        with patch('core.services.program_jobs.os.replace', side_effect=PermissionError), \
                self.assertRaises(PermissionError):
            jobs.write_json(path, {'state': 'error'})
        with patch.object(Path, 'read_text', side_effect=PermissionError), self.assertRaises(PermissionError):
            jobs.read_json(path)
        self.assertEqual(jobs.read_json(path), {'state': 'queued'})
        self.assertEqual([item.name for item in jobs.root().iterdir()], ['state.json'])

    def test_cleanup_handles_links_and_damaged_jobs(self):
        outside = self.work_dir / 'outside'
        outside.mkdir()
        (jobs.root() / ('1' * 32)).symlink_to(outside)
        damaged = jobs.root() / ('2' * 32)
        damaged.mkdir()
        (damaged / 'state.json').write_text('{uszkodzony', encoding='utf-8')
        old = time.time() - jobs.TTL - 60
        for folder in (outside, damaged):
            os.utime(folder, (old, old))
        jobs.cleanup()
        self.assertTrue(outside.exists())
        self.assertFalse(damaged.exists())

    def test_queue_ignores_missing_root_and_unreadable_jobs(self):
        shutil.rmtree(jobs.root())
        self.assertEqual(jobs.active_jobs(), [])
        jobs.root().mkdir()
        (jobs.root() / ('3' * 32)).mkdir()
        self.assertEqual(jobs.active_jobs(), [])

    def test_second_creation_waits_for_the_lock_and_gives_up(self):
        with jobs.creation_lock(), self.assertRaisesMessage(jobs.JobError, 'Serwer przyjmuje teraz inne zadanie'):
            with jobs.creation_lock(wait=0):
                pass
        with jobs.creation_lock(wait=0):
            pass


class JobLifecycleTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Uruchomienie procesu, potwierdzenie przebudowy i przerwanie zadania na każdym etapie."""

    def setUp(self):
        super().setUp()
        self.folder = jobs.root() / ('a' * 32)
        self.folder.mkdir(parents=True)
        jobs.write_json(self.folder / 'request.json', {'user': 1, 'session': 's', 'name': 'Tekst.docx',
                                                       'action': 'clean', 'options': {}})
        (self.folder / 'source.docx').write_bytes(b'docx')

    def state(self):
        return jobs.status(self.folder)

    def test_launch_starts_the_worker_or_records_why_it_could_not(self):
        with patch('core.services.program_jobs.subprocess.Popen') as popen, \
                patch('core.services.program_jobs.threading.Thread') as reaper:
            jobs.launch(self.folder)
        command = popen.call_args.args[0]
        self.assertTrue(command[1].endswith('program_job_worker.py'))
        self.assertEqual(command[2], str(self.folder.resolve()))
        self.assertIn('DJANGO_SETTINGS_MODULE', popen.call_args.kwargs['env'])
        reaper.return_value.start.assert_called_once()
        with patch.object(jobs, 'converter_python', side_effect=ConversionError('brak')), \
                self.assertRaisesMessage(jobs.JobError, 'Nie można uruchomić programu'):
            jobs.launch(self.folder)
        self.assertEqual(self.state()['state'], 'error')

    def test_status_of_stale_or_cancelled_confirmation(self):
        jobs.write_json(self.folder / 'state.json', {'state': 'running', 'heartbeat': time.time() - 1000})
        self.assertIn('Proces nie odpowiedział', self.state()['message'])
        jobs.write_json(self.folder / 'state.json', {'state': 'confirmation', 'message': 'Potwierdź'})
        (self.folder / 'cancel').touch()
        self.assertEqual(self.state()['state'], 'cancelled')

    def test_confirmation_is_accepted_once_and_undone_when_it_cannot_start(self):
        jobs.write_json(self.folder / 'state.json', {'state': 'confirmation', 'message': 'Potwierdź'})
        (self.folder / 'accepted').touch()
        with self.assertRaisesMessage(jobs.JobError, 'Potwierdzenie zostało już wysłane'):
            jobs.confirm(self.folder)
        (self.folder / 'accepted').unlink()
        for patcher in (patch.object(jobs, 'require_capacity', side_effect=jobs.JobError('Kolejka pełna')),
                        patch.object(jobs, 'launch', side_effect=jobs.JobError('Nie można uruchomić'))):
            with self.subTest(patcher=patcher.attribute), patcher, self.assertRaises(jobs.JobError):
                jobs.confirm(self.folder)
            self.assertFalse((self.folder / 'accepted').exists())

    def test_cancel_while_waiting_running_or_after_conversion(self):
        def waiting(source, control, on_wait, **options):
            on_wait()
            self.assertEqual(jobs.read_json(self.folder / 'state.json')['state'], 'queued')
            (self.folder / 'cancel').touch()
            on_wait()

        def running(source, control, on_wait, **options):
            (self.folder / 'cancel').touch()
            control(self.work_dir)

        def finished(source, control, on_wait, **options):
            (self.folder / 'cancel').touch()
            return BytesIO(b'wynik'), 'docx', 'application/octet-stream'

        for converter in (waiting, running, finished):
            with self.subTest(stage=converter.__name__):
                (self.folder / 'source.docx').write_bytes(b'docx')
                clock = count(time.time(), 10)  # każde sprawdzenie „znaku życia” jest 10 s później
                with patch.object(jobs, 'convert_document', side_effect=converter), \
                        patch.object(jobs.time, 'time', side_effect=lambda: next(clock)):
                    jobs.execute(self.folder)
                self.assertEqual(self.state()['state'], 'cancelled')
                self.assertFalse((self.folder / 'result').exists())
                self.assertFalse((self.folder / 'source.docx').exists())
                (self.folder / 'cancel').unlink()

    def test_errors_show_only_converter_messages(self):
        for error, message in ((ConversionError('Konwerter zajęty.'), 'Konwerter zajęty.'),
                               (RuntimeError('/sciezka/tajna'), 'Nie udało się przetworzyć dokumentu.')):
            with self.subTest(error=type(error).__name__), \
                    patch.object(jobs, 'convert_document', side_effect=error), \
                    self.assertLogs('core.services.program_jobs', 'ERROR'):
                jobs.execute(self.folder)
            self.assertTrue(self.state()['message'].startswith(message))
            self.assertNotIn('tajna', self.state()['message'])


class JobAccessTests(IsolatedProgramDirsMixin, SimpleTestCase):
    """Token zadania działa tylko dla właściciela w tej samej sesji i tylko dla prawdziwego folderu zadania."""

    def request(self, user=1, session='nonce'):
        request = RequestFactory().post('/')
        request.user = SimpleNamespace(pk=user)
        request.session = {jobs.SESSION_KEY: session}
        return request

    def test_forged_or_foreign_tokens_are_rejected(self):
        folder = jobs.root() / ('b' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'request.json', {'user': 2, 'session': 'nonce'})
        (jobs.root() / ('c' * 32)).symlink_to(folder)
        for key in ('../../etc', 'c' * 32, 'b' * 32):
            token = signing.dumps({'key': key, 'user': 1, 'session': 'nonce'}, salt=jobs.SALT)
            with self.subTest(key=key), self.assertRaisesMessage(jobs.JobError, 'Zadanie wygasło'):
                jobs.resolve(self.request(), token)

    def test_failed_start_does_not_leave_the_document_on_the_server(self):
        """Gdy procesu nie da się uruchomić, nikt nie dostaje tokenu – kopia dokumentu nie może zostać."""
        with patch.object(jobs, 'converter_python', side_effect=ConversionError('brak')), \
                self.assertRaisesMessage(jobs.JobError, 'Nie można uruchomić programu'):
            jobs.create(self.request(), SimpleUploadedFile('tekst.docx', b'docx'), {}, 'clean')
        self.assertEqual(list(jobs.root().iterdir()), [])

    def test_oversized_upload_is_refused_without_creating_a_job(self):
        upload = SimpleUploadedFile('duzy.docx', b'0' * (PROGRAM_MAX_UPLOAD_BYTES + 1))
        with self.assertRaisesMessage(jobs.JobError, 'limit 5 MB'):
            jobs.create(self.request(), upload, {}, 'clean')
        self.assertEqual(list(jobs.root().iterdir()), [])
