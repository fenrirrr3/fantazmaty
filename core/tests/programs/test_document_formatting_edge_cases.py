"""Ujednolicanie formatowania: nieznane python-docx wyrównanie ze stylu, poziom konspektu wpisany w akapit,
styl akapitu o błędnym typie i poprzednie formatowanie w zmianach formatowania (w:rPrChange, w:pPrChange)."""
from django.test import SimpleTestCase
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from core.services.document_formatting import normalize_docx

from .support import saved


def indent(paragraph):
    node = paragraph._p.find(f'{qn("w:pPr")}/{qn("w:ind")}')
    return None if node is None else node.get(qn('w:firstLine'))


def alignment(paragraph):
    node = paragraph._p.find(f'{qn("w:pPr")}/{qn("w:jc")}')
    return None if node is None else node.get(qn('w:val'))


class FormattingGapTests(SimpleTestCase):
    def test_style_alignment_unknown_to_python_docx_is_inherited(self):
        """Styl z w:jc „end” (do prawej) i „start” (do lewej): tylko ten drugi akapit jest justowany."""
        document = Document()
        for value in ('end', 'start'):
            style = document.styles.add_style(f'Styl {value}', WD_STYLE_TYPE.PARAGRAPH)
            style.element.get_or_add_pPr().append(OxmlElement('w:jc', {qn('w:val'): value}))
            document.add_paragraph(value, style=style)
        right, left = Document(normalize_docx(saved(document), justify=True)).paragraphs
        self.assertEqual((alignment(right), indent(right)), (None, '709'))
        self.assertEqual((alignment(left), indent(left)), ('both', '709'))

    def test_outline_level_in_the_paragraph_keeps_the_heading_indent(self):
        document = Document()
        for level in ('0', '9'):
            paragraph = document.add_paragraph(f'Poziom {level}')
            paragraph._p.get_or_add_pPr().append(OxmlElement('w:outlineLvl', {qn('w:val'): level}))
        heading, plain = Document(normalize_docx(saved(document))).paragraphs
        self.assertIsNone(indent(heading))  # nagłówek zachowuje swoje wcięcie
        self.assertEqual(indent(plain), '709')  # poziom 9 to zwykły tekst

    def test_style_of_unknown_type_is_not_a_heading(self):
        """Styl z błędnym w:type (uszkodzony DOCX): zwykły akapit – z własnym wyrównaniem albo bez niego
        (wcześniej akapit bez własnego wyrównania przerywał ujednolicanie błędem)."""
        cases = {'own alignment': (WD_ALIGN_PARAGRAPH.RIGHT, 'right'), 'no own alignment': (None, 'both')}
        for case, (own, expected) in cases.items():
            document = Document()
            style = document.styles.add_style('Uszkodzony', WD_STYLE_TYPE.PARAGRAPH)
            style.element.set(qn('w:type'), 'nieznany')
            paragraph = document.add_paragraph('Tekst')
            paragraph.alignment = own
            paragraph._p.get_or_add_pPr().insert(0, OxmlElement('w:pStyle', {qn('w:val'): 'Uszkodzony'}))
            with self.subTest(case):
                result = Document(normalize_docx(saved(document), justify=True)).paragraphs[0]
                self.assertEqual((alignment(result), indent(result)), (expected, '709'))

    def test_theme_fonts_are_replaced_by_times_new_roman(self):
        document = Document()
        run = document.add_paragraph().add_run('Motyw')
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:asciiTheme'), 'minorHAnsi')
        fonts = Document(normalize_docx(saved(document))).paragraphs[0].runs[0]._r.rPr.rFonts
        self.assertIsNone(fonts.get(qn('w:asciiTheme')))
        self.assertEqual(fonts.get(qn('w:ascii')), 'Times New Roman')

    def test_previous_formatting_of_tracked_changes_is_normalized_too(self):
        """Poprzednie formatowanie zmiany dostaje to samo ujednolicenie (czcionka, odstępy, wcięcie), więc zmiana
        pokazuje tylko pracę redaktora. Regresja: w:rFonts dopisywany za w:rPrChange (python-docx nie zna jej
        miejsca w schemacie) – Word uznaje taki plik za uszkodzony; pusty w:pPrChange z LibreOffice dostaje w:pPr,
        a jego miejsce przed w:rPr w w:pPr (zapis LibreOffice) jest poprawiane."""
        change = f'w:id="1" w:author="Redaktor" {nsdecls("w")}'
        cases = (
            # nazwa, w:pPr akapitu, w:rPr przebiegu, oczekiwane dzieci w:pPr, w:rPr i poprzednich właściwości
            ('sama zmiana formatowania znaków', '', f'<w:rPrChange {change}><w:rPr><w:i/></w:rPr></w:rPrChange>',
             None, ['rFonts', 'sz', 'szCs', 'rPrChange'], ['rFonts', 'i', 'sz', 'szCs']),
            ('zmiana akapitu z LibreOffice', f'<w:jc w:val="center"/><w:pPrChange {change}/><w:rPr/>', '',
             ['spacing', 'ind', 'contextualSpacing', 'jc', 'rPr', 'pPrChange'], None,
             ['spacing', 'ind', 'contextualSpacing']),
        )
        for label, paragraph_xml, run_xml, paragraph_children, run_children, previous_children in cases:
            with self.subTest(label):
                document = Document()
                paragraph = document.add_paragraph()
                run = paragraph.add_run('Tekst')
                if paragraph_xml:
                    properties = parse_xml(f'<w:pPr {nsdecls("w")}>{paragraph_xml}</w:pPr>')
                    paragraph._p.insert(0, properties)
                if run_xml:
                    run._r.insert(0, parse_xml(f'<w:rPr {nsdecls("w")}>{run_xml}</w:rPr>'))
                result = Document(normalize_docx(saved(document))).paragraphs[0]

                def names(node):
                    return [child.tag.split('}')[1] for child in node]

                if paragraph_children is not None:
                    self.assertEqual(names(result._p.pPr), paragraph_children)
                    previous = result._p.pPr.find(f'{qn("w:pPrChange")}/{qn("w:pPr")}')
                    self.assertEqual(previous.find(qn('w:ind')).get(qn('w:firstLine')), '709')
                else:
                    self.assertEqual(names(result.runs[0]._r.rPr), run_children)
                    previous = result.runs[0]._r.rPr.find(f'{qn("w:rPrChange")}/{qn("w:rPr")}')
                    self.assertEqual(previous.find(qn('w:rFonts')).get(qn('w:ascii')), 'Times New Roman')
                self.assertEqual(names(previous), previous_children)
