"""Słownik zamian Odkurzacza w CMS: dane początkowe i ich aktualizacja, uprawnienia, walidacja i użycie w zadaniach."""
import importlib
from unittest.mock import patch

from django.apps import apps

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, tag
from django.urls import reverse
from docx import Document

from core.models import ReplacementEntry
from core.odkurzacz_forms import REPLACEMENT_ACTIVE_MAX, ReplacementEntryForm
from core.program_test_support import run_program
from core.services import program_jobs as jobs
from core.services.odkurzacz import DEFAULT_REPLACEMENTS, correct_editorial_text
from fantazmaty.testing import create_member, create_superuser

from .support import IsolatedProgramDirsMixin, downloaded_document, saved

INITIAL = importlib.import_module('core.migrations.0047_initial_replacements').INITIAL_REPLACEMENTS
OWNER_LIST = importlib.import_module('core.migrations.0067_replacements_owner_list')
REFORM = importlib.import_module('core.migrations.0068_replacements_reform_2026')
LIST_URL = reverse('core:replacement_list')


def entry_data(**values):
    return {'kind': 'word', 'source': 'kotek', 'target': 'kocik', 'is_active': 'on', 'position': '5', 'note': '', **values}


def owner_list_applied(items):
    """Pozycje (rodzaj, zapis, zamiana) po zmianach migracji 0067 i 0068 – jak w nich, na liście w pamięci."""
    changed = dict(OWNER_LIST.CHANGED)
    result = [changed.get(item, item) for item in items if item not in OWNER_LIST.REMOVED]
    result += [item for item in OWNER_LIST.ADDED if item not in result]
    return [item for item in result if item not in REFORM.REMOVED]


class InitialDictionaryTests(TestCase):
    def test_migrations_fill_the_dictionary_with_the_built_in_list(self):
        """Po migracjach 0047, 0067 i 0068 słownik w bazie ma te same pozycje co lista Odkurzacza w kodzie (programu
        lokalnego), wszystkie aktywne; pozycja pickup → pikap i dwie zamiany sprzeczne z reformą 2026 są usunięte."""
        entries = list(ReplacementEntry.objects.order_by('position', 'pk'))
        self.assertEqual([(entry.kind, entry.source, entry.target) for entry in entries], owner_list_applied(INITIAL))
        self.assertEqual(sorted((entry.kind, entry.source, entry.target) for entry in entries),
                         sorted((item['kind'], item['source'], item['target']) for item in DEFAULT_REPLACEMENTS))
        self.assertFalse([entry for entry in entries if 'pickup' in entry.source])
        self.assertTrue(all(entry.is_active for entry in entries))
        self.assertEqual(ReplacementEntry.active_replacements(), [entry.as_replacement() for entry in entries])

    def test_dictionary_from_the_database_corrects_like_the_built_in_list(self):
        """Kolejność w bazie (nowe pozycje na końcu) nie zmienia wyniku względem listy z kodu."""
        from .test_editorial_rules import OwnerReplacementListTests
        text = ' '.join(source for changed, kept in OwnerReplacementListTests.CASES.values()
                        for source in [*changed, *kept])
        self.assertEqual(
            correct_editorial_text(text, {'user_word_corrections'}, replacements=ReplacementEntry.active_replacements()),
            correct_editorial_text(text, {'user_word_corrections'}))

    def test_owner_list_migration_keeps_coordinator_changes(self):
        """Migracja 0067 usuwa i zmienia tylko pozycje w domyślnej postaci, nie dubluje zapisu dodanego wcześniej
        przez koordynatora (bez względu na wielkość liter), a cofnięcie przywraca poprzedni słownik."""
        pickup, miauczenie = OWNER_LIST.REMOVED[0], OWNER_LIST.CHANGED[3]
        OWNER_LIST.backwards(apps, None)
        self.assertEqual(sorted((entry.kind, entry.source, entry.target) for entry in ReplacementEntry.objects.all()),
                         sorted(INITIAL))
        ReplacementEntry.objects.filter(source=pickup[1]).update(target='furgonetka')
        ReplacementEntry.objects.filter(source=miauczenie[0][1]).update(target=r'miau\g<ending>')
        ReplacementEntry.objects.create(kind='word', source='Księżami', target='księżą', position=1)
        OWNER_LIST.forwards(apps, None)
        entries = {(entry.kind, entry.source): entry.target for entry in ReplacementEntry.objects.all()}
        self.assertEqual(entries[('pattern', pickup[1])], 'furgonetka')
        self.assertEqual(entries[('pattern', miauczenie[0][1])], r'miau\g<ending>')
        self.assertNotIn(('pattern', miauczenie[1][1]), entries)
        self.assertEqual(ReplacementEntry.objects.filter(source__iexact='księżami').count(), 1)
        self.assertEqual(entries[('word', 'Księżami')], 'księżą')
        self.assertEqual(entries[('word', 'dellarte')], 'dell’arte')
        for old, new in OWNER_LIST.CHANGED[:3]:
            self.assertEqual(entries[('pattern', new[1])], new[2])
            self.assertNotIn(('pattern', old[1]), entries)

    def test_reform_migration_removes_only_default_entries_and_can_be_reversed(self):
        """Migracja 0068 usuwa zamiany „niemniejszy → nie mniejszy” i „nienajlepszy → nie najlepszy” (sprzeczne
        z reformą 2026) tylko w domyślnej postaci; cofnięcie dopisuje je na końcu słownika bez dublowania, a słownik
        z bazy nie rozdziela już „nie” ze stopniem wyższym (spójnik „niemniej” i „nie mniej niż” bez zmian)."""
        sources = [source for _, source, _ in REFORM.REMOVED]
        REFORM.backwards(apps, None)
        REFORM.backwards(apps, None)
        entries = list(ReplacementEntry.objects.order_by('position', 'pk'))
        self.assertEqual([(entry.kind, entry.source, entry.target) for entry in entries[-2:]], list(REFORM.REMOVED))
        self.assertEqual(sorted((entry.kind, entry.source, entry.target) for entry in entries),
                         sorted(owner_list_applied(INITIAL) + list(REFORM.REMOVED)))
        text = 'Był niemniejszy i nienajlepszy, niemniej ciekawy, nie mniej niż wczoraj.'
        self.assertEqual(correct_editorial_text(text, {'user_word_corrections'},
                                                replacements=ReplacementEntry.active_replacements()),
                         'Był nie mniejszy i nie najlepszy, niemniej ciekawy, nie mniej niż wczoraj.')
        ReplacementEntry.objects.filter(source=sources[0]).update(target=r'nie mniejs\g<ending> (sprawdź)')
        REFORM.forwards(apps, None)
        self.assertEqual(list(ReplacementEntry.objects.filter(source__in=sources).values_list('source', 'target')),
                         [(sources[0], r'nie mniejs\g<ending> (sprawdź)')])
        ReplacementEntry.objects.filter(source=sources[0]).delete()
        self.assertEqual(correct_editorial_text(text, {'user_word_corrections'},
                                                replacements=ReplacementEntry.active_replacements()), text)

    def test_initial_entries_pass_the_edit_validation(self):
        """Superuser może zapisać każdą pozycję początkową bez poprawiania jej."""
        for entry in ReplacementEntry.objects.all():
            with self.subTest(source=entry.source):
                form = ReplacementEntryForm(entry_data(kind=entry.kind, source=entry.source, target=entry.target),
                                            instance=entry, allow_patterns=True)
                self.assertTrue(form.is_valid(), form.errors)


class DictionaryPageTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('slownik-redaktor', 'Redaktor')
        cls.coordinator = create_member('slownik-koordynator', 'Koordynator')
        cls.admin = create_superuser('slownik-admin')
        cls.pattern = ReplacementEntry.objects.filter(kind='pattern').first()
        cls.word = ReplacementEntry.objects.filter(kind='word').first()

    def test_team_member_sees_the_list_but_cannot_change_it(self):
        self.client.force_login(self.member)
        response = self.client.get(LIST_URL)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'niechcąco')
        self.assertContains(response, 'Słownik mogą zmieniać koordynatorzy i superuser.')
        self.assertIsNone(response.context['form'])
        self.assertNotContains(response, reverse('core:replacement_edit', args=[self.word.pk]))
        count = ReplacementEntry.objects.count()
        self.assertEqual(self.client.post(LIST_URL, entry_data()).status_code, 403)
        url = reverse('core:replacement_edit', args=[self.word.pk])
        for data in ({}, {'action': 'toggle'}, {'action': 'delete'}):
            self.assertEqual(self.client.post(url, data).status_code, 403)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(ReplacementEntry.objects.count(), count)

    def test_coordinator_adds_edits_disables_and_deletes_words(self):
        self.client.force_login(self.coordinator)
        response = self.client.post(LIST_URL, entry_data(note='Zdrobnienie'))
        self.assertRedirects(response, LIST_URL + '#replacement-table', fetch_redirect_response=False)
        entry = ReplacementEntry.objects.get(source='kotek')
        self.assertEqual((entry.target, entry.note, entry.position, entry.is_active), ('kocik', 'Zdrobnienie', 5, True))
        url = reverse('core:replacement_edit', args=[entry.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(url, {**entry_data(target='kicia'), 'action': 'save'})
        entry.refresh_from_db()
        self.assertEqual(entry.target, 'kicia')
        self.client.post(url, {'action': 'toggle'})
        entry.refresh_from_db()
        self.assertFalse(entry.is_active)
        self.assertNotIn(entry.as_replacement(), ReplacementEntry.active_replacements())
        self.client.post(url, {'action': 'toggle'})
        entry.refresh_from_db()
        self.assertTrue(entry.is_active)
        self.assertEqual(self.client.post(url, {'action': 'nieznana'}).status_code, 400)
        self.client.post(url, {'action': 'delete'})
        self.assertFalse(ReplacementEntry.objects.filter(pk=entry.pk).exists())
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_new_entry_suggests_the_next_position(self):
        self.client.force_login(self.coordinator)
        highest = max(ReplacementEntry.objects.values_list('position', flat=True))
        self.assertEqual(self.client.get(LIST_URL).context['form'].initial['position'], highest + 10)

    def test_only_superuser_adds_or_changes_regular_expressions(self):
        self.client.force_login(self.coordinator)
        response = self.client.get(LIST_URL)
        self.assertEqual([value for value, _ in response.context['form'].fields['kind'].choices], ['word'])
        self.assertContains(response, 'Tylko superuser')
        response = self.client.post(LIST_URL, entry_data(kind='pattern', source=r'\bkot(?P<e>a|y)?\b', target=r'pies\g<e>'))
        self.assertEqual(response.status_code, 400)
        self.assertFalse(ReplacementEntry.objects.filter(source__contains='kot(?P').exists())
        url = reverse('core:replacement_edit', args=[self.pattern.pk])
        for data in ({}, {'action': 'toggle'}, {'action': 'delete'}):
            self.assertEqual(self.client.post(url, data).status_code, 403)
        self.assertEqual(self.client.get(url).status_code, 403)
        # Koordynator nie zmieni też słowa we wzorzec.
        word_url = reverse('core:replacement_edit', args=[self.word.pk])
        response = self.client.post(word_url, {**entry_data(kind='pattern', source='abc+', target='x'), 'action': 'save'})
        self.assertEqual(response.status_code, 400)
        self.word.refresh_from_db()
        self.assertEqual(self.word.kind, 'word')

        self.client.force_login(self.admin)
        response = self.client.post(LIST_URL, entry_data(kind='pattern', source=r'\bkot(?P<e>a|y)?\b', target=r'pies\g<e>'))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(ReplacementEntry.objects.filter(kind='pattern', target=r'pies\g<e>').exists())
        self.client.post(url, {'action': 'toggle'})
        self.pattern.refresh_from_db()
        self.assertFalse(self.pattern.is_active)

    def test_switching_on_an_entry_respects_the_active_limit(self):
        self.client.force_login(self.coordinator)
        disabled = ReplacementEntry.objects.create(source='wylaczone', target='x', is_active=False)
        with patch('core.views.replacements.REPLACEMENT_ACTIVE_MAX', ReplacementEntry.objects.filter(is_active=True).count()):
            response = self.client.post(reverse('core:replacement_edit', args=[disabled.pk]), {'action': 'toggle'},
                                        follow=True)
        self.assertContains(response, 'aktywnych pozycji')
        disabled.refresh_from_db()
        self.assertFalse(disabled.is_active)

    def test_programs_page_links_to_the_dictionary_with_the_active_count(self):
        self.client.force_login(self.member)
        ReplacementEntry.objects.filter(pk=self.word.pk).update(is_active=False)
        active = ReplacementEntry.objects.filter(is_active=True).count()
        response = self.client.get(reverse('core:programs'))
        self.assertContains(response, f'Własne zamiany słownikowe – słownik zamian (aktywne pozycje: {active})', count=2)
        self.assertContains(response, f'href="{LIST_URL}"')
        self.assertNotContains(response, '23 pozycje')
        self.assertContains(self.client.get(LIST_URL), f'Aktywne pozycje ({active})')

    def test_dictionary_pages_highlight_programs_in_navigation(self):
        self.client.force_login(self.coordinator)
        for url in (LIST_URL, reverse('core:replacement_edit', args=[self.word.pk])):
            self.assertEqual(self.client.get(url).context['navigation_current'], 'programs')


class EntryValidationTests(TestCase):
    def form(self, allow_patterns=True, instance=None, **values):
        return ReplacementEntryForm(entry_data(**values), allow_patterns=allow_patterns, instance=instance)

    def assertInvalid(self, form, field, message):
        self.assertFalse(form.is_valid())
        self.assertIn(message, ' '.join(form.errors.get(field, [])), form.errors)

    def test_word_entries(self):
        self.assertTrue(self.form(source='  pół-żartem  ', target='półżartem').is_valid())
        form = self.form(source='  kotek  ')
        self.assertTrue(form.is_valid())
        self.assertEqual(form.cleaned_data['source'], 'kotek')
        self.assertInvalid(self.form(source='kot.*'), 'source', 'Słowo może zawierać litery')
        self.assertInvalid(self.form(source='123'), 'source', 'Słowo może zawierać litery')
        self.assertInvalid(self.form(source='a' * 101), 'source', 'najwyżej 100 znaków')
        self.assertInvalid(self.form(target='   '), 'target', 'To pole jest wymagane.')
        self.assertInvalid(self.form(target='a\tb'), 'target', 'znaków sterujących')
        self.assertInvalid(self.form(target='x' * 201), 'target', 'co najwyżej 200')

    def test_duplicates_are_found_regardless_of_letter_case(self):
        """MySQL porównuje bez wielkości liter – duplikaty sprawdzamy w Pythonie, nie unikalnym indeksem."""
        self.assertInvalid(self.form(source='NIECHCĄCO'), 'source', 'Taka pozycja już jest w słowniku.')
        existing = ReplacementEntry.objects.get(source='niechcąco')
        self.assertTrue(self.form(source='Niechcąco', target='niechcący', instance=existing).is_valid())
        # Ten sam zapis jako wzorzec to inna pozycja.
        self.assertTrue(self.form(kind='pattern', source='niechcąco', target='niechcący').is_valid())

    def test_regular_expressions(self):
        self.assertTrue(self.form(kind='pattern', source=r'\bkot(?P<e>a|y)?\b', target=r'pies\g<e>').is_valid())
        self.assertInvalid(self.form(kind='pattern', source='kot(', target='pies'), 'source', 'Nieprawidłowe wyrażenie')
        self.assertInvalid(self.form(kind='pattern', source=r'\bkot\b', target=r'pies\g<e>'), 'target', 'Nieprawidłowa zamiana')
        self.assertInvalid(self.form(kind='pattern', source=r'\bkot\b', target=r'pies\q'), 'target', 'Nieprawidłowa zamiana')
        self.assertInvalid(self.form(kind='pattern', source=r'\bkot\b', target=r'pies\2'), 'target', 'Nieprawidłowa zamiana')
        for empty in ('x*', r'\b', '(?:)', 'a?'):
            with self.subTest(source=empty):
                self.assertInvalid(self.form(kind='pattern', source=empty, target='y'), 'source', 'pustego tekstu')
        self.assertInvalid(self.form(kind='pattern', source='a' * 501, target='b'), 'source', 'co najwyżej 500')

    def test_costly_patterns_are_refused_like_in_the_cleaner(self):
        """Odkurzacz po cichu pomija wzorce o ryzyku bardzo długiego działania – formularz ich nie przyjmuje."""
        for source in ('(a+)+b', r'(\w+\s?)+$', r'\w*\w*\w*\w*!', '(a?){22}a{22}'):
            with self.subTest(source=source):
                self.assertInvalid(self.form(kind='pattern', source=source, target='x'), 'source',
                                   'Odkurzacz nie użyje tego wzorca')

    def test_words_with_dots_hyphens_and_apostrophes_on_the_edges(self):
        for source in ('np.', 'tzw.', '-owy', "rock'n'roll", "o'"):
            with self.subTest(source=source):
                self.assertTrue(self.form(source=source, target='x').is_valid())

    def test_letter_case_and_order_are_explained(self):
        form = ReplacementEntryForm(allow_patterns=True)
        self.assertIn('NATO', form.fields['target'].help_text)
        self.assertIn('nie jest ponownie zamieniany', form.fields['position'].help_text)

    def test_active_entry_limit_is_explained(self):
        ReplacementEntry.objects.bulk_create(ReplacementEntry(source=f'slowo{index}', target='x', position=index)
                                             for index in range(REPLACEMENT_ACTIVE_MAX - ReplacementEntry.objects.count()))
        form = self.form(source='nowe')
        self.assertFalse(form.is_valid())
        self.assertIn(f'Słownik ma już {REPLACEMENT_ACTIVE_MAX} aktywnych pozycji', str(form.non_field_errors()))
        # Wyłączoną pozycję można dodać, a zmiana istniejącej aktywnej pozycji nie liczy jej drugi raz.
        self.assertTrue(self.form(source='nowe', is_active='').is_valid())
        existing = ReplacementEntry.objects.get(source='slowo3')
        self.assertTrue(self.form(source='slowo3', target='y', instance=existing).is_valid())
        disabled = ReplacementEntry.objects.create(source='wylaczone', target='x', is_active=False)
        self.assertFalse(self.form(source='wylaczone', instance=disabled).is_valid())

    def test_patterns_require_permission(self):
        form = self.form(allow_patterns=False, kind='pattern', source=r'\bkot\b', target='pies')
        self.assertFalse(form.is_valid())
        self.assertIn('kind', form.errors)


class DictionaryInCleanerTests(IsolatedProgramDirsMixin, TestCase):
    """Prawdziwy Odkurzacz używa pozycji ze słownika w CMS."""

    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('slownik-odkurzacz', 'Redaktor')

    def setUp(self):
        super().setUp()
        self.client.force_login(self.member)

    def clean(self, text):
        doc = Document()
        doc.add_paragraph(text)
        response = run_program(self.client, {'document': SimpleUploadedFile('tekst.docx', saved(doc).getvalue()),
                                             'rules': ['user_word_corrections']})
        self.assertEqual(response.status_code, 200, getattr(response, 'content', b'')[:300])
        return downloaded_document(response).paragraphs[0].text

    @tag('slow')
    def test_entries_are_applied_disabled_ones_are_not(self):
        ReplacementEntry.objects.create(kind='word', source='kotek', target='kocik', position=1)
        ReplacementEntry.objects.filter(source='zapewnie').update(is_active=False)
        self.assertEqual(self.clean('Niechcąco zapewnie Kotek mentlikiem.'), 'Niechcący zapewnie Kocik mętlikiem.')

    @tag('slow')
    def test_empty_dictionary_changes_nothing(self):
        ReplacementEntry.objects.all().delete()
        self.assertEqual(self.clean('Niechcąco mentlikiem.'), 'Niechcąco mentlikiem.')

    def test_replacements_are_read_when_the_job_is_created(self):
        with patch.object(jobs, 'launch'):
            response = self.client.post(reverse('core:programs'), {
                'document': SimpleUploadedFile('tekst.docx', saved(Document()).getvalue())}, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 202)
        folder = next(jobs.root().iterdir())
        self.assertEqual(len(jobs.read_json(folder / 'request.json')['options']['replacements']),
                         ReplacementEntry.objects.count())
