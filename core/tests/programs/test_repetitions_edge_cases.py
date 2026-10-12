"""Analiza powtórzeń: walidacja parametrów, śledzone zmiany, brak modelu językowego, pamięć lematów,
dodatkowe kontrole wywołane samodzielnie, podobny początek słowa, odległe powtórzenia oraz nietypowe
przebiegi tekstu (puste w:t, hiperłącza, pola, same tła)."""
from collections import OrderedDict
from contextlib import contextmanager
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase, tag
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from core.services import document_repetitions as repetitions
from core.services.document_errors import DocumentInputError
from core.services.document_repetitions import color_document

from .support import (add_complex_field, add_internal_link, add_tracked_insertion, docx_with_paragraph_xml, saved,
                      text_run)

ONLY_REPETITIONS = {'duplicates': False, 'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}


@contextmanager
def simple_language():
    """Lemat = słowo małymi literami, bez lematów z kontekstu i z jednym zdaniem na akapit (bez modelu)."""
    def context(texts, *, sentences, status_callback=None):
        return [({}, [(0, len(text))]) for text in texts]

    with patch.object(repetitions, '_lemma_map', side_effect=lambda words: {w.lower(): w.lower() for w in words}), \
            patch.object(repetitions, '_context_analysis', side_effect=context):
        yield


def marked(source, **options):
    with simple_language():
        return Document(color_document(source, **options))


def coloured(paragraph):
    return [run.text for run in paragraph.runs if run.font.color.rgb]


class ParameterTests(SimpleTestCase):
    def test_parameters_out_of_range_are_refused(self):
        source = saved(Document())
        cases = {
            'zakres jako wartość logiczna': ({'window_size': True}, ValueError),
            'zero słów': ({'min_word_length': '0'}, ValueError),
            'zakres ponad 500': ({'window_size': 501}, DocumentInputError),
            'nieznana kontrola': ({'analysis_options': {'nieznana': True}}, DocumentInputError),
            'próg zdania ponad 10 000': ({'analysis_options': {'sentence_limit': 10_001}}, DocumentInputError),
        }
        for label, (options, error) in cases.items():
            with self.subTest(label), self.assertRaises(error):
                color_document(source, **options)

    def test_document_with_tracked_changes_is_refused(self):
        document = Document()
        add_tracked_insertion(document.add_paragraph('Tekst '), 'dopisany')
        with self.assertRaises(DocumentInputError) as raised:
            color_document(saved(document))
        self.assertEqual(raised.exception.public_code, 'tracked_changes')


class LanguageModelTests(SimpleTestCase):
    def test_missing_polish_model_is_an_import_error(self):
        with patch('spacy.load', side_effect=OSError('brak pakietu')), \
                self.assertRaisesMessage(ImportError, 'Brakuje polskiego modelu analizy.'):
            repetitions._get_nlp.__wrapped__()  # bez pamięci podręcznej prawdziwego modelu

    def test_rule_based_lemmatizer_keeps_the_tagger(self):
        """Lematyzator regułowy potrzebuje części mowy: pomijamy tylko parser, senter i NER."""
        nlp = SimpleNamespace(pipe_names=['tok2vec', 'tagger', 'parser', 'lemmatizer', 'ner'],
                              get_pipe=lambda name: object())
        self.assertEqual(repetitions._unused_pipes(nlp, 'lemmas'), ['parser', 'ner'])
        self.assertEqual(repetitions._lemma_from_doc([], 'słowo'), 'słowo')

    @tag('slow')
    def test_word_lemmas_with_a_bounded_cache(self):
        with patch.object(repetitions, '_LEMMA_CACHE', OrderedDict()) as cache, \
                patch.object(repetitions, '_LEMMA_CACHE_LIMIT', 2):
            self.assertEqual(repetitions.get_lemma('KOTAMI'), 'kot')
            self.assertEqual(repetitions._lemma_map(['psami', 'domami']), {'psami': 'pies', 'domami': 'dom'})
            self.assertEqual(list(cache), [(repetitions.MODEL_NAME, 'psami'), (repetitions.MODEL_NAME, 'domami')])

    @tag('slow')
    def test_extra_checks_can_find_sentences_and_mark_the_document(self):
        """Wywołane samodzielnie: same rozpoznają zdania (z komunikatami postępu) i od razu oznaczają tekst."""
        document = Document()
        document.add_paragraph('Krótko. Bardzo długie zdanie ma tutaj wiele słów.')
        messages = []
        report = repetitions._extra_checks(document.paragraphs, {}, set(), set(),
                                           {**ONLY_REPETITIONS, 'long_sentences': True, 'sentence_limit': 3},
                                           status_callback=messages.append)
        self.assertEqual(report['long_sentences'], 1)
        self.assertEqual(messages, ['Rozpoznawanie zdań i sprawdzanie ich długości…'])
        runs = document.paragraphs[0].runs
        self.assertEqual([run.font.highlight_color for run in runs], [None, WD_COLOR_INDEX.TURQUOISE])
        self.assertEqual(runs[1].text, 'Bardzo długie zdanie ma tutaj wiele słów.')


class GroupingTests(SimpleTestCase):
    def test_similar_beginnings_and_repetitions_outside_the_window(self):
        document = Document()
        document.add_paragraph('Domownik domowy.')
        document.add_paragraph('Lasy szumią wokół, potem lasy.')
        cases = {'bez podobnego początku, okno 1': ({'window_size': 1}, [], []),
                 'podobny początek, okno 35': ({'include_prefix_matches': True}, ['Domownik', 'domowy'],
                                               ['Lasy', 'lasy'])}
        for label, (options, first, second) in cases.items():
            with self.subTest(label):
                result = marked(saved(document), analysis_options=ONLY_REPETITIONS, **options)
                self.assertEqual([coloured(p) for p in result.paragraphs], [first, second])


    def test_neighbouring_duplicate_across_a_paragraph_boundary(self):
        document = Document()
        document.add_paragraph('Widzę kota')
        document.add_paragraph('kota widzę.')
        result = marked(saved(document), analysis_options={**ONLY_REPETITIONS, 'duplicates': True})
        self.assertEqual([[run.text for run in p.runs if run.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN]
                          for p in result.paragraphs], [['kota'], ['kota']])


class UnusualRunTests(SimpleTestCase):
    def test_empty_text_nodes_keep_their_place(self):
        source = docx_with_paragraph_xml(
            '<w:r><w:t/><w:t xml:space="preserve">Koty i koty, </w:t></w:r>',
            '<w:r><w:t xml:space="preserve">lisy </w:t><w:t/><w:t xml:space="preserve">i lisy.</w:t></w:r>')
        paragraph = marked(source, analysis_options=ONLY_REPETITIONS).paragraphs[0]
        self.assertEqual(paragraph.text, 'Koty i koty, lisy i lisy.')
        self.assertEqual(coloured(paragraph), ['Koty', 'koty', 'lisy', 'lisy'])
        self.assertEqual(len(paragraph._p.xpath('.//w:t[not(text())]')), 2)

    def test_links_and_field_results_are_not_words_of_the_text(self):
        document = Document()
        paragraph = document.add_paragraph('Lampa ')
        add_internal_link(paragraph, '_Toc1', 'lampa')
        add_complex_field(paragraph, 'AUTHOR', 'lampa')
        paragraph.add_run(' i lampa.')
        result = marked(saved(document), analysis_options=ONLY_REPETITIONS).paragraphs[0]
        # Kolor dostają tylko słowa tekstu: ani odnośnik, ani wynik pola (między separate i end).
        self.assertEqual(result._p.xpath('.//w:r[w:rPr/w:color]/w:t/text()'), ['Lampa', 'lampa'])
        self.assertEqual(result.runs[-2].text, 'lampa')
        self.assertIsNotNone(result.runs[-2].font.color.rgb)

    def test_runs_with_backgrounds_only(self):
        """Puste nawiasy w osobnym przebiegu dostają samo tło; sąsiednie pary – jeden przebieg."""
        source = docx_with_paragraph_xml(text_run('Tekst '), text_run('()'), text_run(' i '), text_run('[]{}'))
        runs = marked(source, analysis_options={**ONLY_REPETITIONS, 'empty_pairs': True}).paragraphs[0].runs
        self.assertEqual([(run.text, run.font.highlight_color, run.font.color.rgb) for run in runs], [
            ('Tekst ', None, None), ('()', WD_COLOR_INDEX.PINK, None), (' i ', None, None),
            ('[]{}', WD_COLOR_INDEX.PINK, None)])
