"""Odkurzacz: zaimki daleko od początku linii i po adresach, zakresy godzin w podsumowaniu, wielka litera
po wykrzykniku, kontrola wzorców słownika (zagnieżdżenia, pomiar czasu z SIGALRM i bez niego), słownik zamian
oraz śledzone zmiany w nietypowych przebiegach."""
import threading

from django.test import SimpleTestCase
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.services import odkurzacz
from core.services.odkurzacz import CompiledReplacements, _compile_pattern, clean_docx, correct_editorial_text

from .support import add_simple_field, docx_with_paragraph_xml, saved

W = nsdecls('w')


class EndlessPattern:
    """Wzorzec, którego dopasowania nigdy się nie kończą (jak katastrofalne nawroty) – do pomiaru czasu."""

    def finditer(self, sample):
        while True:
            yield sample


class FinitePattern:
    def __init__(self):
        self.searched = []

    def finditer(self, sample):
        self.searched.append(sample)
        yield from sample[:3]


class EditorialTextTests(SimpleTestCase):
    def test_pronouns_far_from_the_line_start_and_after_addresses(self):
        cases = {
            # Ponad 80 znaków odstępu: kontekst sięga początku linii (przecinek – środek zdania).
            'Zdanie,' + ' ' * 90 + 'Jego dom': 'Zdanie,' + ' ' * 90 + 'jego dom',
            # Adres przed zaimkiem jest chroniony, ale nie blokuje dalszych poprawek.
            'Strona www.example.pl, a Jego dom': 'Strona www.example.pl, a jego dom',
        }
        for text, expected in cases.items():
            with self.subTest(text=text.strip()[:30]):
                self.assertEqual(correct_editorial_text(text, {'pronouns_lower'}), expected)

    def test_time_ranges_are_counted_per_rule(self):
        rules = {'dates_times', 'range_dash', 'range_spaces'}
        cases = {
            'Otwarte 8.00 - 16.00 codziennie.': ('Otwarte 08:00–16:00 codziennie.',
                                                 {'dates_times': 2, 'range_dash': 1, 'range_spaces': 2}),
            # Godziny już w dobrym zapisie: liczy się tylko półpauza.
            'Otwarte 08:00-16:00 codziennie.': ('Otwarte 08:00–16:00 codziennie.', {'range_dash': 1}),
        }
        for text, (expected, expected_counts) in cases.items():
            counts = {}
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, rules, counts=counts), expected)
                self.assertEqual(counts, expected_counts)

    def test_capital_letter_after_exclamation_and_question_marks(self):
        self.assertEqual(correct_editorial_text('Hej! jak się masz? dobrze.', {'sentence_case'}),
                         'Hej! Jak się masz? Dobrze.')


class PatternCheckTests(SimpleTestCase):
    def test_nested_repetitions_in_every_construct_are_refused(self):
        for source in (r'x(?:a|(b+)+)', r'(?=(a+)+)x', r'(?>(a+)+)x', r'(a)?(?(1)(b+)+|c)'):
            with self.subTest(source=source):
                self.assertIsNone(_compile_pattern(source))
        # Asercje, grupy atomowe i warunki bez zagnieżdżonych powtórzeń są w porządku.
        for source in (r'(?=a)kot', r'(?>a)kot', r'(a)?(?(1)kot|pies)'):
            with self.subTest(source=source):
                self.assertIsNotNone(_compile_pattern(source))
        for source, target in ((123, ''), ('a' * 501, ''), ('kot', 'x' * 201), ('', '')):
            with self.subTest(source=str(source)[:10], target=target[:10]):
                self.assertIsNone(_compile_pattern(source, target))

    def test_timing_probe_is_interrupted_by_an_alarm(self):
        with self.assertRaises(odkurzacz._PatternTimeout):
            odkurzacz._exhaust(EndlessPattern(), 'aaaa', 0.01)
        self.assertTrue(odkurzacz._too_slow(EndlessPattern()))

    def test_timing_probe_outside_the_main_thread_only_measures(self):
        pattern, results = FinitePattern(), []
        worker = threading.Thread(target=lambda: results.append(odkurzacz._exhaust(pattern, 'abcdef', 0.01)))
        worker.start()
        worker.join()
        self.assertEqual((results, pattern.searched), ([None], ['abcdef']))
        # Wzorzec dopasowujący próbki (pętla po wynikach) przechodzi kontrolę także w głównym wątku.
        self.assertIsNotNone(_compile_pattern(r'\w+'))


class ReplacementDictionaryTests(SimpleTestCase):
    def test_first_entry_of_a_word_wins_and_unchanged_matches_are_not_counted(self):
        compiled = CompiledReplacements([
            {'kind': 'word', 'source': 'pickup', 'target': 'pikap'},
            {'kind': 'word', 'source': 'PICKUP', 'target': 'furgonetka'},
            {'kind': 'pattern', 'source': r'\bkot\b', 'target': 'kot', 'checked': True},
        ])
        record = {}
        self.assertEqual(compiled.apply('Pickup i kot.', record), 'Pikap i kot.')
        self.assertEqual(record, {('Pickup', 'Pikap'): 1})


class CleanDocxGapTests(SimpleTestCase):
    def test_fields_and_empty_text_nodes_split_the_text(self):
        document = Document()
        paragraph = document.add_paragraph('Ala  ma ')
        add_simple_field(paragraph, 'AUTHOR', 'Jan  Nowak')
        paragraph._p.append(parse_xml(f'<w:r {W}><w:t/><w:t xml:space="preserve"> kota  i psa</w:t></w:r>'))
        with clean_docx(saved(document), {'spaces'}) as result:
            cleaned = Document(result).paragraphs[0]._p
        texts = [node.text for node in cleaned.iter(qn('w:t'))]
        self.assertEqual(texts, ['Ala ma ', 'Jan  Nowak', None, ' kota i psa'])

    def test_characters_inserted_beside_a_kept_tabulator(self):
        """Odstępy wokół półpauzy między tabulatorami: tabulatory zostają, nowe spacje trafiają obok nich."""
        source = docx_with_paragraph_xml('<w:r><w:t>Tekst</w:t><w:tab/><w:t>–</w:t><w:tab/><w:t>dalej</w:t></w:r>')
        with clean_docx(source, {'dash_spaces'}) as result:
            paragraph = Document(result).paragraphs[0]
        self.assertEqual(paragraph.text, 'Tekst\t – \tdalej')
        self.assertEqual(len(paragraph._p.findall('.//' + qn('w:tab'))), 2)

    def test_tracked_changes_keep_run_attributes_and_join_neighbouring_deletions(self):
        source = docx_with_paragraph_xml(
            '<w:r w:rsidR="00AB12CD"><w:t xml:space="preserve">Ala </w:t><w:t xml:space="preserve">ma </w:t>'
            '<w:t xml:space="preserve">kota  </w:t><w:tab/><w:t>psa</w:t></w:r>')
        with clean_docx(source, {'spaces', 'tabs'}, track_changes=True) as result:
            paragraph = Document(result).paragraphs[0]._p
        runs = paragraph.xpath('.//w:r')
        self.assertTrue(all(run.get(qn('w:rsidR')) == '00AB12CD' for run in runs))
        deletions = paragraph.xpath('./w:del')
        self.assertEqual(len(deletions), 1)  # spacja i tabulator w jednej śledzonej zmianie
        self.assertEqual([child.tag for run in deletions[0] for child in run], [qn('w:delText'), qn('w:tab')])
        accepted = ''.join(node.text for node in paragraph.xpath('./w:r/w:t | ./w:ins/w:r/w:t'))
        self.assertEqual(accepted, 'Ala ma kota psa')

    def test_tracked_removal_of_empty_paragraphs_with_earlier_revisions(self):
        """Seria pustych akapitów z wcześniejszymi poprawkami Odkurzacza: usuwany jest znak akapitu."""
        mark = 'w:id="1" w:author="Odkurzacz" w:date="2026-01-01T00:00:00Z"'
        document = Document()
        document.add_paragraph('Tekst')
        first = document.add_paragraph('  ')
        first._p.append(parse_xml(f'<w:ins {W} {mark}><w:r><w:t xml:space="preserve"> </w:t></w:r></w:ins>'))
        first._p.get_or_add_pPr().append(parse_xml(f'<w:pPrChange {W} w:id="2" w:author="Redakcja"><w:pPr/>'
                                                   '</w:pPrChange>'))
        second = document.add_paragraph()
        second._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {W}><w:ins {mark.replace("1", "3", 1)}/></w:rPr>'))
        document.add_paragraph()
        inserted = document.add_paragraph()
        inserted._p.append(parse_xml(f'<w:ins {W} {mark.replace("1", "4", 1)}><w:r><w:t>Dopisane</w:t></w:r></w:ins>'))
        document.add_paragraph()
        summary = {}
        with clean_docx(saved(document), {'empty_paragraphs'}, track_changes=True, summary=summary) as result:
            paragraphs = Document(result).paragraphs
        self.assertEqual(summary, {'empty_paragraphs': 2})
        first, second = paragraphs[1]._p, paragraphs[2]._p
        # Tekst pustego akapitu trafia do w:del jako w:delText, a nasza wcześniejsza wstawka znika.
        self.assertEqual(first.xpath('./w:del/w:r/w:delText/text()'), ['  '])
        self.assertEqual(first.xpath('./w:ins'), [])
        properties = first.find(qn('w:pPr'))
        self.assertEqual([child.tag for child in properties], [qn('w:rPr'), qn('w:pPrChange')])
        self.assertIsNotNone(properties.find(f'{qn("w:rPr")}/{qn("w:del")}'))
        self.assertEqual([child.tag for child in second.find(qn('w:pPr')).find(qn('w:rPr'))],
                         [qn('w:ins'), qn('w:del')])
        # Akapit z dopisanym tekstem nie jest pusty i przerywa serię.
        self.assertEqual(paragraphs[4]._p.xpath('.//w:del'), [])
        self.assertEqual(len(paragraphs), 6)
