"""Dopasowanie akapitów DOCX do HTML i części dokumentu: twardy łącznik, pola tekstowe (mc:Choice i
mc:Fallback), podział strony przed dalszymi znacznikami i przed listą, nagłówek wspólny dla kilku sekcji
oraz kopia DOCX dla Mammotha (elementy z rzadszych edytorów, które Mammoth pomija albo skleja)."""
from io import BytesIO

import mammoth
from django.test import SimpleTestCase
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import html

from core.services.document_html import (
    MAMMOTH_NOTES, docx_paragraph_key, mammoth_source, paragraph_queues, split_page_breaks,
)
from core.services.document_stories import document_stories

from .editor_fixtures import R_NS, WORD_PART
from .editor_fixtures_exotic import (
    STYLES, chart_and_smart_art, frames_and_drop_cap, note_reference, notes_part, p, r, word_docx,
)
from .support import add_text_box


def blocks(root):
    return [html.tostring(child, encoding='unicode') for child in root]


class ParagraphKeyTests(SimpleTestCase):
    def test_hyphen_characters_are_part_of_the_key(self):
        paragraph = parse_xml(f'<w:p {nsdecls("w")}><w:r><w:t>e</w:t><w:noBreakHyphen/><w:t>mail</w:t>'
                              '<w:softHyphen/><w:t>owy</w:t></w:r></w:p>')
        self.assertEqual(docx_paragraph_key(paragraph), 'e\u2011mail\u00adowy')

    def test_text_box_is_matched_by_its_fallback_copy_only(self):
        document = Document()
        add_text_box(document.add_paragraph('Przed polem.'), 'W polu')
        queues = paragraph_queues(document)
        # Akapit pola tekstowego liczy się raz (kopia w mc:Choice jest pomijana), a akapit z polem – bez jego treści.
        self.assertEqual({key: len(queue) for key, queue in queues.items()}, {'Przed polem.': 1, 'W polu': 1})


class PageBreakSplitTests(SimpleTestCase):
    def test_markup_after_the_break_moves_to_the_next_block(self):
        root = html.fragment_fromstring('<p>A<br class="page-break"><strong>B</strong> C</p>', create_parent='div')
        split_page_breaks(root)
        self.assertEqual(blocks(root), ['<p>A</p>', '<p data-page-break="1"><strong>B</strong> C</p>'])

    def test_list_after_a_break_starts_the_new_page(self):
        root = html.fragment_fromstring('<p>Koniec<br class="page-break"></p><ul><li>Punkt</li></ul>',
                                        create_parent='div')
        split_page_breaks(root)
        self.assertEqual(blocks(root), ['<p>Koniec</p>', '<ul data-page-break="1"><li>Punkt</li></ul>'])


class DocumentStoriesTests(SimpleTestCase):
    def test_fallback_paragraphs_are_optional_and_shared_parts_are_read_once(self):
        document = Document()
        add_text_box(document.add_paragraph('Przed polem.'), 'W polu')
        header = document.sections[0].header
        header.paragraphs[0].text = 'Nagłówek'
        # Druga relacja do tej samej części (np. ten sam nagłówek dla stron parzystych).
        document.part.rels.add_relationship(RT.HEADER, header.part, 'rIdShared')
        for include_fallback, expected in ((True, ['Przed polem.', 'W polu', 'W polu']),
                                           (False, ['Przed polem.', 'W polu'])):
            with self.subTest(include_fallback=include_fallback):
                stories, _ = document_stories(document, include_fallback=include_fallback)
                self.assertEqual([story.name for story in stories], ['body', 'header'])
                self.assertEqual([p.text for p in stories[0].paragraphs], expected)
                self.assertEqual([p.text for p in stories[1].paragraphs], ['Nagłówek'])


class MammothSourceTests(SimpleTestCase):
    """Kopia DOCX dla Mammotha: treść elementów, które Mammoth pomija albo skleja, jako zwykłe przebiegi; ukryty
    tekst i odsyłacze donikąd usunięte z uwagą; plik bez takich elementów zostaje bez zmian."""

    def html(self, data):
        return mammoth.convert_to_html(BytesIO(data)).value

    def test_markup_mammoth_skips_or_joins_becomes_plain_text(self):
        """Regresja: e-book i PDF gubiły wynik w:fldSimple (numer podpisu z Worda, pole z OnlyOffice) i tekst
        w w:dir/w:bdo, sklejały słowa wokół w:cr i w:ptab, a inicjał (w:dropCap) był osobnym akapitem."""
        ptab = '<w:r><w:ptab w:relativeTo="margin" w:alignment="right" w:leader="none"/></w:r>'
        cases = {
            'pole proste': (p(r('Rysunek '), '<w:fldSimple w:instr=" SEQ Rysunek "><w:fldData>AA==</w:fldData>'
                              + r('1') + '</w:fldSimple>'), '<p>Rysunek 1</p>'),
            'kierunek tekstu': (p('<w:dir w:val="rtl">' + r('עברית') + '</w:dir>', r(' i '),
                                  '<w:bdo w:val="ltr">' + r('LTR') + '</w:bdo>'), '<p>עברית i LTR</p>'),
            'powrót karetki': (p(r('Przed'), '<w:r><w:cr/></w:r>', r('po')), '<p>Przed<br />po</p>'),
            'tabulator położenia': (p(r('Lewa'), ptab, r('prawa')), '<p>Lewa\tprawa</p>'),
            'tag inteligentny': (p('<w:smartTag w:element="City"><w:smartTagPr><w:attr w:name="a" w:val="b"/>'
                                   '</w:smartTagPr>' + r('Warszawa') + '</w:smartTag>'), '<p>Warszawa</p>'),
        }
        for label, (body, expected) in cases.items():
            with self.subTest(label):
                data, notes = mammoth_source(word_docx(body).getvalue())
                self.assertEqual(self.html(data), expected)
                self.assertEqual(notes, [])
        data, _ = mammoth_source(frames_and_drop_cap().getvalue())
        self.assertTrue(self.html(data).startswith('<p>Wielki inicjał zaczyna ten akapit.</p>'))

    def test_hidden_runs_and_dangling_note_references_are_dropped_with_notes(self):
        """Tekst ukryty (własne w:vanish i styl znakowy z w:vanish) nie trafia do e-booka – jak przy wydruku;
        w:vanish="0" i znaczniki pól zostają. Odsyłacz do przypisu, którego nie ma, przerywał konwersję."""
        styles = STYLES + ('<w:style w:type="character" w:styleId="Ukryty"><w:name w:val="Ukryty"/><w:rPr><w:vanish/>'
                           '</w:rPr></w:style>')
        hidden = '<w:rPr><w:vanish/></w:rPr>'
        index_entry = (f'<w:r>{hidden}<w:fldChar w:fldCharType="begin"/></w:r><w:r>{hidden}<w:instrText> XE "kot" '
                       f'</w:instrText></w:r><w:r>{hidden}<w:fldChar w:fldCharType="end"/></w:r>')
        body = p(r('Widać '), r('ukryte ', '<w:vanish/>'), r('stylem ', '<w:rStyle w:val="Ukryty"/>'),
                 r('jawne', '<w:vanish w:val="0"/>'), index_entry, note_reference('footnote', 1),
                 note_reference('footnote', 9), r('.'))
        footnotes = notes_part('footnote', {1: r(' Przypis.')})
        source = word_docx(body, styles=styles, parts={'word/footnotes.xml': footnotes},
                           relationships=[('rIdFoot', f'{R_NS}/footnotes', 'footnotes.xml')],
                           overrides=[('/word/footnotes.xml', WORD_PART.format('footnotes'))])
        data, notes = mammoth_source(source.getvalue())
        html = self.html(data)
        # Odsyłacz w przebiegu z indeksem górnym: Mammoth dokłada własny <sup>.
        self.assertTrue(html.startswith('<p>Widać jawne<sup><sup><a href="#footnote-1" id="footnote-ref-1">[1]</a>'
                                        '</sup></sup>.</p>'), html)
        for hidden_text in ('ukryte', 'stylem', 'footnote-9'):
            self.assertNotIn(hidden_text, html)
        self.assertEqual(notes, [MAMMOTH_NOTES['missing_note'], MAMMOTH_NOTES['hidden'] + ' (3 razy)'])

    def test_charts_are_reported_and_plain_documents_stay_unchanged(self):
        """Wykres i SmartArt Mammoth pomija bez ostrzeżenia – uwaga mówi o nich użytkownikowi; dokument bez
        poprawianych elementów wraca bajt w bajt."""
        charts = chart_and_smart_art().getvalue()
        self.assertEqual(mammoth_source(charts), (charts, [MAMMOTH_NOTES['graphic'] + ' (2 razy)']))
        plain = word_docx(p(r('Zwykły tekst.'))).getvalue()
        self.assertEqual(mammoth_source(plain), (plain, []))
