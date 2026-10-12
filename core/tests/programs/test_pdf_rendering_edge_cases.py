"""Skład PDF: komórki tabel (obraz, listy, sam tekst, nieobsługiwane znaczniki), pola i znaki w nagłówkach
i stopkach, bardzo wysoki nagłówek, odnośniki zewnętrzne, tabela w liście, wiele tabel, za wysokie tabele
po podziale strony i nagłówek rozdziału na dole strony."""
import re
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.test import SimpleTestCase
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fpdf import FPDF
from lxml import etree
from lxml import html as html_parser

from core.services import document_pdf
from core.services.document_pdf import CELL_NOTE, TALL_TABLE_NOTE, render_pdf

from .support import add_complex_field, add_text_box, png


class PdfGapTests(SimpleTestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)

    def render(self, document, content, assets=None):
        """Zwraca (HTML przekazany do fpdf jako drzewo, uwagi, liczba stron, wywołania nagłówków i stopek)."""
        source = self.directory / 'source.docx'
        document.save(source)
        captured, records = [], []
        write_html, draw = FPDF.write_html, document_pdf._draw_lines

        def html_spy(pdf, markup, **kwargs):
            if type(pdf) is not FPDF:  # bez próbnego składu tabel w osobnym PDF
                captured.append((markup, pdf.page_no(), pdf.y))
            return write_html(pdf, markup, **kwargs)

        def draw_spy(pdf, lines, y):
            records.append((pdf.page_no(), [text for text, _, _ in lines]))
            return draw(pdf, lines, y)

        with patch.object(FPDF, 'write_html', html_spy), patch.object(document_pdf, '_draw_lines', draw_spy):
            notes = render_pdf(source, self.directory / 'result.pdf', content, assets or {}, 'Test')
        pdf = (self.directory / 'result.pdf').read_bytes()
        self.captured = captured
        tree = html_parser.fragment_fromstring(''.join(part for part, _, _ in captured), create_parent='div')
        return tree, notes, len(re.findall(rb'/Type /Page\b(?!s)', pdf)), records

    def test_table_cells_with_images_lists_and_plain_text(self):
        image = png(20, 10).getvalue()
        content = (
            '<table><tr>'
            '<td><p><img src="a.png"></p></td>'  # sam obraz w akapicie: komórka zostaje bez zmian
            '<td><p><code>Kod</code></p></td>'  # znacznik spoza listy obsługiwanych: sam tekst
            '<td><ul><li><p>Kropka</p></li></ul><ol><li><p>Pierwszy</p></li><li><p>Drugi</p></li></ol></td>'
            '<td>Luźny <b>tekst</b></td>'  # tekst bez akapitów
            '<td><img src="a.png"><img src="a.png"></td>'  # same obrazy: zostaje pierwszy
            '<td><p></p><p> </p></td>'  # same puste akapity: pusta komórka
            '</tr></table>')
        with self.assertLogs('fpdf.html', 'WARNING'):  # fpdf2 pomija szerokość obrazu w komórce
            tree, notes, _, _ = self.render(Document(), content, {'a.png': image})
        cells = tree.xpath('.//td')
        self.assertEqual(len(cells[0].xpath('./p/img')), 1)
        self.assertEqual([cell.text_content() for cell in cells[1:4]],
                         ['Kod', '• Kropka\n1. Pierwszy\n2. Drugi', 'Luźny tekst'])
        self.assertIsNone(cells[1].xpath('./p')[0].get('align'))
        self.assertEqual([child.tag for child in cells[4]], ['img'])
        self.assertEqual((len(cells[5]), cells[5].text), (0, None))
        self.assertIn(CELL_NOTE, notes)

    def test_header_fields_breaks_text_boxes_and_two_cell_tables(self):
        document = Document()
        header = document.sections[0].header
        paragraph = header.paragraphs[0]
        paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t>e</w:t><w:noBreakHyphen/><w:t>mail</w:t><w:br/>'
                                      '<w:t>dalej </w:t></w:r>'))
        # Pole daty (proste i złożone) zostawia wynik zapisany przez Worda; kod pola po separatorze jest pomijany.
        paragraph._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" DATE "><w:r><w:t>1 maja</w:t></w:r>'
                                      '</w:fldSimple>'))
        paragraph._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" AUTHOR "/>'))  # pole bez wyniku
        add_complex_field(paragraph, 'TIME', ' 12:00')
        paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
        for body in ('<w:instrText>AUTHOR</w:instrText>', '<w:fldChar w:fldCharType="separate"/>',
                     '<w:instrText>zbędny kod</w:instrText>', '<w:t xml:space="preserve"> Anna</w:t>',
                     '<w:fldChar w:fldCharType="end"/>'):
            paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}>{body}</w:r>'))
        add_text_box(header.add_paragraph('Pole: '), 'w ramce')
        header._element.append(etree.Comment('komentarz w XML'))
        table = header.add_table(rows=2, cols=2, width=document.sections[0].page_width)
        table.cell(0, 0).text, table.cell(0, 1).text = 'Lewa', 'Prawa'
        _, _, _, records = self.render(document, '<p>Tekst</p>')
        texts = records[0][1]
        self.assertEqual(texts[0], 'e-mail dalej 1 maja 12:00 Anna')
        self.assertEqual(texts[1:3], ['Pole: ', 'w ramce'])
        self.assertEqual(texts[3], 'Lewa\t\tPrawa')  # dwie komórki: do lewej i do prawej
        self.assertEqual(len(texts), 4)  # pusty wiersz tabeli jest pomijany

    def test_header_taller_than_the_page_does_not_move_the_margins(self):
        document = Document()
        header = document.sections[0].header
        for number in range(46):  # 47 wierszy: jeszcze mieszczą się na stronie Letter szablonu python-docx
            header.add_paragraph(f'Wiersz {number}')
        _, _, pages, records = self.render(document, '<p>Tekst</p>')
        self.assertEqual((pages, len(records[0][1])), (1, 47))
        # Nagłówek i stopka zajęłyby całą stronę: treść zaczyna się na zwykłym marginesie (25,4 mm), jak w Wordzie.
        self.assertAlmostEqual(self.captured[0][2], 25.4, places=1)

    def test_header_taller_than_the_whole_page_still_gives_a_pdf(self):
        """Nagłówek dłuższy niż kartka (wcześniej nieskończona rekurencja i błąd konwersji): PDF powstaje,
        tekst trafia na pierwszą stronę, a nadmiar nagłówka wychodzi poza kartkę."""
        document = Document()
        header = document.sections[0].header
        for number in range(80):
            header.add_paragraph(f'Wiersz {number}')
        _, _, pages, records = self.render(document, '<p>Tekst</p>')
        self.assertEqual((pages, len(records[0][1])), (1, 81))

    def test_external_links_stay_and_tables_in_lists_are_blocks(self):
        tree, _, _, _ = self.render(Document(), '<p><a href="https://example.test/">strona</a></p>'
                                                '<ul><li><table><tr><td>W liście</td></tr></table></li></ul>')
        self.assertEqual(tree.xpath('.//a/@href'), ['https://example.test/'])
        self.assertEqual([child.tag for child in tree], ['p', 'table'])

    def test_font_size_of_the_first_run_sets_the_paragraph_size(self):
        from docx.shared import Pt
        document = Document()
        document.add_paragraph().add_run('Duży').font.size = Pt(20)
        tree, _, _, _ = self.render(document, '<p>Duży</p>')
        self.assertEqual(tree[0].get('data-size'), '20.0')

    def test_scene_break_with_markup_outside_the_font(self):
        tree, _, _, _ = self.render(Document(), '<p>A</p><p><em>❦ ❦</em></p><p>B</p>')
        separator = tree[1]
        self.assertEqual((separator.text, len(separator)), ('* * *', 0))

    def test_page_break_at_the_top_of_a_page_and_heading_at_the_bottom(self):
        """Podział przed pierwszym blokiem nie tworzy pustej strony; nagłówek bez miejsca na dole zaczyna nową."""
        # Strona Letter szablonu python-docx mieści 36 wierszy po 18 pt: nagłówek nie ma już miejsca.
        content = '<p data-page-break="1">Start</p>' + '<p>Wiersz</p>' * 35 + '<h1>Rozdział</h1>'
        _, _, pages, _ = self.render(Document(), content)
        self.assertEqual(pages, 2)
        self.assertEqual([page for markup, page, _ in self.captured if markup.startswith('<h1')], [2])
        self.assertEqual(self.captured[0][1:], (1, 25.4))

    def test_many_tables_and_tall_tables_after_a_page_break(self):
        tall = '<table><tr><td>' + 'Bardzo długa komórka tabeli. ' * 450 + '</td></tr></table>'
        content = ('<table><tr><td>Mała</td></tr></table>' * 51 + '<p>Przed<br class="page-break"></p>'
                   + tall + tall)
        tree, notes, pages, _ = self.render(Document(), content)
        self.assertEqual(notes.count(TALL_TABLE_NOTE), 1)
        self.assertEqual(len(tree.xpath('./table')), 51)
        tall_paragraphs = [node for node in tree if node.tag == 'p' and 'Bardzo długa' in node.text_content()]
        self.assertEqual(len(tall_paragraphs), 2)
        self.assertEqual(tall_paragraphs[0].get('data-page-break'), '1')
        self.assertGreater(pages, 2)
