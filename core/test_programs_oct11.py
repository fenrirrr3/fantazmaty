"""Programy (11.10): reguły Odkurzacza, wszystkie części dokumentu, kolejka,
usuwanie wyniku po pobraniu, uwagi konwersji na stronie i rozdziały EPUB."""
import json
import threading
import time
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from docx import Document

from core.services import program_jobs as jobs
from core.services.document_conversion_worker import polish_warnings, split_chapters
from core.services.document_converter import ConversionError, conversion_slot
from core.services.odkurzacz import (
    DEFAULT_EDITORIAL_RULES, _edit_opcodes, clean_docx, correct_editorial_text,
)

FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">{}</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
)


def with_footnote(document, text):
    """DOCX z przypisem dolnym (python-docx nie umie ich tworzyć)."""
    raw = BytesIO()
    document.save(raw)
    output = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdFoot" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            target.writestr(item, data)
        target.writestr('word/footnotes.xml', FOOTNOTES.format(text))
    output.seek(0)
    return output


def footnote_text(stream):
    with ZipFile(stream) as archive:
        return archive.read('word/footnotes.xml').decode()


class EditorialRuleTests(SimpleTestCase):
    def fix(self, text):
        return correct_editorial_text(text, DEFAULT_EDITORIAL_RULES)

    def test_period_stays_inside_quote_when_the_sentence_continues(self):
        self.assertEqual(self.fix('Napisał „tak.” i wyszedł.'), 'Napisał „tak.” i wyszedł.')
        self.assertEqual(self.fix('Powiedział „tak.”'), 'Powiedział „tak”.')
        self.assertEqual(self.fix('Napisał „tak.” Potem wyszedł.'), 'Napisał „tak”. Potem wyszedł.')

    def test_letter_marks_after_numbers_are_not_units(self):
        self.assertEqual(self.fix('Klasy 2B i 3A wyszły.'), 'Klasy 2B i 3A wyszły.')
        self.assertEqual(self.fix('Telewizor 4K.'), 'Telewizor 4K.')
        self.assertEqual(self.fix('Kabel 2m, 5kg i 20 K.'), 'Kabel 2 m, 5 kg i 20 K.')

    def test_gender_forms_in_brackets_stay_joined(self):
        self.assertEqual(self.fix('Drogi(a) Kliencie, zrobił(a) to.'), 'Drogi(a) Kliencie, zrobił(a) to.')
        self.assertEqual(self.fix('Pies(kundel) szczekał.'), 'Pies (kundel) szczekał.')

    def test_score_keeps_spaces_around_colon(self):
        self.assertEqual(self.fix('Wynik 3 : 1 , a potem :'), 'Wynik 3 : 1, a potem:')

    def test_edit_opcodes_handle_the_longest_paragraph_quickly(self):
        source = ('Ala  ma kota a kot ma Alę... ' * 700)[:20000]
        corrected = self.fix(source)
        started = time.monotonic()
        opcodes = _edit_opcodes(source, corrected)
        self.assertLess(time.monotonic() - started, 2)
        rebuilt = ''.join(source[a:b] if tag == 'equal' else corrected[c:d]
                          for tag, a, b, c, d in opcodes if tag != 'delete')
        self.assertEqual(rebuilt, corrected)


class WholeDocumentTests(SimpleTestCase):
    def document(self):
        document = Document()
        document.add_paragraph('Akapit  główny.')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Lampa świeciła,  lampa gasła.'
        document.sections[0].header.paragraphs[0].text = 'Nagłówek  strony'
        return with_footnote(document, 'Przypis  z lampą i lampą.')

    def test_cleaner_covers_tables_headers_and_footnotes(self):
        result = clean_docx(self.document(), sorted(DEFAULT_EDITORIAL_RULES))
        document = Document(BytesIO(result.getvalue()))
        self.assertEqual(document.paragraphs[0].text, 'Akapit główny.')
        self.assertEqual(document.tables[0].cell(0, 0).text, 'Lampa świeciła, lampa gasła.')
        self.assertEqual(document.sections[0].header.paragraphs[0].text, 'Nagłówek strony')
        self.assertIn('Przypis z lampą i lampą.', footnote_text(BytesIO(result.getvalue())))

    def test_repetitions_are_marked_in_tables_and_footnotes(self):
        from core.services.document_repetitions import color_document
        result = color_document(self.document())
        document = Document(BytesIO(result.getvalue()))
        cell = document.tables[0].cell(0, 0).paragraphs[0]
        self.assertTrue(cell._p.xpath('.//w:color'))
        self.assertEqual(cell.text, 'Lampa świeciła,  lampa gasła.')
        self.assertIn('w:color', footnote_text(BytesIO(result.getvalue())))


class ConverterPresentationTests(SimpleTestCase):
    def test_chapters_follow_the_repeated_heading_level(self):
        content = ('<h1 id="section-1">Tytuł</h1><p>Wstęp</p><h2 id="section-2">Rozdział 1</h2><p>A</p>'
                   '<h3 id="section-3">Część</h3><p>B</p><h2 id="section-4">Rozdział 2</h2><p>C</p>')
        chapters = split_chapters(content, 'Dokument')
        self.assertEqual([title for title, _, _ in chapters], ['Dokument', 'Rozdział 1', 'Rozdział 2'])
        self.assertEqual([anchor for anchor, _ in chapters[1][2]], ['section-2', 'section-3'])

    def test_document_without_headings_is_one_chapter(self):
        self.assertEqual(len(split_chapters('<p>Jeden</p><p>Dwa</p>', 'Dokument')), 1)

    def test_warnings_are_polish_and_grouped(self):
        warnings = polish_warnings([
            "Unrecognised paragraph style: Cytat (Style ID: Cytat)",
            "Unrecognised paragraph style: Cytat (Style ID: Cytat)",
            "Something new",
        ])
        self.assertEqual(warnings[0], 'Styl akapitu „Cytat” nie ma odpowiednika w e-booku – tekst zachowano '
                                      'w zwykłym formacie. (2 razy)')
        self.assertEqual(warnings[1], 'Uwaga konwertera: Something new')


class ConversionQueueTests(SimpleTestCase):
    def test_slot_waits_for_the_other_document(self):
        with TemporaryDirectory() as folder, self.settings(DOCUMENT_CONVERSION_DIR=Path(folder)):
            released = threading.Event()

            def hold():
                with conversion_slot():
                    released.wait(5)

            holder = threading.Thread(target=hold)
            holder.start()
            time.sleep(0.2)
            with self.assertRaises(ConversionError):
                with conversion_slot():
                    pass
            calls = []

            def on_wait():
                calls.append(1)
                released.set()

            with conversion_slot(wait=10, on_wait=on_wait):
                self.assertTrue(calls)
            holder.join()

    def test_waiting_can_be_cancelled(self):
        with TemporaryDirectory() as folder, self.settings(DOCUMENT_CONVERSION_DIR=Path(folder)):
            stop = threading.Event()

            def hold():
                with conversion_slot():
                    stop.wait(5)

            holder = threading.Thread(target=hold)
            holder.start()
            time.sleep(0.2)

            def cancel():
                raise jobs.JobCancelled()

            with self.assertRaises(jobs.JobCancelled):
                with conversion_slot(wait=10, on_wait=cancel):
                    pass
            stop.set()
            holder.join()

    def test_queued_job_with_recent_heartbeat_is_alive(self):
        now = time.time()
        self.assertTrue(jobs.alive({'state': 'queued', 'started': now - 500, 'heartbeat': now}))
        self.assertFalse(jobs.alive({'state': 'running', 'started': now - 500}))


class DownloadRemovalTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        override = self.settings(DOCUMENT_CONVERSION_DIR=Path(self.temp.name))
        override.enable()
        self.addCleanup(override.disable)
        self.client.force_login(get_user_model().objects.create_superuser('downloader', '', 'test'))

    def test_result_and_job_are_removed_after_download(self):
        output = BytesIO()
        Document().save(output)
        upload = BytesIO(output.getvalue())
        upload.name = 'Plik.docx'
        from django.core.files.uploadedfile import SimpleUploadedFile
        with patch.object(jobs, 'launch'):
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'clean', 'document': SimpleUploadedFile('Plik.docx', output.getvalue()),
            }, HTTP_X_PROGRAM_JOB='1')
        url = response.json()['url']
        folder = next(jobs.root().iterdir())
        (folder / 'result').write_bytes(b'wynik')
        jobs.write_json(folder / 'state.json', {
            'state': 'done', 'filename': 'Plik_odkurzony.docx', 'mime': 'application/octet-stream',
            'warnings': ['Uwaga konwertera: test'],
        })
        self.assertContains(self.client.get(url + '?view=page'), 'Uwaga konwertera: test')
        self.assertEqual(self.client.get(url).json()['warnings'], ['Uwaga konwertera: test'])
        download = self.client.get(url + '?download=1')
        self.assertEqual(b''.join(download.streaming_content), b'wynik')
        download.close()
        self.assertFalse(folder.exists())
        again = self.client.get(url + '?download=1')
        self.assertEqual(again.status_code, 404)
        self.assertContains(again, 'Plik został już pobrany i usunięty z serwera', status_code=404)

    def test_old_state_with_a_warning_count_still_renders(self):
        folder = jobs.root() / ('a' * 32)
        folder.mkdir(parents=True)
        jobs.write_json(folder / 'state.json', {'state': 'done', 'warnings': 2})
        self.assertEqual(jobs.status(folder)['warnings'], [])
        json.dumps(jobs.status(folder))
