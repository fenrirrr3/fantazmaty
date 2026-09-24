from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from docx import Document
from docx.shared import RGBColor

from core.services.odkurzacz import clean_docx, EDITORIAL_RULES, correct_editorial_text
from core.odkurzacz_forms import OdkurzaczForm
from people.models import Person


def docx_bytes(text='Ala  ma kota...'):
    document = Document()
    document.add_paragraph(text)
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def upload(data=None, name='tekst.docx'):
    return SimpleUploadedFile(name, docx_bytes() if data is None else data)


class OdkurzaczTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = get_user_model().objects.create_user('member', 'member@example.com')
        Person.objects.create(user=cls.member, first_name='Jan', last_name='Testowy', email=cls.member.email)
        cls.outsider = get_user_model().objects.create_user('outsider', 'outsider@example.com')

    def setUp(self):
        self.url = reverse('core:programs')
        self.client.force_login(self.member)

    def test_page_has_all_rules_and_no_placeholder(self):
        response = self.client.get(self.url)
        self.assertContains(response, 'Odkurzacz')
        self.assertContains(response, 'name="rules"', count=25)
        self.assertEqual(response.content.decode().count(' checked'), 25)
        self.assertNotContains(response, 'gifrific')

    def test_download_uses_selected_rules(self):
        response = self.client.post(self.url, {'document': upload(), 'rules': ['spaces']})
        self.assertEqual(response.status_code, 200)
        self.assertIn('tekst_odkurzony.docx', response['Content-Disposition'])
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        result = Document(BytesIO(b''.join(response.streaming_content)))
        response.close()
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota...')

    def test_no_rules_preserves_text_and_formatting(self):
        document = Document()
        run = document.add_paragraph().add_run('Ala  ma kota...')
        run.bold = True
        run.underline = True
        run.font.color.rgb = RGBColor(12, 34, 56)
        stream = BytesIO()
        document.save(stream)
        response = self.client.post(self.url, {'document': upload(stream.getvalue())})
        result = Document(BytesIO(b''.join(response.streaming_content)))
        response.close()
        run = result.paragraphs[0].runs[0]
        self.assertEqual(run.text, 'Ala  ma kota...')
        self.assertTrue(run.bold)
        self.assertTrue(run.underline)
        self.assertEqual(run.font.color.rgb, RGBColor(12, 34, 56))

    def test_correction_across_runs_preserves_other_parts(self):
        document = Document()
        para = document.add_paragraph()
        para.add_run('Ala ').bold = True
        para.add_run(' ma kota...').underline = True
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Bez  zmian...'
        document.sections[0].header.paragraphs[0].text = 'Nagłówek  bez zmian...'
        stream = BytesIO()
        document.save(stream)
        result = Document(clean_docx(stream, ['spaces', 'ellipsis']))
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota…')
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].underline)
        self.assertEqual(result.tables[0].cell(0, 0).text, 'Bez  zmian...')
        self.assertEqual(result.sections[0].header.paragraphs[0].text, 'Nagłówek  bez zmian...')
        self.assertFalse(result.paragraphs[0]._p.xpath('.//w:color | .//w:highlight'))

    def test_all_rules_run(self):
        result = Document(clean_docx(BytesIO(docx_bytes('  Ala  ma kota...  ')), [key for key, _ in EDITORIAL_RULES]))
        self.assertEqual(result.paragraphs[0].text, 'Ala ma kota…')

    def test_access_and_csrf(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.assertEqual(self.client.post(self.url, {'document': upload()}).status_code, 302)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.post(self.url, {'document': upload()}).status_code, 403)
        secure_client = Client(enforce_csrf_checks=True)
        secure_client.force_login(self.member)
        self.assertEqual(secure_client.post(self.url, {'document': upload()}).status_code, 403)

    def test_invalid_uploads_and_rules(self):
        for data in ({}, {'document': upload(b'bad')}, {'document': upload(name='tekst.txt')},
                     {'document': upload(), 'rules': ['unknown']}):
            with self.subTest(keys=list(data)):
                response = self.client.post(self.url, data)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['form'].errors)
                self.assertFalse(response.streaming)

    def test_zip_expansion_limit(self):
        stream = BytesIO()
        with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', b'a' * (51 * 1024 * 1024))
            archive.writestr('[Content_Types].xml', b'')
        form = OdkurzaczForm({}, {'document': upload(stream.getvalue())})
        self.assertFalse(form.is_valid())
        self.assertIn('50 MB', str(form.errors))

    def test_failure_shows_form_and_retains_selection(self):
        with patch('core.views.programs.clean_docx', side_effect=ValueError('bad package')):
            response = self.client.post(self.url, {'document': upload(), 'rules': ['spaces']})
        self.assertContains(response, 'Nie udało się przetworzyć')
        self.assertEqual(response.context['form']['rules'].value(), ['spaces'])

    def test_oversize_text_is_rejected(self):
        with self.assertRaises(ValueError):
            clean_docx(BytesIO(docx_bytes('a' * 20001)), ['spaces'])

    def test_tabs_are_deleted_without_inserting_spaces(self):
        text = '\tAla\tma \t kota\t\t'
        result = Document(clean_docx(BytesIO(docx_bytes(text)), ['tabs']))
        self.assertEqual(result.paragraphs[0].text, 'Alama  kota')
        self.assertFalse(result.paragraphs[0]._p.xpath('.//w:tab'))
        unchanged = Document(clean_docx(BytesIO(docx_bytes(text)), []))
        self.assertEqual(unchanged.paragraphs[0].text, text)

    def test_removed_rules_and_no_inserted_nbsp(self):
        removed = {'single_nbsp', 'unit_nbsp', 'reference_nbsp', 'dates_times', 'decimal', 'digit_groups'}
        self.assertFalse(removed & {key for key, _ in EDITORIAL_RULES})
        source = '2025-1-2 2/1/2025 o 9.30 9:30 12.50 kg 123456 kg A.Kowalski 20°C'
        result = correct_editorial_text(source)
        self.assertEqual(result, '2025-1-2 2/1/2025 o 9.30 9:30 12.50 kg 123456 kg A. Kowalski 20 °C')
        self.assertNotIn('\u00a0', result)
        self.assertNotIn('\u202f', result)
        for key in removed:
            with self.assertRaises(ValueError):
                correct_editorial_text('Tekst', [key])

    def test_submit_above_rules(self):
        html = self.client.get(self.url).content.decode()
        self.assertLess(html.index('Odkurz i pobierz DOCX'), html.index('<legend>Opcje korekty'))
        self.assertIn('odkurzacz-upload', html)

    def test_dialogue_spaces_survive_internal_docx_markers(self):
        from docx.oxml import OxmlElement
        cases = (
            ('viraptorka', ' – pochwalił.'),
            ('rozumiemy.', ' – W jej głos'),
            ('pszczół.', ' – Roześmiała się'),
            ('viraptorka – ', 'pochwalił.'),
        )
        for marker in ('proofErr', 'bookmarkStart', 'bookmarkEnd'):
            for left, right in cases:
                with self.subTest(marker=marker, left=left):
                    document = Document()
                    paragraph = document.add_paragraph()
                    paragraph.add_run(left).bold = True
                    paragraph._p.append(OxmlElement('w:' + marker))
                    paragraph.add_run(right).italic = True
                    source = BytesIO()
                    document.save(source)
                    result = Document(clean_docx(source, [key for key, _ in EDITORIAL_RULES]))
                    self.assertEqual(result.paragraphs[0].text, left + right)
                    self.assertTrue(result.paragraphs[0].runs[0].bold)
                    self.assertTrue(result.paragraphs[0].runs[-1].italic)

    def test_partial_fragment_is_not_a_new_dialogue(self):
        rules = ['hyphen_dash', 'dash_spaces', 'trim']
        self.assertEqual(correct_editorial_text(' – narracja', rules, trim_start=False), ' – narracja')
        self.assertEqual(correct_editorial_text(' - narracja', rules, trim_start=False), ' – narracja')
        self.assertEqual(correct_editorial_text('tekst – ', rules, trim_end=False), 'tekst – ')
        self.assertEqual(correct_editorial_text('  –Hę?', rules), '– Hę?')
