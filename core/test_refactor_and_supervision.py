"""Regression coverage for preparation, review decisions and anthology tools."""
from io import BytesIO
from itertools import product
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, SimpleTestCase
from django.urls import reverse, resolve, Resolver404
from people.models import Person, Role
from texts.models import Anthology, AnthologyTask, Review, Text, ReviewAssignment
from authors.models import Author
from core.services.reviews import change_review_status, assign_reviewer
from core.supervision import all_duplicates, anthology_checklist, anthology_credit_groups
from core.services.document_preparation import prepare_docx, ALL_EDITORIAL_RULES
from core.services.document_rebuild import rebuild_docx, inspect_docx, RebuildUnsupported
from core.services.document_formatting import normalize_docx
from core.services.odkurzacz import clean_docx
from workflow.completed_import import import_completed_workflow


class TeamToolsTests(TestCase):
    def member(self, name, role, *, first=None):
        user = get_user_model().objects.create_user(name, name+'@example.com')
        person = Person.objects.create(user=user, email=user.email, first_name=first or name, last_name='Test')
        person.roles.add(Role.objects.get_or_create(name=role)[0])
        return user, person

    def setUp(self):
        self.tmp = TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        settings = self.settings(ACTIVITY_SPOOL_DIR=self.tmp.name); settings.enable(); self.addCleanup(settings.disable)
        self.admin = get_user_model().objects.create_superuser('admin', 'admin@example.com', 'unused-password')
        self.coord, self.person = self.member('coord', 'Koordynator recenzji')
        self.reader, self.reader_person = self.member('reader', 'Recenzent')
        self.other, _ = self.member('other', 'Koordynator korekty')
        self.book = Anthology.objects.create(title='Antologia')
        self.review = Review.objects.create(title='Opowieść', anthology=self.book, author_first_name='Jan', author_last_name='Autor', email='jan@example.com', genre='fantasy', length=1000)

    def token(self, url):
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        return response.wsgi_request.edit_version_token

    def test_restore_status_tracks_assignments_without_closing_or_deleting_them(self):
        change_review_status(user=self.coord, review_id=self.review.pk, new_status='to_decide')
        change_review_status(user=self.coord, review_id=self.review.pk, new_status='restore')
        self.review.refresh_from_db(); self.assertEqual(self.review.status, 'new')
        change_review_status(user=self.coord, review_id=self.review.pk, new_status='to_decide')
        assign_reviewer(user=self.reader, review_id=self.review.pk)
        for user in (self.reader, self.other):
            with self.assertRaises(PermissionDenied):
                change_review_status(user=user, review_id=self.review.pk, new_status='restore')
        change_review_status(user=self.coord, review_id=self.review.pk, new_status='restore')
        self.review.refresh_from_db(); self.assertEqual(self.review.status, 'in_review')
        self.assertEqual(self.review.assignments.count(), 1)
        self.assertIsNone(self.review.decision_at)
        with self.assertRaises(ValidationError):
            change_review_status(user=self.coord, review_id=self.review.pk, new_status='restore')

    def test_restore_through_form_and_button_location(self):
        self.client.force_login(self.coord)
        url=reverse('core:assigned_review_detail', args=[self.review.pk])
        page=self.client.get(url)
        self.assertLess(page.content.index(b'Kierowanie do decyzji'), page.content.index(b'Informacje o'))
        endpoint=reverse('core:update_review_status', args=[self.review.pk])
        for status in ('to_decide', 'restore'):
            response=self.client.post(endpoint, {'status':status, '_edit_version':self.token(url)})
            self.assertEqual(response.status_code,302)
        self.review.refresh_from_db();self.assertEqual(self.review.status,'new')

    def test_duplicate_scan_runs_only_on_click_and_accepts_all_scope(self):
        self.client.force_login(self.admin)
        url=reverse('core:data_integrity')
        with patch('core.views.supervision.all_duplicates', return_value=[]) as scan:
            for query in ('?tab=duplicates', '?tab=duplicates&anthology=all'):
                self.assertEqual(self.client.get(url+query).status_code,200)
            scan.assert_not_called()
            self.assertEqual(self.client.get(url+'?tab=duplicates&anthology=all&run=1').status_code,200)
            scan.assert_called_once_with()

    def test_duplicate_all_checks_each_anthology_without_cross_anthology_matches(self):
        author=Author.objects.create(first_name='Jan',last_name='Autor',email='jan@example.com')
        for book in (self.book, Anthology.objects.create(title='Druga')):
            for i in range(2):
                text=Text.objects.create(anthology=book,title='Ten sam tekst',length=1000)
                text.authors.add(author)
        self.assertEqual(len(all_duplicates()),2)
        self.assertEqual(len(all_duplicates(self.book.pk)),1)

    def task_data(self):
        data={}
        for kind, _ in AnthologyTask.TaskType.choices:
            data[kind+'-status']='ready'
            data[kind+'-assigned_to']=self.person.pk
        return data

    def test_checklist_does_not_check_contracts_and_tasks_save_as_one_operation(self):
        self.client.force_login(self.coord)
        url=reverse('core:anthology_detail',args=[self.book.pk])
        data=self.task_data();data['_edit_version']=self.token(url)
        data['blurb-assigned_to']=''
        self.assertEqual(self.client.post(url,data).status_code,400)
        self.assertFalse(AnthologyTask.objects.exclude(status='not_commissioned').exists())
        data=self.task_data();data['_edit_version']=self.token(url)
        self.assertEqual(self.client.post(url,data).status_code,302)
        self.assertEqual(AnthologyTask.objects.filter(status='ready').count(),3)
        author=Author.objects.create(first_name='Jan',last_name='Autor',email='jan@example.com',has_contract=False)
        text=Text.objects.create(title='Tekst',anthology=self.book,length=1000);text.authors.add(author)
        self.assertFalse(any('umow' in str(row).lower() for row in anthology_checklist(self.book)))
        self.assertContains(self.client.get(url),'Kontrola przed składem:')
        self.assertNotContains(self.client.get(url),'CSV')

    def test_task_permissions_and_stale_save(self):
        url=reverse('core:anthology_detail',args=[self.book.pk])
        self.client.force_login(self.reader)
        self.assertNotContains(self.client.get(url),'Zapisz zadania')
        self.assertEqual(self.client.post(url,self.task_data()).status_code,403)
        self.client.force_login(self.coord)
        token=self.token(url)
        task=AnthologyTask.objects.get(anthology=self.book,task_type='blurb')
        task.assigned_to=self.person; task.status='commissioned'; task.save()
        self.assertEqual(self.client.post(url,dict(self.task_data(),_edit_version=token)).status_code,409)
        self.assertEqual(AnthologyTask.objects.count(),3)

    def test_csv_route_removed(self):
        with self.assertRaises(Resolver404):resolve(f'/antologie/{self.book.pk}/stopka.csv')

    def test_grouped_credits_include_imported_work_and_deduplicate(self):
        text=Text.objects.create(title='Tekst',anthology=self.book,length=1000)
        import_completed_workflow(text_id=text.pk, next_stage='ready', stages=[
            {'stage_type':kind,'assigned_to':self.reader.email}
            for kind in ('editing','editing_review','first_proofreading','second_proofreading','fourth_verification')])
        self.review.status='accepted';self.review.copied_text=text;self.review.save()
        ReviewAssignment.objects.create(review=self.review,user=self.reader,position=1,opinion='yes')
        groups={row['label']:row['names'] for row in anthology_credit_groups(self.book)}
        for label in ('Redakcja','Kontrola redakcji','Korekta','Weryfikacja','Recenzje'):
            self.assertEqual(groups[label],[str(self.reader_person)])
        self.assertEqual(groups['Kontrola przed składem'],[])

    def test_role_page_superuser_only_and_sorted_names(self):
        self.member('lukasz','Recenzent',first='Łukasz')
        self.member('zosia','Recenzent',first='Zofia')
        inactive, person=self.member('inactive','Recenzent');inactive.is_active=False;inactive.save()
        role=Role.objects.get(name='Recenzent')
        url=reverse('core:role_names')+'?role='+str(role.pk)
        self.client.force_login(self.coord);self.assertEqual(self.client.get(url).status_code,403)
        self.client.force_login(self.admin);page=self.client.get(url)
        self.assertContains(page,'Łukasz Test, reader Test, Zofia Test')
        self.assertNotContains(page,'inactive Test')
        self.assertNotContains(page,self.reader.email)


class PreparationEquivalenceTests(SimpleTestCase):
    def source(self):
        doc=Document();doc.add_heading('Rozdział',1)
        p=doc.add_paragraph('Tekst  testowy. Dalej.');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.runs[0].italic=True;p.runs[0].bold=True
        doc.add_paragraph();doc.add_paragraph('Środek').alignment=WD_ALIGN_PARAGRAPH.CENTER
        source=BytesIO();doc.save(source);source.seek(0);return source

    def test_all_option_combinations_match_previous_sequential_pipeline(self):
        for rebuild,normalize,clean in product((False,True),repeat=3):
            with self.subTest(rebuild=rebuild,normalize=normalize,clean=clean):
                source=self.source();payload=source.getvalue()
                rules=list(ALL_EDITORIAL_RULES) if clean else []
                # Previous worker sequence, compared with the new shared entry point.
                for enabled, operation, kwargs in (
                    (rebuild,rebuild_docx,{}),(normalize,normalize_docx,{}),(clean,clean_docx,{'rules':rules})):
                    if enabled:
                        with operation(BytesIO(payload),**kwargs) as result:payload=result.read()
                with prepare_docx(source,rebuild=rebuild,normalize_formatting=normalize,cleaner_rules=rules) as prepared, ZipFile(prepared) as actual, ZipFile(BytesIO(payload)) as expected:
                    for name in ('word/document.xml','word/styles.xml'):
                        self.assertEqual(actual.read(name),expected.read(name))

    def test_preflight_does_not_construct_document_and_reports_same_omissions(self):
        doc=Document();doc.add_paragraph('A');doc.add_table(rows=1,cols=1)
        source=BytesIO();doc.save(source)
        with patch('core.services.document_rebuild._new_document',side_effect=AssertionError('must not create output')):
            with self.assertRaises(RebuildUnsupported) as first:inspect_docx(source)
            with self.assertRaises(RebuildUnsupported) as second:rebuild_docx(source)
        self.assertEqual(first.exception.omissions,second.exception.omissions)
        with patch('core.services.document_rebuild._new_document',side_effect=AssertionError('must not create output')):
            inspect_docx(self.source())
