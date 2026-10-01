from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from django.test import SimpleTestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from core.services.document_rebuild import rebuild_docx, RebuildUnsupported
from core.services.document_formatting import normalize_docx
from core.services.document_converter import convert_document, ConversionError
from core.services.mailbox_import import package_messages


def saved(doc):
    out = BytesIO(); doc.save(out); out.seek(0); return out


class RebuildTests(SimpleTestCase):
    def test_text_format_empty_paragraphs_and_metadata(self):
        doc = Document(); doc.core_properties.author = 'Secret'; doc.core_properties.title = 'Private'
        doc.sections[0].header.paragraphs[0].text = 'Private header'
        doc.styles['Normal'].font.italic = True
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        r = p.add_run('Tekst'); r.bold = True
        doc.add_paragraph(); doc.add_paragraph()
        doc.add_paragraph('Środek').alignment = WD_ALIGN_PARAGRAPH.CENTER
        with rebuild_docx(saved(doc), allow_omissions=True) as out:
            result = Document(out)
            self.assertEqual([p.text for p in result.paragraphs], ['Tekst', '', '', 'Środek'])
            self.assertTrue(result.paragraphs[0].runs[0].bold)
            self.assertTrue(result.paragraphs[0].runs[0].italic)
            self.assertEqual(result.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertEqual(result.paragraphs[-1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(result.core_properties.author, '')
            self.assertIsNone(result.core_properties.created)
            with ZipFile(out) as archive:
                self.assertNotIn('word/header1.xml', archive.namelist())
                self.assertNotIn(b'Secret', archive.read('docProps/core.xml'))
            self.assertAlmostEqual(result.sections[0].page_width.cm, 21, places=2)
            self.assertAlmostEqual(result.sections[0].left_margin.cm, 2.5, places=2)

    def test_revisions_and_hidden_text(self):
        doc = Document(); p = doc.add_paragraph('A')
        p.add_run('hidden').font.hidden = True
        for tag, text in [('ins','B'), ('del','C')]:
            wrapper = OxmlElement('w:' + tag); run = OxmlElement('w:r')
            t = OxmlElement('w:t' if tag == 'ins' else 'w:delText'); t.text = text
            run.append(t); wrapper.append(run); p._p.append(wrapper)
        self.assertEqual(Document(rebuild_docx(saved(doc))).paragraphs[0].text, 'AB')

    def test_tables_and_numbering_fail_explicitly(self):
        for kind in ('table','list'):
            doc = Document()
            if kind == 'table': doc.add_table(rows=1, cols=1)
            else: doc.add_paragraph('item', 'List Number')
            with self.assertRaises(RebuildUnsupported): rebuild_docx(saved(doc))

    def test_page_normalization_all_sections(self):
        doc = Document(); doc.add_paragraph('A'); doc.add_section()
        for section in doc.sections:
            section.page_width = Cm(40); section.page_height = Cm(20)
            section.top_margin = section.bottom_margin = Cm(1)
        result = Document(normalize_docx(saved(doc)))
        for section in result.sections:
            self.assertAlmostEqual(section.page_width.cm, 21, places=2)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=2)
            for name in ('left_margin','right_margin','top_margin','bottom_margin'):
                self.assertAlmostEqual(getattr(section,name).cm, 2.5, places=2)

    def test_real_worker_docx_and_mail_default(self):
        doc = Document(); doc.add_paragraph('Tekst'); doc.core_properties.author = 'Secret'
        payload = saved(doc).getvalue()
        row = {'uid': 1, 'folder': 'Story', 'files': [('story.docx', payload)]}
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            out, extension, mime = convert_document(SimpleUploadedFile('a.docx', payload), [], include_docx=True, rebuild=True, normalize=False)
            with out:
                self.assertEqual(extension, 'docx')
                self.assertIn('wordprocessingml', mime)
                self.assertEqual(Document(out).core_properties.author, '')
            for rebuild in (True,False):
                with package_messages([row], clean=False, convert=False, rebuild=rebuild) as archive, ZipFile(archive) as z:
                    content = z.read(z.namelist()[0])
                    if rebuild: self.assertEqual(Document(BytesIO(content)).core_properties.author, '')
                    else: self.assertEqual(content, payload)


from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse

class RebuildViewTests(TestCase):
    def test_button_download_and_unsupported_error(self):
        user = get_user_model().objects.create_superuser('rebuild', 'rebuild@example.com', 'testpassword')
        self.client.force_login(user)
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            for table in (False,True):
                doc = Document(); doc.add_paragraph('Tekst'); doc.core_properties.author = 'Private'
                if table: doc.add_table(rows=1, cols=1)
                original_payload = saved(doc).getvalue()
                response = self.client.post(reverse('core:programs'), {'program_action':'clean', 'rebuild':'on', 'document':SimpleUploadedFile('story.docx', original_payload)})
                self.assertEqual(response.status_code, 200)
                if table:
                    self.assertContains(response, 'W nowym DOCX zostaną pominięte')
                    token = response.context['rebuild_token']
                    accepted = self.client.post(reverse('core:programs'), {'program_action':'clean', 'rebuild':'on', 'allow_rebuild_omissions':'on', 'rebuild_token':token, 'document':SimpleUploadedFile('story.docx', original_payload)})
                    self.assertEqual(accepted.status_code,200)
                    self.assertTrue(accepted.streaming)
                    rebuilt = Document(BytesIO(b''.join(accepted.streaming_content)))
                    self.assertEqual(len(rebuilt.tables),0)
                    self.assertEqual(rebuilt.core_properties.author,'')
                    accepted.close()
                else:
                    self.assertIn('_nowy.docx', response['Content-Disposition'])
                    result = Document(BytesIO(b''.join(response.streaming_content)))
                    self.assertEqual(result.core_properties.author, '')
                    self.assertEqual(result.paragraphs[0].text, 'Tekst')
                response.close()

class ConfirmOmissionsTests(SimpleTestCase):
    def test_confirmation_still_builds_new_docx(self):
        from core.services.document_converter import RebuildConfirmationRequired
        doc = Document(); doc.core_properties.author = 'Private'
        doc.add_paragraph('Zachowaj').runs[0].italic = True
        doc.add_table(rows=1,cols=1).cell(0,0).text = 'Pomiń tabelę'
        doc.add_paragraph('Lista', 'List Number')
        payload = saved(doc).getvalue()
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            with self.assertRaises(RebuildConfirmationRequired) as caught:
                convert_document(BytesIO(payload), [], include_docx=True, rebuild=True)
            self.assertIn('tabele wraz z ich treścią', str(caught.exception))
            self.assertIn('tekst pozycji pozostanie', str(caught.exception))
            out,ext,_ = convert_document(BytesIO(payload), [], include_docx=True, rebuild=True, allow_rebuild_omissions=True)
            with out:
                result = Document(out)
                self.assertEqual(ext,'docx'); self.assertEqual(result.core_properties.author,'')
                self.assertEqual(len(result.tables),0)
                self.assertEqual([p.text for p in result.paragraphs],['Zachowaj','Lista'])
                self.assertTrue(result.paragraphs[0].runs[0].italic)

    def test_batch_collects_warnings_for_all_files(self):
        from core.services.document_converter import RebuildConfirmationRequired
        rows = []
        for uid, kind in enumerate(('table','list'),1):
            doc=Document(); doc.add_paragraph('Tekst')
            if kind=='table':doc.add_table(rows=1,cols=1)
            else:doc.add_paragraph('Pozycja','List Number')
            rows.append({'uid':uid,'folder':kind,'files':[(kind+'.docx',saved(doc).getvalue())]})
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            with self.assertRaises(RebuildConfirmationRequired) as caught:
                package_messages(rows,clean=False,convert=False)
            self.assertIn('table.docx',str(caught.exception));self.assertIn('list.docx',str(caught.exception))
            with package_messages(rows,clean=False,convert=False,allow_rebuild_omissions=True) as out, ZipFile(out) as archive:
                self.assertEqual(len(archive.namelist()),2)

class CleanerCheckboxTests(TestCase):
    def test_checkbox_and_selected_rules(self):
        user = get_user_model().objects.create_superuser('checkbox','checkbox@example.com','testpassword')
        self.client.force_login(user)
        page = self.client.get(reverse('core:programs'))
        self.assertFalse(page.context['form'].fields['rebuild'].initial)
        self.assertNotContains(page, 'value="rebuild"')
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            for enabled in (False,True):
                doc=Document(); doc.add_paragraph('Zdanie. następne zdanie...'); doc.core_properties.author='Private'
                data={'program_action':'clean','rules':['sentence_case'], 'document':SimpleUploadedFile('a.docx',saved(doc).getvalue())}
                if enabled: data['rebuild']='on'
                response=self.client.post(reverse('core:programs'),data)
                self.assertEqual(response.status_code,200)
                result=Document(BytesIO(b''.join(response.streaming_content)))
                self.assertEqual(result.paragraphs[0].text,'Zdanie. Następne zdanie...')
                self.assertEqual(result.core_properties.author,'' if enabled else 'Private')
                response.close()
