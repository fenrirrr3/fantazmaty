from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client, SimpleTestCase
from django.urls import reverse
from docx import Document
from lxml import html

from core.services.document_converter import convert_document, RebuildConfirmationRequired, ConversionError
from core.services.pending_documents import SALT, TTL


def payload(text='Ala  ma kota.'):
    doc = Document()
    doc.add_paragraph(text)
    result = BytesIO()
    doc.save(result)
    return result.getvalue()


class PendingRebuildTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = get_user_model().objects.create_superuser('pending-owner', 'pending@example.test', 'test-only')
        cls.other = get_user_model().objects.create_superuser('pending-other', 'other@example.test', 'test-only')

    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        override = self.settings(DOCUMENT_CONVERSION_DIR=Path(self.temporary.name))
        override.enable()
        self.addCleanup(override.disable)
        self.client.force_login(self.owner)
        self.url = reverse('core:programs')

    def warn(self, name='oryginał.docx', content=None):
        with patch('core.views.programs.convert_document', side_effect=RebuildConfirmationRequired('W nowym DOCX zostaną pominięte: tabele.')):
            response = self.client.post(self.url, {'program_action': 'clean', 'rebuild': 'on',
                'normalize_formatting': 'on', 'remove_soft_whitespace': 'on', 'rules': ['spaces'],
                'document': SimpleUploadedFile(name, content or payload())})
        self.assertEqual(response.status_code, 200)
        token = response.context['rebuild_token']
        self.assertTrue(token)
        doc = html.fromstring(response.content.decode('utf-8'))
        form = doc.xpath('//form[@data-rebuild-confirmation]')[0]
        self.assertFalse(form.xpath('.//input[@type="file"]'))
        self.assertEqual(form.xpath('.//button/@value'), ['clean_confirm'])
        self.assertFalse(doc.xpath('//input[@name="document"]'))
        self.assertNotContains(response, 'Wybierz ponownie ten sam plik')
        return token

    def test_confirmation_needs_no_upload_and_uses_original_file_and_options(self):
        original = payload('Oryginalny tekst')
        token = self.warn(content=original)
        with patch('core.views.programs.convert_document', return_value=(BytesIO(payload('Wynik')), 'docx', 'application/octet-stream')) as converter:
            response = self.client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': token,
                'rules': ['trim'], 'document': SimpleUploadedFile('inny.docx', payload('Podmiana'))})
        self.assertTrue(response.streaming)
        self.assertIn('_nowy.docx', response['Content-Disposition'])
        b''.join(response.streaming_content)
        args, kwargs = converter.call_args
        self.assertEqual(args[0].read(), original)
        self.assertEqual(kwargs['cleaner_rules'], ['spaces'])
        self.assertTrue(kwargs['normalize'])
        self.assertTrue(kwargs['remove_soft_whitespace'])
        self.assertTrue(kwargs['allow_rebuild_omissions'])
        self.assertFalse(list(Path(self.temporary.name).glob('pending_uploads/*/*.docx')))
        with patch('core.views.programs.convert_document') as converter:
            again = self.client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': token})
            converter.assert_not_called()
        self.assertContains(again, 'Potwierdzenie wygasło')

    def test_token_is_bound_to_user_and_session_and_cannot_be_forged(self):
        token = self.warn()
        same_user_other_session = Client()
        same_user_other_session.force_login(self.owner)
        other = Client()
        other.force_login(self.other)
        for client, value in ((same_user_other_session, token), (other, token), (self.client, token + 'x'), (self.client, '')):
            with self.subTest(client=client, value=value[-10:]), patch('core.views.programs.convert_document') as converter:
                response = client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': value})
                converter.assert_not_called()
                self.assertContains(response, 'Potwierdzenie wygasło')

    def test_expired_and_modified_files_are_rejected(self):
        token = self.warn()
        with patch('django.core.signing.time.time', return_value=signing.time.time() + TTL + 2), patch('core.views.programs.convert_document') as converter:
            response = self.client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': token})
            converter.assert_not_called()
        self.assertContains(response, 'Potwierdzenie wygasło')
        token = self.warn()
        data = signing.loads(token, salt=SALT)
        path = Path(self.temporary.name) / 'pending_uploads' / str(self.owner.pk) / (data['key'] + '.docx')
        path.write_bytes(payload('Zmieniony plik'))
        with patch('core.views.programs.convert_document') as converter:
            response = self.client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': token})
            converter.assert_not_called()
        self.assertContains(response, 'Potwierdzenie wygasło')

    def test_busy_converter_keeps_confirmation_for_retry(self):
        token = self.warn()
        with patch('core.views.programs.convert_document', side_effect=ConversionError('Konwerter jest zajęty.')):
            response = self.client.post(self.url, {'program_action': 'clean_confirm', 'rebuild_token': token})
        self.assertContains(response, 'Konwerter jest zajęty.')
        self.assertContains(response, 'Akceptuję i kontynuuję')
        self.assertEqual(response.context['rebuild_token'], token)
        self.assertEqual(len(list(Path(self.temporary.name).glob('pending_uploads/*/*.docx'))), 1)


class LongCleanerDocumentTests(SimpleTestCase):
    def test_real_worker_cleans_567000_characters_without_total_limit(self):
        document = Document()
        paragraph = 'Ala ma kota. ' * 75  # 975 characters, below the per-paragraph guard.
        for _ in range(581):
            document.add_paragraph(paragraph)
        document.add_paragraph('A' * 525)
        self.assertEqual(sum(len(p.text) for p in document.paragraphs), 567000)
        source = BytesIO()
        document.save(source)
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            result, extension, _ = convert_document(source, [], include_docx=True, use_cleaner=True,
                normalize=False, cleaner_rules=['spaces'])
            with result:
                output = Document(result)
                self.assertEqual(extension, 'docx')
                self.assertEqual(sum(len(p.text) for p in output.paragraphs), 567000)

    def test_long_paragraph_has_specific_public_error(self):
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            with self.assertRaisesMessage(ConversionError, 'akapitów przekracza 20 000'):
                convert_document(BytesIO(payload('A' * 20001)), [], include_docx=True, use_cleaner=True, normalize=False)
