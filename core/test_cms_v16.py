from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from datetime import timedelta
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from core.program_test_support import run_program
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_BREAK
from lxml import html
from people.models import Person
from texts.models import Anthology, AnthologyTask, Text
from illustrations.models import Illustration, Illustrator
from illustrations.editing import edit_token
from workflow.tests import create_member
from core.services.document_preparation import prepare_docx


class CMS16Tests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin=get_user_model().objects.create_superuser('admin16','admin16@example.test','test')
        cls.artist=create_member('artist16','Ilustrator')
        cls.artist_contact=Illustrator.objects.create(first_name='Artysta',last_name='Kontakt')
        cls.book=Anthology.objects.create(title='Nowa antologia',has_illustrations=True)
        cls.text=Text.objects.create(title='Nowy tekst',anthology=cls.book,length=20)
        cls.ill=Illustration.objects.get(text=cls.text)

    def setUp(self):
        tmp=TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        settings = self.settings(
            DOCUMENT_CONVERSION_DIR=Path(tmp.name) / "convert",
            ACTIVITY_SPOOL_DIR=Path(tmp.name) / "activity",
        )
        settings.enable()
        self.addCleanup(settings.disable)
        self.client.force_login(self.admin)
        self.url = reverse("illustrations:illustration_detail", args=[self.ill.pk])

    def post_illustration(self, action, **values):
        self.ill.refresh_from_db()
        return self.client.post(
            self.url, {"action": action, "version": edit_token(self.admin, self.ill), **values}
        )

    def test_manual_name_email_without_profile_and_assignment_date(self):
        before = Person.objects.count()
        accounts = get_user_model().objects.count()
        response = self.post_illustration(
            "assignment",
            manual_illustrator_name="Anna Bez Konta",
            manual_illustrator_email="anna@example.test",
            status="assigned",
        )
        self.assertEqual(response.status_code, 302)
        self.ill.refresh_from_db()
        self.assertFalse(self.ill.illustrators.exists())
        self.assertEqual(self.ill.assigned_at, timezone.localdate())
        self.assertEqual(Person.objects.count(), before)
        self.assertEqual(get_user_model().objects.count(), accounts)
        self.assertContains(
            self.client.get(reverse("illustrations:illustration_list")), "Anna Bez Konta"
        )
        old = timezone.localdate() - timedelta(days=10)
        Illustration.objects.filter(pk=self.ill.pk).update(assigned_at=old)
        response = self.post_illustration(
            "assignment",
            manual_illustrator_name="Anna Bez Konta",
            manual_illustrator_email="anna@example.test",
            status="delivered",
        )
        self.assertEqual(response.status_code, 302)
        self.ill.refresh_from_db()
        self.assertEqual(self.ill.assigned_at, old)
        response = self.post_illustration(
            "assignment", manual_illustrator_name="Inna Osoba", status="assigned"
        )
        self.assertEqual(response.status_code, 302)
        self.ill.refresh_from_db()
        self.assertEqual(self.ill.assigned_at, timezone.localdate())

    def test_manual_validation_and_switch_back_to_profile(self):
        for data in (
            {"manual_illustrator_email": "x@example.test"},
            {"manual_illustrator_name": "Anna", "manual_illustrator_email": "wrong"},
            {"manual_illustrator_name": "Anna", "illustrators": self.artist_contact.pk},
        ):
            response = self.post_illustration("assignment", status="assigned", **data)
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.context["forms"]["assignment"].errors)
            self.ill.refresh_from_db()
            self.assertEqual(self.ill.status, "unassigned")
        response = self.post_illustration(
            "assignment", manual_illustrator_name="Anna", status="assigned"
        )
        self.assertEqual(response.status_code, 302)
        response = self.post_illustration(
            "assignment", illustrators=self.artist_contact.pk, status="assigned"
        )
        self.assertEqual(response.status_code, 302)
        self.ill.refresh_from_db()
        self.assertEqual(self.ill.manual_illustrator_name, "")
        self.assertEqual(list(self.ill.illustrators.all()), [self.artist_contact])

    def test_notes_permissions_conflict_and_search_markup(self):
        old = edit_token(self.admin, self.ill)
        response = self.post_illustration(
            "coordinator_notes", coordinator_notes="Proszę o ciemne tło."
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.post(
            self.url,
            {"action": "coordinator_notes", "version": old, "coordinator_notes": "stary zapis"},
        )
        self.assertEqual(response.status_code, 409)
        response = self.client.get(self.url)
        self.assertContains(response, 'id="illustrator-search"')
        self.assertContains(response, "illustrator_search.js")
        self.ill.refresh_from_db()
        self.ill.set_artists([self.artist_contact], status="assigned")
        self.client.force_login(self.artist)
        response = self.client.post(
            self.url,
            {
                "action": "coordinator_notes",
                "version": edit_token(self.artist, self.ill),
                "coordinator_notes": "nie wolno",
            },
        )
        self.assertEqual(response.status_code, 403)
        self.ill.refresh_from_db()
        self.assertEqual(self.ill.coordinator_notes, "Proszę o ciemne tło.")

    def test_tasks_and_audio_description_table_share_assignment(self):
        ready = Anthology.objects.create(title="Gotowa pomijana", status="ready")
        self.assertEqual(
            self.book.production_tasks.filter(task_type="audio_description").count(), 1
        )
        response = self.client.get(reverse("core:task_list"))
        self.assertEqual(response.status_code, 200)
        rows = list(response.context["tasks"])
        self.assertEqual(
            {r["name"] for r in rows},
            {"Blurb", "Banery", "Skład", "Okładka", "Audiodeskrypcja", "Typografia okładki"},
        )
        self.assertNotContains(response, ready.title)
        response = self.client.get(reverse("core:anthology_detail", args=[self.book.pk]))
        self.assertContains(response, "audio_description-status")
        payload = {f"{kind}-status": "not_commissioned" for kind in AnthologyTask.TaskType.values}
        payload.update(
            {
                "audio_description-status": "commissioned",
                "audio_description-assigned_to": self.artist.person_profile.pk,
            }
        )
        token = html.fromstring(response.content).xpath('//input[@name="_edit_version"]/@value')[0]
        response = self.client.post(
            reverse("core:anthology_detail", args=[self.book.pk]),
            {**payload, "_edit_version": token},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            self.book.production_tasks.get(task_type="audio_description").assigned_to_id,
            self.artist.person_profile.pk,
        )
        response = self.client.get(reverse("core:audio_descriptions"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<h1 class="page-title">Audiodeskrypcje</h1>', html=True)
        self.assertTrue(html.fromstring(response.content).xpath("//main//table"))
        rows = list(response.context["rows"])
        assigned = next(row for row in rows if row["book"].pk == self.book.pk)
        self.assertEqual(assigned["task"].assigned_to_id, self.artist.person_profile.pk)
        self.assertFalse(assigned["can_claim"])
        self.assertContains(response, ready.title)

    def test_original_verifier_form_save(self):
        book = Anthology.objects.create(title="Przekład", is_translated=True)
        text = Text.objects.create(title="Tekst obcy", anthology=book, length=10)
        response = self.client.get(reverse("core:translation_detail", args=[text.pk]))
        self.assertContains(response, "Weryfikacja z oryginałem")
        token = html.fromstring(response.content).xpath(
            '//form[contains(@action,"tlumacze")]/input[@name="_edit_version"]/@value'
        )[0]
        response = self.client.post(
            reverse("core:set_translators", args=[text.pk]),
            {"original_verifier": "Jan Sprawdzający", "_edit_version": token},
        )
        self.assertEqual(response.status_code, 302)
        text.translation.refresh_from_db()
        self.assertEqual(text.translation.original_verifier, "Jan Sprawdzający")

    def document(self):
        doc = Document()
        p = doc.add_paragraph()
        r = p.add_run("Ala\u00a0ma")
        r.bold = True
        r.add_break()
        r.add_text("kota\u202f!")
        p.add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("Drugi akapit")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Tabela\u00a0A\nB"
        doc.sections[0].header.paragraphs[0].text = "Nagłówek\u00a0A\nB"
        out = BytesIO()
        doc.save(out)
        return out.getvalue()

    def test_whitespace_keeps_paragraphs_page_breaks_formatting_and_cleans_tables(self):
        with prepare_docx(
            BytesIO(self.document()),
            use_cleaner=True,
            cleaner_rules=["single_nbsp"],
            remove_soft_whitespace=True,
        ) as stream:
            doc = Document(stream)
        self.assertEqual([p.text for p in doc.paragraphs[:2]], ["Ala ma", "kota !"])
        self.assertEqual(len(doc.paragraphs), 3)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertEqual(len(doc.element.xpath('//w:br[@w:type="page"]')), 1)
        self.assertEqual(
            len(doc.element.xpath('//w:br[not(@w:type) or @w:type="textWrapping"]')), 0
        )
        self.assertEqual([p.text for p in doc.tables[0].cell(0, 0).paragraphs], ["Tabela A", "B"])
        self.assertEqual([p.text for p in doc.sections[0].header.paragraphs], ["Nagłówek A", "B"])

    def test_http_cleaner_and_conversion_worker_apply_checkbox(self):
        response = run_program(
            self.client,
            {
                "document": SimpleUploadedFile("tekst.docx", self.document()),
                "remove_soft_whitespace": "on",
            },
        )
        self.assertEqual(response.status_code, 200)
        doc = Document(BytesIO(b"".join(response.streaming_content)))
        self.assertNotIn("\u00a0", doc.paragraphs[0].text)
        response = run_program(
            self.client,
            {
                "program_action": "convert",
                "convert-document": SimpleUploadedFile("tekst.docx", self.document()),
                "convert-formats": ["epub"],
                "convert-remove_soft_whitespace": "on",
            },
        )
        self.assertEqual(response.status_code, 200)
        with ZipFile(BytesIO(b"".join(response.streaming_content))) as archive:
            content = archive.read("EPUB/content.xhtml")
            self.assertEqual(
                [p.text_content().strip() for p in html.fromstring(content).xpath("//p")][:2],
                ["Ala ma", "kota !"],
            )
