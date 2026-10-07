from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import subprocess
import shutil
import time
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, SimpleTestCase
from django.urls import reverse
from docx import Document

from core.services import program_jobs as jobs
from core.services.document_converter import run_converter


def document(header=False):
    doc = Document()
    doc.add_paragraph('Ala  ma kota. Kot odwiedził ogród.')
    if header:
        doc.sections[0].header.paragraphs[0].text = 'Nagłówek'
    output = BytesIO(); doc.save(output)
    return SimpleUploadedFile('Zażółć.docx', output.getvalue())


class ProgramJobTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = get_user_model().objects.create_superuser('job-owner', '', 'test')
        cls.other = get_user_model().objects.create_superuser('job-other', '', 'test')

    def setUp(self):
        self.temp = TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        override = self.settings(DOCUMENT_CONVERSION_DIR=Path(self.temp.name))
        override.enable(); self.addCleanup(override.disable)
        self.client.force_login(self.owner)
        self.url = reverse('core:programs')

    def create(self, **extra):
        with patch.object(jobs, 'launch'):
            response = self.client.post(self.url, {'program_action': 'clean', 'document': document(), **extra}, HTTP_X_PROGRAM_JOB='1')
        self.assertEqual(response.status_code, 202, response.content)
        return response.json()['url'], next(jobs.root().iterdir())

    def test_all_three_forms_create_jobs_and_validate_before_launch(self):
        forms = [
            {'program_action': 'clean', 'document': document()},
            {'program_action': 'convert', 'convert-document': document(), 'convert-formats': ['pdf', 'epub']},
            {'program_action': 'repetitions', 'repetitions-document': document(),
             'repetitions-window_size': 35, 'repetitions-min_word_length': 4,
             'repetitions-sentence_limit': 35, 'repetitions-paragraph_limit': 150},
        ]
        for data in forms:
            with patch.object(jobs, 'launch') as launch:
                response = self.client.post(self.url, data, HTTP_X_PROGRAM_JOB='1')
                self.assertEqual(response.status_code, 202, response.content)
                launch.assert_called_once()
            folder = max(jobs.root().iterdir(), key=lambda p: p.stat().st_mtime_ns)
            jobs.write_json(folder/'state.json', {'state': 'cancelled'})
        with patch.object(jobs, 'launch') as launch:
            bad = self.client.post(self.url, {'program_action': 'convert'}, HTTP_X_PROGRAM_JOB='1')
            self.assertEqual(bad.status_code, 400)
            launch.assert_not_called()

    def test_owner_session_expiration_and_csrf(self):
        url, folder = self.create()
        other = Client(); other.force_login(self.other)
        same = Client(); same.force_login(self.owner)
        for client in (other, same):
            self.assertEqual(client.get(url).status_code, 404)
            self.assertEqual(client.post(url, {'action': 'cancel'}).status_code, 404)
        csrf_client = Client(enforce_csrf_checks=True); csrf_client.force_login(self.owner)
        self.assertEqual(csrf_client.post(url, {'action': 'cancel'}).status_code, 403)
        with patch('django.core.signing.time.time', return_value=time.time() + jobs.TTL + 2):
            self.assertEqual(self.client.get(url).status_code, 404)
        self.assertFalse((folder/'cancel').exists())

    def test_real_cleaner_result_download_and_duplicate_rejected(self):
        url, folder = self.create(rules=['spaces'])
        with patch.object(jobs, 'launch') as launch:
            duplicate = self.client.post(self.url, {'program_action': 'clean', 'document': document()}, HTTP_X_PROGRAM_JOB='1')
            self.assertEqual(duplicate.status_code, 400); launch.assert_not_called()
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'done')
        response = self.client.get(url + '?download=1')
        result = b''.join(response.streaming_content)
        self.assertEqual(Document(BytesIO(result)).paragraphs[0].text, 'Ala ma kota. Kot odwiedził ogród.')
        self.assertFalse((folder/'source.docx').exists())

    def test_cancel_before_start_never_runs_converter(self):
        url, folder = self.create()
        self.assertEqual(self.client.post(url, {'action': 'cancel'}).json()['state'], 'cancelling')
        with patch.object(jobs, 'convert_document') as converter:
            jobs.execute(folder); converter.assert_not_called()
        self.assertEqual(self.client.get(url).json()['state'], 'cancelled')
        self.assertFalse((folder/'source.docx').exists())
        self.assertEqual(self.client.get(url + '?download=1').status_code, 404)

    def test_real_rebuild_confirmation_reuses_original_without_upload(self):
        url, folder = self.create(document=document(header=True), rebuild='on')
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'confirmation')
        self.assertTrue((folder/'source.docx').exists())
        with patch.object(jobs, 'launch') as launch:
            self.assertEqual(self.client.post(url, {'action': 'confirm'}).status_code, 202)
            self.assertEqual(self.client.post(url, {'action': 'confirm'}).status_code, 404)
            launch.assert_called_once()
        jobs.execute(folder)
        self.assertEqual(self.client.get(url).json()['state'], 'done')

    def test_confirmation_renews_access_for_processing_and_download(self):
        url, folder = self.create()
        jobs.write_json(folder/'state.json', {'state': 'confirmation', 'message': 'Potwierdź'})
        now = time.time()
        with patch('django.core.signing.time.time', return_value=now + jobs.TTL - 60), patch.object(jobs, 'launch'):
            response = self.client.post(url, {'action': 'confirm'})
            self.assertEqual(response.status_code, 202)
            renewed = response.json()['url']
            self.assertNotEqual(url, renewed)
        with patch('django.core.signing.time.time', return_value=now + jobs.TTL + 10):
            self.assertEqual(self.client.get(url).status_code, 404)
            self.assertEqual(self.client.get(renewed).status_code, 200)


class CancellationTests(SimpleTestCase):
    @skipUnless(shutil.which('node'), 'Test zachowania JavaScript wymaga Node.js.')
    def test_progress_cancel_confirmation_and_restore_javascript(self):
        script = Path(__file__).with_name('js_tests') / 'program_jobs.cjs'
        result = subprocess.run([shutil.which('node'), str(script)], capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_cancel_kills_and_reaps_the_actual_converter_process(self):
        with TemporaryDirectory() as tmp:
            directory = Path(tmp)
            directory.joinpath('job.json').write_text('{"formats": ["pdf"]}')
            actual = subprocess.Popen
            children = []
            def spawn(*args, **kwargs):
                # A real long-running child makes cancellation deterministic.
                command = [args[0][0], '-c', 'import time; time.sleep(30)']
                process = actual(command, **kwargs); children.append(process); return process
            calls = []
            def control(_):
                calls.append(1)
                if len(calls) >= 2:
                    raise jobs.JobCancelled()
            with patch('core.services.document_converter.subprocess.Popen', side_effect=spawn):
                started = time.monotonic()
                with self.assertRaises(jobs.JobCancelled):
                    run_converter(directory, 10, control=control)
            self.assertLess(time.monotonic() - started, 5)
            self.assertIsNotNone(children[0].poll())
