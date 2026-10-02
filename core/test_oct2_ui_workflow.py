from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase
from django.urls import reverse
from django.utils import timezone
from docx import Document

from core.edit_versions import version_of
from core.odkurzacz_forms import RepetitionsForm
from core.selectors.texts import _annotated_texts, text_detail_context
from core.services.document_converter import convert_document
from core.services.document_repetitions import color_document
from people.models import Person, Role
from texts.models import Anthology, Review, Text
from workflow.admin_stage_edit import edit_stage
from workflow.models import WorkflowStage as S, WorkflowRoleAssignment as A
from workflow.state import current_stage


class AdminShortcutsAndStageDeletionTests(TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        settings = self.settings(ACTIVITY_SPOOL_DIR=temporary.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.admin = get_user_model().objects.create_superuser('admin', 'admin@example.com', 'test')
        self.member = get_user_model().objects.create_user('reader', 'reader@example.com', 'test')
        person = Person.objects.create(user=self.member, email=self.member.email, first_name='Jan', last_name='Test')
        person.roles.add(Role.objects.get_or_create(name='Recenzent')[0])
        self.book = Anthology.objects.create(title='Test')
        self.text = Text.objects.create(title='Tekst', length=100, anthology=self.book)
        self.today = timezone.localdate()
        self.client.force_login(self.admin)

    def completed(self, kind='editing', **kwargs):
        return S.objects.create(text=self.text, stage_type=kind, is_completed=True,
                                started_at=self.today, ended_at=self.today, **kwargs)

    def test_admin_links_only_visible_to_superuser_for_both_details(self):
        self.completed()
        review = Review.objects.create(title='Recenzja', length=100, anthology=self.book,
                                       old_reviews=True)
        for kind, record, route in [('text', self.text, 'assigned_text_detail'),
                                     ('review', review, 'assigned_review_detail')]:
            with self.subTest(kind=kind):
                url = reverse('core:' + route, args=[record.pk])
                admin_url = reverse('admin:texts_' + kind + '_change', args=[record.pk])
                self.client.force_login(self.admin)
                response = self.client.get(url)
                self.assertContains(response, admin_url)
                self.assertContains(response, 'Edytuj w adminie')
                self.client.force_login(self.member)
                response = self.client.get(url)
                self.assertNotContains(response, admin_url)
                self.assertNotContains(response, 'Edytuj w adminie')

    def test_notes_have_accessible_hidden_label_and_short_button(self):
        response = self.client.get(reverse('core:assigned_text_detail', args=[self.text.pk]))
        self.assertNotContains(response, '>Nowa notatka</label>')
        self.assertNotContains(response, '>Dodaj notatkę</button>')
        self.assertContains(response, '>Dodaj</button>')
        self.assertContains(response, '>Treść notatki</label>')

    def test_delete_without_replacement_keeps_completed_last_stage(self):
        last = self.completed('third_proofreading')
        self.completed('editing')  # PK/date must not override workflow order.
        assignment = A.objects.create(text=self.text, role='proofreader_4', assigned_to=self.member)
        removed = S.objects.create(text=self.text, stage_type='fourth_proofreading', assignment=assignment)
        edit_stage(removed.pk, self.admin, version_of(self.text), action='delete')
        self.assertFalse(S.objects.filter(pk=removed.pk).exists())
        self.assertFalse(A.objects.filter(pk=assignment.pk).exists())
        self.assertEqual(current_stage(list(self.text.workflow_stages.all())).pk, last.pk)
        self.assertEqual(_annotated_texts().get(pk=self.text.pk).current_stage_type, last.stage_type)
        last.refresh_from_db()
        self.assertTrue(last.is_completed)
        self.assertEqual(last.ended_at, self.today)
        self.assertEqual(S.objects.filter(text=self.text).count(), 2)
        detail = text_detail_context(user=self.admin, text=self.text)
        self.assertEqual(len(detail['visible_stages']), len({row['pk'] for row in detail['visible_stages']}))

    def test_admin_delete_form_accepts_empty_replacement(self):
        self.completed()
        removed = S.objects.create(text=self.text, stage_type='first_proofreading')
        url = reverse('admin:workflow_stage_correct', args=[removed.pk])
        response = self.client.post(url, {'action': 'delete', 'replacement': '',
                                          'version': version_of(self.text), 'confirm': 'on'})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(S.objects.filter(pk=removed.pk).exists())

    def test_deleting_last_record_leaves_empty_workflow(self):
        removed = S.objects.create(text=self.text, stage_type='ready_for_editing')
        edit_stage(removed.pk, self.admin, version_of(self.text), action='delete')
        self.assertFalse(self.text.workflow_stages.exists())
        self.assertIsNone(_annotated_texts().get(pk=self.text.pk).current_stage_type)

    def test_explicit_replacement_still_opens_chosen_stage(self):
        self.completed()
        removed = S.objects.create(text=self.text, stage_type='first_proofreading')
        edit_stage(removed.pk, self.admin, version_of(self.text), action='delete', replacement='second_proofreading')
        stage = current_stage(list(self.text.workflow_stages.all()))
        self.assertEqual(stage.stage_type, 'second_proofreading')
        self.assertFalse(stage.is_completed)
        self.assertIsNone(stage.started_at)

    def test_stale_form_and_non_superuser_cannot_delete(self):
        removed = S.objects.create(text=self.text, stage_type='ready_for_editing')
        old = version_of(self.text)
        self.completed()
        with self.assertRaises(ValidationError):
            edit_stage(removed.pk, self.admin, old, action='delete')
        with self.assertRaises(PermissionDenied):
            edit_stage(removed.pk, self.member, version_of(self.text), action='delete')
        self.assertTrue(S.objects.filter(pk=removed.pk).exists())

    def test_closed_anthology_still_blocks_deletion(self):
        removed = S.objects.create(text=self.text, stage_type='ready')
        self.book.status = Anthology.Status.READY
        self.book.save()
        with self.assertRaises(ValidationError):
            edit_stage(removed.pk, self.admin, version_of(self.text), action='delete')
        self.assertTrue(S.objects.filter(pk=removed.pk).exists())

    def test_current_stage_ignores_archived_and_unreleased_work(self):
        last = self.completed()
        self.completed('third_proofreading', is_current=False)
        S.objects.create(text=self.text, stage_type='fourth_proofreading', is_released=False)
        self.assertEqual(current_stage(list(self.text.workflow_stages.all())).pk, last.pk)
        self.assertEqual(_annotated_texts().get(pk=self.text.pk).current_stage_type, 'editing')


class RepetitionPaletteTests(SimpleTestCase):
    def document(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Koty').bold = True
        paragraph.add_run(' koty.').italic = True
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def test_both_palettes_keep_text_and_formatting_with_correct_rgb_range(self):
        for palette, low, high in [('dark', 0, 150), ('light', 160, 255)]:
            with self.subTest(palette=palette), patch('core.services.document_repetitions._lemma_map',
                side_effect=lambda words: {word.lower(): word.lower() for word in words}):
                document = Document(color_document(BytesIO(self.document()), color_palette=palette,
                    analysis_options={'duplicates': False, 'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}))
                self.assertEqual(document.paragraphs[0].text, 'Koty koty.')
                colors = [run.font.color.rgb for run in document.paragraphs[0].runs if run.font.color.rgb]
                self.assertTrue(colors)
                self.assertEqual(len(set(colors)), 1)
                self.assertTrue(all(low <= channel <= high for color in colors for channel in color))
                self.assertTrue(document.paragraphs[0].runs[0].bold)
                self.assertTrue(any(run.italic for run in document.paragraphs[0].runs))

    def test_form_accepts_light_and_defaults_missing_choice_to_dark(self):
        data = {'window_size': 35, 'min_word_length': 4, 'sentence_limit': 35, 'paragraph_limit': 150}
        for palette in ('light', ''):
            form = RepetitionsForm({**data, 'color_palette': palette},
                {'document': SimpleUploadedFile('test.docx', self.document())})
            self.assertTrue(form.is_valid(), form.errors)
            self.assertEqual(form.analysis_config()['color_palette'], palette or 'dark')
        form = RepetitionsForm({**data, 'color_palette': 'invalid'},
            {'document': SimpleUploadedFile('test.docx', self.document())})
        self.assertFalse(form.is_valid())
        self.assertIn('color_palette', form.errors)

    def test_invalid_palette_rejected_by_service(self):
        with self.assertRaises(ValueError):
            color_document(BytesIO(self.document()), color_palette='invalid')

    def test_light_palette_reaches_subprocess_and_saved_docx(self):
        with TemporaryDirectory() as temporary, self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary)):
            output, _, _ = convert_document(SimpleUploadedFile('test.docx', self.document()), [],
                include_docx=True, normalize=False, repetitions={'color_palette': 'light',
                    'analysis_options': {'long_sentences': False, 'long_paragraphs': False}})
            with output:
                document = Document(output)
        colors = [run.font.color.rgb for run in document.paragraphs[0].runs if run.font.color.rgb]
        self.assertTrue(colors)
        self.assertTrue(all(160 <= channel <= 255 for color in colors for channel in color))
