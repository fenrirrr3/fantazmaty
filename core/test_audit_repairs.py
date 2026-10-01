"""Durable regressions for the audited permissions, imports and data lifecycle."""
import html
import re
from datetime import datetime, time, timedelta
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

from django.contrib import admin
from django.contrib.admin.options import BaseModelAdmin
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase, SimpleTestCase, RequestFactory
from django.urls import reverse, NoReverseMatch
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from authors.models import Author, BlacklistEntry
from people.models import Person, Role, Vacation
from texts.models import Anthology, Review, ReviewAssignment, Text
from workflow.models import WorkflowStage as Stage, WorkflowRoleAssignment as Assignment


class AuditRepairs(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        override = self.settings(ACTIVITY_SPOOL_DIR=self.temp.name); override.enable(); self.addCleanup(override.disable)
        users = get_user_model()
        self.superuser = users.objects.create_superuser('audit-admin', 'admin@example.test', 'test')
        self.reviewer = users.objects.create_user('audit-member', 'member@example.test', 'test')
        self.person = Person.objects.create(user=self.reviewer, first_name='Jan', last_name='Zeta', email=self.reviewer.email)
        self.person.roles.add(Role.objects.get_or_create(name='Recenzent')[0])
        self.coordinator = users.objects.create_user('audit-coord', 'coordinator@example.test', 'test')
        self.coordinator_profile = Person.objects.create(user=self.coordinator, first_name='Anna', last_name='Alfa', email=self.coordinator.email)
        self.coordinator_profile.roles.add(Role.objects.get_or_create(name='Koordynator recenzji')[0])
        self.book = Anthology.objects.create(title='Test audytu')
        self.review = Review.objects.create(title='Zgłoszenie', anthology=self.book, length=1000, author_first_name='Tajny', author_last_name='Autor', email='secret@example.test', author_message='SECRET_MESSAGE', file_url='https://www.dropbox.com/secret-folder')

    def detail(self): return reverse('core:assigned_review_detail', args=[self.review.pk])

    def token(self):
        page = self.client.get(self.detail())
        return html.unescape(re.search(r'name="_edit_version" value="([^"]+)"', page.content.decode())[1])

    def test_coordinator_manages_folder_without_author_identity(self):
        self.client.force_login(self.coordinator)
        page = self.client.get(self.detail())
        self.assertContains(page, 'review-folder-heading')
        for secret in ('secret@example.test', 'SECRET_MESSAGE', 'Tajny Autor'):
            self.assertNotContains(page, secret)
        result = self.client.post(reverse('core:update_review_file', args=[self.review.pk]), {'file_url':'https://www.dropbox.com/changed-folder', '_edit_version':self.token()})
        self.assertEqual(result.status_code, 302)
        self.review.refresh_from_db(); self.assertIn('changed-folder', self.review.file_url)

    def test_member_conflict_does_not_reveal_private_fields(self):
        ReviewAssignment.objects.create(review=self.review, user=self.reviewer, position=1, opinion='yes', notes='Oddana opinia')
        self.client.force_login(self.reviewer)
        token = self.token()
        self.review.title = 'Zmiana'; self.review.save()
        page = self.client.post(self.detail(), {'opinion':'maybe', 'notes':'Nowa', '_edit_version':token})
        self.assertEqual(page.status_code, 409)
        for secret in ('secret-folder', 'SECRET_MESSAGE', 'secret@example.test'): self.assertNotContains(page, secret, status_code=409)
        self.assertContains(page, 'Oddana opinia', status_code=409)

    def test_submitted_review_cannot_be_relinquished(self):
        from core.services.reviews import unassign_reviewer
        assignment = ReviewAssignment.objects.create(review=self.review, user=self.reviewer, position=1, opinion='yes', notes='Ocena pozostaje')
        with self.assertRaises(ValidationError): unassign_reviewer(user=self.reviewer, review_id=self.review.pk)
        assignment.refresh_from_db(); self.assertEqual(assignment.notes, 'Ocena pozostaje')
        self.client.force_login(self.reviewer)
        page = self.client.get(self.detail())
        self.assertFalse(page.context['can_unassign_review'])
        self.assertNotContains(page, reverse('core:unassign_reviewer',args=[self.review.pk]))

    def test_unsubmitted_review_can_be_relinquished(self):
        from core.services.reviews import unassign_reviewer
        assignment = ReviewAssignment.objects.create(review=self.review, user=self.reviewer, position=1, opinion='reading')
        self.assertTrue(unassign_reviewer(user=self.reviewer, review_id=self.review.pk))
        self.assertFalse(ReviewAssignment.objects.filter(pk=assignment.pk).exists())

    def test_role_change_keeps_own_current_reviews_accessible(self):
        ReviewAssignment.objects.create(review=self.review, user=self.reviewer, position=1)
        self.person.roles.clear()
        self.client.force_login(self.reviewer)
        self.assertEqual(self.client.get(reverse('core:my_reviews')).status_code, 200)
        from core.permissions import can_self_assign_reviews
        self.assertFalse(can_self_assign_reviews(self.reviewer))

    def test_existing_assignment_can_start_without_current_role(self):
        from core.services.texts import start_assigned_stage
        text = Text.objects.create(title='Korekta', length=1, anthology=self.book)
        assignment = Assignment.objects.create(text=text, role='proofreader_1', assigned_to=self.reviewer)
        stage = Stage.objects.create(text=text, stage_type='first_proofreading', assignment=assignment)
        self.person.roles.clear()
        start_assigned_stage(user=self.reviewer, stage_id=stage.pk, started_at=timezone.localdate())
        stage.refresh_from_db(); self.assertEqual(stage.started_at, timezone.localdate())

    def test_new_waiting_stage_gets_queue_date_unknown_old_stage_is_visible(self):
        from core.selectors.reports import _inactivity_stages
        text = Text.objects.create(title='Oczekuje', length=1, anthology=self.book)
        stage = Stage.objects.create(text=text, stage_type='ready_for_editing')
        self.assertEqual(stage.queued_at, timezone.localdate())
        Stage.objects.filter(pk=stage.pk).update(queued_at=None)
        self.assertTrue(_inactivity_stages(timezone.localdate(),28,7,'waiting',[],False).filter(pk=stage.pk).exists())

    def test_midnight_vacation_boundary_does_not_overlap(self):
        today = timezone.localdate()
        midnight = timezone.make_aware(datetime.combine(today + timedelta(days=2), time.min))
        Vacation.objects.create(person=self.person, start_date=today, end_date=midnight)
        second = Vacation(person=self.person, start_date=today+timedelta(days=2), end_date=midnight+timedelta(days=1))
        second.full_clean()

    def test_inactivity_report_renders_populated_rows_and_unknown_dates(self):
        today = timezone.localdate()
        expected = {}
        for title, kind, started, queued in (
            ('Praca 40 dni', 'editing', today-timedelta(days=40), today),
            ('Czeka 12 dni', 'ready_for_editing', None, today-timedelta(days=12)),
            ('Nieznany czas', 'ready_for_editing', None, None),
            ('Krótka praca', 'editing', today-timedelta(days=3), today),
            ('Przyszła praca', 'editing', today+timedelta(days=3), today),
        ):
            text = Text.objects.create(title=title, length=1, anthology=self.book)
            stage = Stage.objects.create(text=text, stage_type=kind, started_at=started)
            Stage.objects.filter(pk=stage.pk).update(queued_at=queued)
            if title in ('Praca 40 dni', 'Czeka 12 dni', 'Nieznany czas'):
                expected[title] = stage.pk
        self.client.force_login(self.superuser)
        url = reverse('core:workflow_inactivity')
        for params, titles in (
            ({}, ['Praca 40 dni', 'Czeka 12 dni', 'Nieznany czas']),
            ({'mode':'active'}, ['Praca 40 dni']),
            ({'mode':'waiting'}, ['Czeka 12 dni', 'Nieznany czas']),
            ({'stage':['editing']}, ['Praca 40 dni']),
            ({'stage':['editing','ready_for_editing'], 'sort':'-days'},
             ['Praca 40 dni', 'Czeka 12 dni', 'Nieznany czas']),
        ):
            with self.subTest(params=params):
                response = self.client.get(url, params)
                self.assertEqual(response.status_code, 200)
                rows = list(response.context['rows'])
                self.assertEqual([row['text']['title'] for row in rows], titles)
                self.assertEqual([row['stage']['pk'] for row in rows], [expected[title] for title in titles])
        response = self.client.get(url)
        self.assertContains(response, 'Nieznany czas')
        self.assertNotContains(response, 'Przyszła praca')
        self.assertNotContains(response, 'Krótka praca')

    def test_inactivity_search_is_polish_and_preserves_author_permissions(self):
        author = Author.objects.create(first_name='Łucja', last_name='Żółć', email='private-author@example.test')
        text = Text.objects.create(title='Zażółć gęślą', length=1, anthology=self.book)
        text.authors.add(author)
        Stage.objects.create(text=text, stage_type='editing',
                             started_at=timezone.localdate()-timedelta(days=40))
        url = reverse('core:workflow_inactivity')
        self.client.force_login(self.superuser)
        for query in ('zazolc gesla', 'LUCJA ZOLC', 'Test audytu redakcja'):
            with self.subTest(query=query):
                response = self.client.get(url, {'q':query})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(len(response.context['rows']), 1)
                self.assertContains(response, text.title)
        self.client.force_login(self.coordinator)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, text.title)
        self.assertNotContains(response, str(author))
        self.assertNotContains(response, author.email)
        self.assertEqual(len(self.client.get(url, {'q':'lucja zolc'}).context['rows']), 0)

    def test_disabled_account_not_in_vacation_choices(self):
        self.reviewer.is_active=False; self.reviewer.save(update_fields=['is_active'])
        self.client.force_login(self.coordinator)
        page=self.client.get(reverse('core:my_vacations'))
        self.assertNotIn(self.person.pk, [p.pk for p in page.context['vacation_people']])

    def test_blacklist_and_mailbox_are_versioned(self):
        from core.edit_versions import version_of
        from core.models import MailboxConnection
        for obj in (BlacklistEntry.objects.create(name='Osoba'), MailboxConnection.objects.create(name='Teksty',host='example.test',username='teksty')):
            before=version_of(obj); obj.save(); self.assertGreater(version_of(obj), before)

    def test_email_uniqueness_is_enforced_outside_forms(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            get_user_model().objects.create_user('duplicate', ' MEMBER@example.test ')
        get_user_model().objects.create_user('blank1', '')
        get_user_model().objects.create_user('blank2', '')

    def test_roles_and_names_have_one_profile_source(self):
        self.person.last_name='Nowe'; self.person.save(update_fields=['last_name'])
        self.reviewer.refresh_from_db(); self.assertEqual(self.reviewer.last_name,'Nowe')
        self.reviewer.groups.add(Group.objects.get_or_create(name='Korektor')[0])
        self.assertTrue(self.person.roles.filter(name='Korektor').exists())
        self.person.roles.remove(*self.person.roles.filter(name='Korektor'))
        self.assertFalse(self.reviewer.groups.filter(name='Korektor').exists())
        self.coordinator_profile.roles.clear()
        from core.permissions import is_coordinator
        self.assertFalse(is_coordinator(self.coordinator))

    def test_custom_admin_leaves_django_base_method_unchanged(self):
        from core.admin_site import CMSAdminSite
        self.assertIsInstance(admin.site, CMSAdminSite)
        self.assertEqual(BaseModelAdmin.formfield_for_foreignkey.__module__, 'django.contrib.admin.options')
        with self.assertRaises(NoReverseMatch): reverse('core:bulk_text_action')

    def test_newsletter_sort_is_global_and_has_all_columns(self):
        from core.models import NewsletterConsent
        from core.pagination import paginate_items
        NewsletterConsent.objects.bulk_create([NewsletterConsent(email=f'{i:02}@example.test',premieres=i%2==0) for i in range(35)])
        request=RequestFactory().get('/',{'sort':'-email'}); request.user=self.superuser
        page=paginate_items(request,NewsletterConsent.objects.all())
        self.assertEqual(page.object_list[0].email,'34@example.test')
        self.assertIn('Zgoda na newsletter o premierach',page.sort_columns)

    def test_report_projection_only_hydrates_one_page(self):
        from core.selectors.reports import reviewer_activity_context
        from core.pagination import paginate_items
        from django.http import QueryDict
        for number in range(32):
            review=Review.objects.create(title=f'Test {number:02}',anthology=self.book,length=1)
            ReviewAssignment.objects.create(review=review,user=self.reviewer,position=1,opinion='yes')
        rows=reviewer_activity_context(user=self.superuser,params=QueryDict())['activity_rows']
        self.assertIsNone(rows.queryset._result_cache)
        request=RequestFactory().get('/',{'sort':'title'});request.user=self.superuser
        page=paginate_items(request,rows)
        self.assertEqual(page.paginator.count,32); self.assertEqual(len(page.object_list),25)
        self.assertEqual(page.object_list[0]['title'],'Test 00')

    def test_workflow_reports_render_assigned_and_unassigned_rows(self):
        self.client.force_login(self.superuser)
        for role, kind, report in (
            ('editor','editing','editor_activity'),
            ('proofreader_1','first_proofreading','proofreader_activity'),
            ('verifier_1','first_verification','verifier_activity'),
        ):
            text = Text.objects.create(title='Etap ' + kind, length=1, anthology=self.book)
            assignment = Assignment.objects.create(text=text, role=role, assigned_to=self.reviewer)
            Stage.objects.create(text=text, stage_type=kind, assignment=assignment)
            other = Text.objects.create(title='Pusty ' + kind, length=1, anthology=self.book)
            Stage.objects.create(text=other, stage_type=kind)
            self.assertEqual(self.client.get(reverse('core:' + report)).status_code, 200)

    def test_polish_search_preserves_accents_and_literal_percent(self):
        author=Author.objects.create(first_name='Łucja', last_name='ŻÓŁĆ', email='lc@example.test')
        self.assertTrue(Author.objects.filter(last_name__plcontains='zolc').filter(pk=author.pk).exists())
        self.assertTrue(Author.objects.filter(first_name__plcontains='lucja').filter(pk=author.pk).exists())
        self.assertFalse(Author.objects.filter(last_name__plcontains='%').filter(pk=author.pk).exists())


class AuditDocuments(SimpleTestCase):
    def saved(self,document):
        stream=BytesIO(); document.save(stream); stream.seek(0); return stream

    def test_all_protected_abbreviations_keep_internal_spacing(self):
        from core.services.odkurzacz import correct_editorial_text, DEFAULT_EDITORIAL_RULES
        for abbreviation in ('ś.p.', 'c.d.n.', 'o.o.', 'j.w.'):
            output=correct_editorial_text(f'To {abbreviation} przykład.',DEFAULT_EDITORIAL_RULES)
            self.assertIn(abbreviation,output)

    def test_abbreviation_protection_does_not_add_sentence_spacing(self):
        from core.services.odkurzacz import correct_editorial_text
        self.assertEqual(correct_editorial_text('To o.o. przykład.', ['after_punct']), 'To o.o. przykład.')

    def test_cyclic_style_chain_is_bounded(self):
        from core.services.document_formatting import normalize_docx
        document=Document(); a=document.styles.add_style('CycleA',1); b=document.styles.add_style('CycleB',1)
        a.base_style=b; b.base_style=a
        document.add_paragraph('Tekst',style=a)
        with normalize_docx(self.saved(document)) as result:
            self.assertEqual(Document(result).paragraphs[0].text,'Tekst')

    def test_nonempty_header_requires_rebuild_acceptance(self):
        from core.services.document_rebuild import inspect_docx, rebuild_docx, RebuildUnsupported
        document=Document();document.add_paragraph('Treść');document.sections[0].header.paragraphs[0].text='Ważny nagłówek'
        with self.assertRaises(RebuildUnsupported) as error: inspect_docx(self.saved(document))
        self.assertIn('nagłówki',str(error.exception))
        with rebuild_docx(self.saved(document),allow_omissions=True) as result:
            self.assertEqual(Document(result).paragraphs[0].text,'Treść')

    def test_mail_csv_matches_manual_import_and_ignores_signature_images(self):
        from core.services.mailbox_import import parse_message
        from core.services.review_import_parser import encode_submission, parse_review_records
        fields=['Jan Autor','Tytuł; część druga','fantasy','przemoc','1000','jan@example.test','123','premierach','Linia pierwsza\nLinia druga; też']
        mail=EmailMessage();mail['Subject']='Nabór: „Test” – Tytuł'
        mail.set_content(encode_submission(fields))
        mail.add_attachment(b'docx',maintype='application',subtype='vnd.openxmlformats-officedocument.wordprocessingml.document',filename='tekst.docx')
        mail.add_attachment(b'logo',maintype='image',subtype='png',filename='logo.png',disposition='inline')
        result=parse_message(1,mail.as_bytes())
        self.assertEqual([name for name,data in result['files']],['tekst.docx'])
        self.assertEqual(result['ignored_files'],['logo.png'])
        direct=parse_review_records(encode_submission(fields))[0][0]
        imported=parse_review_records(result['record'])[0][0]
        self.assertEqual(imported['title'],direct['title']);self.assertEqual(imported['author_message'],direct['author_message'])

    def test_consent_recovery_fetches_body_without_large_attachment(self):
        from core.services.mailbox_body import fetch_text_body
        body=b'Jan Autor;Tytul;fantasy;1000;jan@example.test;123;Nabor;premierach'
        client=MagicMock()
        client.uid.side_effect=[('OK',[f'1 (UID 12 BODYSTRUCTURE (("TEXT" "PLAIN" ("CHARSET" "utf-8") NIL NIL "8BIT" {len(body)} 1 NIL NIL NIL)("APPLICATION" "OCTET-STREAM" NIL NIL NIL "BASE64" 90000000 NIL ("ATTACHMENT" NIL)) "MIXED"))'.encode()]),('OK',[(b'1 (UID 12 BODY[1] {100}',body)])]
        result=fetch_text_body(client,12)
        self.assertIn(body,result)
        self.assertEqual(client.uid.call_args.args[-1],'(UID BODY.PEEK[1]<0.1048577>)')

    def test_pdf_nested_paragraphs_receive_geometry(self):
        from core.services.document_pdf import render_pdf
        from docx.shared import Cm
        from fpdf import FPDF
        doc = Document()
        listed = doc.add_paragraph('Lista', style='List Number')
        listed.alignment = WD_ALIGN_PARAGRAPH.CENTER
        listed.paragraph_format.first_line_indent = Cm(0)
        listed.paragraph_format.line_spacing = 2
        cell = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell.text = 'Tabela'; cell.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        captured = []
        original = FPDF.write_html
        def spy(pdf, content, **kwargs):
            captured.append(content)
            return original(pdf, content, **kwargs)
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / 'source.docx'; doc.save(source)
            with patch.object(FPDF, 'write_html', spy):
                render_pdf(source, Path(tmp)/'result.pdf', '<ol><li>Lista</li></ol><table><tr><td><p>Tabela</p></td></tr></table>', {}, 'Test')
            self.assertTrue((Path(tmp)/'result.pdf').read_bytes().startswith(b'%PDF'))
        from lxml import html as parser
        tree = parser.fragment_fromstring(''.join(captured), create_parent='div')
        listed = tree.xpath('.//li/p')[0]
        self.assertEqual(listed.get('align'), 'center')
        self.assertEqual(listed.get('data-indent'), '0')
        self.assertEqual(listed.get('line-height'), '2.0')
        self.assertEqual(tree.xpath('.//td/p')[0].get('align'), 'right')

    def test_conversion_warnings_are_delivered_in_download(self):
        import json
        from zipfile import ZipFile
        from core.services.document_converter import convert_document
        doc = Document(); doc.add_paragraph('Treść')
        def worker(directory, timeout):
            (directory/'document.epub').write_bytes(b'EPUB test output')
            (directory/'warnings.json').write_text(json.dumps(['Pominięto nieobsługiwany element.']))
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)), patch('core.services.document_converter.run_converter', side_effect=worker):
            result, extension, _ = convert_document(self.saved(doc), ['epub'])
            with result, ZipFile(result) as archive:
                self.assertEqual(extension, 'zip')
                self.assertIn('Pominięto', archive.read('Uwagi_konwersji.txt').decode())

    def test_bootstrap_records_failures_before_worker_initialization(self):
        import json
        import runpy
        import core.services
        runner = runpy.run_path
        bootstrap = Path(core.services.__file__).with_name('document_worker_bootstrap.py')
        with TemporaryDirectory() as tmp:
            with patch('sys.argv', ['bootstrap', tmp]), patch('runpy.run_path', side_effect=ImportError('PRIVATE credential')):
                with self.assertRaises(SystemExit) as caught:
                    runner(str(bootstrap), run_name='__main__')
            self.assertEqual(caught.exception.code, 2)
            report = (Path(tmp)/'error.json').read_text()
            self.assertNotIn('PRIVATE', report)
            self.assertEqual(json.loads(report)['error'], 'ImportError')
