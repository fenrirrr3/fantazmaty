from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from django.test import TestCase, SimpleTestCase
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.http import QueryDict
from django.urls import reverse
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from people.models import Person, Role
from texts.models import Review, Anthology
from core.services.reviews import assign_reviewer, unassign_reviewer, change_review_status
from core.selectors.reviews import review_list_context
from core.services.document_rebuild import rebuild_docx, RebuildUnsupported
from core.services.document_converter import convert_document
from core.services.mailbox_import import package_messages

class ReviewStatusTests(TestCase):
    def member(self, name, role):
        user=get_user_model().objects.create_user(name, name+'@example.com','password')
        person=Person.objects.filter(user=user).first()
        if not person:
            person = Person.objects.create(
                user=user, email=user.email, first_name=name, last_name="Test"
            )
        person.roles.add(Role.objects.get_or_create(name=role)[0])
        return user

    def setUp(self):
        self.coordinator = self.member("coord", "Koordynator recenzji")
        self.reviewer = self.member("reader", "Recenzent")
        self.other = self.member("other", "Koordynator korekty")
        self.review = Review.objects.create(
            anthology=Anthology.objects.create(title="Nabór"),
            title="Tekst",
            author_first_name="Autor",
            author_last_name="Test",
            email="author@example.com",
            genre="fantasy",
            length=1000,
        )

    def test_manual_open_status_and_permissions(self):
        with self.assertRaises(PermissionDenied):
            change_review_status(user=self.other, review_id=self.review.pk, new_status="to_decide")
        change_review_status(
            user=self.coordinator, review_id=self.review.pk, new_status="to_decide"
        )
        self.review.refresh_from_db()
        self.assertIsNone(self.review.decision_at)
        assign_reviewer(user=self.reviewer, review_id=self.review.pk)
        self.review.refresh_from_db()
        self.assertEqual(self.review.status, "to_decide")
        unassign_reviewer(user=self.reviewer, review_id=self.review.pk)
        self.review.refresh_from_db()
        self.assertEqual(self.review.status, "to_decide")
        self.client.force_login(self.reviewer)
        response = self.client.get(reverse("core:assigned_review_detail", args=[self.review.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "wyznaczona osoba")
        self.assertNotContains(response, 'name="status" value="to_decide"')

    def test_assignment_drives_regular_status_and_filters(self):
        self.assertEqual(self.review.get_status_display(), "Do recenzji")
        assign_reviewer(user=self.reviewer, review_id=self.review.pk)
        self.review.refresh_from_db()
        self.assertEqual(self.review.get_status_display(), "W recenzjach")
        self.client.force_login(self.coordinator)
        page = self.client.get(reverse("core:assigned_review_detail", args=[self.review.pk]))
        import re
        import html

        token = html.unescape(
            re.search(r'name="_edit_version" value="([^"]+)"', page.content.decode()).group(1)
        )
        response = self.client.post(
            reverse("core:update_review_status", args=[self.review.pk]),
            {"status": "to_decide", "_edit_version": token},
        )
        self.assertEqual(response.status_code, 302)
        self.review.refresh_from_db()
        self.assertEqual(self.review.status, "to_decide")
        result = review_list_context(user=self.coordinator, params=QueryDict("status=to_decide"))
        self.assertEqual([r["pk"] for r in result["reviews"]], [self.review.pk])
        self.assertIn(("to_decide", "Do decyzji"), result["status_choices"])


class DocumentConsistencyTests(SimpleTestCase):
    def source(self, doc):
        out = BytesIO()
        doc.save(out)
        out.seek(0)
        return out

    def test_loss_warning_sections_links_and_heading_preservation(self):
        doc = Document()
        doc.add_heading("Rozdział", 1)
        doc.add_section()
        p = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "Odnośnik"
        run.append(t)
        link.append(run)
        p._p.append(link)
        with self.assertRaises(RebuildUnsupported) as error:
            rebuild_docx(self.source(doc))
        self.assertIn("sekcje", str(error.exception))
        self.assertIn("hiperłączy", str(error.exception))
        self.assertIn("liczba: 1", str(error.exception))
        rebuilt = Document(rebuild_docx(self.source(doc), allow_omissions=True))
        self.assertEqual(rebuilt.paragraphs[0].style.name, "Heading 1")

    def test_doc_defaults_and_theme(self):
        doc = Document()
        doc.add_paragraph("Test")
        defaults = (
            doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        )
        default_size = defaults.find(qn("w:sz"))
        default_size.set(qn("w:val"), "34")
        font = defaults.find(qn("w:rFonts"))
        for key in list(font.attrib):
            del font.attrib[key]
        font.set(qn("w:ascii"), "Courier New")
        font.set(qn("w:hAnsi"), "Courier New")
        rebuilt = Document(rebuild_docx(self.source(doc)))
        self.assertEqual(rebuilt.paragraphs[0].runs[0].font.size.pt, 17)
        self.assertEqual(rebuilt.paragraphs[0].runs[0].font.name, "Courier New")

    def test_epub_preserves_empty_paragraphs_and_temp_cleanup(self):
        from lxml import etree

        doc = Document()
        doc.add_paragraph("A")
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph("B")
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            out, _, _ = convert_document(self.source(doc), ["epub"])
            with out, ZipFile(out) as archive:
                root = etree.fromstring(archive.read("EPUB/content.xhtml"))
                paragraphs = root.findall(".//{http://www.w3.org/1999/xhtml}p")
                self.assertEqual(len(paragraphs), 4)
                self.assertEqual(
                    ["".join(p.itertext()).strip() for p in paragraphs], ["A", "", "", "B"]
                )
            self.assertFalse(list(Path(tmp).glob("document-*")))

    def test_original_invalid_docx_is_downloadable(self):
        data = b"not a valid docx"
        with (
            package_messages(
                [{"uid": 1, "folder": "Text", "files": [("text.docx", data)]}],
                clean=False,
                convert=False,
                rebuild=False,
            ) as out,
            ZipFile(out) as archive,
        ):
            self.assertEqual(archive.read(archive.namelist()[0]), data)
