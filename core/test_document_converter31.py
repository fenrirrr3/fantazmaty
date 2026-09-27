from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch, MagicMock
from itertools import combinations
from zipfile import ZipFile
import subprocess
from docx import Document
from django.test import TestCase, SimpleTestCase
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from core.odkurzacz_forms import DocumentConversionForm
from core.services.document_converter import convert_document, ConversionError, run_converter, conversion_slot


def sample():
    doc = Document(); doc.add_paragraph('Dwa  słowa')
    doc.add_heading('Zażółć gęślą jaźń', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tabela'
    from PIL import Image
    image = BytesIO(); Image.new('RGB', (30, 30), 'blue').save(image, format='PNG'); image.seek(0)
    doc.add_picture(image)
    output = BytesIO(); doc.save(output)
    return output.getvalue()


def upload():
    return SimpleUploadedFile('Opowiadanie.docx', sample())


class ConverterTests(SimpleTestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        override = self.settings(DOCUMENT_CONVERSION_DIR=Path(self.temp.name)); override.enable(); self.addCleanup(override.disable)

    def test_all_format_combinations_and_optional_cleaning(self):
        original = sample()
        for length in (1, 2):
            for formats in combinations(('pdf', 'epub'), length):
                for clean in (False, True):
                    with self.subTest(formats=formats, clean=clean):
                        result, extension, mime = convert_document(BytesIO(original), formats, use_cleaner=clean)
                        with result:
                            if length == 1:
                                data = {formats[0]: result.read()}
                            else:
                                self.assertEqual(extension, 'zip')
                                with ZipFile(result) as archive:
                                    data = {f: archive.read('document.' + f) for f in formats}
                            if 'pdf' in data:
                                self.assertTrue(data['pdf'].startswith(b'%PDF'))
                            if 'epub' in data:
                                with ZipFile(BytesIO(data['epub'])) as archive:
                                    self.assertEqual(archive.read('mimetype'), b'application/epub+zip')
                                    body = archive.read('EPUB/content.xhtml').decode()
                                    self.assertIn('Dwa słowa' if clean else 'Dwa  słowa', body)
                                    self.assertIn('Zażółć gęślą jaźń', body)
                                    self.assertIn('EPUB/assets/image-1.png', archive.namelist())
                        self.assertFalse(list(Path(self.temp.name).glob('document-*')))

    def test_missing_dependencies_and_no_formats(self):
        process = MagicMock(); process.wait.return_value = 2
        popen = MagicMock(); popen.__enter__.return_value = process
        with patch('core.services.document_converter.subprocess.Popen', return_value=popen), self.assertRaisesMessage(ConversionError, 'Brakuje bibliotek'):
            run_converter(Path(self.temp.name), 1)
        for formats in ([], ['mobi']):
            with self.assertRaises(ConversionError):
                convert_document(BytesIO(sample()), formats)

    def test_timeout_kills_process_group(self):
        process = MagicMock(); process.pid = 12345
        process.wait.side_effect = [subprocess.TimeoutExpired('ebook-convert', 1), 0]
        popen = MagicMock(); popen.__enter__.return_value = process
        with patch('core.services.document_converter.subprocess.Popen', return_value=popen), patch('core.services.document_converter.os.killpg') as kill, self.assertRaisesMessage(ConversionError, '90 sekund'):
            run_converter(Path(self.temp.name), 1)
        kill.assert_called_once()

    def test_failed_conversion_removes_temporary_files(self):
        with patch('core.services.document_converter.run_converter', side_effect=ConversionError('failed')):
            with self.assertRaises(ConversionError):
                convert_document(BytesIO(sample()), ['pdf'])
        self.assertFalse(list(Path(self.temp.name).glob('document-*')))

    def test_parallel_conversion_rejected(self):
        with conversion_slot():
            with self.assertRaisesMessage(ConversionError, 'inny dokument'), conversion_slot():
                pass

    def test_external_images_rejected_but_hyperlinks_allowed(self):
        from core.services.document_converter import validate_resources
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        for rel_type, rejected in ((RT.IMAGE, True), (RT.HYPERLINK, False)):
            doc = Document()
            doc.add_paragraph('Test')
            doc.part.relate_to('https://example.com/file', rel_type, is_external=True)
            data = BytesIO(); doc.save(data); data.seek(0)
            if rejected:
                with self.assertRaisesMessage(ConversionError, 'zewnętrzny'):
                    validate_resources(data)
            else:
                validate_resources(data)
            self.assertEqual(data.tell(), 0)

    def test_formats_required_and_docx_validation_shared(self):
        form = DocumentConversionForm({'formats': []}, {'document': upload()})
        self.assertFalse(form.is_valid()); self.assertIn('formats', form.errors)
        form = DocumentConversionForm({'formats': ['exe']}, {'document': upload()})
        self.assertFalse(form.is_valid())
        form = DocumentConversionForm({'formats': ['epub']}, {'document': SimpleUploadedFile('bad.docx', b'not a zip')})
        self.assertFalse(form.is_valid()); self.assertIn('document', form.errors)


class ProgramsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('admin', 'admin@example.com', 'password')
        self.client.force_login(self.user)
        self.url = reverse('core:programs')

    def test_two_collapsed_tools(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<summary>Odkurzacz</summary>')
        self.assertContains(response, 'name="convert-formats"', count=2)
        self.assertContains(response, 'name="convert-use_cleaner"')
        self.assertNotContains(response, 'program-tool" open')

    def test_conversion_download(self):
        with patch('core.views.programs.convert_document', return_value=(BytesIO(b'epub'), 'epub', 'application/epub+zip')) as convert:
            response = self.client.post(self.url, {'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['epub'], 'convert-use_cleaner': 'on'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/epub+zip')
        self.assertEqual(b''.join(response.streaming_content), b'epub')
        self.assertTrue(convert.call_args.kwargs['use_cleaner'])
        response.close()

    def test_missing_converter_displays_form_error(self):
        with patch('core.views.programs.convert_document', side_effect=ConversionError('Brakuje bibliotek')):
            response = self.client.post(self.url, {'program_action': 'convert', 'convert-document': upload(), 'convert-formats': ['pdf']})
        self.assertContains(response, 'Brakuje bibliotek')
        self.assertContains(response, 'program-tool" open')

    def test_invalid_action_and_anonymous(self):
        self.assertContains(self.client.post(self.url, {'program_action': 'wrong'}), 'Wybierz narzędzie')
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
