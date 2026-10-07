from django.test import SimpleTestCase
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from core.services.document_styles import paragraph_properties, paragraph_property


class DocumentStyleTests(SimpleTestCase):
    def test_direct_zero_and_false_override_inherited_values(self):
        document = Document()
        base = document.styles.add_style('Parent', WD_STYLE_TYPE.PARAGRAPH)
        base.paragraph_format.space_after = Pt(18)
        base.paragraph_format.page_break_before = True
        base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        child = document.styles.add_style('Child', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        child.paragraph_format.line_spacing = 2
        paragraph = document.add_paragraph('Test', style=child)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.page_break_before = False
        expected = {'space_after': Pt(0), 'page_break_before': False,
                    'alignment': WD_ALIGN_PARAGRAPH.RIGHT, 'line_spacing': 2}
        self.assertEqual(paragraph_properties(paragraph, expected), expected)
        for name, value in expected.items():
            self.assertEqual(paragraph_property(paragraph, name), value)

    def test_missing_properties_in_cyclic_styles_remain_unset(self):
        document = Document()
        one = document.styles.add_style('CycleOne', WD_STYLE_TYPE.PARAGRAPH)
        two = document.styles.add_style('CycleTwo', WD_STYLE_TYPE.PARAGRAPH)
        one.base_style = two
        two.base_style = one
        two.paragraph_format.line_spacing = 2
        paragraph = document.add_paragraph('Test', style=one)
        self.assertEqual(paragraph_properties(paragraph, ('line_spacing', 'page_break_before')),
                         {'line_spacing': 2, 'page_break_before': None})
        self.assertIsNone(paragraph_property(paragraph, 'page_break_before'))
