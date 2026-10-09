from io import BytesIO
from email.message import EmailMessage
from unittest.mock import patch
from zipfile import ZipFile
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.core import signing
from docx import Document
from core.models import MailboxConnection, MailboxDownload
from core.services.mailbox import polish_date, story_title, MailboxError
from core.services.mailbox_import import parse_message, package_messages, fetch_messages, mailbox_key, default_mailbox
from texts.models import Anthology, Review


def mail(uid=12, line='Julia Moskalik;Potworna Przystan;dark fantasy;15204;girl@example.com;885711216;Na pokład, psubraty'):
    msg = EmailMessage()
    msg["From"] = "other@example.com"
    msg["Reply-To"] = "reply@example.com"
    msg["Subject"] = "Julia Moskalik – Potworna Przystan"
    msg["Date"] = "Tue, 12 May 2026 10:30:00 +0200"
    msg.set_content(line + "\nTa treść nie może trafić do zgłoszenia.")
    doc = Document()
    doc.add_paragraph("Koniec. następne zdanie.")
    stream = BytesIO()
    doc.save(stream)
    msg.add_attachment(
        stream.getvalue(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="tekst.docx",
    )
    return msg.as_bytes()


class MailParsingTests(TestCase):
    @patch("core.services.mailbox_import.imaplib.IMAP4_SSL")
    def test_partial_fetch_skips_empty_attachment_and_malformed_submission(self, imap):
        empty = EmailMessage()
        empty["Reply-To"] = "empty@example.com"
        empty.set_content("Nieprawidłowe dane")
        empty.add_attachment(
            b"", maintype="application", subtype="octet-stream", filename="pusty.docx"
        )
        raws = {12: empty.as_bytes(), 13: mail(line="Brak zgłoszenia"), 14: mail()}
        client = imap.return_value
        client.select.return_value = ("OK", [b"3"])
        client.response.return_value = ("UIDVALIDITY", [b"7"])

        def fetch(command, uid, fields):
            raw = raws[int(uid)]
            if "RFC822.SIZE" in fields:
                return "OK", [f"1 (UID {uid} RFC822.SIZE {len(raw)})".encode()]
            return "OK", [(f"1 (UID {uid} BODY[] {{{len(raw)}}})".encode(), raw)]

        client.uid.side_effect = fetch
        box = MailboxConnection(host="imap.example.com", username="teksty@example.com")
        box.set_password("test")
        errors = []
        result = fetch_messages(box, 7, [12, 13, 14], errors)
        self.assertEqual([row["uid"] for row in result], [14])
        self.assertEqual(len(errors), 2)
        self.assertIn("empty@example.com", errors[0])
        self.assertIn("pusty (0 bajtów)", errors[0])
        self.assertTrue(client.select.call_args.kwargs["readonly"])
        self.assertFalse(client.store.called)
        self.assertFalse(client.expunge.called)

    def test_fields_ignore_trailing_and_body(self):
        parsed = parse_message(12, mail())
        self.assertEqual(
            parsed["record"],
            "Julia Moskalik;Potworna Przystan;dark fantasy;15204;;girl@example.com;885711216",
        )
        self.assertEqual(parsed["folder"], "Potworna Przystan")
        self.assertNotIn("IGNORUJ", str(parsed["record"]))

    def test_invalid_mail_and_multiple_records(self):
        for line in ("Brak danych", "A B;T;g;2;a@b.pl;;N\nC D;T;g;2;c@d.pl;;N"):
            with self.assertRaises(MailboxError):
                parse_message(1, mail(line=line))

    def test_polish_date_and_title(self):
        self.assertEqual(polish_date("Tue, 12 May 2026 10:30:00 +0200"), "12 maja 2026, 10:30")
        self.assertEqual(polish_date("bad"), "–")
        self.assertEqual(story_title("Autor – Tytuł – część druga"), "Tytuł – część druga")

    def test_original_and_converted_zip(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path

        row = parse_message(12, mail())
        row["folder"] = "../CON"
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            for clean, convert in ((False, False), (True, False), (False, True), (True, True)):
                with (
                    self.subTest(clean=clean, convert=convert),
                    package_messages([row], clean, convert) as out,
                    ZipFile(out) as z,
                ):
                    self.assertEqual(len(z.namelist()), 3 if convert else 1)
                    self.assertTrue(all(".." not in p.split("/") for p in z.namelist()))
                    doc = Document(
                        BytesIO(z.read(next(p for p in z.namelist() if p.endswith(".docx"))))
                    )
                    self.assertEqual(
                        doc.paragraphs[0].text,
                        "Koniec. Następne zdanie." if clean else "Koniec. następne zdanie.",
                    )

    @patch("core.services.mailbox_import.imaplib.IMAP4_SSL")
    def test_fetch_readonly_peek_and_uidvalidity(self, imap):
        raw = mail()
        client = imap.return_value
        client.select.return_value = ("OK", [b"1"])
        client.response.return_value = ("UIDVALIDITY", [b"7"])
        client.uid.side_effect = [
            ("OK", [f"1 (UID 12 RFC822.SIZE {len(raw)})".encode()]),
            ("OK", [(b"1 (UID 12 BODY[] {})", raw)]),
        ]
        box = MailboxConnection(host="imap.example.com", username="teksty@example.com")
        box.set_password("test")
        result = fetch_messages(box, 7, [12])
        self.assertEqual(result[0]["uid"], 12)
        self.assertTrue(client.select.call_args.kwargs["readonly"])
        self.assertIn("BODY.PEEK[]", client.uid.call_args.args[2])
        client.logout.assert_called_once()
        self.assertFalse(client.store.called)
        self.assertFalse(client.expunge.called)
        with self.assertRaises(MailboxError):
            fetch_messages(box, 8, [12])


class MailDownloadTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            "admin", "admin@example.com", "password"
        )
        self.client.force_login(self.user)
        self.box = MailboxConnection.objects.create(
            name="teksty",
            host="imap.example.com",
            username="teksty@example.com",
            encrypted_password="unused",
        )
        self.book = Anthology.objects.create(
            title="Na pokład, psubraty", status=Anthology.Status.IN_PREPARATION
        )
        self.url = reverse("core:review_bulk_import")
        self.selection = signing.dumps(
            {"user": self.user.pk, "mailbox": mailbox_key(self.box), "validity": 7, "uids": [12]},
            salt="mailbox-selection",
        )

    def payload(self):
        return {"action": "download", "selection": self.selection, "uids": ["12"]}

    def mixed_payload(self):
        data = self.payload()
        data.update(
            uids=["12", "13"],
            selection=signing.dumps(
                {
                    "user": self.user.pk,
                    "mailbox": mailbox_key(self.box),
                    "validity": 7,
                    "uids": [12, 13],
                },
                salt="mailbox-selection",
            ),
        )
        return data

    @patch("core.views.mailbox.fetch_messages")
    def test_invalid_data_is_excluded_and_only_valid_preview_is_committed(self, fetch):
        messages = {
            12: parse_message(12, mail()),
            13: parse_message(
                13, mail(line="Jan Test;Błędny;fantasy;0;bad@example.com;;Na pokład, psubraty")
            ),
        }
        fetch.side_effect = lambda config, validity, uids, errors: [
            messages[uid].copy() for uid in uids
        ]
        payload = self.mixed_payload()
        response = self.client.post(self.url, payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected"], [12])
        self.assertTrue(response.context["skipped_errors"])
        self.assertFalse(response.context["hard_errors"])
        self.assertFalse(Review.objects.exists())
        self.assertContains(response, "Zapisz zgłoszenia i pobierz ZIP")
        payload.update(
            action="confirm", uids=["12"], preview=response.context["preview_token"], approve="on"
        )
        response = self.client.post(self.url, payload)
        self.assertTrue(response.streaming)
        with ZipFile(BytesIO(b"".join(response.streaming_content))) as archive:
            self.assertEqual(len(archive.namelist()), 1)
        self.assertEqual(Review.objects.count(), 1)
        self.assertEqual(list(MailboxDownload.objects.values_list("uid", flat=True)), [12])

    @patch("core.views.mailbox.fetch_messages")
    def test_unknown_anthology_does_not_block_valid_submission(self, fetch):
        fetch.return_value = [
            parse_message(12, mail()),
            parse_message(
                13, mail(line="Jan Test;Inny;fantasy;1000;bad@example.com;;Nieistniejąca")
            ),
        ]
        response = self.client.post(self.url, self.mixed_payload())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected"], [12])
        self.assertIn(
            "nie znaleziono jednoznacznej antologii", response.context["skipped_errors"][0]
        )

    @patch("core.views.mailbox.fetch_messages")
    def test_conversion_failure_excludes_entire_message_and_all_its_files(self, fetch):
        from core.services.mailbox_import import _package_messages
        from core.services.document_converter import ConversionError

        good = parse_message(12, mail())
        broken = parse_message(
            13,
            mail(line="Jan Test;Błąd konwersji;fantasy;1000;bad@example.com;;Na pokład, psubraty"),
        )
        broken["files"].append(("drugi.docx", b"broken"))
        fetch.return_value = [good, broken]

        def package(messages, *args, **kwargs):
            if messages[0]["uid"] == 13:
                raise ConversionError("Błąd drugiego załącznika")
            return _package_messages(messages, *args, **kwargs)

        with patch("core.services.mailbox_import._package_messages", side_effect=package):
            response = self.client.post(self.url, self.mixed_payload())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected"], [12])
        self.assertIn("Błąd drugiego załącznika", response.context["skipped_errors"][0])
        self.assertFalse(MailboxDownload.objects.exists())

    @patch("core.views.mailbox.fetch_messages")
    def test_corrupt_docx_is_excluded_during_preview(self, fetch):
        good = parse_message(12, mail())
        broken = parse_message(
            13, mail(line="Jan Test;Uszkodzony;fantasy;1000;bad@example.com;;Na pokład, psubraty")
        )
        broken["files"] = [("uszkodzony.docx", b"this is not a ZIP archive")]
        fetch.return_value = [good, broken]
        payload = self.mixed_payload()
        payload["clean"] = "on"
        response = self.client.post(self.url, payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected"], [12])
        self.assertIn("uszkodzony.docx", response.context["skipped_errors"][0])
        self.assertFalse(Review.objects.exists())

    def test_partial_archive_keeps_equal_titles_in_separate_folders(self):
        messages = [parse_message(12, mail()), parse_message(13, mail())]
        accepted, errors = [], []
        with (
            package_messages(
                messages,
                False,
                False,
                False,
                skipped_errors=errors,
                accepted=accepted,
                rebuild_warnings=[],
            ) as output,
            ZipFile(output) as archive,
        ):
            self.assertEqual(len(archive.namelist()), 2)
            self.assertEqual(len({name.split("/")[0] for name in archive.namelist()}), 2)
        self.assertEqual([message["uid"] for message in accepted], [12, 13])
        self.assertFalse(errors)

    @patch("core.views.mailbox.fetch_messages")
    def test_new_failure_at_confirmation_requires_new_approval(self, fetch):
        from core.services.mailbox_import import _package_messages
        from core.services.document_converter import ConversionError

        messages = {
            12: parse_message(12, mail()),
            13: parse_message(
                13, mail(line="Jan Test;Drugi;fantasy;1000;second@example.com;;Na pokład, psubraty")
            ),
        }
        fetch.side_effect = lambda config, validity, uids, errors: [
            messages[uid].copy() for uid in uids
        ]
        payload = self.mixed_payload()
        preview = self.client.post(self.url, payload)
        payload.update(action="confirm", preview=preview.context["preview_token"], approve="on")

        def package(batch, *args, **kwargs):
            if batch[0]["uid"] == 13:
                raise ConversionError("Nowy błąd konwertera")
            return _package_messages(batch, *args, **kwargs)

        with patch("core.services.mailbox_import._package_messages", side_effect=package):
            response = self.client.post(self.url, payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.context["selected"], [12])
        self.assertFalse(Review.objects.exists())
        self.assertFalse(MailboxDownload.objects.exists())
        payload.update(uids=["12"], preview=response.context["preview_token"])
        response = self.client.post(self.url, payload)
        self.assertTrue(response.streaming)
        b"".join(response.streaming_content)
        self.assertEqual(Review.objects.count(), 1)

    @patch("core.views.mailbox.fetch_messages")
    def test_all_invalid_has_errors_but_no_confirm_button(self, fetch):
        fetch.return_value = [
            parse_message(
                12, mail(line="Jan Test;Błędny;fantasy;0;bad@example.com;;Na pokład, psubraty")
            )
        ]
        response = self.client.post(self.url, self.payload())
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.context["skipped_errors"])
        self.assertContains(response, "Brak poprawnych zgłoszeń", status_code=400)
        self.assertNotContains(response, "Zapisz zgłoszenia i pobierz ZIP", status_code=400)
        self.assertFalse(Review.objects.exists())

    @patch("core.views.mailbox.fetch_messages")
    def test_rebuild_omission_warning_still_requires_explicit_acceptance(self, fetch):
        from core.services.mailbox_import import _package_messages
        from core.services.document_converter import RebuildConfirmationRequired

        raw = mail()
        fetch.side_effect = lambda *args: [parse_message(12, raw)]

        def package(messages, clean, convert, rebuild, allowed, **kwargs):
            if not allowed:
                raise RebuildConfirmationRequired("Pominięta tabela")
            return _package_messages(messages, False, False, False, **kwargs)

        payload = self.payload()
        with patch("core.services.mailbox_import._package_messages", side_effect=package):
            response = self.client.post(self.url, payload)
            self.assertEqual(response.context["selected"], [12])
            self.assertTrue(response.context["rebuild_warning"])
            self.assertFalse(response.context["skipped_errors"])
            payload.update(
                action="confirm", preview=response.context["preview_token"], approve="on"
            )
            response = self.client.post(self.url, payload)
            self.assertFalse(response.streaming)
            self.assertFalse(Review.objects.exists())
            payload.update(preview=response.context["preview_token"], allow_rebuild_omissions="on")
            response = self.client.post(self.url, payload)
            self.assertTrue(response.streaming)
            b"".join(response.streaming_content)
            self.assertEqual(Review.objects.count(), 1)

    @patch("core.views.mailbox.fetch_messages")
    def test_preview_commit_retry(self, fetch):
        raw = mail()
        fetch.side_effect = lambda *args: [parse_message(12, raw)]
        payload = self.payload()
        response = self.client.post(self.url, payload)
        self.assertEqual(response.status_code, 200, response.content[:500])
        self.assertEqual(Review.objects.count(), 0)
        payload.update(action="confirm", preview=response.context["preview_token"], approve="on")
        response = self.client.post(self.url, payload)
        self.assertEqual(
            response.status_code,
            200,
            response.context.get("error") if not response.streaming else "",
        )
        data = b"".join(response.streaming_content)
        with ZipFile(BytesIO(data)) as z:
            self.assertEqual(len(z.namelist()), 1)
        self.assertEqual(Review.objects.count(), 1)
        self.assertEqual(MailboxDownload.objects.count(), 1)
        row = Review.objects.get()
        self.assertEqual(row.length, 15204)
        self.assertEqual(row.content_warnings, "")
        response = self.client.post(self.url, payload)
        self.assertTrue(response.streaming)

        # Compare archive members, not bytes: ZIP entries (also inside DOCX) carry timestamps.
        def members(payload):
            with ZipFile(BytesIO(payload)) as archive:
                return {info.filename: info.file_size for info in archive.infolist()}

        self.assertEqual(members(b"".join(response.streaming_content)), members(data))
        self.assertEqual(Review.objects.count(), 1)

    @patch("core.views.mailbox.fetch_messages")
    def test_conversion_failure_no_records(self, fetch):
        fetch.return_value = [parse_message(12, mail())]
        payload = self.payload()
        preview = self.client.post(self.url, payload).context["preview_token"]
        payload.update(action="confirm", preview=preview, approve="on")
        with patch(
            "core.views.mailbox.package_messages", side_effect=MailboxError("Niepoprawny plik")
        ):
            self.assertEqual(self.client.post(self.url, payload).status_code, 400)
        self.assertFalse(Review.objects.exists())
        self.assertFalse(MailboxDownload.objects.exists())

    @patch("core.views.mailbox.fetch_messages")
    def test_selection_tampering(self, fetch):
        data = self.payload()
        data["uids"] = ["13"]
        self.assertEqual(self.client.post(self.url, data).status_code, 400)
        fetch.assert_not_called()

    @patch("core.views.mailbox.read_headers")
    def test_default_box_and_filtered_headers(self, read):
        MailboxDownload.objects.create(mailbox_key=mailbox_key(self.box), uid_validity=7, uid=12)
        read.return_value = {"rows": [], "validity": 7, "total": 0}
        response = self.client.get(self.url)
        read.assert_not_called()
        self.assertNotContains(response, 'name="mailbox"')
        self.assertContains(response, 'name="clean" checked')
        self.client.post(self.url, {"action": "headers"})
        self.assertEqual(list(read.call_args.kwargs["excluded"](7)), [12])
        self.client.post(self.url, {"action": "headers", "show_downloaded": "on"})
        self.assertIsNone(read.call_args.kwargs["excluded"])

    def test_ambiguous_box(self):
        MailboxConnection.objects.create(
            name="TEKSTY", host="imap.example.com", username="other", encrypted_password="unused"
        )
        with self.assertRaises(MailboxError):
            default_mailbox()

    @patch("core.views.mailbox.fetch_messages")
    def test_changed_content_requires_new_preview(self, fetch):
        raw = mail()
        fetch.return_value = [parse_message(12, raw)]
        payload = self.payload()
        preview = self.client.post(self.url, payload).context["preview_token"]
        fetch.return_value = [
            parse_message(
                12,
                mail(
                    line="Julia Moskalik;Zmieniony;fantasy;100;girl@example.com;;Na pokład, psubraty"
                ),
            )
        ]
        payload.update(action="confirm", preview=preview, approve="on")
        self.assertEqual(self.client.post(self.url, payload).status_code, 400)
        self.assertFalse(Review.objects.exists())

    @patch("core.views.mailbox.fetch_messages")
    def test_non_preparation_anthology_rejected(self, fetch):
        fetch.return_value = [parse_message(12, mail())]
        self.book.status = "ready"
        self.book.save()
        self.assertEqual(self.client.post(self.url, self.payload()).status_code, 400)
        self.assertFalse(Review.objects.exists())

    @patch("core.views.mailbox.fetch_messages")
    def test_duplicate_warning_requires_confirmation_and_token(self, fetch):
        Review.objects.create(
            anthology=self.book,
            title="Potworna Przystan",
            author_first_name="JULIA",
            author_last_name="MOSKALIK",
            email="girl@example.com",
            length=15204,
            genre="dark fantasy",
        )
        raw = mail()
        fetch.side_effect = lambda *args: [parse_message(12, raw)]
        payload = self.payload()
        response = self.client.post(self.url, payload)
        self.assertTrue(response.context["forms"][0].import_warnings)
        payload.update(action="confirm", preview=response.context["preview_token"])
        response = self.client.post(self.url, payload)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Review.objects.count(), 1)
        payload.update(approve="on", preview=response.context["preview_token"])
        response = self.client.post(self.url, payload)
        self.assertTrue(response.streaming)
        with ZipFile(BytesIO(b"".join(response.streaming_content))) as archive:
            self.assertEqual(len(archive.namelist()), 1)
        self.assertEqual(Review.objects.count(), 2)

    def test_non_superuser_forbidden(self):
        user = get_user_model().objects.create_user("member", "member@example.com", "password")
        self.client.force_login(user)
        with patch("core.views.mailbox.fetch_messages") as fetch:
            self.assertEqual(self.client.post(self.url, self.payload()).status_code, 403)
            fetch.assert_not_called()
