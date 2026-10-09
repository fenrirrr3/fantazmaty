from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.test import SimpleTestCase, TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from core.models import MailboxConnection
from core.services.mailbox import read_headers, MailboxError
from core.services.document_converter import convert_document
from core.services.document_pdf import render_pdf
from core.services.document_conversion_worker import convert


class PdfWhitespaceTests(SimpleTestCase):
    def test_empty_paragraphs_survive_cleaning_and_html_conversion(self):
        doc = Document()
        doc.add_paragraph("Pierwszy.")
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph("Drugi.").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        from core.services.document_converter import run_converter

        def inspect_worker(directory, timeout):
            self.assertEqual(
                [p.text for p in Document(directory / "source.docx").paragraphs],
                ["Pierwszy.", "", "", "Drugi."],
            )
            return run_converter(directory, timeout)

        with (
            TemporaryDirectory() as tmp,
            self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)),
            patch("core.services.document_converter.run_converter", side_effect=inspect_worker),
        ):
            output, _, _ = convert_document(stream, ["pdf"], use_cleaner=True)
            with output:
                self.assertTrue(output.read().startswith(b"%PDF"))
            source = Path(tmp) / "input.docx"
            doc.save(source)
            with patch("core.services.document_pdf.render_pdf") as pdf:
                convert(source, Path(tmp), ["pdf"], "Test")
                html = pdf.call_args.args[2]
                from lxml import html as html_parser

                paragraphs = html_parser.fragment_fromstring(html, create_parent="div").xpath(
                    ".//p"
                )
                self.assertEqual(sum(not p.text_content().strip() for p in paragraphs), 2)
                self.assertIn("align-justify", html)

    def test_pdf_receives_justification_and_empty_lines(self):
        from fpdf import FPDF

        doc = Document()
        doc.add_paragraph("Justowany tekst.").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        html = "<p>Justowany tekst.</p><p></p><p></p><p>Koniec</p>"
        original = FPDF.write_html
        rendered = []

        def spy(pdf, html, **kwargs):
            rendered.append((html, pdf.y))
            return original(pdf, html, **kwargs)

        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "input.docx"
            doc.save(source)
            with patch.object(FPDF, "write_html", spy):
                render_pdf(source, Path(tmp) / "output.pdf", html, {}, "Test")
        self.assertIn('align="justify"', rendered[0][0])
        self.assertGreaterEqual(rendered[1][1] - rendered[0][1], 18.9)


class SubjectFilterTests(SimpleTestCase):
    @patch("core.services.mailbox.imaplib.IMAP4_SSL")
    def test_utf8_subject_search_precedes_pagination_and_is_readonly(self, imap):
        config = MailboxConnection(
            host="imap.example.com",
            username="teksty@example.com",
            recruitment_subjects="Na pokład, psubraty",
        )
        client = imap.return_value
        client.select.return_value = ("OK", [b"100"])
        client.response.return_value = ("UIDVALIDITY", [b"1"])
        client.uid.side_effect = [("OK", [b"1 2"]), ("OK", [])]
        with patch.object(config, "get_password", return_value="test"):
            result = read_headers(config, subject_filter="Na pokład, psubraty")
        self.assertEqual(result["total"], 2)
        args = client.uid.call_args_list[0].args
        self.assertEqual(args[:3], ("SEARCH", "CHARSET", "UTF-8"))
        self.assertEqual(args[-1], '"Nabór: „Na pokład, psubraty”"'.encode())
        client.select.assert_called_once_with('"INBOX"', readonly=True)
        self.assertIn("BODY.PEEK", client.uid.call_args_list[1].args[-1])

    def test_unknown_filter_does_not_connect(self):
        config = MailboxConnection(recruitment_subjects="Nabór A")
        with patch("core.services.mailbox.imaplib.IMAP4_SSL") as imap:
            with self.assertRaises(MailboxError):
                read_headers(config, subject_filter="inny")
            imap.assert_not_called()


class MailboxFilterViewTests(TestCase):
    def test_admin_field_and_choice_visible(self):
        from core.admin import MailboxConnectionForm

        user = get_user_model().objects.create_superuser("admin", "a@example.com", "test")
        MailboxConnection.objects.create(
            name="teksty",
            host="imap.example.com",
            username="teksty@example.com",
            recruitment_subjects="Na pokład, psubraty\nDrugi",
        )
        self.assertIn("recruitment_subjects", MailboxConnectionForm().fields)
        self.client.force_login(user)
        response = self.client.get(reverse("core:review_bulk_import"))
        self.assertContains(response, "Na pokład, psubraty")
        self.assertContains(response, 'name="subject_filter"')
