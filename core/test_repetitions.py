from io import BytesIO
from unittest.mock import patch
from django.test import SimpleTestCase, TestCase
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from core.services.document_repetitions import color_document
from core.services.odkurzacz import correct_editorial_text, clean_docx, DEFAULT_EDITORIAL_RULES
from core.services.document_converter import convert_document
from core.odkurzacz_forms import RepetitionsForm
from people.models import Person


def payload(document):
    result = BytesIO(); document.save(result); return result.getvalue()


class RepeatAnalysisTests(SimpleTestCase):
    def test_real_polish_model_flexion_and_run_preservation(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run('Kot').bold = True
        paragraph.add_run(' spaceruje obok ')
        paragraph.add_run('kota').italic = True
        paragraph.add_run('.  Bez zmian!')
        document.add_paragraph('')
        marked = Document(color_document(BytesIO(payload(document)), min_word_length=3,
            analysis_options={'long_sentences': False, 'long_paragraphs': False, 'empty_pairs': False}))
        self.assertEqual([p.text for p in marked.paragraphs], [p.text for p in document.paragraphs])
        self.assertEqual(marked.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        runs = marked.paragraphs[0].runs
        self.assertEqual(runs[0].font.color.rgb, runs[2].font.color.rgb)
        self.assertIsNotNone(runs[0].font.color.rgb)
        self.assertTrue(runs[0].bold); self.assertTrue(runs[2].italic)

    def test_cross_run_word_and_ignored_inflection(self):
        document = Document(); p = document.add_paragraph()
        p.add_run('Ko').bold = True; p.add_run('ty').italic = True
        p.add_run(' kot kota')
        data = payload(document)
        marked = Document(color_document(BytesIO(data), min_word_length=3,
            ignored_words='kota', analysis_options={'long_sentences': False, 'long_paragraphs': False}))
        self.assertEqual(marked.paragraphs[0].text, p.text)
        self.assertTrue(all(r.font.color.rgb is None and r.font.highlight_color is None for r in marked.paragraphs[0].runs))
        marked = Document(color_document(BytesIO(data), min_word_length=3,
            analysis_options={'long_sentences': False, 'long_paragraphs': False}))
        self.assertEqual(marked.paragraphs[0].runs[0].font.color.rgb, marked.paragraphs[0].runs[1].font.color.rgb)
        self.assertIsNotNone(marked.paragraphs[0].runs[0].font.color.rgb)

    def test_priority_tracked_words_and_existing_underline(self):
        document = Document(); p = document.add_paragraph()
        p.add_run('Kot kota').underline = True
        marked = Document(color_document(BytesIO(payload(document)), min_word_length=3,
            tracked_words='kot', analysis_options={'long_sentences': False, 'paragraph_limit': 1}))
        runs = marked.paragraphs[0].runs
        self.assertEqual(runs[0].font.highlight_color, WD_COLOR_INDEX.BRIGHT_GREEN)
        self.assertTrue(all(r.underline for r in runs))

    def test_extra_checks_real_sentence_parser(self):
        document = Document(); document.add_paragraph('Kot spokojnie chodzi po domu. []')
        marked = Document(color_document(BytesIO(payload(document)),
            analysis_options={'sentence_limit': 2, 'paragraph_limit': 2}))
        highlights = {r.font.highlight_color for r in marked.paragraphs[0].runs}
        self.assertIn(WD_COLOR_INDEX.TURQUOISE, highlights)
        self.assertIn(WD_COLOR_INDEX.PINK, highlights)

    def test_worker_roundtrip_and_no_editorial_changes(self):
        document = Document(); document.add_paragraph('Koty  koty...')
        output, extension, _ = convert_document(SimpleUploadedFile('test.docx', payload(document)),
            [], include_docx=True, normalize=False, repetitions={'analysis_options': {'long_sentences': False}})
        with output:
            marked = Document(output)
        self.assertEqual(extension, 'docx')
        self.assertEqual(marked.paragraphs[0].text, 'Koty  koty...')
        self.assertIsNotNone(marked.paragraphs[0].runs[0].font.color.rgb)

    def test_updated_cleaner_defaults_and_explicit_empty_paragraphs(self):
        self.assertEqual(correct_editorial_text('kot\tkota', {'tabs'}), 'kot kota')
        self.assertEqual(correct_editorial_text('1.02.2025', {'dates_times'}), '01.02.2025')
        self.assertNotIn('empty_paragraphs', DEFAULT_EDITORIAL_RULES)
        self.assertNotIn('user_word_corrections', DEFAULT_EDITORIAL_RULES)
        self.assertEqual(correct_editorial_text('np. kot. pies! – powiedział.', {'sentence_case'}), 'np. kot. Pies! – powiedział.')
        document = Document(); document.add_paragraph('A'); document.add_paragraph(); document.add_paragraph(); document.add_paragraph('B')
        marked = Document(clean_docx(BytesIO(payload(document)), ['empty_paragraphs']))
        self.assertEqual([p.text for p in marked.paragraphs], ['A', '', 'B'])
        marked = Document(clean_docx(BytesIO(payload(document)), DEFAULT_EDITORIAL_RULES))
        self.assertEqual(len(marked.paragraphs), 4)

    def test_form_rejects_bad_parameters_and_defaults_match_desktop(self):
        document = Document(); document.add_paragraph('Tekst')
        form = RepetitionsForm({'window_size': '0', 'min_word_length': '4', 'sentence_limit': '35', 'paragraph_limit': '150'},
            {'document': SimpleUploadedFile('test.docx', payload(document))})
        self.assertFalse(form.is_valid()); self.assertIn('window_size', form.errors)
        self.assertTrue(RepetitionsForm().fields['duplicates'].initial)


class RepeatPageTests(TestCase):
    def test_member_download_and_separate_form(self):
        user = get_user_model().objects.create_user('reader', 'reader@example.com')
        Person.objects.create(user=user, first_name='Jan', last_name='Test', email=user.email)
        self.client.force_login(user)
        response = self.client.get(reverse('core:programs'))
        self.assertContains(response, 'Kolorowanie powtórzeń')
        self.assertContains(response, 'name="repetitions-window_size"')
        document = Document(); document.add_paragraph('Koty koty')
        with patch('core.views.programs.convert_document', return_value=(BytesIO(payload(document)), 'docx', 'application/octet-stream')) as convert:
            response = self.client.post(reverse('core:programs'), {
                'program_action': 'repetitions', 'repetitions-document': SimpleUploadedFile('test.docx', payload(document)),
                'repetitions-window_size': 35, 'repetitions-min_word_length': 4,
                'repetitions-sentence_limit': 35, 'repetitions-paragraph_limit': 150})
        self.assertEqual(response.status_code, 200)
        self.assertIn('_powtorzenia.docx', response['Content-Disposition'])
        self.assertFalse(convert.call_args.kwargs['normalize'])
        self.assertNotIn('use_cleaner', convert.call_args.kwargs)
        response.close()
