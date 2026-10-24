"""Updated editorial rules from the desktop cleaner, with web limits."""
import re
from functools import lru_cache
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if __package__:
    from .document_progress import report as report_progress
    from .document_stories import document_stories
else:
    from document_progress import report as report_progress
    from document_stories import document_stories

# Reguły są wykonywane w ustalonej kolejności, niezależnie od kolejności kliknięć.
EDITORIAL_RULES = (
    ('invisible_chars', 'Niewidoczne znaki z kopiowania (zero-width space, miękki łącznik, BOM, znaki kierunku)'),
    ('spaces', 'Podwójne i wielokrotne spacje'),
    ('trim', 'Spacje na początku i końcu akapitu'),
    ('tabs', 'Tabulatory i mieszanki tabulatorów ze spacjami → jedna spacja'),
    ('before_punct', 'Spacje przed znakami interpunkcyjnymi'),
    ('after_punct', 'Brakujące spacje po interpunkcji (z wyjątkami)'),
    ('inside_brackets', 'Spacje wewnątrz nawiasów'),
    ('inside_quotes', 'Spacje wewnątrz cudzysłowów'),
    ('outside_brackets', 'Brakujące spacje wokół nawiasów'),
    ('quotes', 'Proste, angielskie i inne cudzysłowy → polskie „…”, cytat w cytacie «…»'),
    ('apostrophe', 'Apostrof prosty wewnątrz wyrazu → typograficzny ’'),
    ('ellipsis', 'Wielokropek → …'),
    ('duplicate_punct', 'Powtórzone przecinki, średniki, dwukropki i podwójne kropki'),
    ('quote_punct', 'Kropka i przecinek zawsze za cudzysłowem zamykającym: „tak”.'),
    ('note_references', 'Odsyłacze przypisów wg polskich norm: bez spacji przed odsyłaczem; przed kropką, '
                        'przecinkiem, średnikiem i dwukropkiem; za cudzysłowem i nawiasem zamykającym, '
                        'pytajnikiem, wykrzyknikiem i wielokropkiem'),
    ('hyphen_dash', 'Łącznik użyty jako myślnik → półpauza'),
    ('dash_style', 'Pauza, pauza cyfrowa, kreska pozioma i minus → półpauza –; łączniki Unicode → dywiz -; '
                   'powtórzone myślniki → jeden'),
    ('dash_spaces', 'Odstępy przy myślnikach i początku kwestii dialogowej'),
    ('range_dash', 'Łącznik w zakresach liczbowych → półpauza'),
    ('range_spaces', 'Usuwanie spacji wewnątrz zakresów liczbowych'),
    ('unit_nbsp', 'Jedna zwykła spacja między liczbą a jednostką: 5 kg; procent bez spacji: 5%'),
    ('unit_space', 'Brakująca spacja między liczbą a jednostką; m2 → m²'),
    ('reference_nbsp', 'Jedna zwykła spacja w zapisach: 2025 r., s. 15, nr 3'),
    ('reference_space', 'Brakująca spacja: 2025r., s.15, nr3, ok.5; błędne „nr.” → nr'),
    ('initials', 'Odstępy przy inicjałach i nazwisku'),
    ('abbreviations', 'Usuń spacje wewnątrz skrótów: m. in. → m.in., t. j. → tj., t. zw. → tzw., '
                      'n. p. m. → n.p.m., p. p. m. → p.p.m., p. n. e. → p.n.e.'),
    ('dates_times', 'Daty DD.MM.RRRR i godziny HH:MM (jednoznaczne zapisy)'),
    ('decimal', 'Przecinek dziesiętny przy jednostkach i kwotach'),
    ('digit_groups', 'Grupowanie dużych liczb spacją: 12 345'),
    ('temperature', 'Temperatura: 20 °C / 20 °F / 20 K'),
    ('pronouns_lower', 'Zaimki osobowe małą literą – poza początkiem zdania i akapitu'),
    ('user_word_corrections', 'Własne zamiany słownikowe (słownik zamian w CMS)'),
)
# Grupa „Ortografia 2026”: zmiany zasad pisowni obowiązujące od 1 stycznia 2026 r. (komunikat RJP
# z 10 maja 2024 r., zmiany 3, 6, 9a, 10 i 11) – tylko zapisy rozpoznawalne bez rozumienia zdania.
ORTHOGRAPHY_2026_RULES = (
    ('by_separate', 'Ortografia 2026: -by osobno po spójnikach i partykułach (czyby → czy by, alebym → ale bym; '
                    'aby, żeby, gdyby, jakby, niby bez zmian)'),
    ('niby_quasi', 'Ortografia 2026: niby, quasi, pseudo łącznie przed małą literą, z łącznikiem przed wielką '
                   '(niby-artysta → nibyartysta, quasiAnglia → quasi-Anglia)'),
    ('prefix_hyphen', 'Ortografia 2026: łącznik po przedrostku przed wielką literą (superPolak → super-Polak, '
                      'antyPiS → anty-PiS)'),
    ('pol_compounds', 'Ortografia 2026: pół żartem, pół serio → półżartem, półserio; pół Polak, pół Szkot → '
                      'pół-Polak, pół-Szkot („pół godziny” bez zmian)'),
    ('nie_comparative', 'Ortografia 2026: „nie” łącznie ze stopniem wyższym i najwyższym (nie najlepszy → '
                        'nienajlepszy, nie lepiej niż → nielepiej niż; zapisy sporne bez zmian)'),
)
ORTHOGRAPHY_2026_KEYS = frozenset(key for key, _ in ORTHOGRAPHY_2026_RULES)
EDITORIAL_RULES += ORTHOGRAPHY_2026_RULES + (
    ('sentence_case', 'Wielka litera po . ? ! – z wyjątkami dla skrótów, liczb i dialogów'),
    ('empty_paragraphs', 'Wielokrotne puste akapity → jeden pusty akapit'))
ALL_EDITORIAL_RULES = frozenset(key for key, _ in EDITORIAL_RULES)
DEFAULT_EDITORIAL_RULES = ALL_EDITORIAL_RULES - {'pronouns_lower', 'user_word_corrections', 'empty_paragraphs'}
# Formularz Odkurzacza zaznacza domyślnie także zaimki i zamiany słownikowe;
# DEFAULT_EDITORIAL_RULES (bez nich) pozostaje dla automatycznego czyszczenia zgłoszeń.
# Reguły „Ortografia 2026” są w obu zestawach (decyzja właściciela: pewne, domyślnie włączone).
FORM_DEFAULT_EDITORIAL_RULES = ALL_EDITORIAL_RULES - {'empty_paragraphs'}
# Decyzja właściciela: Odkurzacz nigdzie nie wstawia spacji nierozdzielających,
# a istniejące (U+00A0, wąska U+202F, cyfrowa U+2007) zawsze zamienia na zwykłe spacje.
NONBREAKING_SPACES = '\u00a0\u202f\u2007'
# Nietypowe spacje (en, em, cienka, ideograficzna…) także stają się zwykłymi.
UNUSUAL_SPACES = NONBREAKING_SPACES + ''.join(map(chr, range(0x2000, 0x200b))) + '\u205f\u3000\u1680'
NONBREAKING_TABLE = str.maketrans(dict.fromkeys(UNUSUAL_SPACES, ' '))
UNUSUAL_SPACES_RE = re.compile('[' + UNUSUAL_SPACES + ']')
# Znaki bez szerokości z kopiowania (strony WWW, PDF, czytniki): łamią wyszukiwanie i słownik.
INVISIBLE_CHARS = '\u200b\u200c\u200d\u2060\ufeff\u00ad\u200e\u200f\u2061\u2062\u2063\u2064'
INVISIBLE_RE = re.compile('[' + INVISIBLE_CHARS + ']')
# Reguła dash_style: nietypowe kreski → półpauza, łączniki Unicode → dywiz.
DASH_TABLE = str.maketrans({'\u2014': '–', '\u2015': '–', '\u2012': '–', '\u2212': '–', '\u2010': '-', '\u2011': '-'})
# Akapit złożony wyłącznie ze znaków separatora (* * *, - - -, -----) zostaje bez zmian.
SEPARATOR_CHARS = frozenset('-–—*#~_=. \t')
# Skrótowce z kropkami, które nie są inicjałami (nie rozdzielamy ich spacjami).
DOTTED_ACRONYM_RE = re.compile(r'(?<!\w)(?:S\.A\.|P\.S\.|P\.P\.S\.|U\.S\.A\.|N\.N\.|S\.J\.|S\.C\.|A\.D\.|R\.I\.P\.)(?!\w)')
# Data z miesiącem rzymskim (5.X.2024) – chroniona jak data, bez zmiany zapisu.
ROMAN_DATE_RE = re.compile(r'(?<![\w.])\d{1,2}\.[IVX]{1,4}\.\d{4}(?!\w)')
# Skróty, po których zwykle stoi liczba (ok.5 → ok. 5, rys.3 → rys. 3); „nr” i „pkt” nie mają kropki.
NUMBER_ABBREVIATIONS = (r'(?:(?:s|t|cz|poz|rozdz|art|ust|par|zd|godz|str|ok|ur|zm|zał|sygn|wyd|vol|'
                        r'rys|ryc|il|fot|tab)\.|pkt)')
MARKER_CLASS = '\U000F0000-\U000FFFFD'
H = r'[ \u00a0\u202f\u2007]'
LETTERS = 'A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż'
LOWER = 'a-ząćęłńóśźż'
UPPER = 'A-ZĄĆĘŁŃÓŚŹŻ'
UNITS = (r'(?:km/h|m/s|kg|mg|µg|μg|km|dm|cm|mm|µm|nm|ml|mL|'
         r'kHz|MHz|GHz|Hz|kPa|MPa|Pa|kN|kJ|kW|MW|kV|mA|kΩ|'
         r'kB|MB|GB|TB|ms|min|ha|bar|PLN|zł|EUR|USD|'
         r'[kmc]m[²³]|m[²³]|[gtmlLshNJWTAVΩBK])')
# Bez spacji pojedyncza wielka litera to zwykle oznaczenie (klasa 2B, 4K, 3A),
# a nie jednostka, więc reguła „brakująca spacja” jej nie rozdziela.
ATTACHED_UNITS = '(?:' + UNITS.replace('[gtmlLshNJWTAVΩBK]', '[gtmlsh]') + r'|tys\.|mln|mld|proc\.|[kmcd]?m[23])'
QUANTITIES = UNITS + r'|osób|szt\.?|egzemplarzy|mieszkańców|punktów'
NUMBER = r'(?<![\w.,])[-−]?\d+(?:[.,]\d+)?'

PERSONAL_PRONOUNS = frozenset('''
ja mnie mi mną ty ciebie cię tobie ci tobą
on jego go niego jemu mu niemu nim
ona jej niej ją nią ono je nie
my nas nam nami wy was wam wami
oni one ich nich im nimi
'''.split())
PRONOUN_RE = re.compile(r'\b(?:' + '|'.join(
    sorted({form.capitalize() for form in PERSONAL_PRONOUNS}
           | {form.upper() for form in PERSONAL_PRONOUNS}, key=lambda value: (-len(value), value))
) + r')\b')


# Zwroty grzecznościowe (listy, dedykacje, modlitwy): wielka litera w środku zdania jest zamierzona.
HONORIFIC_PRONOUNS = frozenset('Ty Ciebie Cię Tobie Ci Tobą Wy Was Wam Wami'.split())
OPENING_QUOTES = '„«‹‚'
AMBIGUOUS_QUOTES = '"“\'‘'


def _after_opening_quote(text, start):
    """Wyraz tuż po cudzysłowie otwierającym (tytuł, cytat: „Ja, robot”)."""
    window = max(0, start - 4)
    prefix = text[window:start].rstrip(' \u00a0\u202f\u2007')
    if not prefix:
        return False
    if prefix[-1] in OPENING_QUOTES:
        return True
    # Prosty cudzysłów otwiera cytat, gdy nie stoi tuż po literze lub cyfrze.
    position = window + len(prefix) - 1
    return prefix[-1] in AMBIGUOUS_QUOTES and (position == 0 or not text[position - 1].isalnum())


PRONOUN_CONTINUATION_RE = re.compile(r'\b(?:np\.|m\.in\.|tj\.|tzw\.|por\.|zob\.)$', re.IGNORECASE)
WORD_RE = re.compile(r'\w+')
# Kontekst tylko z końca prefiksu: reguły patrzą na kilka znaków, a całe
# prefiksy długich akapitów dawały czas kwadratowy względem długości akapitu.
CONTEXT = 80
OPENING_CHARS = ' \t\r\u00a0\u202f„“”"\'«»()[]{}–—-'


def _significant_prefix(text, start):
    """Koniec bieżącej linii przed `start`, bez znaków okalających pierwsze słowo."""
    line_start = text.rfind('\n', 0, start) + 1
    window_start = max(line_start, start - CONTEXT)
    significant = text[window_start:start].rstrip(OPENING_CHARS)
    if not significant and window_start > line_start:
        significant = text[line_start:start].rstrip(OPENING_CHARS)
    return significant


def _lower_personal_pronouns(text):
    """Reguła pisowni form z listy, bez rozstrzygania ich funkcji gramatycznej.

    Chroni początki zdań, akapitów i wypowiedzi po dwukropku. Przy skrótach
    mogących kończyć zdanie zachowuje wielką literę. Po granicy nietekstowego
    XML również zachowujemy pierwszy zaimek, bo brakuje pełnego kontekstu.
    Formy wersalikami (JEGO, ONA) zostają: to wyróżnienie albo napis.
    """
    if PRONOUN_RE.search(text) is None:
        return text
    protected_spans = [(match.start(), match.end()) for match in TECHNICAL_RE.finditer(text)]
    span_index = 0

    def replace(match):
        nonlocal span_index
        while span_index < len(protected_spans) and protected_spans[span_index][1] <= match.start():
            span_index += 1
        if span_index < len(protected_spans) and protected_spans[span_index][0] < match.end():
            return match[0]
        if match[0].isupper() and len(match[0]) > 1:
            return match[0]  # wersaliki to wyróżnienie albo napis – zostają (decyzja właściciela)
        if match[0] in HONORIFIC_PRONOUNS or _after_opening_quote(text, match.start()):
            return match[0]
        significant = _significant_prefix(text, match.start())
        if not significant or significant[-1] in '!?…:':
            return match[0]
        # Te skróty zazwyczaj wprowadzają dalszy ciąg wypowiedzi.
        if significant.endswith('.') and PRONOUN_CONTINUATION_RE.search(significant) is None:
            return match[0]
        return match[0].lower()

    return PRONOUN_RE.sub(replace, text)

# Chronimy adresy, ścieżki oraz identyfikatory przed wszystkimi regułami tekstowymi.
TECHNICAL_RE = re.compile(
    r'https?://[^\s<>„”"]+|www\.[^\s<>„”"]+|'
    r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|'
    r'\b[A-Za-z]:\\[^\r\n<>"|]+|'
    r'\b(?:[a-z0-9-]+\.)+(?-i:pl|com|org|net|edu|gov|eu|io|co|de|uk)\b'
    r'(?:[/?#][^\s<>„”"]*)?|'
    # Nazwy techniczne z kropkami (moduł.klasa.pole) – co najmniej trzy człony małymi literami.
    r'\b(?-i:[a-z_][a-z0-9_]*(?:\.[a-z_][a-z0-9_]*){2,})\b|'
    # Nazwy technologii z kropką i wersalikami po niej (ASP.NET, Socket.IO); Node.js chroni FILE_NAME_RE.
    r'\b(?-i:[A-Za-z][\w+#]*\.(?:NET|IO))\b|'
    r'\b(?:\d{1,3}\.){3}\d{1,3}\b|'
    # Kod pocztowy przed nazwą miejscowości (00-950 Warszawa) – nie jest zakresem.
    r'\b\d{2}-\d{3}(?=[ \u00a0]+(?-i:[' + UPPER + r']))|'
    r'\b[vV]?\d+(?:\.\d+){2,}\b|'
    r'\b(?:ISBN|ISSN|PESEL|NIP|REGON|IBAN|tel\.|telefon|konto|kod)'
    r'[ \t:]*[+\d][\d \u00a0./-]*\d', re.IGNORECASE)
DATE_RE = re.compile(
    r'(?<![\w.])(?:(?P<iso_y>\d{4})-(?P<iso_m>\d{1,2})-(?P<iso_d>\d{1,2})|'
    r'(?P<d>\d{1,2})(?P<sep>[./-])(?P<m>\d{1,2})(?P=sep)(?P<y>\d{4}))(?!\w|\.\d)')
TIME_RE = re.compile(r'(?<![\w:])\d{1,2}:\d{2}(?::\d{2})?(?![\w:])')
# Zakres godzin: 8:00-16:00, 8.00 – 16.00. Zapis z kropką uznajemy za godziny tylko przy pełnych
# kwadransach i bez jednostki po nim (1.50-2.50 zł, 2.10-3.10 to nie godziny).
TIME_RANGE_RE = re.compile(
    r'(?<![\w.,:])(?P<h1>\d{1,2})(?P<s1>[:.])(?P<m1>\d{2})(?P<left>' + H + r'*)(?P<dash>[-–—])'
    r'(?P<right>' + H + r'*)(?P<h2>\d{1,2})(?P<s2>[:.])(?P<m2>\d{2})(?![\w:]|[.,]\d)')



def _time_range(match, enabled, counts):
    """Zakres godzin: oba końce jak godzina (HH:MM), półpauza bez spacji – według włączonych reguł.

    None, gdy to nie zakres godzin (kwota, wersja) – wtedy działają zwykłe reguły liczbowe.
    """
    hours, minutes = (int(match['h1']), int(match['h2'])), (match['m1'], match['m2'])
    if max(hours) > 24 or max(map(int, minutes)) > 59 or match['s1'] != match['s2']:
        return None
    if match['s1'] == '.':
        after = match.string[match.end():match.end() + 30]
        if any(value not in ('00', '15', '30', '45') for value in minutes) or QUANTITY_AFTER_RE.match(after) \
                or re.match(H + r'*(?:%|°)', after):
            return None
    changes = {}
    ends = [match['h1'] + match['s1'] + match['m1'], match['h2'] + match['s2'] + match['m2']]
    if 'dates_times' in enabled:
        normalized = [f'{hour:02d}:{minute}' for hour, minute in zip(hours, minutes)]
        changes['dates_times'] = sum(a != b for a, b in zip(ends, normalized))
        ends = normalized
    dash, left, right = match['dash'], match['left'], match['right']
    if 'range_dash' in enabled and dash != '–':
        dash, changes['range_dash'] = '–', 1
    if 'range_spaces' in enabled and (left or right):
        left, right, changes['range_spaces'] = '', '', bool(left) + bool(right)
    if counts is not None:
        for key, number in changes.items():
            if number:
                counts[key] = counts.get(key, 0) + number
    return ends[0] + left + dash + right + ends[1]


def _normalize_dates_times(text):
    from datetime import date

    def date_repl(match):
        year = int(match['iso_y'] or match['y'])
        month = int(match['iso_m'] or match['m'])
        day = int(match['iso_d'] or match['d'])
        try:
            date(year, month, day)
        except ValueError:
            return match[0]
        return f'{day:02d}.{month:02d}.{year:04d}'

    text = DATE_RE.sub(date_repl, text)
    # Kropka w godzinie jest rozpoznawana tylko po „godz.” lub „o”.
    text = re.sub(r'\b(godz\.?|o)(' + H + r'+)(\d{1,2})\.(\d{2})(?!\d)',
                  lambda m: f'{m[1]}{m[2]}{int(m[3]):02d}:{m[4]}'
                  if int(m[3]) < 24 and int(m[4]) < 60 else m[0], text)

    def time_repl(match):
        parts = match[0].split(':')
        if int(parts[0]) < 24 and all(int(part) < 60 for part in parts[1:]):
            return ':'.join(f'{int(part):02d}' for part in parts)
        return match[0]

    return TIME_RE.sub(time_repl, text)


# Własna lista użytkownika (lista właściciela). To jawne zamiany redakcyjne, a nie ocena
# poprawności przez słownik Worda; wyłączenie opcji wyłącza całą poniższą listę.
# Warianty docelowe: przekonujący, oddziałujący; „tere-fere” z łącznikiem (nie z półpauzą).
# Fleksja korzysta z jawnych końcówek, ponieważ model językowy nie musi rozpoznawać lematów
# wyrazów zapisanych błędnie. Gdy każdy wyraz zaczynający się od wadliwego tematu jest błędny
# (absorbcja, absorbcyjny…), wzorzec bierze cały wyraz (\w*); gdy obok są poprawne wyrazy
# (niedowierzanie, przysięgły, kreolski, Centkiewicz), końcówki są wyliczone.
USER_WORD_EXACT = {
    'niechcąco': 'niechcący',
    'zaczym': 'za czym',
    'władnie': 'władny',
    'niewiadomo': 'nie wiadomo',
    'terefere': 'tere-fere',
    'inąd': 'skądinąd',
    'wszechczasów': 'wszech czasów',
    'zapewnie': 'zapewne',
    'dellarte': 'dell’arte',
    'księżami': 'księżmi',
    'oldboi': 'oldbojów',
    # nabazgrać: nabazgrzę, nabazgrzesz… (decyzja właściciela; czas przeszły nabazgrał jest poprawny).
    'nabazgram': 'nabazgrzę',
    'nabazgrasz': 'nabazgrzesz',
    'nabazgra': 'nabazgrze',
    'nabazgramy': 'nabazgrzemy',
    'nabazgracie': 'nabazgrzecie',
    'nabazgrają': 'nabazgrzą',
}
# Końcówki rzeczownika żeńskiego na -ka (Azjatka, Mulatka…) po temacie bez „k” i końcówki przymiotnika.
FEMININE_KA = r'k(?:a|i|ę|ą|o|om|ami|ach)|ce|ek'
ADJECTIVE_ENDINGS = r'(?:y|a|e|ego|emu|ej|ą|ym|ymi|ych|i)'


# Skróty z kropką, po których zdanie zwykle trwa dalej. Kropki wewnątrz skrótu nie dostają spacji
# (after_punct), po jego kropce nie wstawiamy wielkiej litery (sentence_case), a odsyłacz przypisu
# zostaje za kropką skrótu (note_references). Wielkość liter nie ma znaczenia (dz.u. = Dz.U.), a skrót
# wieloczłonowy pasuje także z odstępami (p. o., Dz. U.) – istniejący zapis zostaje bez zmian.
# Podkreślnik to spacja w skrócie wielowyrazowym (z_o.o. = „z o.o.”). Wpis o kilku znaczeniach stoi
# w jednej grupie (por. – porównaj i porucznik, pt. – pod tytułem i piątek).
# Celowo pomijamy skróty równe wyrazom, którymi może kończyć się zdanie – po nich nowe zdanie dalej
# dostaje wielką literę: kat., im., cel., nauk., prac., farm., miejsc., rzecz., czas., sport., bot.,
# mat., praw., sąd., post., ser., aut., druk., map., sztab., marsz., pułk., bat., bomb., port., rum.,
# nor., tur., serb., hol., bar., wyż., niż., hm., upadł., tir., a także pr. i hr. (PR, HR) oraz nazwy
# miesięcy (lut., lip., lis., sie.). Pojedynczych liter (b., k., n., ż.) nie dopisujemy: po małej
# literze z kropką sentence_case i tak nie zaczyna zdania.
ABBREVIATION_GROUPS = (
    ('ogólne', """
        np. m.in. tj. tzn. tzw. itd. itp. in. inn. nast. poprz. ew. ewent. prawdop. ok. ca. ww. w.w. jw.
        j.w. cd. cdn. c.d. c.d.n. zob. por. patrz. czyt. cyt. przyp. dot. odn. ds. nt. ws. wg. zał. odp.
        uw. nb. wzgl. zw. zwł. zazw. ozn. zn. tel. wew. b.u. c.o. etc. vs. cf.
    """),
    ('tytuły, stopnie i funkcje', """
        p. pp. ob. sz. szan. sz.p. prof. hab. doc. dr. mgr. inż. lic. lek. stom. dent. med. wet. mec. adw.
        r.pr. dypl. nadzw. dziek. prodz. koresp. czł. kol. tow. red. dyr. wicedyr. kier. zast. przew. sekr.
        prez. min. amb. sen. jun. asyst. ref. nacz. burm. p.o. w.z. z_up.
    """),
    ('kościelne', """
        ks. o. oo. ss. św. śś. bł. śp. ś.p. kard. abp. bp. pw. p.w. prob. wik. kan. prał. diec. archidiec.
        kośc. rz.-kat. gr.-kat. ewang. prawosł.
    """),
    ('wojskowe i służbowe', """
        gen. adm. kontradm. wiceadm. bryg. dyw. płk. ppłk. mjr. kmdr. kpt. rtm. ppor. pchor. chor. st. mł.
        sierż. wachm. ogn. plut. kpr. szer. szt. sekc. asp. kom. podkom. nadkom. insp. podinsp. nadinsp.
    """),
    ('adresowe', """
        ul. al. pl. os. bud. bl. lok. m. pok. kl. wyb. bulw. skw. zauł. przys. woj. pow. gm. m.st. skr.
        poczt. pn. pd. płn. płd. wsch. zach.
    """),
    ('firmy i instytucje', """
        sp. sp.j. sp.k. sp.p. o.o. z_o.o. sp._z_o.o. s.c. likw. oddz. zakł. przeds. spółdz. stow. fund.
        inst. akad. uniw. wydz.
    """),
    ('prawo i urzędy', """
        art. ust. par. lit. zd. poz. k.c. k.k. k.p. k.w. k.m. k.p.a. k.p.c. k.p.k. k.p.w. k.r.o. k.r._i_o.
        k.k.s. k.k.w. k.s.h. k.p.s.w. dz.u. dz.urz. m.p. t.j. jedn. późn. rozp. uchw. orz. wyr. zb. sygn.
        l.dz. spr. prot. zarz. obw. cyw. karn. wiecz.
    """),
    ('bibliografia i wydawnictwa', """
        s. str. t. tt. cz. rozdz. z. zesz. nr. vol. rocz. wyd. wydawn. nakł. oprac. przeł. tłum. przekł.
        wst. posł. prow. popr. uzup. rozsz. przejrz. il. ilustr. rys. ryc. tab. tabl. wykr. fot. bibliogr.
        dz. kk. b.d. b.m. b.m.w. b.r.w. i_n. ibid. sic. op._cit. loc._cit. et_al.
    """),
    ('czas i miary', """
        r. w. b.r. ub.r. r.b. ub.m. bm. br. ub. n.e. p.n.e. godz. sek. tyg. mies. kw. pon. wt. śr. czw. pt.
        sob. niedz. nd. ur. zm. dn. tys. proc. maks. egz. kpl. opak. temp. wys. dł. gł. obj. poj. n.p.m.
        p.p.m.
    """),
    ('językowe', """
        pot. posp. książk. podn. poet. przen. dosł. właśc. daw. przestarz. rzad. arch. hist. żart. iron.
        pej. lekcew. wulg. ekspr. zdr. zdrobn. zgrub. środ. reg. gw. urz. techn. jęz. wym. lp. lm.
        ndm. ndk. dk. przym. przysł. przyim. spój. wykrz. zaim. muz. zool. chem. fiz. biol. geogr. geol.
        filoz. ekon. polit. rel. wojsk. lotn. szk. pol. ang. niem. fr. franc. ros. łac. gr. wł. hiszp.
        czes. słow. ukr. białorus. łot. węg. fiń. szw. duń. bułg. chorw. hebr. chiń. stpol.
    """),
)
SENTENCE_ABBREVIATIONS = frozenset(
    entry.replace('_', ' ') for _, entries in ABBREVIATION_GROUPS for entry in entries.split())


def _abbreviation_pattern(value):
    # Optional spaces belong only between abbreviation parts, never after its
    # final full stop. Otherwise masking also consumes the sentence's spacing.
    parts = value.split('.')
    pattern = ''
    for index, part in enumerate(parts):
        pattern += re.escape(part).replace(r'\ ', H + '+')
        if index < len(parts) - 1:
            pattern += r'\.'
            if index < len(parts) - 2:
                pattern += H + '*'
    return pattern


def _abbreviation_alternatives(values):
    """Wzorce skrótów pogrupowane według pierwszej litery, w grupie dłuższy przed krótszym (p.o. przed p.).

    Przy kilkuset skrótach regex w każdym miejscu tekstu sprawdza wtedy kilkadziesiąt liter,
    a nie wszystkie skróty po kolei.
    """
    groups = {}
    for value in sorted(values, key=lambda x: (-len(x), x)):
        groups.setdefault(value[0], []).append(_abbreviation_pattern(value[1:]))
    return '|'.join(re.escape(first) + '(?:' + '|'.join(rest) + ')' for first, rest in sorted(groups.items()))


# Ciąg pojedynczych liter z kropkami (n.p.m., k.p.c., U.S.A.) jest sprawdzany najpierw, żeby
# jednoliterowy skrót z listy (p., s., w.) nie odcinał początku: p.p. zostaje, a nie „p. p.”.
CAPITALIZATION_ABBREVIATION_RE = re.compile(
    r'(?<![^\W\d])(?:\b(?:[' + UPPER + r']\.){2,}|' + _abbreviation_alternatives(SENTENCE_ABBREVIATIONS) + ')',
    re.IGNORECASE)
ABBREVIATION_RE = CAPITALIZATION_ABBREVIATION_RE
# Reguła abbreviations: tylko jednoznaczne skróty, których rozstrzelony zapis jest błędem. „t.j.” przed
# Dz.U. to tekst jednolity, a nie „tj.”. Pozostałe zapisy ze spacją (p. o., Dz. U., b. r.) zostają.
ABBREVIATION_MERGES = tuple((re.compile(pattern), replacement) for pattern, replacement in (
    (r'\bm\.' + H + r'*in\.', 'm.in.'),
    (r'\bt\.' + H + r'*j\.(?!' + H + r'*[Dd]z\.)', 'tj.'),
    (r'\bt\.' + H + r'*zw\.', 'tzw.'),
    (r'\bn\.' + H + r'*p\.' + H + r'*m\.', 'n.p.m.'),
    (r'\bp\.' + H + r'*p\.' + H + r'*m\.', 'p.p.m.'),
    (r'\bp\.' + H + r'*n\.' + H + r'*e\.', 'p.n.e.'),
))
FILE_NAME_RE = re.compile(r'(?<![\w.])[\w-]+(?:\.[\w-]+)*\.(?:docx?|pdf|epub|mobi|txt|rtf|odt|xlsx?|csv|zip|rar|jpg|jpeg|png|gif|py|html?|js|css)(?!\w)', re.IGNORECASE)


INITIAL_BEFORE_RE = re.compile(r'(?<!\w)[' + UPPER + r']$')
TIME_BEFORE_RE = re.compile(r'(?<![\w:.])\d{1,2}:\d{2}(?::\d{2})?$')


DIGIT_RE = re.compile(r'\d')
# Maski to znaki z prywatnego obszaru Unicode (patrz mask_value).
MARKER_RE = re.compile('[\U000F0000-\U000FFFFD]')
SENTENCE_START_RE = re.compile(r'([.!?]+)(' + H + r'*)([' + LOWER + r'])')


def _merged(spans):
    merged = []
    for begin, end in sorted(spans):
        if merged and begin <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        elif begin < end:
            merged.append([begin, end])
    return merged


def _covers(spans, starts, position):
    from bisect import bisect_right
    index = bisect_right(starts, position) - 1
    return index >= 0 and position < spans[index][1]


def _capitalize_sentences(text):
    # Zwykle poprawny tekst nie ma małej litery po kropce – wtedy nie ma czego chronić.
    if SENTENCE_START_RE.search(text) is None:
        return text
    # Chronione fragmenty jako posortowane przedziały (sprawdzane tylko dla kandydatów).
    spans = []
    for pattern in (TECHNICAL_RE, DATE_RE, TIME_RE, FILE_NAME_RE):
        for match in pattern.finditer(text):
            # URLs can include trailing punctuation. Protect only their body.
            end = match.end()
            if pattern is TECHNICAL_RE:
                end = match.start() + len(match[0].rstrip('.,;:!?)]}'))
            spans.append((match.start(), end))
    spans = _merged(spans)
    starts = [begin for begin, _ in spans]
    abbreviations = [(match.start(), match.end()) for match in CAPITALIZATION_ABBREVIATION_RE.finditer(text)
                     if not _covers(spans, starts, match.start())]
    spans = _merged([*map(tuple, spans), *abbreviations])
    starts = [begin for begin, _ in spans]

    def protected(position):
        return _covers(spans, starts, position)

    def replace(match):
        marks, gap, letter = match.groups()
        start = match.start()
        if any(protected(pos) for pos in range(start, start + len(marks))):
            return match[0]
        if protected(match.start(3)):
            return match[0]  # adres lub nazwa techniczna zaczynająca zdanie
        following = text[match.end():match.end() + 2]
        if following[:1] == ')' or (following[:1] == '.' and not following[1:2].isalnum()):
            return match[0]  # punkt listy: a) b) albo c.
        if '.' in marks:
            if len(marks) != 1 or (start and text[start - 1] in '.…'):
                return match[0]  # ellipsis, not a sentence boundary
            if start and text[start - 1].islower() and (start < 2 or not text[start - 2].isalnum()):
                return match[0]  # punkt listy „c. trzy”
            if start and text[start - 1].isdigit():
                if TIME_BEFORE_RE.search(text, max(0, start - 12), start) is None:
                    return match[0]  # ordinal, version, date
                # Kropka po godzinie (o 9:30.) kończy zdanie.
            if INITIAL_BEFORE_RE.search(text, max(0, start - 1), start):
                return match[0]  # an initial
        return marks + gap + letter.upper()

    # Deliberately do not cross a dash, quotation mark or parenthesis.
    # In particular '. – powiedział' / '? – zapytała' stay untouched.
    return SENTENCE_START_RE.sub(replace, text)


# ---- Ortografia 2026 (reguły działają po przywróceniu masek, więc same omijają adresy i nazwy plików) ----

def _alternatives(words):
    return '|'.join(sorted(words.split(), key=lambda word: (-len(word), word)))


def _orthography_sub(pattern, replace, text):
    """pattern.sub poza adresami, nazwami technicznymi i plików; replace(match) zwraca nowy zapis albo None."""
    spans = starts = None

    def guarded(match):
        nonlocal spans, starts
        if spans is None:
            spans = _merged([(found.start(), found.end()) for protected in (TECHNICAL_RE, FILE_NAME_RE)
                             for found in protected.finditer(text)])
            starts = [begin for begin, _ in spans]
        if _covers(spans, starts, match.start()) or _covers(spans, starts, match.end() - 1):
            return match[0]
        result = replace(match)
        return match[0] if result is None else result
    return pattern.sub(guarded, text)


# Zmiana 3: cząstka -by (-bym, -byś, -byśmy, -byście) osobno po spójnikach i partykułach – zamknięta lista
# połączeń pisanych dotąd łącznie. Zrosty aby, ażeby, żeby, iżby, gdyby, oby, niby, jakby, jakoby, coby, byleby,
# choćby, chociażby, czyżby, czyliżby, gdzieżby, jakżeby, niżby, bodajby i niechby (niechajby, niechżeby) zostają,
# niejednoznaczne toby, jeśliby, jeżeliby, aniżeliby, niżeliby i niźliby – też. „Boby” (bób) i „nimby” (nimb)
# to rzeczowniki, więc po „bo” i „nim” rozdzielamy tylko dłuższe końcówki (bobym → bo bym).
BY_SEPARATE_RE = re.compile(
    r'(?<![\w-])(?P<word>' + _alternatives(
        'ale ależ albo albowiem ani aż bowiem chyba czy czyli dopóki jednak jednakże jeszcze już lecz lub może '
        'natomiast nawet nuż oraz ponieważ póki przecież przeto skoro tedy też toteż tylko więc wszak wszakże '
        'zanim zatem zaś') + r'|bo(?=byś|bym)|nim(?=byś|bym))(?P<by>by(?:śmy|ście|m|ś)?)(?![\w-])', re.IGNORECASE)


def _separate_by(text):
    return _orthography_sub(BY_SEPARATE_RE, lambda match: f"{match['word']} {match['by']}", text)


# Zmiana 10: niby, quasi (i pseudo, pisane łącznie już od 1957 r.) łącznie z wyrazem pisanym małą literą,
# z łącznikiem przed wielką literą (nibyPolak → niby-Polak). Wyraz dalej z łącznikiem (niby-e-book) zostaje.
NIBY_RE = re.compile(r'(?<![\w-])(?P<prefix>(?i:niby|quasi|pseudo))(?:[-\u2010\u2011](?P<word>[' + LETTERS + r']+)'
                     r'|(?P<capital>[' + UPPER + r'][\w-]*))(?![\w-])')
# Zmiana 9a: przedrostek przed wyrazem pisanym wielką literą – z łącznikiem (superPolak → super-Polak,
# antyPiS → anty-PiS, antyNATO → anty-NATO, postPRL-u → post-PRL-u). Zapis od wielkiej litery (MiniMax,
# EkoPark) to zwykle nazwa własna – zostaje; „pan” i „para” to także zwykłe wyrazy (brak spacji: panKowalski),
# więc ich nie ma na liście.
PREFIX_CAPITAL_RE = re.compile(
    r'(?<![\w-])(?P<prefix>' + _alternatives(
        'anty arcy eks ekstra eko hiper kontr makro maksi mega midi mikro mini multi neo post pra pre pro super '
        'turbo ultra wege wice') + r')(?P<capital>[' + UPPER + r'][\w-]*)(?![\w-])')
# Przymiotnik na -owy od nazwy (postSolidarnościowy) piszemy małą literą i łącznie – nie zgadujemy, zostaje.
NAME_ADJECTIVE_RE = re.compile(r'[' + UPPER + r'][' + LOWER + r']+ow(?:y|a|e|ego|ej|emu|ym|ą|ych|ymi|i)')
# Utrwalone terminy i nazwy zapisywane łącznie (złącza, nośniki, technologie, marki) – bez łącznika, także
# z końcówką po łączniku (miniUSB-a). Wersje z micro- (microSD, microUSB) i tak nie pasują do listy przedrostków.
FIXED_PREFIX_TERMS = frozenset('''
    mikroRNA miniDV miniUSB mikroUSB miniSD miniSDHC mikroSD mikroSDHC mikroSDXC miniHDMI mikroHDMI miniDP
    miniPCIe miniLED mikroLED proALPHA
'''.split())


def _capital_after_prefix(match):
    """Łącznik między przedrostkiem pisanym małą literą a wyrazem od wielkiej litery (nie przymiotnikiem na -owy
    ani utrwalonym terminem z FIXED_PREFIX_TERMS)."""
    if (not match['prefix'][0].islower() or NAME_ADJECTIVE_RE.fullmatch(match['capital'])
            or match['prefix'] + match['capital'].split('-')[0] in FIXED_PREFIX_TERMS):
        return None
    return f"{match['prefix']}-{match['capital']}"


# Człon po łączniku wyglądający na formę czasownika zostaje (niby-żyję, quasi-myślę, niby-spał) – decyzja
# właściciela: lepiej zostawić też rzeczownik w bierniku lub narzędniku (quasi-naukę) niż zepsuć czasownik.
VERB_LIKE_RE = re.compile(r'\w*(?:ę|esz|isz|ysz|emy|imy|ymy|ecie|icie|ycie|ą|ał|ała|ało|ali|ały|ać|eć|ić|yć)')


def _niby_quasi(text):
    def replace(match):
        if match['capital'] is not None:
            return _capital_after_prefix(match)
        word = match['word']
        # Z łącznikiem zostaje wyraz pisany wielką literą (niby-Polak); wersaliki całego zapisu to nie nazwa.
        if word[0].isupper() and not (match['prefix'].isupper() and word.isupper()):
            return None
        if VERB_LIKE_RE.fullmatch(word.lower()):
            return None
        return match['prefix'] + word
    return _orthography_sub(NIBY_RE, replace, text)


# Decyzja koordynatora: przymiotnik na -ski, -cki, -dzki po przedrostku zostaje (proRosyjski, antyPolski:
# poprawnie prorosyjski, antypolski, ale to może być też nazwisko – nie zgadujemy).
NAME_SKI_RE = re.compile(r'[' + UPPER + r'][' + LOWER + r']*(?:(?:sk|ck|dzk)(?:i|a|ie|iego|iej|iemu|im|ą|ich|imi)'
                         r'|(?:s|c|dz)cy)')


def _prefix_hyphen(text):
    def replace(match):
        # Pojedyncza litera (superX, miniB) to zwykle oznaczenie lub nazwa – zostaje.
        if len(match['capital'].split('-')[0]) < 2 or NAME_SKI_RE.fullmatch(match['capital']):
            return None
        return _capital_after_prefix(match)
    return _orthography_sub(PREFIX_CAPITAL_RE, replace, text)


# Zmiana 6: pół łącznie w wyrażeniach półżartem, półserio, z łącznikiem przed nazwą osoby pisaną wielką literą
# (pół-Polka, pół-Francuzka). „Pół” ilościowe zostaje (pół godziny, pół Polski, pół Europy), a półzabawa
# czy półspał wymagałyby rozumienia zdania – rozpoznajemy tylko te wyrażenia i zamkniętą listę nazw narodowości
# w mianowniku i narzędniku (był pół-Polakiem, pół-Szkotem).
POL_EXPRESSION_RE = re.compile(r'(?<![\w-])(?P<pol>(?i:pół))[ \t]+(?P<word>(?i:żartem|serio))(?![\w-])')
POL_PEOPLE = (
    'Afrykanin Amerykanin Anglik Arab Argentyńczyk Australijczyk Austriak Azjata Belg Białorusin Brazylijczyk Bułgar '
    'Chińczyk Chorwat Cygan Czech Duńczyk Estończyk Europejczyk Fin Francuz Góral Grek Gruzin Hindus Hiszpan Holender '
    'Indianin Irlandczyk Islandczyk Japończyk Kanadyjczyk Kaszub Koreańczyk Kreol Litwin Łotysz Meksykanin Mulat '
    'Niemiec Norweg Ormianin Persjanin Polak Portugalczyk Rosjanin Rumun Serb Słowak Słoweniec Szkot Szwajcar Szwed '
    'Ślązak Tatar Turek Ukrainiec Walijczyk Węgier Wietnamczyk Włoch Żyd '
    'Afrykanka Amerykanka Angielka Arabka Argentynka Australijka Austriaczka Azjatka Belgijka Białorusinka Brazylijka '
    'Bułgarka Chinka Chorwatka Cyganka Czeszka Dunka Estonka Europejka Finka Francuzka Góralka Greczynka Gruzinka '
    'Hinduska Hiszpanka Holenderka Indianka Irlandka Islandka Japonka Kanadyjka Kaszubka Koreanka Kreolka Litwinka '
    'Łotyszka Meksykanka Mulatka Niemka Norweżka Ormianka Persjanka Polka Portugalka Rosjanka Rumunka Serbka '
    'Słowaczka Słowenka Szkotka Szwajcarka Szwedka Ślązaczka Tatarka Turczynka Ukrainka Walijka Węgierka Wietnamka '
    'Włoszka Żydówka').split()
# Narzędnik nieregularny; pozostałe: -a → -ą, -k/-g → -kiem/-giem, inne → -em.
POL_PEOPLE_INSTRUMENTAL = {'Niemiec': 'Niemcem', 'Ukrainiec': 'Ukraińcem', 'Słoweniec': 'Słoweńcem',
                           'Turek': 'Turkiem', 'Węgier': 'Węgrem', 'Holender': 'Holendrem'}
# Te same zapisy to dopełniacz nazw krain (pół Niemiec, pół Włoch, pół Kaszub) – osoba tylko w parze z inną.
POL_AMBIGUOUS_PEOPLE = frozenset('Niemiec Czech Włoch Węgier Kaszub'.split())
POL_PEOPLE_FORMS = frozenset(POL_PEOPLE) | frozenset(
    POL_PEOPLE_INSTRUMENTAL.get(name) or (name[:-1] + 'ą' if name.endswith('a') else
                                          name + ('iem' if name[-1] in 'kg' else 'em')) for name in POL_PEOPLE)
POL_COMMON_RE = re.compile(r'(?<![\w-])(?i:pół)[ \t]+[' + LOWER + r']')
POL_PERSON_RE = re.compile(r'(?<![\w-])(?P<pol>pół|Pół)[ \t]*(?P<person>'
                           + _alternatives(' '.join(POL_PEOPLE_FORMS)) + r')(?![\w-])')


def _pol_compounds(text):
    from bisect import bisect_left

    def expression(match):
        if match['word'][0].isupper() and not match['pol'].isupper():
            return None
        return match['pol'] + match['word']

    text = _orthography_sub(POL_EXPRESSION_RE, expression, text)
    certain = [match.start() for match in POL_PERSON_RE.finditer(text)
               if match['person'] not in POL_AMBIGUOUS_PEOPLE]

    def same_sentence(positions, start):
        nearby = positions[bisect_left(positions, start - 59):bisect_left(positions, start + 60)]
        return any(set(text[min(other, start):max(other, start)]).isdisjoint('.!?…;\n') for other in nearby)

    common = [match.start() for match in POL_COMMON_RE.finditer(text)]

    def person(match):
        # Zapis niejednoznaczny – tylko w tym samym zdaniu co pewna nazwa osoby (pół Niemiec, pół Polak);
        # para z wyrazem pospolitym („pół Polak, pół byk”) jest sporna – zostaje.
        if match['person'] in POL_AMBIGUOUS_PEOPLE and not same_sentence(certain, match.start()):
            return None
        if same_sentence(common, match.start()):
            return None
        return f"{match['pol']}-{match['person']}"
    return _orthography_sub(POL_PERSON_RE, person, text)


# Zmiana 11: „nie” łącznie z przymiotnikami i przysłówkami odprzymiotnikowymi także w stopniu wyższym
# i najwyższym. Decyzja właściciela: zmieniamy tylko zapisy jednoznaczne, każdy sporny zostaje. Stopień
# rozpoznajemy po budowie wyrazu: przymiotniki na -ejszy i inne na -szy według tematów (lepszy, większy, tańszy,
# milszy, nowszy…) oraz naj- + taki przymiotnik; przysłówki: lepiej, gorzej, prędzej (przykłady RJP), -niej
# (z -źniej), -ściej i krótka lista na -ciej (szybciej) – także z naj-. Ogólne -iej/-ej to często tryb
# rozkazujący (śmiej, siej, wiej, miej, wyśmiej, zasiej) – nie zmieniamy. Bez zmian zostają też miara, czas,
# miejsce i znaczenia niestopniowe (dalej, bliżej, wcześniej, później, dawniej, dłużej, krócej, wyżej, niżej;
# dalszy, bliższy, wcześniejszy, późniejszy, dłuższy, krótszy, wyższy, niższy, starszy, młodszy), liczebniki
# i przysłówki nieodprzymiotnikowe (mniej, więcej, bardziej, najpierw) oraz wyrazy o podobnej budowie
# (dzisiejszy, ostatniej, średniej, trzeciej, złodziej, turniej).
DEGREE_ENDINGS = r'(?:sz(?:y|a|e|ego|ej|emu|ym|ą|ych|ymi)|si)'
COMPARATIVE_RE = re.compile(
    r'(?:\w{2,}ej|\w*(?:[iy]st|[iy]t|[aeiuy]l|ń|ow)|' + _alternatives(
        'blad bled bogat brzyd błah chud chyż cich cięż czyst częst droż dzik elegant gor goręt grub gęst gład głuch '
        'głup głęb kruch lep lich mięk mędr pręd prost płyt rzad skąp sroż such szczer szer szyb słab słod tward '
        'tęp tęż tłust uboż więk węż zacięt zawzięt zażart śwież święt żółt ciekaw jaskraw łaskaw łzaw plugaw żwaw '
        'prawdziw leniw poczciw uczciw chciw godziw sędziw właściw żyw natarczyw uporczyw fałszyw cierpliw '
        'szczęśliw sprawiedliw możliw wrażliw życzliw złośliw troskliw gorliw żarliw') + r')' + DEGREE_ENDINGS
    # Inne tematy na -liw: bez formy -liwszy (domyśliwszy, określiwszy to imiesłowy).
    + r'|\w*[^\Waeiouyąęó]liw(?:sz(?:a|e|ego|ej|emu|ym|ą|ych|ymi)|si)'
    + r'|\w{2,}niej|\w+ściej|' + _alternatives(
        'lepiej gorzej prędzej szybciej bogaciej otwarciej uparciej obficiej pracowiciej znakomiciej przyzwoiciej'))
# Wyrazy o budowie stopnia wyższego, które nim nie są albo są sporne (przysłówki taniej, cieniej, biedniej –
# jednoznaczne; czerwieniej, zieleniej to też tryb rozkazujący – zostają).
NOT_COMPARATIVE_RE = re.compile(
    r'mniej|wcześniej|później|dawniej|naj(?:wyraźniej|widoczniej|prawdopodobniej|pewniej)'
    r'|(?:\w*(?:dzisi|jutrz|gdysi|ongisi|tamt|tut|teraźni)|ichni|nini|(?:z|u|po|poz)mni|dal|wcześni|późni|dawni)'
    r'(?:ej)?' + DEGREE_ENDINGS
    + r'|(?!(?:taniej|cwaniej|cieniej|szaleniej|uczeniej|biedniej|odpowiedniej|wredniej)$)'
    r'\w*(?:[aeiouyąęó]niej|edniej|letniej|ostatniej|wschodniej|zachodniej|wierzchniej|turniej)'
    r'|przechodniej|przewodniej|przyrodniej|dodatniej|sobotniej|dożywotniej|armatniej|herbatniej|bratniej|żytniej'
    r'|zbytniej|ichniej|bliźniej|spodniej|zadniej|wyżniej|niżniej|istniej|zaistniej|bynajmniej|przynajmniej|czerniej'
    r'|\w*niebieściej|zniewieściej')
NIE_DEGREE_RE = re.compile(r'(?<![\w-])(?P<nie>nie|Nie|NIE)[ \t]+(?P<word>[^\W\d_]+)(?![\w-])')
# Bez zmian: zaimek „nie” po przyimku (patrzył na nie łaskawiej); „nie” po „wcale”, „bynajmniej”, „ani”,
# w „…, a nie X”, po „czy” (czy nie przyjemniejsza), w warunku (jeśli, jeżeli, gdyby, o ile – „nie” przeczy
# pominiętemu orzeczeniu: jeśli nie gorzej) i w stanie pośrednim (ale nie, ale i nie, choć nie).
NIE_KEPT_AFTER = frozenset('na w we za przez o po pod nad przed między pomiędzy ponad poza wcale bynajmniej ani a '
                           'czy czyż jeśli jeżeli gdyby ale lecz choć chociaż'.split())
NIE_KEPT_AFTER_PAIRS = frozenset({('o', 'ile'), ('ale', 'i'), ('ale', 'też'), ('ale', 'także'), ('lecz', 'też'),
                                  ('choć', 'też'), ('a', 'też')})
PREVIOUS_WORDS_RE = re.compile(r'(?:(\w+)[ \t]+)?(\w+)[ \t]*$')
# Całe zdanie pytające („Czy to nie najlepszy pomysł?”, „A nie lepiej…?”) – pytajnik w tym samym zdaniu.
QUESTION_REST_RE = re.compile(r'[^.!?\n]{0,200}\?')
# Przeciwstawienie tuż dalej w zdaniu: „nie lepszy, lecz gorszy”, „nie wyższy od niej, tylko najwyższy”
# (ale nie stopniowanie i ciąg dalszy: „…, a nawet”, „…, a potem”, „…, a ja”).
NIE_CONTRAST_RE = re.compile(
    r'[^.!?;:…\n]{0,30}?[\s,](?:ale|lecz|tylko|jeno|a|raczej)[ \t]+(?!(?:nie|też|także|również|nawet|wręcz|i|'
    r'przecież|jednak|potem|później|wtedy|teraz|zarazem|ponadto|do|w|we|na|z|ze|po|przy|od|że|to|co|kiedy|gdy|'
    r'jak|ja|ty|on|ona|ono|my|wy|oni|one|się)\b)\w', re.IGNORECASE)
# Stan pośredni: „nie najlepszy, ale (i) nie najgorszy”, „nie lepszy, choć nie gorszy”.
NIE_BETWEEN_RE = re.compile(r'[^.!?;:…\n]{0,40}?[\s,](?:ale|lecz|choć|chociaż|ani)'
                            r'(?:[ \t]+(?:i|też|także|również))?[ \t]+nie\b', re.IGNORECASE)


def _comparative(word):
    """Przymiotnik albo przysłówek odprzymiotnikowy w stopniu wyższym lub najwyższym (po budowie wyrazu)."""
    word = word.lower()
    if NOT_COMPARATIVE_RE.fullmatch(word):
        return False  # najwyraźniej, najprawdopodobniej, najpewniej – w funkcji partykuły
    if word.startswith('naj') and len(word) > 5:
        word = word[3:]
    return COMPARATIVE_RE.fullmatch(word) is not None and NOT_COMPARATIVE_RE.fullmatch(word) is None


def _nie_comparative(text):
    def replace(match):
        nie, word = match['nie'], match['word']
        if (word.isupper() and len(word) > 1) != (nie == 'NIE') or (nie != 'NIE' and not word.islower()):
            return None
        if not _comparative(word):
            return None
        previous = PREVIOUS_WORDS_RE.search(text, max(0, match.start() - 30), match.start())
        if previous:
            pair = (previous[1] or '').lower(), previous[2].lower()
            if pair[1] in NIE_KEPT_AFTER or pair in NIE_KEPT_AFTER_PAIRS:
                return None
        if (QUESTION_REST_RE.match(text, match.end()) or NIE_CONTRAST_RE.match(text, match.end())
                or NIE_BETWEEN_RE.match(text, match.end())):
            return None
        return nie + word
    return _orthography_sub(NIE_DEGREE_RE, replace, text)


ORTHOGRAPHY_2026_STEPS = (('by_separate', _separate_by), ('niby_quasi', _niby_quasi),
                          ('prefix_hyphen', _prefix_hyphen), ('pol_compounds', _pol_compounds),
                          ('nie_comparative', _nie_comparative))


def _orthography_2026(text, enabled, counts):
    for key, step in ORTHOGRAPHY_2026_STEPS:
        if key in enabled:
            before = text
            text = step(text)
            if counts is not None:
                _count(counts, key, before, text)
    return text



def _replacement_case(source, replacement):
    if source.isupper():
        return replacement.upper()
    if source[:1].isupper():
        return replacement[:1].upper() + replacement[1:]
    return replacement


# Każda para zawiera pełny wzorzec wyrazu i jego docelową postać.
# Granice wyrazu zapobiegają zmianom wewnątrz innych słów.
USER_WORD_PATTERNS = tuple((re.compile(pattern, re.IGNORECASE), replacement)
    for pattern, replacement in (
        (r'\btonicem\b', r'tonikiem'),
        (r'\btonicy\b', r'toniki'),
        (r'\btonic(?P<ending>owi|ów|om|ami|ach|iem|a|u|i)?\b', r'tonik\g<ending>'),
        (r'\bmentlik(?P<ending>owi|ów|om|ami|ach|iem|a|u|i)?\b', r'mętlik\g<ending>'),
        # -ywuj- jest błędne w każdej formie: przekonywujący, przekonywująco, przekonywuje, oddziaływują…
        (r'\b(?P<neg>nie)?przekonywuj(?P<ending>\w*)', r'\g<neg>przekonuj\g<ending>'),
        (r'\b(?P<neg>nie)?oddziaływuj(?P<ending>\w*)', r'\g<neg>oddziałuj\g<ending>'),
        # zadawalać (błędne w każdej formie): zadawalający, zadawalająco, zadawala, zadawalał…
        (r'\b(?P<neg>nie)?zadawala(?P<ending>\w*)', r'\g<neg>zadowala\g<ending>'),
        # Czasownik: miałczeć, miałczę, miałczała, miałczeli, imiesłowy, miałczenie; zamiałczał, pomiałczeć.
        (r'\b(?P<prefix>za|po|roz)?miałcz(?P<ending>\w*)', r'\g<prefix>miaucz\g<ending>'),
        # Zmieniamy wadliwy bezokolicznik oraz formy z wadliwym „-leli”.
        # Poprawne „wymyślał”, „wymyślę”, „wymyślenie” pozostają bez zmian.
        (r'\bwymyśle(?P<ending>ć|liśmy|liście|li|libyśmy|libyście|liby)\b',
         r'wymyśli\g<ending>'),
        (r'\bdomyśle(?P<ending>ć|liśmy|liście|li|libyśmy|libyście|liby)\b',
         r'domyśli\g<ending>'),
        (r'\bdopatrzeć\b', r'dopatrzyć'),
        (r'\bdopatrza(?P<ending>łem|łam|łeś|łaś|ł|ła|ło|łyśmy|łyście|ły|'
         r'łbym|łbyś|łby|łabym|łabyś|łaby|łoby|łybyśmy|łybyście|łyby)\b',
         r'dopatrzy\g<ending>'),
        (r'\bdopatrze(?P<ending>liśmy|liście|li|libyśmy|libyście|liby)\b',
         r'dopatrzy\g<ending>'),
        (r'\bkarnister\b', r'kanister'),
        (r'\bkarnistr(?P<ending>owi|ów|om|ami|ach|em|ze|a|y)\b', r'kanistr\g<ending>'),
        (r'\bmassmedi(?P<ending>ami|ach|ów|om|a)\b', r'mass medi\g<ending>'),
        # Nazwy mieszkańców: wielka litera także w formach odmienionych.
        (r'\beuropejczyk(?P<ending>owi|ów|om|ami|ach|iem|a|u)?\b',
         r'Europejczyk\g<ending>'),
        (r'\beuropejczycy\b', r'Europejczycy'),
        (r'\bazjat(?P<ending>ami|ach|ów|om|a|y|ę|ą|o|' + FEMININE_KA + r')\b', r'Azjat\g<ending>'),
        (r'\bazjaci(?P<ending>e)?\b', r'Azjaci\g<ending>'),
        (r'\beuropej(?P<ending>' + FEMININE_KA + r')\b', r'Europej\g<ending>'),
        (r'\bkreol(?P<ending>a|owi|em|u|e|owie|ów|om|ami|ach|' + FEMININE_KA + r')?\b', r'Kreol\g<ending>'),
        (r'\bmula(?P<ending>t|ta|towi|tem|cie|ci|tów|tom|tami|tach|ty|tk(?:a|i|ę|ą|o|om|ami|ach)|tce|tek)\b',
         r'Mula\g<ending>'),
        # Rzeczownik i przymiotnik: absorbcja, absorbcyjny; mitingowy; cętka, cętkowany (nie Centkiewicz).
        (r'\babsorbc(?P<ending>\w+)', r'absorpc\g<ending>'),
        (r'\bmiting(?P<ending>\w*)', r'mityng\g<ending>'),
        (r'\bcent(?P<ending>k(?:a|i|ę|ą|o|om|ami|ach|owan' + ADJECTIVE_ENDINGS + r')|ce|ek)\b', r'cęt\g<ending>'),
        (r'\bkurestw(?P<ending>\w*)', r'kurewstw\g<ending>'),
        (r'\bindywidulan(?P<ending>\w*)', r'indywidualn\g<ending>'),
        (r'\bprzeciwskaza(?P<ending>\w*)', r'przeciwwskaza\g<ending>'),
        (r'\bśredniozaawansowan(?P<ending>\w*)', r'średnio zaawansowan\g<ending>'),
        # Przymiotnik (grekokatolicki, grekokatoliccy); rzeczownik grekokatolik jest poprawny.
        (r'\bgrekokatolic(?P<ending>[kc]\w*)', r'greckokatolic\g<ending>'),
        # Czasowniki: wszystkie formy, imiesłowy i rzeczowniki odczasownikowe.
        (r'\bwyksztu(?P<ending>[sś]\w*)', r'wykrztu\g<ending>'),
        (r'\bzarzadz(?P<ending>\w*)', r'zarządz\g<ending>'),
        (r'\b(?P<neg>nie)?zcentraliz(?P<ending>\w*)', r'\g<neg>scentraliz\g<ending>'),
        # Decyzja właściciela: nie ma formy dokonanej „skategoryzować” – wszystkie jej formy → kategoryzować.
        (r'\b(?P<neg>nie)?skategoryz(?P<ending>\w*)', r'\g<neg>kategoryz\g<ending>'),
        # Czas przeszły przysiąc; przysięgła, przysięgły (też przymiotnik) są poprawne.
        (r'\b(?P<prefix>za)?przysięgł(?P<ending>em|eś|by|bym|byś)?\b', r'\g<prefix>przysiągł\g<ending>'),
        # nie dowierzać – osobno; niedowierzanie i niedowierzający (imiesłów przymiotnikowy) – razem.
        (r'\bniedowierza(?P<ending>ć|m|sz|my|cie|ją|j|jmy|jcie|jąc|ł\w*|l\w*)?\b', r'nie dowierza\g<ending>'),
    ))
COMMEDIA_RE = re.compile(
    r'\b(?P<noun>komedi(?:ami|ach|om|a|i|ę|ą|o|e))'
    r'(?P<space>[ \t\u00a0\u202f]+)(?P<term>dellarte)\b', re.IGNORECASE)


# Słownik zamian jako dane (do wypełnienia słownika w CMS). Kolejność ma znaczenie:
# kolejne wpisy „word” działają razem jako jedna alternatywa, wzorce – po kolei.
DEFAULT_REPLACEMENTS = (
    [{'kind': 'word', 'source': source, 'target': target} for source, target in USER_WORD_EXACT.items()]
    + [{'kind': 'pattern', 'source': pattern.pattern, 'target': target} for pattern, target in USER_WORD_PATTERNS]
    + [{'kind': 'pattern', 'source': COMMEDIA_RE.pattern, 'target': r'\g<noun>\g<space>dell’arte'}]
)
# Ochrona przed kosztownymi wpisami: moduł re nie ma limitu czasu, więc odrzucamy długie wpisy,
# wzorce z zagnieżdżonym powtórzeniem w rodzaju (a+)+, (a?){22} lub (a|b)* (czas wykładniczy)
# oraz wzorce, które na próbkach złośliwego tekstu działają zbyt wolno (np. \w*\w*\w*\w*!).
MAX_REPLACEMENTS = 1000
MAX_PATTERN_LENGTH = 500
MAX_TARGET_LENGTH = 200
PATTERN_TIME_LIMIT = 0.05  # sekundy na jedną próbkę
# Do 512 znaków: wzorce kwadratowe (\w+ść) mieszczą się z zapasem także przy obciążonym serwerze,
# a sześcienne i gorsze odpadają najpóźniej przy 256 znakach (szacunek wzrostu dla 512).
PROBE_SIZES = (16, 32, 64, 128, 256, 512)
PROBE_UNITS = ('a', 'ą', 'Ż', '1', ' ', '\t', 'ab', 'a ', 'a.', 'x-', 'aB1 ')
_TOKEN_SPLIT_RE = re.compile(r'(\s+)')
_GROUP_REFERENCE_RE = re.compile(r'\\(?:g<[^>]*>|\d+)')
WORD_GAP = '[ \t\u00a0\u202f\u2007]+'


def _apply_replacement_case(source, value, keep):
    """Wielkość liter wyniku zamiany.

    keep (cel zawiera wielkie litery: NATO, iPhone, Polak): zapis celu zostaje, chyba że źródło
    jest całe wielkimi literami – wtedy wynik też. W przeciwnym razie wynik dostaje wielkość
    liter źródła (mentlik → mętlik, Mentlik → Mętlik, MENTLIKIEM → MĘTLIKIEM).
    """
    if keep:
        return value.upper() if source.isupper() else value

    def single(source_part, value_part):
        return _replacement_case(source_part, value_part.lower())

    source_tokens = _TOKEN_SPLIT_RE.split(source)
    value_tokens = _TOKEN_SPLIT_RE.split(value)
    if len(source_tokens) == 1 or len(source_tokens) != len(value_tokens):
        return single(source, value)
    # Kilka wyrazów (komedia dellarte): wyraz, którego zamiana nie zmienia,
    # zostaje w zapisie źródła; pozostałe dostają wielkość liter swojego odpowiednika.
    result = []
    for index, (source_part, value_part) in enumerate(zip(source_tokens, value_tokens)):
        if index % 2 or source_part.lower() == value_part.lower():
            result.append(value_part if index % 2 else source_part)
        else:
            result.append(single(source_part, value_part))
    return ''.join(result)


def _keeps_target_case(target, pattern=False):
    literal = _GROUP_REFERENCE_RE.sub('', target) if pattern else target
    return literal != literal.lower()


def _nested_repeat(parsed, inside_repeat=False):
    """Czy wzorzec ma powtórzenie lub alternatywę wewnątrz powtórzenia (ryzyko czasu wykładniczego)."""
    from re import _constants as const
    repeats = (const.MAX_REPEAT, const.MIN_REPEAT, const.POSSESSIVE_REPEAT)
    for op, value in parsed:
        if op in repeats:
            low, high, sub = value
            # Także opcjonalny element w powtórzeniu: (a?){22} ma 2^22 sposobów dopasowania.
            if inside_repeat:
                return True
            if _nested_repeat(sub, high > 1):
                return True
        elif op is const.BRANCH:
            if inside_repeat:
                return True
            if any(_nested_repeat(branch, inside_repeat) for branch in value[1]):
                return True
        elif op is const.SUBPATTERN:
            if _nested_repeat(value[-1], inside_repeat):
                return True
        elif op in (const.ASSERT, const.ASSERT_NOT):
            if _nested_repeat(value[1], inside_repeat):
                return True
        elif op is const.ATOMIC_GROUP:
            if _nested_repeat(value, inside_repeat):
                return True
        elif op is const.GROUPREF_EXISTS:
            if any(item is not None and _nested_repeat(item, inside_repeat) for item in value[1:]):
                return True
    return False


class _PatternTimeout(Exception):
    pass


def _exhaust(pattern, sample, seconds):
    """Przejdź wszystkie dopasowania; w głównym wątku (worker, uWSGI) z twardym limitem czasu.

    Moduł re sprawdza sygnały w trakcie dopasowania, więc SIGALRM przerywa nawet bardzo wolny
    wzorzec. Poza głównym wątkiem (albo gdy timer jest już zajęty) zostaje samo szacowanie.
    """
    import signal
    import threading
    guarded = (hasattr(signal, 'setitimer') and threading.current_thread() is threading.main_thread()
               and signal.getsignal(signal.SIGALRM) is not None and not signal.getitimer(signal.ITIMER_REAL)[0])
    if not guarded:
        for _ in pattern.finditer(sample):
            pass
        return
    armed = True

    def interrupt(signum, frame):
        if armed:
            raise _PatternTimeout

    previous = signal.signal(signal.SIGALRM, interrupt)
    try:
        signal.setitimer(signal.ITIMER_REAL, seconds)
        for _ in pattern.finditer(sample):
            pass
    finally:
        armed = False
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def _too_slow(pattern):
    r"""Próba czasu na złośliwych napisach (długie ciągi liter, cyfr, spacji) o rosnącej długości.

    Wzorzec, który na próbce przekracza PATTERN_TIME_LIMIT, jest odrzucany. Zanim uruchomimy
    dwa razy dłuższą próbkę, szacujemy jej czas z tempa wzrostu – wzorzec wielomianowo wolny
    (np. \w*\w*\w*\w*!) odpada, zanim zajmie sekundy; w głównym wątku pomiar dodatkowo
    przerywa SIGALRM (_exhaust).
    Pomiar podejrzanie wolny powtarzamy (najwyżej 3 razy, liczy się najkrótszy), żeby
    chwilowe obciążenie serwera nie odrzucało poprawnych wzorców.
    """
    from time import perf_counter

    def run(sample):
        started = perf_counter()
        try:
            _exhaust(pattern, sample, 2 * PATTERN_TIME_LIMIT)
        except _PatternTimeout:
            return float('inf')
        return perf_counter() - started

    def slow(elapsed, previous, last):
        if elapsed > PATTERN_TIME_LIMIT:
            return True
        growth = elapsed / previous if previous and previous > 0.0002 else 8
        return not last and elapsed * max(growth, 2) > PATTERN_TIME_LIMIT

    for unit in PROBE_UNITS:
        for suffix in ('', '!\n'):
            previous = None
            for size in PROBE_SIZES:
                sample = unit * (size // len(unit)) + suffix
                last = size == PROBE_SIZES[-1]
                elapsed = run(sample)
                for _ in range(2):
                    # Wynik kilka razy ponad limit to nie chwilowe obciążenie – bez powtórek.
                    if not slow(elapsed, previous, last) or elapsed > 2 * PATTERN_TIME_LIMIT:
                        break
                    elapsed = min(elapsed, run(sample))
                if slow(elapsed, previous, last):
                    return True
                previous = elapsed
    return False


def _compile_pattern(source, target='', *, probe=True):
    """Skompilowany wzorzec albo None, gdy wpis jest błędny lub zbyt kosztowny.

    Publiczne dla formularza słownika zamian (walidacja wpisu „wzorzec”). probe=False pomija
    próbę czasu (wpis sprawdzony już przy zapisie w CMS); kontrola składni i zagnieżdżeń zostaje.
    """
    from re import _parser
    if not isinstance(source, str) or not isinstance(target, str) or not source \
            or len(source) > MAX_PATTERN_LENGTH or len(target) > MAX_TARGET_LENGTH:
        return None
    try:
        pattern = re.compile(source, re.IGNORECASE)
        if _nested_repeat(_parser.parse(source, re.IGNORECASE)):
            return None
        pattern.sub(target, '')  # sprawdza odwołania \g<nazwa> w targecie
    except (re.error, IndexError, TypeError, ValueError, RecursionError, OverflowError):
        return None
    # Wzorzec dopasowujący pusty napis wstawiałby target między każde dwa znaki.
    if pattern.fullmatch('') is not None or (probe and _too_slow(pattern)):
        return None
    return pattern


class CompiledReplacements:
    """Słownik zamian przygotowany raz na dokument: lista kroków (wzorzec, funkcja zamiany)."""

    def __init__(self, entries):
        self.steps = []
        words = {}

        def flush_words():
            if words:
                alternatives = sorted(words, key=lambda value: (-len(value), value))
                # Granice jak \b, ale działają też dla wpisów z kropką, łącznikiem lub apostrofem
                # na brzegu (wg., itd., 'ten); spacja w źródle pasuje do dowolnego odstępu.
                pattern = re.compile(r'(?<!\w)(?:' + '|'.join(
                    WORD_GAP.join(map(re.escape, value.split(' '))) for value in alternatives
                ) + r')(?!\w)', re.IGNORECASE)
                mapping = dict(words)

                def replace(match):
                    target, keep = mapping[' '.join(match[0].lower().split())]
                    return _apply_replacement_case(match[0], target, keep)
                self.steps.append((pattern, replace))
                words.clear()

        for entry in list(entries or ())[:MAX_REPLACEMENTS]:
            if not isinstance(entry, dict):
                continue
            kind, source, target = entry.get('kind'), entry.get('source'), entry.get('target')
            if (kind not in ('word', 'pattern') or not isinstance(source, str) or not isinstance(target, str)
                    or not source or len(source) > MAX_PATTERN_LENGTH or len(target) > MAX_TARGET_LENGTH):
                continue
            if kind == 'word':
                key = ' '.join(source.lower().split())
                if key and key not in words:
                    words[key] = (target, _keeps_target_case(target))
                continue
            flush_words()
            pattern = _compile_pattern(source, target, probe=entry.get('checked') is not True)
            if pattern is None:
                continue

            def replace(match, template=target, keep=_keeps_target_case(target, pattern=True)):
                # Końcówki są normalizowane do małych liter przed odtworzeniem
                # zapisu całego wyrazu, np. MENTLIKIEM → MĘTLIKIEM.
                return _apply_replacement_case(match[0], match.expand(template), keep)
            self.steps.append((pattern, replace))
        flush_words()

    def apply(self, text, record=None):
        """Zamiany po kolei; record (słownik) dostaje liczby par (znaleziony zapis, wynik)."""
        for pattern, replace in self.steps:
            if record is None:
                text = pattern.sub(replace, text)
                continue

            def counted(match, replace=replace):
                result = replace(match)
                if result != match[0]:
                    record[match[0], result] = record.get((match[0], result), 0) + 1
                return result
            text = pattern.sub(counted, text)
        return text


def compile_replacements(entries=None):
    """None → wbudowana lista DEFAULT_REPLACEMENTS; błędne wpisy są pomijane."""
    if isinstance(entries, CompiledReplacements):
        return entries
    if entries is None:
        global _DEFAULT_COMPILED
        if _DEFAULT_COMPILED is None:
            _DEFAULT_COMPILED = CompiledReplacements(DEFAULT_REPLACEMENTS)
        return _DEFAULT_COMPILED
    return CompiledReplacements(entries)


_DEFAULT_COMPILED = None


def _correct_user_words(text, replacements=None, record=None):
    return compile_replacements(replacements).apply(text, record)


ADDRESS_RE = re.compile(TECHNICAL_RE.pattern.split(r'|\b(?:\d{1,3}\.)')[0], re.IGNORECASE)
QUOTE_OPENERS = ' \t\n(["[{„“«‹–—-'
QUOTE_CLOSERS = ' \t\n)]}.,;:!?…”»›–—-'


def _arithmetic(text, match):
    """„2 - 1 = 1” to odejmowanie, nie zakres: obok stoi znak równości lub plus."""
    window = text[max(0, match.start() - 12):match.end() + 12]
    return '=' in window or '+' in window


def _pair_straight_quotes(text, quote='"'):
    """Cudzysłowy proste parowane stosem: poziom 1 → „…”, cytat w cytacie → «…».

    Otwierający stoi po spacji/nawiasie/na początku i przed znakiem innym niż spacja; zamykający
    stoi po znaku innym niż spacja i przed spacją, interpunkcją lub końcem. Cal (5" x 7") nie ma
    pary. Poziom liczy też otwarte wcześniej „…”. Para musi zawierać literę i mieścić się w linii.
    """
    if quote not in text:
        return text
    stack, pairs = [], []
    for position, char in enumerate(text):
        if char != quote:
            continue
        previous = text[position - 1] if position else ' '
        following = text[position + 1] if position + 1 < len(text) else ' '
        opening = previous in QUOTE_OPENERS and following not in ' \t\n'
        closing = previous not in ' \t\n' and following in QUOTE_CLOSERS
        if opening and (not closing or not stack):
            stack.append(position)
        elif closing and stack:
            start = stack.pop()
            inner = text[start + 1:position]
            if '\n' not in inner and re.search('[' + LETTERS + ']', inner):
                pairs.append((start, position))
    if not pairs:
        return text
    result = list(text)
    for start, end in pairs:
        depth = text.count('„', 0, start) - text.count('”', 0, start)
        depth += sum(1 for a, b in pairs if a < start < b)
        result[start], result[end] = ('«', '»') if depth > 0 else ('„', '”')
    return ''.join(result)


def _nest_polish_quotes(text):
    """„…„…”…” – wewnętrzna para na «…»."""
    if text.count('„') < 2:
        return text
    stack, result = [], list(text)
    for position, char in enumerate(text):
        if char == '„':
            stack.append(position)
        elif char == '”' and stack:
            start = stack.pop()
            if stack and '\n' not in text[start:position]:
                result[start], result[position] = '«', '»'
    return ''.join(result)


def _guillemets_to_polish(text):
    """«tak» jako cudzysłów główny (poza otwartym „…”) → „tak”."""
    def replace(match):
        depth = text.count('„', 0, match.start()) - text.count('”', 0, match.start())
        return match[0] if depth > 0 else f'„{match[1]}”'
    return re.sub(r'«([^«»\n]*[' + LETTERS + r'][^«»\n]*)»', replace, text)


def correct_editorial_text(text, enabled=None, trim_start=True, trim_end=True, *,
                           replacements=None, counts=None):
    """Korekta fragmentu tekstu; puste enabled wyłącza wszystkie reguły.

    Regexy nie ustalają znaczenia zdania. Niejednoznaczne konstrukcje
    pozostają bez zmian, np. 1.234 bez jednostki, samotne cudzysłowy,
    pojedynczy łącznik wewnątrz wyrazu i godzina 9.30 bez kontekstu.
    `replacements`: słownik zamian (lista wpisów lub CompiledReplacements),
    None = DEFAULT_REPLACEMENTS. `counts`: słownik, do którego dopisujemy
    liczbę zmian każdej reguły (liczenie nie wpływa na wynik).
    Reguła note_references działa tylko w DOCX (_correct_noted_group) – tu nic nie zmienia.
    """
    enabled = DEFAULT_EDITORIAL_RULES if enabled is None else frozenset(enabled)
    unknown = enabled - ALL_EDITORIAL_RULES
    if unknown:
        raise ValueError(f'Nieznane reguły korekty: {sorted(unknown)}')
    if not enabled or not text:
        return text
    # Spacje nierozdzielające zawsze stają się zwykłymi (nie liczymy tego jako zmiany reguły;
    # w DOCX robi to wcześniej osobny, cichy krok – patrz _replace_nonbreaking_spaces).
    text = text.translate(NONBREAKING_TABLE)
    if text.strip() and set(text) <= SEPARATOR_CHARS:
        return text  # akapit-separator (* * *, - - -) – reguły nie mają go poprawiać
    if 'invisible_chars' in enabled:
        before = text
        text = INVISIBLE_RE.sub('', text)
        if counts is not None and before != text:
            counts['invisible_chars'] = counts.get('invisible_chars', 0) + len(before) - len(text)

    protected = {}
    occupied = set(text)

    def mask_value(value):
        # Znaki prywatne nie pasują do \w, cyfr ani liter w naszych regexach.
        index = 0xF0000 + len(protected)
        while chr(index) in occupied or chr(index) in protected:
            index += 1
        marker = chr(index)
        protected[marker] = value
        return marker

    def protect_technical(match):
        value = match[0]
        core = value.rstrip('.,;:!?)]}')
        return mask_value(core) + value[len(core):] if core else value

    # Daty przed ochroną wersji (np. 01.02.2025 wygląda też jak numer wersji).
    def date_protect(match):
        value = _normalize_dates_times(match[0]) if 'dates_times' in enabled else match[0]
        if counts is not None and value != match[0]:
            counts['dates_times'] = counts.get('dates_times', 0) + 1
        return mask_value(value)

    # Adresy muszą być chronione zanim rozpoznamy datę wewnątrz URL.
    text = ADDRESS_RE.sub(protect_technical, text)
    # Reguły liczbowe pomijamy w akapitach bez cyfr; żadna reguła cyfr nie dodaje.
    digits = DIGIT_RE.search(text) is not None
    if digits:
        text = DATE_RE.sub(date_protect, text)
        text = ROMAN_DATE_RE.sub(lambda m: mask_value(m[0]), text)
    text = TECHNICAL_RE.sub(protect_technical, text)
    if digits:
        def time_range(match):
            value = _time_range(match, enabled, counts)
            return match[0] if value is None else mask_value(value)
        text = TIME_RANGE_RE.sub(time_range, text)
    if digits and 'dates_times' in enabled:
        before = text
        text = _normalize_dates_times(text)
        if counts is not None:
            _count(counts, 'dates_times', before, text)
    if digits:
        text = TIME_RE.sub(lambda m: mask_value(m[0]), text)
    # Nazwy plików (raport.docx, Node.js) chronimy zawsze, nie tylko przy wielkich literach po kropce.
    text = FILE_NAME_RE.sub(lambda m: mask_value(m[0]), text)
    if 'user_word_corrections' in enabled:
        before = text
        record = {} if counts is not None else None
        text = _correct_user_words(text, replacements, record)
        if counts is not None:
            _count(counts, 'user_word_corrections', before, text)
            # Zamiany słownikowe do podsumowania: (znaleziony zapis, wynik) → liczba.
            pairs = counts.setdefault('replacements', {})
            for (source, target), number in record.items():
                key = (_unmask(source, protected), _unmask(target, protected))
                pairs[key] = pairs.get(key, 0) + number

    if 'tabs' in enabled:
        before = text
        text = re.sub(r'[ \t\u00a0\u202f]*\t[ \t\u00a0\u202f]*', ' ', text)
        if counts is not None:
            _count(counts, 'tabs', before, text)
    if 'spaces' in enabled:
        before = text
        text = re.sub(r' {2,}', ' ', text)
        if counts is not None:
            _count(counts, 'spaces', before, text)
    if 'abbreviations' in enabled:
        before = text
        for pattern, replacement in ABBREVIATION_MERGES:
            text = pattern.sub(replacement, text)
        if counts is not None:
            _count(counts, 'abbreviations', before, text)
    if 'quotes' in enabled:
        before = text
        # Cudzysłów maszynowy z dwóch przecinków i apostrofów: ,,tak'' / ,,tak" → „tak”.
        text = re.sub(r",,([^,„”\"\n]*?[" + LETTERS + r"][^„”\"\n]*?)(?:''|\"|”)", r'„\1”', text)
        # LaTeX i ASCII: ``tak'' oraz ''tak''.
        text = re.sub(r"(?<!\w)(?:``|'')([^`'\n]*[" + LETTERS + r"][^`'\n]*)''(?!\w)", r'„\1”', text)
        text = re.sub(r'“([^“”\n]+)”', r'„\1”', text)
        text = re.sub(r'„([^„“”\n]+)“', r'„\1”', text)
        # Dwa „angielskie zamykające” z autokorekty Worda (”tak”), cudzysłów prosty zamykający po „
        # (otwarty „ bez pary) i odwrotny układ "tak„.
        text = re.sub(r'(?<![\w”])”(?=[' + LETTERS + r'„“"])([^„”“\n]*[' + LETTERS + r'][^„”“\n]*)”(?!\w)',
                      r'„\1”', text)
        text = re.sub(r'„([^„”“"\n]*[' + LETTERS + r'][^„”“"\n]*)"(?!\w)', r'„\1”', text)
        text = re.sub(r'(?<!\w)"([^„”“"\n]*[' + LETTERS + r'][^„”“"\n]*)„(?!\w)', r'„\1”', text)
        # Niemieckie »tak« i francuskie «tak» jako cudzysłów główny (poza otwartym „…”).
        text = re.sub(r'»([^«»\n]*[' + LETTERS + r'][^«»\n]*)«', r'„\1”', text)
        text = _guillemets_to_polish(text)
        text = re.sub(r'‘([^‘’\n]*[' + LETTERS + r'][^‘’\n]*)’', r'„\1”', text)
        text = _pair_straight_quotes(text)
        text = _pair_straight_quotes(text, "'")  # 'tak' jako cudzysłów; apostrof w wyrazie nie ma pary
        text = _nest_polish_quotes(text)
        if counts is not None:
            _count(counts, 'quotes', before, text)
    if 'apostrophe' in enabled:
        before = text
        text = re.sub(r'(?<=[' + LETTERS + r"])'(?=[" + LETTERS + r'])', '’', text)
        if counts is not None:
            _count(counts, 'apostrophe', before, text)
    if 'inside_quotes' in enabled:
        before = text
        text = re.sub(r'([„“])' + H + r'+', r'\1', text)
        text = re.sub(H + r'+(?=”(?![' + LETTERS + r']))', '', text)
        # Cudzysłów prosty tuż po cyfrze to cal (5" x 7"), nie początek cytatu.
        text = re.sub(r'(?<!\w)"([^"\n]+)"', lambda m: '"' + m[1].strip(' \u00a0\u202f') + '"', text)
        if counts is not None:
            _count(counts, 'inside_quotes', before, text)
    if 'ellipsis' in enabled:
        before = text
        text = re.sub(r'\.{3,}', '…', text)
        if counts is not None:
            _count(counts, 'ellipsis', before, text)
    if 'quote_punct' in enabled:
        before = text
        # Polska interpunkcja: przecinek, średnik, dwukropek i kropka stoją zawsze za cudzysłowem
        # zamykającym (pytajnik, wykrzyknik i wielokropek należące do cytatu zostają w środku).
        text = re.sub(r'([,;:])([”»])', r'\2\1', text)
        # Kropka przed cudzysłowem znika, gdy zdanie trwa dalej („tak.” – powiedział; „tak.” i wyszedł),
        # a kończy zdanie za cudzysłowem w pozostałych wypadkach („tak.”. → „tak”.).
        text = re.sub(r'\.([”»])(?=' + H + r'*[–—]|' + H + r'+[' + LOWER + r'])', r'\1', text)
        text = re.sub(r'\.([”»])\.?', r'\1.', text)
        if counts is not None:
            _count(counts, 'quote_punct', before, text)
    if 'inside_brackets' in enabled:
        before = text
        text = re.sub(r'([(\[{])' + H + r'+', r'\1', text)
        text = re.sub(H + r'+([)\]}])', r'\1', text)
        if counts is not None:
            _count(counts, 'inside_brackets', before, text)
    if 'outside_brackets' in enabled:
        before = text
        # Formy w rodzaju „zrobił(a)”, „Drogi(a)”, „pan(i)” zostają złączone.
        text = re.sub(r'(?<=[' + LETTERS + r'”])(?=[(\[{])(?![(\[][' + LOWER + r']{1,4}[)\]])', ' ', text)
        text = re.sub(r'(?<=[)\]}])(?=[' + LETTERS + r'])', ' ', text)
        if counts is not None:
            _count(counts, 'outside_brackets', before, text)
    if 'before_punct' in enabled:
        before = text
        # Wynik lub proporcja „3 : 1” zachowuje symetryczne odstępy.
        def before_punct(match, source=text):
            ratio = (source[match.end()] == ':' and source[max(0, match.start() - 1):match.start()].isdigit()
                     and re.match(':' + H + r'*\d', source[match.end():]))
            if source[match.end()] == '…':
                # Wielokropek otwierający (po myślniku, cudzysłowie, nawiasie, na początku) zostaje po spacji.
                left = source[max(0, match.start() - CONTEXT):match.start()].rstrip()
                if not left or left[-1] in '–—-„“«("[':
                    return match[0]
            return match[0] if ratio else ''
        text = re.sub(H + r'+(?=[,.;:?!…])', before_punct, text)
        if counts is not None:
            _count(counts, 'before_punct', before, text)
    if 'duplicate_punct' in enabled:
        before = text
        text = re.sub(r'([,;:])\1+', r'\1', text)
        # Maszynowy zapis ?.. i !.. to ?… i !…; pozostałe podwójne kropki to jedna kropka.
        text = re.sub(r'(?<=[?!])\.{2}(?!\.)', '…', text)
        text = re.sub(r'(?<!\.)\.{2}(?!\.)', '.', text)
        if counts is not None:
            _count(counts, 'duplicate_punct', before, text)
    if 'hyphen_dash' in enabled:
        before = text
        # Podwójny łącznik z maszynopisu to myślnik: „tak--powiedział”, „tak -- rzekł”, „-- Wiem”.
        # „---” (separatory, znaczniki) i „--opcja” (zapis techniczny) zostają bez zmian.
        text = re.sub(r'(?<=\w)--(?=\w)|(?<![-\S])--(?=' + H + r')', '–', text)
        text = re.sub(r'(?<= )-(?= )', lambda m: m[0] if _arithmetic(text, m) and text[m.start() - 2:m.start() - 1].isdigit()
                      and text[m.end() + 1:m.end() + 2].isdigit() else '–', text)
        # Łącznik tuż po kropce, znaku zapytania lub wykrzyknika przed spacją to myślnik: „Tak.- Nie”.
        text = re.sub(r'(?<=[.?!…])-(?=' + H + r')', ' –', text)
        # Kwestia dialogowa na początku linii; przed liczbą tylko z odstępem (minus pisze się „-5”).
        text = re.sub(
            r'(^|\n)' + H + r'*-(?:' + H + r'*(?=[' + LETTERS + r'„“"])|' + H + r'+(?=\d))',
            lambda match: (match[1] + '– ')
            if trim_start or match[1] == '\n' else match[0], text,
        )
        if counts is not None:
            _count(counts, 'hyphen_dash', before, text)
    if 'dash_style' in enabled:
        before = text
        # Minus (U+2212) i pauza cyfrowa (U+2012) tuż przed cyfrą to znak liczby („−5 °C”, „(−3)”) albo
        # ścisły zapis (555‒1234): półpauza bez odstępów, zamaskowana, żeby dash_spaces nie rozsunął jej
        # jak myślnika. „1990 −1995” (cyfra i spacja przed) zostaje dla reguł zakresów.
        text = re.sub(r'(?<!\d )[\u2012\u2212](?=\d)', lambda m: mask_value('–'), text)
        # Decyzja właściciela: pauza, kreska pozioma, pauza cyfrowa i minus → półpauza w każdym kontekście;
        # łącznik Unicode i nierozdzielający → zwykły dywiz.
        text = text.translate(DASH_TABLE)
        # Powtórzone myślniki (– –, ––) → jeden.
        text = re.sub(r'–(?:' + H + r'*–)+', '–', text)
        if counts is not None:
            _count(counts, 'dash_style', before, text)
    if digits and 'range_dash' in enabled:
        before = text
        text = re.sub(r'(?<![\w.,\-\u2012\u2212])(\d+)' + H + r'*[-\u2012\u2212]' + H + r'*(\d+)(?![\w\-\u2012\u2212]|[.,]\d)',
                      lambda m: m[0] if _arithmetic(text, m) else re.sub('[-\u2012\u2212]', '–', m[0]), text)
        if counts is not None:
            _count(counts, 'range_dash', before, text)
    if digits and 'range_spaces' in enabled:
        before = text
        text = re.sub(r'(?<![\w.,-])(\d+)' + H + r'*([–-])' + H + r'*(\d+)(?![\w-]|[.,]\d)',
                      lambda m: m[0] if _arithmetic(text, m) else m[1] + m[2] + m[3], text)
        if counts is not None:
            _count(counts, 'range_spaces', before, text)
    if 'dash_spaces' in enabled:
        before = text
        def dash_space(match):
            # Wystarczy najbliższy znak niebędący spacją po obu stronach.
            left = text[max(0, match.start() - CONTEXT):match.start()].rstrip(' \u00a0\u202f')
            if not left and match.start() > CONTEXT:
                left = text[:match.start()].rstrip(' \u00a0\u202f')
            right = text[match.end():match.end() + CONTEXT].lstrip(' \u00a0\u202f')
            if not right and len(text) - match.end() > CONTEXT:
                right = text[match.end():].lstrip(' \u00a0\u202f')
            if (not left and not trim_start) or (not right and not trim_end):
                return match[0]
            if left[-1:].isdigit() and right[:1].isdigit():
                return match[0]
            if right[:1].isdigit() and (not left or left[-1:] in '(=:'):
                return match[0]  # potencjalny znak minus
            leading = ' ' if left and not left.endswith('\n') and left[-1] not in '„“«([' else ''
            trailing = ' ' if right and right[0] not in '”»)]–—' else ''
            return leading + match[1] + trailing
        text = re.sub(H + r'*([–—])' + H + r'*', dash_space, text)
        if counts is not None:
            _count(counts, 'dash_spaces', before, text)
    if digits and 'decimal' in enabled:
        before = text
        text = re.sub(r'(?<![\w.,])(\d+)\.(\d{1,2})(?=' + H + r'*(?:' + UNITS + r'|%|°' + H + r'*[CF])(?!\w))',
                      r'\1,\2', text)
        if counts is not None:
            _count(counts, 'decimal', before, text)
    if digits and 'unit_space' in enabled:
        before = text
        text = re.sub('(' + NUMBER + r')(?=' + ATTACHED_UNITS + r'(?!\w))', r'\1 ', text)
        text = re.sub(r'(?<=\d' + H + r')([kmcd]?m)([23])(?!\w)',
                      lambda m: m[1] + {'2': '²', '3': '³'}[m[2]], text)
        if counts is not None:
            _count(counts, 'unit_space', before, text)
    if digits and 'digit_groups' in enabled:
        before = text
        def group_digits(match):
            digits = re.sub(H, '', match[0])
            if digits.startswith('0'):
                return match[0]
            if IDENTIFIER_LABEL_RE.search(text, max(0, match.start() - CONTEXT), match.start()):
                return match[0]
            quantity = QUANTITY_AFTER_RE.match(text, match.end())
            if len(digits) >= 9 and not quantity:
                return match[0]  # możliwy numer telefonu lub identyfikator
            return _group_number(digits)
        text = re.sub(r'(?<![\w.,])(\d{1,3}(?:' + H + r'\d{3})+|\d{5,})(?!\w|\.\d)', group_digits, text)
        if counts is not None:
            _count(counts, 'digit_groups', before, text)
    if digits and 'unit_nbsp' in enabled:
        before = text
        text = re.sub(r'(?<=\d)' + H + r'+(?=' + UNITS + r'(?!\w))', ' ', text)
        text = re.sub(r'(?<=\d)' + H + r'+(?=[%‰])', '', text)
        if counts is not None:
            _count(counts, 'unit_nbsp', before, text)
    if digits and 'reference_space' in enabled:
        before = text
        text = re.sub(r'\b(\d{4})(?=r\.)', r'\1 ', text)
        text = re.sub(r'\bnr\.' + H + r'*(?=\d)', 'nr ', text)
        text = re.sub(r'(?<![^\W\d])(' + NUMBER_ABBREVIATIONS + r'|nr)(?=\d)', r'\1 ', text)
        if counts is not None:
            _count(counts, 'reference_space', before, text)
    if digits and 'reference_nbsp' in enabled:
        before = text
        text = re.sub(r'\b(\d{4})' + H + r'+(?=r\.)', r'\1 ', text)
        text = re.sub(r'(?<![^\W\d])(' + NUMBER_ABBREVIATIONS + r'|nr)' + H + r'+(?=\d)', r'\1 ', text)
        if counts is not None:
            _count(counts, 'reference_nbsp', before, text)
    if 'initials' in enabled:
        # Skrótowce z kropkami (S.A., P.S.) to nie inicjały.
        text = DOTTED_ACRONYM_RE.sub(lambda m: mask_value(m[0]), text)
        before = text
        text = re.sub(r'\b([' + UPPER + r'])' + H + r'+\.', r'\1.', text)
        # Kilka inicjałów rozdzielamy tylko przed nazwiskiem (J.R.R. Tolkien); U.S.A. zostaje.
        text = re.sub(r'\b(?:[' + UPPER + r']\.' + H + r'*){2,}(?=' + H + r'*[' + UPPER + r'][' + LOWER + r']{2,}\b)',
                      lambda m: ' '.join(re.findall(r'[' + UPPER + r']\.', m[0])) + (' ' if m[0][-1] != '.' else ''),
                      text)
        text = re.sub(r'\b([' + UPPER + r']\.)' + H + r'*(?=[' + UPPER + r'][' + LOWER + r']{2,}\b)',
                      r'\1 ', text)
        if counts is not None:
            _count(counts, 'initials', before, text)
    if digits and 'temperature' in enabled:
        before = text
        text = re.sub('(' + NUMBER + r')' + H + r'*°' + H + r'*([CF])(?!\w)', r'\1 °\2', text)
        # Kelwin tylko po odstępie: „4K” to rozdzielczość, nie temperatura.
        text = re.sub('(' + NUMBER + r')' + H + r'+K(?!\w)', r'\1 K', text)
        if counts is not None:
            _count(counts, 'temperature', before, text)
    if 'after_punct' in enabled:
        gaps = 0

        # Zachowaj wewnętrzne kropki skrótów. Liczby nie są celem tej reguły.
        def protect_abbreviation(match):
            nonlocal gaps
            tail = text[match.end():match.end() + 1]
            # Granica po skrócie wymaga spacji; jego wewnętrzne kropki są chronione.
            gap = ' ' if tail and re.fullmatch('[' + LETTERS + '„“\"]', tail) else ''
            gaps += bool(gap)
            return mask_value(match[0]) + gap
        text = ABBREVIATION_RE.sub(protect_abbreviation, text)
        before = text
        # Cudzysłów prosty tylko otwierający (przed literą); zamykający na końcu akapitu zostaje przy kropce.
        text = re.sub(r'(?<=[,;:!?])(?=[' + LETTERS + r'„“]|"[' + LETTERS + r'])', ' ', text)
        text = re.sub(r'(?<=\.)(?=[' + LETTERS + r'„“]|"[' + LETTERS + r'])', ' ', text)
        # Wielokropek między wyrazami (godziny…a potem); „…i” na początku fragmentu zostaje.
        text = re.sub(r'(?<=[' + LETTERS + r']…)(?=[' + LETTERS + r'])', ' ', text)
        # Przecinek przed liczbą wymaga spacji, chyba że tworzy liczbę dziesiętną.
        # Zamaskowany numer (wersja, data) liczy się jak cyfra: 1.234.567,89 zostaje.
        text = re.sub(r'(?<![\d' + MARKER_CLASS + r']),(?=\d)', ', ', text)
        text = re.sub(r'(?<=[;!?])(?=\d)', ' ', text)
        text = re.sub(r'(?<![\d' + MARKER_CLASS + r']):(?=\d)', ': ', text)
        if counts is not None:
            _count(counts, 'after_punct', before, text, gaps)
    text = _unmask(text, protected)
    if 'pronouns_lower' in enabled:
        before = text
        # Po przywróceniu masek dostępne są kropki kończące skróty i zdania.
        text = _lower_personal_pronouns(text)
        if counts is not None:
            _count(counts, 'pronouns_lower', before, text)
    # Po zaimkach: „Był Nie najlepszy” → „nie najlepszy” → „nienajlepszy”; przed wielką literą po kropce.
    if enabled & ORTHOGRAPHY_2026_KEYS:
        text = _orthography_2026(text, enabled, counts)
    if 'sentence_case' in enabled:
        before = text
        text = _capitalize_sentences(text)
        if counts is not None:
            _count(counts, 'sentence_case', before, text)
    if 'trim' in enabled:
        before = text
        if trim_start:
            text = re.sub(r'^' + H + r'+', '', text)
        if trim_end:
            text = re.sub(H + r'+$', '', text)
        if counts is not None:
            _count(counts, 'trim', before, text)
    return text


def _unmask(text, protected):
    # Jedno przejście na raz; powtarzamy, gdyby maska zawierała wcześniejszą maskę.
    while protected:
        restored = MARKER_RE.sub(lambda match: protected.get(match[0], match[0]), text)
        if restored == text:
            break
        text = restored
    return text


IDENTIFIER_LABEL_RE = re.compile(r'\b(?:nr|numer|kod|id|tel\.?)' + H + r'*[:#]?' + H + r'*$', re.I)
QUANTITY_AFTER_RE = re.compile(H + r'*(?:' + QUANTITIES + r')(?!\w)')


def _group_number(digits):
    groups = []
    while digits:
        groups.append(digits[-3:])
        digits = digits[:-3]
    return ' '.join(reversed(groups))


def _edit_opcodes(original, corrected):
    """Najmniejsza lista zmian znaków (jak difflib.get_opcodes).

    rapidfuzz liczy ją w czasie zbliżonym do liniowego także dla akapitów
    po 20 000 znaków; difflib (zapasowo) rośnie z kwadratem długości.
    """
    try:
        from rapidfuzz.distance import Levenshtein
    except ImportError:
        from difflib import SequenceMatcher
        return SequenceMatcher(None, original, corrected, autojunk=False).get_opcodes()
    return [(op.tag, op.src_start, op.src_end, op.dest_start, op.dest_end)
            for op in Levenshtein.opcodes(original, corrected)]


def _is_word_char(char):
    return char.isalnum() or char == '_'


def _change_regions(original, corrected):
    """Zmiany jako przedziały (a, b, c, d): original[a:b] → corrected[c:d].

    Sąsiednie operacje łączymy, a zmianę dotykającą liter lub cyfr rozszerzamy do
    granic wyrazu (mentlik → mętlik to jedna zmiana całego wyrazu, a nie dwie
    zmiany pojedynczych liter). Zmiany samych odstępów i interpunkcji zostają
    najmniejsze, np. „kota...” → „kota…” zmienia tylko kropki.
    """
    regions = []
    for tag, a, b, c, d in _edit_opcodes(original, corrected):
        if tag == 'equal':
            continue
        if regions and regions[-1][1] == a and regions[-1][3] == c:
            regions[-1][1], regions[-1][3] = b, d
        else:
            regions.append([a, b, c, d])
    result = []
    for index, (a, b, c, d) in enumerate(regions):
        if any(map(_is_word_char, original[a:b])) or any(map(_is_word_char, corrected[c:d])):
            lower = result[-1][1] if result else 0
            while a > lower and _is_word_char(original[a - 1]):
                a, c = a - 1, c - 1
            upper = regions[index + 1][0] if index + 1 < len(regions) else len(original)
            while b < upper and _is_word_char(original[b]):
                b, d = b + 1, d + 1
        if result and result[-1][1] >= a:
            result[-1][1], result[-1][3] = b, d
        else:
            result.append([a, b, c, d])
    return result


def _count(counts, key, before, after, extra=0):
    number = extra + (len(_change_regions(before, after)) if before != after else 0)
    if number:
        counts[key] = counts.get(key, 0) + number


TRACK_AUTHOR = 'Odkurzacz'
TRACK_ID_RE = re.compile(rb'\bw:id="(\d+)"')


class _Tracker:
    """Wspólne dane śledzonych zmian: autor, data UTC i kolejne wolne w:id."""

    def __init__(self, document):
        from datetime import datetime, timezone
        self.date = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
        highest = 0
        # w:id musi być unikalny w całej paczce, także względem istniejących
        # zmian, komentarzy i zakładek (w nagłówkach, przypisach itd.).
        for part in document.part.package.iter_parts():
            element = getattr(part, '_element', None)
            if element is not None:
                values = (node.get(qn('w:id')) or '' for node in element.iter())
            elif str(part.partname).endswith('.xml'):
                values = (match.decode() for match in TRACK_ID_RE.findall(part.blob or b''))
            else:
                continue
            highest = max([highest, *(int(value) for value in values if value.isdigit())])
        self.next_id = highest + 1

    def take_id(self):
        self.next_id += 1
        return self.next_id - 1

    def mark(self, tag):
        element = OxmlElement(tag)
        element.set(qn('w:id'), str(self.take_id()))
        element.set(qn('w:author'), TRACK_AUTHOR)
        element.set(qn('w:date'), self.date)
        return element


def _append_text(run, text, deleted=False):
    for index, part in enumerate(text.split('\t')):
        if index:
            run.append(OxmlElement('w:tab'))
        if part:
            node = OxmlElement('w:delText' if deleted else 'w:t')
            node.text = part
            node.set(qn('xml:space'), 'preserve')
            run.append(node)


def _split_tracked_run(run, changed, tracker):
    """Zastąp przebieg ciągiem: zwykłe przebiegi, w:del i w:ins – z tym samym rPr."""
    from copy import deepcopy

    properties = run.find(qn('w:rPr'))
    copies = 0

    def new_run():
        nonlocal copies
        element = OxmlElement('w:r')
        for key, value in run.attrib.items():
            element.set(key, value)
        if properties is not None:
            copied = deepcopy(properties)
            # Zmiana formatowania (w:rPrChange) ma w:id – każda kopia poza pierwszą dostaje nowy.
            if copies:
                for node in copied.iter():
                    if node.get(qn('w:id')) is not None:
                        node.set(qn('w:id'), str(tracker.take_id()))
            copies += 1
            element.append(copied)
        return element

    output, plain = [], None
    for child in list(run):
        if child is properties:
            continue
        if child not in changed:
            if plain is None:
                plain = new_run()
                output.append(plain)
            plain.append(child)
            continue
        for kind, text in changed[child]:
            if kind == 'keep':
                if plain is None:
                    plain = new_run()
                    output.append(plain)
                _append_text(plain, text)
                continue
            tag = 'w:del' if kind == 'del' else 'w:ins'
            if plain is None and output and output[-1].tag == qn(tag):
                wrapper = output[-1]  # ten sam fragment rozbity na węzły w:t i w:tab
            else:
                wrapper = tracker.mark(tag)
                output.append(wrapper)
            piece = new_run()
            _append_text(piece, text, deleted=kind == 'del')
            wrapper.append(piece)
            plain = None
    for element in output:
        run.addprevious(element)
    run.getparent().remove(run)
    for element in output:
        if element.tag == qn('w:ins'):
            _lift_insertion(element, tracker)


INSERTED_CONTAINERS = frozenset(qn('w:' + name) for name in ('ins', 'moveTo'))


def _lift_insertion(insertion, tracker):
    """Nasza wstawka wewnątrz cudzej (w:ins, w:moveTo) staje obok niej, jak u Worda.

    Cudza wstawka dzieli się na część przed naszą i za nią (ta dostaje kopię opakowania z nowym
    w:id, tym samym autorem i datą). Tekst nie zmienia autora, a nasze usunięcia fragmentów cudzej
    wstawki zostają w niej zagnieżdżone (w:ins > w:del) – tak zapisuje je Word.
    """
    parent = insertion.getparent()
    while parent is not None and parent.tag in INSERTED_CONTAINERS:
        before = insertion.getprevious() is not None
        after = list(insertion.itersiblings())
        if before and after:
            rest = parent.makeelement(parent.tag, dict(parent.attrib))
            rest.set(qn('w:id'), str(tracker.take_id()))
            for node in after:
                rest.append(node)
            parent.addnext(rest)
        if before:
            parent.addnext(insertion)
        else:
            parent.addprevious(insertion)
            if not after:
                parent.getparent().remove(parent)
        parent = insertion.getparent()


def _rewrite_tracked(slots, corrected, tracker):
    """Zapisz różnicę jako śledzone zmiany Worda; tekst przebiegów bez zmian zostaje."""
    from bisect import bisect_right

    original = ''.join(value for _, value in slots)
    if original == corrected:
        return
    ends, offset = [], 0
    for _, value in slots:
        offset += len(value)
        ends.append(offset)
    pieces = [[] for _ in slots]

    def spread(kind, start, stop):
        index = bisect_right(ends, start)
        while start < stop:
            end = min(stop, ends[index])
            pieces[index].append((kind, original[start:end]))
            start, index = end, index + 1

    position = 0
    for a, b, c, d in _change_regions(original, corrected):
        spread('keep', position, a)
        # Wstawka dostaje formatowanie przebiegu, w którym zaczyna się zmiana (jak bez śledzenia).
        owner = min(bisect_right(ends, a), len(slots) - 1)
        spread('del', a, b)
        if c < d:
            pieces[owner].append(('ins', corrected[c:d]))
        position = b
    spread('keep', position, len(original))
    changed, runs = {}, []
    for (node, _), node_pieces in zip(slots, pieces):
        if any(kind != 'keep' for kind, _ in node_pieces):
            changed[node] = node_pieces
            run = node.getparent()
            if not any(run is seen for seen in runs):
                runs.append(run)
    for run in runs:
        _split_tracked_run(run, changed, tracker)


def _rewrite_text_nodes(slots, corrected):
    """Zmieniaj wyłącznie tekst i tabulatory; nie przebudowuj całego akapitu."""
    from bisect import bisect_right

    original = ''.join(value for _, value in slots)
    if original == corrected:
        return
    ends, offset = [], 0
    for _, value in slots:
        offset += len(value)
        ends.append(offset)
    chunks = [[] for _ in slots]
    for tag, a, b, c, d in _edit_opcodes(original, corrected):
        if tag == 'equal':
            position = a
            while position < b:
                index = bisect_right(ends, position)
                stop = min(b, ends[index])
                chunks[index].append(original[position:stop])
                position = stop
        elif tag in ('insert', 'replace'):
            index = min(bisect_right(ends, a), len(slots) - 1)
            chunks[index].append(corrected[c:d])
    for (node, before), pieces in zip(slots, chunks):
        after = ''.join(pieces)
        if before == after:
            continue
        if node.tag == qn('w:t'):
            node.text = after
            node.set(qn('xml:space'), 'preserve')
        elif node.tag == qn('w:tab'):
            # Wstawione znaki mogą towarzyszyć zachowanemu tabulatorowi.
            for i, piece in enumerate(after.split('\t')):
                if i:
                    node.addprevious(OxmlElement('w:tab'))
                if piece:
                    replacement = OxmlElement('w:t')
                    replacement.text = piece
                    replacement.set(qn('xml:space'), 'preserve')
                    node.addprevious(replacement)
            node.getparent().remove(node)


def _blank_run(run):
    for node in run:
        if node.tag == qn('w:rPr'):
            continue
        if node.tag == qn('w:t') and not (node.text or '').strip('   \t'):
            continue
        if node.tag == qn('w:tab'):
            continue
        return False
    return True


MARK_CHANGES = frozenset(qn('w:' + name) for name in ('ins', 'del', 'moveFrom', 'moveTo'))


def _mark_changes(paragraph):
    """Śledzone zmiany znaku akapitu (w:pPr/w:rPr: w:ins, w:del, w:moveFrom, w:moveTo)."""
    mark = paragraph.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    return [] if mark is None else [child for child in mark if child.tag in MARK_CHANGES]


def _mark_deleted(paragraph):
    """Znak akapitu usunięty (także przeniesiony): po przyjęciu zmian akapit łączy się z następnym."""
    return (paragraph is not None and paragraph.tag == qn('w:p')
            and any(change.tag in DELETED_CONTAINERS for change in _mark_changes(paragraph)))


def _previous_paragraph(paragraph):
    """Poprzedni akapit tego samego kontenera; znaczniki bez treści między akapitami nie przeszkadzają."""
    for sibling in paragraph.itersiblings(preceding=True):
        if sibling.tag == qn('w:p'):
            return sibling
        if sibling.tag not in PARAGRAPH_MARKERS:
            return None
    return None


def _safe_empty_paragraph(element, tracked=False):
    """Pusty akapit, który reguła empty_paragraphs może usunąć.

    Zostają akapity z treścią, zakotwiczonym komentarzem, zakładką, układem strony i listą,
    a także z usuniętym znakiem akapitu (po przyjęciu zmian i tak łączy się z następnym). Akapit
    z inną śledzoną zmianą (wstawiony znak akapitu, zmiana formatowania) usuwamy tylko jako śledzone
    usunięcie – bez śledzenia zniknęłaby cudza zmiana.
    """
    if element.tag != qn('w:p'):
        return False
    for child in element:
        if child.tag == qn('w:pPr'):
            if any(child.find(qn('w:' + name)) is not None
                   for name in ('sectPr', 'numPr', 'pageBreakBefore')):
                return False
            changes = _mark_changes(element)
            if any(change.tag in DELETED_CONTAINERS for change in changes):
                return False
            if not tracked and (changes or child.find(qn('w:pPrChange')) is not None
                                or child.find(f"{qn('w:rPr')}/{qn('w:rPrChange')}") is not None):
                return False
        elif child.tag == qn('w:r'):
            if not _blank_run(child):
                return False
        elif tracked and child.tag in (qn('w:del'), qn('w:ins')) and child.get(qn('w:author')) == TRACK_AUTHOR:
            # Nasze śledzone zmiany: liczy się tekst po ich przyjęciu.
            if child.tag == qn('w:ins') and not all(
                    run.tag == qn('w:r') and _blank_run(run) for run in child):
                return False
        else:
            return False
    return True


def _delete_paragraph_tracked(paragraph, tracker):
    """Pusty akapit jako śledzone usunięcie: treść w w:del i usunięty znak akapitu."""
    for child in list(paragraph):
        if child.tag == qn('w:r'):
            wrapper = tracker.mark('w:del')
            child.addprevious(wrapper)
            wrapper.append(child)
            for node in child.iter(qn('w:t')):
                node.tag = qn('w:delText')
        elif child.tag == qn('w:ins'):
            paragraph.remove(child)  # nasza wstawka w usuwanym akapicie znika bez śladu
    properties = paragraph.find(qn('w:pPr'))
    if properties is None:
        properties = OxmlElement('w:pPr')
        paragraph.insert(0, properties)
    mark_properties = properties.find(qn('w:rPr'))
    if mark_properties is None:
        mark_properties = OxmlElement('w:rPr')
        following = [properties.find(qn('w:' + name)) for name in ('sectPr', 'pPrChange')]
        following = [node for node in following if node is not None]
        if following:
            following[0].addprevious(mark_properties)
        else:
            properties.append(mark_properties)
    if mark_properties.find(qn('w:del')) is not None:
        return  # znak akapitu był już usunięty (wcześniejsza śledzona zmiana)
    inserted = mark_properties.find(qn('w:ins'))
    if inserted is not None:
        inserted.addnext(tracker.mark('w:del'))
    else:
        mark_properties.insert(0, tracker.mark('w:del'))


# Znaczniki bez treści (sprawdzanie pisowni, zakładki, zakresy komentarzy i przeniesień, układ strony
# zapamiętany przez Worda) nie przerywają tekstu: korekta widzi zdanie w całości,
# a znaczniki zostają na swoich miejscach między węzłami tekstu.
PARAGRAPH_MARKERS = frozenset(qn('w:' + name) for name in (
    'pPr', 'proofErr', 'bookmarkStart', 'bookmarkEnd', 'commentRangeStart', 'commentRangeEnd',
    'permStart', 'permEnd', 'moveFromRangeStart', 'moveFromRangeEnd', 'moveToRangeStart', 'moveToRangeEnd',
    'customXmlInsRangeStart', 'customXmlInsRangeEnd', 'customXmlDelRangeStart', 'customXmlDelRangeEnd',
    'customXmlMoveFromRangeStart', 'customXmlMoveFromRangeEnd', 'customXmlMoveToRangeStart',
    'customXmlMoveToRangeEnd'))
# Odwołanie do komentarza (dymek na marginesie) też jest tylko znacznikiem: przebieg zostaje na miejscu.
RUN_MARKERS = frozenset(qn('w:' + name) for name in ('rPr', 'lastRenderedPageBreak', 'commentReference'))
# Odsyłacze przypisów w tekście nie przerywają zdania: w grupie są nieruchomymi kotwicami.
NOTE_REFERENCES = frozenset(qn('w:' + name) for name in ('footnoteReference', 'endnoteReference'))
# Kropka skrótu lub inicjału przed odsyłaczem zostaje na miejscu (itd.¹, r.¹, J.¹).
NOTE_ABBREVIATION_RE = re.compile(
    r'(?:' + CAPITALIZATION_ABBREVIATION_RE.pattern + r'|(?<!\w)(?-i:[' + UPPER + r'])\.)$', re.IGNORECASE)


def _note_reference_run(run):
    """Przebieg z samym odsyłaczem (i formatowaniem). Odsyłacz z własnym znakiem (customMarkFollows – znak stoi
    jako tekst tuż za nim, także w osobnym przebiegu) pozostaje barierą, żeby reguły go nie oddzieliły od znaku."""
    return all(node.tag in RUN_MARKERS or (node.tag in NOTE_REFERENCES
                                          and node.get(qn('w:customMarkFollows')) not in ('1', 'true', 'on'))
               for node in run)


@lru_cache(maxsize=8)
def _note_patterns(marker):
    """Wzorce reguły note_references dla tekstu, w którym `marker` zastępuje odsyłacz."""
    mark = re.escape(marker)
    return (
        # Spacje między wyrazem (lub interpunkcją) a odsyłaczem; początek linii, myślnik i otwarcie zostają.
        re.compile(r'(?<=[^\s–—\-(\[{„«‚‹' + mark + r'])[ \t]+(?=' + mark + ')'),
        # Odsyłacz przed cudzysłowem lub nawiasem zamykającym → za nim.
        re.compile('(' + mark + r'+)([”»)\]]|[’"](?!\w))'),
        # Odsyłacz za kropką, przecinkiem, średnikiem lub dwukropkiem → przed nimi (kropki „...” zostają).
        re.compile(r'(?<!\.)([.,;:])(' + mark + '+)(?!\\.)'),
    )


def _place_note_references(text, marker, *, spaces_only=False):
    """Reguła note_references (do punktu stałego) na tekście z odsyłaczami zastąpionymi przez `marker`.

    Odsyłacz stojący przed ? ! … zostaje (autor mógł odnieść przypis do wyrazu); za nimi – także.
    spaces_only: tylko usunięcie odstępu przed odsyłaczem (zawsze, także przy wyłączonej regule).
    """
    spaces, closers, punctuation = _note_patterns(marker)
    if spaces_only:
        return spaces.sub('', text)

    def before_punct(match):
        if match[1] == '.':
            prefix = match.string[max(0, match.start() - CONTEXT):match.end(1)].replace(marker, '')
            if NOTE_ABBREVIATION_RE.search(prefix):
                return match[0]
        return match[2] + match[1]

    for _ in range(len(text) + 1):
        previous = text
        text = punctuation.sub(before_punct, closers.sub(r'\2\1', spaces.sub('', text)))
        if text == previous:
            break
    return text


def _note_positions(original, corrected, positions):
    """Pozycje odsyłaczy w poprawionym tekście: tuż za obrazem znaku, za którym stały w oryginale."""
    image = [None] * len(original)
    for tag, a, b, c, d in _edit_opcodes(original, corrected):
        if tag in ('equal', 'replace'):
            for shift in range(min(b - a, d - c)):
                image[a + shift] = c + shift
    result, floor = [], 0
    for position in positions:
        target = next((image[i] + 1 for i in range(position - 1, -1, -1) if image[i] is not None), 0)
        # Odstępy po obu stronach odsyłacza złączone w jeden: odstęp zostaje za odsyłaczem.
        if (position < len(original) and original[position] in ' \t' and 0 < target < len(corrected)
                and corrected[target - 1] in ' \t' and corrected[target] not in ' \t'):
            target -= 1
        floor = max(floor, target)
        result.append(floor)
    return result


def _moved_note_blocks(before, after, marker):
    """Liczba bloków odsyłaczy (¹²), których sąsiedztwo zmieniła reguła note_references."""
    blocks = re.compile(re.escape(marker) + '+')

    def contexts(text):
        found = []
        for match in blocks.finditer(text):
            context = (text[match.start() - 1:match.start()], text[match.end():match.end() + 1])
            found.extend([(context, match.start())] * len(match[0]))
        return found

    return len({block for (context, block), (previous, _) in zip(contexts(after), contexts(before))
                if context != previous})


def _new_text_slot(group, anchor, after):
    """Pusty przebieg tekstu tuż za (after) albo przed odsyłaczem, z formatowaniem najbliższego tekstu grupy."""
    from copy import deepcopy
    position = next(index for index, (node, _) in enumerate(group) if node is anchor)
    before = [node for node, value in group[:position] if value is not None and node.tag not in DELETED_CONTAINERS]
    behind = [node for node, value in group[position + 1:] if value is not None and node.tag not in DELETED_CONTAINERS]
    source = (before[-1] if before and (after or not behind) else behind[0]).getparent()
    run = OxmlElement('w:r')
    properties = source.find(qn('w:rPr'))
    if properties is not None:
        properties = deepcopy(properties)
        for change in properties.findall(qn('w:rPrChange')):
            properties.remove(change)  # w:id zmiany formatowania musi być unikalny
        run.append(properties)
    node = OxmlElement('w:t')
    node.set(qn('xml:space'), 'preserve')
    node.text = ''
    run.append(node)
    if after:
        anchor.getparent().addnext(run)
    else:
        anchor.getparent().addprevious(run)
    return node, ''


def _correct_noted_group(group, enabled, tracker, options, boundary=None):
    """Korekta grupy z odsyłaczami przypisów (sloty z wartością None).

    Reguły widzą tekst bez odsyłaczy, odsyłacze wracają za znak, za którym stały, a reguła
    note_references przesuwa je w tekście. Odsyłacz zostaje tym samym elementem w tym samym
    miejscu XML – zmienia się wyłącznie tekst po obu jego stronach (także jako śledzone zmiany).
    """
    segments, anchors = [[]], []
    for node, value in group:
        if value is None:
            anchors.append(node)
            segments.append([])
        else:
            segments[-1].append((node, value))
    texts = [''.join(value for _, value in segment) for segment in segments]
    original = ''.join(texts)
    if not original:
        return
    # Odsyłacz na brzegu akapitu: tekst za nim (przed nim) nie zaczyna (nie kończy) akapitu.
    options = {**options, 'trim_start': options['trim_start'] and bool(texts[0]),
               'trim_end': options['trim_end'] and bool(texts[-1])}
    corrected = correct_editorial_text(original, enabled, **options)
    ends, offset = [], 0
    for text in texts[:-1]:
        offset += len(text)
        ends.append(offset)
    marker = next(chr(code) for code in range(0xE000, 0xF900) if chr(code) not in corrected)
    pieces, start = [], 0
    for position in _note_positions(original, corrected, ends):
        pieces.append(corrected[start:position])
        start = position
    marked = marker.join([*pieces, corrected[start:]])
    # Odsyłacza nigdy nie poprzedza odstęp; pełne miejsce wg norm – tylko z regułą note_references
    # (bez niej usunięty odstęp liczy reguła spaces).
    full = 'note_references' in enabled
    placed = _place_note_references(marked, marker, spaces_only=not full)
    counts = options['counts']
    changed = _moved_note_blocks(marked, placed, marker)
    if counts is not None and changed:
        key = 'note_references' if full else 'spaces'
        counts[key] = counts.get(key, 0) + changed
    marked = placed
    for index, (segment, text) in enumerate(zip(segments, marked.split(marker))):
        if not segment:
            if not text:
                continue
            segment = [_new_text_slot(group, anchors[index - 1] if index else anchors[0], after=index > 0)]
        _rewrite(segment, text, tracker, boundary)


FALLBACK_TAG = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback'
# Podział strony lub kolumny nie kończy zdania ani akapitu: tekst za nim na początku
# akapitu traktujemy jak początek akapitu (spacja wiodąca, kwestia dialogowa).
PAGE_BREAK_TYPES = ('page', 'column')


def _in_fallback(element):
    """Kopia zapasowa pola tekstowego (mc:Fallback) – poprawiamy ją, ale nie liczymy drugi raz."""
    return any(node.tag == FALLBACK_TAG for node in element.iterancestors())


def _replace_nonbreaking_spaces(stories, counts):
    """Spacje nierozdzielające i nietypowe (wąska, półfiretowa…) → zwykłe we wszystkich częściach,
    bez śledzenia zmian.

    Zamiana prawie nie zmienia wyglądu tekstu, więc w trybie śledzenia nie tworzymy dla niej
    (często tysięcy) niewidocznych zmian; podsumowanie podaje ich liczbę.
    """
    replaced = 0
    for story in stories:
        for node in story.root.iter(qn('w:t')):
            text = node.text or ''
            found = len(UNUSUAL_SPACES_RE.findall(text))
            if found and _deleted(node):
                continue  # tekst przeniesiony (w:moveFrom ma w:t) albo usunięty zostaje nietknięty
            if found:
                if counts is not None and not _in_fallback(node):
                    replaced += found
                node.text = text.translate(NONBREAKING_TABLE)
                node.set(qn('xml:space'), 'preserve')
    if counts is not None and replaced:
        counts['nonbreaking_spaces'] = counts.get('nonbreaking_spaces', 0) + replaced


def _replacement_summary(pairs):
    rows = [{'source': source, 'target': target, 'count': count} for (source, target), count in pairs.items()]
    return sorted(rows, key=lambda row: (-row['count'], row['source'].lower(), row['source'], row['target']))


def apply_editorial_corrections(doc, enabled=None, *, track_changes=False, replacements=None, summary=None):
    """Popraw tekst dokumentu w miejscu.

    track_changes: zmiany jako śledzone zmiany Worda (autor „Odkurzacz”).
    replacements: lista zamian słownikowych (None = DEFAULT_REPLACEMENTS).
    summary: słownik, do którego trafia liczba zmian każdej reguły (tylko > 0), liczba spacji
    nierozdzielających zamienionych na zwykłe („nonbreaking_spaces”) oraz lista zamian
    słownikowych „replacements”: [{"source", "target", "count"}], malejąco po count.
    Spacje nierozdzielające są zamieniane zawsze, także bez żadnej reguły.

    Istniejące zmiany śledzone i komentarze zostają (decyzja właściciela): tekst usunięty (w:del,
    w:moveFrom) nie jest zmieniany ani liczony, tekst wstawiony (w:ins, w:moveTo) jest poprawiany
    w swoim opakowaniu, a znaczniki komentarzy, zmian formatowania i znaku akapitu stoją na miejscu.
    Treści komentarzy (comments.xml) Odkurzacz nie poprawia – to uwagi do tekstu, nie tekst.
    """
    enabled = DEFAULT_EDITORIAL_RULES if enabled is None else frozenset(enabled)
    if enabled - ALL_EDITORIAL_RULES:
        raise ValueError('Nieznana reguła korekty edytorskiej.')
    counts = {} if summary is not None else None
    # Treść, tabele, pola tekstowe, nagłówki, stopki i przypisy.
    stories, commit = document_stories(doc)
    _replace_nonbreaking_spaces(stories, counts)
    if not enabled:
        commit()
        _store_summary(summary, counts)
        return
    compiled = compile_replacements(replacements) if 'user_word_corrections' in enabled else None
    tracker = _Tracker(doc) if track_changes else None
    field_depth = 0
    paragraphs = [(number, para) for number, story in enumerate(stories) for para in story.paragraphs]
    current_story = 0
    for paragraph_index, (story_number, para) in enumerate(paragraphs):
        if story_number != current_story:
            current_story, field_depth = story_number, 0  # pola nie przechodzą między częściami
        if paragraph_index % 10 == 0:
            report_progress('Odkurzanie – akapity', paragraph_index, len(paragraphs))
        groups, slots = [], []
        paragraph_counts = None if counts is None or _in_fallback(para._p) else counts

        def barrier(soft=False):
            if slots:
                groups.append(list(slots))
                slots.clear()
            groups.append(soft if soft is LINE_BARRIER else (SOFT_BARRIER if soft else None))

        # Korekta widzi sąsiednie runy (także w cudzych wstawkach, hiperłączach i kontrolkach),
        # ale nie ingeruje w pola, rysunki, przypisy ani znaczniki dokumentu.
        for child in _paragraph_runs(para._p):
            if child.tag in DELETED_CONTAINERS:
                # Usunięty tekst nie przerywa zdania (po przyjęciu zmian go nie ma); granicę zapamiętuje
                # pusty slot – patrz _deletion_boundaries.
                slots.append((child, ''))
                continue
            if child.tag != qn('w:r'):
                barrier()
                continue
            for node in child:
                if node.tag in RUN_MARKERS:
                    continue
                if node.tag == qn('w:fldChar'):
                    kind = node.get(qn('w:fldCharType'))
                    if kind == 'begin':
                        field_depth += 1
                    elif kind == 'end':
                        field_depth = max(0, field_depth - 1)
                    barrier()
                elif field_depth:
                    barrier()
                elif node.tag == qn('w:t'):
                    if node.text:
                        slots.append((node, node.text))
                elif node.tag == qn('w:tab'):
                    slots.append((node, '\t'))
                elif node.tag in NOTE_REFERENCES and _note_reference_run(child):
                    slots.append((node, None))  # odsyłacz: nieruchoma kotwica, nie bariera
                elif node.tag == qn('w:cr') or (node.tag == qn('w:br')
                                                and node.get(qn('w:type'), 'textWrapping') == 'textWrapping'):
                    barrier(soft=LINE_BARRIER)  # miękki enter: koniec i początek linii
                else:
                    barrier(soft=node.tag == qn('w:br') and node.get(qn('w:type')) in PAGE_BREAK_TYPES)
        if slots:
            groups.append(slots)
        # Usunięty (śledzony) znak akapitu: po przyjęciu zmian akapit łączy się z następnym, więc jego
        # koniec i początek następnego nie są brzegami akapitu (spacja między nimi zostaje).
        joined_after = _mark_deleted(para._p)
        joined_before = _mark_deleted(_previous_paragraph(para._p))
        for index, group in enumerate(groups):
            if group and group not in (SOFT_BARRIER, LINE_BARRIER):
                group, boundary = _deletion_boundaries(group)
                if not group:
                    continue
                # Brzeg akapitu: przed (po) grupą są najwyżej podziały strony lub kolumny;
                # miękki enter tuż obok daje brzeg linii (spacje, kwestia dialogowa).
                options = dict(
                    trim_start=(index > 0 and groups[index - 1] is LINE_BARRIER)
                    or (not joined_before and all(item is SOFT_BARRIER for item in groups[:index])),
                    trim_end=(index + 1 < len(groups) and groups[index + 1] is LINE_BARRIER)
                    or (not joined_after and all(item is SOFT_BARRIER for item in groups[index + 1:])),
                    replacements=compiled, counts=paragraph_counts)
                if any(value is None for _, value in group):
                    _correct_noted_group(group, enabled, tracker, options, boundary)
                    continue
                corrected = correct_editorial_text(''.join(value for _, value in group), enabled, **options)
                _rewrite(group, corrected, tracker, boundary)
    commit()
    if 'empty_paragraphs' in enabled:
        previous_empty, removed = False, 0
        elements = list(doc._element.body)
        for index, element in enumerate(elements):
            empty = _safe_empty_paragraph(element, tracked=tracker is not None)
            if empty and previous_empty:
                removed += 1
                if tracker is None:
                    element.getparent().remove(element)
                else:
                    # Usuwamy znak poprzedniego pustego akapitu: po przyjęciu zmian
                    # puste akapity łączą się w jeden (ostatni z serii).
                    _delete_paragraph_tracked(elements[index - 1], tracker)
            else:
                previous_empty = empty
        if counts is not None and removed:
            counts['empty_paragraphs'] = removed
    _store_summary(summary, counts)


SOFT_BARRIER = object()
LINE_BARRIER = object()
# Opakowania tekstu, do których korekta wchodzi jak do zwykłych przebiegów.
TRANSPARENT_CONTAINERS = frozenset(qn('w:' + name) for name in (
    'ins', 'moveTo', 'hyperlink', 'sdt', 'sdtContent', 'smartTag', 'customXml', 'dir', 'bdo'))
CONTAINER_PROPERTIES = frozenset(qn('w:' + name) for name in ('sdtPr', 'sdtEndPr', 'smartTagPr', 'customXmlPr'))
DELETED_CONTAINERS = frozenset(qn('w:' + name) for name in ('del', 'moveFrom'))


def _deleted(node):
    """Węzeł wewnątrz usunięcia (w:del) albo źródła przeniesienia (w:moveFrom)."""
    return any(ancestor.tag in DELETED_CONTAINERS for ancestor in node.iterancestors())


def _paragraph_runs(paragraph):
    """Dzieci akapitu w kolejności tekstu: przebiegi także z wnętrza wstawek, hiperłączy
    i kontrolek; usunięty tekst (w:del, w:moveFrom) jako cały element; reszta jako bariery."""
    for child in paragraph:
        if child.tag in PARAGRAPH_MARKERS or child.tag in CONTAINER_PROPERTIES:
            continue
        if child.tag in DELETED_CONTAINERS:
            yield child
        elif child.tag in TRANSPARENT_CONTAINERS:
            yield from _paragraph_runs(child)
        else:
            yield child


def _deleted_text(element):
    """Tekst usuniętego fragmentu (w:delText, w:t w źródle przeniesienia, tabulatory)."""
    return ''.join('\t' if node.tag == qn('w:tab') else node.text or ''
                   for node in element.iter(qn('w:delText'), qn('w:t'), qn('w:tab')))


def _deletion_boundaries(group):
    """Granice cudzych usunięć w grupie (sloty (w:del albo w:moveFrom, '')) → (grupa, separator albo None).

    Odstęp tekstu widocznego, który styka się z wyrazem usuniętego tekstu, musi zostać: po odrzuceniu
    zmian bez niego wyrazy by się skleiły („i ” + usunięte „usunięty” + „ oraz” → „i usuniętyoraz”).
    Taka granica dostaje znak-separator spoza tekstu: reguły nie zbiją odstępów po obu jej stronach ani
    nie usuną odstępu przy niej, a tekst jest zapisywany osobno po każdej stronie (_rewrite). Granica bez
    odstępu obok znika z grupy – zdanie i wyraz rozcięty usunięciem są poprawiane w całości.
    """
    merged = []
    for node, value in group:
        deleted = value == '' and node.tag in DELETED_CONTAINERS
        if deleted and merged and merged[-1][0] is None:
            merged[-1][1].append(node)  # kilka usunięć tuż obok siebie to jedna granica
        elif deleted:
            merged.append((None, [node]))
        else:
            merged.append((node, value))
    visible = ''.join(value for node, value in merged if node is not None and value)
    boundary = next(chr(code) for code in range(0xE000, 0xF900) if chr(code) not in visible)
    result, used = [], False

    def neighbour(items):
        for node, value in items:
            if node is None or value == '':
                continue
            return 'x' if value is None else value  # odsyłacz przypisu to nie odstęp
        return ''

    for index, (node, value) in enumerate(merged):
        if node is not None:
            result.append((node, value))
            continue
        removed = ''.join(_deleted_text(element) for element in value)
        left, right = neighbour(reversed(merged[:index]))[-1:], neighbour(merged[index + 1:])[:1]
        if removed and ((left.isspace() and not removed[0].isspace())
                        or (right.isspace() and not removed[-1].isspace())):
            result.append((value[0], boundary))
            used = True
    return result, boundary if used else None


def _rewrite(slots, corrected, tracker, boundary=None):
    """Zapisz poprawiony tekst w węzłach (ze śledzeniem zmian albo bez).

    Separator granicy cudzego usunięcia (_deletion_boundaries) dzieli tekst na części zapisywane osobno:
    nic nie przechodzi na drugą stronę usuniętego fragmentu.
    """
    parts = [[]]
    for node, value in slots:
        if boundary is not None and value == boundary and node.tag in DELETED_CONTAINERS:
            parts.append([])
        else:
            parts[-1].append((node, value))
    pieces = corrected.split(boundary) if boundary is not None else [corrected]
    if len(pieces) != len(parts):  # reguła usunęła separator (nie powinna) – zapis bez podziału
        parts, pieces = [[slot for part in parts for slot in part]], [corrected.replace(boundary, '')]
    for index, part in enumerate(parts):
        if not part and pieces[index]:
            # Tekst dopisany po stronie bez węzłów (np. usunięcie na brzegu akapitu) – do sąsiedniej części.
            neighbour = index + 1 if index + 1 < len(parts) else index - 1
            pieces[neighbour] = (pieces[index] + pieces[neighbour] if neighbour > index
                                 else pieces[neighbour] + pieces[index])
            pieces[index] = ''
    for part, piece in zip(parts, pieces):
        if not part:
            continue
        if tracker is not None:
            _rewrite_tracked(part, piece, tracker)
        else:
            _rewrite_text_nodes(part, piece)


def _store_summary(summary, counts):
    if summary is None:
        return
    pairs = counts.pop('replacements', None)
    summary.update({key: value for key, value in counts.items() if value})
    if pairs:
        summary['replacements'] = _replacement_summary(pairs)


def clean_docx(source, rules, *, track_changes=False, replacements=None, summary=None):
    """Zwróć dokument w pamięci; nie zapisuj tekstu użytkownika na serwerze.

    track_changes=True zapisuje poprawki jako śledzone zmiany Worda,
    replacements zastępuje wbudowaną listę zamian słownikowych,
    summary (słownik) dostaje liczbę zmian każdej reguły.
    """
    source.seek(0)
    document = Document(source)
    stories, _ = document_stories(document)
    if any(len(paragraph.text) > 20_000 for story in stories for paragraph in story.paragraphs):
        if __package__:
            from .document_errors import DocumentInputError
        else:
            from document_errors import DocumentInputError
        raise DocumentInputError('cleaner_paragraph_limit')
    apply_editorial_corrections(document, frozenset(rules), track_changes=track_changes,
                                replacements=replacements, summary=summary)
    report_progress('Zapis odkurzonego DOCX')
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
