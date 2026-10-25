"""Recreate supported prose using new DOCX objects, never copying source parts.

Śledzone zmiany (wstawki, usunięcia, przeniesienia, znaki akapitu, zmiany formatowania) i komentarze (zakresy,
odpowiedzi, stan „rozwiązany”) są odtwarzane w nowym dokumencie z autorami i datami źródła oraz nowymi w:id –
zostają do przejrzenia w Wordzie albo LibreOffice. Pominięcia zgłaszają tylko to, co faktycznie ginie.
"""
from io import BytesIO
from collections import Counter
from copy import deepcopy
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.shared import RGBColor
from lxml import etree
from docx import Document
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

if __package__:
    from .document_styles import document_outline, merge_drop_caps, style_chain, unknown_format_value
    from .document_xml import CHOICE, FALLBACK, MC, mc, parse_part
else:
    from document_styles import document_outline, merge_drop_caps, style_chain, unknown_format_value
    from document_xml import CHOICE, FALLBACK, MC, mc, parse_part


class RebuildUnsupported(ValueError):
    def __init__(self, omissions):
        self.omissions = [f"{label} (liczba: {count})" for label, count in sorted(omissions.items())]
        super().__init__("; ".join(self.omissions))


THEME_FONT_OMISSION = "nierozpoznany krój czcionki z motywu (zastosujemy domyślny)"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
# Znaczniki bez treści (zakładki, ślady sprawdzania pisowni) – także między akapitami; znikają bez uwagi.
CONTENT_FREE = {
    qn("w:" + tag)
    for tag in (
        "bookmarkStart", "bookmarkEnd", "proofErr", "permStart", "permEnd",
        "customXmlInsRangeStart", "customXmlInsRangeEnd",
        "customXmlDelRangeStart", "customXmlDelRangeEnd", "customXmlMoveFromRangeStart",
        "customXmlMoveFromRangeEnd", "customXmlMoveToRangeStart", "customXmlMoveToRangeEnd",
    )
}
# Zakresy komentarzy i przeniesień: odtwarzane na swoich miejscach (w akapicie albo między akapitami).
RANGE_MARKERS = {
    qn("w:" + tag): tag
    for tag in (
        "commentRangeStart", "commentRangeEnd", "moveFromRangeStart", "moveFromRangeEnd",
        "moveToRangeStart", "moveToRangeEnd",
    )
}
# Śledzone zmiany treści (w kolejności schematu w znaku akapitu w:pPr/w:rPr); usunięcia i przeniesienia
# skąd niosą tekst, który znika po przyjęciu zmian.
REVISION_TAGS = {qn("w:" + tag): tag for tag in ("ins", "del", "moveFrom", "moveTo")}
DELETING = ("del", "moveFrom")
# Zapis poprzedniego formatowania przy śledzonych zmianach; przyjęcie zmian go usuwa.
FORMAT_HISTORY = tuple(
    qn("w:" + tag)
    for tag in (
        "rPrChange", "pPrChange", "sectPrChange", "tblPrChange", "tblPrExChange",
        "tblGridChange", "trPrChange", "tcPrChange", "numberingChange",
    )
)
FORMAT_HISTORY_OMISSION = "historia zmian formatowania (zmiany zostaną przyjęte)"
# Poprzednie formatowanie, którego przebudowa nie odtworzy (jak te same właściwości obecnego formatowania,
# zgłaszane w pominięciach): zmiana formatowania z nimi jest przyjmowana.
RUN_LOST = frozenset(qn("w:" + tag) for tag in ("spacing", "shd", "bdr"))
PARAGRAPH_LOST = frozenset(qn("w:" + tag) for tag in ("pBdr", "shd", "framePr", "textDirection"))
LOST_COMMENTS_OMISSION = "komentarze do pomijanych fragmentów (tabel, pól tekstowych, nagłówków i stopek)"
COMMENTS_EXTENDED = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"


MISSING = object()  # właściwość nieokreślona na tym poziomie: decyduje kolejny (styl, ustawienia domyślne)


class Omissions(Counter):
    def add(self, label):
        self[label] += 1


def _hex_byte(value):
    """Wartość themeShade/themeTint (00–FF); uszkodzona jest pomijana jak w Wordzie."""
    try:
        number = int(value, 16) if value else None
    except ValueError:
        return None
    return number if number is not None and 0 <= number <= 255 else None


def _format_resolver(original, omissions):
    # Resolve document defaults and theme fonts/colors without copying source parts.
    default_p = OxmlElement("w:p")
    default_r = OxmlElement("w:r")
    defaults = original.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        for path, target in (("w:rPrDefault/w:rPr", default_r), ("w:pPrDefault/w:pPr", default_p)):
            node = defaults
            for part in path.split("/"):
                node = node.find(qn(part)) if node is not None else None
            if node is not None:
                target.append(deepcopy(node))
    default_font = Run(default_r, Paragraph(default_p, original)).font
    default_format = Paragraph(default_p, original).paragraph_format
    theme = None
    for rel in original.part.rels.values():
        if rel.reltype.endswith("/theme") and not rel.is_external:
            theme = parse_part(rel.target_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

    def theme_font(key):
        if theme is None:
            return None
        family = "majorFont" if key.startswith("major") else "minorFont"
        nodes = theme.findall(".//a:fontScheme/a:" + family + "/a:latin", ns)
        return nodes[0].get("typeface") if nodes else None

    def font_name(fonts):
        for font in fonts:
            props = font._element.rPr
            node = props.find(qn("w:rFonts")) if props is not None else None
            if node is not None:
                key = node.get(qn("w:asciiTheme")) or node.get(qn("w:hAnsiTheme"))
                if key:
                    value = theme_font(key)
                    if value:
                        return value
                    # Brak czcionki motywu dotyczy całego dokumentu: jedna uwaga, nie jedna na fragment.
                    omissions[THEME_FONT_OMISSION] = 1
                value = node.get(qn("w:ascii")) or node.get(qn("w:hAnsi"))
                if value:
                    return value
        return None

    def font_color(fonts):
        for font in fonts:
            props = font._element.rPr
            node = props.find(qn("w:color")) if props is not None else None
            if node is None:
                continue
            value = node.get(qn("w:val"))
            key = node.get(qn("w:themeColor"))
            if key and key.isalnum() and theme is not None:
                key = {
                    "text1": "dk1",
                    "text2": "dk2",
                    "background1": "lt1",
                    "background2": "lt2",
                }.get(key, key)
                color = theme.find(".//a:clrScheme/a:" + key, ns)
                if color is not None and len(color):
                    value = color[0].get("lastClr") or color[0].get("val")
            if value and len(value) == 6 and all(c in "0123456789abcdefABCDEF" for c in value):
                rgb = [int(value[i : i + 2], 16) for i in (0, 2, 4)]
                shade = _hex_byte(node.get(qn("w:themeShade")))
                tint = _hex_byte(node.get(qn("w:themeTint")))
                if shade is not None:
                    rgb = [round(c * shade / 255) for c in rgb]
                if tint is not None:
                    rgb = [round(c + (255 - c) * (1 - tint / 255)) for c in rgb]
                return RGBColor(*rgb)
            if value == "auto":
                return None
        return MISSING

    return default_font, default_format, font_name, font_color


def _unwrap(node, content):
    parent = node.getparent()
    index = parent.index(node)
    for child in list(content) if content is not None else ():
        parent.insert(index, child)
        index += 1
    parent.remove(node)


HYPERLINK_OMISSION = "odnośniki hiperłączy (widoczny tekst pozostanie)"
TOC_OMISSION = "spis treści (wstaw go ponownie w Wordzie: Odwołania → Spis treści)"


PLACEHOLDER_STYLES = frozenset({"PlaceholderText"})
PLACEHOLDER_OMISSION = "oznaczenie tekstu zastępczego kontrolek z wpisaną treścią (treść pozostanie – sprawdź ją)"


def placeholder_styles(document):
    """Identyfikatory stylu „Tekst zastępczy” (w polskim Wordzie np. „Tekstzastpczy”)."""
    ids = set(PLACEHOLDER_STYLES)
    for style in document.styles.element.findall(qn("w:style")):
        name = style.find(qn("w:name"))
        if name is not None and (name.get(qn("w:val")) or "").casefold() == "placeholder text":
            ids.add(style.get(qn("w:styleId")))
    return ids


def _only_placeholder(node, styles):
    """Cała widoczna treść kontrolki ma styl tekstu zastępczego (albo kontrolka jest pusta)."""
    content = node.find(qn("w:sdtContent"))
    if content is None:
        return True
    for text in content.iter(qn("w:t")):
        if not (text.text or "").strip():
            continue
        run = text.getparent()
        style = run.find(qn("w:rPr") + "/" + qn("w:rStyle"))
        if style is None or style.get(qn("w:val")) not in styles:
            return False
    return True


def _unwrap_content_controls(body, omissions, styles=PLACEHOLDER_STYLES):
    """Kontrolki zawartości (w:sdt) znikają, a ich treść z formatowaniem zostaje na swoim miejscu.

    Dotyczy kontrolek blokowych, w akapicie, w wierszach i komórkach tabel. Kontrolka pokazująca tekst
    zastępczy („Kliknij tutaj…”) znika razem z nim – ale tylko wtedy, gdy cała jej treść ma styl tekstu
    zastępczego i nie ma w niej innych kontrolek. Kontrolka z oznaczeniem w:showingPlcHdr, w którą wpisano
    treść (np. program ją wypełnił albo grupa z wypełnionymi polami), jest rozpakowana z uwagą.
    Spis treści z galerii Worda (w:docPartObj) usuwamy w całości: jego numery stron nie pasowałyby
    do nowego dokumentu.
    """
    controls = list(body.iter(qn("w:sdt")))
    # Przed rozpakowaniem: kontrolki zawierające inne kontrolki (np. grupa pól).
    nested = {id(node) for node in controls if any(True for _ in node.iterdescendants(qn("w:sdt")))}
    # Od końca: zagnieżdżona kontrolka jest rozpakowana, zanim zajmiemy się zewnętrzną.
    for node in reversed(controls):
        parent = node.getparent()
        if parent is None:
            continue
        properties = node.find(qn("w:sdtPr"))
        # Elementy kontrolki nie mają klas python-docx (xpath bez prefiksów), więc szukamy przez find.
        gallery = node.find("/".join(qn(tag) for tag in ("w:sdtPr", "w:docPartObj", "w:docPartGallery")))
        if gallery is not None and (gallery.get(qn("w:val")) or "").startswith("Table of Contents"):
            omissions.add(TOC_OMISSION)
            parent.remove(node)
        elif properties is not None and properties.find(qn("w:showingPlcHdr")) is not None:
            if id(node) not in nested and _only_placeholder(node, styles):
                parent.remove(node)
            else:
                omissions.add(PLACEHOLDER_OMISSION)
                _unwrap(node, node.find(qn("w:sdtContent")))
        else:
            _unwrap(node, node.find(qn("w:sdtContent")))


def _unwrap_simple_fields(body, omissions, label):
    """Pole proste (w:fldSimple) zostawia swój ostatni wynik jako zwykły tekst – jak pole złożone."""
    fields = list(body.iter(qn("w:fldSimple")))
    for node in fields:
        if not any(parent.tag == node.tag for parent in node.iterancestors()):
            omissions.add(label)
    for node in fields:
        for data in node.findall(qn("w:fldData")):
            node.remove(data)
        _unwrap(node, node)


DROP_CAP_OMISSION = "inicjały (litera zostanie na początku akapitu, bez powiększenia)"
# Znaczniki własnego XML (Word 2007) i tagi inteligentne: Word 2010+ usuwa je przy otwarciu, a treść zostawia.
MARKUP_WRAPPERS = (qn("w:customXml"), qn("w:smartTag"))
MARKUP_PROPERTIES = (qn("w:customXmlPr"), qn("w:smartTagPr"))


def _sanitize_body(body, omissions, styles=PLACEHOLDER_STYLES):
    for node in reversed(list(body.iter(*MARKUP_WRAPPERS))):
        for properties in [child for child in node if child.tag in MARKUP_PROPERTIES]:
            node.remove(properties)
        _unwrap(node, node)
    drop_caps = merge_drop_caps(body)
    if drop_caps:
        omissions[DROP_CAP_OMISSION] += drop_caps
    # Treść alternatywna (mc:AlternateContent): bierzemy wersję zapasową dla prostszych programów,
    # a bez niej pierwszy wariant. Inaczej pole tekstowe liczyło się podwójnie (DrawingML + VML).
    for node in list(body.iter(mc("AlternateContent"))):
        if node.getparent() is not None:
            _unwrap(node, node.find(FALLBACK) if node.find(FALLBACK) is not None else node.find(CHOICE))
    _unwrap_content_controls(body, omissions, styles)
    # Śledzone zmiany w akapitach zostają (odtwarza je _paragraph_specs). Zmiana obejmująca całe akapity
    # (niezgodna ze schematem, ale bywa w generatorach) jest przyjmowana jak dawniej.
    for node in list(body.iter(*REVISION_TAGS)):
        if node.getparent() is not None and node.find(qn("w:p")) is not None:
            if REVISION_TAGS[node.tag] in DELETING:
                node.getparent().remove(node)
            else:
                _unwrap(node, node)
    unsupported = {
        "tbl": "tabele wraz z ich treścią",
        "drawing": "ilustracje i pola tekstowe wraz z ich treścią",
        "pict": "ilustracje i pola tekstowe wraz z ich treścią",
        "object": "obiekty osadzone",
        "fldChar": "mechanizm pól automatycznych (zwykły tekst wyniku może pozostać)",
        # Instrukcja (także usunięta śledzoną zmianą) jest częścią tego samego pola (liczonego przy fldChar).
        "instrText": None,
        "delInstrText": None,
        "altChunk": "dołączona treść zewnętrzna",
        "sym": "symbole specjalne zapisane jako obiekty",
    }
    for tag, label in unsupported.items():
        for node in list(body.iter(qn("w:" + tag))):
            # Jedna tabela czy ilustracja to jedna pozycja, także gdy zawiera zagnieżdżone; pole – jego początek.
            counted = label and not any(parent.tag == node.tag for parent in node.iterancestors())
            if tag == "fldChar":
                counted = node.get(qn("w:fldCharType")) == "begin"
            if counted:
                omissions.add(label)
            if node.getparent() is not None:
                node.getparent().remove(node)
    # Po usunięciu tabel i ilustracji: pole proste w nich znika razem z nimi i nie jest liczone.
    _unwrap_simple_fields(body, omissions, unsupported["fldChar"])
    math = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    for node in list(body.iter()):
        parent = node.getparent()
        # Jedno równanie to jeden element najwyższego poziomu; jego części nie są liczone osobno.
        if node.tag.startswith(math) and parent is not None and not parent.tag.startswith(math):
            omissions.add("równania matematyczne")
            parent.remove(node)
    # Po usunięciu tabel: zmiany ich formatowania znikają razem z nimi, bez osobnej uwagi.
    _format_history(body, omissions)
    for paragraph in body.iter(qn("w:p")):
        for path, label in (
            ("./w:pPr/w:pBdr", "obramowania akapitów"),
            ("./w:pPr/w:shd", "cieniowanie akapitów"),
            ("./w:pPr/w:framePr", "ramki akapitów"),
            ("./w:pPr/w:textDirection", "niestandardowy kierunek tekstu"),
            (".//w:rPr/w:spacing", "rozstrzelenie znaków"),
            (".//w:rPr/w:shd", "cieniowanie znaków"),
            (".//w:rPr/w:bdr", "obramowania znaków"),
        ):
            for _ in paragraph.xpath(path):
                omissions.add(label)


def _format_history(body, omissions):
    """Zmiany formatowania znaków i akapitów (w:rPrChange, w:pPrChange) zostają, gdy przebudowa odtworzy poprzednie
    formatowanie; z obramowaniem, cieniowaniem, rozstrzeleniem, ramką albo kierunkiem tekstu są przyjmowane z uwagą
    (poprzednie formatowanie z numeracją – w paragraph_spec, według reguły list). Zmiana układu sekcji też jest
    przyjmowana z uwagą (sekcja powstaje od nowa), a pozostałe (numeracja, tabele) – bez uwagi, jak dotąd."""
    reported = (qn("w:rPrChange"), qn("w:pPrChange"), qn("w:sectPrChange"))
    for node in list(body.iter(*FORMAT_HISTORY)):
        parent = node.getparent()
        if parent is None:
            continue
        if node.tag == qn("w:rPrChange"):
            previous = node.find(qn("w:rPr"))
            if previous is None or not any(child.tag in RUN_LOST for child in previous):
                continue
        elif node.tag == qn("w:pPrChange"):
            previous = node.find(qn("w:pPr"))
            if previous is None or not any(child.tag in PARAGRAPH_LOST for child in previous):
                continue
        if node.tag in reported:
            omissions.add(FORMAT_HISTORY_OMISSION)
        parent.remove(node)


LANGUAGE_ATTRIBUTES = ("val", "eastAsia", "bidi")


def _document_language(original):
    """Atrybuty w:lang z ustawień domyślnych dokumentu (np. pl-PL) – słownik i dzielenie wyrazów."""
    nodes = original.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang")
    if not nodes:
        return {}
    return {name: nodes[0].get(qn("w:" + name)) for name in LANGUAGE_ATTRIBUTES if nodes[0].get(qn("w:" + name))}


def _language(elements):
    """Język tekstu łacińskiego: fragment, potem styl znakowy i styl akapitu."""
    for element in elements:
        values = element.xpath("./w:rPr/w:lang/@w:val")
        if values:
            return values[0]
    return None


UNKNOWN_FORMAT_OMISSION = "nierozpoznana wartość formatowania (zastosujemy ustawienie ze stylu lub domyślne)"


def read_format(obj, name, omissions=None):
    """Właściwość formatowania python-docx; wartość spoza jego słownika (ValueError) nie przerywa
    przebudowy: „start”/„end” wyrównania są mapowane (document_styles), reszta jest pomijana i liczona
    w pominięciach."""
    try:
        return getattr(obj, name)
    except ValueError:
        value = unknown_format_value(obj, name)
        if value is None and omissions is not None:
            omissions.add(UNKNOWN_FORMAT_OMISSION)
        return value


def effective(objects, name, omissions=None):
    for obj in objects:
        value = read_format(obj, name, omissions)
        if value is not None:
            return value
    return None


PARAGRAPH_FORMAT = (
    "alignment",
    "first_line_indent",
    "left_indent",
    "right_indent",
    "space_before",
    "space_after",
    "line_spacing",
    "line_spacing_rule",
    "keep_together",
    "keep_with_next",
    "page_break_before",
    "widow_control",
)
RUN_FORMAT = (
    "size",
    "bold",
    "italic",
    "underline",
    "strike",
    "double_strike",
    "superscript",
    "subscript",
    "small_caps",
    "all_caps",
    "highlight_color",
)
BREAKS = {"page": WD_BREAK.PAGE, "column": WD_BREAK.COLUMN}


def _tab_stops(paragraph_format, omissions=None):
    """Tabulatory akapitu; tabulator o nieznanym wyrównaniu lub wypełnieniu jest pomijany i liczony."""
    stops = []
    for stop in paragraph_format.tab_stops:
        try:
            stops.append((stop.position, stop.alignment, stop.leader))
        except ValueError:
            if omissions is not None:
                omissions.add(UNKNOWN_FORMAT_OMISSION)
    return stops


class _StyleChains:
    """Formatowanie wynikające ze stylów i ustawień domyślnych, liczone raz na styl.

    Fragment dziedziczy po: własnym formatowaniu, stylu znakowym (z jego bazowymi), stylu akapitu
    (z bazowymi) i ustawieniach domyślnych dokumentu. Stylów jest kilkadziesiąt, fragmentów –
    dziesiątki tysięcy, więc część wspólną liczymy raz, a dla fragmentu czytamy tylko jego własny w:rPr."""

    def __init__(self, original, omissions):
        self.part = original.part
        self.omissions = omissions
        self.default_font, self.default_format, self.font_name, self.font_color = _format_resolver(
            original, omissions
        )
        self._styles = {}
        self._paragraphs = {}
        self._runs = {}

    def _style(self, style_id, kind):
        key = (style_id, kind)
        if key not in self._styles:
            self._styles[key] = list(style_chain(self.part.get_style(style_id, kind)))
        return self._styles[key]

    def paragraph(self, style_id):
        if style_id not in self._paragraphs:
            chain = self._style(style_id, WD_STYLE_TYPE.PARAGRAPH)
            formats = [s.paragraph_format for s in chain] + [self.default_format]
            fonts = [s.font for s in chain] + [self.default_font]
            tabs = next((_tab_stops(f, self.omissions) for f in formats if len(f.tab_stops)), [])
            self._paragraphs[style_id] = {
                "styles": chain,
                "format": {name: effective(formats, name, self.omissions) for name in PARAGRAPH_FORMAT},
                "tabs": tabs,
                "mark_size": effective(fonts, "size", self.omissions),
                "mark_name": self.font_name(fonts),
            }
        return self._paragraphs[style_id]

    def run(self, run_style_id, paragraph_style_id):
        key = (run_style_id, paragraph_style_id)
        if key not in self._runs:
            chain = self._style(run_style_id, WD_STYLE_TYPE.CHARACTER) + self.paragraph(
                paragraph_style_id
            )["styles"]
            fonts = [s.font for s in chain] + [self.default_font]
            color = self.font_color(fonts)
            self._runs[key] = {
                "format": {name: effective(fonts, name, self.omissions) for name in RUN_FORMAT + ("hidden",)},
                "name": self.font_name(fonts),
                "color": None if color is MISSING else color,
                "lang": _language([s.element for s in chain]),
            }
        return self._runs[key]


def _style_id(properties, tag):
    node = properties.find(qn(tag)) if properties is not None else None
    return node.get(qn("w:val")) if node is not None else None


NOTE_KINDS = ("footnote", "endnote")
NOTE_REFERENCE_TAGS = {qn("w:" + kind + "Reference"): kind for kind in NOTE_KINDS}
NOTE_MARK_TAGS = frozenset(qn("w:" + kind + "Ref") for kind in NOTE_KINDS)
NOTE_MISSING_OMISSION = "odsyłacze do nieistniejących przypisów"
# Style Worda (identyfikator, nazwa) tekstu przypisu i odsyłacza.
NOTE_STYLES = {
    "footnote": (("FootnoteText", "footnote text"), ("FootnoteReference", "footnote reference")),
    "endnote": (("EndnoteText", "endnote text"), ("EndnoteReference", "endnote reference")),
}
NOTE_RELATIONSHIPS = {"footnote": RT.FOOTNOTES, "endnote": RT.ENDNOTES}
# Style komentarzy Worda (identyfikator, nazwa): tekst komentarza i odwołanie do komentarza.
COMMENT_STYLES = (("CommentText", "annotation text"), ("CommentReference", "annotation reference"))


class _SourceNotes:
    """Przypisy dolne i końcowe źródła, odtwarzane w kolejności odsyłaczy (nowe w:id od 1).

    Treść przypisu przechodzi tę samą sanityzację i to samo odtwarzanie akapitów co treść główna,
    a pominięcia w niej liczą się razem z resztą. Przypis bez odsyłacza w tekście znika bez uwagi;
    odsyłacz do nieistniejącego przypisu znika z uwagą. Odsyłacz z własnym znakiem (customMarkFollows)
    zachowuje ten atrybut – sam znak to zwykły tekst za odsyłaczem i zostaje jak każdy tekst."""

    def __init__(self, original, omissions, placeholders):
        self.omissions = omissions
        self.placeholders = placeholders
        self.source = {kind: {} for kind in NOTE_KINDS}
        for rel in original.part.rels.values():
            for kind in NOTE_KINDS:
                if not rel.is_external and rel.reltype == NOTE_RELATIONSHIPS[kind]:
                    for note in parse_xml(rel.target_part.blob).findall(qn("w:" + kind)):
                        if note.get(qn("w:type"), "normal") == "normal":
                            self.source[kind].setdefault(note.get(qn("w:id")), note)
        self.notes = {kind: [] for kind in NOTE_KINDS}
        self._built = {}

    def reference(self, node, build):
        """Nowy przypis (rodzaj, w:id, własny znak) dla odsyłacza albo None, gdy przypisu brak."""
        kind = NOTE_REFERENCE_TAGS[node.tag]
        key = (kind, node.get(qn("w:id")))
        note = self.source[kind].get(key[1])
        if note is None:
            self.omissions.add(NOTE_MISSING_OMISSION)
            return None
        custom = node.get(qn("w:customMarkFollows")) in ("1", "true", "on")
        if key not in self._built:
            _sanitize_body(note, self.omissions, self.placeholders)
            for _ in note.iter(qn("w:hyperlink")):
                self.omissions.add(HYPERLINK_OMISSION)
            marks = []
            paragraphs = _annotation_paragraphs(build(note, marks), NOTE_STYLES[kind][0][0])
            # Znak przypisu zawsze na początku treści; przy własnym znaku tylko wtedy, gdy był w źródle.
            if marks or not custom:
                first = next(item for item in paragraphs if "marker" not in item)
                first["runs"].insert(0, {**(marks[0] if marks else {"font": {}}), "content": [("mark", kind)]})
            self._built[key] = paragraphs
        self.notes[kind].append(self._built[key])
        return kind, len(self.notes[kind]), custom


def _annotation_paragraphs(items, style):
    """Akapity przypisu albo komentarza: własny styl (także w poprzednim formatowaniu), bez nagłówków;
    pusta treść (albo same znaczniki zakresów) dostaje pusty akapit."""
    if all("marker" in item for item in items):
        items.append({"heading": None, "format": {}, "tabs": [], "runs": [], "mark": {}})
    for item in items:
        if "marker" not in item:
            item["heading"] = None
            item["style"] = style
            if "change" in item:
                item["change"]["style"] = style
    return items


def _change_attributes(node):
    """Autor i data śledzonej zmiany (w:author jest w schemacie wymagany, w:date – nie)."""
    return node.get(qn("w:author")) or "", node.get(qn("w:date"))


class _SourceComments:
    """Komentarze źródła odtwarzane w nowej części comments.xml (i commentsExtended.xml dla odpowiedzi i stanu
    „rozwiązany”), w kolejności pierwszego wystąpienia w tekście (nowe w:id od 0).

    Komentarz zostaje, gdy zostaje jego odwołanie (w:commentReference) – w treści albo w przypisie. Brakujący
    koniec zakresu staje przed odwołaniem, koniec bez początku znika (komentarz punktowy). Komentarz, którego
    odwołanie było w źródle, ale znika z pomijaną treścią (tabela, pole tekstowe, nagłówek), jest liczony
    w pominięciach. Treść komentarza przechodzi tę samą sanityzację i to samo odtwarzanie akapitów co treść
    główna. Części commentsIds, commentsExtensible i people nie są potrzebne do odpowiedzi ani rozwiązania
    (Word odtwarza je sam), a people zawiera identyfikatory kont autorów – nie są przenoszone.
    """

    def __init__(self, original, omissions, placeholders):
        self.omissions, self.placeholders = omissions, placeholders
        self.source, self.threads = {}, {}
        referenced = {node.get(qn("w:id")) for node in original.element.body.iter(qn("w:commentReference"))}
        stories = (RT.HEADER, RT.FOOTER, RT.FOOTNOTES, RT.ENDNOTES)
        for rel in original.part.rels.values():
            if rel.is_external:
                continue
            if rel.reltype == RT.COMMENTS:
                for comment in parse_xml(rel.target_part.blob).findall(qn("w:comment")):
                    self.source.setdefault(comment.get(qn("w:id")), comment)
            elif rel.reltype == COMMENTS_EXTENDED:
                for entry in parse_part(rel.target_part.blob).iter(W15 + "commentEx"):
                    self.threads[entry.get(W15 + "paraId")] = (
                        entry.get(W15 + "paraIdParent"), entry.get(W15 + "done") in ("1", "true", "on"))
            elif rel.reltype in stories:
                root = parse_part(rel.target_part.blob)
                referenced.update(node.get(qn("w:id")) for node in root.iter(qn("w:commentReference")))
        self.referenced_in_source = referenced & set(self.source)
        self.found = {}  # w:id źródła → {"start", "end", "referenced"} w kolejności pierwszego wystąpienia
        self.moves = {}  # (rodzaj przeniesienia, w:id źródła) → {"start", "end"}
        self.closed = False

    def _comment(self, node):
        comment_id = node.get(qn("w:id"))
        if self.closed or comment_id not in self.source:
            return None
        return self.found.setdefault(comment_id, {"start": False, "end": False, "referenced": False}), comment_id

    def marker(self, node):
        """Opis znacznika zakresu (komentarza albo przeniesienia) do odtworzenia albo None."""
        tag = RANGE_MARKERS[node.tag]
        if tag.startswith("comment"):
            found = self._comment(node)
            if found is None:
                return None
            state, comment_id = found
            state["start" if tag.endswith("Start") else "end"] = True
            return ("comment", tag, comment_id)
        kind, edge = ("moveFrom" if tag.startswith("moveFrom") else "moveTo"), tag.endswith("Start")
        key = (kind, node.get(qn("w:id")))
        state = self.moves.setdefault(key, {"start": False, "end": False, "name": None})
        state["start" if edge else "end"] = True
        attributes = (node.get(qn("w:name")) or "",) + (_change_attributes(node) if edge else ("", None))
        if edge:
            state["name"] = attributes[0]
        return ("move", tag, key, attributes)

    def reference(self, node):
        """w:id źródła dla odwołania do istniejącego komentarza (przebieg z odwołaniem) albo None."""
        found = self._comment(node)
        if found is None:
            return None
        found[0]["referenced"] = True
        return found[1]

    def finish(self, build):
        """Treść zachowanych komentarzy (po odczycie całego tekstu i przypisów) i uwaga o komentarzach utraconych."""
        self.closed = True
        kept = [comment_id for comment_id, state in self.found.items() if state["referenced"]]
        lost = self.referenced_in_source - set(kept)
        if lost:
            self.omissions[LOST_COMMENTS_OMISSION] += len(lost)
        para_ids = {}
        for comment_id in kept:
            paragraphs = self.source[comment_id].findall(qn("w:p"))
            if paragraphs:
                para_ids[paragraphs[-1].get(qn("w14:paraId"))] = comment_id
        comments = []
        for comment_id in kept:
            source = self.source[comment_id]
            last = source.findall(qn("w:p"))
            parent, done = self.threads.get(last[-1].get(qn("w14:paraId")) if last else None, (None, False))
            _sanitize_body(source, self.omissions, self.placeholders)
            for _ in source.iter(qn("w:hyperlink")):
                self.omissions.add(HYPERLINK_OMISSION)
            paragraphs = _annotation_paragraphs(build(source, []), COMMENT_STYLES[0][0])
            first = next(item for item in paragraphs if "marker" not in item)
            first["runs"].insert(0, {"font": {}, "content": [("annotation", None)]})
            comments.append({
                "source": comment_id, "author": source.get(qn("w:author")) or "", "date": source.get(qn("w:date")),
                "initials": source.get(qn("w:initials")), "paragraphs": paragraphs,
                "parent": para_ids.get(parent) if parent else None, "done": done, "range": self.found[comment_id],
            })
        # Zakres przeniesienia zostaje z parą (początek i koniec) i drugą stroną przeniesienia o tej samej nazwie.
        named = {(key[0], state["name"]) for key, state in self.moves.items() if state["start"] and state["end"]}
        other = {"moveFrom": "moveTo", "moveTo": "moveFrom"}
        moves = {key: (key[0], state["name"]) in named and (other[key[0]], state["name"]) in named
                 for key, state in self.moves.items()}
        return {"comments": comments, "threads": bool(self.threads), "moves": moves}


def _paragraph_specs(original, omissions, notes, comments):
    """Opisy akapitów treści głównej; odsyłacze przypisów (notes) budują też opisy treści przypisów,
    a odwołania do komentarzy (comments) – na końcu opisy treści zachowanych komentarzy.

    Przebieg (i znacznik zakresu) niesie ścieżkę śledzonych zmian, w których leży w źródle: (klucz zmiany,
    rodzaj, (autor, data)) od zewnętrznej – sąsiednie przebiegi tej samej zmiany trafiają do jednego
    nowego w:ins/w:del, a zagnieżdżenia (usunięcie cudzej wstawki) zostają."""
    chains = _StyleChains(original, omissions)
    outline = document_outline(original)
    font_name, font_color = chains.font_name, chains.font_color
    document_language = _document_language(original).get("val")
    revision_keys = {}

    def blocks(container, marks=None):
        """Akapity kontenera (treść, przypis albo komentarz) i znaczniki zakresów między nimi; marks zbiera
        formatowanie znaku przypisu (w komentarzu – pusta lista: bez odsyłaczy przypisów)."""
        paragraphs = []
        for element in container:
            if element.tag == qn("w:sectPr"):
                continue
            if element.tag in RANGE_MARKERS:
                marker = comments.marker(element)
                if marker is not None:
                    paragraphs.append({"marker": marker})
                continue
            if element.tag in CONTENT_FREE:
                continue
            if element.tag != qn("w:p"):
                omissions.add("blok " + etree.QName(element).localname + " wraz z treścią")
                continue
            paragraphs.append(paragraph_spec(element, marks))
        return paragraphs

    def revisions(node):
        """Śledzone zmiany obejmujące węzeł w akapicie, od zewnętrznej."""
        path = []
        for ancestor in node.iterancestors():
            if ancestor.tag == qn("w:p"):
                break
            if ancestor.tag in REVISION_TAGS:
                key = revision_keys.setdefault(ancestor, len(revision_keys))
                path.append((key, REVISION_TAGS[ancestor.tag], _change_attributes(ancestor)))
        return tuple(reversed(path))

    def run_format(element_run, paragraph_style_id, old):
        """Formatowanie przebiegu (własne, ze stylu znakowego, stylu akapitu i ustawień domyślnych) albo None,
        gdy przebieg jest ukryty."""
        run_properties = element_run.rPr
        shared = chains.run(_style_id(run_properties, "w:rStyle"), paragraph_style_id)
        font = Run(element_run, old).font if run_properties is not None else None
        hidden = read_format(font, "hidden", omissions) if font is not None else None
        if hidden if hidden is not None else shared["format"]["hidden"]:
            return None
        formatting = {"font": {}}
        for name in RUN_FORMAT:
            value = read_format(font, name, omissions) if font is not None else None
            if value is None:
                value = shared["format"][name]
            if value is not None:
                formatting["font"][name] = value
        name = (font_name([font]) if font is not None else None) or shared["name"]
        if name:
            formatting["font"]["name"] = name
        color = font_color([font]) if font is not None else MISSING
        color = shared["color"] if color is MISSING else color
        if color is not None:
            formatting["color"] = color
        language = _language([element_run]) or shared["lang"]
        if language and language != document_language:
            formatting["lang"] = language
        return formatting

    def paragraph_format(element, old, inherited):
        """Formatowanie akapitu (własne albo ze stylu) i jego tabulatory."""
        properties = element.pPr
        own = old.paragraph_format if properties is not None else None
        result = {}
        for name in PARAGRAPH_FORMAT:
            value = read_format(own, name, omissions) if own is not None else None
            if value is None:
                value = inherited["format"][name]
            if value is not None:
                result[name] = value
        tabs = _tab_stops(own, omissions) if own is not None and len(own.tab_stops) else list(inherited["tabs"])
        return result, tabs

    def template_heading(element):
        """Styl szablonu „Heading N” dla nagłówka poziomu N (document_styles.ParagraphOutline), inaczej None."""
        level = outline.heading_level(element)
        return f"Heading {level}" if level else None

    def previous(change, tag):
        """Element z poprzednimi właściwościami ze zmiany formatowania, opakowany jak w źródle (w:r albo w:p)."""
        wrapper = OxmlElement(tag)
        properties = change.find(qn("w:rPr" if tag == "w:r" else "w:pPr"))
        if properties is not None:
            wrapper.append(deepcopy(properties))
        return wrapper

    def mark_font(properties, inherited, old):
        """Rozmiar i krój znaku akapitu (wyznaczają wysokość pustego akapitu, np. odstępu między scenami)."""
        mark_size, mark_name = inherited["mark_size"], inherited["mark_name"]
        if properties is not None:
            mark = OxmlElement("w:r")
            mark.append(deepcopy(properties))
            mark_font = Run(mark, old).font
            mark_size = mark_font.size if mark_font.size is not None else mark_size
            mark_name = font_name([mark_font]) or mark_name
        result = {}
        if mark_size is not None:
            result["size"] = mark_size
        if mark_name:
            result["name"] = mark_name
        return result

    def paragraph_spec(element, marks):
        old = Paragraph(element, original)
        properties = element.pPr
        paragraph_style_id = _style_id(properties, "w:pStyle")
        inherited = chains.paragraph(paragraph_style_id)
        if outline.is_list_item(element):
            omissions.add("automatyczne numerowanie i punktory list (tekst pozycji pozostanie)")
        paragraph = {"heading": template_heading(element), "format": {}, "tabs": [], "runs": [], "mark": {}}
        mark_properties = properties.find(qn("w:rPr")) if properties is not None else None
        paragraph["mark"] = mark_font(mark_properties, inherited, old)
        if mark_properties is not None:
            # Wstawiony, usunięty albo przeniesiony znak akapitu: akapity zostają rozdzielone jak w źródle.
            paragraph["revisions"] = sorted(
                ((REVISION_TAGS[child.tag], _change_attributes(child))
                 for child in mark_properties if child.tag in REVISION_TAGS),
                key=lambda item: tuple(REVISION_TAGS.values()).index(item[0]))
            change = mark_properties.find(qn("w:rPrChange"))
            if change is not None:
                paragraph["mark_change"] = {"attributes": _change_attributes(change),
                                            "mark": mark_font(change.find(qn("w:rPr")), inherited, old)}
        paragraph["format"], paragraph["tabs"] = paragraph_format(element, old, inherited)
        change = properties.find(qn("w:pPrChange")) if properties is not None else None
        if change is not None:
            before = previous(change, "w:p")
            before_style = chains.paragraph(_style_id(before.pPr, "w:pStyle"))
            if outline.is_list_item(before):
                omissions.add(FORMAT_HISTORY_OMISSION)  # numeracji po odrzuceniu zmiany nie będzie
            else:
                before_format, before_tabs = paragraph_format(before, Paragraph(before, original), before_style)
                paragraph["change"] = {"attributes": _change_attributes(change), "heading": template_heading(before),
                                       "format": before_format, "tabs": before_tabs}
        runs = paragraph["runs"]
        for element_run in element.iter(qn("w:r"), *RANGE_MARKERS):
            path = revisions(element_run)
            if element_run.tag != qn("w:r"):
                marker = comments.marker(element_run)
                if marker is not None:
                    runs.append({"marker": marker, "revision": path})
                continue
            formatting = run_format(element_run, paragraph_style_id, old)
            if formatting is None:
                continue
            formatting["revision"] = path
            deleted = any(kind in DELETING for _, kind, _ in path)
            change = element_run.rPr.find(qn("w:rPrChange")) if element_run.rPr is not None else None
            if change is not None:
                # Tekst ukryty przed zmianą formatowania: po jej odrzuceniu znów ukryty.
                before = run_format(previous(change, "w:r"), paragraph_style_id, old) or {"font": {"hidden": True}}
                formatting["change"] = {**before, "attributes": _change_attributes(change)}
            # Odsyłacz przypisu, odwołanie do komentarza i znak przypisu dostają własny przebieg (styl odsyłacza
            # nie obejmie tekstu obok); tekst za nimi trafia do kolejnego przebiegu o tym samym formatowaniu.
            new_run = {**formatting, "content": []}
            runs.append(new_run)
            for child in element_run:
                tag = child.tag
                if tag == qn("w:rPr"):
                    continue
                # own: własny przebieg (odsyłacz, odwołanie), None – element bez przebiegu, False – zwykła treść.
                own = None
                if tag in NOTE_MARK_TAGS and marks is not None:
                    marks.append(formatting)
                elif tag in NOTE_REFERENCE_TAGS and marks is None:
                    reference = notes.reference(child, blocks)
                    own = reference and ("note", reference)
                elif tag == qn("w:commentReference"):
                    reference = comments.reference(child)  # odwołanie do brakującego komentarza znika
                    own = None if reference is None else ("comment", reference)
                elif tag != qn("w:annotationRef"):  # znacznik komentarza: odtwarzany na początku jego treści
                    own = False
                if own is not False:
                    if new_run is not None and not new_run["content"]:
                        runs.pop()
                    if own:
                        runs.append({**formatting, "content": [own]})
                    new_run = None
                    continue
                if new_run is None:
                    new_run = {**formatting, "content": []}
                    runs.append(new_run)
                if tag == qn("w:t"):
                    new_run["content"].append(("text", child.text or ""))
                elif tag == qn("w:delText") and deleted:
                    new_run["content"].append(("deleted", child.text or ""))
                elif tag == qn("w:tab"):
                    new_run["content"].append(("tab", None))
                elif tag in (qn("w:br"), qn("w:cr")):
                    kind = child.get(qn("w:type"), "textWrapping")
                    new_run["content"].append(("break", BREAKS.get(kind, WD_BREAK.LINE)))
                elif tag == qn("w:noBreakHyphen"):
                    new_run["content"].append(("text", "‑"))
                elif tag == qn("w:softHyphen"):
                    new_run["content"].append(("text", "­"))
                elif tag == qn("w:lastRenderedPageBreak"):
                    continue
                else:
                    omissions.add("element " + etree.QName(child).localname + " wewnątrz tekstu")
        return paragraph

    paragraphs = blocks(original.element.body)
    return paragraphs, comments.finish(blocks)


PAGE_DIMENSIONS = (
    "page_width", "page_height", "orientation", "top_margin", "bottom_margin", "left_margin",
    "right_margin", "header_distance", "footer_distance", "gutter",
)
COLUMN_ATTRIBUTES = ("num", "space", "equalWidth", "sep")
# Ustawienia dokumentu dotyczące marginesów (strony lustrzane w książce, margines na oprawę u góry).
PAGE_SETTINGS = ("mirrorMargins", "gutterAtTop")


def _page_format(original):
    """Format strony pierwszej sekcji źródła: wymiary, orientacja, marginesy i kolumny.

    Przenosimy same wartości (liczby i atrybuty), budując nowe elementy – nie kopiujemy części źródła.
    """
    section = original.sections[0]
    columns = section._sectPr.find(qn("w:cols"))
    return {
        "dimensions": {name: getattr(section, name) for name in PAGE_DIMENSIONS},
        "columns": None if columns is None else (
            {name: columns.get(qn("w:" + name)) for name in COLUMN_ATTRIBUTES if columns.get(qn("w:" + name))},
            [
                {name: column.get(qn("w:" + name)) for name in ("w", "space") if column.get(qn("w:" + name))}
                for column in columns.findall(qn("w:col"))
            ],
        ),
        "settings": [
            name for name in PAGE_SETTINGS
            if original.settings.element.xpath("./w:%s[not(@w:val) or not(@w:val='0' or @w:val='false'"
                                               " or @w:val='off')]" % name)
        ],
    }


def _analyze_docx(source, keep_page_format=False):
    source.seek(0)
    original = Document(source)
    body = original.element.body
    omissions = Omissions()
    for rel in original.part.rels.values():
        if not rel.is_external and rel.reltype.endswith(("/header", "/footer")):
            root = parse_part(rel.target_part.blob)
            # Znak wodny i grafika w nagłówku nie mają w:t (np. VML v:textpath), a też znikną.
            if any((node.text or "").strip() for node in root.iter(qn("w:t"))) or any(
                True for _ in root.iter(qn("w:drawing"), qn("w:pict"), qn("w:object"))
            ):
                omissions.add("niepuste nagłówki lub stopki")
    placeholders = placeholder_styles(original)
    # Przed sanityzacją: odwołania do komentarzy w treści, która zniknie (tabele, pola tekstowe).
    comments = _SourceComments(original, omissions, placeholders)
    _sanitize_body(body, omissions, placeholders)
    # Po rozpakowaniu kontrolek zawartości: podział sekcji bywa w akapicie wewnątrz kontrolki.
    if len(original.sections) > 1:
        label = "o układzie pierwszej sekcji" if keep_page_format else "A4"
        omissions[f"podziały na sekcje i ustawienia ich układu (pozostanie jedna sekcja {label})"] = (
            len(original.sections) - 1
        )
    elif not keep_page_format and len(original.sections) and _custom_page_layout(original.sections[0]):
        # Marginesy nie są zgłaszane (Word i LibreOffice mają różne domyślne);
        # format strony, orientacja i kolumny zmieniają wygląd całej książki.
        omissions.add("format strony, orientacja lub kolumny (wynik: A4 pionowo, jedna kolumna)")
    # Liczymy po usunięciu tabel i pól tekstowych: odnośnik z nich nie zostawia „widocznego tekstu”.
    for link in body.iter(qn("w:hyperlink")):
        omissions.add(HYPERLINK_OMISSION)
    notes = _SourceNotes(original, omissions, placeholders)
    paragraphs, review = _paragraph_specs(original, omissions, notes, comments)
    layout = {
        "paragraphs": paragraphs,
        "notes": notes.notes,
        "review": review,
        # Dokument bez w:sectPr (minimalne eksportery) dostaje stronę A4 jak przy zwykłej przebudowie.
        "page": _page_format(original) if keep_page_format and len(original.sections) else None,
        "language": _document_language(original),
        "default_tab": original.settings.element.xpath("./w:defaultTabStop/@w:val"),
        "defaults": {
            etree.QName(node).localname
            for node in original.styles.element.xpath(
                "./w:docDefaults | ./w:docDefaults/w:rPrDefault | ./w:docDefaults/w:pPrDefault"
            )
        },
    }
    return layout, omissions


def inspect_docx(source, *, keep_page_format=False):
    """Return exactly the rebuild warnings without creating or saving an output DOCX."""
    _, omissions = _analyze_docx(source, keep_page_format)
    if omissions:
        raise RebuildUnsupported(omissions)


def _neutral_styles(result, language, present):
    """Szablon python-docx nie może dokładać własnego wyglądu (10 pt odstępu, 1,15 interlinii, 11 pt,
    Calibri, niebieskich nagłówków): formatowanie źródła przenosimy jawnie, a czego źródło nie
    określało, zostaje przy wartościach bazowych Worda – tak jak w oryginale."""
    styles = result.styles.element
    for node in styles.xpath("./w:docDefaults/w:rPrDefault/w:rPr | ./w:docDefaults/w:pPrDefault/w:pPr"):
        for child in list(node):
            node.remove(child)
    # Programy różnie traktują brak ustawień domyślnych (np. LibreOffice przyjmuje wtedy wartości
    # Worda 2007), więc odtwarzamy też samą obecność tych elementów.
    for name in ("rPrDefault", "pPrDefault"):
        if name not in present:
            for node in styles.xpath("./w:docDefaults/w:" + name):
                node.getparent().remove(node)
    if "docDefaults" not in present:
        for node in styles.xpath("./w:docDefaults"):
            styles.remove(node)
    if language and styles.xpath("./w:docDefaults/w:rPrDefault/w:rPr"):
        properties = styles.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
        node = OxmlElement("w:lang")
        for name, value in language.items():
            node.set(qn("w:" + name), value)
        properties.append(node)
    for style in styles.xpath("./w:style[@w:type='paragraph']"):
        for node in style.xpath("./w:rPr"):
            style.remove(node)
        for node in style.xpath("./w:pPr/*[not(self::w:outlineLvl)]"):
            node.getparent().remove(node)


def _new_document(layout):
    paragraphs = layout["paragraphs"]
    result = Document()
    _neutral_styles(result, layout["language"], layout["defaults"])
    settings = result.settings.element
    for node in settings.xpath("./w:proofState"):
        # Szablon deklaruje „sprawdzono pisownię”; nowy dokument Word ma sprawdzić od nowa.
        settings.remove(node)
    if layout["default_tab"]:
        settings.xpath("./w:defaultTabStop")[0].set(qn("w:val"), layout["default_tab"][0])
    for section in result.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = (
            Cm(2.5)
        )
        if layout["page"]:
            _apply_page_format(result, section, layout["page"])
    heading_styles = {None: None}

    def heading_style(heading):
        # Wyszukanie stylu po nazwie przegląda wszystkie style; robimy to raz na nazwę, nie na akapit.
        if heading not in heading_styles:
            heading_styles[heading] = result.styles[heading] if heading in result.styles else None
        return heading_styles[heading]

    review = _ReviewIds(layout["review"])
    body = result.element.body
    for item in paragraphs:
        if "marker" in item:
            node = review.marker(item["marker"])
            if node is not None:
                body.insert_element_before(node, "w:sectPr")
            continue
        if "change" in item and "style" not in item["change"]:
            style = heading_style(item["change"]["heading"])
            item["change"]["style"] = style.style_id if style is not None else None
        _fill_paragraph(result.add_paragraph(style=heading_style(item["heading"])), item, review)
    _add_notes(result, layout["notes"], review)
    _add_comments(result, layout["review"], review)
    props = result.core_properties
    for name in (
        "author",
        "last_modified_by",
        "title",
        "subject",
        "comments",
        "keywords",
        "category",
        "content_status",
        "identifier",
        "language",
        "version",
    ):
        setattr(props, name, "")
    for tag in ("dcterms:created", "dcterms:modified", "cp:lastPrinted"):
        for node in list(props._element.findall(qn(tag))):
            props._element.remove(node)
    props.revision = 1
    return result


class _ReviewIds:
    """Identyfikatory w nowym dokumencie: komentarze od 0 (w kolejności z _SourceComments), śledzone zmiany,
    zmiany formatowania i zakresy przeniesień – kolejne liczby za nimi, unikalne w całym pakiecie."""

    def __init__(self, review):
        comments = review["comments"]
        self.comments = {comment["source"]: number for number, comment in enumerate(comments)}
        self.ranges = {comment["source"]: comment["range"] for comment in comments}
        self.pairs = review["moves"]
        self.moves = {}
        # Znaczniki komentarza już zapisane: treść przypisu z dwoma odsyłaczami jest odtwarzana dwa razy.
        self.written = set()
        self.next = len(comments)
        self.paragraphs = 0

    def take(self):
        self.next += 1
        return str(self.next - 1)

    def para_id(self):
        """Kolejny w14:paraId (8 cyfr szesnastkowych, mniej niż 0x80000000)."""
        self.paragraphs += 1
        return f"{self.paragraphs:08X}"

    def revision(self, tag, attributes):
        """Nowy element śledzonej zmiany (w:ins, w:del, w:rPrChange…) z autorem i datą źródła."""
        author, date = attributes
        node = OxmlElement("w:" + tag)
        node.set(qn("w:id"), self.take())
        node.set(qn("w:author"), author)
        if date:
            node.set(qn("w:date"), date)
        return node

    def marker(self, marker):
        """Znacznik zakresu w nowym dokumencie albo None (komentarz pominięty, para niepełna)."""
        if marker[0] == "comment":
            _, tag, source = marker
            if (source not in self.comments or (tag == "commentRangeEnd" and not self.ranges[source]["start"])
                    or (tag, source) in self.written):
                return None
            self.written.add((tag, source))
            node = OxmlElement("w:" + tag)
            node.set(qn("w:id"), str(self.comments[source]))
            return node
        _, tag, key, (name, author, date) = marker
        if not self.pairs.get(key):
            return None  # zakres bez pary (druga połowa w pomijanej treści)
        if key not in self.moves:
            self.moves[key] = self.take()
        node = OxmlElement("w:" + tag)
        node.set(qn("w:id"), self.moves[key])
        if tag.endswith("Start"):
            node.set(qn("w:author"), author)
            if date:
                node.set(qn("w:date"), date)
            node.set(qn("w:name"), name)
        return node

    def comment_reference(self, run, source):
        """Odwołanie do komentarza w przebiegu; brakujący koniec zakresu staje tuż przed przebiegiem."""
        if ("commentReference", source) in self.written:
            return
        self.written.add(("commentReference", source))
        state = self.ranges[source]
        if state["start"] and not state["end"]:
            end = OxmlElement("w:commentRangeEnd")
            end.set(qn("w:id"), str(self.comments[source]))
            run.addprevious(end)
        run.style = COMMENT_STYLES[1][0]
        reference = OxmlElement("w:commentReference")
        reference.set(qn("w:id"), str(self.comments[source]))
        run.append(reference)


def _mark_properties(node, mark):
    """Krój i rozmiar znaku akapitu (w:rFonts, w:sz) dopisane do w:rPr w kolejności schematu."""
    if "name" in mark:
        fonts = OxmlElement("w:rFonts")
        for key in ("w:ascii", "w:hAnsi"):
            fonts.set(qn(key), mark["name"])
        node.append(fonts)
    if "size" in mark:
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(round(mark["size"].pt * 2)))
        node.append(size)
    return node


def _apply_run_format(run, specification):
    """Formatowanie przebiegu z opisu (font, kolor, język)."""
    for name, value in specification["font"].items():
        setattr(run.font, name, value)
    if "color" in specification:
        run.font.color.rgb = specification["color"]
    if "lang" in specification:
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), specification["lang"])
        run._r.get_or_add_rPr().insert_element_before(
            language, "w:eastAsianLayout", "w:specVanish", "w:oMath"
        )


def _previous_run_format(paragraph, change, review):
    """w:rPrChange z poprzednim formatowaniem przebiegu (odtworzonym jak obecne)."""
    element = OxmlElement("w:r")
    _apply_run_format(Run(element, paragraph), change)
    node = review.revision("rPrChange", change["attributes"])
    properties = OxmlElement("w:rPr")
    for child in list(element.rPr) if element.rPr is not None else ():
        properties.append(child)
    node.append(properties)
    return node


def _previous_paragraph_format(paragraph, change, review):
    """w:pPrChange z poprzednim formatowaniem akapitu (styl, wyrównanie, wcięcia, odstępy, tabulatory)."""
    element = OxmlElement("w:p")
    before = Paragraph(element, paragraph._parent)
    if change.get("style"):
        element.style = change["style"]
    for name, value in change["format"].items():
        setattr(before.paragraph_format, name, value)
    for stop in change["tabs"]:
        before.paragraph_format.tab_stops.add_tab_stop(*stop)
    node = review.revision("pPrChange", change["attributes"])
    properties = OxmlElement("w:pPr")
    for child in list(element.pPr) if element.pPr is not None else ():
        properties.append(child)
    node.append(properties)
    return node


def _revision_container(paragraph, stack, path, review):
    """Element, do którego trafia przebieg o ścieżce zmian path: otwarte w:ins/w:del wspólne z poprzednim
    przebiegiem zostają, pozostałe są zamykane, a brakujące – tworzone (stack: otwarte (klucz, element))."""
    common = 0
    while common < min(len(stack), len(path)) and stack[common][0] == path[common][0]:
        common += 1
    del stack[common:]
    container = stack[-1][1] if stack else paragraph
    for key, tag, attributes in path[common:]:
        wrapper = review.revision(tag, attributes)
        container.append(wrapper)
        stack.append((key, wrapper))
        container = wrapper
    return container


def _fill_paragraph(paragraph, item, review):
    """Formatowanie, przebiegi, śledzone zmiany i znaczniki akapitu według opisu z _paragraph_specs (treść
    główna, przypisy i komentarze)."""
    if "style" in item:
        paragraph._p.style = item["style"]
    for name, value in item["format"].items():
        setattr(paragraph.paragraph_format, name, value)
    for stop in item["tabs"]:
        paragraph.paragraph_format.tab_stops.add_tab_stop(*stop)
    revisions, mark_change = item.get("revisions", ()), item.get("mark_change")
    if item["mark"] or revisions or mark_change:
        # Kolejność schematu znaku akapitu: zmiany (w:ins, w:del…), formatowanie, w:rPrChange.
        mark = OxmlElement("w:rPr")
        for tag, attributes in revisions:
            mark.append(review.revision(tag, attributes))
        _mark_properties(mark, item["mark"])
        if mark_change:
            change = review.revision("rPrChange", mark_change["attributes"])
            change.append(_mark_properties(OxmlElement("w:rPr"), mark_change["mark"]))
            mark.append(change)
        paragraph._p.get_or_add_pPr().append(mark)
    if item.get("change"):
        paragraph._p.get_or_add_pPr().append(_previous_paragraph_format(paragraph, item["change"], review))
    stack = []
    for specification in item["runs"]:
        container = _revision_container(paragraph._p, stack, specification.get("revision", ()), review)
        if "marker" in specification:
            node = review.marker(specification["marker"])
            if node is not None:
                container.append(node)
            continue
        element = OxmlElement("w:r")
        container.append(element)
        run = Run(element, paragraph)
        for kind, value in specification["content"]:
            if kind == "text":
                run.add_text(value)
            elif kind == "deleted":
                node = OxmlElement("w:delText")
                node.text = value
                node.set(qn("xml:space"), "preserve")
                element.append(node)
            elif kind == "tab":
                run.add_tab()
            elif kind == "note":
                note_kind, number, custom = value
                element.style = NOTE_STYLES[note_kind][1][0]
                reference = OxmlElement("w:" + note_kind + "Reference")
                if custom:
                    reference.set(qn("w:customMarkFollows"), "1")
                reference.set(qn("w:id"), str(number))
                element.append(reference)
            elif kind == "mark":
                element.style = NOTE_STYLES[value][1][0]
                element.append(OxmlElement("w:" + value + "Ref"))
            elif kind == "annotation":
                element.style = COMMENT_STYLES[1][0]
                element.append(OxmlElement("w:annotationRef"))
            elif kind == "comment":
                review.comment_reference(element, value)
            else:
                run.add_break(value)
        _apply_run_format(run, specification)
        if "change" in specification:
            # Zmiana formatowania na końcu w:rPr (python-docx nie zna jej miejsca w schemacie).
            element.get_or_add_rPr().append(_previous_run_format(paragraph, specification["change"], review))


# Elementy ustawień (schemat w:settings), przed którymi stoją w:footnotePr i w:endnotePr.
SETTINGS_AFTER_NOTES = tuple(
    qn(tag)
    for tag in (
        "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr", "w:attachedSchema", "w:themeFontLang",
        "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
        "w:captions", "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults",
        "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
    )
)


def _add_style(result, kind, style_id, name, properties=""):
    """Neutralny styl (akapitowy albo znakowy) o identyfikatorze i nazwie wbudowanej Worda, gdy go brak."""
    if result.styles.element.get_by_id(style_id) is None:
        based = "Normal" if kind == "paragraph" else "DefaultParagraphFont"
        result.styles.element.append(parse_xml(
            f'<w:style {nsdecls("w")} w:type="{kind}" w:styleId="{style_id}"><w:name w:val="{name}"/>'
            f'<w:basedOn w:val="{based}"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/>{properties}</w:style>'))


def _fill_story(result, container, items, review):
    """Akapity (i znaczniki zakresów między nimi) przypisu albo komentarza w nowym elemencie container."""
    for item in items:
        if "marker" in item:
            node = review.marker(item["marker"])
            if node is not None:
                container.append(node)
            continue
        element = OxmlElement("w:p")
        container.append(element)
        _fill_paragraph(Paragraph(element, result._body), item, review)


def _add_notes(result, notes, review):
    """Nowe części przypisów dolnych i końcowych z separatorami jak w Wordzie (w:id -1 i 0).

    Style tekstu przypisu i odsyłacza są neutralne jak pozostałe (_neutral_styles): formatowanie
    źródła niosą akapity i przebiegi, a styl odsyłacza daje tylko indeks górny, jak w Wordzie.
    """
    settings = result.settings.element
    for kind in NOTE_KINDS:
        if not notes.get(kind):
            continue
        (text_id, text_name), (reference_id, reference_name) = NOTE_STYLES[kind]
        _add_style(result, "paragraph", text_id, text_name)
        _add_style(result, "character", reference_id, reference_name,
                   '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>')
        separators = "".join(
            f'<w:{kind} w:type="{name}" w:id="{number}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
            f'w:lineRule="auto"/></w:pPr><w:r><w:{name}/></w:r></w:p></w:{kind}>'
            for number, name in ((-1, "separator"), (0, "continuationSeparator"))
        )
        root = parse_xml(f'<w:{kind}s {nsdecls("w")}>{separators}</w:{kind}s>')
        for number, paragraphs in enumerate(notes[kind], 1):
            note = OxmlElement("w:" + kind)
            note.set(qn("w:id"), str(number))
            root.append(note)
            _fill_story(result, note, paragraphs, review)
        part = XmlPart(
            PackURI(f"/word/{kind}s.xml"),
            f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
            root,
            result.part.package,
        )
        result.part.relate_to(part, NOTE_RELATIONSHIPS[kind])
        properties = parse_xml(
            f'<w:{kind}Pr {nsdecls("w")}><w:{kind} w:id="-1"/><w:{kind} w:id="0"/></w:{kind}Pr>'
        )
        anchor = next((node for node in settings if node.tag in SETTINGS_AFTER_NOTES), None)
        if anchor is None:
            settings.append(properties)
        else:
            anchor.addprevious(properties)


def _add_comments(result, review_layout, review):
    """Nowa część comments.xml (autor, inicjały, data i treść każdego komentarza) i – gdy źródło miało wątki
    lub stan „rozwiązany” – commentsExtended.xml: w15:paraIdParent wskazuje komentarz nadrzędny, w15:done
    rozwiązanie. Obie części odwołują się do akapitów komentarzy przez nowe w14:paraId."""
    comments = review_layout["comments"]
    if not comments:
        return
    _add_style(result, "paragraph", *COMMENT_STYLES[0])
    _add_style(result, "character", *COMMENT_STYLES[1])
    root = parse_xml(f'<w:comments {nsdecls("w", "w14")} xmlns:mc="{MC}" mc:Ignorable="w14"/>')
    last = []
    for number, comment in enumerate(comments):
        node = OxmlElement("w:comment")
        node.set(qn("w:id"), str(number))
        node.set(qn("w:author"), comment["author"])
        if comment["date"]:
            node.set(qn("w:date"), comment["date"])
        if comment["initials"] is not None:
            node.set(qn("w:initials"), comment["initials"])
        root.append(node)
        _fill_story(result, node, comment["paragraphs"], review)
        for paragraph in node.findall(qn("w:p")):
            paragraph.set(qn("w14:paraId"), review.para_id())
        last.append(paragraph.get(qn("w14:paraId")))
    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.{}+xml"
    part = XmlPart(PackURI("/word/comments.xml"), content_type.format("comments"), root, result.part.package)
    result.part.relate_to(part, RT.COMMENTS)
    if not review_layout["threads"]:
        return
    extended = etree.Element(W15 + "commentsEx", nsmap={"w15": W15[1:-1], "mc": MC})
    extended.set(mc("Ignorable"), "w15")
    sources = {comment["source"]: number for number, comment in enumerate(comments)}
    for number, comment in enumerate(comments):
        entry = etree.SubElement(extended, W15 + "commentEx")
        entry.set(W15 + "paraId", last[number])
        if comment["parent"] in sources:
            entry.set(W15 + "paraIdParent", last[sources[comment["parent"]]])
        entry.set(W15 + "done", "1" if comment["done"] else "0")
    part = XmlPart(PackURI("/word/commentsExtended.xml"), content_type.format("commentsExtended"), extended,
                   result.part.package)
    result.part.relate_to(part, COMMENTS_EXTENDED)


def _apply_page_format(result, section, page):
    """Format strony źródła; czego źródło nie określało, zostaje jak przy zwykłej przebudowie."""
    for name, value in page["dimensions"].items():
        if value is not None:
            setattr(section, name, value)
    if page["columns"] is not None:
        attributes, widths = page["columns"]
        columns = section._sectPr.find(qn("w:cols"))
        if columns is None:
            columns = OxmlElement("w:cols")
            section._sectPr.insert_element_before(columns, "w:formProt", "w:vAlign", "w:noEndnote",
                                                   "w:titlePg", "w:textDirection", "w:bidi",
                                                   "w:rtlGutter", "w:docGrid", "w:printerSettings",
                                                   "w:sectPrChange")
        for key in list(columns.attrib):
            del columns.attrib[key]
        for child in list(columns):
            columns.remove(child)
        for name, value in attributes.items():
            columns.set(qn("w:" + name), value)
        for width in widths:
            column = OxmlElement("w:col")
            for name, value in width.items():
                column.set(qn("w:" + name), value)
            columns.append(column)
    settings = result.settings.element
    anchor = settings.find(qn("w:zoom"))
    index = settings.index(anchor) + 1 if anchor is not None else 0
    for name in page["settings"]:
        settings.insert(index, OxmlElement("w:" + name))
        index += 1


def _custom_page_layout(section):
    """Strona inna niż A4 lub Letter pionowo albo w kilku kolumnach (tolerancja 2 mm).

    Letter (domyślny w angielskim Wordzie) zamieniony na A4 nie zmienia układu książki,
    więc nie wymaga potwierdzenia; A5, B5, własny format książkowy i orientacja pozioma tak.
    """
    width, height = section.page_width, section.page_height
    if width is not None and height is not None and not any(
        abs(width - Cm(w)) <= Cm(0.2) and abs(height - Cm(h)) <= Cm(0.2)
        for w, h in ((21, 29.7), (21.59, 27.94))
    ):
        return True
    columns = section._sectPr.find(qn("w:cols"))
    try:
        return columns is not None and int(columns.get(qn("w:num"), "1") or 1) > 1
    except ValueError:
        return False


def rebuild_docx(source, *, allow_omissions=False, keep_page_format=False):
    layout, omissions = _analyze_docx(source, keep_page_format)
    if omissions and not allow_omissions:
        raise RebuildUnsupported(omissions)
    result = _new_document(layout)
    output = BytesIO()
    result.save(output)
    output.seek(0)
    return output
