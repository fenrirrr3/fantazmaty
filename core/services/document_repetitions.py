"""Repeat analysis from the desktop application, independent of Django/GUI."""
import random
import re
from collections import OrderedDict
from threading import RLock
from copy import deepcopy
from dataclasses import dataclass
from functools import lru_cache
from typing import Optional
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.run import Run
from docx.enum.text import WD_COLOR_INDEX

# The converter worker runs this module without the package context.
try:
    from .document_progress import report as report_progress
    from .document_errors import DocumentInputError
    from .document_stories import document_stories
except ImportError:
    from document_progress import report as report_progress
    from document_errors import DocumentInputError
    from document_stories import document_stories

WINDOW_SIZE = 35
MIN_WORD_LENGTH = 4
MODEL_NAME = 'pl_core_news_sm'
BATCH_SIZE = 128


@lru_cache(maxsize=1)
def _get_nlp():
    import spacy
    # NER nie jest używany; zachowujemy parser i zależności lematyzatora.
    try:
        return spacy.load(MODEL_NAME, exclude=['ner'])
    except OSError as error:
        raise ImportError('Brakuje polskiego modelu analizy.') from error


def _lemma_from_doc(doc, fallback: str) -> str:
    if not doc or not doc[0].lemma_:
        return fallback
    return doc[0].lemma_.lower()


def _get_lemma_lower(word: str) -> str:
    return _lemma_map([word])[word.lower()]


def get_lemma(word: str) -> str:
    return _get_lemma_lower(word.lower())


def dark_rgb() -> RGBColor:
    return RGBColor(*(random.randint(0, 150) for _ in range(3)))


def light_rgb() -> RGBColor:
    # Avoid near-white text on the white pages of generated documents.
    return RGBColor(*(random.randint(160, 220) for _ in range(3)))


def first5(word: str) -> str:
    return word.lower()[:5]


@dataclass
class Token:
    text: str
    is_word: bool
    lemma: Optional[str] = None
    stem5: Optional[str] = None
    word_len: int = 0
    color: Optional[RGBColor] = None
    highlight: Optional[int] = None
    ignored: bool = False


def _find_repeat_groups(words, progress=None, window_size=WINDOW_SIZE,
                        include_prefix_matches=False):
    """Indeks pozycji; sąsiednie pozycje wystarczą do wykrycia członków grupy.

    Kolejność grup odpowiada pierwszej parze w dawnej pętli i,j.
    Wyjątki zajmują pozycje w oknie, ale nie trafiają do indeksu.
    """
    buckets = {}
    for i, token in enumerate(words):
        if not token.ignored:
            if token.lemma:
                buckets.setdefault(('lemma', token.lemma, None), []).append(i)
            if include_prefix_matches and token.stem5:
                category = 5 if token.word_len == 5 else 6
                buckets.setdefault(('stem5', token.stem5, category), []).append(i)
        if progress is not None:
            progress.update(1)
    groups, first_pairs = {}, {}
    for (kind, value, _), positions in buckets.items():
        key = (kind, value)
        for left, right in zip(positions, positions[1:]):
            if right - left <= window_size:
                groups.setdefault(key, set()).update((left, right))
                pair = (left, right, 0 if kind == 'lemma' else 1)
                first_pairs[key] = min(first_pairs.get(key, pair), pair)
    return {key: groups[key] for key in sorted(groups, key=first_pairs.__getitem__)}


def _assign_colors(words, groups, color_palette='dark'):
    for indices in groups.values():
        color = light_rgb() if color_palette == 'light' else dark_rgb()
        for idx in indices:
            if words[idx].color is None:
                words[idx].color = color


def _color_spans(tokens):
    """Scal sąsiednie fragmenty o tym samym kolorze w obrębie runu."""
    spans = []
    end = 0
    for token in tokens:
        if not token.text:
            continue
        end += len(token.text)
        if spans and spans[-1][1:] == (token.color, token.highlight):
            spans[-1] = (end, token.color, token.highlight)
        else:
            spans.append((end, token.color, token.highlight))
    return spans


def _child_text(child):
    if child.tag == qn('w:t'):
        return child.text or ''
    # Używamy interpretacji tekstu bieżącej wersji python-docx, także dla
    # tabulatorów, podziałów wiersza/strony i nierozdzielających łączników.
    probe = OxmlElement('w:r')
    probe.append(deepcopy(child))
    return probe.text or ''


def _apply_run_colors(run, tokens):
    spans = _color_spans(tokens)
    protected_nodes = any(child.tag not in (qn('w:rPr'), qn('w:t'), qn('w:tab'))
                          for child in run._r)
    if not spans or (len(spans) <= 1 and not protected_nodes):
        if spans and spans[0][1] is not None:
            run.font.color.rgb = spans[0][1]
        if spans and spans[0][2] is not None:
            run.font.highlight_color = spans[0][2]
        return

    original = run._r
    children = [(child, _child_text(child)) for child in original
                if child.tag != qn('w:rPr')]
    if ''.join(text for _, text in children) != run.text:
        raise DocumentInputError('unsupported_run')

    rpr = original.rPr
    replacements = []

    def new_run(color, highlight):
        element = OxmlElement('w:r')
        element.attrib.update(original.attrib)
        if rpr is not None:
            element.append(deepcopy(rpr))
        wrapper = Run(element, run._parent)
        if color is not None:
            wrapper.font.color.rgb = color
        if highlight is not None:
            wrapper.font.highlight_color = highlight
        replacements.append(element)
        return element

    span_idx, offset = 0, 0
    current, current_style = None, None
    for child, text in children:
        if child.tag not in (qn('w:t'), qn('w:tab')):
            # Odwołania, rysunki i znaczniki pól zachowują oryginalne rPr,
            # aby kolor sąsiedniego słowa nie zmieniał np. numeru przypisu.
            isolated = new_run(None, None)
            isolated.append(deepcopy(child))
            offset += len(text)
            current, current_style = None, None
            continue
        if not text:
            if current is None:
                current = new_run(None, None)
                current_style = (None, None)
            current.append(deepcopy(child))
            continue
        consumed = 0
        while consumed < len(text):
            while offset >= spans[span_idx][0]:
                span_idx += 1
            style = spans[span_idx][1:]
            if current is None or current_style != style:
                current = new_run(*style)
                current_style = style
            take = min(len(text) - consumed, spans[span_idx][0] - offset)
            copied = deepcopy(child)
            if child.tag == qn('w:t'):
                copied.text = text[consumed:consumed + take]
                copied.set(qn('xml:space'), 'preserve')
            elif consumed or take != len(text):
                raise DocumentInputError('unsupported_run')
            current.append(copied)
            consumed += take
            offset += take

    for element in replacements:
        original.addprevious(element)
    original.getparent().remove(original)


def _positive_integer(value, label):
    if isinstance(value, bool):
        raise ValueError(f'{label}: wpisz dodatnią liczbę całkowitą.')
    text = str(value).strip()
    if not re.fullmatch(r'[0-9]+', text) or int(text) < 1:
        raise ValueError(f'{label}: wpisz dodatnią liczbę całkowitą.')
    return int(text)


LEXICAL_RE = re.compile(r"[^\W\d_]+(?:[-’'][^\W\d_]+)*", re.UNICODE)
EMPTY_PAIRS_RE = re.compile(r'\([ \t\u00a0]*\)|\[[ \t\u00a0]*\]|\{[ \t\u00a0]*\}|„[ \t\u00a0]*”|“[ \t\u00a0]*”|"[ \t\u00a0]*"')
CHECK_DEFAULTS = {
    'duplicates': True, 'long_sentences': True, 'long_paragraphs': True,
    'empty_pairs': True, 'sentence_limit': 35, 'paragraph_limit': 150,
}
MARK_STYLES = {
    'long_paragraphs': (1, WD_COLOR_INDEX.GRAY_25, 'Długie akapity'),
    'long_sentences': (2, WD_COLOR_INDEX.TURQUOISE, 'Długie zdania'),
    'tracked': (3, WD_COLOR_INDEX.YELLOW, 'Własne słowa'),
    'duplicates': (4, WD_COLOR_INDEX.BRIGHT_GREEN, 'Sąsiednie powtórzenia'),
    'empty_pairs': (5, WD_COLOR_INDEX.PINK, 'Puste nawiasy / cudzysłowy'),
}


def parse_word_list(value):
    """Jedno słowo na wiersz lub słowa rozdzielone spacjami, przecinkami, średnikami."""
    text = value if isinstance(value, str) else ' '.join(value or ())
    return tuple(dict.fromkeys(match[0].lower() for match in LEXICAL_RE.finditer(text)))


_LEMMA_CACHE_LIMIT = 50000
_LEMMA_CACHE = OrderedDict()
_LEMMA_LOCK = RLock()


def _lemma_map(forms):
    unique = dict.fromkeys(form.lower() for form in forms)
    if not unique:
        return {}
    # Klucz zawiera nazwę modelu. Cache jest współdzielony z get_lemma().
    with _LEMMA_LOCK:
        result, missing = {}, []
        for word in unique:
            key = (MODEL_NAME, word)
            if key in _LEMMA_CACHE:
                result[word] = _LEMMA_CACHE[key]
                _LEMMA_CACHE.move_to_end(key)
            else:
                missing.append(word)
        if missing:
            nlp = _get_nlp()
            # Parser jest potrzebny przy zdaniach, nie przy izolowanych słowach.
            disabled = [name for name in ('parser', 'senter', 'ner') if name in nlp.pipe_names]
            for word, parsed in zip(missing, nlp.pipe(missing, batch_size=BATCH_SIZE,
                                                    disable=disabled)):
                lemma = _lemma_from_doc(parsed, word)
                result[word] = lemma
                _LEMMA_CACHE[(MODEL_NAME, word)] = lemma
                while len(_LEMMA_CACHE) > _LEMMA_CACHE_LIMIT:
                    _LEMMA_CACHE.popitem(last=False)
        return result


def _paragraph_layout(para, field_state=None):
    """Tekst analizy oraz mapowanie (run, start, end, run_start, run_end).

    Znaki barier nie należą do tekstu Worda. Pozycje runów odwołują się
    do oryginalnego run.text, więc oznaczenia nie zmieniają zawartości XML.
    field_state pozwala chronić pola przechodzące przez granicę akapitu.
    """
    state = field_state if field_state is not None else [0]
    text, pieces, offset = [], [], 0

    def barrier():
        nonlocal offset
        text.append('\ufffc')
        offset += 1

    for child in para._p:
        if child.tag == qn('w:pPr'):
            continue
        if child.tag != qn('w:r'):
            barrier()
            continue
        run, run_offset = Run(child, para), 0
        for node in child:
            if node.tag == qn('w:rPr'):
                continue
            value = _child_text(node)
            if node.tag == qn('w:fldChar'):
                kind = node.get(qn('w:fldCharType'))
                if kind == 'begin':
                    state[0] += 1
                elif kind == 'end':
                    state[0] = max(0, state[0] - 1)
                barrier()
            elif state[0] or node.tag not in (qn('w:t'), qn('w:tab')):
                barrier()
            elif value:
                pieces.append((run, offset, offset + len(value),
                               run_offset, run_offset + len(value)))
                text.append(value)
                offset += len(value)
            run_offset += len(value)
        if run_offset != len(run.text or ''):
            raise DocumentInputError('unsupported_run')
    return ''.join(text), pieces


def _document_layout(paragraphs):
    state = [0]
    return [_paragraph_layout(para, state) for para in paragraphs]


def _resolved_spans(text, marks):
    """Rozstrzygnij kolor liter i priorytet tła jednym przebiegiem."""
    import heapq
    events = {}
    for ident, (start, end, color, highlight, priority) in enumerate(marks):
        if not 0 <= start < end <= len(text):
            continue
        events.setdefault(start, []).append((True, ident, color, highlight, priority))
        events.setdefault(end, []).append((False, ident, color, highlight, priority))
    if not events:
        return []
    active, colors, backgrounds, spans = set(), [], [], []
    boundaries = sorted({0, len(text), *events})
    for start, end in zip(boundaries, boundaries[1:]):
        for adding, ident, color, highlight, priority in events.get(start, ()):
            if adding:
                active.add(ident)
                if color is not None:
                    heapq.heappush(colors, (ident, color))
                if highlight is not None:
                    heapq.heappush(backgrounds, (-priority, ident, highlight))
            else:
                active.discard(ident)
        while colors and colors[0][0] not in active:
            heapq.heappop(colors)
        while backgrounds and backgrounds[0][1] not in active:
            heapq.heappop(backgrounds)
        spans.append((start, end, colors[0][1] if colors else None,
                      backgrounds[0][2] if backgrounds else None))
    return spans


def _apply_layout_marks(layout, marks):
    """Przenieś oznaczenia na istniejące runy; podziel każdy najwyżej raz."""
    from bisect import bisect_right
    text, pieces = layout
    spans = _resolved_spans(text, marks)
    if not spans:
        return
    ends = [end for _, end, _, _ in spans]
    run_marks = {}
    for run, start, end, run_start, _ in pieces:
        index = bisect_right(ends, start)
        while index < len(spans) and spans[index][0] < end:
            a, b, color, highlight = spans[index]
            if color is not None or highlight is not None:
                entry = run_marks.setdefault(id(run._r), (run, []))
                entry[1].append((run_start + max(start, a) - start,
                                 run_start + min(end, b) - start, color, highlight))
            index += 1
    for run, ranges in run_marks.values():
        value, tokens, position = run.text or '', [], 0
        for start, end, color, highlight in ranges:
            if position < start:
                tokens.append(Token(value[position:start], False))
            tokens.append(Token(value[start:end], False, color=color, highlight=highlight))
            position = end
        if position < len(value):
            tokens.append(Token(value[position:], False))
        _apply_run_colors(run, tokens)


def _extra_checks(paragraphs, lemmas, ignored, tracked, options, status_callback=None,
                  layouts=None, apply_marks=True):
    checks = {**CHECK_DEFAULTS, **options}
    layouts = _document_layout(paragraphs) if layouts is None else layouts
    texts = [layout[0] for layout in layouts]
    marks = [[] for _ in paragraphs]
    report = {key: 0 for key in MARK_STYLES}
    previous = None
    for p_index, text in enumerate(texts):
        lexemes = list(LEXICAL_RE.finditer(text))
        if checks['long_paragraphs'] and len(lexemes) > checks['paragraph_limit']:
            marks[p_index].append((0, len(text), 'long_paragraphs'))
            report['long_paragraphs'] += 1
        for word in lexemes:
            form = word[0].lower()
            lemma = lemmas.get(form, form)
            if lemma in tracked and lemma not in ignored:
                marks[p_index].append((word.start(), word.end(), 'tracked'))
                report['tracked'] += 1
            if checks['duplicates'] and previous is not None:
                old_p, old_word, old_lemma = previous
                # Sąsiedztwo przez białe znaki, również przez granicę akapitu.
                if old_p == p_index:
                    separator = text[old_word.end():word.start()]
                else:
                    separator = texts[old_p][old_word.end():] + '\n' + text[:word.start()]
                    if paragraphs[old_p]._p.getnext() is not paragraphs[p_index]._p:
                        separator += '\ufffc'  # np. tabela pomiędzy akapitami
                if separator and separator.isspace() and lemma == old_lemma and lemma not in ignored:
                    marks[old_p].append((old_word.start(), old_word.end(), 'duplicates'))
                    marks[p_index].append((word.start(), word.end(), 'duplicates'))
                    report['duplicates'] += 1
            previous = (p_index, word, lemma)
        # Pusty akapit lub nietekstowy blok przerywa kontrolę sąsiedztwa.
        if not lexemes:
            previous = None
        if checks['empty_pairs']:
            for match in EMPTY_PAIRS_RE.finditer(text):
                marks[p_index].append((match.start(), match.end(), 'empty_pairs'))
                report['empty_pairs'] += 1
    if checks['long_sentences']:
        if status_callback:
            status_callback('Rozpoznawanie zdań i sprawdzanie ich długości…')
        # Model rozpoznaje zdania w kontekście, zamiast dzielić przy każdej kropce.
        nlp = _get_nlp()
        disabled = [name for name in ('lemmatizer', 'ner') if name in nlp.pipe_names]
        for index, parsed in enumerate(nlp.pipe(texts, batch_size=BATCH_SIZE, disable=disabled)):
            for sentence in parsed.sents:
                if len(LEXICAL_RE.findall(sentence.text)) > checks['sentence_limit']:
                    marks[index].append((sentence.start_char, sentence.end_char, 'long_sentences'))
                    report['long_sentences'] += 1
    if apply_marks:
        for layout, ranges in zip(layouts, marks):
            converted = [(a, b, None, MARK_STYLES[k][1], MARK_STYLES[k][0]) for a, b, k in ranges]
            _apply_layout_marks(layout, converted)
        return report
    return report, marks


def color_document(source, *, window_size=35, min_word_length=4,
                   ignored_words='', tracked_words='', analysis_options=None,
                   include_prefix_matches=False, color_palette='dark'):
    """Return a marked DOCX; never change its words or editorial formatting."""
    if color_palette not in ('dark', 'light'):
        raise DocumentInputError('analysis_options')
    window_size = _positive_integer(window_size, 'Zakres wyszukiwania')
    min_word_length = _positive_integer(min_word_length, 'Minimalna długość słowa')
    if window_size > 500 or min_word_length > 100:
        raise DocumentInputError('analysis_options')
    options = {**CHECK_DEFAULTS, **(analysis_options or {})}
    if set(options) - set(CHECK_DEFAULTS):
        raise DocumentInputError('analysis_options')
    for key in ('sentence_limit', 'paragraph_limit'):
        options[key] = _positive_integer(options[key], key)
        if options[key] > 10000:
            raise DocumentInputError('analysis_options')
    source.seek(0)
    doc = Document(source)
    # Treść główna z tabelami i polami tekstowymi, nagłówki, stopki i przypisy.
    # Każda część jest analizowana osobno: odległość między słowami liczy się
    # tylko w obrębie tej samej części.
    stories, commit = document_stories(doc, include_fallback=False)
    stories = [story for story in stories if story.paragraphs]
    lengths = [len(p.text) for story in stories for p in story.paragraphs]
    if sum(lengths) > 500000 or any(n > 20000 for n in lengths):
        raise DocumentInputError('analysis_limit')
    from lxml import etree
    namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    if any(etree._Element.xpath(story.root, './/w:ins | .//w:del | .//w:moveFrom | .//w:moveTo',
                                namespaces=namespaces) for story in stories):
        raise DocumentInputError('tracked_changes')
    layouts = [_document_layout(story.paragraphs) for story in stories]
    story_words = []
    for story_layouts in layouts:
        words, locations = [], []
        for p_index, (text, _) in enumerate(story_layouts):
            for match in re.finditer(r'\w+', text):
                value = match[0]
                if len(value) >= min_word_length:
                    words.append(Token(value, True, word_len=len(value),
                                       stem5=first5(value) if len(value) >= 5 else None))
                    locations.append((p_index, match.start(), match.end()))
        story_words.append((words, locations))
    ignored_forms, tracked_forms = parse_word_list(ignored_words), parse_word_list(tracked_words)
    forms = [w.text.lower() for words, _ in story_words for w in words]
    if tracked_forms or options['duplicates']:
        forms.extend(m[0].lower() for story_layouts in layouts for text, _ in story_layouts
                     for m in LEXICAL_RE.finditer(text))
    forms.extend(ignored_forms)
    forms.extend(tracked_forms)
    report_progress("Analiza językowa – rozpoznawanie odmian słów")
    lemmas = _lemma_map(forms)
    ignored = {lemmas.get(w, w) for w in ignored_forms}
    tracked = {lemmas.get(w, w) for w in tracked_forms}
    report_progress("Wyszukiwanie powtórzeń")
    total = sum(len(story.paragraphs) for story in stories)
    done = 0
    for story, story_layouts, (words, locations) in zip(stories, layouts, story_words):
        for token in words:
            token.lemma = lemmas.get(token.text.lower(), token.text.lower())
            token.ignored = token.lemma in ignored
        groups = _find_repeat_groups(
            words, window_size=window_size, include_prefix_matches=include_prefix_matches
        )
        _assign_colors(words, groups, color_palette)
        marks = [[] for _ in story.paragraphs]
        for token, (index, start, end) in zip(words, locations):
            if token.color is not None:
                marks[index].append((start, end, token.color, None, 0))
        _, extra_marks = _extra_checks(
            story.paragraphs, lemmas, ignored, tracked, options, layouts=story_layouts, apply_marks=False
        )
        for target, ranges in zip(marks, extra_marks):
            target.extend((a, b, None, MARK_STYLES[k][1], MARK_STYLES[k][0]) for a, b, k in ranges)
        for layout, ranges in zip(story_layouts, marks):
            if done % 10 == 0:
                report_progress("Kolorowanie – akapity", done, total)
            done += 1
            _apply_layout_marks(layout, ranges)
    commit()
    report_progress("Zapis oznaczonego DOCX")
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
