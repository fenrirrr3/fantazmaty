from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from unittest.mock import patch
from django.test import SimpleTestCase, TestCase, RequestFactory
from django.contrib.auth import get_user_model
from django.http import QueryDict
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from core.services.document_formatting import normalize_docx
from core.services.document_converter import convert_document
from core.models import WorkflowEvent
from core.selectors.texts import text_list_context
from core.table_sorting import prepare_table_sort
from texts.models import Text


def sample():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('Kursywa ').italic = True
    p.add_run('pogrubienie').bold = True
    p.paragraph_format.space_after = Pt(20)
    p.paragraph_format.line_spacing = 2
    center = doc.add_paragraph('Środek', style='Title')
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Prawo').alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_table(rows=1, cols=1).cell(0,0).text = 'Tabela'
    stream=BytesIO();doc.save(stream);stream.seek(0)
    return stream


class FormattingTests(SimpleTestCase):
    def test_typography_preserves_runs_and_inherited_alignment(self):
        with normalize_docx(sample()) as out:
            doc=Document(out)
        first,center,right=doc.paragraphs
        self.assertTrue(first.runs[0].italic)
        self.assertTrue(first.runs[1].bold)
        self.assertEqual(first.runs[0].font.name,'Times New Roman')
        self.assertEqual(first.runs[0].font.size.pt,12)
        self.assertEqual(first.paragraph_format.line_spacing,1.5)
        self.assertEqual(first.paragraph_format.space_after.pt,0)
        self.assertEqual(first.paragraph_format.space_before.pt,0)
        self.assertAlmostEqual(first.paragraph_format.first_line_indent.cm,1.25,places=2)
        self.assertEqual(center.paragraph_format.first_line_indent,0)
        self.assertEqual(center.style.paragraph_format.alignment,WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(right.alignment,WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(doc.tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt,12)

    def test_conversion_epub_and_pdf_preserve_formatting(self):
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)):
            result,ext,_=convert_document(sample(),['pdf','epub'])
            with result,ZipFile(result) as archive:
                self.assertTrue(archive.read('document.pdf').startswith(b'%PDF'))
                with ZipFile(BytesIO(archive.read('document.epub'))) as epub:
                    content=epub.read('EPUB/content.xhtml').decode()
                    css=epub.read('EPUB/style.css').decode()
                    self.assertIn('align-center',content)
                    self.assertIn('<strong>',content)
                    self.assertIn('<em>',content)
                    self.assertIn('1.25cm',css)
                    self.assertIn('Times New Roman',css)

    def test_formatting_happens_before_cleaning(self):
        from core.services.odkurzacz import clean_docx
        observed=[]
        def cleaner(source,rules):
            doc=Document(source)
            observed.append(doc.paragraphs[0].runs[0].font.size.pt)
            source.seek(0)
            return clean_docx(source,rules)
        with TemporaryDirectory() as tmp, self.settings(DOCUMENT_CONVERSION_DIR=Path(tmp)), patch('core.services.document_converter.clean_docx',side_effect=cleaner):
            result,_,_=convert_document(sample(),['epub'],use_cleaner=True)
            result.close()
        self.assertEqual(observed,[12])


class StatusSortingTests(TestCase):
    def test_status_date_excludes_non_status_events_and_sorts_before_pagination(self):
        user=get_user_model().objects.create_superuser('admin','admin@example.com','test')
        one=Text.objects.create(title='A',length=100)
        two=Text.objects.create(title='B',length=100)
        missing=Text.objects.create(title='C',length=100)
        def event(text,previous,next_):
            return WorkflowEvent.objects.create(text=text,title=text.title,actor=user,actor_name='A',previous_status=previous,next_status=next_,channel='')
        first=event(one,'Redakcja','Weryfikacja')
        second=event(two,'Redakcja','Korekta')
        event(one,'Weryfikacja','Weryfikacja')
        for direction,expected in [('-',[two.pk,one.pk,missing.pk]),('',[one.pk,two.pk,missing.pk])]:
            request=RequestFactory().get('/',{'sort':direction+'last_status_change'})
            request.user=user
            rows=text_list_context(user=user,params=request.GET)['texts']
            rows,columns=prepare_table_sort(request,rows)
            self.assertEqual([r['pk'] for r in rows],expected)
            self.assertEqual(columns['Ostatnia zmiana statusu'],'last_status_change')
            values={r['pk']:r['last_status_change'] for r in rows}
            self.assertEqual(values[one.pk],first.created_at)
            self.assertEqual(values[two.pk],second.created_at)
