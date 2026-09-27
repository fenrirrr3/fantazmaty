from io import BytesIO
from django.test import SimpleTestCase
from docx import Document
from core.services.odkurzacz import (
    correct_editorial_text, clean_docx, ALL_EDITORIAL_RULES, SENTENCE_ABBREVIATIONS,
)


class SentenceCaseTests(SimpleTestCase):
    def test_basic_and_optional(self):
        self.assertIn('sentence_case', ALL_EDITORIAL_RULES)
        text = 'Zdanie. drugie? trzecie!żółw.'
        self.assertEqual(correct_editorial_text(text, {'sentence_case'}), 'Zdanie. Drugie? Trzecie!Żółw.')
        self.assertEqual(correct_editorial_text(text, set()), text)
        self.assertEqual(correct_editorial_text('Koniec.żółw! już?', {'sentence_case', 'after_punct'}), 'Koniec. Żółw! Już?')

    def test_every_abbreviation(self):
        for value in SENTENCE_ABBREVIATIONS:
            for variant in (value, value.upper(), value.replace('.', '. ')):
                text = 'Przykład ' + variant + ' dalej'
                with self.subTest(value=variant):
                    self.assertEqual(correct_editorial_text(text, {'sentence_case'}), text)
        self.assertEqual(correct_editorial_text('To m.in. kot. dalej', {'sentence_case'}), 'To m.in. kot. Dalej')

    def test_dialogues_ellipsis_and_numbers(self):
        for text in ('– Co? – zapytała.', '– Idź! – krzyknął.', 'Koniec. – powiedział.',
                     'Co?\u00a0– zapytała.', 'Co?! – zapytała.', 'Chciał... ale nie mógł.',
                     'Chciał… ale nie mógł.', 'W 3. rozdziale', 'A. nowak',
                     '27.09.2026 r. rano', 'Wersja 1.2.3 działa.', 'O 12.30 przyjdę.'):
            with self.subTest(text=text):
                self.assertEqual(correct_editorial_text(text, {'sentence_case'}), text)

    def test_files_urls_and_new_sentence(self):
        text = 'Plik tekst.docx i adres x@y.pl oraz https://example.org/a.b. dalej'
        expected = text[:-5] + 'Dalej'
        self.assertEqual(correct_editorial_text(text, {'sentence_case'}), expected)
        self.assertEqual(correct_editorial_text('Plik tekst.docx. dalej', {'sentence_case', 'after_punct'}), 'Plik tekst.docx. Dalej')

    def test_docx_runs_preserved(self):
        doc = Document(); p = doc.add_paragraph(); p.add_run('Koniec. ')
        p.add_run('żółw').italic = True
        p.add_run('? – zapytał. '); p.add_run('dalej').bold = True
        source = BytesIO(); doc.save(source); source.seek(0)
        with clean_docx(source, {'sentence_case'}) as result:
            changed = Document(result).paragraphs[0]
        self.assertEqual(changed.text, 'Koniec. Żółw? – zapytał. Dalej')
        self.assertTrue(changed.runs[1].italic)
        self.assertTrue(changed.runs[-1].bold)
