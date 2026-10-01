import html
import re
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.test import TestCase, SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from people.models import Person, Role
from texts.models import Anthology, Review
from core.services.reviews import copy_review_to_text
from core.services.document_preparation import prepare_docx
from core.services.mailbox_import import package_messages
from texts.admin import review_text_initial


@override_settings(DROPBOX_CHOOSER_APP_KEY='public-test-key')
class ReviewFolderTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.admin = users.objects.create_superuser('admin-folder', 'admin-folder@example.com', 'test')
        self.member = users.objects.create_user('member-folder', 'member-folder@example.com', 'test')
        person = Person.objects.create(user=self.member, first_name='Jan', last_name='Test', email=self.member.email)
        role, _ = Role.objects.get_or_create(name='Recenzent')
        person.roles.add(role)
        self.book = Anthology.objects.create(title='Nabór folderów')
        self.review = Review.objects.create(title='Zgłoszenie', anthology=self.book,
            author_first_name='Jan', author_last_name='Autor', email='author@example.com', length=1234,
            status='accepted', author_notified_at=timezone.localdate(), old_reviews=True,
            file_url='https://www.dropbox.com/sh/original/link')
        self.detail = reverse('core:assigned_review_detail', args=[self.review.pk])
        self.save = reverse('core:update_review_file', args=[self.review.pk])

    def token(self):
        response = self.client.get(self.detail)
        return html.unescape(re.search(r'name="_edit_version" value="([^"]+)"', response.content.decode())[1])

    def test_superuser_picker_and_save_clear_invalid_url(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.detail)
        self.assertContains(response, 'data-app-key="public-test-key"')
        self.assertContains(response, 'review-folder-layout')
        from lxml import html as parser
        page = parser.fromstring(response.content)
        self.assertEqual(len(page.xpath('//*[@class="review-folder-layout"]/section')), 2)
        self.assertEqual(len(page.xpath('//*[@id="dropbox-folder-choose"]/parent::div/button')), 2)
        response = self.client.post(self.save, {'file_url': 'https://www.dropbox.com/sh/new/link', '_edit_version': self.token()})
        self.assertEqual(response.status_code, 302)
        self.review.refresh_from_db(); self.assertIn('/new/', self.review.file_url)
        response = self.client.post(self.save, {'file_url': 'javascript:alert(1)', '_edit_version': self.token()})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'Wybierz folder Dropbox', status_code=400)
        self.review.refresh_from_db(); self.assertIn('/new/', self.review.file_url)
        response = self.client.post(self.save, {'file_url': '', '_edit_version': self.token()})
        self.assertEqual(response.status_code, 302)
        self.review.refresh_from_db(); self.assertEqual(self.review.file_url, '')

    def test_member_neither_sees_folder_nor_can_save(self):
        self.client.force_login(self.member)
        response = self.client.get(self.detail)
        self.assertEqual(response.status_code, 200)
        for value in ('dropins.js', 'public-test-key', '/original/link', 'review-folder-heading', 'author@example.com'):
            self.assertNotContains(response, value)
        self.assertEqual(self.client.post(self.save, {'file_url': ''}).status_code, 403)
        self.review.refresh_from_db(); self.assertIn('/original/', self.review.file_url)

    def test_stale_or_missing_token_does_not_overwrite_folder(self):
        self.client.force_login(self.admin)
        token = self.token()
        self.review.title = 'Zmieniony tytuł'; self.review.save()
        for data in ({'file_url': '', '_edit_version': token}, {'file_url': ''}):
            self.assertEqual(self.client.post(self.save, data).status_code, 409)
        self.review.refresh_from_db(); self.assertIn('/original/', self.review.file_url)

    def test_transfer_and_admin_popup_carry_folder(self):
        self.assertEqual(review_text_initial(self.review)['file_url'], self.review.file_url)
        text = copy_review_to_text(user=self.admin, review_id=self.review.pk, contract_received=True)
        self.assertEqual(text.file_url, self.review.file_url)
        self.assertEqual(copy_review_to_text(user=self.admin, review_id=self.review.pk).pk, text.pk)
        self.client.force_login(self.admin)
        page = self.client.get(reverse('core:assigned_text_detail', args=[text.pk]))
        self.assertContains(page, 'class="secondary-button"\n                    >\n                        Otwórz zgłoszenie')
        self.assertContains(page, text.file_url)

    @override_settings(DROPBOX_CHOOSER_APP_KEY='')
    def test_manual_fallback_without_key(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.detail)
        self.assertContains(response, 'data-input-id="id_file_url" disabled')
        self.assertNotContains(response, 'dropins.js')

    def test_dashboard_buttons_and_program_instructions(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse('core:home'))
        from lxml import html as parser
        tree = parser.fromstring(page.content)
        links = tree.xpath('//a[contains(text(), "Obsłuż zgłoszenia")]/parent::div/a')
        self.assertEqual(len(links), 2)
        self.assertTrue(all('button' in link.attrib['class'] for link in links))
        page = self.client.get(reverse('core:programs'))
        self.assertEqual(page.status_code, 200)
        for phrase in ('uruchamiamy', 'zachowujemy', 'ustawiamy', 'pomijamy', 'przenosimy', 'poprawiamy'):
            self.assertNotContains(page, phrase)


def sample():
    document = Document()
    paragraph = document.add_paragraph('Zwykły tekst ')
    paragraph.add_run('kursywa').italic = True
    paragraph.add_run(' gruby').bold = True
    document.add_paragraph('Do prawej').alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph('Pośrodku').alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('Środek ze stylu', style='Title')
    document.add_paragraph('')
    stream = BytesIO(); document.save(stream); stream.seek(0)
    return stream


class JustificationTests(SimpleTestCase):
    def check_document(self, stream):
        doc = Document(stream)
        self.assertEqual(len(doc.paragraphs), 5)
        self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertEqual(doc.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        for paragraph in doc.paragraphs[2:4]:
            self.assertEqual(paragraph.paragraph_format.first_line_indent, 0)
            self.assertTrue(paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER or paragraph.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(doc.paragraphs[0].runs[1].italic)
        self.assertTrue(doc.paragraphs[0].runs[2].bold)
        self.assertEqual(doc.paragraphs[4].text, '')

    def test_shared_preparation_only_justifies_when_requested(self):
        with prepare_docx(sample(), normalize_formatting=True, justify=True) as output:
            self.check_document(output)
        with prepare_docx(sample(), normalize_formatting=True) as output:
            self.assertEqual(Document(output).paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_mail_conversion_justifies_docx_pdf_and_epub(self):
        row = {'uid': 1, 'folder': 'Tekst', 'files': [('tekst.docx', sample().getvalue())]}
        with TemporaryDirectory() as directory, self.settings(DOCUMENT_CONVERSION_DIR=Path(directory)):
            with package_messages([row], clean=True, convert=True) as result, ZipFile(result) as archive:
                self.check_document(BytesIO(archive.read('Tekst/01_tekst.docx')))
                self.assertTrue(archive.read('Tekst/01_tekst.pdf').startswith(b'%PDF'))
                with ZipFile(BytesIO(archive.read('Tekst/01_tekst.epub'))) as epub:
                    content = epub.read('EPUB/content.xhtml').decode()
                    self.assertIn('align-justify', content)
                    self.assertIn('align-center', content)
