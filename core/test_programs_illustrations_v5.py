from datetime import timedelta
from importlib import import_module
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile, ZIP_STORED

from django.apps import apps
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from lxml import html

from authors.models import Author
from core.odkurzacz_forms import OdkurzaczForm, DocumentConversionForm, RepetitionsForm
from core.services.document_converter import convert_document, conversion_filename
from illustrations.models import Illustration, Illustrator
from texts.models import Anthology, Text, Review
from workflow.tests import create_member


def document_bytes():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run('Ala  ma kota.').italic = True
    paragraph.runs[0].font.size = Pt(17)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph('Środek').alignment = WD_ALIGN_PARAGRAPH.CENTER
    result = BytesIO()
    doc.save(result)
    return result.getvalue()


def sized_docx(size):
    stream = BytesIO(document_bytes())
    with ZipFile(stream, 'a', compression=ZIP_STORED) as archive:
        overhead = 30 + 46 + 2 * len('padding.bin')
        archive.writestr('padding.bin', b'0' * (size - len(stream.getvalue()) - overhead))
    assert len(stream.getvalue()) == size
    return SimpleUploadedFile('tekst.docx', stream.getvalue())


class ProgramChangesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.member = create_member('programy', 'Redaktor')

    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        override = self.settings(DOCUMENT_CONVERSION_DIR=Path(temporary.name)/'convert',
                                 ACTIVITY_SPOOL_DIR=Path(temporary.name)/'activity')
        override.enable(); self.addCleanup(override.disable)
        self.client.force_login(self.member)
        self.url = reverse('core:programs')

    def test_all_program_forms_accept_two_mb_and_reject_one_extra_byte(self):
        for form_class in (OdkurzaczForm, DocumentConversionForm, RepetitionsForm):
            data = {'formats': ['pdf'], 'window_size': 35, 'min_word_length': 4,
                    'sentence_limit': 35, 'paragraph_limit': 150}
            with self.subTest(form=form_class.__name__):
                form = form_class(data, {'document': sized_docx(2*1024*1024)})
                self.assertTrue(form.is_valid(), form.errors)
                form = form_class(data, {'document': sized_docx(2*1024*1024+1)})
                self.assertFalse(form.is_valid())
                self.assertIn('2 MB', str(form.errors['document']))

    def test_oversized_http_upload_never_starts_conversion(self):
        for action, field, extra in (
            ('clean', 'document', {}),
            ('convert', 'convert-document', {'convert-formats': ['pdf']}),
            ('repetitions', 'repetitions-document', {'repetitions-window_size':35,
             'repetitions-min_word_length':4,'repetitions-sentence_limit':35,'repetitions-paragraph_limit':150}),
        ):
            with self.subTest(action=action), patch('core.views.programs.convert_document') as converter:
                response = self.client.post(self.url, {'program_action':action,
                    field:sized_docx(2*1024*1024+1), **extra})
                self.assertContains(response, 'Maksymalny rozmiar to 2 MB')
                converter.assert_not_called()

    def test_mailbox_validator_can_keep_its_existing_ten_mb_limit(self):
        form = OdkurzaczForm({}, {'document':sized_docx(2*1024*1024+1)}, max_document_bytes=10*1024*1024)
        self.assertTrue(form.is_valid(), form.errors)

    def test_conversion_http_preserves_name_for_single_and_multiple_formats(self):
        for formats, ext in ((['pdf'],'pdf'), (['pdf','epub'],'zip')):
            with self.subTest(formats=formats):
                name = 'Żółw i smok.v2.docx'
                response = self.client.post(self.url, {'program_action':'convert',
                    'convert-formats':formats, 'convert-document':SimpleUploadedFile(name,document_bytes())})
                self.assertEqual(response.status_code,200)
                self.assertIn('Żółw i smok.v2.'+ext, unquote(response['Content-Disposition']))
                payload = b''.join(response.streaming_content)
                if ext == 'zip':
                    with ZipFile(BytesIO(payload)) as archive:
                        self.assertEqual(set(archive.namelist()), {'Żółw i smok.v2.pdf','Żółw i smok.v2.epub'})
                else:
                    self.assertTrue(payload.startswith(b'%PDF'))

    def test_warning_archive_also_keeps_the_original_name(self):
        import json
        def worker(directory, timeout):
            (directory/'document.pdf').write_bytes(b'%PDF-test')
            (directory/'warnings.json').write_text(json.dumps(['Uwaga testowa']),encoding='utf8')
        with patch('core.services.document_converter.run_converter', side_effect=worker):
            output,ext,_ = convert_document(SimpleUploadedFile('Zażółć.docx',document_bytes()),
                                            ['pdf'],preserve_filename=True)
            self.assertEqual(ext,'zip')
            with output, ZipFile(output) as archive:
                self.assertEqual(set(archive.namelist()),{'Zażółć.pdf','Uwagi_konwersji.txt'})

    def test_names_cannot_escape_zip_or_inject_response_headers(self):
        self.assertEqual(conversion_filename(r'C:\folder\Tytuł.v1.docx','pdf'),'Tytuł.v1.pdf')
        self.assertEqual(conversion_filename('../../tekst\r\n.docx','epub'),'tekst.epub')

    def test_cleaner_boolean_applies_mailbox_typography_with_and_without_rebuild(self):
        for rebuild in (False,True):
            with self.subTest(rebuild=rebuild):
                data = {'document':SimpleUploadedFile('tekst.docx',document_bytes()),
                        'normalize_formatting':'on','rules':['spaces']}
                if rebuild: data['rebuild']='on'
                response = self.client.post(self.url,data)
                self.assertEqual(response.status_code,200)
                doc = Document(BytesIO(b''.join(response.streaming_content)))
                first, center = doc.paragraphs
                self.assertEqual(first.text,'Ala ma kota.')
                self.assertEqual(first.runs[0].font.name,'Times New Roman')
                self.assertEqual(first.runs[0].font.size.pt,12)
                self.assertTrue(first.runs[0].italic)
                self.assertEqual(first.alignment,WD_ALIGN_PARAGRAPH.JUSTIFY)
                self.assertEqual(first.paragraph_format.line_spacing,1.5)
                self.assertEqual(center.alignment,WD_ALIGN_PARAGRAPH.CENTER)
                self.assertEqual(center.paragraph_format.first_line_indent,0)

    def test_cleaner_without_boolean_preserves_formatting_and_other_forms_omit_it(self):
        response = self.client.post(self.url, {'document':SimpleUploadedFile('tekst.docx',document_bytes()),'rules':['spaces']})
        doc = Document(BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt,17)
        self.assertEqual(doc.paragraphs[0].alignment,WD_ALIGN_PARAGRAPH.LEFT)
        self.assertNotIn('normalize_formatting',DocumentConversionForm().fields)
        self.assertNotIn('normalize_formatting',RepetitionsForm().fields)


class IllustrationWorkspaceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser('admin-v5','admin-v5@example.test','test-only')
        cls.artist = create_member('ilustrator-v5','Ilustrator')
        cls.other = create_member('inny-ilustrator-v5','Ilustrator')
        cls.editor = create_member('redaktor-v5','Redaktor')
        cls.artist_contact=Illustrator.objects.create(first_name='Artysta',last_name='Pierwszy')
        cls.other_contact=Illustrator.objects.create(first_name='Artysta',last_name='Drugi')
        cls.book = Anthology.objects.create(title='Ilustrowana',has_illustrations=True)
        cls.text = Text.objects.create(title='Tytuł ilustracji',anthology=cls.book,length=100,genre='Science fantasy',content_warnings='Ostrzeżenia tekstu')
        cls.text.authors.add(Author.objects.create(first_name='Jan',last_name='Autor'))
        cls.illustration = Illustration.objects.get(text=cls.text)
        cls.review = Review.objects.create(title=cls.text.title,anthology=cls.book,length=100,
            email='autor@example.test',author_first_name='Jan',author_last_name='Autor',
            genre='Science fantasy',copied_text=cls.text)

    def setUp(self):
        temporary = TemporaryDirectory(); self.addCleanup(temporary.cleanup)
        override = self.settings(ACTIVITY_SPOOL_DIR=temporary.name)
        override.enable(); self.addCleanup(override.disable)
        self.client.force_login(self.admin)
        self.url = reverse('illustrations:illustration_detail',args=[self.illustration.pk])

    def token(self, client=None, action='assignment'):
        response = (client or self.client).get(self.url)
        self.assertEqual(response.status_code,200)
        return html.fromstring(response.content).xpath('//form[input[@name="action"][@value=$action]]/input[@name="version"]/@value',action=action)[0]

    def post(self, action, **data):
        return self.client.post(self.url, {'action':action,'version':self.token(action=action),**data})

    def assigned(self):
        Illustration.objects.filter(pk=self.illustration.pk).update(
            illustrator=self.artist_contact,status='assigned',assigned_at=timezone.localdate()-timedelta(days=20))

    def test_list_loads_text_metadata_and_legacy_warnings_as_fallback(self):
        self.illustration.trigger_warnings='Stare ostrzeżenia';self.illustration.save()
        response=self.client.get(reverse('illustrations:illustration_list'))
        for value in ('Ilustrowana','Jan Autor','Science fantasy','Ostrzeżenia tekstu'):
            self.assertContains(response,value)
        self.assertNotContains(response,'Stare ostrzeżenia')
        Text.objects.filter(pk=self.text.pk).update(content_warnings='')
        self.assertContains(self.client.get(self.url),'Stare ostrzeżenia')

    def test_only_in_preparation_and_illustrated_anthologies_are_visible(self):
        for field,value in (('status','ready'),('has_illustrations',False)):
            with self.subTest(field=field):
                Anthology.objects.filter(pk=self.book.pk).update(**{field:value})
                self.assertNotContains(self.client.get(reverse('illustrations:illustration_list')),self.text.title)
                self.assertEqual(self.client.get(self.url).status_code,404)
                Anthology.objects.filter(pk=self.book.pk).update(status='in_preparation',has_illustrations=True)

    def test_saves_assignment_status_and_date(self):
        response=self.post('assignment',illustrator=self.artist_contact.pk,status='assigned')
        self.assertEqual(response.status_code,302)
        self.illustration.refresh_from_db()
        self.assertEqual(self.illustration.illustrator_id,self.artist_contact.pk)
        self.assertEqual(self.illustration.assigned_at,timezone.localdate())
        self.assertEqual(self.post('assignment',illustrator=self.artist_contact.pk,status='delivered').status_code,302)

    def test_corrections_do_not_reset_date_but_new_performer_does(self):
        self.assigned(); original=timezone.localdate()-timedelta(days=20)
        for status in ('delivered','in_corrections','assigned'):
            self.assertEqual(self.post('assignment',illustrator=self.artist_contact.pk,status=status).status_code,302)
            self.illustration.refresh_from_db();self.assertEqual(self.illustration.assigned_at,original)
        self.post('assignment',illustrator=self.other_contact.pk,status='assigned')
        self.illustration.refresh_from_db();self.assertEqual(self.illustration.assigned_at,timezone.localdate())

    def test_unassignment_clears_date_and_missing_illustrator_is_rejected(self):
        self.assertEqual(self.post('assignment',status='assigned',illustrator='').status_code,200)
        self.illustration.refresh_from_db();self.assertEqual(self.illustration.status,'unassigned')
        self.assigned()
        self.assertEqual(self.post('assignment',status='unassigned',illustrator='').status_code,302)
        self.illustration.refresh_from_db();self.assertIsNone(self.illustration.assigned_at)

    def test_link_and_long_excerpt_save_separately_without_changing_assignment(self):
        self.assigned()
        self.assertEqual(self.post('link',story_url='https://example.org/opowiadanie?dl=1').status_code,302)
        fragment=('Długi fragment z polskimi znakami.\n'*500)
        self.assertEqual(self.post('excerpt',illustrated_excerpt=fragment).status_code,302)
        self.illustration.refresh_from_db()
        self.assertEqual(self.illustration.illustrated_excerpt,fragment.strip())
        self.assertEqual(self.illustration.story_url,'https://example.org/opowiadanie?dl=1')
        self.assertEqual(self.illustration.status,'assigned')
        self.assertEqual(self.illustration.assigned_at,timezone.localdate()-timedelta(days=20))
        self.assertNotContains(self.client.get(self.url),'Wybierz folder Dropbox')

    def test_invalid_link_and_html_excerpt_are_safe(self):
        response=self.post('link',story_url='javascript:alert(1)')
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response,'href="javascript:')
        self.illustration.refresh_from_db();self.assertEqual(self.illustration.story_url,'')
        self.post('excerpt',illustrated_excerpt='<script>alert(1)</script>')
        self.assertContains(self.client.get(self.url),'&lt;script&gt;alert(1)&lt;/script&gt;')

    def test_old_or_missing_token_does_not_overwrite_newer_data(self):
        old=self.token(action='link')
        self.illustration.story_url='https://example.org/new';self.illustration.save()
        response=self.client.post(self.url,{'action':'link','version':old,'story_url':'https://example.org/old'})
        self.assertEqual(response.status_code,409)
        self.assertContains(response,'https://example.org/old',status_code=409)
        self.assertContains(response,'href="https://example.org/new"',status_code=409)
        self.illustration.refresh_from_db();self.assertEqual(self.illustration.story_url,'https://example.org/new')
        self.assertEqual(self.client.post(self.url,{'action':'excerpt','illustrated_excerpt':'brak tokena'}).status_code,409)

    def test_artist_account_does_not_grant_contact_edit_rights(self):
        self.assigned()
        for user in (self.artist,self.other):
            self.client.force_login(user)
            self.assertNotContains(self.client.get(self.url),'Zapisz fragment')
            self.assertEqual(self.client.post(self.url,{'action':'excerpt','illustrated_excerpt':'Obcy fragment'}).status_code,403)
            self.assertEqual(self.client.post(self.url,{'action':'assignment','illustrator':self.other_contact.pk,'status':'delivered'}).status_code,403)
        self.illustration.refresh_from_db()
        self.assertEqual(self.illustration.illustrator_id,self.artist_contact.pk)
        self.assertEqual(self.illustration.status,'assigned')


    def test_editor_without_illustration_access_and_csrf_are_blocked(self):
        self.client.force_login(self.editor)
        self.assertEqual(self.client.get(self.url).status_code,403)
        self.assertEqual(self.client.get(reverse('illustrations:illustration_list')).status_code,403)
        csrf_client=Client(enforce_csrf_checks=True);csrf_client.force_login(self.admin)
        self.assertEqual(csrf_client.post(self.url,{'action':'excerpt','illustrated_excerpt':'Nie'}).status_code,403)

    def test_migration_creates_only_missing_required_rows_and_is_repeatable(self):
        self.assigned()
        other=Text.objects.create(title='Brakująca ilustracja',anthology=self.book,length=100)
        Illustration.objects.filter(text=other).delete()
        skipped_book=Anthology.objects.create(title='Bez ilustracji',has_illustrations=False)
        skipped=Text.objects.create(title='Pomijany',anthology=skipped_book,length=100)
        before=Illustration.objects.get(pk=self.illustration.pk)
        migrate=import_module('illustrations.migrations.0002_prepare_illustration_workspace').create_missing
        for _ in range(2): migrate(apps,SimpleNamespace(connection=connection))
        self.assertEqual(Illustration.objects.filter(text=other).count(),1)
        self.assertFalse(Illustration.objects.filter(text=skipped).exists())
        after=Illustration.objects.get(pk=self.illustration.pk)
        self.assertEqual((before.status,before.illustrator_id,before.assigned_at),
                         (after.status,after.illustrator_id,after.assigned_at))
